"""Verificación de fórmulas del papel de PCE por cohortes (AUD-ECL-01) con LibreOffice.

No hay Excel real de Windows en este entorno, así que se usa el fallback documentado en
el contrato: se arma un libro con las cédulas de ``hojas()`` (cada celda calculada como
fórmula), se recalcula con ``soffice --headless`` y se compara el valor recalculado con
el que devolvió Python. Debe terminar en «DIFERENCIAS: 0».

Es una prueba previa útil; no reemplaza la de Excel real (queda como pendiente declarado
hasta contar con Windows + Excel).
"""
from __future__ import annotations

import datetime as dt
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from backend.app.aud.niif.procesadores import pce_cohortes_niif9 as m  # noqa: E402

FILA0 = m.FILA0
_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _celda_valor(fmt, c):
    """Valor a escribir en la celda (fórmula, número, fecha o texto), como haría libro.xlsx."""
    if isinstance(c, dict):
        return "=" + c["f"]
    if fmt in ("d", "x") and isinstance(c, str) and _ISO.match(c):
        y, mo, da = map(int, c.split("-"))
        return dt.date(y, mo, da)
    return c


def _armar(res, ruta):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for h in m.hojas(res):
        ws = wb.create_sheet(h["name"][:31])
        fmts = [c[1] for c in h["cols"]]
        for j, (titulo, _f) in enumerate(h["cols"]):
            ws.cell(row=FILA0 - 1, column=j + 1, value=titulo)
        filas = list(h["rows"]) + ([h["total"]] if h.get("total") else [])
        for i, fila in enumerate(filas):
            for j, c in enumerate(fila):
                v = _celda_valor(fmts[j] if j < len(fmts) else "x", c)
                if v is not None and v != "":
                    ws.cell(row=FILA0 + i, column=j + 1, value=v)
    wb.save(ruta)


def _recalcular(ruta: Path) -> Path:
    salida = ruta.with_name(ruta.stem + "_calc.xlsx")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to",
                        "xlsx:Calc MS Excel 2007 XML", "--outdir", tmp, str(ruta)],
                       check=True, capture_output=True, env={"HOME": tmp, "PATH": "/usr/bin:/bin"})
        conv = Path(tmp) / (ruta.stem + ".xlsx")
        salida.write_bytes(conv.read_bytes())
    return salida


def main() -> int:
    res = m.ejecutar(m.EJEMPLO["datasets"], m.EJEMPLO["parametros"], m.EJEMPLO["corte"])
    hojas = m.hojas(res)
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "pce_cohortes.xlsx"
        _armar(res, ruta)
        calc = _recalcular(ruta)
        wb = openpyxl.load_workbook(calc, data_only=True)
        difs, n = 0, 0
        for h in hojas:
            ws = wb[h["name"][:31]]
            filas = list(h["rows"]) + ([h["total"]] if h.get("total") else [])
            for i, fila in enumerate(filas):
                for j, c in enumerate(fila):
                    if not isinstance(c, dict):
                        continue
                    n += 1
                    xl = ws.cell(row=FILA0 + i, column=j + 1).value
                    py = c["v"]
                    if isinstance(xl, dt.datetime):
                        xl = xl.date().isoformat()
                    ok = ((py in (None, "") and xl in (None, "")) or
                          (isinstance(py, str) and str(py) == str(xl)) or
                          (isinstance(py, (int, float)) and isinstance(xl, (int, float))
                           and abs(py - xl) <= (1e-6 if abs(py) <= 2 else 0.005)))
                    if not ok:
                        difs += 1
                        print(f"  ✗ {h['name']} f{FILA0 + i}c{j + 1}: py={py!r} calc={xl!r} | {c['f'][:70]}")
        print(f"\nFórmulas comprobadas: {n} · DIFERENCIAS: {difs}")
        return 1 if difs else 0


if __name__ == "__main__":
    sys.exit(main())
