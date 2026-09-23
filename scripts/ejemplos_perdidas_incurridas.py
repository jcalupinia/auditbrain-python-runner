"""Escribe los archivos de EJEMPLO (formato válido con datos ficticios) del
requerimiento de la herramienta «Deterioro de cuentas por cobrar · pérdidas
incurridas (PYMES)» (procesador ``perdidas_incurridas_s11``).

Los DATOS salen de ``backend/app/aud/niif/ejemplos_pi.py`` (única fuente de
verdad, compartida con el endpoint del «Ejercicio modelo»). Este script solo
arma los libros/documentos.

Salida:
    frontend/public/ejemplos/perdidas_incurridas_s11/
        RQ-001_cartera_2025.xlsx                 (anexo a3, cartera del corte)
        RQ-002_cartera_2024.xlsx                 (anexo a2, ejercicio anterior)
        RQ-003_cartera_2023.xlsx                 (anexo a1, dos ejercicios antes)
        RQ-004_provision_inicial_2025.xlsx       (provisión y diferido por factura)
        RQ-005_movimiento_provision_3_ejercicios.xlsx
        RQ-006_cobros_posteriores_al_cierre.xlsx
        RQ-007_politica_credito_cobranza.docx
        RQ-008_ventas_por_factura_3_ejercicios.xlsx

Cada libro trae los datos en la hoja «Datos» (los rótulos son los de ``CAMPOS``
del procesador, para que el mapeo del ciclo sea automático) y una hoja aparte
«Léame» con la nota «EJEMPLO · datos ficticios» (en su propia hoja para NO
estorbar la lectura: el lector del ciclo toma la hoja «Datos»).

Determinista: reescribe cada zip con fechas fijas, así los archivos comiteados
no cambian entre corridas.

Uso:
    python scripts/ejemplos_perdidas_incurridas.py            # escribe los 8 archivos
    python scripts/ejemplos_perdidas_incurridas.py --verificar  # corre el procesador y reporta métricas, sin escribir
"""
from __future__ import annotations

import os
import re
import sys
import zipfile
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.app.aud.niif import ejemplos_pi as ej  # noqa: E402
from backend.app.aud.niif.procesadores import perdidas_incurridas_s11 as pi  # noqa: E402

SALIDA = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "frontend", "public", "ejemplos", "perdidas_incurridas_s11",
)
NOTA = "EJEMPLO · datos ficticios"
FECHA_FIJA = date(2026, 1, 15)
_FIX_DT = (2026, 1, 15, 0, 0, 0)
_FIX = b"2026-01-15T00:00:00Z"

COLUMNAS = {
    "cartera": ["N° de factura", "Cliente", "RUC / identificación", "Fecha de emisión",
                "Fecha de vencimiento", "Importe original", "Saldo por cobrar"],
    "provision": ["N° de factura", "Cliente", "Provisión / deterioro", "Impuesto diferido"],
    "movimiento": ["Año", "Provisión inicial", "Gasto del año", "Castigos", "Recuperaciones"],
    "cobros": ["Fecha de cobro", "N° de factura", "Cliente", "Valor cobrado", "Referencia bancaria"],
    "ventas": ["N° de factura", "Cliente", "Fecha de emisión", "Importe facturado"],
}


def _celda(v):
    return v.strftime("%d/%m/%Y") if isinstance(v, date) else v


def _normalizar(path: str):
    """Reescribe el .xlsx/.docx (que es un zip) con fechas fijas en las entradas y
    en ``docProps/core.xml`` (openpyxl/python-docx pisan «modified» con la hora
    actual). Así el archivo comiteado no cambia entre corridas."""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, i.external_attr, z.read(i.filename)) for i in z.infolist()]
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, attr, data in items:
            if nombre == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                              rb"\g<1>" + _FIX + rb"\g<2>", data)
            zi = zipfile.ZipInfo(nombre, date_time=_FIX_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = attr
            z.writestr(zi, data)
    os.replace(tmp, path)


def _libro_datos(columnas, filas, titulo, explicacion):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(columnas)
    for fila in filas:
        ws.append([_celda(fila.get(c, "")) for c in columnas])
    lee = wb.create_sheet("Léame")
    lee["A1"] = NOTA
    lee["A3"] = titulo
    lee["A5"] = explicacion
    lee["A7"] = ("Este archivo es un formato de ejemplo con datos ficticios para mostrar la estructura "
                 "esperada. La hoja «Datos» tiene los rótulos que la prueba reconoce automáticamente.")
    wb.active = 0
    creado = datetime(FECHA_FIJA.year, FECHA_FIJA.month, FECHA_FIJA.day)
    wb.properties.creator = "AuditConsulting Auditores Cía. Ltda. · Ejercicio modelo"
    wb.properties.created = creado
    wb.properties.modified = creado
    return wb


def _guardar(wb, nombre):
    ruta = os.path.join(SALIDA, nombre)
    wb.save(ruta)
    _normalizar(ruta)


def _docx_politica(nombre):
    from docx import Document

    doc = Document()
    doc.add_heading("Política de crédito y cobranza (EJEMPLO · datos ficticios)", level=0)
    doc.add_paragraph(NOTA + " — documento de muestra para el requerimiento RQ-007.")
    doc.add_heading("1. Otorgamiento de crédito", level=1)
    doc.add_paragraph("El crédito se aprueba por el Comité de Crédito según la calificación del cliente. "
                      "Plazos estándar: 30, 45, 60 o 90 días. El límite por cliente se revisa cada semestre.")
    doc.add_heading("2. Gestión de cobranza", level=1)
    doc.add_paragraph("Recordatorio a los 5 días de vencido; gestión telefónica a los 30; carta de "
                      "requerimiento a los 60; y traslado a cobranza judicial sobre los 180 días.")
    doc.add_heading("3. Deterioro y castigo", level=1)
    doc.add_paragraph("La provisión por deterioro se estima con el modelo de pérdida incurrida de la Sección 11 "
                      "de la NIIF para las PYMES, por tramos de antigüedad. Se propone castigo cuando la mora "
                      "supera 730 días y se agotó la gestión de cobro documentada.")
    doc.core_properties.author = "AuditConsulting Auditores Cía. Ltda. · Ejercicio modelo"
    ruta = os.path.join(SALIDA, nombre)
    doc.save(ruta)
    _normalizar(ruta)


def escribir():
    os.makedirs(SALIDA, exist_ok=True)
    facturas = ej.generar_facturas()
    prov = ej.provision_inicial(facturas)
    _guardar(_libro_datos(COLUMNAS["cartera"], ej.cartera_de(facturas, 2025),
                          "Cartera por factura al 31/12/2025 (anexo del corte)",
                          "Una fila por factura pendiente al cierre; sin filas de total."), "RQ-001_cartera_2025.xlsx")
    _guardar(_libro_datos(COLUMNAS["cartera"], ej.cartera_de(facturas, 2024),
                          "Cartera por factura al 31/12/2024 (ejercicio anterior)",
                          "Sirve para medir cuánto de cada tramo no se recuperó al año siguiente."), "RQ-002_cartera_2024.xlsx")
    _guardar(_libro_datos(COLUMNAS["cartera"], ej.cartera_de(facturas, 2023),
                          "Cartera por factura al 31/12/2023 (dos ejercicios antes)",
                          "Segunda ventana de evidencia histórica de recuperación."), "RQ-003_cartera_2023.xlsx")
    _guardar(_libro_datos(COLUMNAS["provision"], prov,
                          "Provisión y diferido por factura al inicio de 2025",
                          "Solo facturas que también están en la cartera (cruzan): reversión o baja según la evidencia."),
             "RQ-004_provision_inicial_2025.xlsx")
    _guardar(_libro_datos(COLUMNAS["movimiento"], ej.movimiento(prov),
                          "Movimiento de la provisión (3 ejercicios) según el mayor",
                          "El saldo final de un año es el inicial del siguiente; la inicial 2025 iguala el total del RQ-004."),
             "RQ-005_movimiento_provision_3_ejercicios.xlsx")
    _guardar(_libro_datos(COLUMNAS["cobros"], ej.cobros_posteriores(facturas),
                          "Cobros posteriores al cierre (I trimestre 2026)",
                          "Evidencia de recuperabilidad posterior de la cartera del corte."),
             "RQ-006_cobros_posteriores_al_cierre.xlsx")
    _docx_politica("RQ-007_politica_credito_cobranza.docx")
    _guardar(_libro_datos(COLUMNAS["ventas"], ej.ventas_por_factura(facturas),
                          "Ventas por factura de los 3 ejercicios",
                          "Contexto de rotación de la cartera; no interviene en el cálculo."),
             "RQ-008_ventas_por_factura_3_ejercicios.xlsx")
    print(f"Escritos 8 archivos de ejemplo en: {SALIDA}")


def verificar() -> bool:
    m = ej.ejercicio_modelo()
    res = pi.ejecutar(m["datasets"], m["parametros"], m["corte"])
    a3 = pi._cartera(m["datasets"]["a3"])
    total = sum(f["saldo"] for f in a3)
    perdida = float(res["totals"]["perdida"])
    pct = perdida / total * 100 if total else 0
    problemas = [p["code"] for p in res["exceptions"]]
    print("=== Verificación del ejercicio modelo (pérdidas incurridas) ===")
    print(f"poblaciones: {{k: len}} = {{'a1':{len(m['datasets']['a1'])},'a2':{len(m['datasets']['a2'])},"
          f"'a3':{len(m['datasets']['a3'])},'provision':{len(m['datasets']['provision'])},'movimiento':{len(m['datasets']['movimiento'])}}}")
    print(f"cartera total a3: {total:,.2f}")
    print(f"pérdida recalculada: {perdida:,.2f}  ({pct:.1f} % de la cartera)")
    print(f"problemas: {problemas}")
    ok = 5 <= pct <= 15 and "PROVISION_PENDIENTE" not in problemas and "CONCILIACION_INICIAL" not in problemas
    print("RESULTADO:", "OK" if ok else "REVISAR CALIBRACIÓN")
    return ok


def main():
    if "--verificar" in sys.argv:
        sys.exit(0 if verificar() else 1)
    escribir()


if __name__ == "__main__":
    main()
