"""Verificación con LibreOffice (sin Excel) de los datos del cliente dentro del libro.

Uso:  python scripts/verificar_datos_cliente.py [<id_procesador> ...]
Sin argumentos revisa las 19 herramientas que usan ``procesadores/datos_cliente.py``.
Para cada una arma el papel del ejercicio modelo (con sus hojas «Datos del cliente»), lo
recalcula con LibreOffice sin interfaz y compara:

- cada celda enlazada a una hoja de datos (``'D1_…'!F7``) contra el valor que traía la cédula;
- cada fórmula del libro contra el valor que calculó Python;
- que ninguna celda quede en error (#¡VALOR!, #N/D, #¡REF!…).

Imprime cuántas celdas se enlazaron por herramienta y termina con «DIFERENCIAS: 0».
"""
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook  # noqa: E402

from backend.app.aud.niif import ejercicio_modelo as em  # noqa: E402
from backend.app.aud.niif.procesadores import PROCESADORES, datos_cliente, libro  # noqa: E402

FILA0 = 5
ERRORES = ("#VALUE!", "#N/A", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "Err:")


def igual(py, xl) -> bool:
    if py in (None, "") and xl in (None, ""):
        return True
    if isinstance(py, str) and len(py) == 10 and py[4] == "-" and hasattr(xl, "year"):
        return f"{xl.year:04d}-{xl.month:02d}-{xl.day:02d}" == py
    if isinstance(py, bool):
        return bool(xl) == py
    if isinstance(py, (int, float)) and isinstance(xl, (int, float)):
        return abs(py - xl) <= (1e-6 if abs(py) <= 2 else 0.005)
    return str(py) == str(xl)


def verificar(pid: str, tmp: str) -> tuple[int, int, int]:
    mod = PROCESADORES[pid]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    reg = em._reg(d, mod, ds, par, corte)
    hojas = libro.cedulas(d, reg, [], 1, "VERIFICACION")
    titulos = libro._titulos_unicos(hojas)
    ruta = os.path.join(tmp, f"{pid}.xlsx")
    with open(ruta, "wb") as fh:
        fh.write(libro.xlsx(d, reg, [], 1, "VERIFICACION"))
    out = os.path.join(tmp, "recalc")
    os.makedirs(out, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, ruta],
                   check=True, capture_output=True, timeout=300)
    wb = load_workbook(os.path.join(out, f"{pid}.xlsx"), data_only=True)
    fallas = n = enlazadas = 0
    for h, t in zip(hojas, titulos):
        ws = wb[t]
        filas = h["rows"] + ([h["total"]] if h.get("total") else [])
        for i, fila in enumerate(filas):
            for j, c in enumerate(fila):
                if not isinstance(c, dict) or "f" not in c:
                    continue
                n += 1
                if re.fullmatch(r"(?:VALUE\(|IF\()?'D\d+_[^']+'![A-Z]+\d+.*", c["f"]):
                    enlazadas += 1
                xl = ws.cell(row=FILA0 + i, column=j + 1).value
                if not igual(c["v"], xl):
                    fallas += 1
                    print(f"[{pid}] {t}!{ws.cell(row=FILA0 + i, column=j + 1).coordinate}: "
                          f"Python={c['v']!r} LibreOffice={xl!r}  ={c['f'][:110]}")
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith(ERRORES):
                    fallas += 1
                    print(f"[{pid}] {ws.title}!{c.coordinate}: {c.value}")
    datos = [h["name"] for h in hojas if re.match(r"D\d+_", h["name"])]
    print(f"[{pid}] hojas de datos: {', '.join(datos) or '—'} · celdas enlazadas: {enlazadas} · fórmulas comparadas: {n} · diferencias: {fallas}")
    return fallas, n, enlazadas


if __name__ == "__main__":
    ids = sys.argv[1:] or [p for p in PROCESADORES if p not in datos_cliente.PROPIAS]
    tf = tn = te = 0
    with tempfile.TemporaryDirectory() as tmp:
        for pid in ids:
            f, n, e = verificar(pid, tmp)
            tf, tn, te = tf + f, tn + n, te + e
    print(f"HERRAMIENTAS: {len(ids)} · ENLAZADAS: {te} · FÓRMULAS: {tn} · DIFERENCIAS: {tf}")
    sys.exit(1 if tf else 0)
