"""Motor del agente guía «NIIF Piloto» (skill .claude/skills/niif-piloto).

Un solo motor para TODAS las pruebas NIIF del Command Center (las 20 herramientas
del catálogo + la planificación NIA). No hay un agente por prueba: el skill lee de
aquí qué datos pide cada prueba, se los solicita al usuario de a uno, y luego ejecuta
el procesador determinista que YA existe y arma el papel de trabajo en los cinco
formatos (Excel con fórmulas, HTML autónomo, Word, PowerPoint y PDF).

El motor NO calcula nada por su cuenta: reutiliza el contrato oficial de
``backend/app/aud/niif`` (``procesadores.PROCESADORES``, ``mod.ejecutar``,
``mod.definicion``, ``ejercicio_modelo`` y ``libro.{xlsx,html,docx,pptx,pdf}``),
el mismo que usan el router del ciclo y ``scripts/papeles_muestra.py``.

Subcomandos:
  listar                         Todas las pruebas disponibles (id, rubro, marco, datos).
  requisitos <id>                Qué datos necesita una prueba (datasets, parámetros, requerimientos).
  plantilla  <id> <carpeta>      Escribe el ejemplo realista como CSV+JSON para que el usuario lo llene.
  ejemplo    <id> <carpeta>      Ejecuta la prueba con el ejemplo del manifiesto y arma los papeles (demo/verificación).
  ejecutar   <id> <carpeta_salida> (--datos datos.json | --carpeta carpeta_plantilla)
                                 Ejecuta con los datos reales del usuario, arma los papeles y VERIFICA el Excel.

Todo en español (regla del proyecto). Salida en JSON con ``--json`` para que el skill la consuma.

Uso desde la raíz del repo:
  python scripts/piloto/piloto.py listar
  python scripts/piloto/piloto.py requisitos arrendamientos
  python scripts/piloto/piloto.py plantilla arrendamientos /tmp/arr
  python scripts/piloto/piloto.py ejecutar arrendamientos /tmp/salida --carpeta /tmp/arr
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys

# La raíz del repo, para importar backend.* sin instalar el paquete.
_RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)

from backend.app.aud.niif import ejercicio_modelo, procesadores  # noqa: E402
from backend.app.aud.niif.procesadores import libro  # noqa: E402

# (extensión de archivo, nombre de la función en libro, etiqueta legible).
FORMATOS = (
    ("xlsx", "xlsx", "Excel con fórmulas"),
    ("html", "html", "HTML autónomo"),
    ("docx", "docx", "Word"),
    ("pptx", "pptx", "PowerPoint"),
    ("pdf", "pdf", "PDF"),
)


def _mod(pid: str):
    """El módulo del procesador, o aborta con un mensaje claro."""
    if pid not in procesadores.PROCESADORES:
        disponibles = ", ".join(sorted(procesadores.PROCESADORES))
        _salir(f"Prueba desconocida: «{pid}». Disponibles: {disponibles}")
    return procesadores.PROCESADORES[pid]


def _salir(mensaje: str, codigo: int = 2):
    print(json.dumps({"ok": False, "error": mensaje}, ensure_ascii=False, indent=2))
    sys.exit(codigo)


def _campos(mod) -> dict:
    """{dataset: [campos]} declarados por la prueba (o {} si es solo procesador)."""
    return dict(getattr(mod, "CAMPOS", {}) or {})


def _resumen_prueba(pid: str, mod) -> dict:
    d = mod.definicion()
    campos = _campos(mod)
    return {
        "id": pid,
        "rubro": getattr(mod, "RUBRO", "") or "(sin rubro de catálogo)",
        "nombre": d.get("name", pid),
        "area": d.get("area", ""),
        "marcos": d.get("frameworks", []),
        "principal": getattr(mod, "PRINCIPAL", ""),
        "datasets": {ds: len(cs) for ds, cs in campos.items()},
        "parametros": sorted((getattr(mod, "PARAMETROS", {}) or {}).keys()),
        "n_requerimientos": len(d.get("requests", [])),
        "tiene_ejemplo": ejercicio_modelo.disponible(mod),
    }


# --------------------------------------------------------------------------- listar
def cmd_listar(args):
    filas = [_resumen_prueba(pid, procesadores.PROCESADORES[pid]) for pid in procesadores.PROCESADORES]
    if args.json:
        print(json.dumps({"ok": True, "pruebas": filas}, ensure_ascii=False, indent=2))
        return
    print(f"\n{len(filas)} pruebas NIIF disponibles (el orden es la matriz del socio):\n")
    for f in filas:
        marca = "✓ ejemplo" if f["tiene_ejemplo"] else "sin ejemplo"
        print(f"  • {f['id']:<26} [{f['rubro']:<20}] {f['nombre']}  — {marca}")
    print("\nUsá:  requisitos <id>   para ver qué datos pide cada prueba.\n")


# ----------------------------------------------------------------------- requisitos
def cmd_requisitos(args):
    mod = _mod(args.id)
    d = mod.definicion()
    campos = _campos(mod)
    datasets = []
    for ds, cs in campos.items():
        datasets.append({
            "dataset": ds,
            "es_principal": ds == getattr(mod, "PRINCIPAL", None),
            "campos": [
                {"key": c["key"], "label": c["label"], "tipo": c.get("type", "text"),
                 "requerido": c.get("required", True), "ejemplo": c.get("example")}
                for c in cs
            ],
        })
    parametros = getattr(mod, "PARAMETROS", {}) or {}
    requerimientos = [
        {"id": r["id"], "documento": r.get("document", ""), "requerido": r.get("required", True),
         "uso": r.get("use", ""), "formatos": r.get("formats", []), "alimenta_dataset": r.get("dataset")}
        for r in d.get("requests", [])
    ]
    payload = {
        "ok": True,
        "id": args.id,
        "nombre": d.get("name", args.id),
        "marcos": d.get("frameworks", []),
        "resumen": d.get("summary", ""),
        "corte_ejemplo": (ejercicio_modelo.escenario(mod) or (None, None, None))[2],
        "datasets": datasets,
        "parametros_por_defecto": parametros,
        "requerimientos_al_cliente": requerimientos,
        "nota": ("El motor calcula sobre 'datasets' (una fila por registro). Los requerimientos "
                 "marcados 'soporte' sustentan la evidencia pero no se cargan como filas."),
    }
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return
    print(f"\n=== {payload['nombre']} ({args.id}) ===")
    print(f"Marcos: {', '.join(payload['marcos'])}")
    print(f"\n{payload['resumen']}\n")
    for ds in datasets:
        estrella = " (PRINCIPAL: población a recalcular)" if ds["es_principal"] else ""
        print(f"— Anexo «{ds['dataset']}»{estrella}: una fila por registro, columnas:")
        for c in ds["campos"]:
            req = "obligatorio" if c["requerido"] else "opcional"
            ej = f"  ej: {c['ejemplo']}" if c["ejemplo"] is not None else ""
            print(f"     · {c['key']:<20} {c['label']}  [{c['tipo']}, {req}]{ej}")
    if parametros:
        print("\n— Parámetros del auditor (tienen valor por defecto; ajustá solo lo que aplique):")
        for k, v in parametros.items():
            print(f"     · {k} = {v!r}")
    print("\n— Requerimientos formales al cliente (NIA 500):")
    for r in requerimientos:
        req = "obligatorio" if r["requerido"] else "opcional"
        print(f"     · {r['id']} [{req}] {r['documento']}")
    print("\nGenerá el molde con:  plantilla", args.id, "<carpeta>\n")


# ------------------------------------------------------------------------ plantilla
def _escribir_csv(ruta: str, filas: list[dict], columnas: list[str]):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=columnas, extrasaction="ignore")
        w.writeheader()
        for fila in filas:
            w.writerow({c: fila.get(c, "") for c in columnas})


def cmd_plantilla(args):
    mod = _mod(args.id)
    esc = ejercicio_modelo.escenario(mod)
    if esc is None:
        _salir(f"La prueba «{args.id}» no trae ejemplo de manifiesto; cargá los datos a mano según 'requisitos'.")
    datasets, param, corte = esc
    campos = _campos(mod)
    os.makedirs(args.carpeta, exist_ok=True)
    escritos = []
    for ds, filas in datasets.items():
        columnas = [c["key"] for c in campos.get(ds, [])] or (list(filas[0].keys()) if filas else [])
        ruta = os.path.join(args.carpeta, f"{ds}.csv")
        _escribir_csv(ruta, filas, columnas)
        escritos.append({"archivo": os.path.basename(ruta), "dataset": ds, "filas": len(filas), "columnas": columnas})
    with open(os.path.join(args.carpeta, "parametros.json"), "w", encoding="utf-8") as fh:
        json.dump({**(getattr(mod, "PARAMETROS", {}) or {}), **(param or {})}, fh, ensure_ascii=False, indent=2)
    with open(os.path.join(args.carpeta, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump({"id": args.id, "corte": corte}, fh, ensure_ascii=False, indent=2)
    payload = {"ok": True, "id": args.id, "carpeta": args.carpeta, "corte": corte, "csv": escritos,
               "aviso": ("Son los datos del EJEMPLO ficticio. Reemplazalos por los del cliente conservando "
                         "las columnas y el nombre de cada CSV. Luego: ejecutar " + args.id + " <salida> --carpeta " + args.carpeta)}
    print(json.dumps(payload, ensure_ascii=False, indent=2))


# ---------------------------------------------------------------- lectura de datos
def _leer_carpeta(carpeta: str, mod) -> tuple[dict, dict, str]:
    meta_p = os.path.join(carpeta, "meta.json")
    if not os.path.exists(meta_p):
        _salir(f"Falta {meta_p}. Generá la carpeta con 'plantilla' primero.")
    meta = json.load(open(meta_p, encoding="utf-8"))
    corte = meta.get("corte")
    param = {}
    param_p = os.path.join(carpeta, "parametros.json")
    if os.path.exists(param_p):
        param = json.load(open(param_p, encoding="utf-8"))
    datasets = {}
    for ds in _campos(mod):
        ruta = os.path.join(carpeta, f"{ds}.csv")
        if os.path.exists(ruta):
            with open(ruta, encoding="utf-8-sig", newline="") as fh:
                datasets[ds] = [dict(r) for r in csv.DictReader(fh)]
    return datasets, param, corte


def _leer_json(ruta: str) -> tuple[dict, dict, str]:
    obj = json.load(open(ruta, encoding="utf-8"))
    return obj.get("datasets", {}), obj.get("parametros", {}), obj.get("corte")


# ------------------------------------------------------------ verificación del Excel
_PREFIJOS_PELIGROSOS = ("=", "+", "-", "@")


def verificar_xlsx(contenido: bytes) -> dict:
    """Reabre el Excel (regla suprema del CLAUDE.md, punto 6): que NO levante el cuadro
    «reparaciones». Devuelve hojas y celdas de texto que Excel podría leer como fórmula."""
    import openpyxl  # dependencia del backend
    wb = openpyxl.load_workbook(io.BytesIO(contenido))  # si estuviera corrupto, lanza aquí
    sospechosas = []
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                if celda.data_type == "s" and isinstance(celda.value, str) and celda.value[:1] in _PREFIJOS_PELIGROSOS:
                    sospechosas.append(f"{ws.title}!{celda.coordinate}={celda.value[:20]!r}")
    return {"reabre": True, "hojas": wb.sheetnames, "n_hojas": len(wb.sheetnames),
            "celdas_texto_riesgosas": sospechosas[:20]}


# ------------------------------------------------------------------- ejecutar/ejemplo
def _armar_papeles(mod, d, reg, salida: str, base: str) -> dict:
    os.makedirs(salida, exist_ok=True)
    hechos, errores = [], []
    contenido_xlsx = None
    pdf_no_disp = getattr(libro, "PDFNoDisponible", None)
    for ext, fn, etq in FORMATOS:
        try:
            contenido = getattr(libro, fn)(d, reg, [], 1, "PILOTO")
        except Exception as exc:  # noqa: BLE001 — se reporta, no se oculta
            if pdf_no_disp is not None and isinstance(exc, pdf_no_disp):
                errores.append({"formato": ext, "motivo": "PDF no disponible en este entorno (falta WeasyPrint); "
                                                          "el usuario lo obtiene con «Guardar como PDF» del navegador desde el HTML"})
            else:
                errores.append({"formato": ext, "motivo": f"{type(exc).__name__}: {exc}"})
            continue
        ruta = os.path.join(salida, f"{base}.{ext}")
        with open(ruta, "wb") as fh:
            fh.write(contenido)
        if ext == "xlsx":
            contenido_xlsx = contenido
        hechos.append({"formato": ext, "etiqueta": etq, "archivo": os.path.basename(ruta), "bytes": len(contenido)})
    verificacion = verificar_xlsx(contenido_xlsx) if contenido_xlsx else {"reabre": False}
    run = reg["run"]
    prim = run.get("primary")
    return {
        "papeles": hechos,
        "errores": errores,
        "verificacion_excel": verificacion,
        "resultado": {
            "principal": prim,
            "etiqueta_principal": run.get("labels", {}).get(prim, prim),
            "valor_principal": run.get("totals", {}).get(prim),
            "n_problemas": len(run.get("exceptions", [])),
            "filas_por_dataset": {k: len(v) for k, v in (reg["datasets"] or {}).items()},
        },
    }


def cmd_ejemplo(args):
    mod = _mod(args.id)
    esc = ejercicio_modelo.escenario(mod)
    if esc is None:
        _salir(f"La prueba «{args.id}» no trae ejemplo de manifiesto.")
    datasets, param, corte = esc
    d = mod.definicion()
    reg = ejercicio_modelo._reg(d, mod, datasets, param, corte)
    out = _armar_papeles(mod, d, reg, args.carpeta, f"{args.id}__ejemplo")
    print(json.dumps({"ok": not out["errores"] or bool(out["papeles"]), "id": args.id, "corte": corte,
                      "modo": "ejemplo ficticio del manifiesto", **out}, ensure_ascii=False, indent=2))


def cmd_ejecutar(args):
    mod = _mod(args.id)
    if bool(args.datos) == bool(args.carpeta):
        _salir("Indicá exactamente uno: --datos <datos.json>  o  --carpeta <carpeta_plantilla>.")
    datasets, param, corte = _leer_json(args.datos) if args.datos else _leer_carpeta(args.carpeta, mod)
    if not corte:
        _salir("Falta la fecha de corte (meta.json 'corte' o clave 'corte' del JSON).")
    # Parámetros del usuario sobre los valores por defecto de la prueba.
    parametros = {**(getattr(mod, "PARAMETROS", {}) or {}), **(param or {})}
    d = mod.definicion()
    try:
        reg = ejercicio_modelo._reg(d, mod, datasets, parametros, corte)
    except Exception as exc:  # noqa: BLE001
        _salir(f"El procesador no pudo calcular: {type(exc).__name__}: {exc}", codigo=1)
    out = _armar_papeles(mod, d, reg, args.carpeta_salida, f"{args.id}__cliente")
    print(json.dumps({"ok": bool(out["papeles"]) and out["verificacion_excel"].get("reabre"),
                      "id": args.id, "corte": corte, "modo": "datos del cliente", **out},
                     ensure_ascii=False, indent=2))


def main(argv=None):
    p = argparse.ArgumentParser(description="Motor del agente guía NIIF Piloto.")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("listar", help="Todas las pruebas disponibles.")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_listar)

    s = sub.add_parser("requisitos", help="Qué datos necesita una prueba.")
    s.add_argument("id")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_requisitos)

    s = sub.add_parser("plantilla", help="Escribe el ejemplo como CSV+JSON para llenar.")
    s.add_argument("id")
    s.add_argument("carpeta")
    s.set_defaults(func=cmd_plantilla)

    s = sub.add_parser("ejemplo", help="Ejecuta con el ejemplo del manifiesto (demo/verificación).")
    s.add_argument("id")
    s.add_argument("carpeta")
    s.set_defaults(func=cmd_ejemplo)

    s = sub.add_parser("ejecutar", help="Ejecuta con datos reales y arma los papeles.")
    s.add_argument("id")
    s.add_argument("carpeta_salida")
    s.add_argument("--datos", help="Ruta a datos.json {corte, parametros, datasets}.")
    s.add_argument("--carpeta", help="Carpeta con <dataset>.csv + parametros.json + meta.json (de 'plantilla').")
    s.set_defaults(func=cmd_ejecutar)

    args = p.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
