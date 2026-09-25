"""Verificación del papel de las pruebas DECLARATIVAS con LibreOffice (sin Excel).

Uso:  python scripts/verificar_papel_declarativo.py [carga.json ...]
Sin argumentos usa los ejemplos de ``tests/fixtures/papel_declarativo``. Arma el Excel con el
diseño nuevo (``procesadores/declarativo.py``), lo recalcula con LibreOffice sin interfaz y
compara cada celda con fórmula contra el valor que calculó el motor del sitio. Revisa también
que ninguna celda del libro quede en error (#¡VALOR!, #N/D, #¡REF!…) y que las tarjetas de la
portada coincidan con el panel del HTML. Debe terminar con «DIFERENCIAS: 0».
"""
import glob
import json
import os
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook  # noqa: E402

from backend.app.aud.niif.procesadores import declarativo, graficos  # noqa: E402

FILA0 = 5
ERRORES = ("#VALUE!", "#N/A", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "Err:")


def recalcular(ruta: str, destino: str) -> str:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", destino, ruta],
                   check=True, capture_output=True, timeout=180)
    return os.path.join(destino, os.path.basename(ruta))


def igual(py, xl) -> bool:
    if py in (None, "") and xl in (None, ""):
        return True
    if isinstance(py, str) and len(py) == 10 and py[4] == "-" and hasattr(xl, "year"):
        return f"{xl.year:04d}-{xl.month:02d}-{xl.day:02d}" == py
    if isinstance(py, (int, float)) and isinstance(xl, (int, float)):
        return abs(py - xl) <= (1e-6 if abs(py) <= 2 else 0.005)
    return str(py) == str(xl)


def verificar(ruta_carga: str, tmp: str) -> tuple[int, int]:
    carga = json.load(open(ruta_carga, encoding="utf-8"))
    nombre = os.path.splitext(os.path.basename(ruta_carga))[0]
    d, reg, _, _, _ = declarativo.armar(carga)
    xlsx = os.path.join(tmp, f"{nombre}.xlsx")
    with open(xlsx, "wb") as fh:
        fh.write(declarativo.archivo(carga, "xlsx"))
    out = os.path.join(tmp, "recalc")
    os.makedirs(out, exist_ok=True)
    wb = load_workbook(recalcular(xlsx, out), data_only=True)
    fallas = n = 0
    for h in reg["run"]["hojas"]:
        ws = wb[h["name"][:31]]
        for i, fila in enumerate(h["rows"]):
            for j, c in enumerate(fila):
                if not isinstance(c, dict):
                    continue
                n += 1
                xl = ws.cell(row=FILA0 + i, column=j + 1).value
                if not igual(c["v"], xl):
                    fallas += 1
                    print(f"[{nombre}] {h['name']}!{ws.cell(row=FILA0 + i, column=j + 1).coordinate}: "
                          f"motor={c['v']!r} LibreOffice={xl!r}  ={c['f'][:120]}")
    for ws in wb.worksheets:                     # ninguna celda en error, en ninguna hoja
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith(ERRORES):
                    fallas += 1
                    print(f"[{nombre}] {ws.title}!{c.coordinate}: {c.value}")
    # Tarjetas de la portada = panel del HTML (misma cifra).
    p = graficos.panel(graficos.modulo(d), reg["run"], reg["run"]["hojas"])
    esperado = {"principal": p["principal"]["valor"], "registrado": p["registrado"]["valor"],
                "recalculado": p["recalculado"]["valor"], "problemas": p["problemas"]["valor"]}
    ini = wb["00_Inicio"]
    rotulos = {str(ini.cell(row=13, column=c).value or "").upper(): ini.cell(row=14, column=c).value for c in range(2, 7)}
    for clave, rot in (("principal", p["principal"]["rotulo"]), ("registrado", p["registrado"]["rotulo"]),
                       ("recalculado", p["recalculado"]["rotulo"]), ("problemas", p["problemas"]["rotulo"])):
        if rot.upper() in rotulos:
            n += 1
            if not igual(float(esperado[clave] or 0), rotulos[rot.upper()]):
                fallas += 1
                print(f"[{nombre}] tarjeta «{rot}»: panel={esperado[clave]!r} Excel={rotulos[rot.upper()]!r}")
    print(f"[{nombre}] celdas comparadas: {n}")
    return fallas, n


if __name__ == "__main__":
    base = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tests", "fixtures", "papel_declarativo")
    rutas = sys.argv[1:] or sorted(glob.glob(os.path.join(base, "*.json")))
    total_f = total_n = 0
    with tempfile.TemporaryDirectory() as tmp:
        for r in rutas:
            f, n = verificar(r, tmp)
            total_f += f
            total_n += n
    print(f"CELDAS: {total_n} · DIFERENCIAS: {total_f}")
    sys.exit(1 if total_f else 0)
