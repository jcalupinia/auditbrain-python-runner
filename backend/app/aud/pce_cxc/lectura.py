"""Lectura de los archivos de cartera del cliente.

Los archivos llegan con el formato regional del sistema que los generó. Aquí se
normalizan importes y fechas antes de que el motor vea un solo número.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import load_workbook

from backend.app.aud.pce_cxc.bandas import clasificar
from backend.app.aud.pce_cxc.motor import redondear
from backend.app.aud.pce_cxc.texto import limpiar_texto

_PATRON_DMY = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$")
_PATRON_ISO = re.compile(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})")


def _signo_delante(s: str) -> bool:
    """¿El primer guion de la cadena es un signo menos y no un separador?

    Lo es cuando delante no hay ninguna cifra y el carácter inmediatamente
    anterior no es alfanumérico: "-987,65", "USD -1.500,00" y "$-1500" son
    negativos; "F-1234", "1234-5678" y "USD 1,500.00 s/n-c" no lo son.
    """
    i = s.find("-")
    if i < 0:
        return False
    previo = s[:i]
    if previo and previo[-1].isalnum():
        return False
    return not any(c.isdigit() for c in previo)


def a_numero(valor: Any) -> float:
    """Convierte a número respetando cualquier formato regional.

    - Si aparecen ambos separadores ("." y ","): el último que aparece es el
      decimal, el otro es de miles.
    - Si aparece un solo tipo de separador:
      - más de una vez -> todos son de miles ("8.917.458" -> 8917458);
      - una sola vez y le siguen exactamente tres dígitos -> es de miles
        ("1.500" -> 1500);
      - una sola vez y no le siguen exactamente tres dígitos -> es decimal
        ("1234.56" -> 1234.56).

    El importe es negativo solo si trae un SIGNO de verdad: envuelto en
    paréntesis ("(1.500,00)"), con el guion delante del primer dígito
    ("-987,65", "USD -1.500,00") o con el guion al final, como lo escriben
    algunos ERP ("1.500,00-"). Un guion INTERIOR entre dígitos es parte del
    dato, no un signo: "1234-5678" vale 12.345.678, no -12.345.678. Como el
    mapeo de columnas es automático, tratar cualquier guion como signo convertía
    una columna mal mapeada (números de documento tipo "F-1234") en una
    exposición negativa enorme.
    """
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)
    s = str(valor).strip()
    if not s:
        return 0.0
    negativo = ((s.startswith("(") and s.endswith(")"))
                or _signo_delante(s)
                or s.endswith("-"))
    s = re.sub(r"[^0-9.,]", "", s)
    if not s:
        return 0.0
    tiene_punto = "." in s
    tiene_coma = "," in s
    entero, decimal = s, ""
    if tiene_punto and tiene_coma:
        i = max(s.rfind(","), s.rfind("."))
        entero, decimal = s[:i], s[i + 1:]
    elif tiene_punto or tiene_coma:
        sep = "." if tiene_punto else ","
        if s.count(sep) == 1:
            i = s.rfind(sep)
            cola = s[i + 1:]
            if len(cola) != 3:
                entero, decimal = s[:i], cola
    entero = re.sub(r"[.,]", "", entero)
    try:
        n = float(f"{entero}.{decimal}" if decimal else entero or "0")
    except ValueError:
        return 0.0
    return -n if negativo and n > 0 else n


def inferir_formato_fecha(valores: Iterable[Any]) -> str:
    """Deduce si las fechas del archivo son día/mes/año o mes/día/año.

    Recorre todos los valores (no se detiene en el primero) porque un Excel
    puede mezclar celdas con formato de fecha nativo y celdas de texto:

    - solo fechas nativas -> "nativo";
    - fechas nativas y además cadenas con patrón de fecha -> "inconsistente"
      (el archivo mezcla formatos);
    - solo cadenas -> "dmy"/"mdy"/"inconsistente" según qué componente supere
      12, o "ambiguo" si ninguno lo delata;
    - nada parseable -> "ambiguo" (no se sabe, no se asume "nativo").
    """
    dmy = mdy = total = 0
    hay_nativas = False
    hay_texto_con_patron = False
    for v in valores:
        if isinstance(v, (date, datetime)):
            hay_nativas = True
            continue
        m = _PATRON_DMY.match(str(v or "").strip())
        if not m:
            continue
        hay_texto_con_patron = True
        total += 1
        if int(m.group(1)) > 12:
            dmy += 1
        if int(m.group(2)) > 12:
            mdy += 1
    if hay_nativas and hay_texto_con_patron:
        return "inconsistente"
    if hay_nativas:
        return "nativo"
    if not total:
        return "ambiguo"
    if dmy and mdy:
        return "inconsistente"
    if dmy:
        return "dmy"
    if mdy:
        return "mdy"
    return "ambiguo"


def a_fecha(valor: Any, formato: str = "dmy") -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if valor is None or valor == "":
        return None
    s = str(valor).strip()
    m = _PATRON_DMY.match(s)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        anio = int(m.group(3))
        anio += 2000 if anio < 100 else 0
        dia, mes = (b, a) if formato == "mdy" else (a, b)
        try:
            return date(anio, mes, dia)
        except ValueError:
            return None
    m = _PATRON_ISO.match(s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


# Pistas para reconocer cada columna en el archivo del cliente.
CAMPOS: dict[str, list[str]] = {
    "cliente": ["cliente", "razon", "razón", "nombre", "deudor"],
    "documento": ["documento", "comprobante", "factura", "numero", "número", "nro", "n°"],
    "tipo": ["tipo", "relacion", "relación", "clasif", "categoria", "categoría"],
    "emision": ["emision", "emisión", "fecha emis", "f. emis"],
    "vencimiento": ["vencimiento", "vence", "venc"],
    "saldo": ["saldo", "monto", "valor", "importe", "total", "cuentas por cobrar"],
}
_OBLIGATORIOS = ("documento", "vencimiento", "saldo")

#: Motivo de descarte que NO representa cartera perdida: la fila idéntica
#: repetida se depura a propósito y su importe ya está contado una vez.
MOTIVO_FILA_REPETIDA = "fila idéntica repetida"

_SEPARADORES = re.compile(r"[^0-9A-ZÁÉÍÓÚÜÑ]+")


def _es_relacionada(tipo: Any, clave: str) -> bool:
    """¿El tipo de cliente del archivo corresponde a una parte relacionada?

    Normaliza cualquier separador antes de comparar, porque el archivo del
    cliente escribe lo mismo de muchas formas: "NO-RELACIONADOS",
    "NO RELACIONADOS", "NO_RELACIONADOS", "NO.RELACIONADOS" y
    "NORELACIONADOS" son todos TERCEROS. Reconocer solo dos de esas formas
    mudaba toda la cartera de terceros al segmento de relacionadas y dejaba
    mal el anclaje de los dos segmentos a la vez.
    """
    texto = _SEPARADORES.sub(" ", str(tipo or "").upper()).strip()
    patron = _SEPARADORES.sub(" ", str(clave or "").upper()).strip()
    if not patron or patron.replace(" ", "") not in texto.replace(" ", ""):
        return False
    # "NO" pegado o separado del patrón, al inicio o tras un espacio, niega la
    # relación ("TERCEROS NO RELACIONADOS"); "INTERNO RELACIONADAS" no, porque
    # ahí el "NO" es el final de otra palabra.
    partes = r"\s*".join(re.escape(p) for p in patron.split())
    return not re.search(rf"(?:^|\s)NO\s*{partes}", texto)


def _detectar_encabezado(filas: list[tuple]) -> int:
    mejor, puntaje = -1, 0
    for i, fila in enumerate(filas[:45]):
        celdas = [str(c or "").lower() for c in fila]
        p = sum(1 for pistas in CAMPOS.values() if any(any(h in c for h in pistas) for c in celdas))
        if p > puntaje:
            mejor, puntaje = i, p
    return mejor if puntaje >= 3 else -1


def _mapear(encabezado: tuple) -> dict[str, int]:
    celdas = [str(c or "").lower().strip() for c in encabezado]
    mapeo: dict[str, int] = {}
    usadas: set[int] = set()
    for campo, pistas in CAMPOS.items():
        for pista in pistas:
            idx = next((j for j, c in enumerate(celdas) if pista in c and j not in usadas), None)
            if idx is not None:
                mapeo[campo] = idx
                usadas.add(idx)
                break
    return mapeo


#: Lo único que este lector sabe abrir. Aceptar `.xls` exigiría una
#: dependencia nueva (`xlrd`) y `.csv` cambiaría el perfil de memoria con el
#: que está calibrado el límite por archivo, así que la pantalla ofrece
#: exactamente esto y el backend lo dice cuando llega otra cosa.
EXTENSIONES_LEIBLES = (".xlsx", ".xlsm")


def _instruccion_por_extension(nombre: str) -> str:
    """Qué hacer con este archivo, según lo que parece ser.

    Un `.csv` no es un libro dañado y un `.xls` no es un `.xlsx` roto: cada uno
    tiene su propio arreglo, y un mensaje genérico obliga al auditor a
    adivinarlo.
    """
    bajo = nombre.lower()
    if bajo.endswith((".csv", ".txt", ".tsv")):
        return ("Es un archivo de texto separado por delimitadores, no un libro de Excel. Ábralo "
                "en Excel y guárdelo con «Guardar como → Libro de Excel (*.xlsx)».")
    if bajo.endswith(".xls"):
        return ("Está en el formato binario antiguo de Excel (.xls), que esta herramienta no lee. "
                "Ábralo en Excel y guárdelo con «Guardar como → Libro de Excel (*.xlsx)».")
    return ("No se pudo abrir como libro de Excel (.xlsx): puede estar dañado, protegido con "
            "contraseña o no ser realmente un .xlsx pese a su nombre. Ábralo en Excel, compruebe "
            "que se ve bien y guárdelo de nuevo con «Guardar como → Libro de Excel (*.xlsx)».")


def _abrir_libro(contenido: bytes, nombre: str):
    """Abre el análisis de antigüedad, o explica por qué no se pudo.

    `openpyxl` levanta `zipfile.BadZipFile` -que no es `ValueError`- con un
    `.csv`, con un `.xls` binario y con un `.xlsx` corrupto, y también puede
    levantar `KeyError` o `OSError` con un zip válido que no es un libro. Sin
    esta traducción, los tres formatos que la pantalla llegó a ofrecer salían
    por el 500 genérico: un error de entrada sin una sola instrucción.
    """
    try:
        return load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - la librería no declara un tipo común
        raise ValueError(
            f"{nombre}: {_instruccion_por_extension(nombre)} La herramienta lee libros "
            f"{' o '.join(EXTENSIONES_LEIBLES)}."
        ) from e


def leer_cartera(contenido: bytes, nombre: str, corte, bandas, hoja=None, mapeo=None,
                 clave_relacionadas: str = "RELACIONAD",
                 fila_encabezado: int | None = None) -> dict:
    """Lee un análisis de antigüedad y devuelve sus documentos clasificados por mora.

    `fila_encabezado` es opcional y se numera como lo ve el usuario (1 es la
    primera fila de la hoja). Si viene, manda sobre la detección automática y
    sobre el encabezado en la primera fila que se asume cuando se pasa `mapeo`
    sin indicarlo: el mapeo manual suele usarse justo cuando el encabezado no
    está en la primera fila, que es cuando la detección automática falla.
    """
    wb = _abrir_libro(contenido, nombre)
    if hoja and hoja not in wb.sheetnames:
        disponibles = ", ".join(wb.sheetnames)
        wb.close()
        raise ValueError(
            f"{nombre}: el libro no tiene ninguna hoja llamada '{hoja}'. Hojas disponibles: "
            f"{disponibles}. Elija una de esas, o deje la hoja sin indicar para leer la primera."
        )
    ws = wb[hoja] if hoja else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    if fila_encabezado is not None:
        i_enc = fila_encabezado - 1
        # Valida que fila_encabezado esté dentro del rango válido (1 <= fila_encabezado <= len(filas))
        if i_enc < 0 or i_enc >= len(filas):
            raise ValueError(f"{nombre}: fila_encabezado {fila_encabezado} fuera de rango (la hoja tiene {len(filas)} filas)")
    else:
        i_enc = 0 if mapeo else _detectar_encabezado(filas)
        if i_enc < 0:
            raise ValueError(f"{nombre}: no se identificó la fila de encabezados")
    cols = mapeo or _mapear(filas[i_enc])
    faltantes = [c for c in _OBLIGATORIOS if c not in cols]
    if faltantes:
        raise ValueError(f"{nombre}: no se encontraron las columnas {', '.join(faltantes)}")

    cuerpo = filas[i_enc + 1:]
    formato = inferir_formato_fecha([f[cols["vencimiento"]] for f in cuerpo[:4000]
                                     if len(f) > cols["vencimiento"]])

    salida, descartados = [], []
    vistas: set[tuple] = set()
    documentos: set[str] = set()
    dup_exactos = repetidos = 0

    def valor(fila, campo):
        j = cols.get(campo)
        return fila[j] if j is not None and j < len(fila) else None

    for n, fila in enumerate(cuerpo, start=i_enc + 2):
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            continue
        # El texto del cliente se limpia AL ENTRAR, no al escribirlo: un
        # carácter de control que el XML de Excel no admite hace que guardar el
        # libro levante una excepción -un 500 por un dato de entrada- y, si el
        # resultado ya se guardó con él, la descarga falla para siempre. El
        # dato no se altera de ninguna otra forma (ver `texto.limpiar_texto`).
        documento = limpiar_texto(valor(fila, "documento") or "").strip()
        saldo = a_numero(valor(fila, "saldo"))
        vencimiento = a_fecha(valor(fila, "vencimiento"), formato)
        if not documento:
            descartados.append({"fila_origen": n, "motivo": "sin número de documento", "saldo": saldo})
            continue
        if vencimiento is None:
            descartados.append({"fila_origen": n, "motivo": "sin fecha de vencimiento", "saldo": saldo})
            continue
        if abs(saldo) < 0.005:
            descartados.append({"fila_origen": n, "motivo": "saldo cero", "saldo": saldo})
            continue
        cliente = limpiar_texto(valor(fila, "cliente") or "").strip()
        # `motor.redondear`, nunca `round()`: la regla del proyecto para todo
        # importe contable. `round()` redondea al par más cercano y arrastra la
        # representación binaria del float, así que 2,675 y 2,674 daban los dos
        # 2,67 y dos saldos DISTINTOS al centavo compartían firma: el segundo
        # documento se descartaba como «fila idéntica repetida» y su importe
        # salía de la medición.
        firma = (documento, redondear(saldo), vencimiento, cliente)
        if firma in vistas:
            dup_exactos += 1
            descartados.append({"fila_origen": n, "motivo": "fila idéntica repetida", "saldo": saldo})
            continue
        vistas.add(firma)
        repetidos += 1 if documento in documentos else 0
        documentos.add(documento)
        es_rel = _es_relacionada(valor(fila, "tipo"), clave_relacionadas)
        dias = (corte - vencimiento).days
        salida.append({
            "fila_origen": n, "cliente": cliente, "documento": documento,
            "segmento": "RELACIONADOS" if es_rel else "NO-RELACIONADOS",
            "emision": a_fecha(valor(fila, "emision"), formato),
            "vencimiento": vencimiento, "saldo": saldo, "dias": dias,
            "banda": clasificar(dias, bandas), "repetido": False,
        })

    # Marca TODAS las filas de un documento que aparece más de una vez (la
    # primera incluida), no solo la segunda en adelante: un auditor que filtre
    # por "repetido" debe encontrar el par completo.
    conteo_documentos: dict[str, int] = {}
    for f in salida:
        conteo_documentos[f["documento"]] = conteo_documentos.get(f["documento"], 0) + 1
    for f in salida:
        f["repetido"] = conteo_documentos[f["documento"]] > 1

    # El importe de lo descartado se totaliza y se desglosa por motivo: guardar
    # solo el conteo hacía desaparecer el dinero que no se pudo leer, y con
    # anclaje a los estados financieros ese hueco se reparte sobre las filas que
    # sí entraron (se convierte en exposición inventada).
    acumulado: dict[str, dict[str, Any]] = {}
    for d in descartados:
        m = acumulado.setdefault(d["motivo"], {"motivo": d["motivo"], "filas": 0, "importe": 0.0})
        m["filas"] += 1
        m["importe"] += d["saldo"]
    por_motivo = [{**m, "importe": redondear(m["importe"])}
                  for m in sorted(acumulado.values(), key=lambda m: m["motivo"])]
    descartados_importe = redondear(sum(d["saldo"] for d in descartados))
    cartera_no_leida = redondear(sum(d["saldo"] for d in descartados
                                    if d["motivo"] != MOTIVO_FILA_REPETIDA))

    return {"filas": salida, "hoja": ws.title, "fila_encabezado": i_enc + 1, "mapeo": cols,
            "formato_fecha": formato, "duplicados_exactos": dup_exactos,
            "documentos_repetidos": repetidos, "descartados": descartados,
            "descartados_importe": descartados_importe,
            "descartados_por_motivo": por_motivo,
            "cartera_no_leida": cartera_no_leida,
            "total_saldo": redondear(sum(f["saldo"] for f in salida))}
