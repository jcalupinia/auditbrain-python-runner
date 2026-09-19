"""Extracción acotada. No ejecutar fórmulas, macros ni extraer ZIP al disco."""
import csv
import hashlib
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
import xml.etree.ElementTree as ET
from .engine import MAX_ROWS

MAX_BYTES=8*1024*1024
MAX_EXPANDED=24*1024*1024
MAX_COLS=40


def bounded_table(name, raw):
    values=[]
    for row in raw:
        if not any(v is not None and str(v).strip() for v in row):continue
        if len(row)>MAX_COLS or len(values)>MAX_ROWS:raise ValueError('Tabla supera 2000 filas o 40 columnas; divida la población.')
        values.append(['' if v is None else str(v) for v in row])
    if not values:return None
    return {'name':name,'headers':values[0],'rows':values[1:]}


def checked_zip(data):
    z=ZipFile(BytesIO(data)); entries=z.infolist()
    if len(entries)>200 or sum(i.file_size for i in entries)>MAX_EXPANDED:
        z.close();raise ValueError('ZIP supera límites de expansión.')
    for i in entries:
        path=PurePosixPath(i.filename.replace('\\','/'))
        if path.is_absolute() or '..' in path.parts or i.flag_bits&1 or i.file_size>MAX_BYTES*3:
            z.close();raise ValueError('ZIP contiene rutas o archivos no admitidos.')
        if i.file_size>1024*1024 and i.file_size/max(i.compress_size,1)>200:
            z.close();raise ValueError('ZIP con compresión excesiva.')
    return z


def extract(name,data,*,nested=False):
    if not data or len(data)>MAX_BYTES:raise ValueError('Archivo vacío o mayor a 8 MB.')
    name=PurePosixPath(name.replace('\\','/')).name[:200]
    ext=PurePosixPath(name).suffix.lower()
    result={'name':name,'sha256':hashlib.sha256(data).hexdigest(),'size':len(data),'tables':[],'text':'','note':'Revise la extracción y asigne columnas antes de procesar.'}
    try:
        if ext=='.csv':
            text=data.decode('utf-8-sig')
            try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            t=bounded_table(name,csv.reader(StringIO(text),dialect))
            if t:result['tables'].append(t)
        elif ext=='.xlsx':
            with checked_zip(data):pass
            import openpyxl
            wb=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
            try:
                if len(wb.worksheets)>20:raise ValueError('Excel supera 20 hojas.')
                for ws in wb.worksheets:
                    if ws.max_row>MAX_ROWS+1 or ws.max_column>MAX_COLS:raise ValueError('Hoja supera límites; use tabla de hasta 2000 filas y 40 columnas.')
                    t=bounded_table(ws.title,ws.iter_rows(values_only=True))
                    if t:result['tables'].append(t)
            finally:wb.close()
            result['note']='Fórmulas de origen no se ejecutan. Exporte valores o transcriba y revise las celdas con fórmulas.'
        elif ext=='.xml':
            if b'\x00' in data or b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():raise ValueError('XML con DTD o entidades no admitido.')
            root=ET.fromstring(data)
            records=list(root)
            if len(records)>MAX_ROWS:raise ValueError('XML supera 2000 registros.')
            if records and all(list(r) for r in records):
                headers=list(dict.fromkeys(c.tag for r in records for c in r))
                t=bounded_table(name,[headers]+[[r.findtext(k,'') for k in headers] for r in records])
                if t:result['tables'].append(t)
            else:result['text']=data.decode('utf-8')[:20000]
        elif ext=='.docx':
            with checked_zip(data):pass
            from docx import Document
            doc=Document(BytesIO(data))
            if len(doc.tables)>20:raise ValueError('Documento supera 20 tablas.')
            for i,table in enumerate(doc.tables):
                t=bounded_table(f'Tabla {i+1}',([c.text for c in r.cells] for r in table.rows))
                if t:result['tables'].append(t)
            result['text']='\n'.join(p.text for p in doc.paragraphs)[:20000]
        elif ext=='.pdf':
            from pypdf import PdfReader
            pdf=PdfReader(BytesIO(data))
            if pdf.is_encrypted or len(pdf.pages)>60:raise ValueError('PDF cifrado o mayor a 60 páginas no admitido.')
            # Text is supporting evidence, never guessed numeric inventory rows.
            result['text']='\n'.join((p.extract_text() or '') for p in pdf.pages)[:20000]
            result['note']='Texto de soporte; transcripción manual y revisión de datos numéricos. PDF escaneado requiere OCR externo.'
        elif ext in ('.png','.jpg','.jpeg','.webp'):
            result['note']='Imagen recibida como soporte; requiere transcripción manual u OCR externo. No se han extraído importes.'
        elif ext=='.txt':result['text']=data.decode('utf-8-sig')[:20000]
        elif ext=='.zip':
            if nested:raise ValueError('No se admiten ZIP anidados.')
            with checked_zip(data) as z:
                files=[i for i in z.infolist() if not i.is_dir()]
                if len(files)>20:raise ValueError('ZIP: máximo 20 documentos.')
                children=[extract(i.filename,z.read(i),nested=True) for i in files]
                if sum(len(t['rows']) for c in children for t in c['tables'])>MAX_ROWS:raise ValueError('ZIP supera 2000 filas extraídas.')
            result['tables']=[{**t,'name':c['name']+' / '+t['name']} for c in children for t in c['tables']]
            result['text']='\n'.join(c['name']+': '+c['text']+' '+c['note'] for c in children)[:20000]
            result['children']=[{k:v for k,v in c.items() if k not in ('tables','text')} for c in children]
        else:raise ValueError('Use XLSX, CSV, XML, DOCX, PDF, TXT, ZIP o imagen PNG/JPG/WEBP. Convierta XLS/DOC antiguos.')
    except ValueError:raise
    except Exception as exc:raise ValueError('No se pudo leer el documento; compruebe formato e integridad.') from exc
    return result
