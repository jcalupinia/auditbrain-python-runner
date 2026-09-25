"""Word y PowerPoint del papel de trabajo NIIF con el diseño del HTML.

Pedido del dueño (2026-09-25): «haz el Word y el PowerPoint igual al HTML».

- **PowerPoint = el HTML en pantalla (tema «Ejecutivo»)**: fondo azul marino, portada con
  los logotipos, título y chips; tarjetas KPI con sus colores; los 4 gráficos del panel
  (los mismos SVG del HTML dibujados como imagen, ``svg_png``); cédulas en tablas oscuras
  con cabecera y total con filete dorado.
- **Word = el HTML impreso (tema «Claro», el mismo del PDF)**: membrete con los dos
  logotipos y filete dorado, título, chips, tarjetas, los 4 gráficos y cada cédula con su
  bloque «ⓘ Cómo se calcula esta hoja» y su tabla (cabecera gris con filete dorado, filas
  alternas, total en negrita). Un documento para imprimir no lleva fondo oscuro: Word no
  imprime el color de página y el texto claro quedaría invisible.

Las cifras, rótulos y variaciones de las tarjetas salen de ``html_ejecutivo.kpis_datos``
y los gráficos de ``graficos.panel``: los tres formatos leen de la misma fuente.
"""
from __future__ import annotations

import html as _html
import io

from backend.app.aud.niif.procesadores import estilo_ejecutivo as est
from backend.app.aud.niif.procesadores import graficos
from backend.app.aud.niif.procesadores import html_ejecutivo as hx
from backend.app.aud.niif.procesadores import marca
from backend.app.aud.niif.procesadores import svg_png

CLARO = {k: v.lstrip("#").upper() for k, v in hx.TEMAS["claro"][1].items() if isinstance(v, str) and v.startswith("#")}
OSCURO = {k: v.lstrip("#").upper() for k, v in hx.TEMAS["ejecutivo"][1].items() if isinstance(v, str) and v.startswith("#")}
# «zebra» del tema oscuro es blanco al 3,5 % sobre la tarjeta: en PowerPoint va ya mezclado.
OSCURO["zebra"] = "132B49"
_CLASE = {"k-oro": "oro-txt", "k-azul": "k-azul", "k-verde": "k-verde", "k-ambar": "k-ambar",
          "k-alto": "alta", "k-medio": "media", "k-bajo": "baja"}
PIE = ("AuditConsulting Auditores Cía. Ltda. · AUDIT-IA · Papel de trabajo generado por la herramienta; "
       "las cifras son de la prueba ejecutada.")


def _panel(definicion, reg):
    run = reg.get("run") or {}
    return graficos.panel(graficos.modulo(definicion), run, run.get("hojas") or [])


def _chips(p, e, version, estado):
    riesgo = {"alto": "▲ Riesgo alto", "medio": "▲ Riesgo medio", "bajo": "▲ Riesgo bajo"}[p["riesgo"]]
    return [(riesgo, {"alto": "alta", "medio": "media", "bajo": "baja"}[p["riesgo"]]),
            (f"Corte {hx._fecha(e.get('cutoff')) or 'pendiente'}", None), (f"Versión {version}", None),
            (est.estado_es(estado), None), (str(e.get("framework") or "Marco pendiente"), None)]


def _var(k):
    """(texto con flecha, clave de color) de la línea de variación de una tarjeta."""
    if "var" in k:
        v, texto = k["var"]
        if v is None:
            return "", None
        return (f"{'↗' if v >= 0 else '↘'} {hx.gs.pct(abs(v))} {texto}", "sube" if v >= 0 else "baja-txt")
    return k.get("nota") or "", None


# =============================================================================================
# Word (tema Claro)
# =============================================================================================

def _w():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    return OxmlElement, qn


_SUC_TC = {"tcBorders": ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"),
           "shd": ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"),
           "tcMar": ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")}
_SUC_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
             "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
             "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
             "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_SUC_TBLBORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _tc(celda, nombre):
    """Elemento hijo de tcPr (se crea en el orden que exige el esquema de Word)."""
    OxmlElement, qn = _w()
    tcPr = celda._tc.get_or_add_tcPr()
    el = tcPr.find(qn(f"w:{nombre}"))
    if el is None:
        el = OxmlElement(f"w:{nombre}")
        tcPr.insert_element_before(el, *_SUC_TC[nombre])
    return el


def _sombra(celda, color):
    _, qn = _w()
    shd = _tc(celda, "shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)


def _bordes(celda, **lados):
    """lados: top/left/bottom/right = (grosor en octavos de punto, color) o None (sin borde)."""
    OxmlElement, qn = _w()
    tb = _tc(celda, "tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        if lado not in lados:
            continue
        viejo = tb.find(qn(f"w:{lado}"))
        if viejo is not None:
            tb.remove(viejo)
        el = OxmlElement(f"w:{lado}")
        if lados[lado] is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, color = lados[lado]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        tb.append(el)
    # Orden del esquema dentro de tcBorders: top, left (start), bottom, right (end)…
    orden = {qn("w:top"): 0, qn("w:left"): 1, qn("w:bottom"): 2, qn("w:right"): 3}
    for el in sorted(list(tb), key=lambda x: orden.get(x.tag, 9)):
        tb.remove(el)
        tb.append(el)


def _margen(celda, pt=5):
    OxmlElement, qn = _w()
    mar = _tc(celda, "tcMar")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(int(pt * 20)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)


def _sin_bordes(tabla):
    OxmlElement, qn = _w()
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblPr.insert_element_before(b, *_SUC_TBLBORDERS)


def _filete(par, color, sz=12, lado="bottom"):
    OxmlElement, qn = _w()
    pPr = par._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    pb.append(el)
    pPr.insert_element_before(pb, *_SUC_PBDR)


def _no_partir(fila):
    OxmlElement, qn = _w()
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _texto(par, txt, sz, color, negrita=False, fuente=None, cursiva=False):
    from docx.shared import Pt, RGBColor

    r = par.add_run(txt)
    r.font.size = Pt(sz)
    r.font.bold = negrita
    r.font.italic = cursiva
    r.font.color.rgb = RGBColor.from_string(color)
    if fuente:
        r.font.name = fuente
    return r


def _parrafo(contenedor, txt="", sz=9, color=None, negrita=False, antes=0, despues=2, alinear=None, fuente=None):
    from docx.shared import Pt

    par = contenedor.add_paragraph()
    par.paragraph_format.space_before = Pt(antes)
    par.paragraph_format.space_after = Pt(despues)
    if alinear is not None:
        par.alignment = alinear
    if txt:
        _texto(par, txt, sz, color or CLARO["texto"], negrita, fuente)
    return par


def _celda_par(celda, primero=True):
    return celda.paragraphs[0] if primero and not celda.paragraphs[0].text else celda.add_paragraph()


def docx(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    from backend.app.aud.niif.procesadores import libro as L

    C = CLARO
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(1.6))
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
    sec.header_distance, sec.footer_distance = Cm(0.7), Cm(0.7)
    ancho = sec.page_width - sec.left_margin - sec.right_margin
    normal = doc.styles["Normal"]
    normal.font.name = est.FONT_TEXTO
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2)
    e = reg.get("engagement") or {}
    p = _panel(definicion, reg)

    # Membrete (como el del PDF): logos a izquierda y derecha y filete dorado.
    enc = sec.header
    t = enc.add_table(rows=1, cols=2, width=ancho)
    _sin_bordes(t)
    iz, de = t.rows[0].cells
    iz.paragraphs[0].add_run().add_picture(marca.flujo("auditconsulting_oscuro"), height=Cm(1.15))
    de.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    de.paragraphs[0].add_run().add_picture(marca.flujo("audit_ia"), height=Cm(1.15))
    # El párrafo vacío que trae el encabezado quedaba ARRIBA de los logos: se quita y el filete
    # dorado va en un párrafo debajo de ellos (como la línea del membrete del HTML impreso).
    vacio = enc.paragraphs[0]._p
    vacio.getparent().remove(vacio)
    linea = enc.add_paragraph()
    linea.paragraph_format.space_after = Pt(0)
    _filete(linea, C["oro"], 12, "top")
    pie = sec.footer.paragraphs[0]
    _texto(pie, PIE, 7, C["texto2"])

    # Panel: título, cliente, chips.
    _parrafo(doc, definicion.get("name", ""), 20, C["texto"], True, despues=2)
    _parrafo(doc, f"{e.get('client', '')} · RUC {e.get('ruc', '')} · {definicion.get('area', '')}", 10, C["texto2"], despues=6)
    chips = _chips(p, e, version, estado)
    tc = doc.add_table(rows=1, cols=len(chips))
    _sin_bordes(tc)
    for celda, (txt, color) in zip(tc.rows[0].cells, chips):
        celda.width = Cm(4.2)
        _sombra(celda, C["card2"])
        borde = (6, C[color] if color else C["borde"])
        _bordes(celda, top=borde, left=borde, bottom=borde, right=borde)
        par = celda.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _texto(par, txt, 8.5, C[color] if color else C["texto2"], True)
    _parrafo(doc, "", despues=4)

    # Tarjetas KPI (las 5 del HTML).
    kp = hx.kpis_datos(p)
    tk = doc.add_table(rows=1, cols=len(kp))
    _sin_bordes(tk)
    tk.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, k in zip(tk.rows[0].cells, kp):
        celda.width = int(ancho / len(kp))
        _sombra(celda, C["card"])
        b = (8, C["borde"])
        _bordes(celda, top=b, left=b, bottom=b, right=b)
        _margen(celda, 6)
        _texto(celda.paragraphs[0], k["rotulo"].upper(), 7.5, C["texto2"], True)
        _texto(_celda_par(celda, False), k["valor"], 15, C[_CLASE[k["clase"]]], True)
        if k["exacto"]:
            _texto(_celda_par(celda, False), k["exacto"], 8, C["muted"], fuente="Consolas")
        txt, color = _var(k)
        if txt:
            _texto(_celda_par(celda, False), txt, 8, C[color] if color else C["texto2"], True)
    _parrafo(doc, "", despues=6)

    # Los 4 gráficos del panel del HTML, en rejilla 2×2 dentro de tarjetas.
    graf = svg_png.graficos_panel(p, "claro")
    tg = doc.add_table(rows=2, cols=2)
    _sin_bordes(tg)
    ancho_img = int(ancho / 2) - Cm(0.9)
    for k, g in enumerate(graf):
        celda = tg.rows[k // 2].cells[k % 2]
        _sombra(celda, C["card"])
        b = (8, C["borde"])
        _bordes(celda, top=b, left=b, bottom=b, right=b)
        _margen(celda, 7)
        _texto(celda.paragraphs[0], g["titulo"], 10.5, C["texto"], True)
        _texto(_celda_par(celda, False), g["sub"], 8, C["texto2"])
        if g["png"]:
            _celda_par(celda, False).add_run().add_picture(io.BytesIO(g["png"]), width=ancho_img)
        else:
            _texto(_celda_par(celda, False), "Sin datos para graficar en este ejemplo.", 8.5, C["muted"], cursiva=True)
    for fila in tg.rows:
        _no_partir(fila)

    # Cada cédula en su página, como las secciones del HTML impreso.
    hojas = L.cedulas(definicion, reg, eventos, version, estado)
    for h in hojas:
        doc.add_page_break()
        cab = _parrafo(doc, h["label"], 15, C["texto"], True, despues=6)
        _filete(cab, C["oro"], 16)
        bloque = L.como_se_calcula(h, hojas)
        if bloque:
            _calc_word(doc, bloque, ancho)
        _tabla_word(doc, h, L)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def _calc_word(doc, bloque, ancho):
    """«ⓘ Cómo se calcula esta hoja» (el recuadro del HTML): fondo gris claro, título dorado."""
    from docx.shared import Pt

    C = CLARO
    caja = doc.add_table(rows=1, cols=1)
    _sin_bordes(caja)
    celda = caja.rows[0].cells[0]
    _sombra(celda, C["card2"])
    b = (8, C["borde"])
    _bordes(celda, top=b, left=b, bottom=b, right=b)
    _margen(celda, 6)
    _texto(celda.paragraphs[0], "ⓘ  Cómo se calcula esta hoja", 10, C["oro-txt"], True)
    cols = ["Columna", "Fórmula", "Cómo se calcula", "Ejemplo con números reales", "De dónde viene"]
    t = celda.add_table(rows=1 + len(bloque), cols=len(cols))
    _sin_bordes(t)
    for j, nombre in enumerate(cols):
        c = t.rows[0].cells[j]
        _bordes(c, bottom=(12, C["oro"]))
        _texto(c.paragraphs[0], nombre, 8, C["texto"], True)
    for i, blq in enumerate(bloque, start=1):
        for j, val in enumerate([blq["columna"], blq["formula"], blq["explicacion"], blq["ejemplo"], blq["origen"]]):
            c = t.rows[i].cells[j]
            _bordes(c, bottom=(4, C["borde"]))
            _texto(c.paragraphs[0], str(val), 7.5, C["texto"], j == 0, "Consolas" if j in (1, 3) else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _tabla_word(doc, h, L):
    """Tabla de la cédula con el estilo del HTML: cabecera gris con filete dorado, filas
    alternas, cifras a la derecha y total en negrita con filete dorado arriba."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    C = CLARO
    filas = L._filas(h)
    n = len(h["cols"])
    sz = 8 if n <= 7 else 7 if n <= 11 else 6.5
    t = doc.add_table(rows=1 + len(filas), cols=n)
    _sin_bordes(t)
    for j, (nombre, _) in enumerate(h["cols"]):
        c = t.rows[0].cells[j]
        _sombra(c, C["card2"])
        _bordes(c, bottom=(16, C["oro"]))
        _texto(c.paragraphs[0], nombre, sz, C["texto"], True)
    for i, (fila, total) in enumerate(filas, start=1):
        for j, ((_, fmt), v) in enumerate(zip(h["cols"], fila)):
            c = t.rows[i].cells[j]
            if total:
                _sombra(c, C["card2"])
                _bordes(c, top=(16, C["oro"]), bottom=(4, C["borde"]))
            else:
                if i % 2 == 0:
                    _sombra(c, C["zebra"])
                _bordes(c, bottom=(4, C["borde"]))
            num = fmt in ("n", "p", "i", "g")
            par = c.paragraphs[0]
            if num:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _texto(par, _html.unescape(L._celda(v, fmt)), sz, C["texto"], total, "Consolas" if num else None)
    for fila in t.rows:
        _no_partir(fila)


# =============================================================================================
# PowerPoint (tema Ejecutivo)
# =============================================================================================

def _rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)


def _fondo(diap):
    f = diap.background.fill
    f.solid()
    f.fore_color.rgb = _rgb(OSCURO["bg"])


def _caja(diap, x, y, w, h, relleno, borde=None, redondeado=True):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    s = diap.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if redondeado else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if redondeado:
        s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(relleno)
    if borde:
        s.line.color.rgb = _rgb(borde)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _tx(diap, x, y, w, h, lineas, alinear=None, ancla="top"):
    """Cuadro de texto: lineas = [(texto, tamaño, color, negrita[, fuente])]."""
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Pt

    tb = diap.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[ancla]
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Pt(2))
    for i, ln in enumerate(lineas):
        txt, sz, color, negrita = ln[:4]
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if alinear is not None:
            par.alignment = alinear
        r = par.add_run()
        r.text = txt
        r.font.size = Pt(sz)
        r.font.bold = negrita
        r.font.color.rgb = _rgb(color)
        r.font.name = ln[4] if len(ln) > 4 and ln[4] else est.FONT_TEXTO
    return tb


def _titulo_diap(diap, prs, texto, sub=None):
    from pptx.util import Inches

    T = OSCURO
    _tx(diap, Inches(0.55), Inches(0.35), prs.slide_width - Inches(1.1), Inches(0.6), [(texto, 24, T["texto"], True)])
    linea = _caja(diap, Inches(0.55), Inches(0.98), prs.slide_width - Inches(1.1), Inches(0.03), T["oro"], redondeado=False)
    linea.line.fill.background()
    if sub:
        _tx(diap, Inches(0.55), Inches(1.03), prs.slide_width - Inches(1.1), Inches(0.35), [(sub, 11, T["texto2"], False)])


def _pie(diap, prs):
    from pptx.util import Inches

    alto = 0.34
    diap.shapes.add_picture(marca.flujo("auditconsulting_blanco"),
                            prs.slide_width - Inches(0.35) - Inches(marca.ancho_para("auditconsulting_blanco", alto)),
                            prs.slide_height - Inches(0.14) - Inches(alto), height=Inches(alto))


def _tarjeta_grafico(diap, g, x, y, w, h):
    from pptx.util import Inches

    T = OSCURO
    _caja(diap, x, y, w, h, T["card"], T["borde"])
    _tx(diap, x + Inches(0.18), y + Inches(0.1), w - Inches(0.36), Inches(0.62),
        [(g["titulo"], 13, T["texto"], True), (g["sub"], 9, T["texto2"], False)])
    if not g["png"]:
        _tx(diap, x + Inches(0.18), y + Inches(0.9), w - Inches(0.36), Inches(0.4),
            [("Sin datos para graficar en este ejemplo.", 10, T["muted"], False)])
        return
    disp_w, disp_h = w - Inches(0.3), h - Inches(0.85)
    ratio = g["ancho"] / g["alto"]
    iw = min(disp_w, int(disp_h * ratio))
    ih = int(iw / ratio)
    diap.shapes.add_picture(io.BytesIO(g["png"]), x + int((w - iw) / 2), y + Inches(0.78) + int((disp_h - ih) / 2), iw, ih)


def pptx(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    from backend.app.aud.niif.procesadores import libro as L

    T = OSCURO
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    vacia = prs.slide_layouts[6]
    e = reg.get("engagement") or {}
    p = _panel(definicion, reg)

    # 1. Portada: barra superior con los logos, título, cliente y chips.
    s = prs.slides.add_slide(vacia)
    _fondo(s)
    _caja(s, 0, 0, W, Inches(1.05), T["card"], redondeado=False)
    _caja(s, 0, Inches(1.05), W, Inches(0.02), T["borde"], redondeado=False)
    s.shapes.add_picture(marca.flujo("auditconsulting_blanco"), Inches(0.45), Inches(0.2), height=Inches(0.66))
    x_ia = Inches(0.45) + Inches(marca.ancho_para("auditconsulting_blanco", 0.66)) + Inches(0.25)
    s.shapes.add_picture(marca.flujo("audit_ia"), x_ia, Inches(0.2), height=Inches(0.66))
    x_txt = x_ia + Inches(marca.ancho_para("audit_ia", 0.66)) + Inches(0.3)
    _tx(s, x_txt, Inches(0.22), Inches(6), Inches(0.66),
        [("AuditConsulting Auditores", 16, T["texto"], True), ("AUDIT-IA · Papel de trabajo NIIF", 11, T["texto2"], False)], ancla="middle")
    _tx(s, Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(1.5), [(definicion.get("name", ""), 38, T["texto"], True)])
    _tx(s, Inches(0.8), Inches(3.85), W - Inches(1.6), Inches(0.5),
        [(f"{e.get('client', '')} · RUC {e.get('ruc', '')} · {definicion.get('area', '')}", 16, T["texto2"], False)])
    x = Inches(0.8)
    for txt, color in _chips(p, e, version, estado):
        ancho = Inches(0.35 + 0.095 * len(txt))
        chip = _caja(s, x, Inches(4.6), ancho, Inches(0.42), T["card2"], T[color] if color else T["borde"])
        chip.adjustments[0] = 0.5
        _tx(s, x, Inches(4.6), ancho, Inches(0.42), [(txt, 11, T[color] if color else T["texto2"], True)],
            alinear=PP_ALIGN.CENTER, ancla="middle")
        x += ancho + Inches(0.12)
    _tx(s, Inches(0.8), H - Inches(0.6), W - Inches(1.6), Inches(0.35), [(PIE, 9, T["muted"], False)])

    # 2-3. Panel: tarjetas KPI y los 4 gráficos del HTML.
    graf = svg_png.graficos_panel(p, "ejecutivo")
    s = prs.slides.add_slide(vacia)
    _fondo(s)
    _titulo_diap(s, prs, "Panel", f"{e.get('client', '')} · corte {hx._fecha(e.get('cutoff'))}")
    kp = hx.kpis_datos(p)
    margen, sep = Inches(0.55), Inches(0.14)
    wk = int((W - 2 * margen - sep * (len(kp) - 1)) / len(kp))
    for i, k in enumerate(kp):
        xk = margen + i * (wk + sep)
        _caja(s, xk, Inches(1.5), wk, Inches(1.55), T["card"], T["borde"])
        txt, color = _var(k)
        lineas = [(k["rotulo"].upper(), 9, T["texto2"], True), (k["valor"], 18, T[_CLASE[k["clase"]]], True)]
        if k["exacto"]:
            lineas.append((k["exacto"], 10, T["muted"], False, "Consolas"))
        if txt:
            lineas.append((txt, 10, T[color] if color else T["texto2"], True))
        _tx(s, xk + Inches(0.12), Inches(1.57), wk - Inches(0.24), Inches(1.45), lineas)
    wg = int((W - 2 * margen - sep) / 2)
    for k, g in enumerate(graf[:2]):
        _tarjeta_grafico(s, g, margen + k * (wg + sep), Inches(3.2), wg, Inches(3.75))
    _pie(s, prs)
    s = prs.slides.add_slide(vacia)
    _fondo(s)
    _titulo_diap(s, prs, "Panel · distribución y severidad")
    for k, g in enumerate(graf[2:]):
        _tarjeta_grafico(s, g, margen + k * (wg + sep), Inches(1.45), wg, Inches(5.45))
    _pie(s, prs)

    # 4+. Cédulas de lectura ejecutiva en tablas con el estilo del HTML (el detalle completo, en el Excel).
    for h in L.cedulas(definicion, reg, eventos, version, estado):
        if not L._en_ppt(h["name"]):
            continue
        filas = L._filas(h)
        recorte = len(filas) > L._MAX_FILAS_PPT
        if recorte:
            filas = filas[:L._MAX_FILAS_PPT - 1] + ([f for f in filas if f[1]] or [])
        s = prs.slides.add_slide(vacia)
        _fondo(s)
        _titulo_diap(s, prs, h["label"], "Primeras filas; el detalle completo está en el Excel." if recorte else None)
        _tabla_ppt(s, prs, h, filas, L)
        _pie(s, prs)
    salida = io.BytesIO()
    prs.save(salida)
    return salida.getvalue()


def _borde_celda(celda, lado, color, ancho_emu=19050):
    """Filete de una celda de tabla (lnT/lnB) — python-pptx no lo expone."""
    from pptx.oxml.ns import qn
    from lxml import etree

    tcPr = celda._tc.get_or_add_tcPr()
    tag = {"top": "a:lnT", "bottom": "a:lnB", "left": "a:lnL", "right": "a:lnR"}[lado]
    for viejo in tcPr.findall(qn(tag)):
        tcPr.remove(viejo)
    ln = etree.SubElement(tcPr, qn(tag), w=str(ancho_emu if color else 0), cap="flat", cmpd="sng", algn="ctr")
    if color:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    else:
        etree.SubElement(ln, qn("a:noFill"))   # sin línea vertical, como las tablas del HTML
    # Orden del esquema: lnL, lnR, lnT, lnB, …, relleno (solidFill) al final.
    orden = [qn(t) for t in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr", "a:cell3D")]
    hijos = list(tcPr)
    for el in hijos:
        tcPr.remove(el)
    for el in sorted(hijos, key=lambda x: orden.index(x.tag) if x.tag in orden else len(orden)):
        tcPr.append(el)


def _tabla_ppt(diap, prs, h, filas, L):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    T = OSCURO
    n = len(h["cols"])
    sz = 10 if n <= 7 else 9 if n <= 10 else 8
    alto_fila = Inches(0.3)
    forma = diap.shapes.add_table(1 + len(filas), n, Inches(0.55), Inches(1.5), prs.slide_width - Inches(1.1),
                                  alto_fila * (1 + len(filas)))
    t = forma.table
    # Ancho de columna según su contenido (rótulo y celdas), como la tabla del HTML.
    largos = []
    for j, (nombre, fmt) in enumerate(h["cols"]):
        cont = [len(_html.unescape(L._celda(f[j], fmt))) for f, _ in filas if j < len(f)]
        largos.append(min(42, max(6, len(nombre) + 2, *(cont or [0]))))
    total_w = prs.slide_width - Inches(1.1)
    for j, lg in enumerate(largos):
        t.columns[j].width = int(total_w * lg / sum(largos))
    t.first_row = True
    t.horz_banding = False
    # Sin estilo de tabla del tema (evita bordes blancos): el relleno y los filetes van por celda.
    tblPr = t._tbl.tblPr
    for el in list(tblPr):
        if el.tag.endswith("tableStyleId"):
            tblPr.remove(el)

    def pinta(celda, txt, color, relleno, negrita=False, derecha=False, fuente=None):
        celda.fill.solid()
        celda.fill.fore_color.rgb = _rgb(relleno)
        celda.margin_left = celda.margin_right = Pt(5)
        celda.margin_top = celda.margin_bottom = Pt(3)
        tf = celda.text_frame
        tf.word_wrap = True
        par = tf.paragraphs[0]
        par.alignment = PP_ALIGN.RIGHT if derecha else PP_ALIGN.LEFT
        r = par.add_run()
        r.text = txt
        r.font.size = Pt(sz)
        r.font.bold = negrita
        r.font.color.rgb = _rgb(color)
        r.font.name = fuente or est.FONT_TEXTO

    for j, (nombre, _) in enumerate(h["cols"]):
        c = t.cell(0, j)
        pinta(c, nombre, T["texto"], T["card2"], True)
        _borde_celda(c, "bottom", T["oro"], 25400)
    for i, (fila, total) in enumerate(filas, start=1):
        for j, ((_, fmt), v) in enumerate(zip(h["cols"], fila)):
            c = t.cell(i, j)
            num = fmt in ("n", "p", "i", "g")
            relleno = T["card2"] if total else (T["zebra"] if i % 2 == 0 else T["card"])
            pinta(c, _html.unescape(L._celda(v, fmt)), T["texto"], relleno, total, num, "Consolas" if num else None)
            if total:
                _borde_celda(c, "top", T["oro"], 25400)
            _borde_celda(c, "bottom", T["borde"], 9525)
    for fila in t.rows:
        for c in fila.cells:
            _borde_celda(c, "left", None)
            _borde_celda(c, "right", None)
