"""Portada ``00_Inicio`` del Excel = el panel del HTML (tema «Ejecutivo»).

El dueño pidió (2026-09-25) que el Excel muestre los MISMOS gráficos del HTML, con
sus colores de fondo y sus botones, y no otros:

- Fondo azul marino (#071B2F), barra superior con los logotipos, título, chips
  (riesgo, corte, versión, estado, marco) y las 5 tarjetas KPI del HTML con sus
  colores (oro, azul, verde, ámbar y el color del riesgo).
- Los 4 gráficos del panel del HTML, nativos de Excel: «Composición del resultado»
  (dona), «Registrado vs recalculado», la distribución de la población y
  «Problemas por severidad» (barras + línea, colores por barra del HTML).
- Los tableros adicionales del ``PANEL`` (``"tableros"``: columnas agrupadas anterior vs actual),
  debajo, en la misma rejilla.
- Botones con el estilo del HTML (fondo #0E2C50, borde #1B3A60) agrupados por sección.

Todas las cifras son fórmulas (regla «Sin cifras calculadas pegadas»): las
tarjetas remiten a la cédula que calcula cada importe y los datos de los gráficos
(hoja oculta ``00_Datos_graficos``) son SUMIFS/COUNTIFS sobre las cédulas con la
misma agrupación que el HTML (``graficos.serie_spec``, ``graficos.severidad``).
"""
from __future__ import annotations

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend.app.aud.niif.procesadores import estilo_ejecutivo as est
from backend.app.aud.niif.procesadores import graficos
from backend.app.aud.niif.procesadores import html_ejecutivo as hx
from backend.app.aud.niif.procesadores import marca

# Colores del tema «Ejecutivo» del HTML (una sola fuente: html_ejecutivo.TEMAS).
_T = hx.TEMAS["ejecutivo"][1]
H = lambda k: _T[k].lstrip("#").upper()  # noqa: E731
BG, CARD, CARD2, BORDE, TEXTO, TEXTO2, MUTED = (H(k) for k in ("bg", "card", "card2", "borde", "texto", "texto2", "muted"))
ORO, ORO_TXT = H("oro"), H("oro-txt")
KPI_COLOR = {"principal": H("oro-txt"), "poblacion": H("k-azul"), "recalculado": H("k-verde"), "registrado": H("k-ambar"),
             "alto": H("alta"), "medio": H("media"), "bajo": H("baja")}
SERIE = {k: v.lstrip("#").upper() for k, v in hx._SERIES_OSCURO.items()}
C_REG, C_REC, C_SERIE, C_LINEA = H("c-registrado"), H("c-recalculado"), H("c-serie"), H("c-linea")
C_SEV = [H("c-alta"), H("c-media"), H("c-baja"), H("c-informativa")]

COLS = "BCDEF"
ANCHO_COL = 27                 # caracteres ≈ 194 px
PX_COL = 194
PX_PANEL = PX_COL * len(COLS)  # 970 px
GAP = 12
R0 = 12                        # fila del rótulo «INDICADORES CLAVE»; las cifras van en R0 + 2
FMT_USD = '"USD "#,##0.00;"USD "-#,##0.00;"USD "0.00'
FMT_CORTO = '[>=1000000]#,##0.0,," M";[>=1000]#,##0.0," mil";#,##0.00'

_fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
_lado = lambda c, s="thin": Side(style=s, color=c)  # noqa: E731


def _f(sz, color, bold=False, italic=False):
    return Font(name=est.FONT_TEXTO, size=sz, bold=bold, italic=italic, color=color)


# --- Texto de los gráficos -------------------------------------------------------------------

def _cp(color, sz, b=False):
    return CharacterProperties(sz=int(sz * 100), b=b, solidFill=color, latin=DFont(typeface=est.FONT_TEXTO))


def _txpr(color, sz, b=False):
    cp = _cp(color, sz, b)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(texto, sub=None):
    """Título de tarjeta del HTML: rótulo en negrita y, debajo, el subtítulo en gris azulado."""
    p = [Paragraph(pPr=ParagraphProperties(algn="l", defRPr=_cp(TEXTO, 11, True)),
                   r=[RegularTextRun(rPr=_cp(TEXTO, 11, True), t=texto)])]
    if sub:
        p.append(Paragraph(pPr=ParagraphProperties(algn="l", defRPr=_cp(TEXTO2, 8)),
                           r=[RegularTextRun(rPr=_cp(TEXTO2, 8), t=sub)]))
    t = Title(tx=Text(rich=RichText(p=p)), overlay=False)
    t.layout = Layout(manualLayout=ManualLayout(x=0.025, y=0.025, xMode="edge", yMode="edge"))
    return t


def _superficie(ch):
    """Tarjeta del HTML: fondo #0A2342 con borde #1B3A60; área de trazado transparente."""
    ch.graphical_properties = GraphicalProperties(solidFill=CARD, ln=LineProperties(solidFill=BORDE, w=9525))
    ch.roundedCorners = True
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))


def _puntos(serie, colores, borde=None):
    for i, c in enumerate(colores):
        gp = GraphicalProperties(solidFill=c, ln=LineProperties(solidFill=borde, w=19050) if borde else LineProperties(noFill=True))
        serie.dPt.append(DataPoint(idx=i, spPr=gp))


def _refs(wd, bloque):
    fila0, n = bloque["fila"], len(bloque["items"])
    cats = Reference(wd, min_col=1, min_row=fila0 + 1, max_row=fila0 + n)
    vals = Reference(wd, min_col=2, min_row=fila0, max_row=fila0 + n)
    return cats, vals


def grafico_dona(wd, bloque, titulo, sub):
    cats, vals = _refs(wd, bloque)
    ch = DoughnutChart(holeSize=58, firstSliceAng=0)
    ch.varyColors = True
    ch.add_data(vals, titles_from_data=True)
    s = ch.series[0]
    s.cat = AxDataSource(strRef=StrRef(f=str(cats)))
    roles = [("otros" if e.startswith("Otros") else f"s{min(i, 7) + 1}") for i, (e, _) in enumerate(bloque["items"])]
    _puntos(s, [SERIE[r] for r in roles], borde=CARD)   # 2 px de superficie entre porciones, como el HTML
    total = sum(abs(v) for _, v in bloque["items"]) or 1.0
    dl = DataLabelList(showPercent=True, showVal=False, showCatName=False, showSerName=False, showLegendKey=False,
                       numFmt="0.0%", txPr=_txpr("FFFFFF", 9, True))
    # Porciones chicas sin rótulo (en el HTML el % está en la leyenda): no se amontonan.
    dl.dLbl = [DataLabel(idx=i, showPercent=False, showVal=False, showCatName=False, showSerName=False, showLegendKey=False)
               for i, (_, v) in enumerate(bloque["items"]) if abs(v) / total < 0.06]
    s.dLbls = dl
    ch.legend = Legend(legendPos="r", txPr=_txpr(TEXTO, 9))
    ch.title = _titulo(titulo, sub)
    _superficie(ch)
    return ch


def grafico_barras_linea(wd, bloque, titulo, sub, colores, enteros=False):
    """Barras + línea (la vista por defecto del HTML): una barra por categoría con su color
    y la línea dorada que une las cimas; cifra corta encima de cada barra."""
    cats, vals = _refs(wd, bloque)
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 90
    ch.varyColors = True
    ch.add_data(vals, titles_from_data=True)
    s = ch.series[0]
    s.cat = AxDataSource(strRef=StrRef(f=str(cats)))
    s.invertIfNegative = False
    _puntos(s, colores[:len(bloque["items"])] if len(colores) > 1 else colores * len(bloque["items"]))
    s.dLbls = DataLabelList(showVal=True, showPercent=False, showCatName=False, showSerName=False, showLegendKey=False,
                            numFmt="#,##0" if enteros else FMT_CORTO, dLblPos="outEnd", txPr=_txpr(TEXTO, 9, True))
    ln = LineChart()
    ln.add_data(vals, titles_from_data=True)
    sl = ln.series[0]
    sl.cat = AxDataSource(strRef=StrRef(f=str(cats)))
    sl.smooth = False
    sl.graphicalProperties.line.solidFill = C_LINEA
    sl.graphicalProperties.line.width = 22225
    sl.marker = Marker(symbol="circle", size=6, spPr=GraphicalProperties(solidFill=C_LINEA, ln=LineProperties(solidFill=C_LINEA)))
    for c in (ch, ln):
        c.y_axis.delete = True                 # el HTML no muestra eje de valores
        c.y_axis.majorGridlines = None
        c.x_axis.delete = False
        c.x_axis.tickLblPos = "low"
        c.x_axis.txPr = _txpr(TEXTO2, 9)
        c.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=BORDE))
    if all(v >= 0 for _, v in bloque["items"]):
        ch.y_axis.scaling.min = 0              # un eje que no empieza en 0 exagera las diferencias
        ln.y_axis.scaling.min = 0
    ch.legend = None
    ch += ln
    ch.title = _titulo(titulo, sub)
    _superficie(ch)
    return ch


FMT_TABLERO = {"veces": "#,##0.00", "días": "#,##0", "%": "#,##0.00",
               "USD": '[>=1000000]#,##0.00,," M";[>=1000]#,##0," mil";#,##0'}   # como ``graficos_svg.corto``


def grafico_agrupadas(wd, bloque, titulo, sub, colores, fmt):
    """Tablero del HTML (``graficos_svg.agrupadas``): columnas agrupadas, una serie por columna del
    bloque (anterior y actual), cifra encima de cada barra y la leyenda debajo."""
    fila0, n, k = bloque["fila"], len(bloque["items"]), bloque["k"]
    cats = Reference(wd, min_col=1, min_row=fila0 + 1, max_row=fila0 + n)
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.gapWidth = 80
    ch.overlap = 0
    for j in range(k):
        ch.add_data(Reference(wd, min_col=2 + j, min_row=fila0, max_row=fila0 + n), titles_from_data=True)
    for s, c in zip(ch.series, colores):
        s.cat = AxDataSource(strRef=StrRef(f=str(cats)))
        s.graphicalProperties = GraphicalProperties(solidFill=c, ln=LineProperties(noFill=True))
        s.invertIfNegative = False
        s.dLbls = DataLabelList(showVal=True, showPercent=False, showCatName=False, showSerName=False, showLegendKey=False,
                                numFmt=fmt, dLblPos="outEnd", txPr=_txpr(TEXTO, 8, True))
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.txPr = _txpr(TEXTO2, 9)
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=BORDE))
    if all(v >= 0 for _, vs in bloque["items"] for v in vs if v is not None):
        ch.y_axis.scaling.min = 0
    ch.legend = Legend(legendPos="b", txPr=_txpr(TEXTO2, 9))
    ch.title = _titulo(titulo, sub)
    _superficie(ch)
    return ch


def tableros_formulas(t, hojas, titulos):
    """Filas del bloque de datos de un tablero: [(categoría, [«=fórmula» por serie], [valor por serie])],
    cada fórmula a la celda de la cédula que ya calcula el índice o el importe. None si la cédula no está."""
    from openpyxl.utils import get_column_letter

    from backend.app.aud.niif.procesadores import libro as L

    i = next((k for k, h in enumerate(hojas) if h.get("name") == t["hoja"]), None)
    if i is None:
        return None
    nombres = [c[0] for c in hojas[i].get("cols") or []]
    if any(c not in nombres for c in t["columnas"]):
        return None
    letras = [get_column_letter(nombres.index(c) + 1) for c in t["columnas"]]
    q = L._q(titulos[i])
    return [(cat, [f"={q}${le}${5 + fila}" for le in letras], [vs[k] for _, vs in t["series"]])
            for k, (cat, fila) in enumerate(zip(t["categorias"], t["filas"]))]


# --- Datos de los gráficos (fórmulas) ---------------------------------------------------------

class Datos:
    """Hoja oculta ``00_Datos_graficos``: un bloque por gráfico (rótulo | fórmula)."""

    def __init__(self, wd):
        self.wd, self.r = wd, 3
        wd.column_dimensions["A"].width = 46
        wd.column_dimensions["B"].width = 20
        wd.cell(row=1, column=1, value="Datos de los gráficos del panel (fórmulas a las cédulas)").font = _f(9, "4B5563", italic=True)

    def bloque(self, titulo, items, fmt=None):
        """items: [(rótulo, «=fórmula», valor esperado)]."""
        wd, fila = self.wd, self.r
        wd.cell(row=fila, column=1, value=titulo).font = _f(9, "4B5563", True)
        wd.cell(row=fila, column=2, value="Importe (USD)").font = _f(9, "4B5563", True)
        for k, (rot, f, _) in enumerate(items, start=1):
            wd.cell(row=fila + k, column=1, value=rot).font = _f(9, "374151")
            c = wd.cell(row=fila + k, column=2, value=f)
            c.font = _f(9, "374151")
            c.number_format = fmt or est.FMT["n"]
        self.r = fila + len(items) + 2
        return {"fila": fila, "items": [(rot, v) for rot, _, v in items]}

    def bloque_series(self, titulo, nombres, filas, fmt=None):
        """Bloque de varias series (tableros): rótulo | una columna de fórmulas por serie.
        filas: [(rótulo, [«=fórmula»], [valor esperado])]."""
        wd, fila = self.wd, self.r
        wd.cell(row=fila, column=1, value=titulo).font = _f(9, "4B5563", True)
        for j, nombre in enumerate(nombres):
            wd.cell(row=fila, column=2 + j, value=nombre).font = _f(9, "4B5563", True)
            if not wd.column_dimensions[chr(66 + j)].width:
                wd.column_dimensions[chr(66 + j)].width = 20
        for k, (rot, fs, _) in enumerate(filas, start=1):
            wd.cell(row=fila + k, column=1, value=rot).font = _f(9, "374151")
            for j, f in enumerate(fs):
                c = wd.cell(row=fila + k, column=2 + j, value=f)
                c.font = _f(9, "374151")
                c.number_format = fmt or est.FMT["n"]
        self.r = fila + len(filas) + 2
        return {"fila": fila, "k": len(nombres), "items": [(rot, vs) for rot, _, vs in filas]}

    def celda(self, fila_rel, bloque):
        return f"'{self.wd.title}'!$B${bloque['fila'] + fila_rel}"


def _grupos_crudos(h, spec):
    """rótulo normalizado (como ``serie_spec``) → conjunto de rótulos tal como están en la cédula."""
    je = graficos._cols_idx(h, spec.get("etiqueta"))
    jv = graficos._cols_idx(h, spec.get("valor"))
    out: dict[str, set] = {}
    for f in graficos._filas(h, spec):
        if graficos._num(f[jv] if jv < len(f) else None) is None:
            continue
        crudo = f[je] if je < len(f) else None
        crudo = crudo.get("v") if isinstance(crudo, dict) else crudo
        rot = " ".join(str(crudo if crudo not in (None, "") else "(sin rótulo)").split())
        out.setdefault(rot, set()).add(crudo if crudo not in (None, "") else None)
    return out


def serie_formulas(spec, run, hojas, titulos, absoluto):
    """Los mismos ítems que el HTML (``graficos.serie_spec``) con una fórmula por ítem:
    [(rótulo, «=fórmula», valor)]. None si alguno no se puede expresar con fórmulas."""
    from backend.app.aud.niif.procesadores import libro as L

    mapa = {h["name"]: h for h in hojas}
    items = graficos.serie_spec(spec, mapa, absoluto=absoluto, run=run)
    if not items:
        return None
    envolver = (lambda f: f"=ABS({f})") if absoluto else (lambda f: "=" + f)
    base = lambda rot: rot[4:] if absoluto and rot.startswith("(−) ") else rot  # noqa: E731
    if spec.get("totales"):
        claves = {r: c for r, c in spec["totales"]}
        out = []
        for rot, v in items:
            clave = claves.get(base(rot))
            val = graficos._num((run.get("totals") or {}).get(clave)) if clave else None
            f = L._formula_por_valor(hojas, titulos, val, (run.get("labels") or {}).get(clave)) if val is not None else None
            if not f:
                return None
            out.append((rot, envolver(f), v))
        return out
    i = next((k for k, h in enumerate(hojas) if h["name"] == spec.get("hoja")), None)
    if i is None:
        return None
    rng_e = L._rango_col(hojas, titulos, i, spec.get("etiqueta"))
    rng_v = L._rango_col(hojas, titulos, i, spec.get("valor"))
    cr = L._criterios(spec, hojas, titulos, i)
    if rng_e is None or rng_v is None or cr is None:
        return None
    grupos = _grupos_crudos(hojas[i], spec)
    suma = lambda g: "+".join(L._suma_si(rng_v, cr[0], cr[1], [(rng_e, L._crit(c))])  # noqa: E731
                              for c in sorted(grupos[g], key=lambda x: str(x)))
    out, nombrados = [], []
    for rot, v in items:
        if rot.startswith("Otros ("):
            continue
        g = base(rot)
        if g not in grupos:
            return None
        nombrados.append(g)
        out.append((rot, envolver(suma(g)), v))
    for rot, v in items:
        if rot.startswith("Otros ("):
            # Resto = total − los nombrados (en la dona, las partidas que restan nunca van a «Otros»).
            f = f"{L._suma_si(rng_v, *cr)}-(" + "+".join(suma(g) for g in nombrados) + ")"
            out.append((rot, envolver(f), v))
    return out


def severidad_formulas(hojas, titulos, base_ref, run, base_val):
    """«Problemas por severidad» del HTML con COUNTIFS sobre la hoja de problemas:
    ≥ 5 % de la población alta, ≥ 1 % media, el resto con importe baja; sin importe, informativa."""
    from backend.app.aud.niif.procesadores import libro as L
    from backend.app.aud.niif.procesadores import problemas

    hp = problemas.hoja_problemas(hojas)
    if hp is None:
        return None
    ip, imp = hp
    n = max(len(hojas[ip].get("rows") or []), 1)
    q = L._q(titulos[ip])
    A, C = f"{q}$A$5:$A${4 + n}", f"{q}${imp}$5:${imp}${4 + n}"
    b = base_ref or "0"

    def cuenta(pct):
        u = f"MAX(0.005,{pct}*ABS({b}))"
        return f'IF(ABS({b})=0,0,COUNTIFS({C},">="&{u})+COUNTIFS({C},"<="&-{u}))'

    con_importe = f'(COUNTIFS({C},">=0.005")+COUNTIFS({C},"<=-0.005"))'
    alta, al_menos_media = cuenta(0.05), cuenta(0.01)
    esperado = graficos.severidad(run, base_val)
    return [
        ("Alta", f"={alta}", float(esperado["Alta"])),
        ("Media", f"={al_menos_media}-({alta})", float(esperado["Media"])),
        ("Baja", f"={con_importe}-({al_menos_media})", float(esperado["Baja"])),
        ("Informativa", f"=COUNTA({A})-{con_importe}", float(esperado["Informativa"])),
    ]


# --- Portada -----------------------------------------------------------------------------------

def _pinta_fondo(ws, hasta_fila, hasta_col=26):
    """Fondo azul marino del HTML en toda la vista; respeta lo que ya tiene relleno
    (barra, tarjetas, chips, botones)."""
    f = _fill(BG)
    for r in range(1, hasta_fila + 1):
        for c in range(1, hasta_col + 1):
            cel = ws.cell(row=r, column=c)
            if cel.fill is None or cel.fill.fill_type is None:
                cel.fill = f


def _caja(ws, fila1, fila2, col1, col2, relleno, borde=BORDE):
    """Marco de tarjeta: relleno y borde exterior."""
    for r in range(fila1, fila2 + 1):
        for c in range(col1, col2 + 1):
            cel = ws.cell(row=r, column=c)
            cel.fill = _fill(relleno)
            cel.border = Border(left=_lado(borde) if c == col1 else None, right=_lado(borde) if c == col2 else None,
                                top=_lado(borde) if r == fila1 else None, bottom=_lado(borde) if r == fila2 else None)


def boton_html(cel, texto, destino, primario=False):
    """Botón del HTML: fondo #0E2C50, borde #1B3A60 y texto claro; el primario, dorado."""
    cel.value = texto
    cel.hyperlink = destino
    cel.font = Font(name=est.FONT_TITULO, size=10, bold=True, color=BG if primario else TEXTO, underline=None)
    cel.fill = _fill(ORO if primario else CARD2)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lado = _lado(ORO if primario else BORDE)
    cel.border = Border(left=lado, right=lado, top=lado, bottom=Side(style="medium", color=ORO if primario else BORDE))


def portada(ws, wd, definicion, reg, hojas, titulos, estado, version, grupos_nav, secciones):
    """Arma ``00_Inicio`` como el panel del HTML. Devuelve la última fila usada."""
    from backend.app.aud.niif.procesadores import libro as L

    e = reg.get("engagement") or {}
    run = reg.get("run") or {}
    mod = graficos.modulo(definicion)
    spec = getattr(mod, "PANEL", None) or {}
    p = graficos.panel(mod, run, run.get("hojas") or [])
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for c in COLS:
        ws.column_dimensions[c].width = ANCHO_COL
    ws.column_dimensions["G"].width = 3

    # Barra superior (como la del HTML): logotipos y la marca.
    for r, alto in ((1, 8), (2, 27), (3, 27), (4, 8)):
        ws.row_dimensions[r].height = alto
    _caja(ws, 2, 3, 2, 6, CARD)
    ws.merge_cells("D2:F2")
    ws.merge_cells("D3:F3")
    ws["D2"].value = "AuditConsulting Auditores Cía. Ltda."
    ws["D2"].font = Font(name=est.FONT_TITULO, size=13, bold=True, color=TEXTO)
    ws["D2"].alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    ws["D3"].value = "AUDIT-IA · Papel de trabajo NIIF"
    ws["D3"].font = _f(10, TEXTO2)
    ws["D3"].alignment = Alignment(horizontal="left", vertical="top", indent=1)
    alto_logo = 54
    firma = marca.imagen_excel("auditconsulting_blanco", alto_logo)
    L._ancla(ws, firma, 1, 1, 12, 9)
    ia = marca.imagen_excel("audit_ia", alto_logo)
    L._ancla(ws, ia, 2, 1, 16, 9)

    # Encabezado: título, cliente y chips.
    ws.row_dimensions[6].height = 30
    ws.merge_cells("B6:F6")
    ws["B6"].value = L._seguro(definicion.get("name", ""))
    ws["B6"].font = Font(name=est.FONT_TITULO, size=20, bold=True, color=TEXTO)
    ws["B6"].alignment = Alignment(vertical="center")
    ws.merge_cells("B7:F7")
    ws["B7"].value = L._seguro(f"{e.get('client', '')} · RUC {e.get('ruc', '')} · {definicion.get('area', '')}")
    ws["B7"].font = _f(10, TEXTO2)
    riesgo = {"alto": "▲ Riesgo alto", "medio": "▲ Riesgo medio", "bajo": "▲ Riesgo bajo"}[p["riesgo"]]
    chips = [(riesgo, KPI_COLOR[p["riesgo"]]), (f"Corte {hx._fecha(e.get('cutoff')) or 'pendiente'}", TEXTO2), (f"Versión {version}", TEXTO2),
             (est.estado_es(estado), TEXTO2), (str(e.get("framework") or "Marco pendiente"), TEXTO2)]
    ws.row_dimensions[9].height = 20
    for col, (txt, color) in zip(COLS, chips):
        c = ws[f"{col}9"]
        c.value = L._seguro(txt)
        c.font = _f(9, color, True)
        c.fill = _fill(CARD2)
        c.alignment = Alignment(horizontal="center", vertical="center")
        lado = _lado(color if color != TEXTO2 else BORDE)
        c.border = Border(left=lado, right=lado, top=lado, bottom=lado)
    # Datos del encargo (compactos, como la carátula).
    datos = [("Marco contable", e.get("framework")), ("Fecha de corte", e.get("cutoff")),
             ("Preparó", e.get("preparer")), ("Revisó", e.get("reviewer"))]
    for k, (etq, val) in enumerate(datos):
        c = ws.cell(row=10, column=2 + k, value=L._seguro(f"{etq}: {val if val not in (None, '') else '—'}"))
        c.font = _f(8.5, MUTED)

    # Tarjetas KPI (las 5 del HTML).
    datos_g = Datos(wd)
    r0 = R0
    ws[f"B{r0}"].value = "INDICADORES CLAVE"
    ws[f"B{r0}"].font = Font(name=est.FONT_TITULO, size=9, bold=True, color=ORO_TXT)
    kpis = L._kpis_panel(definicion, reg, hojas, titulos)
    por = {k["clave"]: k for k in kpis}
    f_pob = L._formula_spec(spec.get("poblacion"), run, hojas, titulos)
    base = datos_g.bloque("Base de la severidad: población", [("Población", "=" + f_pob[0], f_pob[1])]) if f_pob else None
    base_ref = datos_g.celda(1, base) if base else None
    sev_items = severidad_formulas(hojas, titulos, base_ref, run, p["poblacion"]["valor"])
    sev = datos_g.bloque("Problemas por severidad", sev_items, fmt="0") if sev_items else None
    tarjetas = _tarjetas(p, por, f_pob, spec, run, hojas, titulos, datos_g, sev, base_ref)
    celdas = {}
    ws.row_dimensions[r0 + 1].height = 30
    ws.row_dimensions[r0 + 2].height = 32
    ws.row_dimensions[r0 + 3].height = 20
    for col, t in zip(COLS, tarjetas):
        _caja(ws, r0 + 1, r0 + 3, ord(col) - 64, ord(col) - 64, CARD)
        c = ws[f"{col}{r0 + 1}"]
        c.value = t["rotulo"].upper()
        c.font = Font(name=est.FONT_TITULO, size=8.5, bold=True, color=TEXTO2)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        v = ws[f"{col}{r0 + 2}"]
        v.value = t["valor"]
        v.number_format = t["fmt"]
        v.font = Font(name=est.FONT_TITULO, size=17, bold=True, color=t["color"])
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True)
        s = ws[f"{col}{r0 + 3}"]
        s.value = t.get("sub")
        s.number_format = t.get("fmt_sub") or "General"
        s.font = _f(9, t.get("color_sub") or TEXTO2, True)
        s.alignment = Alignment(horizontal="left", vertical="top", indent=1, shrink_to_fit=True)
        celdas[t["clave"]] = f"${col}${r0 + 2}"

    # Los 4 gráficos del panel del HTML, en rejilla 2×2.
    q = L._q(ws.title)
    graf = []
    comp = serie_formulas(spec.get("composicion"), run, hojas, titulos, True) if spec.get("composicion") else None
    if comp:
        b = datos_g.bloque(p["composicion"]["rotulo"], comp)
        graf.append(grafico_dona(wd, b, "Composición del resultado", p["composicion"]["rotulo"]))
    f_reg = f"={q}{celdas['registrado']}" if "registrado" in celdas else None
    f_rec = f"={q}{celdas['recalculado']}" if "recalculado" in celdas else None
    if f_reg and f_rec:
        cmp_ = p["comparativo"]
        b = datos_g.bloque(cmp_["rotulo"], [(p["registrado"]["rotulo"], f_reg, p["registrado"]["valor"] or 0.0),
                                            (p["recalculado"]["rotulo"], f_rec, p["recalculado"]["valor"] or 0.0)])
        graf.append(grafico_barras_linea(wd, b, cmp_["rotulo"], cmp_.get("sub") or graficos.TEXTOS["comparativo_sub"], [C_REG, C_REC]))
    dist = serie_formulas(spec.get("distribucion"), run, hojas, titulos, False) if spec.get("distribucion") else None
    if dist:
        b = datos_g.bloque(p["distribucion"]["rotulo"], dist)
        graf.append(grafico_barras_linea(wd, b, p["distribucion"]["rotulo"], "Distribución de la población (USD).", [C_SERIE]))
    if sev:
        graf.append(grafico_barras_linea(wd, sev, "Problemas por severidad", graficos.REGLA_SEVERIDAD, C_SEV, enteros=True))

    fila = r0 + 5
    filas_graf = 19                     # 19 filas de 15 pt ≈ 380 px
    for k, ch in enumerate(graf):
        fila_g = fila + (k // 2) * (filas_graf + 1)
        _ancla_grafico(ws, ch, izquierda=(k % 2 == 0), fila=fila_g - 1, filas=filas_graf)
    fila += ((len(graf) + 1) // 2) * (filas_graf + 1) + 1
    # Impresión: la segunda fila de gráficos y la navegación empiezan página (no se corta un gráfico).
    from openpyxl.worksheet.pagebreak import Break

    if len(graf) > 2:
        ws.row_breaks.append(Break(id=r0 + 5 + filas_graf))

    # Tableros adicionales del PANEL (p. ej. índices por grupo y analítico de la planificación),
    # como en el HTML: debajo de los 4 gráficos, en la misma rejilla, con sus datos por fórmula.
    tabs = []
    for t in p.get("tableros") or []:
        filas_t = tableros_formulas(t, hojas, titulos)
        if not filas_t:
            continue
        b = datos_g.bloque_series(t["rotulo"], [n for n, _ in t["series"]], filas_t, FMT_TABLERO.get(t.get("unidad"), est.FMT["n"]))
        tabs.append(grafico_agrupadas(wd, b, t["rotulo"], t.get("sub"), [SERIE["s1"], SERIE["s3"], SERIE["s2"], SERIE["s4"]],
                                      FMT_TABLERO.get(t.get("unidad"), "#,##0.00")))
    if tabs:
        ws.row_breaks.append(Break(id=fila - 1))
        ws[f"B{fila}"].value = "TABLEROS DEL ANÁLISIS"
        ws[f"B{fila}"].font = Font(name=est.FONT_TITULO, size=9, bold=True, color=ORO_TXT)
        fila += 2
        for k, ch in enumerate(tabs):
            fila_g = fila + (k // 2) * (filas_graf + 1)
            _ancla_grafico(ws, ch, izquierda=(k % 2 == 0), fila=fila_g - 1, filas=filas_graf)
            if k % 4 == 3 and k + 1 < len(tabs):   # dos filas de tableros por página impresa
                ws.row_breaks.append(Break(id=fila_g + filas_graf))
        fila += ((len(tabs) + 1) // 2) * (filas_graf + 1) + 1
        ws.row_breaks.append(Break(id=fila - 1))

    # Navegación por sección con los botones del HTML.
    ws[f"B{fila}"].value = "NAVEGAR POR SECCIÓN"
    ws[f"B{fila}"].font = Font(name=est.FONT_TITULO, size=9, bold=True, color=ORO_TXT)
    fila += 1
    for sec, items in grupos_nav.items():
        if not items:
            continue
        nombre, color, _, desc = secciones[sec]
        ws.merge_cells(f"B{fila}:F{fila}")
        banda = ws[f"B{fila}"]
        banda.value = f"{nombre}  ·  {desc}"
        banda.font = Font(name=est.FONT_TITULO, size=11, bold=True, color=TEXTO)
        banda.alignment = Alignment(horizontal="left", vertical="center")
        for c in COLS:   # como la pestaña activa del HTML: filete dorado debajo
            ws[f"{c}{fila}"].border = Border(bottom=Side(style="medium", color=ORO))
        ws.row_dimensions[fila].height = 24
        fila += 1
        for i, (etq, t, _) in enumerate(items):
            row = fila + i // 5
            boton_html(ws[f"{COLS[i % 5]}{row}"], etq, L._ref(t), primario=(sec == 0 and i == 0))
            ws.row_dimensions[row].height = 32
        fila += (len(items) - 1) // 5 + 2
    _pinta_fondo(ws, fila + 20)
    return fila


def _ancla_grafico(ws, ch, izquierda, fila, filas):
    """Ancla de dos celdas: el gráfico ocupa media rejilla B:F (hasta la mitad de la columna D)
    y se estira con las columnas, así se ve igual en Excel, LibreOffice o un visor."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.utils.units import pixels_to_EMU

    medio = PX_COL // 2
    if izquierda:
        desde, hasta = AnchorMarker(col=1, colOff=0, row=fila), AnchorMarker(col=3, colOff=pixels_to_EMU(medio - GAP // 2), row=fila + filas)
    else:
        desde, hasta = AnchorMarker(col=3, colOff=pixels_to_EMU(medio + GAP // 2), row=fila), AnchorMarker(col=6, colOff=0, row=fila + filas)
    ch.anchor = TwoCellAnchor(editAs="oneCell", _from=desde, to=hasta)
    ws.add_chart(ch)


def _tarjetas(p, por, f_pob, spec, run, hojas, titulos, datos_g, sev, base_ref=None):
    """Las 5 tarjetas del HTML con su color, su cifra (fórmula) y la línea de variación (fórmula)."""
    txt = p.get("textos") or graficos.textos(spec)
    from backend.app.aud.niif.procesadores import libro as L

    out = []
    pr = por.get("principal")
    if pr:
        out.append({"clave": "principal", "rotulo": pr["rotulo"], "valor": "=" + pr["f"], "fmt": FMT_USD,
                    "color": KPI_COLOR["principal"]})
    po = p["poblacion"]
    if f_pob:
        i = next((k for k, h in enumerate(hojas) if h["name"] == (spec.get("poblacion") or {}).get("hoja")), None)
        n_f = None
        if i is not None:
            rng = L._rango_col(hojas, titulos, i, spec["poblacion"].get("col"))
            cr = L._criterios(spec["poblacion"], hojas, titulos, i)
            if rng and cr is not None:
                pares, matriz = cr
                n_f = (f"ROWS({rng})" if not pares else
                       ("SUM(" if matriz else "") + "COUNTIFS(" + ",".join(f"{r},{c}" for r, c in pares) + ")" + (")" if matriz else ""))
        if po.get("igual_registrado") and n_f:
            out.append({"clave": "poblacion", "rotulo": po["rotulo"], "valor": "=" + n_f, "fmt": '#,##0" partidas"',
                        "color": KPI_COLOR["poblacion"], "sub": "misma base que el saldo registrado"})
        else:
            t = {"clave": "poblacion", "rotulo": po["rotulo"], "valor": "=" + f_pob[0], "fmt": FMT_USD, "color": KPI_COLOR["poblacion"]}
            if n_f:
                t.update(sub="=" + n_f, fmt_sub='#,##0" registros"')
            out.append(t)
    for clave in ("recalculado", "registrado"):
        k = por.get(clave)
        if k:
            nota = txt["nota_registrado"] if clave == "registrado" else txt["nota_recalculado"]
            out.append({"clave": clave, "rotulo": k["rotulo"], "valor": "=" + k["f"], "fmt": FMT_USD, "color": KPI_COLOR[clave],
                        "sub": nota})
    pb = por.get("problemas")
    if pb:
        t = {"clave": "problemas", "rotulo": pb["rotulo"], "valor": "=" + pb["f"], "fmt": "#,##0", "color": KPI_COLOR[p["riesgo"]]}
        if sev:
            a, m, b = (datos_g.celda(k, sev) for k in (1, 2, 3))
            t["sub"] = f'={a}&" alta · "&{m}&" media · "&{b}&" baja"'
        out.append(t)
    # Variaciones con fórmula entre tarjetas (como las flechas ↗/↘ del HTML).
    pos = {t["clave"]: f"${COLS[i]}$" for i, t in enumerate(out)}
    fila_val = "{fila}"
    for t in out:
        if t["clave"] == "principal" and base_ref:
            t["sub"] = f"=IFERROR({pos['principal']}{fila_val}/ABS({base_ref}),0)"
            t["fmt_sub"] = '"↗ "0.0 %" de la población";"↘ "0.0 %" de la población";"0.0 % de la población"'
            t["color_sub"] = H("sube")
        if t["clave"] == "recalculado" and "registrado" in pos and not txt["nota_recalculado"]:
            t["sub"] = f"=IFERROR(({pos['recalculado']}{fila_val}-{pos['registrado']}{fila_val})/ABS({pos['registrado']}{fila_val}),0)"
            t["fmt_sub"] = f'"↗ "0.0 %" {txt["vs"]}";"↘ "0.0 %" {txt["vs"]}";"{txt["igual"]}"'
            t["color_sub"] = H("sube")
    for t in out:
        if isinstance(t.get("sub"), str):
            t["sub"] = t["sub"].replace(fila_val, str(R0 + 2))   # fila de las cifras
    return out[:5]
