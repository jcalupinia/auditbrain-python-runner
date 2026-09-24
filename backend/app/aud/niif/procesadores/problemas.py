"""Importe de cada problema como fórmula a la celda de la cédula que lo calcula.

La hoja de problemas (columnas «Código», «Descripción», «Importe») la arma cada
procesador con el importe que calculó Python. Para que el papel sea trazable,
ese importe no puede ir pegado como valor: debe remitir a la celda de la cédula
donde se origina (una diferencia, un ajuste, una pérdida), y cambiar si esa
celda cambia.

Cada procesador declara en ``REF_PROBLEMAS`` de dónde sale el importe de cada
código de problema:

- ``(hoja, columna)``: la celda de esa columna cuya fila corresponde al
  problema. Si varias filas tienen el mismo importe, se elige la fila cuyo
  identificador (cliente, factura, contrato…) aparece en la descripción.
- ``(hoja, columna, "total")``: la fila TOTAL de esa columna.
- una función ``f(hojas, e) -> (fórmula sin «=», valor)`` para importes que son
  combinación de varias celdas (por ejemplo, un SUMIF).

El enlace SOLO se escribe si el valor de la celda coincide con el importe del
problema (con su signo o con el signo contrario, que se refleja con «-»). Si no
coincide o el código no está declarado, el importe queda como valor y se reporta
en ``pendientes()``: así un mapeo equivocado nunca cambia una cifra.
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter

FILA0 = 5
COLS = ["Código", "Descripción", "Importe"]
TOL = 0.006

EXPLICA_IMPORTE = ("Trae el importe desde la celda de la cédula donde se detecta el problema (la diferencia, "
                   "el ajuste o la pérdida que lo origina): si cambia ese cálculo, el importe del problema cambia con él.")


def _num(v):
    if isinstance(v, dict):
        v = v.get("v")
    if isinstance(v, bool) or v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _texto(v) -> str:
    if isinstance(v, dict):
        v = v.get("v")
    return v.strip() if isinstance(v, str) else ""


def es_hoja_problemas(h: dict) -> bool:
    return [c[0] for c in h.get("cols") or []] == COLS


def _q(nombre: str) -> str:
    return "'" + nombre[:31].replace("'", "''") + "'!"


def celda(hojas: list[dict], hoja: str, columna: str, fila_idx: int) -> str:
    """Referencia A1 a la celda (fila_idx 0-based sobre filas + total) de la columna."""
    h = next(x for x in hojas if x["name"] == hoja)
    j = [c[0] for c in h["cols"]].index(columna)
    return f"{_q(h['name'])}{get_column_letter(j + 1)}{FILA0 + fila_idx}"


def _resolver(hojas: list[dict], ref, e: dict, importe: float):
    """(fórmula, valor) o None."""
    if callable(ref):
        r = ref(hojas, e)
        if not r:
            return None
        formula, valor = r
        v = _num(valor)
        if v is None:
            return None
        if abs(v - importe) < TOL:
            return formula, importe
        if abs(v + importe) < TOL:
            return f"-({formula})", importe
        return None
    hoja, columna = ref[0], ref[1]
    solo_total = len(ref) > 2 and ref[2] == "total"
    h = next((x for x in hojas if x["name"] == hoja), None)
    if h is None:
        return None
    nombres = [c[0] for c in h["cols"]]
    if columna not in nombres:
        return None
    j = nombres.index(columna)
    filas = list(h.get("rows") or [])
    idx_total = None
    if h.get("total"):
        idx_total = len(filas)
        filas.append(h["total"])
    candidatos = []
    for i, f in enumerate(filas):
        if solo_total and i != idx_total:
            continue
        if j >= len(f):
            continue
        v = _num(f[j])
        if v is None:
            continue
        if abs(v - importe) < TOL:
            candidatos.append((i, f, 1))
        elif abs(v + importe) < TOL:
            candidatos.append((i, f, -1))
    if not candidatos:
        return None
    if len(candidatos) > 1:
        msg = e.get("message") or ""
        con_id = [c for c in candidatos if any(t and len(t) > 2 and t in msg for t in map(_texto, c[1]))]
        if con_id:
            candidatos = con_id
    i, _, signo = candidatos[0]
    ref_a1 = f"{_q(h['name'])}{get_column_letter(j + 1)}{FILA0 + i}"
    return (ref_a1 if signo == 1 else f"-{ref_a1}"), importe


def enlazar(hojas: list[dict], refs: dict, excepciones: list | None = None) -> tuple[list[dict], list[dict]]:
    """Devuelve (hojas con la hoja de problemas enlazada, pendientes).

    ``excepciones`` (run["exceptions"]) aporta la descripción completa para
    ubicar la fila; si falta, se usa la columna «Descripción» de la hoja."""
    salida, pendientes = [], []
    for h in hojas:
        if not es_hoja_problemas(h) or not h.get("rows"):
            salida.append(h)
            continue
        filas = []
        for k, fila in enumerate(h["rows"]):
            fila = list(fila)
            codigo, desc, imp = (fila + [None, None, None])[:3]
            importe = _num(imp)
            if importe is None or abs(importe) < TOL or isinstance(imp, dict):
                filas.append(fila)
                continue
            e = {"code": codigo, "message": desc if isinstance(desc, str) else ""}
            if excepciones and k < len(excepciones) and excepciones[k].get("code") == codigo:
                e = excepciones[k]
            ref = refs.get(codigo)
            r = _resolver(hojas, ref, e, importe) if ref else None
            if r:
                fila[2] = {"f": r[0], "v": imp if not isinstance(imp, dict) else imp.get("v")}
            else:
                pendientes.append({"hoja": h["name"], "codigo": codigo, "importe": importe,
                                   "motivo": "sin mapeo" if not ref else "la celda declarada no tiene ese importe"})
            filas.append(fila)
        nueva = dict(h)
        nueva["rows"] = filas
        nueva["explica"] = {**(h.get("explica") or {}), "Importe": EXPLICA_IMPORTE}
        salida.append(nueva)
    return salida, pendientes


def refs_de(definicion: dict) -> dict:
    """``REF_PROBLEMAS`` del procesador de la definición ({} si no declara)."""
    pid = (definicion or {}).get("processor")
    if not pid:
        return {}
    from backend.app.aud.niif.procesadores import PROCESADORES  # import diferido: evita el ciclo

    mod = PROCESADORES.get(pid)
    return dict(getattr(mod, "REF_PROBLEMAS", {}) or {}) if mod else {}


def pendientes_de(mod, datasets, parametros, corte) -> list[dict]:
    """Importes de problemas que quedarían como valor (para pruebas y revisión)."""
    run = mod.ejecutar(datasets, parametros, corte)
    _, pend = enlazar(mod.hojas(run), dict(getattr(mod, "REF_PROBLEMAS", {}) or {}), run.get("exceptions"))
    return pend
