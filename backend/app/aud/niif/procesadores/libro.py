"""Papel de trabajo (Excel con fórmulas, HTML autónomo, Word y PowerPoint) de una prueba con procesador.

El exportador del sitio arma el papel de las pruebas declarativas en el
navegador; el de un procesador lo arma aquí el servidor con las cédulas que el
propio procesador calculó (``run["hojas"]``), más carátula, programa, fuentes,
conclusión y bitácora del registro.
"""
from __future__ import annotations

import html as _html
import io
import math
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from backend.app.aud.niif.procesadores import estilo_ejecutivo as est
from backend.app.aud.niif.procesadores import graficos
from backend.app.aud.niif.procesadores import problemas

NAVY, GOLD, BLANCO, CELESTE = "0A2342", "C7A83C", "FFFFFF", "DCE6F1"
_FINO = Side(style="thin", color="B7C0CC")
_DOBLE = Side(style="double", color="0A2342")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
_BORDE_TOTAL = Border(left=_FINO, right=_FINO, top=_DOBLE, bottom=_DOBLE)
_FMT = {"n": "#,##0.00", "p": "0.00%", "i": "#,##0", "a": "0", "d": "yyyy-mm-dd"}
# «g»: número con 2 a 6 decimales (tasas, factores y precios unitarios de las pruebas declarativas).
FMT_GENERAL = "#,##0.00####"


_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _excel(v, fmt):
    """Valor que se escribe en la celda: la fórmula (editable y trazable) si la
    hay; las fechas ISO como fecha de Excel, para que las fórmulas operen con ellas."""
    if isinstance(v, dict):
        return "=" + v["f"]
    if fmt in ("d", "x") and isinstance(v, str) and _ISO.match(v):
        return date.fromisoformat(v)
    return _seguro(v)


def _valor(v):
    return v.get("v") if isinstance(v, dict) else v


def _seguro(v):
    """Texto que Excel no confunda con una fórmula (regla del proyecto)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _contexto(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> list[dict]:
    e = reg.get("engagement") or {}
    run = reg.get("run") or {}
    caratula = [
        ["Firma", e.get("firm") or "AuditConsulting Auditores Cía. Ltda."], ["Cliente", e.get("client")], ["RUC", e.get("ruc")],
        ["Ejercicio", e.get("year")], ["Fecha de corte", e.get("cutoff")], ["Marco contable", e.get("framework")],
        ["Herramienta", definicion.get("name")], ["Rubro", definicion.get("area")], ["Motor", run.get("engine")],
        ["Versión del papel", f"v{version}"], ["Estado", est.estado_es(estado)], ["Preparó", e.get("preparer")], ["Revisó", e.get("reviewer")],
        ["Aprobó", reg.get("approvedBy") or ""], ["Fecha de aprobación", (reg.get("approvedAt") or "")[:10]],
        ["Huella de la ejecución (SHA-256)", reg.get("runHash") or ""],
    ]
    if reg.get("taxApplicable"):
        caratula.append(["Tratamiento tributario revisado", reg.get("taxScope") or ""])
        meta = reg.get("taxScopeMeta") or {}
        if meta:
            caratula.append(["Sustento tributario · revisión",
                             f"{meta.get('resumen', '')} · {meta.get('actor', '')} · {(meta.get('fecha') or '')[:19].replace('T', ' ')}"])
    fuentes = [[s.get("category"), s.get("document"), s.get("section"), s.get("url"), "Sí" if s.get("verified") else "No"]
               for s in reg.get("sources") or []]
    fuentes += [["NIA · " + x.get("document", ""), x.get("requirement", ""), x.get("section", ""), "", ""]
                for x in definicion.get("nia") or [] if isinstance(x, dict)]
    c = reg.get("reconciliation") or {}
    if c.get("total") is None and c.get("ledger") is None:
        # Aún no se registra (se hace al revisar la cobertura): nunca «None».
        conciliacion = "Pendiente: se registra al revisar la cobertura (población según el papel vs saldo del mayor)."
    else:
        conciliacion = (f"Población {c.get('total')} · mayor {c.get('ledger')} · diferencia {c.get('difference')}"
                        + (f" · aceptación: {c.get('acceptance')}" if c.get("acceptance") and not c.get("within") else ""))
    cierre = [
        ["Conciliación con el mayor", conciliacion],
        ["Evaluación de excepciones", reg.get("exceptionReview") or ""],
        ["Análisis", reg.get("analysis") or ""], ["Conclusión", reg.get("conclusion") or ""],
    ]
    return [
        {"name": "00_Caratula", "label": "Carátula", "cols": [["Concepto", "t"], ["Detalle", "t"]], "rows": caratula, "total": None},
        {"name": "00_Programa", "label": "Programa", "total": None,
         "cols": [["Código", "t"], ["Objetivo", "t"], ["Afirmación", "t"], ["Procedimiento", "t"], ["Evidencia", "t"], ["Criterio", "t"], ["Referencia", "t"]],
         "rows": [[p.get("code"), p.get("objective"), p.get("assertion"), p.get("procedure"), p.get("evidence"), p.get("criterion"), p.get("reference")]
                  for p in reg.get("program") or []]},
        {"name": "00_Fuentes", "label": "Base técnica", "total": None,
         "cols": [["Categoría", "t"], ["Documento", "t"], ["Párrafos", "t"], ["URL", "t"], ["Verificada", "t"]], "rows": fuentes},
    ], [
        {"name": "13_Conclusion", "label": "Conclusión", "cols": [["Concepto", "t"], ["Detalle", "t"]], "rows": cierre, "total": None},
        {"name": "14_Control_Revision", "label": "Control de revisión", "total": None,
         "cols": [["Fecha", "t"], ["Acción", "t"], ["Estado anterior", "t"], ["Estado nuevo", "t"], ["Actor", "t"], ["Comentario", "t"]],
         "rows": [[(x.get("fecha") or "")[:19].replace("T", " "), est.accion_es(x.get("accion")),
                   est.estado_es(x.get("estado_anterior")), est.estado_es(x.get("estado_nuevo")),
                   x.get("actor"), x.get("comentario")] for x in eventos]},
    ]


def cedulas(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> list[dict]:
    antes, despues = _contexto(definicion, reg, eventos, version, estado)
    run = reg.get("run") or {}
    if definicion.get("declarativa"):
        # Prueba declarativa (``declarativo``): el exportador del sitio ya trae programa, fuentes,
        # conclusión y control de revisión; su portada la reemplazan 00_Inicio y la carátula.
        return antes[:1] + list(run.get("hojas") or [])
    # El importe de cada problema remite por fórmula a la celda de la cédula que lo calcula.
    propias, _ = problemas.enlazar(run.get("hojas") or [], problemas.refs_de(definicion), run.get("exceptions"))
    # El código técnico del problema se lee como texto («TRAMO_NO_MEDIBLE» → «Tramo no medible»).
    propias = [{**h, "rows": [[graficos._humaniza(f[0]), *f[1:]] for f in h["rows"]]} if problemas.es_hoja_problemas(h) else h
               for h in propias]
    return antes + propias + despues


def _titulos_unicos(hojas: list[dict]) -> list[str]:
    """Nombres de hoja (≤31) únicos, en el orden de las cédulas."""
    titulos, vistos = [], set()
    for h in hojas:
        base = h["name"][:31]
        t, k = base, 1
        while t in vistos:
            k += 1
            t = f"{base[:28]}_{k}"
        vistos.add(t)
        titulos.append(t)
    return titulos


# Secciones del libro: nombre, color de pestaña y botón, color del texto y qué reúne.
SECCIONES = [
    ("RESULTADO", GOLD, NAVY, "resumen, problemas, asientos y conclusión"),
    ("CÓMO SE CALCULÓ", NAVY, BLANCO, "parámetros y cédulas de cálculo con fórmulas"),
    ("DATOS DEL CLIENTE", "0D7377", BLANCO, "lo que entregó el cliente, fila por fila y con su archivo de origen"),
    ("DOCUMENTACIÓN", "5B6472", BLANCO, "carátula, programa, base técnica, anexo técnico y control"),
]
HOJA_DATOS_GRAFICOS = "00_Datos_graficos"
HOJA_ANEXO = "00_Anexo_tecnico"


def _seccion(h: dict, es_resumen: bool = False) -> int:
    """Índice en SECCIONES de una cédula."""
    if "seccion" in h:           # la cédula declara su sección (pruebas declarativas)
        return h["seccion"]
    n = h.get("name", "")
    if es_resumen or problemas.es_hoja_problemas(h) or n.startswith("13_") or re.search(r"Asiento|Ajuste", n):
        return 0
    if re.match(r"D\d_", n):
        return 2
    if n.startswith(("00_", "14_")):
        return 3
    return 1


def _secciones_de(hojas: list[dict]) -> list[int]:
    i_res = next((i for i, h in enumerate(hojas) if not h["name"].startswith("00_")
                  and [c[1] for c in h.get("cols", [])] == ["t", "n"]), None)
    return [_seccion(h, i == i_res) for i, h in enumerate(hojas)]


def _ref(hoja: str, celda: str = "A1") -> str:
    return f"#'{hoja}'!{celda}"


def _ancla(ws, img, col: int, fila: int, dx_px: int, dy_px: int):
    """Coloca la imagen con desplazamiento dentro de la celda (col y fila 0-based)."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU

    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col, colOff=pixels_to_EMU(dx_px), row=fila, rowOff=pixels_to_EMU(dy_px)),
        ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)))
    ws.add_image(img)


def _grupos_nav(hojas, titulos, extra=()):
    """Botones de la portada agrupados por sección: {sección: [(rótulo, hoja destino, nombre)]}."""
    grupos = {i: [] for i in range(len(SECCIONES))}
    for (h, t), sec in zip(zip(hojas, titulos), _secciones_de(hojas)):
        etq = h.get("label", t)
        grupos[sec].append((etq.replace("Datos del cliente · ", "") if sec == 2 else etq, t, h["name"]))
    for etq, t, sec in extra:
        grupos[sec].append((etq, t, t))
    orden_res = lambda x: (0 if x[2][:3] in ("01_",) else 1 if "Problema" in x[0] else 3 if x[2].startswith("13_") else 2)  # noqa: E731
    grupos[0].sort(key=orden_res)
    return grupos


def _formula_por_valor(hojas, titulos, valor, etiqueta=None):
    """Fórmula (sin «=») a la celda que ya tiene ese importe: primero la fila del Resumen con
    el mismo rótulo, luego cualquier fila del Resumen con el mismo importe y, por último, la fila
    TOTAL de una cédula. None si ninguna celda lo tiene (nunca se pega el valor)."""
    if valor is None:
        return None
    if etiqueta:
        f = _formula_kpi(hojas, titulos, etiqueta, valor, "n")
        if f:
            return f[1:]
    for i, h in enumerate(hojas):
        if [c[1] for c in h.get("cols", [])] != ["t", "n"]:
            continue
        for k, fila in enumerate(h.get("rows") or []):
            if len(fila) > 1 and graficos._num(fila[1]) is not None and abs(graficos._num(fila[1]) - valor) < 0.006:
                return f"{_q(titulos[i])}B{5 + k}"
    for i, h in enumerate(hojas):
        tot = h.get("total")
        if not tot:
            continue
        n = len(h.get("rows") or [])
        for j, v in enumerate(tot):
            if graficos._num(v) is not None and abs(graficos._num(v) - valor) < 0.006:
                return f"{_q(titulos[i])}{get_column_letter(j + 1)}{5 + n}"
    return None


def _crit(v) -> str:
    """Criterio de SUMIFS que iguala el contenido tal cual (sin comodines)."""
    if v in (None, ""):
        return '"="'
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')
    return f'"={t}"'


def _rango_col(hojas, titulos, i, col):
    h = hojas[i]
    j = next((k for k, c in enumerate(h.get("cols") or []) if c[0] == col), None)
    n = len(h.get("rows") or [])
    if j is None or n == 0:
        return None
    L = get_column_letter(j + 1)
    return f"{_q(titulos[i])}${L}$5:${L}${4 + n}"


def _criterios(spec, hojas, titulos, i):
    """Pares (rango, criterio) de ``donde`` y ``con_valor``, y si hace falta envolver en SUM
    (lista de admitidos como constante matricial). None si no se puede expresar."""
    pares, matriz, sep = [], False, [",", ";"]
    for col, admitidos in (spec.get("donde") or {}).items():
        rng = _rango_col(hojas, titulos, i, col)
        if rng is None:
            return None
        adm = list(admitidos)
        if len(adm) == 1:
            pares.append((rng, _crit(adm[0])))
        else:
            if not sep:
                return None            # más de dos listas: no cabe en una constante 2D
            pares.append((rng, "{" + sep.pop(0).join(_crit(a) for a in adm) + "}"))
            matriz = True
    if spec.get("con_valor"):
        rng = _rango_col(hojas, titulos, i, spec["con_valor"])
        if rng is None:
            return None
        pares.append((rng, '">-1E+307"'))  # solo filas con número en esa columna
    return pares, matriz


def _suma_si(rng_val, pares, matriz, extra=()):
    todos = list(pares) + list(extra)
    if not todos:
        return f"SUM({rng_val})"
    f = f"SUMIFS({rng_val}," + ",".join(f"{r},{c}" for r, c in todos) + ")"
    return f"SUM({f})" if matriz else f


def _formula_spec(spec, run, hojas, titulos):
    """(fórmula sin «=», valor) de un indicador del ``PANEL`` o None."""
    if not spec:
        return None
    if "total" in spec:
        v = graficos._num((run.get("totals") or {}).get(spec["total"]))
        f = _formula_por_valor(hojas, titulos, v, (run.get("labels") or {}).get(spec["total"]))
        return (f, v) if f else None
    i = next((k for k, h in enumerate(hojas) if h["name"] == spec.get("hoja")), None)
    if i is None:
        return None
    rng = _rango_col(hojas, titulos, i, spec.get("col"))
    cr = _criterios(spec, hojas, titulos, i)
    if rng is None or cr is None:
        return None
    v = graficos._suma_col(hojas[i], spec.get("col"), spec)
    return _suma_si(rng, *cr), v


def _kpis_panel(definicion, reg, hojas, titulos):
    """Indicadores de la portada = los del panel del HTML (resultado principal, población,
    recalculado, registrado y problemas). Cada uno es una fórmula; el que no tiene celda de
    origen no se muestra (nunca un valor pegado)."""
    run = reg.get("run") or {}
    spec = getattr(graficos.modulo(definicion), "PANEL", None) or {}
    tot, etq = run.get("totals") or {}, run.get("labels") or {}
    prim = run.get("primary")
    out = []
    if prim in tot:
        v = graficos._num(tot[prim])
        f = _formula_por_valor(hojas, titulos, v, etq.get(prim, prim))
        if f:
            out.append({"clave": "principal", "rotulo": etq.get(prim, prim), "valor": v, "fmt": "n", "f": f})
    for clave, defecto in (("poblacion", "Población"), ("recalculado", "Recalculado"), ("registrado", "Registrado")):
        r = _formula_spec(spec.get(clave), run, hojas, titulos)
        if not r or r[1] is None:
            continue
        # La población de una prueba de saldo es el mismo saldo registrado: no se repite la tarjeta.
        if clave == "registrado" and any(k["clave"] == "poblacion" and abs(k["valor"] - r[1]) < 0.005 for k in out):
            out = [k for k in out if k["clave"] != "poblacion"]
        out.append({"clave": clave, "rotulo": (spec.get(clave) or {}).get("rotulo", defecto), "valor": r[1], "fmt": "n", "f": r[0]})
    n_prob = len(run.get("exceptions") or [])
    ip = (problemas.hoja_problemas(hojas) or (None,))[0]
    if ip is not None:
        n = len(hojas[ip].get("rows") or [])
        out.append({"clave": "problemas", "rotulo": graficos.textos(spec)["problemas"], "valor": n_prob, "fmt": "i",
                    "semaforo": est.color_semaforo(n_prob), "f": f"COUNTA({_q(titulos[ip])}A5:A{4 + max(n, 1)})"})
    return out[:5]


def _formula_kpi(hojas, titulos, etiqueta, valor, fmt):
    """Fórmula del indicador de la portada: el conteo de la hoja de problemas o la
    fila del Resumen con el mismo rótulo e importe. None si no hay celda de origen."""
    if fmt == "i":
        ip = (problemas.hoja_problemas(hojas) or (None,))[0]
        if ip is None:
            return None
        n = len(hojas[ip].get("rows") or [])
        return f"=COUNTA({_q(titulos[ip])}A5:A{4 + max(n, 1)})"
    try:
        objetivo = float(str(valor).replace(",", ""))
    except (TypeError, ValueError):
        return None
    for i, h in enumerate(hojas):
        if [c[1] for c in h.get("cols", [])] != ["t", "n"]:
            continue
        for k, f in enumerate(h.get("rows") or []):
            if len(f) > 1 and _valor(f[0]) == etiqueta and graficos._num(f[1]) is not None \
                    and abs(graficos._num(f[1]) - objetivo) < 0.006:
                return f"={_q(titulos[i])}B{5 + k}"
    return None


def _print_setup(ws, e):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    firma = e.get("firm") or "AuditConsulting Auditores Cía. Ltda."
    ws.oddFooter.left.text = firma
    ws.oddFooter.center.text = str(e.get("client") or "")
    ws.oddFooter.right.text = "Página &P de &N"


def _hoja_ejecutiva(ws, S, h, titulo_prueba, nav, hojas=None, anexo=None):
    ws.sheet_view.showGridLines = False
    ancho = max(6, len(h["cols"]))
    # Fila 1: botonera arriba a la izquierda (siempre visible: el panel se congela hasta la fila 4) y la prueba.
    # Filas 1-3: la franja del HTML (fondo azul marino, título claro y botones del HTML).
    from backend.app.aud.niif.procesadores import panel_excel as px

    for r in (1, 2, 3):
        for j in range(1, ancho + 1):
            ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor=px.BG)
    for col, clave, texto in (("A", "inicio", "⟵ Inicio"), ("B", "anterior", "◀ Anterior"), ("C", "siguiente", "Siguiente ▶")):
        if nav.get(clave):
            px.boton_html(ws[f"{col}1"], texto, _ref(nav[clave]), primario=(clave == "inicio"))
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=ancho)
    ws["D1"].value = _seguro(titulo_prueba)
    ws["D1"].font = Font(name=est.FONT_TEXTO, size=9, color=px.TEXTO2)
    ws["D1"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24
    # Fila 2: título de la cédula.
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    ws["A2"].value = _seguro(h["label"])
    ws["A2"].font = Font(name=est.FONT_TITULO, size=15, bold=True, color=px.TEXTO)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 28
    if h.get("guia"):
        # «¿De dónde saco este dato?»: qué documento, reporte o cuenta alimenta la hoja.
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(6, len(h["cols"])))
        g = ws["A3"]
        g.value = _seguro("¿De dónde saco este dato?  " + h["guia"])
        g.font = Font(name=est.FONT_TEXTO, size=9, italic=True, color=px.TEXTO2)
        g.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[3].height = 15 * max(2, math.ceil(len(h["guia"]) / 150))
    # Cabecera de la tabla
    fila_enc = 4
    anchos = [len(c[0]) + 2 for c in h["cols"]]
    for j, (nombre, _) in enumerate(h["cols"], start=1):
        c = ws.cell(row=fila_enc, column=j, value=nombre)
        c.font = S["encabezado"]
        c.fill = S["fill_encabezado"]
        c.alignment = S["centro"]
        c.border = S["borde_enc"]
    ws.row_dimensions[fila_enc].height = 30
    # Subrayado dorado bajo el título de la cédula
    for j in range(1, max(6, len(h["cols"])) + 1):
        ws.cell(row=2, column=j).border = S["filete_oro"]
    from backend.app.aud.niif.procesadores.base import estilo_fila

    filas = [(r, False) for r in h["rows"]] + ([(h["total"], True)] if h.get("total") else [])
    for i, (fila, total) in enumerate(filas, start=fila_enc + 1):
        ef = {} if total else estilo_fila(h, i - fila_enc - 1)
        tipo = ef.get("tipo")
        for j, ((nombre_col, fmt), v) in enumerate(zip(h["cols"], fila), start=1):
            c = ws.cell(row=i, column=j, value=_excel(v, fmt))
            es_num = fmt in est.FMT or (fmt in ("x", "g") and isinstance(_valor(v), (int, float)))
            fuerte = total or tipo in ("titulo", "total")
            c.font = S["total"] if fuerte and es_num else (
                Font(name=est.FONT_TEXTO, size=9, bold=True, color=est.NAVY) if fuerte else (S["cifra"] if es_num else S["dato"]))
            if tipo == "control":
                c.font = Font(name=est.FONT_TEXTO, size=8.5, italic=True, color="4B5563")
            c.border = S["borde_total"] if total or tipo == "total" else S["borde_fila"]
            if total or tipo == "total":
                c.fill = S["fill_total"]
            elif tipo == "titulo":
                c.fill = PatternFill("solid", fgColor="E8EEF7")   # rubro de la sumaria
            elif (i - fila_enc) % 2 == 0 and not ef:
                c.fill = S["fill_zebra"]   # filas alternas
            if isinstance(c.value, date) or (isinstance(v, dict) and isinstance(v.get("v"), str) and _ISO.match(v["v"])):
                c.number_format = est.FMT["d"]   # también las fórmulas que devuelven una fecha
                c.alignment = S["centro"]
            elif fmt in est.FMT:
                c.number_format = est.FMT[fmt]
                c.alignment = S["der"]
            elif fmt == "g" and es_num:
                c.number_format = FMT_GENERAL
                c.alignment = S["der"]
            elif es_num:
                c.alignment = S["der"]
            else:
                c.alignment = S["izq"]
            if ef.get("sangria") and nombre_col == ef.get("col"):
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1 + 2 * int(ef["sangria"]))
            vista = _valor(v)
            anchos[j - 1] = max(anchos[j - 1], min(60, len(str(vista if vista is not None else "")) + 2))
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(13 if j <= 3 else 12, min(60, w))
    for j in range(len(anchos) + 1, 4):   # la botonera ocupa A:C aunque la tabla tenga menos columnas
        ws.column_dimensions[get_column_letter(j)].width = 13
    for j in range(len(anchos) + 1, max(6, len(anchos)) + 1):   # columnas del bloque «Cómo se calcula» fuera de la tabla
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, 28)
    # Columnas técnicas (p. ej. claves de cruce): agrupadas y ocultas; el «+» de Excel las muestra.
    nombres = [c[0] for c in h["cols"]]
    for nombre in h.get("ocultas") or []:
        if nombre in nombres:
            dim = ws.column_dimensions[get_column_letter(nombres.index(nombre) + 1)]
            dim.hidden = True
            dim.outlineLevel = 1
    ws.freeze_panes = f"A{fila_enc + 1}"
    _bloque_como_se_calcula(ws, S, h, len(filas) + fila_enc + 2, hojas, anexo, ws.title)
    _print_setup(ws, {})



def _q(titulo: str) -> str:
    """Nombre de hoja citado para una fórmula: 'Hoja con espacios'!"""
    return "'" + titulo.replace("'", "''") + "'!"


def _bloque_como_se_calcula(ws, S, h, fila_inicio, hojas=None, anexo=None, titulo_hoja=None):
    """Debajo de la tabla, «ⓘ Cómo se calcula esta hoja» en lenguaje sencillo: por cada columna
    calculada, qué hace y de dónde viene el dato. La fórmula de Excel y el ejemplo con números van
    al «Anexo técnico» (``anexo``), para el auditor que quiera revisarlos."""
    bloque = como_se_calcula(h, hojas)
    if not bloque:
        return
    ancho = max(6, len(h["cols"]))
    corte = max(3, ancho - 2)                    # B..corte: explicación; corte+1..ancho: de dónde viene
    ancho_de = lambda a, b: sum((ws.column_dimensions[get_column_letter(k)].width or 10) for k in range(a, b + 1))  # noqa: E731
    r = fila_inicio
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ancho)
    t = ws.cell(row=r, column=1, value="ⓘ  Cómo se calcula esta hoja")
    t.font = S["subtitulo"]
    t.fill = S["fill_gold"]
    t.alignment = S["izq"]
    r += 1
    for c1, c2, txt in ((1, 1, "Columna"), (2, corte, "Cómo se calcula"), (corte + 1, ancho, "De dónde viene el dato")):
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        c = ws.cell(row=r, column=c1, value=txt)
        c.font = S["encabezado"]
        c.fill = S["fill_encabezado"]
        c.alignment = S["centro"]
        for k in range(c1, c2 + 1):
            ws.cell(row=r, column=k).fill = S["fill_encabezado"]
    for b in bloque:
        r += 1
        for c1, c2, txt in ((1, 1, b["columna"]), (2, corte, b["explicacion"]), (corte + 1, ancho, b["origen"])):
            if c2 > c1:
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            c = ws.cell(row=r, column=c1, value=_seguro(txt))
            c.font = Font(name=est.FONT_TEXTO, size=9, bold=(c1 == 1), color=NAVY if c1 == 1 else "1A1A1A")
            c.alignment = S["izq"]
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).fill = S["fill_panel"]
                ws.cell(row=r, column=k).border = S["borde_fila"]
        lineas = max(math.ceil(len(b["explicacion"]) * 1.15 / max(ancho_de(2, corte), 10)),
                     math.ceil(len(b["origen"]) * 1.15 / max(ancho_de(corte + 1, ancho), 10)), 1)
        ws.row_dimensions[r].height = 13 * lineas + 4
        if anexo is not None:
            anexo.append({"hoja": titulo_hoja or h["name"], "cedula": h["label"], "columna": b["columna"],
                          "formula": b["formula"], "ejemplo": b["ejemplo"], "norma": b.get("norma") or "—"})
    r += 1
    nota = ws.cell(row=r, column=1, value="La fórmula de Excel de cada columna y un ejemplo con números están en la hoja «Anexo técnico».")
    nota.font = S["nota"]
    nota.hyperlink = _ref(HOJA_ANEXO)
    ws.print_area = None


def _anexo_tecnico(ws, S, anexo, titulo_prueba):
    """Hoja «Anexo técnico»: la fórmula de Excel de cada columna calculada, con un ejemplo con
    números reales y la norma, para el auditor que revisa el papel (el resto del libro va en sencillo)."""
    from backend.app.aud.niif.procesadores import panel_excel as px

    ws.sheet_view.showGridLines = False
    for r in (1, 2, 3):
        for j in range(1, 6):
            ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor=px.BG)
    px.boton_html(ws["A1"], "⟵ Inicio", _ref("00_Inicio"), primario=True)
    ws.row_dimensions[1].height = 24
    ws.merge_cells("B1:E1")
    ws["B1"].value = _seguro(titulo_prueba)
    ws["B1"].font = Font(name=est.FONT_TEXTO, size=9, color=px.TEXTO2)
    ws["B1"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.merge_cells("A2:E2")
    ws["A2"].value = "Anexo técnico · fórmulas de Excel de cada cédula"
    ws["A2"].font = Font(name=est.FONT_TITULO, size=15, bold=True, color=px.TEXTO)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 28
    for j in range(1, 6):
        ws.cell(row=2, column=j).border = S["filete_oro"]
    cab = ["Hoja", "Columna", "Fórmula de Excel (primera fila)", "Ejemplo con números reales", "Norma"]
    for j, txt in enumerate(cab, start=1):
        c = ws.cell(row=4, column=j, value=txt)
        c.font = S["encabezado"]
        c.fill = S["fill_encabezado"]
        c.alignment = S["centro"]
        c.border = S["borde_enc"]
    for i, a in enumerate(anexo, start=5):
        # La fórmula se muestra sin el «=» inicial: así Excel la lee como texto (sin apóstrofo visible).
        vals = [a["cedula"], a["columna"], _seguro(a["formula"].lstrip("=")), _seguro(a["ejemplo"]), a["norma"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = S["cifra"] if j == 3 else S["dato"]
            c.alignment = S["izq"]
            c.border = S["borde_fila"]
            if (i - 4) % 2 == 0:
                c.fill = S["fill_zebra"]
        ws.cell(row=i, column=1).hyperlink = _ref(a["hoja"])
        ws.row_dimensions[i].height = 13 * max(1, math.ceil(max(len(vals[2]), len(vals[3])) / 70)) + 4
    for col, w in zip("ABCDE", (30, 24, 70, 70, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    _print_setup(ws, {})


def xlsx(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    """Papel de trabajo ejecutivo: panel 00_Inicio con marca, KPIs y navegación,
    y cada cédula como hoja «tipo software» (sin cuadrícula, botones de
    navegación, semáforo, impresión horizontal). Las celdas con fórmula se
    escriben como fórmula viva (trazable)."""
    S = est.estilos(Font, PatternFill, Border, Side, Alignment)
    wb = Workbook()
    wb.remove(wb.active)
    hojas = cedulas(definicion, reg, eventos, version, estado)
    titulos = _titulos_unicos(hojas)
    titulo_prueba = f"{definicion.get('name', '')} · {(reg.get('engagement') or {}).get('client', '')} · corte {(reg.get('engagement') or {}).get('cutoff', '')}"

    inicio = wb.create_sheet("00_Inicio")
    inicio.sheet_properties.tabColor = GOLD
    datos_graf = wb.create_sheet(HOJA_DATOS_GRAFICOS)
    datos_graf.sheet_state = "hidden"
    _print_setup(datos_graf, {})
    # Portada = el panel del HTML (tema «Ejecutivo»): mismos KPI, gráficos, colores y botones.
    from backend.app.aud.niif.procesadores import panel_excel

    fin_panel = panel_excel.portada(inicio, datos_graf, definicion, reg, hojas, titulos, estado, version,
                                    _grupos_nav(hojas, titulos, [("Anexo técnico (fórmulas)", HOJA_ANEXO, 3)]), SECCIONES)
    inicio.print_options.horizontalCentered = True
    _print_setup(inicio, reg.get("engagement") or {})
    inicio.print_area = f"A1:G{fin_panel}"

    anexo = []
    secciones = _secciones_de(hojas)
    for idx, (h, t) in enumerate(zip(hojas, titulos)):
        ws = wb.create_sheet(t)
        ws.sheet_properties.tabColor = SECCIONES[secciones[idx]][1]   # pestaña del color de su sección
        nav = {"inicio": "00_Inicio",
               "anterior": titulos[idx - 1] if idx > 0 else None,
               "siguiente": titulos[idx + 1] if idx < len(titulos) - 1 else None}
        _hoja_ejecutiva(ws, S, h, titulo_prueba, nav, hojas, anexo)
    # Anexo técnico junto a la documentación (después de la base técnica).
    pos = next((wb.sheetnames.index(t) + 1 for t in titulos if t == "00_Fuentes"), 2)
    tec = wb.create_sheet(HOJA_ANEXO, pos)
    tec.sheet_properties.tabColor = SECCIONES[3][1]
    _anexo_tecnico(tec, S, anexo, titulo_prueba)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # el gráfico y las fórmulas se calculan al abrir
    salida = io.BytesIO()
    wb.save(salida)
    return _imagenes_con_marco(salida.getvalue())


_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _imagenes_con_marco(datos: bytes) -> bytes:
    """Los logotipos con posición y tamaño explícitos (``<a:xfrm>``), como los guarda Excel.

    openpyxl escribe la imagen solo con su ancla; Excel la muestra, pero varios visores
    (vista previa del celular, visores web) necesitan el marco dentro de ``spPr`` y, sin él,
    no la dibujan: el usuario veía el Excel «sin logos»."""
    import zipfile

    patron = re.compile(r'(<ext cx="(\d+)" cy="(\d+)"/><pic>.*?<spPr>)', re.S)

    def marco(m):
        return (m.group(1) + f'<a:xfrm xmlns:a="{_A}"><a:off x="0" y="0"/>'
                f'<a:ext cx="{m.group(2)}" cy="{m.group(3)}"/></a:xfrm>')

    ent, sal = io.BytesIO(datos), io.BytesIO()
    with zipfile.ZipFile(ent) as zin, zipfile.ZipFile(sal, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            contenido = zin.read(item.filename)
            if item.filename.startswith("xl/drawings/drawing") and item.filename.endswith(".xml"):
                contenido = patron.sub(marco, contenido.decode("utf-8")).encode("utf-8")
            zout.writestr(item, contenido)
    return sal.getvalue()


def _celda(v, fmt) -> str:
    v = _valor(v)
    if v is None or v == "":
        return ""
    if fmt in ("n", "p", "i", "a", "g") and not isinstance(v, (int, float)):
        return _html.escape(str(v))          # texto en una columna numérica («No aplica», «—»)
    if fmt == "g":                               # 2 a 6 decimales, como el exportador del sitio
        entero, dec = f"{float(v):,.6f}".split(".")
        return (entero.replace(",", ".") + "," + dec.rstrip("0").ljust(2, "0"))
    if fmt == "n" or (fmt == "x" and isinstance(v, (int, float))):
        return f"{float(v):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    if fmt == "p":
        return f"{float(v) * 100:.2f} %".replace(".", ",")
    if fmt == "i":
        return f"{int(v):,}".replace(",", ".")
    if fmt == "a":                               # año: 2025, nunca «2.025»
        return str(int(v))
    return _html.escape(str(v))


def _filas(h):
    return [(r, False) for r in h["rows"]] + ([(h["total"], True)] if h.get("total") else [])


# --- «Cómo se calcula esta hoja» ---------------------------------------------
# Se GENERA desde la misma definición de la cédula que produce las fórmulas
# (h["cols"]/h["rows"] con celdas {"f":…,"v":…}). Una sola fuente: si la fórmula
# cambia, la explicación cambia.
_REF = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)")
# Referencia a celda o a rango, con hoja opcional: 'Hoja'!$A$5:$A$24 · B5 · $B$8
_REF_RANGO = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
_TEXTO_FORMULA = re.compile(r'"[^"]*"')
_COLUMNA_COMPLETA = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3}):\$?([A-Z]{1,3})(?![0-9A-Za-z])")
FILA_DATOS = 5  # los datos de toda cédula arrancan en la fila 5 (fila 4 = encabezado)

# Plantilla de respaldo cuando una columna calculada NO tiene explicación escrita en
# la definición de su cédula. Nunca debe llegar al papel: el test
# tests/test_aud_explicaciones_humanas.py falla si alguna la usa.
PLANTILLA_GENERICA = "se obtiene con la fórmula indicada"


def _fmt_num(v, fmt) -> str:
    return _celda(v, fmt if fmt in ("n", "p", "i", "a") else "n")


def _mapa_hojas(hojas) -> dict:
    return {h["name"]: h for h in (hojas or [])}


def _celda_de(h: dict, letra: str, num: int):
    """(valor, formato, encabezado) de la celda letra+num de la cédula h."""
    j = column_index_from_string(letra) - 1
    cols = h.get("cols") or []
    fmt = cols[j][1] if 0 <= j < len(cols) else "n"
    enc = cols[j][0] if 0 <= j < len(cols) else letra
    i = num - FILA_DATOS
    filas = h.get("rows") or []
    if 0 <= i < len(filas):
        fila = filas[i]
    elif i == len(filas) and h.get("total"):
        fila = h["total"]
    else:
        return None, fmt, enc
    return (_valor(fila[j]) if 0 <= j < len(fila) else None), fmt, enc


def _muestra(v, fmt) -> str:
    """Valor de una celda para leerlo dentro del ejemplo (texto entre comillas)."""
    if v is None or v == "":
        return "(vacío)"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _fmt_num(v, fmt)
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return f"«{_html.unescape(str(v))}»"


def _ejemplo_fila1(formula: str, h: dict, mapa: dict, fmt: str) -> str:
    """El cálculo de la FILA 1 con sus números reales: cada celda de la fórmula se
    reemplaza por su valor (de esta hoja o de la hoja citada) y cada rango por el
    nombre de la columna y la hoja de donde sale. Siempre devuelve el cálculo y el
    resultado; nunca queda vacío."""
    textos = []

    def guarda(m):  # protege los textos literales ("<=0", "Sí") de la sustitución
        textos.append(m.group(0))
        return f"\x00{len(textos) - 1}\x00"

    expr = _TEXTO_FORMULA.sub(guarda, formula)
    # Separador de argumentos «;» (como Excel en español): con decimales es-EC la
    # coma ya es parte de los números. Se hace ANTES de sustituir los valores.
    expr = expr.replace(",", "; ")

    def sustituye(m):
        hoja_ref, l1, n1, l2, n2_ = m.group(1), m.group(2), int(m.group(3)), m.group(4), m.group(5)
        destino = mapa.get(hoja_ref, h) if hoja_ref else h
        nombre_hoja = (destino.get("label") or hoja_ref) if hoja_ref else None
        if l2:  # rango → «Columna» de Hoja (n filas)
            _, _, enc = _celda_de(destino, l1, n1)
            n = int(n2_) - n1 + 1
            donde = f" de «{nombre_hoja}»" if nombre_hoja else ""
            return f"[«{enc}»{donde}, {n} fila{'s' if n != 1 else ''}]"
        v, f_cel, _ = _celda_de(destino, l1, n1)
        if v is None and hoja_ref and hoja_ref not in mapa:
            return m.group(0)
        return _muestra(v, f_cel)

    def columna_completa(m):  # 'Hoja'!$F:$F → [«Columna» de «Hoja», columna completa]
        hoja_ref, l1 = m.group(1), m.group(2)
        destino = mapa.get(hoja_ref, h) if hoja_ref else h
        _, _, enc = _celda_de(destino, l1, FILA_DATOS)
        donde = f" de «{destino.get('label') or hoja_ref}»" if hoja_ref else ""
        return f"[«{enc}»{donde}, columna completa]"

    expr = _COLUMNA_COMPLETA.sub(columna_completa, expr)
    expr = _REF_RANGO.sub(sustituye, expr)
    expr = re.sub("\x00(\\d+)\x00", lambda m: textos[int(m.group(1))], expr)
    return expr


def como_se_calcula(h: dict, hojas: list | None = None) -> list[dict]:
    """Filas del bloque «Cómo se calcula esta hoja»: una por columna CALCULADA.

    - ``explicacion``: la escrita por una persona en la definición de la cédula
      (``hoja(..., explica={...})``). Si falta, se usa la plantilla genérica, que
      los tests rechazan.
    - ``ejemplo``: el cálculo de la fila 1 con sus números reales y el resultado.
    - ``origen``: las columnas de ESTA hoja y las hojas de donde salen los datos
      (una referencia a otra hoja se nombra por esa hoja, no por una columna de
      esta, que era el error anterior).
    ``hojas`` (todas las cédulas del papel) permite leer valores de otras hojas."""
    filas = h.get("rows") or []
    if not filas:
        return []
    mapa = _mapa_hojas(hojas)
    mapa.setdefault(h["name"], h)
    cols = [c[0] for c in h["cols"]]
    header_por_letra = {get_column_letter(j + 1): cols[j] for j in range(len(cols))}
    escritas = h.get("explica") or {}
    bloque = []
    for j, (nombre, fmt) in enumerate(h["cols"]):
        # La primera fila donde la columna tiene fórmula (en casi todas es la fila 1;
        # en otras —p. ej. «Haber» de los asientos— la fila 1 es un dato).
        k = next((i for i, f in enumerate(filas) if j < len(f) and isinstance(f[j], dict) and "f" in f[j]), None)
        if k is None:
            continue
        v0 = filas[k][j]
        formula = v0["f"]
        origen_cols, origen_hojas = [], []
        for m in _REF_RANGO.finditer(_TEXTO_FORMULA.sub('""', formula)):
            hoja_ref, letra = m.group(1), m.group(2)
            if hoja_ref:
                etq = (mapa.get(hoja_ref) or {}).get("label") or hoja_ref
                if etq not in origen_hojas:
                    origen_hojas.append(etq)
            elif letra in header_por_letra and header_por_letra[letra] not in origen_cols:
                origen_cols.append(header_por_letra[letra])
        explic = escritas.get(nombre)
        if not explic:
            explic = f"«{nombre}» {PLANTILLA_GENERICA}."
        resultado = _valor(v0)
        if resultado in (None, ""):
            res_txt = "(en blanco)"
        elif isinstance(resultado, (int, float)) and not isinstance(resultado, bool):
            res_txt = _fmt_num(resultado, fmt)
        else:
            res_txt = _muestra(resultado, fmt).strip("«»")
        ejemplo = f"Fila {k + 1}: {_ejemplo_fila1(formula, h, mapa, fmt)} → {res_txt}"
        origen = [f"«{c}» (esta hoja)" for c in origen_cols] + [f"hoja «{x}»" for x in origen_hojas]
        bloque.append({"columna": nombre, "formula": "=" + formula, "explicacion": explic, "ejemplo": ejemplo,
                       "origen": (h.get("origen") or {}).get(nombre) or (", ".join(origen) if origen else "datos cargados del cliente"),
                       "destino": h.get("label", h["name"]), "norma": h.get("norma") or "",
                       "escrita": nombre in escritas})
    return bloque


# Cédulas que van a la presentación (PowerPoint): las de lectura ejecutiva. Se comparan por el nombre sin el número,
# porque una herramienta puede renumerar sus cédulas al insertar una nueva.
_EN_PPT = ("Resumen", "Matriz_deterioro", "Por_cliente", "Fiscal", "Asientos", "Problemas", "Conclusion")


def _en_ppt(nombre: str) -> bool:
    return nombre.split("_", 1)[-1] in _EN_PPT
_MAX_FILAS_PPT = 14


def _rgb(RGBColor, hex6):
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


def csv_zip(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    """Un CSV por cédula dentro de un ZIP. UTF-8 con BOM y separador «;» para que
    Excel en español lo abra bien; solo datos limpios (valores, no fórmulas)."""
    import csv
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for h in cedulas(definicion, reg, eventos, version, estado):
            sio = io.StringIO()
            w = csv.writer(sio, delimiter=";")
            w.writerow([c[0] for c in h["cols"]])
            for fila, _total in _filas(h):
                w.writerow([_html.unescape(_celda(v, fmt)) for (_, fmt), v in zip(h["cols"], fila)])
            z.writestr(f"{h['name']}.csv", ("﻿" + sio.getvalue()).encode("utf-8"))
    return buf.getvalue()


def docx(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    """Word del papel con el diseño del HTML impreso (tema Claro): ver ``papel_office``."""
    from backend.app.aud.niif.procesadores import papel_office

    return papel_office.docx(definicion, reg, eventos, version, estado)


def pptx(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    """PowerPoint del papel con el diseño del HTML en pantalla (tema Ejecutivo): ver ``papel_office``."""
    from backend.app.aud.niif.procesadores import papel_office

    return papel_office.pptx(definicion, reg, eventos, version, estado)


def html(definicion: dict, reg: dict, eventos: list, version: int, estado: str, para_pdf: bool = False) -> bytes:
    """HTML autónomo con el dashboard ejecutivo (``html_ejecutivo``): funciona sin
    internet (sin fuentes, scripts ni estilos externos) y trae dentro el Excel con
    fórmulas, el Word, el PowerPoint y el CSV para descargarlos. ``para_pdf=True``
    devuelve la versión estática (tema Claro, todo visible, sin JS) para el PDF."""
    from backend.app.aud.niif.procesadores import html_ejecutivo

    hojas = cedulas(definicion, reg, eventos, version, estado)
    adjuntos = [] if para_pdf else [
        ("xlsx", "Excel con fórmulas", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         xlsx(definicion, reg, eventos, version, estado)),
        ("docx", "Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         docx(definicion, reg, eventos, version, estado)),
        ("pptx", "PowerPoint", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
         pptx(definicion, reg, eventos, version, estado)),
        ("zip", "CSV (ZIP)", "application/zip", csv_zip(definicion, reg, eventos, version, estado)),
    ]
    calculadora = None
    if definicion.get("declarativa") and not para_pdf:
        # Prueba declarativa: la «Calculadora reutilizable» del sitio, con su motor portable.
        from backend.app.aud.niif.procesadores import declarativo

        calculadora = declarativo.calculadora(definicion, reg)
    return html_ejecutivo.render(definicion, reg, eventos, version, estado, hojas, adjuntos, _celda, como_se_calcula,
                                 para_pdf=para_pdf, calculadora=calculadora).encode("utf-8")


class PDFNoDisponible(ValueError):
    """WeasyPrint (o sus librerías nativas Pango/cairo) no está disponible en
    este entorno. El endpoint lo traduce a un aviso, no a un error 500."""


def pdf(definicion: dict, reg: dict, eventos: list, version: int, estado: str) -> bytes:
    """PDF ejecutivo generado en el servidor (WeasyPrint) a partir del HTML
    estático: fiel a la vista, horizontal, con marca, KPIs, cédulas y el bloque
    «Cómo se calcula». No usa la impresión del navegador.

    En entornos sin las librerías nativas de WeasyPrint (p. ej. un servicio
    Render ``env: python``) degrada con ``PDFNoDisponible``; el papel sigue
    disponible en Excel/Word/HTML y el HTML trae «Guardar como PDF»."""
    try:
        import weasyprint  # requiere Pango/cairo/gdk-pixbuf nativos
    except Exception as e:  # ImportError o OSError al cargar las libs nativas
        raise PDFNoDisponible(
            "El PDF en el servidor no está disponible en este entorno (faltan las "
            "librerías nativas de WeasyPrint). Descargue el papel en Excel, Word o "
            "HTML, o use «Guardar como PDF» desde el HTML."
        ) from e
    contenido = html(definicion, reg, eventos, version, estado, para_pdf=True)
    return weasyprint.HTML(string=contenido.decode("utf-8")).write_pdf()
