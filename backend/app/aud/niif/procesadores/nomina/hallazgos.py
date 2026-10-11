# -*- coding: utf-8 -*-
"""Sujetos de los hallazgos derivados de los datos del cliente (no fijos).

Antes los generadores de referencia traían los nombres de los trabajadores de
PROPHAR escritos a mano en las narrativas. Aquí se calculan desde la población,
de modo que el código fuente no contiene datos personales de ningún cliente y el
texto se arma en tiempo de ejecución con los nombres y los importes que tocan."""
from __future__ import annotations
import collections
import openpyxl
from backend.app.aud.niif.procesadores.nomina import motor


def cifras(R, rutas, comp=None) -> dict:
    """Importes de los hallazgos y metas de control, derivados de los datos del
    cliente (no fijos). Reemplaza las ~393 cifras de PROPHAR que los generadores
    de referencia traían escritas a mano. Se calcula del motor, del anexo RQ-002,
    del rol, de los mayores RQ-004 y de los comprobantes RQ-005."""
    from backend.app.aud.niif.procesadores.nomina import cedulas_nuevas, series, mayores
    det = R['det']; cons = R['cons']
    S = lambda col, pred=lambda d: True: round(sum(d[col] for d in det if pred(d)), 2)
    SC = lambda col: round(sum(r[col] for r in cons), 2)

    pas, gas = cedulas_nuevas.leer_anexo(rutas.anexo)
    def gasto_concepto(concepto):  # gasto libros por concepto DO-2 (misma clasificación por palabra clave)
        return round(sum(s25 for cod, desc, s24, s25 in gas if cedulas_nuevas._concepto_do2(desc) == concepto), 2)
    sal_ini = next((abs(s24) for cod, desc, s24, s25 in pas if 'vacacion' in desc.lower()), 0.0)
    vac_libros = next((abs(s25) for cod, desc, s24, s25 in pas if 'vacacion' in desc.lower()), 0.0)
    pasivo_vac_rep = round(sal_ini + S('R') - S('vac_gozada') - S('vac_finiquito'), 2)

    base_emp = S('base', lambda d: d['X'])
    rol_base = round(sum(series._rol_base_mes(rutas.rol).values()), 2)
    sue_rol = round(sum(series.rol_mes_df(rutas.rol)['SUELDO']), 2)
    base_pas = S('base', lambda d: not d['X'])
    estip = round(sum(mayores.gasto_mes('EST')), 2)  # estipendio de pasantes (mayor RQ-004)
    # cifras del mayor de beneficios RQ-004 (si está cargado)
    cxc = ajustes_vac = gasto_solo_mayor = 0.0
    try:
        B = mayores.B
        m = B[(B.cta == '5BS.21204.001') & B['concepto'].astype(str).str.upper().str.contains('CXC|CIERR', regex=True)]
        cxc = round(float(m.Importe.abs().sum()), 2)
        av = B[(B.cta == '5BS.21204.003') & (B['clase'] == 'AJUSTES')]
        ajustes_vac = round(float(av.Importe.abs().sum()), 2)
        ctas_anexo = {c for c, *_ in gas}
        G = mayores.G
        gasto_solo_mayor = round(float(G[~G.cta.isin(ctas_anexo)].Importe.sum()), 2)
    except Exception:
        pass
    # gasto de vacaciones y FR en libros (mayor / anexo) y pasivos D13/D14
    vac_gasto_libros = round(sum(mayores.gasto_mes('VAC')), 2)
    fr_gasto_libros = gasto_concepto('Fondos de Reserva')
    d13_pas_libros = next((abs(s25) for cod, desc, s24, s25 in pas if 'tercer' in desc.lower()), 0.0)
    d14_pas_libros = next((abs(s25) for cod, desc, s24, s25 in pas if 'cuarto' in desc.lower()), 0.0)
    pasivo_d13_rep = S('AJ', lambda d: d['mes'] == 12)
    pasivo_d14_rep = S('AK', lambda d: d['mes'] >= 8)
    d14_gasto = gasto_concepto('Décimo Cuarto Sueldo'); d14_dev = S('Q_bruto')
    sue_gasto = gasto_concepto('Sueldos')
    # saldo de sueldos por pagar: cuenta de pasivo 5BS.21204.004 (o desc 'sueldo'+'pagar')
    sueldo_pend = next((abs(s25) for cod, desc, s24, s25 in pas if cod.strip() == '5BS.21204.004'),
                       round(sum(abs(s25) for cod, desc, s24, s25 in pas if 'sueldo' in desc.lower() and 'pagar' in desc.lower()), 2))

    # --- metas de control de los generadores de referencia (derivadas) ---
    rm_df = series.rol_mes_df(rutas.rol)
    vac_tomadas = round(float(rm_df['VAC TOMADAS'].sum()), 2)
    fr_rol = round(float(rm_df['FOND.RESERV.'].sum()), 2)
    vac_finiquito_rol = round(float(rm_df['VAC.FINIQUIT'].sum()), 2)
    mens_total = round(float(rm_df['D13 MENS'].sum() + rm_df['D14 MENS SIE'].sum()), 2)
    he_gasto = gasto_concepto('Horas Extras'); bonos_gasto = gasto_concepto('Bono Desempeño')
    comisiones_gasto = round(sum(mayores.gasto_mes('COM')), 2)
    premios_gasto = round(sum(mayores.gasto_mes('PRE')), 2)
    base_iess = SC('base'); consolidado_aportes = round(SC('pat') + SC('ind') + SC('ccc'), 2)
    aporte_patronal_total = gasto_concepto('Aporte Patronal')
    aporte_ccc_aud = round(SC('pat') + SC('ccc'), 2)
    pasantes_aporte = S('ind', lambda d: not d['X'])
    d13_gasto = gasto_concepto('Décimo Tercer Sueldo')
    esf_do2 = round(sum(abs(s25) for cod, desc, s24, s25 in pas if cod.strip() in ('5BS.21204.001', '5BS.21204.002', '5BS.21204.003')), 2)
    d13_bruto = S('P_bruto')
    fr_iess = prestamos = 0.0
    if comp is not None and len(comp):
        fr_iess = round(float(comp[comp.tipo == 'Fondos de reserva'].total.sum()), 2)
        prestamos = round(float(comp[comp.tipo.isin(['Préstamo hipotecario', 'Préstamo quirografario'])].total.sum()), 2)

    return {
        'vac_libros': round(vac_libros, 2),
        'vac_reproc': pasivo_vac_rep,
        'vac_exceso': round(vac_libros - pasivo_vac_rep, 2),
        'excluido_fr': S('AU'),
        'rol_menos_iess': round(rol_base - base_emp, 2),
        'd14_gasto': d14_gasto, 'd14_dev': round(d14_dev, 2),
        'd14_exceso': round(d14_gasto - d14_dev, 2),
        'sue_rol': sue_rol, 'sue_gasto': sue_gasto,
        'sue_exceso': round(sue_rol - sue_gasto, 2),
        'v04_total': round(sum(d['ab'] for d in det if d['ab'] > 0 and d['base'] == 0), 2),
        'sueldo_pend': round(sueldo_pend, 2),
        'base_pas': round(base_pas, 2), 'estipendio': round(estip, 2),
        'base_pas_dif': round(base_pas - estip, 2),
        'provision_vac_anual': S('R'),
        'cxc_compensacion': cxc, 'ajustes_vac': ajustes_vac, 'gasto_solo_mayor': gasto_solo_mayor,
        'n_acum_d13': len({d['nom'] for d in det if d['P'] > 0}),
        'n_acum_d14': len({d['nom'] for d in det if d['Q'] > 0}),
        'n_pasantes': len({d['nom'] for d in det if not d['X']}),
        'vac_gasto_libros': vac_gasto_libros, 'fr_gasto_libros': fr_gasto_libros,
        'd13_pas_libros': round(d13_pas_libros, 2), 'd14_pas_libros': round(d14_pas_libros, 2),
        'pasivo_d13_rep': pasivo_d13_rep, 'pasivo_d14_rep': pasivo_d14_rep,
        'd13_pas_dif': round(pasivo_d13_rep - d13_pas_libros, 2),
        'd14_pas_dif': round(pasivo_d14_rep - d14_pas_libros, 2),
        'sal_ini_vac': round(sal_ini, 2),
        'vac_tomadas': vac_tomadas, 'fr_rol': fr_rol, 'vac_finiquito_rol': vac_finiquito_rol,
        'mens_total': mens_total, 'he_gasto': he_gasto, 'bonos_gasto': bonos_gasto,
        'comisiones_gasto': comisiones_gasto, 'premios_gasto': premios_gasto,
        'base_iess': base_iess, 'consolidado_aportes': consolidado_aportes,
        'aporte_patronal_total': aporte_patronal_total, 'aporte_ccc_aud': aporte_ccc_aud,
        'pasantes_aporte': round(pasantes_aporte, 2), 'd13_gasto': d13_gasto,
        'esf_do2': esf_do2, 'd13_bruto': d13_bruto, 'd14_gasto_rq002': d14_gasto,
        'fr_iess': fr_iess, 'prestamos': prestamos, 'rol_base': rol_base,
    }


def _lista(nombres):
    nombres = [n for n in nombres if n]
    if not nombres:
        return ""
    if len(nombres) == 1:
        return nombres[0]
    return ", ".join(nombres[:-1]) + " y " + nombres[-1]


def sujetos(R, rutas) -> dict:
    det = R['det']
    # V04: FR pagado en rol en un mes con base IESS 0
    v04 = collections.defaultdict(float)
    for d in det:
        if d['ab'] > 0 and d['base'] == 0:
            v04[d['nom']] += d['ab']
    v04_nombres = sorted(v04)
    v04_total = round(sum(v04.values()), 2)

    # Excluidos por el auditor (nombres de la hoja de exclusiones)
    exc_ced, exc_nom = motor.leer_exclusiones(rutas.exclusiones) if rutas.exclusiones else ({}, {})
    excluidos = sorted(exc_nom)

    # Código 06 reportado al IESS en meses de pasantía (fuera del rol normal)
    cod06 = sorted({d['nom'] for d in det if d['cod'] == '06' and not d['X']})

    # Pasantes que exceden 6 meses (meses afiliados en la pestaña 'pasantes')
    exceso = []
    try:
        wb = openpyxl.load_workbook(rutas.rol, read_only=True, data_only=True)
        if 'pasantes' in wb.sheetnames:
            meses = collections.defaultdict(set)
            for r in wb['pasantes'].iter_rows(min_row=3, values_only=True):
                if r[1] in (None, ''):
                    continue
                m = motor.mnum(r[0])
                if m:
                    meses[str(r[1]).strip()].add(m)
            exceso = sorted((n, len(ms)) for n, ms in meses.items() if len(ms) > 6)
    except Exception:
        pass

    # Saldo de sueldos por pagar al corte (anexo RQ-002, cuenta de pasivo con 'sueldo')
    sueldo_pend = 0.0
    try:
        from backend.app.aud.niif.procesadores.nomina import cedulas_nuevas
        pas, _ = cedulas_nuevas.leer_anexo(rutas.anexo)
        sueldo_pend = round(sum(abs(s25) for cod, desc, s24, s25 in pas if 'sueldo' in desc.lower()), 2)
    except Exception:
        pass

    return {
        'v04_nombres': _lista(v04_nombres), 'v04_total': v04_total,
        'excluidos': _lista(excluidos),
        'cod06': _lista(cod06),
        'pasante_exceso': _lista([f"{n} ({m})" for n, m in exceso]),
        'pasante_exceso_nombres': _lista([n for n, _ in exceso]),
        'sueldo_pend': sueldo_pend,
    }
