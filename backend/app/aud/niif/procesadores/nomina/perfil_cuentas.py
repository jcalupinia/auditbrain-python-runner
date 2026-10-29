# -*- coding: utf-8 -*-
"""Adaptador de plan de cuentas para la herramienta de nómina (soporte multi-cliente).

El motor fue escrito para el plan de cuentas del encargo de referencia (IESS por pagar
``5BS.21203.*``, beneficios por pagar ``5BS.21204.*``, gastos por concepto). Otros clientes
usan otros códigos (p. ej. ``2.1.7.1.*`` / ``2.1.7.6.*`` pasivos, ``5.2.1.2.1.*`` gastos) y
pueden traer el pasivo en UN solo mayor.

Este módulo traduce los mayores de un cliente a la "forma" que el motor ya entiende,
clasificando CADA cuenta por el NOMBRE de la cuenta (no por su código), de modo que funciona
con distintos planes de cuentas sin cablear cada cliente. Lo que NO reconoce no se clasifica a
ciegas: queda en el reporte ``no_reconocidas`` para que el auditor lo revise.

Qué hace:
- Gasto: fija ``concepto`` por el nombre de la cuenta (→ textos de ``GKEY``), para que
  ``gasto_mes``/``gasto_cta`` del motor funcionen igual.
- Pasivo (un solo mayor): remapea el ``cta`` al código de referencia equivalente y fija
  ``clase`` desde la columna de referencia/observación del cliente
  (provisión/pagos/finiquitos/goces/saldo inicial). Se usa el MISMO df como I y B porque el
  motor filtra por código (``5BS.21203.*`` para IESS, ``5BS.21204.*`` para beneficios).

VERIFICACIÓN (regla de la casa): el mapeo se valida contra el propio anexo del cliente
(saldo anterior ± movimientos = saldo al corte). El desglose fino que el cliente no tagea
(p. ej. GOCES de vacaciones) queda para validación del socio, no se inventa.
"""
import datetime
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook


def _n(s) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


# Nombre de cuenta de GASTO → concepto del motor (texto de GKEY). Orden: específico primero;
# «sueldo» va al final porque es amplio (p. ej. «décimo tercer sueldo» no debe caer en Sueldos).
_GASTO = [
    (("aporte", "patronal"), "Aporte Patronal"),
    (("decimo", "tercer"), "Décimo Tercer Sueldo"),
    (("decimo", "cuarto"), "Décimo Cuarto Sueldo"),
    (("fondo", "reserva"), "Fondos de Reserva"),
    (("vacacion",), "Vacaciones"),
    (("comision",), "Comisiones"),
    (("hora", "extra"), "Horas Extras"),
    (("bono", "familiar"), "Bono Familiar"),
    (("premio",), "Premios"),
    (("estipendio",), "Estipendio Pasantes"),
    (("pasante",), "Estipendio Pasantes"),
    (("bono",), "Bono Desempeño"),
    (("sueldo",), "Sueldos"),
]

# Nombre de cuenta de PASIVO → código de referencia que usa el motor.
_PASIVO = [
    (("decimo", "tercer"), "5BS.21204.001"),
    (("decimo", "cuarto"), "5BS.21204.002"),
    (("vacacion",), "5BS.21204.003"),
    (("sueldo", "pagar"), "5BS.21204.004"),
    (("fondo", "reserva"), "5BS.21203.004"),
    (("quirografario",), "5BS.21203.003"),
    (("hipotecario",), "5BS.21203.005"),
    (("aporte",), "5BS.21203.001"),  # aportes por pagar (patronal + individual, el motor los suma)
]


def _match(nombre, reglas):
    t = _n(nombre)
    for claves, destino in reglas:
        if all(k in t for k in claves):
            return destino
    return None


def _clase(ref) -> str:
    """Categoría de movimiento del motor a partir de la referencia/observación del cliente."""
    r = _n(ref)
    if "provision" in r:
        return "PROVISION"
    if "finiquito" in r:
        return "FINIQUITOS"
    if "goce" in r or ("vac" in r and "tomad" in r):
        return "GOCES"
    if "pago" in r:
        return "PAGOS"
    if "inicial" in r:  # «saldo inicial»: la apertura la toma el motor del anexo, no del mayor
        return "SALDO_INICIAL"
    return "AJUSTES"  # «otros» o desconocido


def _col_nombre(df):
    """Columna con el NOMBRE de la cuenta (no el código ni la glosa del asiento)."""
    for c in df.columns:
        if _n(c) in ("cuenta", "cuenta contable", "nombre de cuenta", "nombre cuenta"):
            return c
    return None


def es_plan_de_referencia(df) -> bool:
    """¿El mayor ya viene en el plan de cuentas del motor (no hay que adaptar)?"""
    cta = df["cta"].astype(str) if "cta" in df.columns else df["Nºcuenta"].astype(str)
    return cta.str.startswith(("5BS.", "5PYG", "5D")).any()


def adaptar_gasto(G):
    """Fija ``concepto`` (y ``mqla``) en el mayor de gasto por el nombre de cada cuenta.
    Devuelve (G, no_reconocidas: list[(cta, nombre)])."""
    col = _col_nombre(G)
    G = G.copy()
    no_rec = []
    if col is None:  # sin columna de nombre no se puede clasificar por nombre
        G["concepto"] = G["desc"].astype(str)
        return G, [("(sin columna de nombre de cuenta)", "")]
    concepto, nombres = [], G[col].astype(str)
    for cta, nom in zip(G["cta"].astype(str), nombres):
        c = _match(nom, _GASTO)
        concepto.append(c if c is not None else nom)
        if c is None and str(nom).strip() and str(nom).lower() != "nan":
            no_rec.append((cta, nom))
    G["concepto"] = concepto
    G["mqla"] = False  # el prefijo 'MQLA' es del cliente de referencia; otros no lo usan
    # cuentas no reconocidas: únicas
    vistos, uniq = set(), []
    for cta, nom in no_rec:
        if cta not in vistos:
            vistos.add(cta); uniq.append((cta, nom))
    return G, uniq


def adaptar_pasivo(P):
    """Remapea el mayor de pasivo combinado al plan del motor y fija ``clase``.
    Devuelve (df_remapeado, no_reconocidas: list[(cta, nombre)])."""
    col = _col_nombre(P)
    P = P.copy()
    no_rec = []
    if col is None:
        raise ValueError("El Mayor de pasivo no trae una columna con el NOMBRE de la cuenta; "
                         "no se puede mapear por nombre.")
    nuevo_cta, nombres = [], P[col].astype(str)
    for cta, nom in zip(P["cta"].astype(str), nombres):
        dest = _match(nom, _PASIVO)
        if dest is None:
            no_rec.append((cta, nom)); nuevo_cta.append(cta)  # se deja tal cual (no se inventa)
        else:
            nuevo_cta.append(dest)
    P["cta"] = nuevo_cta
    P["Nºcuenta"] = nuevo_cta
    # clase desde la columna de referencia/observación del cliente
    ref_col = next((c for c in P.columns if _n(c) in ("referencia", "ref", "explicacion observacion", "observacion")), None)
    P["clase"] = (P[ref_col].map(_clase) if ref_col is not None else "AJUSTES")
    vistos, uniq = set(), []
    for cta, nom in no_rec:
        if cta not in vistos:
            vistos.add(cta); uniq.append((cta, nom))
    return P, uniq


def adaptar(G, P):
    """Adapta un cliente con mayor de gasto (G) y mayor de pasivo combinado (P) al plan del
    motor. Devuelve (G, I, B, no_reconocidas). I y B son el MISMO pasivo remapeado: el motor
    filtra por código (5BS.21203.* IESS, 5BS.21204.* beneficios)."""
    G2, nr_g = adaptar_gasto(G)
    P2, nr_p = adaptar_pasivo(P)
    I = P2.copy()
    B = P2.copy()
    B["concepto"] = B["clase"]  # compat: algún paso leía B.concepto; la clase ya está fijada
    return G2, I, B, {"gasto": nr_g, "pasivo": nr_p}


def _es_fecha(v):
    return isinstance(v, (pd.Timestamp, datetime.date, datetime.datetime))


def reescribir_anexo(anexo_path, destino) -> str:
    """Reescribe el anexo (RQ-002) de un cliente al LAYOUT de referencia que esperan los ~10
    lectores del motor, remapeando SOLO los códigos de pasivo (los de gasto ya empiezan por «5»
    y se clasifican por el nombre). Layout de salida (1-based):
      fila 1: título · fila 2: fechas (col C anterior, col F corte) · fila 3: encabezados ·
      fila 4+: cuentas de pasivo (código de referencia 5BS.212xx), una fila en blanco, gastos.
    Devuelve la ruta del archivo escrito. Las cuentas de pasivo cuyo nombre no se reconoce se
    OMITEN (no se inventan); el motor avisará si falta algún saldo."""
    df = pd.read_excel(anexo_path, header=None)
    ncols = df.shape[1]
    # fila de fechas: la primera (entre las 5 primeras) con una fecha en col C o F
    d_ant = d_cor = None
    for i in range(min(5, len(df))):
        v2 = df.iloc[i, 2] if ncols > 2 else None
        v5 = df.iloc[i, 5] if ncols > 5 else None
        if _es_fecha(v2) or _es_fecha(v5):
            d_ant, d_cor = v2, v5
            break

    pas, gas = [], []
    for _, r in df.iterrows():
        cod = str(r[0]).strip()
        if not cod or not cod[0].isdigit() or cod.count(".") < 2:
            continue  # no es fila de cuenta (título, total, fechas, en blanco)
        nombre = str(r[1]).strip() if ncols > 1 else ""
        s_ant = float(r[2]) if (ncols > 2 and isinstance(r[2], (int, float)) and pd.notna(r[2])) else 0.0
        s_cor = float(r[5]) if (ncols > 5 and isinstance(r[5], (int, float)) and pd.notna(r[5])) else 0.0
        # Pasivo vs gasto SE DECIDE POR EL CÓDIGO, no por el nombre: en el PUC de Ecuador el
        # pasivo empieza por «2» y el gasto/resultado por «5/6/7». (Clave: un gasto como «Aporte
        # patronal IESS» contiene «aporte» y, si se mapeara por nombre, se colaría como pasivo.)
        if cod.startswith("2"):
            ref = _match(nombre, _PASIVO)
            if ref is not None:
                pas.append((ref, nombre, s_ant, s_cor))     # pasivo de nómina → código de referencia
            # un pasivo cuyo nombre no es de nómina se omite (no se inventa)
        elif cod.startswith(("5", "6", "7")):
            gas.append((cod, nombre, s_ant, s_cor))          # gasto/resultado: ya empieza por 5, se deja igual

    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "Anexo de nómina — saldos por cuenta")
    if d_ant is not None:
        ws.cell(2, 3, d_ant)
    if d_cor is not None:
        ws.cell(2, 6, d_cor)
    for col, txt in ((1, "Código"), (2, "Descripción"), (3, "Saldo anterior"), (6, "Saldo corte")):
        ws.cell(3, col, txt)
    fila = 4
    for cod, nom, sa, sc in pas:
        ws.cell(fila, 1, cod); ws.cell(fila, 2, nom); ws.cell(fila, 3, sa); ws.cell(fila, 6, sc)
        fila += 1
    fila += 1  # fila en blanco que separa pasivo de gasto (el motor la preserva)
    for cod, nom, sa, sc in gas:
        ws.cell(fila, 1, cod); ws.cell(fila, 2, nom); ws.cell(fila, 3, sa); ws.cell(fila, 6, sc)
        fila += 1
    wb.save(destino)
    return destino


def reescribir_consolidado(cons_path, destino) -> str:
    """Reescribe el Consolidado del IESS (RQ-003) del cliente al formato que lee el motor:
    hoja ``detallePlanillaConsolidado``, datos desde la fila 3, con los campos CORRIDOS a col B
    (el motor lee ``mes`` en col B, …, ``valor CCC`` en la col 13 y ``total aporte`` en la 14).
    Vale para el export estándar del IESS en una sola hoja con encabezados arriba. Devuelve la
    ruta escrita. NB: el export del IESS ya trae patronal/individual/CCC/total por cédula y mes;
    solo se reubican las columnas (no se recalcula nada)."""
    import openpyxl

    ws_in = openpyxl.load_workbook(cons_path, read_only=True, data_only=True).active
    filas = list(ws_in.iter_rows(values_only=True))
    # fila de encabezados: la primera (de las 5 primeras) con «cédula»/«periodo»
    hi = 0
    for i, r in enumerate(filas[:5]):
        celdas = [_n(c) for c in (r or []) if c is not None]
        if any("cedula" in c for c in celdas) or any("periodo" in c for c in celdas):
            hi = i
            break
    wb = Workbook()
    ws = wb.active
    ws.title = "detallePlanillaConsolidado"
    ws.cell(1, 1, "Planilla consolidada IESS (adaptada al motor)")
    for j, h in enumerate(filas[hi] or []):          # encabezados en la fila 2, corridos a col B
        ws.cell(2, j + 2, h)
    r = 3
    for fila in filas[hi + 1:]:
        if not fila or len(fila) < 2 or fila[1] is None or str(fila[1]).strip() == "":
            continue  # fila sin cédula (total/vacía)
        for j, v in enumerate(fila):
            ws.cell(r, j + 2, v)
        r += 1
    wb.save(destino)
    return destino


# --- Rol de pagos (RQ-006) -------------------------------------------------------------------
# El motor lee el rol de la hoja «normal» (empleados + jubilados) con un layout de 67 columnas en
# posiciones exactas. Otros clientes traen el rol en otra hoja y otro layout (secciones Datos
# Generales / Ingresos / Egresos / Provisiones…). `reescribir_rol` mapea por SECCIÓN + NOMBRE de
# campo (robusto) al layout del motor. Vacaciones GOZADAS y FINIQUITOS no suelen venir en el rol
# del cliente: se dejan en 0 y quedan PENDIENTES DE VALIDACIÓN DEL SOCIO (decisión del dueño), no
# se inventan. Admite rol de un mes (se pasa `mes`) o combinado (columna de período por fila).

# Posición (0-based) de cada campo en la hoja «normal» que lee motor.leer_rol_normal.
_ROL_POS = {"mes": 0, "nombre": 1, "cargo": 2, "ingreso": 3, "salida": 4, "sueldo": 5,
            "vac_gozada": 6, "vac_finiquito": 18, "fr_rol": 24, "d13mens": 25, "d14mens": 26,
            "prov_d13": 63, "prov_d14": 64, "prov_vac": 65, "prov_fr": 66}


# Campos del rol que el motor necesita de verdad para las cédulas (los demás son opcionales).
_ROL_CAMPOS_CLAVE = ("nombre", "sueldo", "fr_rol", "d13mens", "d14mens",
                     "prov_d13", "prov_d14", "prov_vac", "prov_fr")


def _mapa_rol(filas):
    """Localiza la fila de encabezados y mapea cada dato del motor a su columna en el rol del
    cliente, de forma GENÉRICA (los roles varían). Usa la SECCIÓN (fila superior, si existe) para
    desambiguar los nombres repetidos (p. ej. «Fondos de reserva» aparece en Ingresos, Egresos y
    Provisiones); si no hay secciones, cae a nombre + heurística. Devuelve (hi, campos, faltantes):
    faltantes = campos clave que NO se reconocieron (se reportan; no se inventan)."""
    hi = 0
    for i, r in enumerate(filas[:12]):
        celdas = [_n(c) for c in (r or []) if c is not None]
        tiene_id = any("cedula" in c for c in celdas) or any(c in ("empleado", "nombre") for c in celdas)
        if tiene_id and any(c == "sueldo" for c in celdas):
            hi = i
            break
    sec_row = filas[hi - 1] if hi > 0 else [None] * len(filas[hi])
    secciones, actual = [], ""
    for s in sec_row:
        actual = _n(s) if s not in (None, "") else actual
        secciones.append(actual)
    heads = [_n(h) if h is not None else "" for h in filas[hi]]

    def buscar(*, nombre, seccion=None, excluir_sec=(), excluir_nom=()):
        """Mejor columna: primero con la sección pedida; si no, por nombre evitando secciones
        a excluir (p. ej. provisión/egreso para un ingreso). `nombre` es una función(h)->bool."""
        if seccion is not None:
            for j, h in enumerate(heads):
                sec = secciones[j] if j < len(secciones) else ""
                if h and nombre(h) and seccion in sec:
                    return j
        for j, h in enumerate(heads):
            sec = secciones[j] if j < len(secciones) else ""
            if h and nombre(h) and not any(e in sec for e in excluir_sec) \
                    and not any(e in h for e in excluir_nom):
                return j
        return None

    campos = {}
    campos["nombre"] = buscar(nombre=lambda h: h in ("empleado", "nombre", "nombres", "trabajador",
                              "apellidos y nombres", "nombres y apellidos")
                              or "apellidos y nombres" in h or "nombres y apellidos" in h
                              or (("apellido" in h or "nombre" in h) and "fecha" not in h))
    campos["cargo"] = buscar(nombre=lambda h: h == "cargo")
    campos["ingreso"] = buscar(nombre=lambda h: "fecha de registro" in h or "fecha de ingreso" in h or h == "ingreso")
    # Ingresos (lo pagado en el mes): preferir la sección «ingreso», evitando egresos/provisiones.
    campos["sueldo"] = buscar(nombre=lambda h: h == "sueldo", seccion="ingreso",
                              excluir_sec=("egreso", "provision"), excluir_nom=("variable",))
    campos["fr_rol"] = buscar(nombre=lambda h: "fondos" in h and "reserva" in h, seccion="ingreso",
                              excluir_sec=("egreso", "provision"))
    campos["d13mens"] = buscar(nombre=lambda h: "decimo tercero" in h or "decimo tercer" in h,
                               seccion="ingreso", excluir_sec=("provision",))
    campos["d14mens"] = buscar(nombre=lambda h: "decimo cuarto" in h, seccion="ingreso",
                               excluir_sec=("provision",))
    # Provisiones de beneficios sociales.
    campos["prov_d13"] = buscar(nombre=lambda h: "decimo tercer" in h or "decimo tercero" in h, seccion="provision")
    campos["prov_d14"] = buscar(nombre=lambda h: "decimo cuarto" in h, seccion="provision")
    campos["prov_vac"] = buscar(nombre=lambda h: "vacacion" in h, seccion="provision")
    campos["prov_fr"] = buscar(nombre=lambda h: "fondos" in h and "reserva" in h, seccion="provision")
    campos["mes"] = buscar(nombre=lambda h: h in ("mes", "periodo", "período"))
    faltantes = [c for c in _ROL_CAMPOS_CLAVE if campos.get(c) is None]
    return hi, campos, faltantes


# Campos del rol no reconocidos en el último reescribir_rol (para el reporte al auditor).
ROL_CAMPOS_FALTANTES: list = []


_MESES_NOMBRE = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
                 "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
                 "noviembre": 11, "diciembre": 12}


def mes_de_nombre(nombre) -> int | None:
    """Número de mes a partir del NOMBRE del archivo del rol (p. ej. «Periodo_enero_S_2» → 1).
    Busca el nombre del mes en español (no un número suelto como «S_2», que sería ambiguo)."""
    t = _n(nombre)
    for nom, num in _MESES_NOMBRE.items():
        if nom in t:
            return num
    return None


def _filas_rol(rol_path, mes=None):
    """Filas de un rol del cliente en el layout de 67 columnas del motor (una por empleado).
    Devuelve (filas, faltantes). `mes` = nº de mes cuando el rol es de un período; si el rol trae
    columna de período, se usa esa por fila. Vacaciones gozadas y finiquitos quedan en 0."""
    import openpyxl

    ws_in = openpyxl.load_workbook(rol_path, read_only=True, data_only=True).active
    filas = list(ws_in.iter_rows(values_only=True))
    hi, campos, faltantes = _mapa_rol(filas)
    if campos["nombre"] is None or campos["sueldo"] is None:
        raise ValueError("No reconozco el rol: no encuentro las columnas de empleado y de sueldo. "
                         "El adaptador mapea por nombre de campo (y sección si existe); revise que "
                         "el rol traiga al menos columnas de empleado y sueldo.")
    from backend.app.aud.niif.procesadores.nomina.motor import mnum

    def val(fila, campo):
        j = campos.get(campo)
        return fila[j] if (j is not None and j < len(fila)) else None

    out = []
    for fila in filas[hi + 1:]:
        nombre = val(fila, "nombre")
        if nombre is None or str(nombre).strip() == "" or str(nombre).strip().lower().startswith("total"):
            continue
        m = mnum(val(fila, "mes")) if campos["mes"] is not None else mes
        fila_out = [None] * 67
        fila_out[_ROL_POS["mes"]] = m if m is not None else 1
        fila_out[_ROL_POS["nombre"]] = str(nombre).strip()
        for campo in ("cargo", "ingreso", "sueldo", "fr_rol", "d13mens", "d14mens",
                      "prov_d13", "prov_d14", "prov_vac", "prov_fr"):
            v = val(fila, campo)
            if v is not None:
                fila_out[_ROL_POS[campo]] = v
        # vac_gozada / vac_finiquito: PENDIENTES DEL SOCIO → 0 (no se inventan)
        out.append(fila_out)
    return out, faltantes


def _escribir_rol(filas_rol, destino) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "normal"
    for c in range(1, 68):
        ws.cell(5, c, f"c{c}")  # encabezados (el motor lee desde la fila 6)
    r = 6
    for fila_out in filas_rol:
        for c, v in enumerate(fila_out, start=1):
            if v is not None:
                ws.cell(r, c, v)
        r += 1
    # Hoja «pasantes» (este cliente no la separa): con fila de encabezado para que los lectores
    # que leen la fila 2 (series.rol_mes_df) no fallen; sin empleados.
    wp = wb.create_sheet("pasantes")
    wp.cell(1, 1, "pasantes")
    wp.cell(2, 1, "Mes Nº")
    wb.save(destino)
    return destino


def reescribir_rol(rol_path, destino, mes=None) -> str:
    """Reescribe UN rol del cliente a la hoja «normal» del layout del motor."""
    global ROL_CAMPOS_FALTANTES
    filas, ROL_CAMPOS_FALTANTES = _filas_rol(rol_path, mes)
    return _escribir_rol(filas, destino)


def reescribir_roles(items, destino) -> str:
    """Combina VARIOS roles mensuales del cliente en una sola hoja «normal» (una fila por empleado
    y mes). `items` = lista de (ruta, mes). El mes de cada archivo lo decide el llamador (p. ej.
    del nombre del archivo con `mes_de_nombre`)."""
    global ROL_CAMPOS_FALTANTES
    todas, faltan = [], []
    for ruta, mes in items:
        filas, f = _filas_rol(ruta, mes)
        todas.extend(filas)
        faltan = f  # misma estructura mes a mes; basta el último reporte
    ROL_CAMPOS_FALTANTES = faltan
    return _escribir_rol(todas, destino)
