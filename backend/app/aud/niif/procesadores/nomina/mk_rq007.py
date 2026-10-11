# -*- coding: utf-8 -*-
from backend.app.aud.niif.procesadores.nomina import rutas as _R
"""RQ-007 — Muestreo de solicitudes de acumulación de décimos (D13 y D14) — DO-4 V08.
Desde 2015 los décimos se pagan mensualizados salvo solicitud escrita del trabajador para acumularlos (arts. 111 y 113 CT).
Población: trabajadores que acumulan D13 (DO-4 Detalle RQ-003): 147, de los cuales 146 acumulan también D14.
Muestra única para ambos certificados: % configurable (10% por defecto) = partidas clave (mayores D13) + selección sistemática."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (_R.CTX.out + '/RQ-007_Muestreo_Solicitudes_Decimos.xlsx')
pop = pd.read_pickle((_R.CTX.home + '/pop_acum.pkl')).sort_values('Nombre').reset_index(drop=True)
N = len(pop); MAXS = 60

NAVY = PatternFill('solid', fgColor='FF0F2341'); YEL = PatternFill('solid', fgColor='FFFFF2CC'); INP = PatternFill('solid', fgColor='FFFFFFE0')
SELF = PatternFill('solid', fgColor='FFE2EFDA')
AR = Font(name='Arial', size=10); ARB = Font(name='Arial', size=10, bold=True)
WH = Font(name='Arial', size=10, bold=True, color='FFFFFFFF'); WH12 = Font(name='Arial', size=12, bold=True, color='FFFFFFFF')
BLUE = Font(name='Arial', size=10, color='FF0000FF'); GREEN = Font(name='Arial', size=10, color='FF008000')
RED = Font(name='Arial', size=10, bold=True, color='FFC00000'); IT = Font(name='Arial', size=8, italic=True)
thin = Side(style='thin', color='FF000000'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N2 = '#,##0.00;(#,##0.00);"-"'; DT = 'dd/mm/yyyy'
C = Alignment(horizontal='center', vertical='center', wrap_text=True); W = Alignment(wrap_text=True, vertical='top')

def cell(ws, ref, v, font=AR, fmt=None, fill=None, al=None, box=False):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if box: c.border = BOX
    return c

def th(ws, row, col0, labels, height=40):
    for i, t in enumerate(labels):
        c = ws.cell(row, col0 + i, t); c.font = WH; c.fill = NAVY; c.alignment = C; c.border = BOX
    ws.row_dimensions[row].height = height

wb = Workbook()
wf = wb.active; wf.title = 'Ficha'
cell(wf, 'B2', 'FICHA DEL ENCARGO Y PARÁMETROS DEL MUESTREO', ARB)
pars = [('Cliente', _R.CTX.cliente, None, BLUE), ('Año auditado', _R.CTX.ejercicio, '0', BLUE), ('Fecha de corte', '=DATE(C5,12,31)', DT, AR),
        ('Preparado por', _R.CTX.preparado, None, BLUE), ('Revisado por', _R.CTX.revisado, None, BLUE), ('Fecha', '=TODAY()', DT, AR),
        ('Tamaño de muestra (% de la población)', 0.10, '0%', BLUE),
        ('Partidas clave: mayores importes de D13 (n.º)', 5, '0', BLUE),
        ('Arranque aleatorio de la selección sistemática (1 a intervalo)', 2, '0', BLUE)]
for i, (k, v, f, fo) in enumerate(pars):
    cell(wf, f'B{4+i}', k, ARB, box=True, al=W); cell(wf, f'C{4+i}', v, fo, f, box=True)
cell(wf, 'B14', 'Población del año auditado 2026; datos de la plantilla 2025 (DO-4 Detalle RQ-003). El auditor puede cambiar el %, las partidas clave y el arranque.', IT)
wf.column_dimensions['B'].width = 60; wf.column_dimensions['C'].width = 16
PCT, KEY, ST = 'Ficha!$C$10', 'Ficha!$C$11', 'Ficha!$C$12'

# ---- Población
wp = wb.create_sheet('Población')
cell(wp, 'B2', 'POBLACIÓN — trabajadores que acumulan D13 (no mensualizan) · fuente DO-4 Detalle RQ-003', ARB)
th(wp, 4, 2, ['N°', 'Cédula', 'Nombre', 'D13 acumulado', 'D14 acumulado', 'Meses', 'Activo al corte', 'Rango D13', '¿Clave?',
              'Posición entre no clave', '¿Sistemática?', '¿Seleccionado?', 'N.º en muestra'], 44)
P0, P1 = 5, 4 + N
for i, w in pop.iterrows():
    r = P0 + i
    cell(wp, f'B{r}', i + 1, AR, box=True); cell(wp, f'C{r}', w['Cédula'], AR, box=True); cell(wp, f'D{r}', w['Nombre'], AR, box=True)
    cell(wp, f'E{r}', round(float(w['d13']), 2), GREEN, N2, box=True); cell(wp, f'F{r}', round(float(w['d14']), 2), GREEN, N2, box=True)
    cell(wp, f'G{r}', int(w['meses']), AR, '0', box=True); cell(wp, f'H{r}', w['Activo al 31-dic'], AR, box=True)
    cell(wp, f'I{r}', f'=RANK(E{r},$E${P0}:$E${P1},0)+COUNTIF($E${P0}:E{r},E{r})-1', AR, '0', box=True)
    cell(wp, f'J{r}', f'=IF(I{r}<={KEY},"Sí","No")', AR, box=True)
    cell(wp, f'K{r}', f'=IF(J{r}="No",COUNTIF($J${P0}:J{r},"No"),"")', AR, '0', box=True)
    cell(wp, f'L{r}', f'=IF(J{r}="No",IF(COUNTIF(Parámetros!$C$12:$C${11+MAXS},K{r})>0,"Sí","No"),"No")', AR, box=True)
    cell(wp, f'M{r}', f'=IF(OR(J{r}="Sí",L{r}="Sí"),"Sí","No")', ARB, box=True)
    cell(wp, f'N{r}', f'=IF(M{r}="Sí",COUNTIF($M${P0}:M{r},"Sí"),"")', AR, '0', box=True)
cell(wp, f'D{P1+1}', 'TOTAL', RED, box=True)
for col in 'EF': cell(wp, f'{col}{P1+1}', f'=SUM({col}{P0}:{col}{P1})', ARB, N2, box=True)
for k, v in {'A': 2, 'B': 5, 'C': 12, 'D': 38, 'E': 13, 'F': 13, 'G': 7, 'H': 11, 'I': 8, 'J': 8, 'K': 10, 'L': 11, 'M': 12, 'N': 10}.items():
    wp.column_dimensions[k].width = v
wp.freeze_panes = 'E5'

# ---- Parámetros de selección (cálculo del tamaño y posiciones sistemáticas)
wq = wb.create_sheet('Parámetros')
cell(wq, 'B2', 'CÁLCULO DEL TAMAÑO DE MUESTRA Y POSICIONES SISTEMÁTICAS', ARB)
calc = [('Población (N)', f'=COUNTA(Población!C{P0}:C{P1})', '0'),
        ('Población que también acumula D14', f'=COUNTIF(Población!F{P0}:F{P1},">0")', '0'),
        ('Tamaño de muestra n = REDONDEAR.MAS(N × %)', f'=MIN(ROUNDUP(C4*{PCT},0),{MAXS})', '0'),
        ('Partidas clave (mayores D13)', f'=MIN({KEY},C6)', '0'),
        ('Selección sistemática (n − clave)', '=C6-C7', '0'),
        ('Población remanente (N − clave)', '=C4-C7', '0'),
        ('Intervalo de selección', '=IF(C8>0,C9/C8,0)', '0.00')]
for i, (k, f, fm) in enumerate(calc):
    cell(wq, f'B{4+i}', k, ARB, box=True); cell(wq, f'C{4+i}', f, AR, fm, box=True)
th(wq, 11, 2, ['j', 'Posición sistemática'], 30)
for j in range(1, MAXS + 1):
    r = 11 + j
    cell(wq, f'B{r}', j, AR, '0', box=True)
    cell(wq, f'C{r}', f'=IF({j}<=$C$8,MIN($C$9,ROUND(MIN({ST},$C$10)+({j}-1)*$C$10,0)),"")', AR, '0', box=True)
wq.column_dimensions['B'].width = 46; wq.column_dimensions['C'].width = 18

# ---- Cédula de muestreo
ws = wb.create_sheet('DO-4 V08 Muestreo', 0)
for r in range(1, 4):
    for col in range(1, 13): ws.cell(r, col).fill = NAVY
ws.row_dimensions[1].height = 34; ws.row_dimensions[2].height = 22
img = XLImage(_R.CTX.logo); img.height = 60; img.width = 160; ws.add_image(img, 'A1')
cell(ws, 'F1', 'OBLIGACIONES ACUMULADAS', WH12); cell(ws, 'F3', 'Solicitudes de acumulación de décimos — muestreo de atributos', WH)
cell(ws, 'A4', 'Nombre del cliente:', ARB, box=True); cell(ws, 'C4', '=Ficha!C4', AR, box=True)
cell(ws, 'F4', 'Período terminado:', ARB, box=True); cell(ws, 'H4', '=Ficha!C6', AR, DT, box=True)
cell(ws, 'J4', 'Año auditado:', ARB, box=True); cell(ws, 'K4', '=Ficha!C5', RED, '0', box=True)
cell(ws, 'A5', 'Preparado por:', ARB, box=True); cell(ws, 'C5', '=Ficha!C7', AR, box=True)
cell(ws, 'D5', 'Revisado:', ARB, box=True); cell(ws, 'E5', '=Ficha!C8', AR, box=True)
cell(ws, 'F5', 'Fecha:', ARB, box=True); cell(ws, 'H5', '=Ficha!C9', AR, DT, box=True)
cell(ws, 'J5', 'Referencia de PT', ARB, box=True); cell(ws, 'K5', 'DO-4 V08', RED, box=True)
ws.merge_cells('A6:L6'); b = ws['A6']
b.value = ('Muestra única: las mismas personas se verifican para el certificado de acumulación de D13 y el de D14. Celdas amarillas: el auditor marca '
           'Sí / No según el documento físico (RQ-007). N/A en D14 cuando el trabajador no acumula D14.')
b.font = Font(name='Arial', size=9, bold=True, color='FFC00000'); b.fill = YEL; b.alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[6].height = 28

res = [('Población (trabajadores que acumulan D13)', '=Parámetros!C4', '0'),
       ('… de los cuales acumulan también D14', '=Parámetros!C5', '0'),
       ('Tamaño de muestra (% de Ficha)', '=Parámetros!C6', '0'),
       ('Partidas clave + sistemáticas', '=Parámetros!C7&" + "&Parámetros!C8', None),
       ('Documentos revisados', f'=COUNTIF(F22:G{21+MAXS},"Sí")+COUNTIF(F22:G{21+MAXS},"No")', '0'),
       ('Excepciones D13 (sin solicitud)', f'=COUNTIF(F22:F{21+MAXS},"No")', '0'),
       ('Excepciones D14 (sin solicitud)', f'=COUNTIF(G22:G{21+MAXS},"No")', '0'),
       ('Tasa de desviación de la muestra', f'=IF(E14=0,0,(E15+E16)/E14)', '0.0%'),
       ('Pendientes de marcar', f'=E12*2-COUNTIF(G22:G{21+MAXS},"N/A")-E14', '0'),
       ('Conclusión', '=IF(E18>0,"PENDIENTE — faltan "&E18&" marcas del auditor",IF(E15+E16=0,"CONFORME — todas las personas de la muestra tienen solicitud de acumulación de D13 y D14","CON EXCEPCIONES — "&(E15+E16)&" solicitudes faltantes: comunicar en DO-12 H-8 y evaluar ampliar la muestra"))', None)]
cell(ws, 'A8', 'RESUMEN', WH, fill=NAVY); ws.merge_cells('A8:E8')
for i, (t, f, fm) in enumerate(res):
    r = 10 + i
    ws.merge_cells(f'A{r}:D{r}'); cell(ws, f'A{r}', t, ARB if i == 9 else AR, box=True)
    cell(ws, f'E{r}', f, RED if i in (5, 6, 9) else AR, fm, box=True)
ws.merge_cells('E19:L19')
th(ws, 21, 1, ['N.º', 'Cédula', 'Nombre', 'D13 acumulado', 'D14 acumulado', 'Solicitud D13 (Sí/No)', 'Solicitud D14 (Sí/No/N/A)',
               'Tipo de selección', 'Fecha del documento', 'Firma del trabajador conforme', 'Observación', 'Marca'], 44)
dv = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True); ws.add_data_validation(dv)
dv2 = DataValidation(type='list', formula1='"Sí,No,N/A"', allow_blank=True); ws.add_data_validation(dv2)
PR = f'Población!$N${P0}:$N${P1}'
for j in range(1, MAXS + 1):
    r = 21 + j
    m = f'MATCH({j},{PR},0)'
    cell(ws, f'A{r}', f'=IF({j}<=$E$12,{j},"")', AR, '0', box=True)
    cell(ws, f'B{r}', f'=IF(A{r}="","",INDEX(Población!$C${P0}:$C${P1},{m}))', GREEN, box=True)
    cell(ws, f'C{r}', f'=IF(A{r}="","",INDEX(Población!$D${P0}:$D${P1},{m}))', GREEN, box=True)
    cell(ws, f'D{r}', f'=IF(A{r}="","",INDEX(Población!$E${P0}:$E${P1},{m}))', GREEN, N2, box=True)
    cell(ws, f'E{r}', f'=IF(A{r}="","",INDEX(Población!$F${P0}:$F${P1},{m}))', GREEN, N2, box=True)
    cell(ws, f'F{r}', None, BLUE, fill=INP, box=True, al=Alignment(horizontal='center'))
    cell(ws, f'G{r}', None, BLUE, fill=INP, box=True, al=Alignment(horizontal='center'))
    cell(ws, f'H{r}', f'=IF(A{r}="","",IF(INDEX(Población!$J${P0}:$J${P1},{m})="Sí","Clave (mayor D13)","Sistemática"))', AR, box=True)
    for col in 'IJK': cell(ws, f'{col}{r}', None, BLUE, DT if col == 'I' else None, INP, box=True)
    cell(ws, f'L{r}', f'=IF(A{r}="","",IF(OR(F{r}="No",G{r}="No"),"‡",IF(F{r}="Sí","✓","")))', RED, box=True, al=Alignment(horizontal='center'))
    dv.add(f'F{r}'); dv2.add(f'G{r}')
t = 21 + MAXS + 2
txt = [('Objetivo:', ARB),
       ('Verificar que los trabajadores que reciben el décimo tercero y el décimo cuarto acumulados (no mensualizados) cuentan con la solicitud escrita '
        'de acumulación exigida por los arts. 111 y 113 del Código del Trabajo, con prueba de atributos sobre una muestra.', AR),
       ('Procedimiento:', ARB),
       ('1. Población: trabajadores que acumulan D13 según el reproceso DO-4 (Detalle RQ-003), después de exclusiones.', AR),
       ('2. Muestra única para ambos décimos: % configurable en Ficha (10% por defecto), con partidas clave de mayor D13 y selección sistemática con arranque aleatorio.', AR),
       ('3. Se solicita al cliente la solicitud de acumulación de cada persona seleccionada (RQ-007) y se marca su existencia para D13 y D14.', AR),
       ('4. Excepción = persona sin solicitud. Con excepciones se comunica en DO-12 (H-8) y se evalúa ampliar la muestra o extender a la población.', AR),
       ('Marcas:  ✓ Solicitud verificada · ‡ Excepción (sin solicitud).', AR)]
for i, (s, f) in enumerate(txt):
    ws.merge_cells(f'A{t+i}:L{t+i}'); cell(ws, f'A{t+i}', s, f, al=W)
    if len(s) > 120: ws.row_dimensions[t + i].height = 28
for k, v in {'A': 6, 'B': 12, 'C': 36, 'D': 13, 'E': 13, 'F': 12, 'G': 13, 'H': 17, 'I': 12, 'J': 14, 'K': 28, 'L': 7}.items():
    ws.column_dimensions[k].width = v
ws.sheet_view.showGridLines = False; ws.freeze_panes = 'D22'
ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

# ---- Controles
wc = wb.create_sheet('Controles')
cell(wc, 'B2', 'CONTROLES DE INTEGRIDAD', ARB)
th(wc, 4, 2, ['Control', 'Valor A', 'Valor B', 'Diferencia', 'Estado'], 30)
V = "'DO-4 V08 Muestreo'"
ctl = [('Población = trabajadores que acumulan D13 (DO-4: 147)', '=Parámetros!C4', 147, 0),
       ('Acumulan también D14 (DO-4: 146)', '=Parámetros!C5', 146, 0),
       ('Seleccionados en Población = tamaño de muestra', f'=COUNTIF(Población!M{P0}:M{P1},"Sí")', '=Parámetros!C6', 0),
       ('Filas listadas en la cédula = tamaño de muestra', f'=COUNT({V}!A22:A{21+MAXS})', '=Parámetros!C6', 0),
       ('Posiciones sistemáticas sin duplicados', '=SUMPRODUCT((Parámetros!C12:C71<>"")/COUNTIF(Parámetros!C12:C71,Parámetros!C12:C71&""))', '=Parámetros!C8', 0)]
for i, (tt, a, b2, tol) in enumerate(ctl):
    r = 5 + i
    cell(wc, f'B{r}', tt, AR, box=True); cell(wc, f'C{r}', a, GREEN, '0', box=True); cell(wc, f'D{r}', b2, GREEN, '0', box=True)
    cell(wc, f'E{r}', f'=ROUND(C{r}-D{r},2)', AR, '0', box=True)
    cell(wc, f'F{r}', f'=IF(ABS(E{r})<={tol},"OK","REVISAR")', ARB, box=True, al=Alignment(horizontal='center'))
n = len(ctl); cell(wc, f'B{6+n}', f'=COUNTIF(F5:F{4+n},"OK")&" de {n} controles OK"', ARB)
wc.column_dimensions['B'].width = 56
for col in 'CDEF': wc.column_dimensions[col].width = 14
wb.active = 0
wb.save(OUT); print('ok', OUT, N)
