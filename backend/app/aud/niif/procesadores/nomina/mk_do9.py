# -*- coding: utf-8 -*-
from backend.app.aud.niif.procesadores.nomina import rutas as _R
"""DO-9 IESS por pagar: libros vs comprobantes de pago (RQ-005) vs Consolidado (RQ-003), con saldos por pagar a diciembre.
Formato: pestaña «DO9 Provisiones Ben vs Comp. Pa» del papel de referencia."""
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

OUT = (_R.CTX.out + '/DO-9_IESS_por_Pagar_Comprobantes.xlsx')
SRC = _R.CTX.papel
comp = pd.read_pickle((_R.CTX.home + '/comprobantes.pkl')).sort_values(['mes', 'tipo', 'comprobante']).reset_index(drop=True)
cons = pd.read_pickle((_R.CTX.home + '/consol.pkl'))
MESES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
ANEXO = {str(r[0]).strip(): abs(r[5]) for r in load_workbook(_R.CTX.anexo, data_only=True).active.iter_rows(min_row=4, values_only=True)
         if r[0] and isinstance(r[5], (int, float))}

from backend.app.aud.niif.procesadores.nomina import mayores as MY
_I = MY.I.copy(); _I['n'] = _I['Nombre alfa explicación'].astype(str).str.upper()
def _sub(r):
    if r.clase != 'PAGOS': return r.clase
    if 'CANCELACION' in r.n: return 'PAGO IESS'
    if 'CIERR.CXC' in r.n: return 'NC'
    return 'OTROS PAGOS'
_I['sub'] = _I.apply(_sub, axis=1)
def _mes(ctas, subs):
    s = _I[_I.cta.isin(ctas) & _I['sub'].isin(subs)].groupby('m').Importe.sum()
    return [round(-float(s.get(m, 0)), 2) for m in range(1, 13)]
APO_C = ['5BS.21203.001', '5BS.21203.002']; PRE_C = ['5BS.21203.003', '5BS.21203.005']
lib = {'APO': _mes(APO_C, ['PROVISION', 'FINIQUITOS']), 'FR': _mes(['5BS.21203.004'], ['PROVISION']), 'PRE': _mes(PRE_C, ['PROVISION', 'FINIQUITOS'])}
# pagos al IESS por período (mes de pago m → período m−1)
pag = {k: [round(float(_I[_I.cta.isin(c) & (_I['sub'] == 'PAGO IESS') & (_I.m == m)].Importe.sum()), 2) for m in range(1, 13)]
       for k, c in (('APO', APO_C), ('FR', ['5BS.21203.004']), ('QUI', ['5BS.21203.003']), ('HIP', ['5BS.21203.005']))}

NAVY = PatternFill('solid', fgColor='FF0F2341'); SUB = PatternFill('solid', fgColor='FFDDE6F3'); YEL = PatternFill('solid', fgColor='FFFFF2CC')
AR = Font(name='Arial', size=10); ARB = Font(name='Arial', size=10, bold=True)
WH = Font(name='Arial', size=10, bold=True, color='FFFFFFFF'); WH12 = Font(name='Arial', size=12, bold=True, color='FFFFFFFF')
BLUE = Font(name='Arial', size=10, color='FF0000FF'); GREEN = Font(name='Arial', size=10, color='FF008000'); ORANGE = Font(name='Arial', size=10, color='FFC65911')
RED = Font(name='Arial', size=10, bold=True, color='FFC00000'); RU = Font(name='Arial', size=10, bold=True, color='FFC00000', underline='single')
NAVYF = Font(name='Arial', size=10, bold=True, color='FF1F3864'); IT = Font(name='Arial', size=8, italic=True)
thin = Side(style='thin', color='FF000000'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N2 = '#,##0.00;(#,##0.00);"-"'
C = Alignment(horizontal='center', vertical='center', wrap_text=True)

def cell(ws, ref, v, font=AR, fmt=None, fill=None, al=None, box=False):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if box: c.border = BOX
    return c

def th(ws, row, col0, labels, height=30):
    for i, t in enumerate(labels):
        c = ws.cell(row, col0 + i, t); c.font = WH; c.fill = NAVY; c.alignment = C; c.border = BOX
    ws.row_dimensions[row].height = height

wb = Workbook()
wf = wb.active; wf.title = 'Ficha'
cell(wf, 'B2', 'FICHA DEL ENCARGO (encabezado de las cédulas)', ARB)
for i, (k, v) in enumerate([('Cliente', _R.CTX.cliente), ('Período terminado', pd.Timestamp(_R.CTX.corte).to_pydatetime()),
                            ('Preparado por', _R.CTX.preparado), ('Revisado por', _R.CTX.revisado), ('Fecha', '=TODAY()')]):
    cell(wf, f'B{4+i}', k, ARB, box=True); c = cell(wf, f'C{4+i}', v, BLUE if i < 4 else AR, box=True)
    if i in (1, 4): c.number_format = 'dd/mm/yyyy'
wf.column_dimensions['B'].width = 22; wf.column_dimensions['C'].width = 30

# ---------------- Comprobantes RQ-005 (detalle)
wcmp = wb.create_sheet('Comprobantes RQ-005')
cell(wcmp, 'B2', 'COMPROBANTES DE PAGO IESS — extracción de los PDF del RQ-005 (85 comprobantes)', ARB)
cols = ['Archivo', 'Tipo', 'N.º comprobante', 'Concepto', 'Mes', 'Emitido', 'Vigencia de pago', 'Valor', 'Intereses mora',
        'Seguro salud T. parcial', 'Honorarios', 'Gastos adm.', 'Subtotal', 'Notas de crédito', 'Total', 'Días de margen', 'V10 Plazo']
th(wcmp, 4, 2, cols, 40)
for i, rw in comp.iterrows():
    r = 5 + i
    vals = [rw.archivo, rw.tipo, rw.comprobante, rw.concepto, int(rw.mes), pd.Timestamp(rw.emitido).to_pydatetime(), pd.Timestamp(rw.vigencia).to_pydatetime(),
            rw.valor, rw.intereses, rw.seguro_tp, rw.honorarios, rw.gastos_adm, rw.subtotal, rw.notas_credito, rw.total]
    for j, v in enumerate(vals):
        c = wcmp.cell(r, 2 + j, v); c.font = GREEN if j >= 7 else AR; c.border = BOX
        if j in (5, 6): c.number_format = 'dd/mm/yyyy'
        if j >= 7: c.number_format = N2
    wcmp.cell(r, 17, f'=H{r}-G{r}').font = AR; wcmp.cell(r, 17).border = BOX
    wcmp.cell(r, 18, f'=IF(G{r}>H{r},"FUERA DE PLAZO","OK")').font = ARB; wcmp.cell(r, 18).border = BOX
    wcmp.cell(r, 19, f'=ROUND(N{r}-(I{r}+J{r}+K{r}+L{r}+M{r}),2)+ROUND(P{r}-(N{r}-O{r}),2)').font = IT
NC = 4 + len(comp)
cell(wcmp, f'S4', 'Control aritmético', WH, fill=NAVY, al=C, box=True)
cell(wcmp, f'B{NC+1}', 'Σ', RED)
for col in 'IJKLMNOP': cell(wcmp, f'{col}{NC+1}', f'=SUM({col}5:{col}{NC})', ARB, N2)
cell(wcmp, f'B{NC+3}', 'Fuente: PDF «Comprobante de pago» del IESS (RQ-005), leídos automáticamente. «Emitido» es la fecha de generación del comprobante; '
                       'la fecha efectiva de pago se confirma con el estado bancario. Préstamos clasificados por el concepto del comprobante '
                       '(«PAGO CREDITOS HIPOTECARIOS» = hipotecario; «PAG.DIV.PRESTAMOS» = quirografario).', IT)
for col, w in zip('ABCDEFGHIJKLMNOPQRS', [2, 44, 22, 18, 36, 6, 11, 11, 12, 10, 10, 10, 10, 12, 11, 12, 9, 13, 10]): wcmp.column_dimensions[col].width = w
wcmp.freeze_panes = 'D5'; wcmp.auto_filter.ref = f'B4:S{NC}'
CR = lambda col: f"'Comprobantes RQ-005'!${col}$5:${col}${NC}"

# ---------------- Insumos (libros, consolidado)
wi = wb.create_sheet('Insumos')
cell(wi, 'B2', 'INSUMOS MENSUALES', ARB)
cell(wi, 'B3', 'Libros (azul): Mayor de pasivos RQ-004 (5BS.21203), créditos del mes por provisión y finiquitos; fondos de reserva netos de reversiones. Consolidado IESS (verde): RQ-003.', IT)
th(wi, 5, 2, ['Mes', 'Libros: aportes', 'Libros: fondos de reserva', 'Libros: préstamos IESS', 'Consolidado: aportes (pers.+patr.+CCC)', 'Consolidado: FR a depositar'], 44)
for m in range(12):
    r = 6 + m; cell(wi, f'B{r}', MESES[m], AR, box=True)
    for j, k in enumerate(['APO', 'FR', 'PRE']): cell(wi, f'{chr(67+j)}{r}', lib[k][m], BLUE, N2, box=True)
    cell(wi, f'F{r}', float(cons.loc[float(m + 1), 'tot']), GREEN, N2, box=True)
    cell(wi, f'G{r}', float(cons.loc[float(m + 1), 'fr']), GREEN, N2, box=True)
cell(wi, 'B18', 'Σ', RED, box=True)
for col in 'CDEFG': cell(wi, f'{col}18', f'=SUM({col}6:{col}17)', ARB, N2, box=True)
cell(wi, 'B21', 'Saldos al 31-dic-2025 (Anexo RQ-002)', ARB)
th(wi, 22, 2, ['Cuenta', 'Nombre', 'Saldo'], 20)
SALDOS = [('5BS.21203.001', 'Aporte patronal por pagar'), ('5BS.21203.002', 'Aporte individual por pagar'), ('5BS.21203.003', 'Préstamos quirografarios'),
          ('5BS.21203.004', 'Fondos de reserva'), ('5BS.21203.005', 'Préstamo hipotecario'), ('5BS.21204.004', 'Sueldos por pagar')]
srow = {}
for i, (cod, nom) in enumerate(SALDOS):
    r = 23 + i; srow[cod] = r
    cell(wi, f'B{r}', cod, AR, box=True); cell(wi, f'C{r}', nom, AR, box=True); cell(wi, f'D{r}', ANEXO[cod], BLUE, N2, box=True)
cell(wi, 'B30', 'Rol de diciembre (DO-3.1, pestaña «normal») — líquido a recibir', ARB)
cell(wi, 'B31', 'Empleados y jubilados', AR, box=True); cell(wi, 'D31', 163774.65, GREEN, N2, box=True)
cell(wi, 'B32', 'Estipendio de pasantes (Mayor de gastos RQ-004, diciembre)', AR, box=True); cell(wi, 'D32', MY.gasto_mes('EST')[11], BLUE, N2, box=True)
for col, w in zip('ABCDEFG', [2, 16, 26, 18, 20, 20, 18]): wi.column_dimensions[col].width = w
IN = "'Insumos'"

# ---------------- DO-9
ws = wb.create_sheet('DO9 IESS por Pagar', 0)
for r in range(1, 4):
    for col in range(1, 16): ws.cell(r, col).fill = NAVY
ws.row_dimensions[1].height = 34; ws.row_dimensions[2].height = 22
img = XLImage(_R.CTX.logo); img.height = 60; img.width = 160; ws.add_image(img, 'A1')
cell(ws, 'G1', 'OBLIGACIONES ACUMULADAS', WH12); cell(ws, 'G3', 'IESS por pagar — libros vs comprobantes de pago', WH)
cell(ws, 'A4', 'Nombre del cliente:', ARB, box=True); cell(ws, 'B4', '=Ficha!C4', AR, box=True)
cell(ws, 'F4', 'Período terminado:', ARB, box=True); cell(ws, 'G4', '=Ficha!C5', AR, 'dd/mm/yyyy', box=True)
cell(ws, 'A5', 'Preparado por:', ARB, box=True); cell(ws, 'B5', '=Ficha!C6', AR, box=True)
cell(ws, 'C5', 'Revisado:', ARB, box=True); cell(ws, 'D5', '=Ficha!C7', AR, box=True)
cell(ws, 'F5', 'Fecha:', ARB, box=True); cell(ws, 'G5', '=Ficha!C8', AR, 'dd/mm/yyyy', box=True)
cell(ws, 'I5', 'Referencia de PT', ARB, box=True); cell(ws, 'J5', 'DO-9', RED, box=True)
ws.merge_cells('A6:O6'); b = ws['A6']
b.value = ('Comprobantes: RQ-005 (85 PDF leídos). Saldos: Anexo RQ-002. Libros mensuales: Mayor de pasivos RQ-004 (5BS.21203), conciliado con el anexo (pestaña «Mayor RQ-004»).')
b.font = Font(name='Arial', size=9, bold=True, color='FFC00000'); b.fill = YEL; b.alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[6].height = 24

# Cabeceras
for rng_, t in (('B8:D8', 'Según libros'), ('F8:J8', 'Según comprobantes de pago IESS (RQ-005)'), ('L8:N8', 'Diferencias (libros − comprobantes)')):
    ws.merge_cells(rng_); c = ws[rng_.split(':')[0]]; c.value = t; c.font = WH; c.fill = NAVY; c.alignment = C
th(ws, 9, 1, ['Meses', 'Aportes', 'Fondos de reserva', 'Préstamos IESS'], 44)
th(ws, 9, 6, ['Aportes', 'Fondos de reserva', 'Préstamos hipotecarios', 'Extensión salud cónyuges', 'Préstamos quirografarios'], 44)
th(ws, 9, 12, ['Aportes', 'Fondos de reserva', 'Préstamos IESS'], 44)
th(ws, 9, 15, ['Notas de crédito IESS (aportes)'], 44)
for m in range(12):
    r = 10 + m; ir = 6 + m; mm = m + 1
    cell(ws, f'A{r}', MESES[m], AR, box=True)
    cell(ws, f'B{r}', f'={IN}!C{ir}', GREEN, N2, box=True); cell(ws, f'C{r}', f'={IN}!D{ir}', GREEN, N2, box=True); cell(ws, f'D{r}', f'={IN}!E{ir}', GREEN, N2, box=True)
    cell(ws, f'E{r}', '√', RED)
    for col, tipo in (('F', 'Aportes'), ('G', 'Fondos de reserva'), ('H', 'Préstamo hipotecario'), ('I', 'Extensión de salud cónyuges'), ('J', 'Préstamo quirografario')):
        cell(ws, f'{col}{r}', f'=SUMIFS({CR("P")},{CR("C")},"{tipo}",{CR("F")},{mm})', GREEN, N2, box=True)
    cell(ws, f'L{r}', f'=B{r}-F{r}-I{r}', AR, N2, box=True)
    cell(ws, f'M{r}', f'=C{r}-G{r}', AR, N2, box=True)
    cell(ws, f'N{r}', f'=D{r}-H{r}-J{r}', AR, N2, box=True)
    cell(ws, f'O{r}', f'=SUMIFS({CR("O")},{CR("C")},"Aportes",{CR("F")},{mm})', GREEN, N2, box=True)
for col in 'BCDFGHIJLMNO':
    cell(ws, f'{col}22', f'=SUM({col}10:{col}21)', WH, N2, NAVY, box=True)
cell(ws, 'A22', 'Total', WH, fill=NAVY, box=True)
for col in 'BCDFGHJLMN': cell(ws, f'{col}23', 'Σ', RED, al=Alignment(horizontal='center'))
cell(ws, 'A24', 'Aportes: la diferencia libros − comprobantes se explica principalmente por las notas de crédito que el IESS descuenta al cobrar (col. O); '
                'la empresa registra el aporte bruto. Fondos de reserva: enero (238,18) es la nota de crédito del comprobante de ajuste. El Mayor RQ-004 muestra que los 223,57 de febrero del papel anterior eran una provisión revertida en el mismo mes (REV.FOND.RESERVA/5DFIN): diferencia resuelta.', IT)
ws.merge_cells('A24:O25'); ws['A24'].alignment = Alignment(wrap_text=True, vertical='top')

# Cruce a tres vías aportes y FR
cell(ws, 'A27', 'CRUCE A TRES VÍAS — APORTES Y FONDOS DE RESERVA (anual)', NAVYF)
th(ws, 28, 1, ['Concepto', 'Libros', 'Comprobantes (total)', '+ Notas de crédito', 'Comprobantes (bruto)', 'Consolidado IESS', 'Comp. bruto − Consolidado', 'Libros − Comp. bruto'], 36)
three = [('Aportes (incl. extensión de salud)', '=B22', '=F22+I22', f'=SUMIFS({CR("O")},{CR("C")},"Aportes")+SUMIFS({CR("O")},{CR("C")},"Extensión de salud cónyuges")', f'{IN}!F18'),
         ('Fondos de reserva', '=C22', '=G22', f'=SUMIFS({CR("O")},{CR("C")},"Fondos de reserva")', f'{IN}!G18')]
for i, (t, a, b2, nc, co) in enumerate(three):
    r = 29 + i
    cell(ws, f'A{r}', t, AR, box=True); cell(ws, f'B{r}', a, AR, N2, box=True); cell(ws, f'C{r}', b2, AR, N2, box=True)
    cell(ws, f'D{r}', nc, GREEN, N2, box=True); cell(ws, f'E{r}', f'=C{r}+D{r}', ARB, N2, box=True)
    cell(ws, f'F{r}', '=' + co, GREEN, N2, box=True); cell(ws, f'G{r}', f'=E{r}-F{r}', ARB, N2, box=True); cell(ws, f'H{r}', f'=B{r}-E{r}', ARB, N2, box=True)
cell(ws, 'A31', 'Comprobantes bruto − Consolidado: seguro de salud tiempo parcial (≈ AD 4,41% JP, 349,18) y ajustes de planilla. '
                'Libros − comprobantes bruto: según el Mayor RQ-004, aporte de pasantes pagado fuera de planilla (330,88), pagos de ajuste y diferencias de redondeo (ver pestaña «Mayor RQ-004»).', IT)
ws.merge_cells('A31:O31')
cell(ws, 'A32', 'Seguro de salud tiempo parcial en comprobantes (= «Otros gastos de nómina» AD 4,41% JP de la DO-2):', AR)
cell(ws, 'H32', f'=SUM({CR("K")})', GREEN, N2); cell(ws, 'I32', '=IF(ABS(H32-349.18)<0.01,"Cuadra con DO-2 ✓","Revisar")', IT)

# Saldos por pagar a diciembre
cell(ws, 'A35', 'SALDOS POR PAGAR AL 31-DIC-2025', NAVYF)
th(ws, 36, 1, ['Concepto', 'Comprobante / rol de diciembre', 'Saldo en libros (RQ-002)', 'Diferencia', 'Observación'], 30)
sal = [('Aportes por pagar (patronal + individual)', '=F21', f"={IN}!D{srow['5BS.21203.001']}+{IN}!D{srow['5BS.21203.002']}", 'Planilla de diciembre pagada en enero 2026'),
       ('Fondos de reserva por pagar', '=G21', f"={IN}!D{srow['5BS.21203.004']}", 'Depósito de diciembre en enero 2026'),
       ('Préstamos quirografarios por pagar', '=J21', f"={IN}!D{srow['5BS.21203.003']}", 'Clasificación por concepto del comprobante'),
       ('Préstamos hipotecarios por pagar', '=H21', f"={IN}!D{srow['5BS.21203.005']}", 'Clasificación por concepto del comprobante'),
       ('Sueldos por pagar', f"={IN}!D31+{IN}!D32", f"={IN}!D{srow['5BS.21204.004']}", 'Rol de diciembre pagado dentro del mes; saldo corresponde al sueldo de diciembre de un trabajador: confirmar pago posterior')]
for i, (t, a, b2, obs) in enumerate(sal):
    r = 37 + i
    cell(ws, f'A{r}', t, AR, box=True); cell(ws, f'B{r}', a, GREEN, N2, box=True); cell(ws, f'C{r}', b2, GREEN, N2, box=True)
    cell(ws, f'D{r}', f'=B{r}-C{r}' if i < 4 else '="Pagado en el mes"', ARB, N2, box=True)
    cell(ws, f'E{r}', obs, IT if i < 4 else BLUE, box=True); ws.merge_cells(f'E{r}:O{r}')
cell(ws, 'A42', 'Sueldos por pagar: el rol «normal» de diciembre más el estipendio de pasantes reproduce lo pagado en el mes; '
                'el saldo contable del sueldo de diciembre debe sustentarse con el pago en el mes siguiente.', IT)
ws.merge_cells('A42:O42')

# Corrección préstamos
cell(ws, 'A44', 'CORRECCIÓN AL PAPEL ANTERIOR — CLASIFICACIÓN DE PRÉSTAMOS', NAVYF)
th(ws, 45, 1, ['Tipo', 'Papel anterior', 'Según comprobantes', 'Diferencia'], 22)
cell(ws, 'A46', 'Préstamos hipotecarios', AR, box=True); cell(ws, 'B46', 72981.65, BLUE, N2, box=True); cell(ws, 'C46', '=H22', AR, N2, box=True); cell(ws, 'D46', '=B46-C46', ARB, N2, box=True)
cell(ws, 'A47', 'Préstamos quirografarios', AR, box=True); cell(ws, 'B47', 64716.90, BLUE, N2, box=True); cell(ws, 'C47', '=J22', AR, N2, box=True); cell(ws, 'D47', '=B47-C47', ARB, N2, box=True)
cell(ws, 'A48', 'Total préstamos', ARB, box=True)
for col in 'BCD': cell(ws, f'{col}48', f'=SUM({col}46:{col}47)', ARB, N2, box=True)
cell(ws, 'A49', 'De enero a agosto el papel anterior tenía cruzadas las columnas hipotecario / quirografario; el total no cambia.', IT)
ws.merge_cells('A49:O49')
cell(ws, 'A50', 'Los aportes se concilian con el total de cada comprobante (no solo el «valor»), considerando las extensiones de salud compensadas con notas de crédito.', IT)
ws.merge_cells('A50:O50')

# Validación V10
cell(ws, 'A52', 'V10 — PLAZO DE PAGO', NAVYF)
cell(ws, 'A53', 'Comprobantes emitidos después de la fecha de vigencia de pago:', AR); cell(ws, 'H53', f'=COUNTIF({CR("R")},"FUERA DE PLAZO")', ARB)
cell(ws, 'A54', 'Margen mínimo (días) entre emisión y vigencia:', AR); cell(ws, 'H54', f'=MIN({CR("Q")})', ARB)
cell(ws, 'A55', 'La fecha efectiva de pago se confirma con el estado bancario; el comprobante evidencia la generación dentro del plazo.', IT); ws.merge_cells('A55:O55')

txt = [(58, 'Objetivo:', RU),
       (59, 'Verificar que las obligaciones con el IESS (aportes, fondos de reserva y préstamos retenidos) registradas en libros se hayan pagado oportunamente '
            'según los comprobantes de pago, y que los saldos por pagar al 31 de diciembre de 2025 estén sustentados.', AR),
       (61, 'Procedimiento:', RU),
       (62, '1. Se leyeron los 85 comprobantes de pago del IESS (RQ-005) y se clasificaron por tipo y período. 2. Se compararon mes a mes con las obligaciones '
            'registradas en libros. 3. Se conciliaron aportes y fondos de reserva a tres vías con el Consolidado IESS (RQ-003). 4. Se cotejaron los saldos de '
            'diciembre con el Anexo RQ-002 y los sueldos por pagar con el rol (DO-3.1). 5. Se evaluó el plazo de pago (V10).', AR),
       (64, 'Marcas:', RED), (65, 'Σ  Sumado, totalizado.', ARB), (66, '√  Verificado, comprobado, cotejado (con libros).', ARB), (67, '‡  Diferencia determinada.', ARB),
       (69, 'Comentario de auditoría:', RED)]
for r, t, f in txt:
    cell(ws, f'A{r}', t, f)
    if f is AR: ws.merge_cells(f'A{r}:O{r}'); ws[f'A{r}'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 32
cell(ws, 'A70', '="Aportes: libros − comprobantes "&TEXT(L22,"#,##0.00")&", explicado en "&TEXT(O22,"#,##0.00")&" por notas de crédito del IESS. Fondos de reserva: "&TEXT(M22,"#,##0.00")&" (238,18 nota de crédito de enero; los 223,57 de febrero del papel anterior eran una reversión según el Mayor RQ-004). '
                'Préstamos: cuadran en total; el papel anterior tenía cruzada la clasificación de enero a agosto. Saldos de diciembre sustentados con comprobantes. Sin pagos fuera de plazo. Mayor RQ-004: los 5 saldos se reconstruyen desde el saldo inicial (diferencia "&TEXT(\'Mayor RQ-004\'!K12,"#,##0.00")&") y cada pago al IESS se coteja con su comprobante (diferencia total "&TEXT(\'Mayor RQ-004\'!G33,"#,##0.00")&")."', AR)
ws.merge_cells('A70:O71'); ws['A70'].alignment = Alignment(wrap_text=True, vertical='top'); ws['A70'].border = BOX; ws.row_dimensions[70].height = 30
for col, wd in zip('ABCDEFGHIJKLMNO', [34, 14, 14, 14, 3, 14, 13, 13, 12, 13, 3, 13, 12, 12, 13]): ws.column_dimensions[col].width = wd
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

# ---------------- Mayor RQ-004 (5BS.21203)
wm = wb.create_sheet('Mayor RQ-004')
cell(wm, 'B2', 'MAYOR DE PASIVOS RQ-004 — IESS POR PAGAR (5BS.21203): movimiento del año y pagos por período', ARB)
cell(wm, 'B3', 'Signo del mayor: débito +, crédito −. Saldo inicial y final: Anexo RQ-002. Clasificación: columna REF del cliente; pagos separados por la explicación del asiento.', IT)
th(wm, 5, 2, ['Cuenta', 'Nombre', 'Saldo inicial (RQ-002)', 'Provisión', 'Finiquitos', 'Ajustes', 'Pagos al IESS', 'Notas de crédito (permisos médicos)',
              'Otros pagos', 'Saldo final calculado', 'Saldo ESF (RQ-002)', 'Diferencia'], 44)
NOM = {'5BS.21203.001': 'Aporte patronal', '5BS.21203.002': 'Aporte individual', '5BS.21203.003': 'Préstamos quirografarios',
       '5BS.21203.004': 'Fondos de reserva', '5BS.21203.005': 'Préstamo hipotecario'}
for i, cta in enumerate(NOM):
    r = 6 + i; x = _I[_I.cta == cta]
    cell(wm, f'B{r}', cta, AR, box=True); cell(wm, f'C{r}', NOM[cta], AR, box=True)
    cell(wm, f'D{r}', MY.ANEXO[cta][0], BLUE, N2, box=True)
    for col, sb in zip('EFGHIJ', ['PROVISION', 'FINIQUITOS', 'AJUSTES', 'PAGO IESS', 'NC', 'OTROS PAGOS']):
        cell(wm, f'{col}{r}', round(float(x[x['sub'] == sb].Importe.sum()), 2), BLUE, N2, box=True)
    cell(wm, f'K{r}', f'=SUM(D{r}:J{r})', ARB, N2, box=True); cell(wm, f'L{r}', MY.ANEXO[cta][1], BLUE, N2, box=True)
    cell(wm, f'M{r}', f'=ROUND(K{r}-L{r},2)', ARB, N2, box=True)
cell(wm, 'B11', 'TOTAL', WH, fill=NAVY, box=True); cell(wm, 'C11', '', WH, fill=NAVY, box=True)
for col in 'DEFGHIJKLM': cell(wm, f'{col}11', f'=SUM({col}6:{col}10)', WH, N2, NAVY, box=True)
cell(wm, 'B12', 'Diferencia total saldo calculado − ESF:', ARB); cell(wm, 'K12', '=M11', RED, N2)
cell(wm, 'B13', 'Otros pagos: aporte de pasantes pagado fuera de planilla (junio 330,88), ajustes de aportes y fondos pagados al IESS (17,90) y descuento mal aplicado de extensión de salud (27,28).', IT)
cell(wm, 'B16', 'PAGOS AL IESS SEGÚN MAYOR vs COMPROBANTES DE PAGO (RQ-005) — por período de la planilla', ARB)
th(wm, 17, 2, ['Período', 'Mes de pago', 'Pago aportes (mayor)', 'Comprobante aportes + ext. salud', 'Diferencia aportes', 'Pago FR (mayor)', 'Comprobante FR',
               'Pago préstamos (mayor)', 'Comprobantes préstamos', 'Diferencia FR + préstamos'], 44)
PER = ['Dic-2024'] + [f'{m[:3]}-2025' for m in MESES[:11]]
for k in range(12):
    r = 18 + k; mm = k  # período: 0 = dic-2024 (saldo inicial)
    cell(wm, f'B{r}', PER[k], AR, box=True); cell(wm, f'C{r}', MESES[k], AR, box=True)
    cell(wm, f'D{r}', pag['APO'][k], BLUE, N2, box=True)
    if k == 0:
        cell(wm, f'E{r}', '=-(D6+D7)', GREEN, N2, box=True); cell(wm, f'H{r}', '=-D9', GREEN, N2, box=True); cell(wm, f'J{r}', '=-(D8+D10)', GREEN, N2, box=True)
    else:
        cell(wm, f'E{r}', f'=SUMIFS({CR("P")},{CR("C")},"Aportes",{CR("F")},{mm})+SUMIFS({CR("P")},{CR("C")},"Extensión de salud cónyuges",{CR("F")},{mm})', GREEN, N2, box=True)
        cell(wm, f'H{r}', f'=SUMIFS({CR("P")},{CR("C")},"Fondos de reserva",{CR("F")},{mm})', GREEN, N2, box=True)
        cell(wm, f'J{r}', f'=SUMIFS({CR("P")},{CR("C")},"Préstamo hipotecario",{CR("F")},{mm})+SUMIFS({CR("P")},{CR("C")},"Préstamo quirografario",{CR("F")},{mm})', GREEN, N2, box=True)
    cell(wm, f'F{r}', f'=ROUND(D{r}-E{r},2)', AR, N2, box=True)
    cell(wm, f'G{r}', pag['FR'][k], BLUE, N2, box=True)
    cell(wm, f'I{r}', round(pag['QUI'][k] + pag['HIP'][k], 2), BLUE, N2, box=True)
    cell(wm, f'K{r}', f'=ROUND(G{r}-H{r}+I{r}-J{r},2)', AR, N2, box=True)
cell(wm, 'B30', 'TOTAL', WH, fill=NAVY, box=True); cell(wm, 'C30', '', WH, fill=NAVY, box=True)
for col in 'DEFGHIJK': cell(wm, f'{col}30', f'=SUM({col}18:{col}29)', WH, N2, NAVY, box=True)
cell(wm, 'B31', 'La planilla de diciembre de 2025 se paga en enero de 2026 y es el saldo por pagar al cierre (ver DO-9, saldos por pagar). El período dic-2024 se coteja con el saldo inicial del anexo. Diferencias de aportes (ene, sep–nov): extensión de salud y notas de crédito descontadas en el comprobante; FR ene: ajuste pagado 1,16.', IT)
cell(wm, 'B33', 'Diferencia total pagos − comprobantes (aportes + FR + préstamos):', ARB); cell(wm, 'G33', '=F30+K30', RED, N2)
for col, w in zip('ABCDEFGHIJKLM', [2, 15, 22, 15, 15, 13, 14, 15, 14, 14, 15, 15, 12]): wm.column_dimensions[col].width = w
wm.freeze_panes = 'D6'

# ---------------- Controles
wc = wb.create_sheet('Controles')
cell(wc, 'B2', 'CONTROLES DE INTEGRIDAD', ARB)
th(wc, 4, 2, ['Control', 'Valor A', 'Valor B', 'Diferencia', 'Estado'])
D9 = "'DO9 IESS por Pagar'"
ctl = [('Comprobantes leídos = 85', f'=COUNTA({CR("D")})', 85, 0),
       ('Control aritmético de cada comprobante (subtotal y total)', f'=SUMPRODUCT(ABS({CR("S")}))', 0, 0.01),
       ('Fondos de reserva: comprobantes', f'={D9}!G22', _R.CIFRAS['fr_iess'], 0.01),
       ('Préstamos totales: comprobantes = libros', f'={D9}!H22+{D9}!J22', f'={D9}!D22', 0.01),
       ('Aportes dic = saldo libros aportes por pagar', f'={D9}!B37', f'={D9}!C37', 0.01),
       ('FR dic = saldo libros', f'={D9}!B38', f'={D9}!C38', 0.01),
       ('Quirografario dic = saldo libros', f'={D9}!B39', f'={D9}!C39', 0.01),
       ('Hipotecario dic = saldo libros', f'={D9}!B40', f'={D9}!C40', 0.01),
       ('Consolidado aportes = DO-3', f"={IN}!F18", _R.CIFRAS['consolidado_aportes'], 0.01),
       ('Seguro salud tiempo parcial = AD 4,41% JP DO-2 (349,18)', f'={D9}!H32', 349.18, 0.01),
       ('Sin comprobantes fuera de plazo (V10)', f'={D9}!H53', 0, 0),
       ('Mayor RQ-004: saldo inicial + movimientos = saldo ESF (5 cuentas)', "='Mayor RQ-004'!M11", 0, 0.01),
       ('Mayor RQ-004: aportes provisionados = libros del cruce mensual', "='Mayor RQ-004'!E6+'Mayor RQ-004'!E7+'Mayor RQ-004'!F6+'Mayor RQ-004'!F7", f'=-{D9}!B22', 0.01)]
for i, (t, a, b2, tol) in enumerate(ctl):
    r = 5 + i
    cell(wc, f'B{r}', t, AR, box=True); cell(wc, f'C{r}', a, GREEN, N2, box=True); cell(wc, f'D{r}', b2, GREEN, N2, box=True)
    cell(wc, f'E{r}', f'=ROUND(C{r}-D{r},2)', AR, N2, box=True)
    cell(wc, f'F{r}', f'=IF(ABS(E{r})<={tol},"OK","REVISAR")', ARB, box=True, al=Alignment(horizontal='center'))
n = len(ctl); cell(wc, f'B{6+n}', f'=COUNTIF(F5:F{4+n},"OK")&" de {n} controles OK"', ARB)
wc.column_dimensions['B'].width = 60
for col in 'CDEF': wc.column_dimensions[col].width = 16
wb.active = 0
_wf = wb['Ficha']
_r = max(_wf.max_row + 2, 14)
_wf.cell(_r, 2, 'Año auditado 2026. Cédula construida con datos de la plantilla 2025 (rol, IESS, anexo y comprobantes); se reemplazan con la información 2026 del cliente.').font = Font(name='Arial', size=8, italic=True, color='FFC00000')
wb.save(OUT); print('ok', OUT)
