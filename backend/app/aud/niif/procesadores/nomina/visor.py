# -*- coding: utf-8 -*-
"""Visor de cédulas de nómina: convierte los libros que genera el motor en datos JSON
que el panel muestra en pantalla, pestaña por pestaña (valor calculado + fórmula).

Uso:  python3 visor.py <carpeta_con_las_cedulas_xlsx> <salida.json>

Requisito: los libros deben estar recalculados (LibreOffice / recalc.py) para que
cada fórmula tenga su valor en caché. No contiene cifras ni nombres de ningún cliente:
todo sale de los libros que recibe.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

# Cédula → (código, título) en el orden del panel. Las claves son prefijos de archivo.
CEDULAS = [
    ("DO-0_", "DO-0", "Programa de auditoría de nómina"),
    ("DO-2_", "DO-2", "Cédula sumaria de nómina"),
    ("DO-3_", "DO-3", "Resumen de planillas IESS"),
    ("DO-3.1_", "DO-3.1", "Resumen de roles de pago"),
    ("DO-4_", "DO-4", "Reproceso de beneficios sociales"),
    ("DO-4C_DO-5_", "DO-4C / DO-5", "Cruce global y mensual del pasivo"),
    ("DO-6_DO-7_", "DO-6 / DO-7", "Cruce global y mensual del gasto"),
    ("DO-8_", "DO-8", "Base imponible IESS vs libros"),
    ("DO-9_", "DO-9", "IESS por pagar vs comprobantes"),
    ("DO-10_", "DO-10", "Roles de pago vs libros"),
    ("DO-12_", "DO-12", "Hallazgos y ajustes"),
    ("VAC_", "DO-4V", "Vacaciones no gozadas"),
    ("RQ-007_", "DO-4 V08", "Muestreo de solicitudes de décimos"),
    ("RQ-004_", "RQ-004", "Mayores de nómina"),
    ("RQ-008_", "RQ-008", "Kárdex de vacaciones"),
    ("RQ-003_", "RQ-003", "Exclusiones de beneficios"),
]


def _num(v: float, dec: int) -> str:
    s = f"{abs(v):,.{dec}f}".replace(",", "§").replace(".", ",").replace("§", ".")
    return ("-" if v < 0 and round(abs(v), dec) != 0 else "") + s


def _fmt(v, nf: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float)):
        nf = nf or "General"
        if "%" in nf:
            m = re.search(r"0\.(0+)%", nf)
            return _num(v * 100, len(m.group(1)) if m else 0) + "%"
        if re.search(r"[dmy]{2}", nf.lower()) and "0" not in nf:
            try:
                d = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
                return d.strftime("%d/%m/%Y")
            except Exception:  # noqa: BLE001
                pass
        m = re.search(r"0\.(0+)", nf)
        if m:
            return _num(v, len(m.group(1)))
        if "#,##0" in nf or nf.strip() == "0":
            return _num(v, 0)
        if isinstance(v, float) and not v.is_integer():
            # Formato «General» de Excel: muestra los decimales que tiene el número (máx. 10).
            txt = f"{v:.10f}".rstrip("0")
            return _num(v, max(1, len(txt.split(".")[1])) if "." in txt else 0)
        return str(int(v)) if isinstance(v, float) else str(v)
    return str(v)


def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb[:2] != "00":
            return "#" + c.rgb[2:]
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb not in ("00000000",):
            return "#" + c.rgb[2:]
    except Exception:  # noqa: BLE001
        return None
    return None


def hoja(ws_f, ws_v, estilos: dict) -> dict:
    max_r, max_c = ws_f.max_row, ws_f.max_column
    ocultas = {c for c in range(1, max_c + 1)
               if ws_f.column_dimensions[get_column_letter(c)].hidden}
    cols = [c for c in range(1, max_c + 1) if c not in ocultas]
    filas, ult_c, ri = [], 0, []
    for r in range(1, max_r + 1):
        if ws_f.row_dimensions[r].hidden:
            continue
        ri.append(r)
        fila = []
        for i, c in enumerate(cols):
            cf = ws_f.cell(r, c)
            vf = cf.value
            v = ws_v.cell(r, c).value
            formula = vf if isinstance(vf, str) and vf.startswith("=") else None
            txt = _fmt(v if formula else vf, cf.number_format)
            fill = _color(cf.fill.fgColor) if cf.fill and cf.fill.fill_type == "solid" else None
            key = (bool(cf.font and cf.font.b), _color(cf.font.color) if cf.font else None, fill,
                   (cf.alignment.horizontal or "") if cf.alignment else "",
                   isinstance(v if formula else vf, (int, float)) and not isinstance(v, bool))
            if not txt and not fill and not formula:
                fila.append(None)
                continue
            sid = estilos.setdefault(key, len(estilos))
            cell = [txt, sid]
            if formula:
                cell.append(formula)
            fila.append(cell)
            ult_c = max(ult_c, i + 1)
        while fila and fila[-1] is None:
            fila.pop()
        filas.append(fila)
    while filas and not filas[-1]:
        filas.pop()
    anchos = [round((ws_f.column_dimensions[get_column_letter(c)].width or 9) * 7) for c in cols[:ult_c]]
    pos = {c: i for i, c in enumerate(cols)}
    rpos = {r: i for i, r in enumerate(ri[:len(filas)])}
    merges = []
    for m in ws_f.merged_cells.ranges:
        if m.min_col in pos and m.max_col in pos and m.min_row in rpos and m.max_row in rpos:
            merges.append([rpos[m.min_row], pos[m.min_col], rpos[m.max_row], pos[m.max_col]])
    ri = ri[:len(filas)]
    return {"n": ws_f.title, "w": anchos, "r": filas, "m": merges,
            "fz": ws_f.freeze_panes or "", "ci": cols[:ult_c],
            "ri": None if ri == list(range(1, len(ri) + 1)) else ri}


def construir(carpeta: str) -> dict:
    estilos: dict = {}
    libros = []
    archivos = sorted(os.listdir(carpeta))
    for pref, cod, tit in CEDULAS:
        f = next((a for a in archivos if a.startswith(pref) and a.endswith(".xlsx")), None)
        if not f:
            continue
        ruta = os.path.join(carpeta, f)
        wf = openpyxl.load_workbook(ruta)
        wv = openpyxl.load_workbook(ruta, data_only=True)
        hojas = [hoja(ws, wv[ws.title], estilos) for ws in wf.worksheets if ws.sheet_state == "visible"]
        libros.append({"f": f, "c": cod, "t": tit, "h": hojas})
    tabla = [None] * len(estilos)
    for (b, col, fill, al, num), i in estilos.items():
        tabla[i] = {"b": int(b), "c": col, "f": fill, "a": al or ("right" if num else "")}
    return {"v": 1, "estilos": tabla, "libros": libros}


def empaquetar(carpeta: str) -> str:
    """JSON del visor comprimido en gzip y codificado en base64 (para inyectar en el panel)."""
    import base64
    import gzip
    crudo = json.dumps(construir(carpeta), ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return base64.b64encode(gzip.compress(crudo, 9, mtime=0)).decode("ascii")


if __name__ == "__main__":
    datos = construir(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as fh:
        json.dump(datos, fh, ensure_ascii=False, separators=(",", ":"))
    print(len(datos["libros"]), "libros ·", sum(len(x["h"]) for x in datos["libros"]), "hojas ·",
          os.path.getsize(sys.argv[2]) // 1024, "KB")
