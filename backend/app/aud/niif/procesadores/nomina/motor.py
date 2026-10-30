"""Motor de reproceso de nómina del Command Center.

Reproduce las series compartidas que consumen las cédulas DO-2..DO-12,
siguiendo el algoritmo de las fórmulas de DO-3.1 y DO-4. No toca el repo;
sirve para probar paridad al centavo contra valores_control/.
"""
from __future__ import annotations
import openpyxl, collections, datetime, calendar, re

MESES_TXT = {m: i + 1 for i, m in enumerate(
    ['january','february','march','april','may','june','july','august','september','october','november','december'])}
MESES_TXT.update({m: i + 1 for i, m in enumerate(
    ['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre'])})

SBU_2025 = 470.0
PCT_FR = 0.0833
EJ = 2025        # año del ejercicio auditado; lo fija motor.correr() desde la ficha
MES_CORTE = 12   # mes de corte (1-12); 12 = cierre calendario, otro = visita preliminar
VISITA = ""      # "preliminar" | "final" (de la ficha)


def mnum(v):
    if isinstance(v, (int, float)):
        n = int(v); return n if 1 <= n <= 12 else None
    s = str(v).strip().lower()
    if s in MESES_TXT: return MESES_TXT[s]
    m = re.search(r'\b(1[0-2]|[1-9])\b', s); return int(m.group(1)) if m else None


def num(v):
    return v if isinstance(v, (int, float)) else 0.0


def fdate(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    return None


def edate(d, months):
    # mismo día, +months meses (ajusta fin de mes)
    y = d.year + (d.month - 1 + months) // 12
    m = (d.month - 1 + months) % 12 + 1
    last = calendar.monthrange(y, m)[1]
    return datetime.date(y, m, min(d.day, last))


def codigo_rel(rel):
    s = str(rel or '').strip()
    return s.split('-')[0].strip() if s else ''


# ---------------------------------------------------------------- lectores
def leer_consolidado(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['detallePlanillaConsolidado']
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[1] is None or str(r[1]).startswith('Totales'): continue
        out.append(dict(mes=mnum(r[1]), ced=str(r[2]).strip(), nombre=str(r[3]).strip(),
                        rel=r[4], base=num(r[5]), dias=num(r[6]), pat=num(r[7]),
                        ind=num(r[8]), ccc=num(r[12]), total=num(r[13])))
    return out


def leer_rol_normal(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['normal']
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] in (None, '') or str(r[0]).startswith('Total') or r[1] in (None, ''): continue
        m = mnum(r[0])
        if m is None: continue
        cargo = str(r[2] or '')
        tipo = 'Jubilado' if 'JUBIL' in cargo.upper() else 'Empleado'
        out.append(dict(
            mes=m, nombre=str(r[1]).strip(), cargo=cargo, tipo=tipo,
            ingreso=fdate(r[3]), salida=fdate(r[4]), sueldo=num(r[5]),
            vac_gozada=num(r[6]), fr_rol=num(r[24]), d13mens=num(r[25]), d14mens=num(r[26]),
            vac_finiquito=num(r[18]),
            finiquito=num(r[18]) + num(r[19]) + num(r[20]) + num(r[27]) + num(r[28]),
            prov_d13=num(r[63]), prov_d14=num(r[64]), prov_vac=num(r[65]), prov_fr=num(r[66]),
        ))
    return out


def leer_exclusiones(path):
    try:
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['Exclusiones']
    except Exception:
        return {}, {}
    por_ced, por_nom = {}, {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        ced = r[1]; nom = r[2]
        if ced in (None, '', '(opcional)') and nom in (None, '', '(copie de la hoja «Nombres del IESS»)'):
            continue
        if not ced and not nom: continue
        d = dict(d13=str(r[3] or '').strip().upper() == 'X',
                 d14=str(r[4] or '').strip().upper() == 'X',
                 vac=str(r[5] or '').strip().upper() == 'X',
                 fr=str(r[6] or '').strip().upper() == 'X',
                 motivo=r[7])
        if ced not in (None, ''): por_ced[str(ced).strip()] = d
        if nom not in (None, ''): por_nom[str(nom).strip()] = d
    return por_ced, por_nom


# ---------------------------------------------------------------- motor
def correr(rutas):
    # Año y mes de corte del ejercicio auditado (de la ficha). El período auditado va de
    # enero al mes de corte; en diciembre es el cierre calendario (comportamiento histórico).
    global EJ, MES_CORTE, VISITA
    EJ = rutas.anio_corte
    MES_CORTE = rutas.mes_corte
    VISITA = (rutas.visit or "").strip().lower()
    cons = leer_consolidado(rutas.consolidado)
    rol = leer_rol_normal(rutas.rol)
    exc_ced, exc_nom = leer_exclusiones(rutas.exclusiones) if rutas.exclusiones else ({}, {})

    rol_emp = [x for x in rol if x['tipo'] != 'Jubilado']

    # Insumo DO-3.1: por (mes, nombre)
    insumo = {}
    for x in rol_emp:
        k = (x['mes'], x['nombre'])
        d = insumo.setdefault(k, dict(d13mens=0.0, d14mens=0.0, fr_rol=0.0, vac_gozada=0.0,
                                      finiquito=0.0, vac_finiquito=0.0, prov_d13=0.0, prov_d14=0.0,
                                      prov_vac=0.0, prov_fr=0.0, ingreso=None, salida=None))
        for c in ('d13mens', 'd14mens', 'fr_rol', 'vac_gozada', 'finiquito', 'vac_finiquito',
                  'prov_d13', 'prov_d14', 'prov_vac', 'prov_fr'):
            d[c] += x[c]
        if x['ingreso'] and not d['ingreso']: d['ingreso'] = x['ingreso']
        if x['salida']: d['salida'] = x['salida']

    # presencia en rol y salida por nombre (cualquier mes)
    en_rol = set(insumo.keys())
    salida_nom = {}
    ingreso_nom = {}
    for x in rol_emp:
        if x['salida']: salida_nom[x['nombre']] = x['salida']
        if x['ingreso'] and x['nombre'] not in ingreso_nom: ingreso_nom[x['nombre']] = x['ingreso']

    def exc_de(ced, nom):
        return exc_ced.get(ced) or exc_nom.get(nom)

    # ---- DO-4 detalle por registro IESS
    det = []
    for r in cons:
        mes, nom, ced = r['mes'], r['nombre'], r['ced']
        cod = codigo_rel(r['rel'])
        X = (mes, nom) in en_rol
        aplica_tabla = cod in ('06', '16')
        L = aplica_tabla and X
        ingreso = insumo.get((mes, nom), {}).get('ingreso') if X else None
        activo = X and (nom not in salida_nom)
        ins = insumo.get((mes, nom), {})
        z = ins.get('d13mens', 0.0); aa = ins.get('d14mens', 0.0); ab = ins.get('fr_rol', 0.0)
        mens_d13 = z > 0; mens_d14 = aa > 0
        ex = exc_de(ced, nom) or {}
        exD13, exD14, exVac, exFR = ex.get('d13'), ex.get('d14'), ex.get('vac'), ex.get('fr')
        base, dias = r['base'], r['dias']
        # devengados
        P_bruto = base / 12 if (L and not exD13) else 0.0
        P = base / 12 if (L and not mens_d13 and not exD13) else 0.0
        P_mens = base / 12 if (L and mens_d13 and not exD13) else 0.0
        Q_bruto = SBU_2025 / 12 * dias / 30 if (L and not exD14) else 0.0
        Q_mens = SBU_2025 / 12 * dias / 30 if (L and mens_d14 and not exD14) else 0.0
        Q = SBU_2025 / 12 * dias / 30 if (L and not mens_d14 and not exD14) else 0.0
        R = base / 24 if (L and not exVac) else 0.0
        S = fr_calc(L, ingreso, mes, base, exFR)
        AI = 0.0 if ab > 0 else S
        # excluidos
        AR = base / 12 if (L and not mens_d13 and exD13) else 0.0
        AS = SBU_2025 / 12 * dias / 30 if (L and not mens_d14 and exD14) else 0.0
        AT = base / 24 if (L and exVac) else 0.0
        AU = fr_calc(L, ingreso, mes, base, False) if (L and ingreso and exFR) else 0.0
        AJ = P if activo else 0.0
        AK = Q if activo else 0.0
        AL = AI if activo else 0.0
        det.append(dict(mes=mes, nom=nom, ced=ced, base=base, dias=dias, cod=cod, X=X, L=L,
                        activo=activo, ingreso=ingreso, z=z, aa=aa, ab=ab,
                        pat=r['pat'], ind=r['ind'], ccc=r['ccc'],
                        P=P, P_bruto=P_bruto, P_mens=P_mens, Q=Q, Q_bruto=Q_bruto, Q_mens=Q_mens,
                        R=R, S=S, AI=AI, AR=AR, AS=AS, AT=AT, AU=AU,
                        AJ=AJ, AK=AK, AL=AL, mens_d13=mens_d13, mens_d14=mens_d14,
                        vac_gozada=ins.get('vac_gozada', 0.0),
                        vac_finiquito=ins.get('vac_finiquito', 0.0),
                        finiquito=ins.get('finiquito', 0.0)))

    return dict(cons=cons, rol=rol, rol_emp=rol_emp, insumo=insumo, det=det,
                salida_nom=salida_nom, ingreso_nom=ingreso_nom)


def fr_calc(L, ingreso, mes, base, exFR):
    if not L or ingreso is None or exFR:
        return 0.0
    aniv = edate(ingreso, 12)
    dim = calendar.monthrange(EJ, mes)[1]
    first = datetime.date(EJ, mes, 1)
    last = datetime.date(EJ, mes, dim)
    if first >= aniv:
        return base * PCT_FR
    if last >= aniv:
        return base * PCT_FR * ((last - aniv).days + 1) / dim
    return 0.0
