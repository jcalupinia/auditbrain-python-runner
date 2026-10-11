# -*- coding: utf-8 -*-
"""Puente motor → series compartidas (los pickles que consumen los generadores DO-*).

Reproduce, desde la salida del motor (``motor.correr``), las tablas intermedias
``consol``, ``rol_mes``, ``g``, ``g2``, ``g3``, ``cal_base`` y ``pop_acum`` con las
MISMAS columnas que esperaban los generadores de referencia. Validado al centavo
(≤0,01) contra ``valores_control/intermedios/``."""
from __future__ import annotations
import collections
import openpyxl
import pandas as pd
from backend.app.aud.niif.procesadores.nomina import motor

GRAV_IDX = [5, 6, 7, 8, 9, 10, 11, 12, 30, 31, 29]  # columnas gravadas del rol (base gravada)
ROL_MES_CONC = {  # concepto rol_mes -> índice en el rol 'normal'
    'SUELDO': 5, 'VAC TOMADAS': 6, 'REC.NOCTURNO': 7, 'H. EXT. 50%': 8, 'H. EXT. 100%': 9,
    'DESEMPEÑO': 10, 'BONIFICACION': 29, 'COMISIONES': 11, 'PREMIOS': 12, 'FOND.RESERV.': 24,
    'D13 MENS': 25, 'D14 MENS SIE': 26, 'LIQ.DEC13': 27, 'LIQ.DEC14S': 28, 'VAC.FINIQUIT': 18,
    'PENSION JUB': 21}


def _por_mes(det, fn):
    a = collections.defaultdict(float)
    for d in det:
        a[d['mes']] += fn(d)
    return {m: round(a[m], 2) for m in range(1, 13)}


def consol_df(R):
    det, cons = R['det'], R['cons']
    pat = _por_mes(cons, lambda r: r['pat']); ind = _por_mes(cons, lambda r: r['ind'])
    ccc = _por_mes(cons, lambda r: r['ccc']); fr = _por_mes(det, lambda d: d['AI'])
    filas = []
    for m in range(1, 13):
        ap = round(pat[m] + ind[m], 2); tot = round(ap + ccc[m], 2)
        filas.append(dict(Mes=float(m), aporte=ap, ccc=ccc[m], pat=pat[m], ind=ind[m], fr=fr[m], tot=tot, chk=tot))
    return pd.DataFrame(filas).set_index('Mes')


def _rol_rows(rol_path):
    ws = openpyxl.load_workbook(rol_path, read_only=True, data_only=True)['normal']
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] in (None, '') or str(r[0]).startswith('Total') or r[1] in (None, ''): continue
        m = motor.mnum(r[0])
        if m is None: continue
        out.append((m, 'JUBIL' in str(r[2] or '').upper(), r))
    return out


def rol_mes_df(rol_path):
    rows = _rol_rows(rol_path)
    agg = collections.defaultdict(lambda: collections.defaultdict(float))
    for m, _jub, r in rows:
        for k, idx in ROL_MES_CONC.items():
            agg[m][k] += motor.num(r[idx]) if idx < len(r) else 0
    # PAS = Total INGRESOS de la pestaña pasantes
    wb = openpyxl.load_workbook(rol_path, read_only=True, data_only=True)
    pas = collections.defaultdict(float)
    if 'pasantes' in wb.sheetnames:
        wp = wb['pasantes']; hdr = list(wp.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        ti = next((i for i, v in enumerate(hdr) if v and 'Total INGRESOS' in str(v)), None)
        if ti is not None:
            for r in wp.iter_rows(min_row=3, values_only=True):
                if r[0] in (None, '') or str(r[0]).startswith('Total'): continue
                mm = motor.mnum(r[0])
                if mm: pas[mm] += motor.num(r[ti])
    filas = []
    for m in range(1, 13):
        row = {'Mes Nº': m}
        for k in ROL_MES_CONC: row[k] = round(agg[m][k], 2)
        row['PAS'] = round(pas[m], 2)
        filas.append(row)
    return pd.DataFrame(filas).set_index('Mes Nº')


def _rol_base_mes(rol_path):
    rows = _rol_rows(rol_path); a = collections.defaultdict(float)
    for m, jub, r in rows:
        if not jub:
            a[m] += sum(motor.num(r[i]) for i in GRAV_IDX)
    return {m: round(a[m], 2) for m in range(1, 13)}


def _g_cols(R):
    det = R['det']
    return {
        'd13b': _por_mes(det, lambda d: d['P_bruto']),
        'D13 mens. pagado (rol)': _por_mes(det, lambda d: d['z']),
        'd13f': _por_mes(det, lambda d: d['P'] if not d['activo'] else 0.0),
        'd13m': _por_mes(det, lambda d: d['P_mens']),
        'd14b': _por_mes(det, lambda d: d['Q_bruto']),
        'D14 mens. pagado (rol)': _por_mes(det, lambda d: d['aa']),
        'd14f': _por_mes(det, lambda d: d['Q'] if not d['activo'] else 0.0),
        'd14m': _por_mes(det, lambda d: d['Q_mens']),
        'Vacaciones': _por_mes(det, lambda d: d['R']),
        'Vac. gozadas (rol)': _por_mes(det, lambda d: d['vac_gozada']),
        'Vac. finiquito (rol)': _por_mes(det, lambda d: d['vac_finiquito']),
        'D13 por pagar (activos)': _por_mes(det, lambda d: d['AJ']),
        'D14 por pagar (activos)': _por_mes(det, lambda d: d['AK']),
    }


def _g2_extra(R):
    det = R['det']
    return {
        'pat_emp': _por_mes(det, lambda d: (d['pat'] + d['ccc']) if d['X'] else 0.0),
        'pat_pas': _por_mes(det, lambda d: 0.0),
        'pat_only_pas': _por_mes(det, lambda d: 0.0),
        'Fondos de reserva': _por_mes(det, lambda d: d['S']),
        'FR excluido': _por_mes(det, lambda d: d['AU']),
        'FR pagado en rol': _por_mes(det, lambda d: d['ab']),
        'FR a depositar IESS': _por_mes(det, lambda d: d['AI']),
        'pas_aporte': _por_mes(det, lambda d: d['ind'] if not d['X'] else 0.0),
    }


def _df_mes(colmap):
    filas = []
    for m in range(1, 13):
        row = {'Mes': float(m)}
        for c, serie in colmap.items(): row[c] = serie[m]
        filas.append(row)
    return pd.DataFrame(filas).set_index('Mes')


def g_df(R):
    return _df_mes(_g_cols(R))


def g2_df(R):
    c = _g_cols(R); c.update(_g2_extra(R))
    return _df_mes(c)


def g3_df(R, rol_path):
    det = R['det']
    c = _g_cols(R); c.update(_g2_extra(R))
    c['base_pas'] = _por_mes(det, lambda d: d['base'] if not d['X'] else 0.0)
    c['base_emp'] = _por_mes(det, lambda d: d['base'] if d['X'] else 0.0)
    c['rol_base'] = _rol_base_mes(rol_path)
    return _df_mes(c)


def cal_base_df(R):
    det = R['det']
    agg = collections.defaultdict(lambda: dict(base12=0.0, meses=0, base_dic=0.0, ced=None, nom=None, activo=False, ingreso=None))
    for d in det:
        a = agg[d['nom']]; a['ced'] = d['ced']; a['nom'] = d['nom']
        if d['base'] > 0: a['meses'] += 1
        if d['X']:
            a['base12'] += d['base']
            if d['mes'] == 12: a['base_dic'] += d['base']
            a['activo'] = d['activo']; a['ingreso'] = d['ingreso']
    filas = []
    for a in agg.values():
        if not a['activo']: continue
        m = a['meses'] or 1
        filas.append(dict(Cédula=a['ced'], Nombre=a['nom'], ingreso=pd.Timestamp(a['ingreso']) if a['ingreso'] else pd.NaT,
                          base12=round(a['base12'], 2), vac_prov=round(a['base12'] / 24, 6),
                          base_dic=round(a['base_dic'], 2), meses=a['meses'], prom=round(a['base12'] / m, 2)))
    filas.sort(key=lambda x: x['Nombre'])
    return pd.DataFrame(filas)


def pop_acum_df(R):
    det = R['det']
    agg = collections.defaultdict(lambda: dict(d13=0.0, d14=0.0, meses=0, ced=None, nom=None, activo=False))
    for d in det:
        if d['P'] > 0:
            a = agg[d['ced']]; a['d13'] += d['P']; a['d14'] += d['Q']; a['meses'] += 1
            a['ced'] = d['ced']; a['nom'] = d['nom']; a['activo'] = d['activo']
    filas = [dict(Cédula=a['ced'], Nombre=a['nom'], d13=round(a['d13'], 2), meses=a['meses'],
                  d14=round(a['d14'], 2), **{'Activo al 31-dic': 'Sí' if a['activo'] else 'No'})
             for a in agg.values()]
    filas.sort(key=lambda x: x['Nombre'])
    return pd.DataFrame(filas)


def exportar_pickles(R, rutas):
    """Escribe en ``rutas.home`` los pickles que leen los generadores de referencia."""
    import os
    os.makedirs(rutas.home, exist_ok=True)
    consol_df(R).to_pickle(f'{rutas.home}/consol.pkl')
    rol_mes_df(rutas.rol).to_pickle(f'{rutas.home}/rol_mes.pkl')
    g_df(R).to_pickle(f'{rutas.home}/g.pkl')
    g2_df(R).to_pickle(f'{rutas.home}/g2.pkl')
    g3_df(R, rutas.rol).to_pickle(f'{rutas.home}/g3.pkl')
    cal_base_df(R).to_pickle(f'{rutas.home}/cal_base.pkl')
    pop_acum_df(R).to_pickle(f'{rutas.home}/pop_acum.pkl')
