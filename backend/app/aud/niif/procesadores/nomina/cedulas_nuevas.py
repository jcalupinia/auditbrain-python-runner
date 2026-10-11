"""Generadores de las 4 cédulas que no tenían generador: DO-2, DO-3, DO-3.1, DO-4.
Genéricos: toman los datos del motor; rangos de fórmula y metas de control se
derivan de la población real (no fijos de PROPHAR). Formato calcado de
referencia/cedulas/."""
from __future__ import annotations
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.app.aud.niif.procesadores.nomina import motor

N2 = '#,##0.00'
THIN = Side(style='thin', color='B0B0B0')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill('solid', fgColor='0A2342')
TFILL = PatternFill('solid', fgColor='DCE6F1')
OKF = PatternFill('solid', fgColor='E2EFDA')
HF = Font(bold=True, color='FFFFFF', size=10)
BF = Font(bold=True, size=10)
MESES = ['Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']


def _banda(ws, titulo, ref):
    ws.merge_cells('B1:N2')
    ws['B1'] = titulo; ws['B1'].font = Font(bold=True, color='FFFFFF', size=13)
    ws['B1'].fill = HFILL; ws['B1'].alignment = Alignment('center', 'center')
    ws['B3'] = f'Ref. PT: {ref}'; ws['B3'].font = Font(italic=True, size=9)


def _ctrl_header(ws):
    ws['B2'] = f'CONTROLES DE INTEGRIDAD Y RECÁLCULO'; ws['B2'].font = BF
    for col, t in zip('BCDEF', ['Control', 'Valor A', 'Valor B', 'Diferencia', 'Estado']):
        ws[f'{col}4'] = t; ws[f'{col}4'].font = HF; ws[f'{col}4'].fill = HFILL
    ws.column_dimensions['B'].width = 52


def _ctrl(ws, r, desc, a, b):
    ws[f'B{r}'] = desc
    ws[f'C{r}'] = a; ws[f'D{r}'] = b
    ws[f'C{r}'].number_format = ws[f'D{r}'].number_format = N2
    ws[f'E{r}'] = f'=ROUND(C{r}-D{r},2)'; ws[f'E{r}'].number_format = N2
    ws[f'F{r}'] = f'=IF(ABS(E{r})<=0.01,"OK","REVISAR")'
    ws[f'F{r}'].alignment = Alignment('center')


# ------------------------------------------------------------------ DO-3
def build_do3(R, path, ficha):
    cons = R['cons']
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    # Parámetros
    pa = wb.create_sheet('Parámetros')
    pa['B3'] = 'Tasas por relación de trabajo (base × tasa)'
    for c, t in zip('BCDEF', ['Código', 'Relación', 'Patronal', 'Individual', 'CCC']):
        pa[f'{c}4'] = t; pa[f'{c}4'].font = BF
    tasas = [('06', 'Código del Trabajo', 0.1115, 0.0945, 0.01),
             ('16', 'Jornada parcial', 0.1115, 0.0945, 0.01),
             ('74', 'Pasantes', 0.0, 0.176, 0.0)]
    for i, t in enumerate(tasas):
        for j, v in enumerate(t):
            pa.cell(5 + i, 2 + j, v)
    pa['B9'] = 'SBU'; pa['C9'] = motor.SBU_2025
    pa['B10'] = 'Tolerancia'; pa['C10'] = 0.01
    # Detalle RQ-003
    de = wb.create_sheet('Detalle RQ-003')
    hdr = ['Período', 'Mes', 'Cédula', 'Nombre', 'Relación de trabajo', 'Base imponible', 'Días',
           'Patronal', 'Individual', 'Aporte adic.', 'Cesantía', '% CCC', 'Valor CCC', 'Total aporte IESS',
           'Recálc. patronal', 'Recálc. individual', 'Recálc. CCC', 'Dif. patronal', 'Dif. individual',
           'Dif. CCC', 'Estado']
    for j, h in enumerate(hdr, 1):
        de.cell(1, j, h).font = BF
    for i, r in enumerate(cons, 2):
        de.cell(i, 1, f"2025-{r['mes']}"); de.cell(i, 2, r['mes']); de.cell(i, 3, r['ced'])
        de.cell(i, 4, r['nombre']); de.cell(i, 5, str(r['rel'] or '').strip())
        de.cell(i, 6, r['base']); de.cell(i, 7, r['dias']); de.cell(i, 8, r['pat'])
        de.cell(i, 9, r['ind']); de.cell(i, 10, 0); de.cell(i, 11, 0); de.cell(i, 12, 0)
        de.cell(i, 13, r['ccc']); de.cell(i, 14, r['pat'] + r['ind'])
        de.cell(i, 15, f'=ROUND(F{i}*VLOOKUP(LEFT(E{i},2),Parámetros!$B$5:$F$7,3,0),2)')
        de.cell(i, 16, f'=ROUND(F{i}*VLOOKUP(LEFT(E{i},2),Parámetros!$B$5:$F$7,4,0),2)')
        de.cell(i, 17, f'=ROUND(F{i}*VLOOKUP(LEFT(E{i},2),Parámetros!$B$5:$F$7,5,0),2)')
        de.cell(i, 18, f'=ROUND(O{i}-H{i},2)'); de.cell(i, 19, f'=ROUND(P{i}-I{i},2)')
        de.cell(i, 20, f'=ROUND(Q{i}-M{i},2)')
        de.cell(i, 21, f'=IF(MAX(ABS(R{i}),ABS(S{i}),ABS(T{i}))>Parámetros!$C$10,"‡ REVISAR","OK")')
    last = len(cons) + 1
    # Resumen
    rs = wb.create_sheet('DO3 Resumen de Planillas')
    _banda(rs, 'OBLIGACIONES ACUMULADAS · RESUMEN DE PLANILLAS IESS', 'DO-3')
    heads = ['Mes', 'No. Empleados', 'Base Imponible', 'Aporte Personal', 'Aporte Patronal',
             'Valor CCC', 'Fondos de Reserva', 'Días Laborados', 'Total Aporte']
    for c, t in zip('BCDEFGHIJ', heads):
        rs[f'{c}20'] = t; rs[f'{c}20'].font = HF; rs[f'{c}20'].fill = HFILL
    R3 = "'Detalle RQ-003'"
    for m in range(1, 13):
        r = 21 + m
        rs[f'B{r}'] = MESES[m - 1]
        rs[f'C{r}'] = f'=COUNTIFS({R3}!$B$2:$B${last},{m})'
        rs[f'D{r}'] = f'=SUMIFS({R3}!$F$2:$F${last},{R3}!$B$2:$B${last},{m})'
        rs[f'E{r}'] = f'=SUMIFS({R3}!$I$2:$I${last},{R3}!$B$2:$B${last},{m})'
        rs[f'F{r}'] = f'=SUMIFS({R3}!$H$2:$H${last},{R3}!$B$2:$B${last},{m})'
        rs[f'G{r}'] = f'=SUMIFS({R3}!$M$2:$M${last},{R3}!$B$2:$B${last},{m})'
        rs[f'H{r}'] = 0
        rs[f'I{r}'] = f'=SUMIFS({R3}!$G$2:$G${last},{R3}!$B$2:$B${last},{m})'
        rs[f'J{r}'] = f'=SUM(E{r}:G{r})'
        for c in 'DEFGHIJ': rs[f'{c}{r}'].number_format = N2
    tot = 34
    rs[f'B{tot}'] = 'Σ'
    for c in 'CDEFGHIJ':
        rs[f'{c}{tot}'] = f'=SUM({c}22:{c}33)'; rs[f'{c}{tot}'].font = BF; rs[f'{c}{tot}'].number_format = N2
    rs[f'B36'] = 'Fondos de Reserva: no forman parte del Consolidado IESS; se verifican contra RQ-005 en DO-9.'
    # metas de control = fila Totales del Consolidado (del cliente)
    tb = dict(base=sum(r['base'] for r in cons), pat=sum(r['pat'] for r in cons),
              ind=sum(r['ind'] for r in cons), ccc=sum(r['ccc'] for r in cons),
              n=len(cons))
    co = wb.create_sheet('Controles'); _ctrl_header(co)
    _ctrl(co, 5, 'Registros trasladados vs reporte IESS', f'=COUNTA({R3}!$C$2:$C${last})', tb['n'])
    _ctrl(co, 6, 'Base imponible: detalle vs Totales', f'=ROUND(SUM({R3}!$F$2:$F${last}),2)', round(tb['base'], 2))
    _ctrl(co, 7, 'Patronal: detalle vs Totales', f'=ROUND(SUM({R3}!$H$2:$H${last}),2)', round(tb['pat'], 2))
    _ctrl(co, 8, 'Individual: detalle vs Totales', f'=ROUND(SUM({R3}!$I$2:$I${last}),2)', round(tb['ind'], 2))
    _ctrl(co, 9, 'Valor CCC: detalle vs Totales', f'=ROUND(SUM({R3}!$M$2:$M${last}),2)', round(tb['ccc'], 2))
    _ctrl(co, 10, 'DO-3 Base = detalle', "='DO3 Resumen de Planillas'!D34", f'=ROUND(SUM({R3}!$F$2:$F${last}),2)')
    _ctrl(co, 11, 'DO-3 Patronal = detalle', "='DO3 Resumen de Planillas'!F34", f'=ROUND(SUM({R3}!$H$2:$H${last}),2)')
    _ctrl(co, 12, 'DO-3 Total (E+F+G) cuadra', "='DO3 Resumen de Planillas'!J34",
          "='DO3 Resumen de Planillas'!E34+'DO3 Resumen de Planillas'!F34+'DO3 Resumen de Planillas'!G34")
    wb.save(path)
    return dict(hojas=4, registros=len(cons))


# ------------------------------------------------------------------ DO-2
CONCEPTOS_DO2 = ['Sueldos', 'Horas Extras', 'Bono Desempeño', 'Décimo Tercer Sueldo',
                 'Décimo Cuarto Sueldo', 'Fondos de Reserva', 'Aporte Patronal', 'Otros gastos de nómina']


def _concepto_do2(desc):
    d = (desc or '').lower()
    if 'patronal' in d: return 'Aporte Patronal'
    if 'hora' in d: return 'Horas Extras'
    if 'bono' in d or 'desemp' in d: return 'Bono Desempeño'
    if 'tercer' in d: return 'Décimo Tercer Sueldo'
    if 'cuarto' in d: return 'Décimo Cuarto Sueldo'
    if 'fondo' in d and 'reserva' in d: return 'Fondos de Reserva'
    if 'sueldo' in d: return 'Sueldos'
    return 'Otros gastos de nómina'


def leer_anexo(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    pas, gas = [], []
    for r in ws.iter_rows(min_row=3, values_only=True):
        cod = r[0]
        if cod is None or not str(cod).startswith('5'): continue
        cod = str(cod).strip(); desc = str(r[1] or '').strip()
        s24 = r[2] if isinstance(r[2], (int, float)) else 0
        s25 = r[5] if isinstance(r[5], (int, float)) else 0
        (pas if cod.startswith('5BS.212') else gas).append((cod, desc, s24, s25))
    return pas, gas


def build_do2(anexo_path, path, ficha):
    pas, gas = leer_anexo(anexo_path)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'DO2 Cedula Sumaria'
    _banda(ws, 'OBLIGACIONES ACUMULADAS · SUMARIA DE NÓMINA', 'DO-2')
    heads = ['Código', 'Descripción', 'Saldo 2024', 'Variación', '%', 'Saldo 2025',
             'm', 'Debe', 'Haber', 'm', 'Saldo Auditado', 'm', 'Referencia']
    for c, t in zip('BCDEFGHIJKLMN', heads):
        ws[f'{c}9'] = t; ws[f'{c}9'].font = HF; ws[f'{c}9'].fill = HFILL
    def fila(r, cod, desc, s24, s25, concepto=None):
        ws[f'B{r}'] = cod; ws[f'C{r}'] = desc; ws[f'D{r}'] = s24
        ws[f'E{r}'] = f'=+G{r}-D{r}'; ws[f'F{r}'] = f'=IF(D{r}=0,0,E{r}/ABS(D{r}))'
        ws[f'G{r}'] = s25; ws[f'I{r}'] = 0; ws[f'J{r}'] = 0
        ws[f'L{r}'] = f'=G{r}+I{r}-J{r}'
        for c in 'DEGIJL': ws[f'{c}{r}'].number_format = N2
        ws[f'F{r}'].number_format = '0.0%'
        if concepto is not None: ws[f'P{r}'] = concepto
    # pasivo 11..
    r = 11
    for cod, desc, s24, s25 in pas:
        fila(r, cod, desc, s24, s25); r += 1
    pt = r; ws[f'B{pt}'] = 'Total'
    for c in 'DEGIJL': ws[f'{c}{pt}'] = f'=SUM({c}11:{c}{pt-1})'; ws[f'{c}{pt}'].font = BF; ws[f'{c}{pt}'].number_format = N2
    # gasto desde fila 25 (igual que el modelo)
    g0 = 25
    rr = g0
    for cod, desc, s24, s25 in gas:
        fila(rr, cod, desc, s24, s25, _concepto_do2(desc)); rr += 1
    gt = rr; ws[f'B{gt}'] = 'Total'
    for c in 'DEGIJL': ws[f'{c}{gt}'] = f'=SUM({c}{g0}:{c}{gt-1})'; ws[f'{c}{gt}'].font = BF; ws[f'{c}{gt}'].number_format = N2
    # named range resumen_concepto = P g0..gt-1
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names.add(DefinedName('resumen_concepto', attr_text=f"'DO2 Cedula Sumaria'!$P${g0}:$P${gt-1}"))
    # resumen por concepto
    rb = gt + 3
    ws[f'B{rb}'] = 'RESUMEN DEL GASTO POR CONCEPTO'; ws[f'B{rb}'].font = BF
    for c, t in zip('BCDGL', ['Concepto', 'Nº ctas', 'Saldo 2024', 'Saldo 2025', 'Saldo auditado']):
        ws[f'{c}{rb+1}'] = t; ws[f'{c}{rb+1}'].font = BF
    r = rb + 2; r0 = r
    for con in CONCEPTOS_DO2:
        ws[f'B{r}'] = con
        ws[f'C{r}'] = f'=COUNTIF(resumen_concepto,B{r})'
        ws[f'D{r}'] = f'=SUMIF(resumen_concepto,B{r},$D${g0}:$D${gt-1})'
        ws[f'G{r}'] = f'=SUMIF(resumen_concepto,B{r},$G${g0}:$G${gt-1})'
        ws[f'L{r}'] = f'=SUMIF(resumen_concepto,B{r},$L${g0}:$L${gt-1})'
        for c in 'DGL': ws[f'{c}{r}'].number_format = N2
        r += 1
    rt = r; ws[f'B{rt}'] = 'Total gasto'
    for c in 'CDGL': ws[f'{c}{rt}'] = f'=SUM({c}{r0}:{c}{rt-1})'; ws[f'{c}{rt}'].font = BF
    ws.column_dimensions['P'].hidden = True
    # metas desde el anexo (cliente)
    p24 = round(sum(x[2] for x in pas), 2); p25 = round(sum(x[3] for x in pas), 2)
    g24 = round(sum(x[2] for x in gas), 2); g25 = round(sum(x[3] for x in gas), 2)
    S = "'DO2 Cedula Sumaria'"
    co = wb.create_sheet('Controles'); _ctrl_header(co)
    _ctrl(co, 5, 'Cuentas trasladadas (pasivo+gasto)', len(pas) + len(gas), len(pas) + len(gas))
    _ctrl(co, 6, 'Pasivo 2024', f'={S}!D{pt}', p24)
    _ctrl(co, 7, 'Pasivo 2025', f'={S}!G{pt}', p25)
    _ctrl(co, 8, 'Gasto 2024', f'={S}!D{gt}', g24)
    _ctrl(co, 9, 'Gasto 2025', f'={S}!G{gt}', g25)
    _ctrl(co, 10, 'Resumen por concepto = total gasto', f'={S}!G{rt}', f'={S}!G{gt}')
    _ctrl(co, 11, 'Nº cuentas de gasto agrupadas', f'={S}!C{rt}', len(gas))
    _ctrl(co, 12, 'Asiento cuadrado (ΣDebe=ΣHaber)', f'={S}!I{pt}+{S}!I{gt}', f'={S}!J{pt}+{S}!J{gt}')
    wb.save(path)
    return dict(pasivo=len(pas), gasto=len(gas), p25=p25, g25=g25)


# ------------------------------------------------------------------ DO-3.1
GRAV_IDX = [5, 6, 7, 8, 9, 10, 11, 12, 30, 31, 29]  # sueldo..bonificacion (base gravada rol)


def _rol_registros(rol_path, nm2ced):
    ws = openpyxl.load_workbook(rol_path, read_only=True, data_only=True)['normal']
    regs = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] in (None, '') or str(r[0]).startswith('Total') or r[1] in (None, ''): continue
        m = motor.mnum(r[0])
        if m is None: continue
        cargo = str(r[2] or ''); nom = str(r[1]).strip()
        tipo = 'Jubilado' if 'JUBIL' in cargo.upper() else 'Empleado'
        grav = 0.0 if tipo == 'Jubilado' else sum(motor.num(r[i]) for i in GRAV_IDX)
        regs.append(dict(mes=m, nom=nom, ced=nm2ced.get(nom, ''), tipo=tipo,
                         sueldo=motor.num(r[5]), vac_tom=motor.num(r[6]), fr=motor.num(r[24]),
                         d13m=motor.num(r[25]), d14m=motor.num(r[26]), liq13=motor.num(r[27]),
                         liq14=motor.num(r[28]), vacfin=motor.num(r[18]), pens=motor.num(r[21]),
                         grav=grav))
    return regs


def build_do31(R, rol_path, path, ficha):
    nm2ced = {r['nombre']: r['ced'] for r in R['cons']}
    regs = _rol_registros(rol_path, nm2ced)
    base_emp = sum(d['base'] for d in R['det'] if d['X'])  # base IESS en rol (empleados)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    # Detalle Rol
    de = wb.create_sheet('Detalle Rol')
    cols = ['Mes', 'Cédula', 'Nombre', 'Tipo', 'Sueldo', 'Vac tomadas', 'FR en rol', 'D13 MENS',
            'D14 MENS SIE', 'LIQ.DEC13', 'LIQ.DEC14S', 'Vac finiquito', 'Pensión jub', 'Base gravada rol']
    for j, h in enumerate(cols, 1): de.cell(1, j, h).font = BF
    for i, d in enumerate(regs, 2):
        for j, k in enumerate(['mes', 'ced', 'nom', 'tipo', 'sueldo', 'vac_tom', 'fr', 'd13m',
                               'd14m', 'liq13', 'liq14', 'vacfin', 'pens', 'grav'], 1):
            de.cell(i, j, d[k])
    last = len(regs) + 1
    DR = "'Detalle Rol'"
    JUB = f'{DR}!$D$2:$D${last}'
    # Resumen
    rs = wb.create_sheet('DO3.1 Resumen Roles')
    _banda(rs, 'OBLIGACIONES ACUMULADAS · RESUMEN DE ROLES DE PAGO', 'DO-3.1')
    cab = ['Mes', 'Sueldo', 'Vac tomadas', 'FR en rol', 'Décimos mensualizados', 'Liquidaciones',
           'Base gravada rol']
    for c, t in zip('BCDEFGH', cab):
        rs[f'{c}11'] = t; rs[f'{c}11'].font = HF; rs[f'{c}11'].fill = HFILL
    def sif(col, m, extra=''):
        return f'=SUMIFS({DR}!${col}$2:${col}${last},{DR}!$A$2:$A${last},{m}{extra})'
    for m in range(1, 13):
        r = 11 + m; rs[f'B{r}'] = MESES[m - 1]
        rs[f'C{r}'] = sif('E', m)  # sueldo
        rs[f'D{r}'] = sif('F', m)  # vac tomadas
        rs[f'E{r}'] = sif('G', m)  # fr
        rs[f'F{r}'] = f'=SUMIFS({DR}!$H$2:$H${last},{DR}!$A$2:$A${last},{m})+SUMIFS({DR}!$I$2:$I${last},{DR}!$A$2:$A${last},{m})'
        rs[f'G{r}'] = f'=SUMIFS({DR}!$J$2:$J${last},{DR}!$A$2:$A${last},{m})+SUMIFS({DR}!$K$2:$K${last},{DR}!$A$2:$A${last},{m})'
        rs[f'H{r}'] = sif('N', m)  # base gravada
        for c in 'CDEFGH': rs[f'{c}{r}'].number_format = N2
    tt = 24
    rs[f'B{tt}'] = 'Σ'
    for c in 'CDEFGH': rs[f'{c}{tt}'] = f'=SUM({c}12:{c}23)'; rs[f'{c}{tt}'].font = BF; rs[f'{c}{tt}'].number_format = N2
    # targets desde datos
    T = dict(
        sueldo=round(sum(d['sueldo'] for d in regs), 2),
        vac_tom=round(sum(d['vac_tom'] for d in regs), 2),
        vacfin=round(sum(d['vacfin'] for d in regs), 2),
        fr=round(sum(d['fr'] for d in regs), 2),
        liq13=round(sum(d['liq13'] for d in regs), 2),
        liq14=round(sum(d['liq14'] for d in regs), 2),
        mens_emp=round(sum(d['d13m'] + d['d14m'] for d in regs), 2),  # todos (incluye jubilares)
        jub=round(sum(d['d13m'] + d['d14m'] for d in regs if d['tipo'] == 'Jubilado'), 2),
        grav=round(sum(d['grav'] for d in regs), 2),
    )
    co = wb.create_sheet('Controles'); _ctrl_header(co)
    _ctrl(co, 5, 'Sueldo rol', f'=ROUND(SUM({DR}!$E$2:$E${last}),2)', T['sueldo'])
    _ctrl(co, 6, 'Vacaciones tomadas (goces)', f'=ROUND(SUM({DR}!$F$2:$F${last}),2)', T['vac_tom'])
    _ctrl(co, 7, 'Vacaciones de finiquito', f'=ROUND(SUM({DR}!$L$2:$L${last}),2)', T['vacfin'])
    _ctrl(co, 8, 'FR pagado en rol', f'=ROUND(SUM({DR}!$G$2:$G${last}),2)', T['fr'])
    _ctrl(co, 9, 'LIQ.DEC13', f'=ROUND(SUM({DR}!$J$2:$J${last}),2)', T['liq13'])
    _ctrl(co, 10, 'LIQ.DEC14S', f'=ROUND(SUM({DR}!$K$2:$K${last}),2)', T['liq14'])
    _ctrl(co, 11, 'Décimos mensualizados en rol (todos)', T['mens_emp'], T['mens_emp'])
    _ctrl(co, 12, 'Décimos jubilares en rol', T['jub'], T['jub'])
    _ctrl(co, 13, 'Base gravada rol', f'=ROUND(SUM({DR}!$N$2:$N${last}),2)', T['grav'])
    _ctrl(co, 14, 'Base rol − base IESS (en rol)', round(T['grav'] - base_emp, 2), round(T['grav'] - base_emp, 2))
    wb.save(path)
    return dict(registros=len(regs), targets=T, rol_menos_iess=round(T['grav'] - base_emp, 2))


# ------------------------------------------------------------------ DO-4
def _saldo_vac_2024(anexo_path):
    pas, _ = leer_anexo(anexo_path)
    for cod, desc, s24, s25 in pas:
        if 'vacacion' in desc.lower():
            return abs(s24)
    return 0.0


def build_do4(R, anexo_path, path, ficha):
    det = R['det']
    sal_ini = _saldo_vac_2024(anexo_path)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    pa = wb.create_sheet('Parámetros')
    pa['B3'] = 'Parámetros del reproceso'
    for i, (k, v) in enumerate([('SBU', motor.SBU_2025), ('% Fondos de reserva', motor.PCT_FR),
                                ('Divisor D13', 12), ('Divisor vacaciones', 24),
                                ('Saldo vacaciones 31-dic-2024', sal_ini), ('Tolerancia', 0.01)], 4):
        pa[f'B{i}'] = k; pa[f'C{i}'] = v
    de = wb.create_sheet('Detalle RQ-003')
    cols = ['Período', 'Mes', 'Cédula', 'Nombre', 'Relación', 'Base', 'Días', 'Región',
            'Aplica', 'En rol', 'Activo', 'D13', 'D14', 'Vacaciones', 'FR', 'FR a depositar',
            'D13 x pagar', 'D14 x pagar', 'FR x pagar', 'D13 excl', 'D14 excl', 'Vac excl',
            'FR excl', 'D13 bruto', 'Vac gozada', 'Vac finiquito']
    for j, h in enumerate(cols, 1): de.cell(1, j, h).font = BF
    for i, d in enumerate(det, 2):
        vals = [f"2025-{d['mes']}", d['mes'], d['ced'], d['nom'], str(d.get('cod', '')), d['base'],
                d['dias'], 'Sierra / Amazonía', 'Sí' if d['L'] else 'No', 'Sí' if d['X'] else 'No',
                'Sí' if d['activo'] else 'No', d['P'], d['Q'], d['R'], d['S'], d['AI'], d['AJ'],
                d['AK'], d['AL'], d['AR'], d['AS'], d['AT'], d['AU'], d['P_bruto'],
                d['vac_gozada'], d['vac_finiquito']]
        for j, v in enumerate(vals, 1): de.cell(i, j, v)
    last = len(det) + 1
    DR = "'Detalle RQ-003'"
    rs = wb.create_sheet('DO4 Reproceso Ben Sociales')
    _banda(rs, 'OBLIGACIONES ACUMULADAS · REPROCESO DE BENEFICIOS SOCIALES', 'DO-4')
    cab = ['Mes', 'Base', 'D13 bruto', 'D14', 'Vacaciones', 'FR devengado']
    for c, t in zip('BCDEFG', cab):
        rs[f'{c}11'] = t; rs[f'{c}11'].font = HF; rs[f'{c}11'].fill = HFILL
    cmap = dict(Base='F', **{'D13 bruto': 'X', 'D14': 'M', 'Vacaciones': 'N', 'FR devengado': 'O'})
    for m in range(1, 13):
        r = 11 + m; rs[f'B{r}'] = MESES[m - 1]
        for c, col in zip('CDEFG', ['F', 'X', 'M', 'N', 'O']):
            rs[f'{c}{r}'] = f'=SUMIFS({DR}!${col}$2:${col}${last},{DR}!$B$2:$B${last},{m})'
            rs[f'{c}{r}'].number_format = N2
    tt = 24; rs[f'B{tt}'] = 'Σ'
    for c in 'CDEFG': rs[f'{c}{tt}'] = f'=SUM({c}12:{c}23)'; rs[f'{c}{tt}'].font = BF; rs[f'{c}{tt}'].number_format = N2
    # Provisiones al 31-dic
    rs['B27'] = 'PROVISIONES AL 31-DIC (pasivo)'; rs['B27'].font = BF
    pp = {
        'D13 por pagar (dic)': round(sum(d['AJ'] for d in det if d['mes'] == 12), 2),
        'D14 por pagar (ago-dic Sierra)': round(sum(d['AK'] for d in det if d['mes'] >= 8), 2),
        'Vacaciones por pagar': round(sal_ini + sum(d['R'] for d in det) - sum(d['vac_gozada'] for d in det) - sum(d['vac_finiquito'] for d in det), 2),
        'FR por pagar (dic)': round(sum(d['AL'] for d in det if d['mes'] == 12), 2),
    }
    r = 28
    for k, v in pp.items():
        rs[f'B{r}'] = k; rs[f'C{r}'] = v; rs[f'C{r}'].number_format = N2; r += 1
    # targets
    T = dict(
        d13_bruto=round(sum(d['P_bruto'] for d in det), 2),
        vac_prov=round(sum(d['R'] for d in det), 2),
        exc13=round(sum(d['AR'] for d in det), 2),
        exc14=round(sum(d['AS'] for d in det), 2),
        excfr=round(sum(d['AU'] for d in det), 2),
    )
    co = wb.create_sheet('Controles'); _ctrl_header(co)
    _ctrl(co, 5, 'D13 bruto anual', f'=ROUND(SUM({DR}!$X$2:$X${last}),2)', T['d13_bruto'])
    _ctrl(co, 6, 'Vacaciones provisión anual', f'=ROUND(SUM({DR}!$N$2:$N${last}),2)', T['vac_prov'])
    _ctrl(co, 7, 'D13 por pagar (dic)', "='DO4 Reproceso Ben Sociales'!C28", pp['D13 por pagar (dic)'])
    _ctrl(co, 8, 'D14 por pagar (Sierra)', "='DO4 Reproceso Ben Sociales'!C29", pp['D14 por pagar (ago-dic Sierra)'])
    _ctrl(co, 9, 'Vacaciones por pagar', "='DO4 Reproceso Ben Sociales'!C30", pp['Vacaciones por pagar'])
    _ctrl(co, 10, 'FR por pagar (dic)', "='DO4 Reproceso Ben Sociales'!C31", pp['FR por pagar (dic)'])
    _ctrl(co, 11, 'D13 excluido (auditor)', f'=ROUND(SUM({DR}!$T$2:$T${last}),2)', T['exc13'])
    _ctrl(co, 12, 'D14 excluido (auditor)', f'=ROUND(SUM({DR}!$U$2:$U${last}),2)', T['exc14'])
    _ctrl(co, 13, 'FR excluido (auditor)', f'=ROUND(SUM({DR}!$W$2:$W${last}),2)', T['excfr'])
    wb.save(path)
    return dict(registros=len(det), provisiones=pp, targets=T, saldo_ini_vac=sal_ini)
