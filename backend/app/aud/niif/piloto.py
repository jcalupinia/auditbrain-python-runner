"""Servicio del agente guía «NIIF Piloto» para la plataforma AUDIT-IA.

Un solo servicio para TODAS las pruebas NIIF del Command Center (las 20 herramientas
del catálogo + la planificación NIA). No hay un servicio por prueba: el catálogo se
lee en runtime de ``procesadores.PROCESADORES`` y el cálculo lo hace el procesador
determinista que ya existe. Este módulo solo:

  1. dice qué pruebas hay (``listar``),
  2. dice qué datos necesita cada una (``requisitos``),
  3. prepara el cálculo con los datos del cliente (``preparar``),
  4. arma el papel en cada formato (``papel``) y
  5. verifica empíricamente el Excel (``verificar_excel``, regla suprema del CLAUDE.md).

Reutiliza el contrato oficial (``mod.ejecutar``, ``mod.definicion``,
``libro.{xlsx,html,docx,pptx,pdf}`` y ``datos_cliente.con_datos``), el mismo que usan
el router del ciclo y ``scripts/papeles_muestra.py``. El cálculo NUNCA usa LLM.

Lo consumen: el router HTTP ``piloto_router`` (plataforma) y la CLI
``scripts/piloto/piloto.py`` (sesiones de Claude Code). Una sola fuente de verdad.
"""
from __future__ import annotations

import io

from backend.app.aud.niif import ejercicio_modelo, procesadores
from backend.app.aud.niif.procesadores import datos_cliente, libro

# (extensión, nombre de la función en libro, etiqueta legible, media type).
FORMATOS = (
    ("xlsx", "xlsx", "Excel con fórmulas",
     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("html", "html", "HTML autónomo", "text/html; charset=utf-8"),
    ("docx", "docx", "Word",
     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("pptx", "pptx", "PowerPoint",
     "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    ("pdf", "pdf", "PDF", "application/pdf"),
)
TIPOS = {ext: mime for ext, _fn, _etq, mime in FORMATOS}
_FN = {ext: fn for ext, fn, _etq, _mime in FORMATOS}
_ETIQUETAS = {ext: etq for ext, _fn, etq, _mime in FORMATOS}

# Prefijos que Excel podría interpretar como fórmula en una celda de TEXTO.
_PREFIJOS_PELIGROSOS = ("=", "+", "-", "@")


class PruebaDesconocida(ValueError):
    """El id no corresponde a ningún procesador registrado."""


def _mod(pid: str):
    if pid not in procesadores.PROCESADORES:
        raise PruebaDesconocida(f"Prueba desconocida: «{pid}».")
    return procesadores.PROCESADORES[pid]


def _campos(mod) -> dict:
    return dict(getattr(mod, "CAMPOS", {}) or {})


# --------------------------------------------------------------------------- catálogo
def resumen_prueba(pid: str, mod=None) -> dict:
    mod = mod or _mod(pid)
    d = mod.definicion()
    campos = _campos(mod)
    return {
        "id": pid,
        "rubro": getattr(mod, "RUBRO", "") or "",
        "nombre": d.get("name", pid),
        "area": d.get("area", ""),
        "marcos": d.get("frameworks", []),
        "principal": getattr(mod, "PRINCIPAL", ""),
        "datasets": {ds: len(cs) for ds, cs in campos.items()},
        "parametros": sorted((getattr(mod, "PARAMETROS", {}) or {}).keys()),
        "n_requerimientos": len(d.get("requests", [])),
        "tiene_ejemplo": ejercicio_modelo.disponible(mod),
    }


def listar() -> list[dict]:
    """Todas las pruebas disponibles, en el orden de la matriz del socio."""
    return [resumen_prueba(pid, mod) for pid, mod in procesadores.PROCESADORES.items()]


def requisitos(pid: str) -> dict:
    """Qué datos necesita una prueba: datasets con sus columnas, parámetros del
    auditor (con su valor por defecto) y requerimientos formales al cliente."""
    mod = _mod(pid)
    d = mod.definicion()
    campos = _campos(mod)
    datasets = [
        {
            "dataset": ds,
            "es_principal": ds == getattr(mod, "PRINCIPAL", None),
            "campos": [
                {"key": c["key"], "label": c["label"], "tipo": c.get("type", "text"),
                 "requerido": c.get("required", True), "ejemplo": c.get("example")}
                for c in cs
            ],
        }
        for ds, cs in campos.items()
    ]
    corte_ej = (ejercicio_modelo.escenario(mod) or (None, None, None))[2]
    return {
        "id": pid,
        "nombre": d.get("name", pid),
        "marcos": d.get("frameworks", []),
        "resumen": d.get("summary", ""),
        "corte_ejemplo": corte_ej,
        "datasets": datasets,
        "parametros_por_defecto": getattr(mod, "PARAMETROS", {}) or {},
        "requerimientos_al_cliente": [
            {"id": r["id"], "documento": r.get("document", ""), "requerido": r.get("required", True),
             "uso": r.get("use", ""), "formatos": r.get("formats", []), "alimenta_dataset": r.get("dataset")}
            for r in d.get("requests", [])
        ],
    }


def plantilla(pid: str) -> dict | None:
    """El ejemplo realista del manifiesto como {corte, parametros, datasets}, para
    que la plataforma precargue un molde. None si la prueba no trae ejemplo."""
    mod = _mod(pid)
    esc = ejercicio_modelo.escenario(mod)
    if esc is None:
        return None
    datasets, param, corte = esc
    return {"id": pid, "corte": corte,
            "parametros": {**(getattr(mod, "PARAMETROS", {}) or {}), **(param or {})},
            "datasets": datasets}


# --------------------------------------------------------------------------- ejecución
def _engagement(pid: str, mod, param: dict, corte: str, encargo: dict | None) -> dict:
    encargo = encargo or {}
    marco = (encargo.get("framework") or param.get("_marco")
             or (mod.definicion().get("frameworks") or ["NIIF para las PYMES"])[0])
    return {
        "client": encargo.get("client") or "Cliente",
        "ruc": encargo.get("ruc") or "",
        "cutoff": corte,
        "year": (corte or "")[:4],
        "framework": marco,
        "firm": encargo.get("firm") or "AuditConsulting Auditores Cía. Ltda.",
        "preparer": encargo.get("preparer") or "NIIF Piloto",
        "reviewer": encargo.get("reviewer") or "Pendiente",
    }


def preparar(pid: str, datasets: dict, parametros: dict | None, corte: str, encargo: dict | None = None):
    """Calcula la prueba con los datos del cliente y devuelve ``(definicion, reg)``
    listos para armar cualquier formato del papel. No arma archivos todavía."""
    mod = _mod(pid)
    d = mod.definicion()
    param = {**(getattr(mod, "PARAMETROS", {}) or {}), **(parametros or {})}
    res = mod.ejecutar(datasets or {}, param, corte)
    res["hojas"] = datos_cliente.con_datos(mod, res, datasets or {})
    reg = {
        "run": res,
        "datasets": datasets or {},
        "parameters": param,
        "program": [{**x, "reference": x.get("source", "")} for x in d.get("program", [])],
        "sources": [],
        "engagement": _engagement(pid, mod, param, corte, encargo),
    }
    return d, reg


def papel(d: dict, reg: dict, formato: str, version: int = 1, estado: str = "PILOTO") -> bytes:
    """Bytes del papel de trabajo en un formato. Propaga ``libro.PDFNoDisponible``."""
    if formato not in _FN:
        raise ValueError(f"Formato no disponible: {formato!r}.")
    return getattr(libro, _FN[formato])(d, reg, [], version, estado)


def resumen_run(reg: dict) -> dict:
    """Lo que la plataforma reporta tras ejecutar: resultado principal y problemas."""
    run = reg["run"]
    prim = run.get("primary")
    return {
        "principal": prim,
        "etiqueta_principal": run.get("labels", {}).get(prim, prim),
        "valor_principal": run.get("totals", {}).get(prim),
        "n_problemas": len(run.get("exceptions", [])),
        "problemas": run.get("exceptions", []),
        "filas_por_dataset": {k: len(v) for k, v in (reg["datasets"] or {}).items()},
    }


def verificar_excel(contenido: bytes) -> dict:
    """Reabre el Excel (regla suprema, punto 6): que NO levante «reparaciones».
    Reporta hojas y celdas de TEXTO que Excel podría leer como fórmula."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido))  # corrupto -> lanza aquí
    riesgosas = []
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                if (celda.data_type == "s" and isinstance(celda.value, str)
                        and celda.value[:1] in _PREFIJOS_PELIGROSOS):
                    riesgosas.append(f"{ws.title}!{celda.coordinate}={celda.value[:20]!r}")
    return {"reabre": True, "hojas": wb.sheetnames, "n_hojas": len(wb.sheetnames),
            "celdas_texto_riesgosas": riesgosas[:20]}
