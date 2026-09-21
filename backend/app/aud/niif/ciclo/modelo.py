"""Modelo Excel de un requerimiento de cálculo (E10).

El auditor se lo envía al cliente para que todos entreguen la información en
el mismo formato: la hoja «Datos» tiene exactamente las columnas de la ficha
(con el nombre que reconoce «Procesar») y la hoja «Instrucciones» explica qué
documento es y cómo llenar cada columna. Si el requerimiento llega por partes
(meses, bodegas), se usa el mismo modelo en cada una y «Procesar» las une.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, GOLD, BLANCO = "0A2342", "C7A83C", "FFFFFF"
_FINO = Side(style="thin", color="B7C0CC")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
_FILAS_CON_FORMATO = 1000

FLOW_FIELDS = [
    {"key": "id", "label": "Contrato", "type": "text", "required": True},
    {"key": "fecha", "label": "Fecha del pago", "type": "date", "required": True},
    {"key": "importe", "label": "Importe", "type": "number", "required": True},
]

_TIPO = {"text": "Texto", "number": "Número", "date": "Fecha"}
_FORMATO = {
    "text": "Texto libre. En «id», un código único por partida.",
    "number": "Número con punto decimal, sin símbolos ni separadores de miles. Celda de tipo número.",
    "date": "Fecha de Excel o texto AAAA-MM-DD.",
}
_EJEMPLO = {"text": "A-001", "number": "1250.00", "date": "2025-12-31"}


def requerimientos_de_calculo(requests: list, definicion: dict) -> dict:
    """{id de requerimiento: campos} de los requerimientos que alimentan el
    cálculo: los marcados `use: calculo` o, si ninguno lo está, el primero
    (el RQ-001 genérico es la población). El segundo de cálculo es el
    calendario de pagos cuando la ficha tiene flujos."""
    reqs = requests or []
    calculo = [r for r in reqs if r.get("use") == "calculo"] or reqs[:1]
    salida = {}
    if calculo:
        salida[calculo[0]["id"]] = definicion.get("fields") or []
    if "flows" in definicion and len(calculo) > 1:
        salida[calculo[1]["id"]] = FLOW_FIELDS
    return salida


def _seguro(v) -> str:
    """Texto que Excel no confunda con una fórmula (regla del proyecto)."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def construir(definicion: dict, req: dict, campos: list, encargo: dict) -> bytes:
    wb = Workbook()
    datos = wb.active
    datos.title = "Datos"
    for i, f in enumerate(campos, start=1):
        c = datos.cell(row=1, column=i, value=f.get("label") or f["key"])
        c.font = Font(name="Calibri", size=10, bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDE
        letra = get_column_letter(i)
        datos.column_dimensions[letra].width = max(14, min(40, len(c.value) + 4))
        # Sin preformatear filas vacías: Excel las contaría como usadas y quien
        # pegue o agregue datos «al final» caería en la fila 1002. El tipo de
        # cada columna lo cuida la validación de datos, que no crea celdas.
        if f.get("type") in ("number", "date"):
            dv = DataValidation(
                type="decimal" if f["type"] == "number" else "date",
                operator="greaterThan" if f["type"] == "number" and f.get("positive") else "between",
                formula1="0" if f["type"] == "number" and f.get("positive") else ("-999999999999" if f["type"] == "number" else "1"),
                formula2=None if f["type"] == "number" and f.get("positive") else ("999999999999" if f["type"] == "number" else "73050"),
                allow_blank=True, showErrorMessage=True,
                errorTitle="Valor no válido",
                error="Ingrese un número." if f["type"] == "number" else "Ingrese una fecha.",
            )
            dv.add(f"{letra}2:{letra}{_FILAS_CON_FORMATO + 1}")
            datos.add_data_validation(dv)
    datos.freeze_panes = "A2"
    datos.row_dimensions[1].height = 30

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 30
    for col, w in zip("BCDEF", (12, 13, 46, 18, 34)):
        ins.column_dimensions[col].width = w
    titulo = ins.cell(row=1, column=1, value=f"MODELO DE ENTREGA · {req['id']} · {definicion.get('name', '')}")
    titulo.font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    ins.cell(row=2, column=1, value=_seguro(
        f"{encargo.get('firm') or 'AuditConsulting'} · Cliente: {encargo.get('client') or ''} · Corte: {encargo.get('cutoff') or ''}"
    )).font = Font(name="Calibri", size=10, color=GOLD, bold=True)
    fila = 4
    for etiqueta, valor in (
        ("Documento", req.get("document")),
        ("Para qué se usa", req.get("purpose")),
        ("Reporte de origen", req.get("report")),
        ("Período que debe cubrir", req.get("timing") or req.get("period")),
        ("Contenido mínimo", req.get("content")),
        ("Partes", ", ".join(req.get("components") or []) or "Un solo archivo"),
    ):
        if not valor:
            continue
        ins.cell(row=fila, column=1, value=etiqueta).font = Font(name="Calibri", size=10, bold=True)
        c = ins.cell(row=fila, column=2, value=_seguro(valor))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ins.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=6)
        fila += 1
    fila += 1
    reglas = (
        "Llene la hoja «Datos»: una fila por partida, desde la fila 2.",
        "No cambie ni reordene los nombres de las columnas de la fila 1.",
        "No agregue filas de total ni de subtotal, ni filas en blanco entre datos.",
        "Si entrega por partes (meses, bodegas, sucursales), use este mismo modelo en cada archivo.",
        "Los importes van como número (no como texto), con punto decimal.",
    )
    for r in reglas:
        ins.cell(row=fila, column=1, value="• " + r).font = Font(name="Calibri", size=10)
        ins.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        fila += 1
    fila += 1
    for i, h in enumerate(("Columna", "Tipo", "Obligatoria", "Formato", "Ejemplo", "También puede llamarse"), start=1):
        c = ins.cell(row=fila, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = _BORDE
    for f in campos:
        fila += 1
        valores = (
            f.get("label") or f["key"], _TIPO.get(f.get("type"), "Texto"),
            "Sí" if f.get("required", True) else "No", _FORMATO.get(f.get("type"), ""),
            _seguro(f.get("example") or _EJEMPLO.get(f.get("type"), "")),
            ", ".join(f.get("aliases") or []),
        )
        for i, v in enumerate(valores, start=1):
            c = ins.cell(row=fila, column=i, value=v)
            c.font = Font(name="Calibri", size=9)
            c.border = _BORDE
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.active = 0
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
