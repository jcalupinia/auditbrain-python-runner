"""Extracción por IA de documentos narrativos (PDF/Word) a las filas que un
procesador de planificación espera.

Caso de uso: la **carta de control interno** (RQ-004) y el **informe de auditoría
del año anterior** (RQ-005) llegan como documentos firmados (PDF o Word), no como
una tabla transcrita en Excel. Este módulo lee el texto del documento y le pide al
modelo de IA que devuelva **las mismas filas** que hoy se transcriben a mano (una
fila por hallazgo / por asunto del informe, con sus columnas), como JSON estricto.

**Proveedor de IA:** usa el cliente compartido ``backend.app.chat.providers``
(``chat_complete``), cuyo orden de preferencia pone **el servidor de IA LOCAL
primero** (privacidad + costo cero) y cae a los proveedores de nube solo como
respaldo. Así respeta ``AUDITBRAIN_LLM_PROVIDER`` y la misma política que el resto
de la plataforma; no llama a ningún proveedor directamente.

Reglas que este módulo respeta:

- **La IA solo transcribe lo que el documento dice.** No inventa calificaciones ni
  cifras: si un dato no está en el documento, la celda queda vacía (regla M22 del
  proyecto). El auditor revisa y confirma la tabla extraída ANTES de que alimente la
  herramienta (la IA no decide sola).
- **Degradación elegante.** Si no hay ningún proveedor LLM configurado (ni el local
  ni uno de nube) o la extracción está apagada (``NIIF_EXTRACCION_ENABLED=false``),
  se levanta :class:`ExtraccionNoDisponible` con un mensaje claro para que el auditor
  caiga al respaldo (subir la tabla en Excel/CSV).

El motor es **genérico**: recibe la lista de ``campos`` (del procesador) y no conoce
la semántica de la planificación. La normalización y validación finas (p. ej. el
tipo del informe contra ``TIPOS_INFORME``) las hace el consumidor con
``proc.validar_filas``.
"""
from __future__ import annotations

import json
import logging
import os
import re
import time
from typing import Any, Callable, Optional

log = logging.getLogger(__name__)

# Recorte del texto que se manda al modelo (los documentos del año anterior son
# cortos; este tope evita costos/lentitud si alguien sube un PDF gigante por error).
MAX_CHARS = int(os.getenv("NIIF_EXTRACCION_MAX_CHARS", "60000"))
MAX_RETRIES = int(os.getenv("NIIF_EXTRACCION_MAX_RETRIES", "2"))
EXTRACCION_ENABLED = os.getenv("NIIF_EXTRACCION_ENABLED", "true").lower() in ("true", "1", "yes")

# Literal que `providers._call_openai_compatible`/`_call_anthropic` ponen cuando el
# modelo contesta con contenido vacío. Lo espejamos aquí para detectar esa respuesta
# vacía y dar un error accionable en vez del críptico «Expecting value» del JSON.
_RESPUESTA_VACIA_PROVEEDOR = "(respuesta vacía del proveedor)"


class ExtraccionError(Exception):
    """La extracción no se pudo completar (texto ilegible, respuesta inválida...)."""


class ExtraccionNoDisponible(ExtraccionError):
    """La IA de extracción no está disponible (sin proveedor LLM o apagada)."""


# --------------------------------------------------------------------------- #
#  1 · Texto del documento (PDF / Word / texto plano)                          #
# --------------------------------------------------------------------------- #
def _extension(nombre: str) -> str:
    return (os.path.splitext(str(nombre or ""))[1] or "").lower().lstrip(".")


def texto_de_documento(nombre: str, datos: bytes) -> str:
    """Devuelve el texto plano de un PDF, Word (.docx) o archivo de texto.

    Levanta :class:`ExtraccionError` si el formato no se soporta o si no se pudo
    extraer texto (p. ej. un PDF escaneado sin capa de texto).
    """
    ext = _extension(nombre)
    if ext == "pdf":
        texto = _texto_pdf(datos)
    elif ext == "docx":
        texto = _texto_docx(datos)
    elif ext == "doc":
        raise ExtraccionError(
            "El formato .doc antiguo no se puede leer automáticamente; guárdelo como .docx o .pdf y vuelva a subirlo."
        )
    elif ext in ("txt", "md", "csv"):
        texto = datos.decode("utf-8", errors="replace")
    elif ext == "xlsx":
        texto = _texto_xlsx(datos)
    else:
        raise ExtraccionError(
            f"Formato no soportado para extracción por IA: «.{ext}». Use PDF, Word (.docx) o Excel (.xlsx)."
        )
    texto = (texto or "").strip()
    if not texto:
        raise ExtraccionError(
            "No se pudo extraer texto del documento. Si es un PDF escaneado (imagen), súbalo con capa de texto "
            "o transcriba los datos en la plantilla Excel."
        )
    return texto


def _texto_xlsx(datos: bytes) -> str:
    """Convierte un .xlsx a texto plano (una línea por fila, celdas unidas por « | »)
    para que la IA lea una carta/informe/notas transcritas en Excel igual que si fueran
    un documento firmado. Así un requerimiento extraíble subido en Excel se lee solo en
    vez de pedir el mapeo manual de columnas. Levanta :class:`ExtraccionError` si el
    archivo no es un XLSX válido."""
    import io

    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:
        raise ExtraccionError(f"No se pudo leer el Excel para extracción por IA ({e}).") from e
    partes: list[str] = []
    try:
        for ws in wb.worksheets:
            filas: list[str] = []
            for fila in ws.iter_rows(values_only=True):
                celdas = ["" if c is None else str(c).strip() for c in fila]
                while celdas and not celdas[-1]:   # recorta las celdas vacías finales (rango usado del ERP)
                    celdas.pop()
                if celdas:
                    filas.append(" | ".join(celdas))
                if len(filas) >= 2000:   # tope de seguridad; MAX_CHARS recorta aguas abajo
                    break
            if filas:
                partes.append(f"# Hoja: {ws.title}\n" + "\n".join(filas))
    finally:
        wb.close()
    return "\n\n".join(partes)


def _texto_pdf(datos: bytes) -> str:
    import io

    partes: list[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(datos)) as pdf:
            for pagina in pdf.pages:
                partes.append(pagina.extract_text() or "")
        if any(p.strip() for p in partes):
            return "\n".join(partes)
    except Exception as e:  # pragma: no cover - depende de pdfplumber en runtime
        log.warning("pdfplumber no pudo leer el PDF (%s); intento pypdf", e)
    # Respaldo: pypdf.
    texto = ""
    try:
        from pypdf import PdfReader

        lector = PdfReader(io.BytesIO(datos))
        texto = "\n".join((p.extract_text() or "") for p in lector.pages)
    except Exception as e:  # pragma: no cover
        log.warning("pypdf no pudo leer el PDF (%s); intento OCR", e)
    if texto.strip():
        return texto
    # Respaldo final: PDF escaneado (sin capa de texto) → OCR, si está disponible.
    # Degrada con elegancia: si OCR no está configurado o falla, devuelve lo que
    # haya (vacío) y texto_de_documento levanta el aviso de «súbalo con capa de
    # texto o transcriba la plantilla».
    return _texto_pdf_ocr(datos) or texto


def _texto_pdf_ocr(datos: bytes) -> str:
    """Intenta OCR (Google Vision) sobre un PDF escaneado. Nunca levanta: si el
    OCR no está disponible o falla, devuelve cadena vacía para que el flujo caiga
    al respaldo Excel/CSV del ciclo."""
    try:
        from backend.app.utils import ocr
    except Exception as e:  # pragma: no cover - import defensivo
        log.warning("No se pudo importar el módulo OCR (%s)", e)
        return ""
    if not ocr.is_available():
        return ""
    try:
        resultado = ocr.ocr_pdf_bytes(datos)
    except Exception as e:  # pragma: no cover - depende de la API de Vision
        log.warning("OCR no pudo leer el PDF escaneado (%s)", e)
        return ""
    return (resultado.get("text") or "").strip()


def _texto_docx(datos: bytes) -> str:
    import io

    try:
        from docx import Document
    except Exception as e:  # pragma: no cover
        raise ExtraccionError("Falta la librería python-docx para leer archivos Word.") from e
    doc = Document(io.BytesIO(datos))
    lineas: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Las cartas de control interno suelen traer los hallazgos en tablas: se vuelcan
    # fila por fila (celdas separadas por « | ») para que el modelo las vea.
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                lineas.append(" | ".join(celdas))
    return "\n".join(lineas)


# --------------------------------------------------------------------------- #
#  2 · Prompt y parseo de JSON                                                 #
# --------------------------------------------------------------------------- #
def _columnas(campos: list, enums: dict) -> str:
    filas = []
    for c in campos:
        linea = f'- "{c["key"]}": {c.get("label") or c["key"]}'
        if c.get("type") == "number":
            linea += " (número)"
        if not c.get("required", True):
            linea += " [opcional]"
        if c["key"] in enums:
            linea += " — valores permitidos: " + ", ".join(str(v) for v in enums[c["key"]])
        elif c.get("example") is not None:
            linea += f" — ej.: {c['example']}"
        filas.append(linea)
    return "\n".join(filas)


_SISTEMA = (
    "Eres un asistente de auditoría que TRANSCRIBE datos de un documento a una tabla estructurada. "
    "Respondes ÚNICAMENTE con JSON válido, sin texto adicional, sin explicaciones y sin markdown."
)


def _prompt(campos: list, instrucciones: str, contexto: str, enums: dict, texto: str) -> str:
    partes = [
        "Transcribe el documento a una tabla. Reglas estrictas:",
        "1. Transcribe SOLO lo que el documento dice de forma explícita. NO inventes, deduzcas ni completes datos.",
        "2. Si un dato no aparece en el documento, usa null en ese campo (no pongas 0 ni un valor inventado).",
        "3. Una entrada por cada fila/ítem real del documento; no agregues filas de total ni de resumen.",
        "4. Respeta los valores permitidos de cada campo cuando se indiquen.",
        "",
        "Columnas de cada fila (usa exactamente estas claves):",
        _columnas(campos, enums),
    ]
    if instrucciones:
        partes += ["", instrucciones]
    if contexto:
        partes += ["", f"Contexto del encargo: {contexto}"]
    partes += [
        "",
        'Devuelve un objeto JSON con esta forma exacta: {"filas": [ { … una fila … }, … ]}. '
        "Si no hay filas, devuelve {\"filas\": []}.",
        "",
        "Documento:",
        '"""',
        texto,
        '"""',
    ]
    return "\n".join(partes)


def _salvar_filas(texto: str) -> list:
    """Rescata los objetos de fila de un texto con JSON MAL FORMADO.

    Los modelos locales (gpt-oss, etc.) a veces devuelven un JSON con un error de
    sintaxis menor —coma faltante entre dos filas, salida truncada a mitad de la
    última fila— que hace fallar `json.loads` entero y perder TODAS las filas,
    aunque casi todas estuvieran bien. Esto escanea el array ``"filas"`` y extrae
    cada bloque ``{...}`` balanceado por separado (respetando las comillas y los
    escapes), parsea cada uno de forma independiente y descarta solo el que esté
    roto. Así una coma faltante o un corte al final cuestan a lo sumo UNA fila, no
    la extracción completa.
    """
    # Empezar dentro del array de "filas" si existe; si no, desde el primer '['.
    m = re.search(r'"filas"\s*:\s*\[', texto)
    if m:
        inicio = m.end()
    else:
        corchete = texto.find("[")
        inicio = corchete + 1 if corchete != -1 else 0

    filas: list = []
    i, n = inicio, len(texto)
    while i < n:
        if texto[i] != "{":
            if texto[i] == "]":  # fin del array de filas
                break
            i += 1
            continue
        # Escanear desde '{' hasta su '}' balanceado, respetando strings.
        depth, en_str, escape, j = 0, False, False, i
        while j < n:
            c = texto[j]
            if en_str:
                if escape:
                    escape = False
                elif c == "\\":
                    escape = True
                elif c == '"':
                    en_str = False
            elif c == '"':
                en_str = True
            elif c == "{":
                depth += 1
            elif c == "}":
                depth -= 1
                if depth == 0:
                    break
            j += 1
        if depth != 0 or j >= n:
            break  # bloque no balanceado (truncado): no hay más filas rescatables
        try:
            obj = json.loads(texto[i:j + 1])
            if isinstance(obj, dict):
                filas.append(obj)
        except json.JSONDecodeError:
            pass  # fila rota: se descarta solo ella
        i = j + 1
    return filas


def _json_de_texto(texto: str) -> dict:
    """Extrae el objeto JSON de la respuesta del modelo, tolerando cercas de código
    y texto alrededor. Si el JSON está mal formado (coma faltante, truncado),
    rescata las filas válidas una por una en vez de perder toda la extracción."""
    t = (texto or "").strip()
    # Quitar cercas ```json ... ```
    m = re.search(r"```(?:json)?\s*(.+?)```", t, re.DOTALL)
    if m:
        t = m.group(1).strip()
    # Quedarse con el primer objeto {...} balanceado por sus llaves extremas.
    ini, fin = t.find("{"), t.rfind("}")
    if ini != -1 and fin != -1 and fin > ini:
        t = t[ini:fin + 1]
    try:
        return json.loads(t)
    except json.JSONDecodeError:
        # Rescate: el modelo devolvió JSON con un error de sintaxis. Salvar las
        # filas que sí se puedan leer en lugar de fallar toda la extracción.
        filas = _salvar_filas(t)
        if filas:
            log.warning("JSON de la IA mal formado; se rescataron %d filas válidas.", len(filas))
            return {"filas": filas}
        raise


# --------------------------------------------------------------------------- #
#  3 · Cliente de chat (cadena de proveedores, local primero; inyectable)     #
# --------------------------------------------------------------------------- #
def _chat_por_defecto() -> Callable:
    """Devuelve la función de chat compartida (servidor local primero, con
    respaldo a la nube). Levanta :class:`ExtraccionNoDisponible` si no hay ningún
    proveedor configurado o si la extracción está apagada.

    Usa :func:`providers.completar_para_extraccion`, afinada para esta tarea:
    ``temperature=0`` (transcripción determinista) y **streaming** cuando el
    proveedor primario lo soporta (el servidor local no entrega nada hasta
    terminar de generar; en streaming los tokens fluyen y la conexión no se corta
    por timeout en documentos largos). Cae al no-streaming automáticamente."""
    if not EXTRACCION_ENABLED:
        raise ExtraccionNoDisponible(
            "La extracción por IA está deshabilitada (NIIF_EXTRACCION_ENABLED=false). "
            "Suba la tabla en Excel/CSV con la plantilla del requerimiento."
        )
    from backend.app.chat import providers

    if providers.available_provider() is None:
        raise ExtraccionNoDisponible(
            "No hay ningún proveedor de IA configurado en el servidor (ni el local ni uno de nube): "
            "la extracción por IA no está disponible. Suba la tabla en Excel/CSV con la plantilla del requerimiento."
        )
    return providers.completar_para_extraccion


# --------------------------------------------------------------------------- #
#  4 · API pública                                                            #
# --------------------------------------------------------------------------- #
def extraer_filas(
    campos: list,
    texto: str,
    *,
    instrucciones: str = "",
    enums: Optional[dict] = None,
    contexto: str = "",
    chat: Optional[Callable] = None,
) -> dict:
    """Extrae filas estructuradas del ``texto`` según los ``campos`` del procesador,
    usando el servidor de IA local primero (cadena de proveedores compartida).

    Devuelve ``{"rows": [...], "modelo": str, "n": int}``. Cada fila es un dict con
    las mismas claves que produciría la transcripción en Excel (más ``_row``,
    el número de fila 1-based que usa la validación). Levanta
    :class:`ExtraccionNoDisponible` si la IA no está disponible, o
    :class:`ExtraccionError` si la respuesta es inválida.
    """
    enums = enums or {}
    chat = chat or _chat_por_defecto()
    texto = texto[:MAX_CHARS]
    prompt = _prompt(campos, instrucciones, contexto, enums, texto)
    claves = {c["key"] for c in campos}
    numericos = {c["key"] for c in campos if c.get("type") == "number"}

    error: Optional[Exception] = None
    for intento in range(MAX_RETRIES):
        try:
            resp = chat([{"role": "user", "content": prompt}], system=_SISTEMA)
            contenido = (getattr(resp, "content", "") or "").strip()
            # Respuesta vacía del modelo: un modelo de razonamiento (gpt-oss del
            # servidor local) a veces gasta todo su presupuesto de tokens pensando
            # y no llega a escribir el JSON; el proveedor devuelve "" o el literal
            # placeholder. En vez del críptico «Expecting value: line 1 column 1»,
            # damos un mensaje accionable (se reintenta igual con backoff).
            if not contenido or contenido == _RESPUESTA_VACIA_PROVEEDOR:
                raise ExtraccionError(
                    "El modelo de IA devolvió una respuesta vacía (no transcribió nada). "
                    "Suele pasar cuando el modelo de razonamiento del servidor local agota "
                    "su presupuesto de tokens razonando sin llegar a responder. Opciones: "
                    "reintente; suba AUDITBRAIN_LLM_MAX_TOKENS_EXTRACCION en Render; o cargue "
                    "la carta como tabla en Excel/CSV con la plantilla del requerimiento."
                )
            datos = _json_de_texto(contenido)
            filas = datos.get("filas") if isinstance(datos, dict) else datos
            if not isinstance(filas, list):
                raise ExtraccionError("La IA no devolvió una lista de filas.")
            rows = [_normalizar_fila(f, claves, numericos, i) for i, f in enumerate(filas, start=1)
                    if isinstance(f, dict)]
            return {"rows": rows, "modelo": getattr(resp, "model", "") or "", "n": len(rows)}
        except ExtraccionNoDisponible:
            raise
        except Exception as e:  # reintento con backoff exponencial 1s/2s/...
            error = e
            log.warning("Extracción IA falló (intento %d/%d): %s", intento + 1, MAX_RETRIES, e)
            if intento < MAX_RETRIES - 1:
                time.sleep(2 ** intento)
    raise ExtraccionError(f"No se pudo extraer la información con la IA: {error}")


def _normalizar_fila(fila: dict, claves: set, numericos: set, n: int) -> dict:
    """Se queda con las claves conocidas, vacía los null y numera la fila."""
    out: dict[str, Any] = {"_row": n}
    for k in claves:
        v = fila.get(k)
        if v is None:
            out[k] = ""
        elif k in numericos and isinstance(v, float) and v.is_integer():
            out[k] = int(v)
        else:
            out[k] = v
    return out
