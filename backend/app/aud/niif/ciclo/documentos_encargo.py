"""Documentos del encargo que genera la plataforma (decisión del dueño, 2026-09-26: «la plataforma los genera»).

- Carta de encargo (NIA 210): de la ficha del encargo y de los registros (socio que aceptó).
- Acta de la discusión del equipo (NIA 315 y NIA 240): asistentes registrados con un clic y los riesgos e indicios de
  fraude de la última planificación ejecutada (hojas 13 y 26).
- Carta de planificación al gobierno de la entidad (NIA 260): los asuntos de la hoja 32 y los riesgos significativos.

Son MODELOS para revisar antes de firmar: el texto normativo lleva «VERIFICAR» donde cita párrafos, y lo que no consta
en la plataforma queda «[PENDIENTE]» (no se completa por inferencia). La firma se registra luego con un clic
(carta firmada, comunicación enviada) y ese registro es el que entra al papel de la planificación.
"""
from __future__ import annotations

import datetime
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from backend.app.aud.niif.procesadores import marca

TIPOS = {
    "carta_encargo": "Carta_de_encargo",
    "acta_discusion": "Acta_discusion_equipo",
    "carta_planificacion": "Carta_de_planificacion",
}
PENDIENTE = "[PENDIENTE]"
FIRMA_LEGAL = {"Audit Consulting": "AuditConsulting Auditores Cía. Ltda."}
NAVY = RGBColor(0x0A, 0x23, 0x42)
GRIS = RGBColor(0x6B, 0x72, 0x80)
AVISO = ("Documento generado por AUDIT-IA a partir de la ficha del encargo, los registros de la plataforma y la "
         "planificación ejecutada. Es un modelo: el socio lo revisa y ajusta antes de firmarlo. Las citas de párrafos "
         "marcadas «VERIFICAR» se confirman contra la norma vigente.")


def _v(c):
    return c.get("v") if isinstance(c, dict) else c


def _txt(x) -> str:
    if x is None:
        return ""
    if isinstance(x, float):
        return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(x)


def firma_legal(ficha: dict) -> str:
    f = str((ficha or {}).get("firm") or "").strip()
    return FIRMA_LEGAL.get(f, f or "AuditConsulting Auditores Cía. Ltda.")


def hoja(run: dict | None, prefijo: str) -> dict | None:
    for h in (run or {}).get("hojas") or []:
        if str(h.get("name", "")).startswith(prefijo):
            return h
    return None


def filas(h: dict | None) -> list[dict]:
    """Las filas de una hoja guardada como diccionarios por encabezado (valores, no fórmulas)."""
    if not h:
        return []
    cols = [c[0] for c in h.get("cols") or []]
    return [{cols[i]: _v(c) for i, c in enumerate(f) if i < len(cols)} for f in h.get("rows") or []]


# --- armado del Word ---------------------------------------------------------------------------------------------------

def _doc(titulo: str, sub: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    enc = sec.header
    t = enc.add_table(rows=1, cols=2, width=sec.page_width - sec.left_margin - sec.right_margin)
    iz, de = t.rows[0].cells
    iz.paragraphs[0].add_run().add_picture(marca.flujo("auditconsulting_oscuro"), height=Cm(1.1))
    de.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    de.paragraphs[0].add_run().add_picture(marca.flujo("audit_ia"), height=Cm(1.1))
    vacio = enc.paragraphs[0]._p
    vacio.getparent().remove(vacio)
    pie = sec.footer.paragraphs[0]
    r = pie.add_run(AVISO)
    r.font.size, r.font.italic, r.font.color.rgb = Pt(7.5), True, GRIS
    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), NAVY
    if sub:
        p = doc.add_paragraph()
        r = p.add_run(sub)
        r.font.size, r.font.color.rgb = Pt(10), GRIS
    return doc


def _h(doc, texto: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(texto)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11.5), NAVY


def _p(doc, texto: str, negrita: bool = False):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = negrita
    return p


def _items(doc, lista: list[str]):
    for x in lista:
        doc.add_paragraph(x, style="List Bullet")


def _tabla(doc, encabezados: list[str], datos: list[list], anchos: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.autofit = False
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(e)
        r.bold, r.font.size = True, Pt(9)
    for f in datos:
        celdas = t.add_row().cells
        for i, x in enumerate(f):
            celdas[i].text = ""
            celdas[i].paragraphs[0].add_run(_txt(x)).font.size = Pt(9)
    if anchos:
        # LibreOffice y Word leen la cuadrícula (w:gridCol); Word además, el ancho de cada celda.
        for col, a in zip(t.columns, anchos):
            col.width = Cm(a)
        for fila in t.rows:
            for c, a in zip(fila.cells, anchos):
                c.width = Cm(a)
    doc.add_paragraph()
    return t


def _firmas(doc, izquierda: tuple[str, str], derecha: tuple[str, str] | None):
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2)
    for celda, (nombre, cargo) in zip(t.rows[0].cells, [izquierda, derecha or ("", "")]):
        if not nombre and not cargo:
            continue
        celda.text = ""
        celda.paragraphs[0].add_run("_______________________________")
        celda.add_paragraph(nombre).runs[0].bold = True
        celda.add_paragraph(cargo)


def _bytes(doc) -> bytes:
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def _fecha(f) -> str:
    return str(f or "")[:10] or PENDIENTE


# --- documentos --------------------------------------------------------------------------------------------------------

def carta_encargo(ficha: dict, encargo: dict, hoy: datetime.date | None = None) -> bytes:
    f = ficha or {}
    reg = (encargo or {}).get("registros") or {}
    socio = next((x["integrante"] for x in reg.get("equipo") or [] if x.get("rol") == "Socio"), "") or PENDIENTE
    firma = firma_legal(f)
    cliente = f.get("client") or PENDIENTE
    marco = f.get("framework") or PENDIENTE
    corte = f.get("cutoff") or PENDIENTE
    doc = _doc("Carta de encargo de auditoría", f"NIA 210 · {firma}")
    _p(doc, f"Fecha: {(hoy or datetime.date.today()).isoformat()}")
    _p(doc, f"Señores\nJunta General de Accionistas y Administración\n{cliente}\nRUC {f.get('ruc') or PENDIENTE}")
    _h(doc, "Objetivo y alcance de la auditoría")
    _p(doc, f"Ustedes nos han solicitado auditar los estados financieros de {cliente}, que comprenden el estado de situación "
            f"financiera al {corte}, el estado de resultados integral, el estado de cambios en el patrimonio y el estado de "
            f"flujos de efectivo del ejercicio terminado en esa fecha, y las notas explicativas. Nos complace confirmar mediante "
            f"esta carta nuestra aceptación y nuestro entendimiento de este encargo (NIA 210 párr. 10 — VERIFICAR).")
    _p(doc, "El objetivo de nuestra auditoría es obtener una seguridad razonable de que los estados financieros en su conjunto "
            "están libres de incorrección material, debida a fraude o error, y emitir un informe que contenga nuestra opinión.")
    _h(doc, "Responsabilidades del auditor")
    _items(doc, [
        "Realizaremos la auditoría de conformidad con las Normas Internacionales de Auditoría (NIA), que exigen cumplir los "
        "requerimientos de ética y planificar y ejecutar la auditoría para obtener una seguridad razonable.",
        "Por las limitaciones inherentes de una auditoría y del control interno, existe un riesgo inevitable de que no se "
        "detecten algunas incorrecciones materiales, aun cuando la auditoría se planifique y ejecute adecuadamente.",
        "Les comunicaremos por escrito las deficiencias significativas del control interno que identifiquemos (NIA 265 — VERIFICAR).",
    ])
    _h(doc, "Responsabilidades de la dirección")
    _p(doc, "La auditoría se realiza sobre la premisa de que la dirección reconoce y comprende su responsabilidad de:")
    _items(doc, [
        f"preparar los estados financieros de conformidad con {marco};",
        "mantener el control interno que considere necesario para que los estados financieros estén libres de incorrección "
        "material, debida a fraude o error;",
        "proporcionarnos acceso a toda la información relevante, la información adicional que solicitemos y acceso ilimitado "
        "a las personas de la entidad de quienes consideremos necesario obtener evidencia;",
        "entregarnos al término de la auditoría las manifestaciones escritas que solicitemos (NIA 580 — VERIFICAR).",
    ])
    _h(doc, "Marco de información financiera e informe")
    _p(doc, f"Marco aplicable: {marco}. Emitiremos un informe de auditoría cuya forma y contenido pueden requerir "
            "modificaciones en función de los hallazgos de la auditoría.")
    lim = (reg.get("carta") or {}).get("detalle") or ""
    if lim:
        _h(doc, "Limitaciones conocidas al alcance")
        _p(doc, lim)
    _h(doc, "Otros términos")
    _p(doc, f"Honorarios, calendario de visitas y equipo asignado: {PENDIENTE} (condiciones comerciales de la propuesta aceptada).")
    _p(doc, "Les rogamos firmar y devolver una copia de esta carta en señal de conformidad con sus términos.")
    _firmas(doc, (socio, f"Socio del encargo · {firma}"), (PENDIENTE, f"Representante legal · {cliente}"))
    return _bytes(doc)


def acta_discusion(ficha: dict, encargo: dict, run: dict | None) -> bytes:
    f = ficha or {}
    reg = (encargo or {}).get("registros") or {}
    asist = reg.get("asistencia") or []
    doc = _doc("Acta de la discusión del equipo del encargo",
               f"NIA 315 (Revisada 2019) párr. 17 y NIA 240 párr. 15 (VERIFICAR) · {f.get('client') or PENDIENTE} · corte {f.get('cutoff') or PENDIENTE}")
    _h(doc, "Fecha y asistentes")
    fechas = sorted({x.get("fecha") for x in asist if x.get("fecha")})
    _p(doc, f"Fecha: {', '.join(fechas) if fechas else PENDIENTE}. Asistencia registrada con un clic por cada integrante.")
    _tabla(doc, ["Integrante", "Rol", "Fecha del registro"],
           [[x.get("integrante"), x.get("rol"), x.get("fecha")] for x in asist] or [[PENDIENTE, "", ""]])
    if not any(x.get("rol") == "Socio" for x in asist):
        _p(doc, "Advertencia: no consta la participación del socio del encargo (NIA 315 párr. 17 — VERIFICAR).", True)
    _h(doc, "Susceptibilidad de los estados financieros a incorrección material")
    riesgos = [x for x in filas(hoja(run, "13_")) if x.get("¿Se presenta?") == "Sí"]
    if riesgos:
        _tabla(doc, ["Código", "Área", "Condición observada", "Posible riesgo", "Severidad"],
               [[x.get("Código"), x.get("Rubro o área"), x.get("Condición observada"), x.get("Posible riesgo"), x.get("Severidad")]
                for x in riesgos], [1.8, 3.0, 4.4, 4.6, 2.6])
    else:
        _p(doc, f"{PENDIENTE}: ejecute la planificación para incorporar los riesgos identificados.")
    _h(doc, "Fraude: indicios y respuesta")
    fraude = [x for x in filas(hoja(run, "26_")) if str(x.get("Código") or "").startswith(("FRA", "DIS"))]
    if fraude:
        _tabla(doc, ["Código", "Aspecto", "Resultado", "Estado"],
               [[x.get("Código"), x.get("Aspecto"), x.get("Resultado"), x.get("Estado")] for x in fraude], [1.6, 6.4, 6.2, 2.4])
    _p(doc, "Conclusión: el equipo discutió cómo y dónde los estados financieros podrían contener una incorrección material "
            "debida a fraude o error; los riesgos anteriores se responden en el programa de auditoría (hoja 19 del papel).")
    socio = next((x.get("integrante") for x in asist if x.get("rol") == "Socio"), "") or PENDIENTE
    _firmas(doc, (socio, "Socio del encargo"), None)
    return _bytes(doc)


def carta_planificacion(ficha: dict, encargo: dict, run: dict | None, hoy: datetime.date | None = None) -> bytes:
    f = ficha or {}
    reg = (encargo or {}).get("registros") or {}
    firma = firma_legal(f)
    cliente = f.get("client") or PENDIENTE
    doc = _doc("Comunicación de la planificación de la auditoría", f"NIA 260 (Revisada) · {firma}")
    _p(doc, f"Fecha: {(hoy or datetime.date.today()).isoformat()}")
    _p(doc, f"Señores\nResponsables del gobierno de la entidad\n{cliente}")
    _p(doc, f"Les comunicamos el alcance y el momento de realización de la auditoría de los estados financieros al "
            f"{f.get('cutoff') or PENDIENTE} y los riesgos significativos identificados en la planificación.")
    asuntos = filas(hoja(run, "32_"))
    if asuntos:
        # La propia carta no se lista (su envío se registra con un clic después de mandarla).
        _tabla(doc, ["Asunto", "Contenido", "Norma"],
               [[x.get("Asunto"), x.get("Contenido"), x.get("Norma")] for x in asuntos
                if x.get("Estado") != "No aplica" and not str(x.get("Asunto") or "").startswith("Envío de la carta")],
               [4.6, 7.4, 4.6])
    else:
        _p(doc, f"{PENDIENTE}: ejecute la planificación para incorporar los asuntos a comunicar.")
    _p(doc, "Esta comunicación no incluye procedimientos detallados, para no reducir su eficacia. Quedamos a su disposición "
            "para comentar su contenido.")
    socio = next((x["integrante"] for x in reg.get("equipo") or [] if x.get("rol") == "Socio"), "") or PENDIENTE
    _firmas(doc, (socio, f"Socio del encargo · {firma}"), None)
    return _bytes(doc)


def generar(tipo: str, ficha: dict, encargo: dict, run: dict | None) -> bytes:
    if tipo == "carta_encargo":
        return carta_encargo(ficha, encargo)
    if tipo == "acta_discusion":
        return acta_discusion(ficha, encargo, run)
    if tipo == "carta_planificacion":
        return carta_planificacion(ficha, encargo, run)
    raise ValueError("Documento desconocido.")
