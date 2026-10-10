"""DM2 · Cédula Sumaria de las cuentas de impuestos.

Por cada cuenta clasificada: el saldo del cierre anterior (el asiento de
apertura del mayor), el saldo al corte (saldo neto acumulado = debe − haber) y
la variación del período. Equivale a la hoja "DM2 Cédula Sumaria" del papel de
trabajo de la firma.

Los saldos se calculan a partir de los movimientos del mayor (que
``armar_libro`` ya tiene en memoria), no por fórmula: el saldo contable es el
NETO (débito − crédito), distinto del lado bruto que muestra "Mayores
homologados" para los cruces, así que no hay una celda única que referenciar.
El saldo del cierre anterior se identifica por el asiento de apertura del
propio mayor (descripción tipo "SALDOS INICIALES" / "APERTURA"); si la cuenta
no trae apertura, queda en cero y toda la variación es del período.
"""

from __future__ import annotations

import re
from collections import defaultdict

from openpyxl.utils import get_column_letter

from backend.app.aud.obligaciones_fiscales.libro.estilos import (
    BORDE, FONT_DATA, FONT_TOTAL, FORMATO_NUM,
    RELLENO_TOTAL, escribir_encabezado_cedula, escribir_leyenda_marcas,
    estilo_encabezado_tabla,
)
from backend.app.aud.obligaciones_fiscales.mayor.catalogo import CATEGORIAS

SHEET_DM2 = "DM2 Cédula Sumaria"

# Asiento de apertura: el saldo que viene del cierre del ejercicio anterior.
# Mismo criterio que el motor de referencia; cubre "REG. SALDOS INICIALES",
# "SALDO ANTERIOR" y "ASIENTO DE APERTURA".
_RE_APERTURA = re.compile(
    r"SALDOS?\s*INIC|SALDO\s*ANTERIOR|APERTURA|ASIENTO\s*DE\s*APERTURA",
    re.IGNORECASE,
)

COL_CODIGO, COL_NOMBRE, COL_ANTERIOR, COL_VARIACION, COL_PCT, COL_CORTE, COL_MARCA = range(1, 8)


def _orden_categoria(categoria: str | None) -> int:
    cat = CATEGORIAS.get(categoria or "")
    return cat.orden if cat else 99


def _saldos_por_cuenta(movimientos) -> dict[str, tuple[float, float]]:
    """{codigo → (saldo_anterior, saldo_corte)} en NETO (débito − crédito)."""
    anterior: dict[str, float] = defaultdict(float)
    total: dict[str, float] = defaultdict(float)
    for m in movimientos or []:
        neto = round(m.debe - m.haber, 2)
        total[m.codigo] += neto
        if _RE_APERTURA.search(m.descripcion or ""):
            anterior[m.codigo] += neto
    return {
        cod: (round(anterior.get(cod, 0.0), 2), round(total.get(cod, 0.0), 2))
        for cod in total
    }


def _fila_valores(ws, fila, *, codigo, nombre, anterior, corte, total=False):
    variacion = round(corte - anterior, 2)
    pct = round(variacion / anterior, 4) if anterior else None
    font = FONT_TOTAL if total else FONT_DATA
    ws.cell(fila, COL_CODIGO, codigo).font = font
    ws.cell(fila, COL_NOMBRE, nombre).font = font
    valores = {
        COL_ANTERIOR: anterior, COL_VARIACION: variacion, COL_CORTE: corte,
    }
    for col, val in valores.items():
        c = ws.cell(fila, col, val)
        c.font = font
        c.number_format = FORMATO_NUM
        c.border = BORDE
    # El porcentaje solo si hay base: evita el "#DIV/0!" del modelo manual.
    cpct = ws.cell(fila, COL_PCT, pct if pct is not None else "—")
    cpct.font = font
    if pct is not None:
        cpct.number_format = "0.00%"
    cpct.border = BORDE
    if total:
        for col in (COL_CODIGO, COL_NOMBRE, COL_ANTERIOR, COL_VARIACION, COL_PCT, COL_CORTE):
            ws.cell(fila, col).fill = RELLENO_TOTAL
    else:
        ws.cell(fila, COL_MARCA, "ü").font = font


def build_dm2(
    wb,
    *,
    clasificacion,
    movimientos,
    cliente: str = "",
    periodo: str = "",
    preparado_por: str | None = None,
    revisado_por: str | None = None,
) -> dict:
    """Construye DM2. Devuelve {("saldo_corte", código) → dirección} para que
    DM3 cruce el saldo al corte contra la misma cifra de la sumaria."""
    if SHEET_DM2 in wb.sheetnames:
        del wb[SHEET_DM2]
    ws = wb.create_sheet(SHEET_DM2)

    escribir_encabezado_cedula(
        ws, titulo="Cédula Sumaria", referencia="DM2",
        cliente=cliente, periodo=periodo,
        preparado_por=preparado_por, revisado_por=revisado_por,
    )

    saldos = _saldos_por_cuenta(movimientos)

    fila = 12
    encabezado = ("Código", "Descripción", "Saldo cierre anterior", "Variación",
                  "%", "Saldo al corte", "")
    for j, texto in enumerate(encabezado):
        estilo_encabezado_tabla(ws.cell(fila, 1 + j, texto))
    fila += 1

    # Agrupadas por categoría (mismo orden que "Mayores homologados"); dentro,
    # por código. Un subtotal por categoría y un total general al final.
    por_categoria: dict[str, list] = defaultdict(list)
    for f in clasificacion:
        por_categoria[f.categoria_final or "SIN_CLASIFICAR"].append(f)

    # Direcciones del "Saldo al corte" por cuenta, para que DM3 cruce contra
    # la MISMA cifra de la sumaria (ata las dos cédulas por fórmula).
    salida: dict[tuple[str, str], str] = {}
    col_corte = get_column_letter(COL_CORTE)

    tot_anterior = tot_corte = 0.0
    for categoria in sorted(por_categoria, key=lambda c: (_orden_categoria(c), c)):
        filas = sorted(por_categoria[categoria], key=lambda f: f.codigo_cuenta)
        sub_anterior = sub_corte = 0.0
        for f in filas:
            anterior, corte = saldos.get(f.codigo_cuenta, (0.0, 0.0))
            _fila_valores(ws, fila, codigo=f.codigo_cuenta, nombre=f.nombre_cuenta,
                          anterior=anterior, corte=corte)
            salida[("saldo_corte", f.codigo_cuenta)] = f"'{SHEET_DM2}'!{col_corte}{fila}"
            sub_anterior += anterior
            sub_corte += corte
            fila += 1
        _fila_valores(ws, fila, codigo="", nombre=f"Subtotal {categoria}",
                      anterior=round(sub_anterior, 2), corte=round(sub_corte, 2),
                      total=True)
        # Dirección del saldo al corte del SUBTOTAL de la categoría, para que
        # DM3 cruce el saldo al corte por categoría contra lo declarado.
        salida[("saldo_corte_categoria", categoria)] = f"'{SHEET_DM2}'!{col_corte}{fila}"
        fila += 2
        tot_anterior += sub_anterior
        tot_corte += sub_corte

    _fila_valores(ws, fila, codigo="", nombre="TOTAL",
                  anterior=round(tot_anterior, 2), corte=round(tot_corte, 2),
                  total=True)
    fila += 2

    escribir_leyenda_marcas(ws, fila=fila)

    anchos = {COL_CODIGO: 16, COL_NOMBRE: 46, COL_ANTERIOR: 18, COL_VARIACION: 16,
              COL_PCT: 10, COL_CORTE: 18, COL_MARCA: 6}
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    return salida
