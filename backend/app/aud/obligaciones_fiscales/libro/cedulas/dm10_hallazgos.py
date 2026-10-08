"""DM10 · Hoja de hallazgos.

Consolida, en el formato del papel de trabajo de la firma, los puntos a revisar
de cada cédula (DM3..DM8). Las diferencias concretas quedan determinadas por
fórmula dentro de cada cédula (filas "Diferencia"); esta hoja es el resumen
donde el auditor registra la observación y la recomendación de las que superen
su criterio. Se genera siempre (una fila guía por cédula con diferencias) para
que el auditor no tenga que armarla a mano.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from backend.app.aud.obligaciones_fiscales.libro.estilos import (
    BORDE, FONT_DATA, FONT_TITULO_CEDULA, estilo_encabezado_tabla,
    escribir_encabezado_cedula,
)

SHEET_DM10 = "DM10 Hoja de hallazgos"

# Cada cédula que determina diferencias, con la guía de qué revisar en ella.
CEDULAS_REVISABLES = (
    ("DM3", "Revisar las diferencias de crédito tributario, IVA diferido y "
            "obligaciones con el SRI al corte."),
    ("DM4", "Revisar las diferencias del IVA en compras (libros vs F-104)."),
    ("DM5", "Revisar las diferencias de ventas y del IVA en ventas (libros vs F-104)."),
    ("DM6", "Revisar las diferencias de la conciliación de IVA del período."),
    ("DM7", "Revisar las diferencias de retenciones de IVA y de renta (vs F-104/F-103)."),
    ("DM8", "Revisar las diferencias contra el Anexo Transaccional (ATS)."),
)


def build_dm10(
    wb,
    *,
    cliente: str = "",
    periodo: str = "",
    preparado_por: str | None = None,
    revisado_por: str | None = None,
) -> dict:
    """Construye DM10. No publica direcciones: nada la consume por fórmula."""
    if SHEET_DM10 in wb.sheetnames:
        del wb[SHEET_DM10]
    ws = wb.create_sheet(SHEET_DM10)

    escribir_encabezado_cedula(
        ws, titulo="Hoja de Hallazgos", referencia="DM10",
        cliente=cliente, periodo=periodo,
        preparado_por=preparado_por, revisado_por=revisado_por,
    )

    fila = 12
    encabezado = ("Nº", "Referencia de PT", "Observación / Condición", "Recomendación")
    for j, texto in enumerate(encabezado):
        estilo_encabezado_tabla(ws.cell(fila, 1 + j, texto))
    fila += 1

    for i, (ref, guia) in enumerate(CEDULAS_REVISABLES, start=1):
        ws.cell(fila, 1, str(i)).font = FONT_DATA
        ws.cell(fila, 2, ref).font = FONT_DATA
        ws.cell(fila, 3, guia).font = FONT_DATA
        ws.cell(fila, 4, "").font = FONT_DATA  # recomendación: juicio del auditor
        for col in (1, 2, 3, 4):
            ws.cell(fila, col).border = BORDE
        fila += 1

    fila += 1
    ws.cell(fila, 1, "Procedimiento").font = FONT_TITULO_CEDULA
    fila += 1
    ws.cell(fila, 1,
            "Detallar las observaciones sobre las diferencias determinadas en "
            "las cédulas DM3 a DM8 que superen el criterio del auditor, su "
            "impacto tributario y la recomendación. Dejar constancia de las "
            "diferencias no significativas revisadas y depuradas.").font = FONT_DATA

    anchos = {1: 6, 2: 16, 3: 70, 4: 50}
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    return {}
