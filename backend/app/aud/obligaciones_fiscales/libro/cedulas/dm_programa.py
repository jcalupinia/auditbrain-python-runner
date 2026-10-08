"""DM · Programa de Auditoría de Obligaciones Fiscales.

Hoja de portada del papel de trabajo: naturaleza de la cuenta, objetivos de la
prueba y la lista de procedimientos con su referencia a cada cédula. Es
estándar (no depende de cifras del cliente), así que se arma solo con la
cabecera del encargo. Equivale a la hoja "DM Programa" del papel de trabajo de
la firma.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from backend.app.aud.obligaciones_fiscales.libro.estilos import (
    BORDE, FONT_DATA, FONT_ENCABEZADO_TABLA, FONT_ETIQUETA, FONT_TITULO_CEDULA,
    escribir_encabezado_cedula,
)

SHEET_DM_PROGRAMA = "DM  Programa de Auditoria"

NATURALEZA = (
    "Naturaleza de la cuenta",
    "Se debita: las compras de bienes y servicios gravados, el IVA pagado en "
    "compras e importaciones, las retenciones de renta e IVA que le efectuaron "
    "y el pago de las obligaciones al SRI.",
    "Se acredita: el IVA cobrado en ventas, las retenciones de renta e IVA "
    "efectuadas a terceros por pagar al fisco y el impuesto a la renta causado.",
)

OBJETIVOS = (
    "Objetivos",
    "1. Analizar los riesgos tributarios del período y el control interno sobre "
    "las obligaciones fiscales.",
    "2. Confirmar que lo registrado en libros concuerda con lo declarado al SRI "
    "(F-104, F-103) y con el Anexo Transaccional (ATS).",
    "3. Determinar diferencias en IVA, retenciones en la fuente y crédito "
    "tributario, y su efecto en las declaraciones.",
    "4. Dejar documentada la evidencia y los hallazgos para el informe.",
)

# (nº, procedimiento, referencia a la cédula). DM1 (cuestionario) no lo genera
# el sistema todavía; se omite de la lista para no referenciar una hoja que no
# existe.
PROCEDIMIENTOS = (
    ("1", "Elaborar la cédula sumaria de las cuentas de impuestos: saldo del "
          "cierre anterior, saldo al corte y variación.", "DM2"),
    ("2", "Revisar los saldos de crédito tributario, IVA diferido y obligaciones "
          "con el SRI al corte.", "DM3"),
    ("3", "Conciliar el IVA en compras de libros contra el F-104.", "DM4"),
    ("4", "Conciliar las ventas y el IVA en ventas de libros contra el F-104.", "DM5"),
    ("5", "Elaborar el cuadro de conciliación de IVA del período (crédito, "
          "retenciones, saldo a favor o impuesto a pagar).", "DM6"),
    ("6", "Conciliar las retenciones en la fuente de IVA y de renta contra el "
          "F-104 y el F-103.", "DM7"),
    ("7", "Conciliar las compras, ventas y retenciones contra el Anexo "
          "Transaccional (ATS).", "DM8"),
    ("8", "Resumir en la hoja de hallazgos las diferencias determinadas y las "
          "recomendaciones.", "DM10"),
)


def build_dm_programa(
    wb,
    *,
    cliente: str = "",
    periodo: str = "",
    preparado_por: str | None = None,
    revisado_por: str | None = None,
) -> dict:
    """Construye la hoja 'DM Programa de Auditoría'. No publica direcciones."""
    if SHEET_DM_PROGRAMA in wb.sheetnames:
        del wb[SHEET_DM_PROGRAMA]
    ws = wb.create_sheet(SHEET_DM_PROGRAMA)

    escribir_encabezado_cedula(
        ws, titulo="Programa de Auditoría de Obligaciones Fiscales",
        referencia="DM", cliente=cliente, periodo=periodo,
        preparado_por=preparado_por, revisado_por=revisado_por,
    )

    fila = 12
    for bloque in (NATURALEZA, OBJETIVOS):
        ws.cell(fila, 1, bloque[0]).font = FONT_TITULO_CEDULA
        fila += 1
        for texto in bloque[1:]:
            ws.cell(fila, 1, texto).font = FONT_DATA
            fila += 1
        fila += 1

    ws.cell(fila, 1, "Procedimientos").font = FONT_TITULO_CEDULA
    fila += 1
    encabezado = ("Nº", "Procedimiento", "Referencia")
    for j, texto in enumerate(encabezado):
        c = ws.cell(fila, 1 + j, texto)
        c.font = FONT_ENCABEZADO_TABLA
        c.border = BORDE
    fila += 1
    for numero, proc, ref in PROCEDIMIENTOS:
        ws.cell(fila, 1, numero).font = FONT_DATA
        ws.cell(fila, 2, proc).font = FONT_DATA
        ws.cell(fila, 3, ref).font = FONT_ETIQUETA
        for col in (1, 2, 3):
            ws.cell(fila, col).border = BORDE
        fila += 1

    ws.column_dimensions[get_column_letter(1)].width = 6
    ws.column_dimensions[get_column_letter(2)].width = 90
    ws.column_dimensions[get_column_letter(3)].width = 14
    return {}
