"""Exportaciones compartidas: cartas (DOCX), registro de circularización (XLSX) y
vista de consulta (HTML). Ningún dato del cliente entra al Excel como fórmula."""
from io import BytesIO
from decimal import Decimal
from html import escape
from pathlib import Path
from base64 import b64encode
import xlsxwriter

from .plantillas import SHEETS, TITLES, NOTES, REFERENCES, TYPES, TYPE_LABEL_ES, METHOD_LABEL, firm_name


def f(formula, value):
    return {'formula': formula, 'value': value}


def n(value):
    return float(Decimal(str(value if value not in (None, '') else '0')))


def cell_value(value):
    return value.get('value', '') if isinstance(value, dict) else value


def logo_path(r):
    name = 'logo_partner_auditing.png' if r['context']['firm'] == 'partner_auditing' else 'logo_audit_consulting.png'
    return Path(__file__).parent / 'assets' / name


def schedules(r):
    ctx = r['context']
    tables = []

    def add(headers, values):
        tables.append({'name': SHEETS[len(tables)], 'title': TITLES[len(tables)],
                       'headers': headers, 'rows': values, 'note': NOTES[len(tables)]})

    add(['Dato', 'Valor'], [
        ['Cliente', ctx['client']], ['RUC', ctx.get('client_ruc', '')],
        ['País', ctx['country']], ['Moneda', ctx['currency']], ['Ejercicio', str(ctx['year'])],
        ['Corte', ctx['cutoff']], ['Visita', ctx['visit']],
        ['Vuelta', ctx['request_round']], ['Fecha límite de respuesta', ctx['response_deadline']],
        ['Firma auditora', firm_name(ctx['firm'])], ['Correo del auditor', ctx['auditor_email']],
        ['Firmante (cliente)', ctx['signatory']], ['Cargo', ctx.get('signatory_role', '')],
        ['Preparado por', ctx['preparer']], ['Revisado por', ctx['reviewer']],
        ['Estado', r['status']], ['Versión', r['version']], ['SHA256 solicitud', r['input_sha256']],
    ])
    add(['ID', 'Rubro', 'Entidad', 'Contacto', 'Correo', 'Referencia', 'Método', 'Saldo'],
        [[it['id'], TYPE_LABEL_ES[it['type']], it['entity'], it['contact_name'], it['contact_email'],
          it['reference'] or it['account_ref'], it['method'],
          it['amount'] if it['amount'] is not None else '—'] for it in r['items']])
    add(['ID', 'Rubro', 'Entidad', 'Método', 'Asunto'],
        [[l['id'], l['type_label'], l['entity'], l['method'], l['subject']] for l in r['letters']])
    add(['ID', 'Para (correo)', 'Asunto', 'Método', 'Estado', 'SHA256 cuerpo'],
        [[d['id'], d['to'] or 'SIN CORREO', d['subject'], d['method'], d['status'], d['body_sha256']]
         for d in r['dispatch']])
    # 04 Control de confirmaciones — formato de la firma (agrupado por rubro).
    control_rows = []
    groups = {}
    for it in r['items']:
        groups.setdefault(it['type'], []).append(it)
    excel_row = 6  # el cuerpo de la tabla arranca en la fila 6 de Excel
    for rtype in sorted(groups):
        control_rows.append([TYPE_LABEL_ES[rtype].upper(), '', '', '', '', '', '', '', ''])
        excel_row += 1
        for it in groups[rtype]:
            libros = n(it['amount']) if it['amount'] is not None else ''
            dif = f(f'=IF(OR(E{excel_row}="",F{excel_row}=""),"",E{excel_row}-F{excel_row})', '')
            control_rows.append(['', it['entity'], '', '', libros, '', dif, '', ''])
            excel_row += 1
    add(['Descripción', 'Detalle', 'Enviado (Sí/No)', 'Recibido (Sí/No)', 'Libros',
         'Confirmación', 'Diferencia', 'Gestión con cliente', 'Observación'], control_rows)
    coverage_rows = []
    for i, c in enumerate(r['coverage'], 6):
        if c['ledger'] is not None:
            coverage_rows.append([c['label'], c['count'], n(c['sampled']), n(c['ledger']),
                                  f(f'=D{i}-C{i}', n(c['difference'])),
                                  f(f'=IF(D{i}>0,ROUND(C{i}/D{i}*100,2),0)', n(c['coverage_pct'])),
                                  'Dentro de tolerancia' if c['within_tolerance'] else 'Revisar diferencia'])
        else:
            coverage_rows.append([c['label'], c['count'], n(c['sampled']), '—', '—', '—',
                                  'Sin saldo del mayor'])
    add(['Rubro', 'Nº cartas', 'Circularizado', 'Mayor', 'Diferencia', 'Cobertura %', 'Estado'],
        coverage_rows or [['Sin cobertura declarada', 0, 0, '—', '—', '—', 'Declare saldos del mayor por rubro']])
    t = r['totals']
    add(['Aspecto', 'Resultado'], [
        ['Conclusión', 'Pendiente de evaluación de las respuestas'],
        ['Cartas generadas', t['count']],
        ['Con correo / sin correo', f"{t['with_email']} / {t['without_email']}"],
        ['Importe total circularizado', n(t['total_sampled'])],
        ['Por rubro', ', '.join(f'{TYPE_LABEL_ES[k]}: {v}' for k, v in sorted(t['by_type'].items()))],
        ['Por método', ', '.join(f'{METHOD_LABEL[k].split(" (")[0]}: {v}' for k, v in sorted(t['by_method'].items()))],
        ['Procedimientos', 'NIA 505: control del auditor sobre el envío y la respuesta; procedimientos alternativos ante no respuestas y evaluación de discrepancias.'],
        ['Excepciones', 'Las diferencias y no respuestas se detallan en 04 Control; la cobertura, en 05 Cobertura.'],
    ])
    add(['Control', 'Estado'], [
        ['Preparado por', ctx['preparer']], ['Revisado por', ctx['reviewer']],
        ['Control del envío', 'El auditor conserva el control del envío y la recepción (NIA 505)'],
        ['Aprobación profesional', 'PENDIENTE — nombres declarados no son firmas'],
        ['Conclusión del revisor', 'Completar en papeles de trabajo'],
        ['Trazabilidad', r['input_sha256']],
    ])
    return tables


def build_xlsx(r):
    buf = BytesIO()
    wb = xlsxwriter.Workbook(buf, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    wb.set_properties({'title': 'Confirmaciones · ' + r['context']['client'], 'author': r['context']['firm'],
                       'comments': 'AuditBrain Confirmaciones ' + r['version']})
    title = wb.add_format({'bold': True, 'font_size': 18, 'font_color': '#FFFFFF', 'bg_color': '#102D4D', 'valign': 'vcenter'})
    header = wb.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#217A55', 'border': 1, 'text_wrap': True})
    text = wb.add_format({'font_name': 'Calibri', 'font_size': 10, 'border': 1, 'border_color': '#D9E3E9', 'valign': 'top', 'text_wrap': True})
    numeric = wb.add_format({'font_name': 'Calibri', 'font_size': 10, 'border': 1, 'border_color': '#D9E3E9', 'num_format': '#,##0.00;[Red](#,##0.00)', 'valign': 'top'})
    formula = wb.add_format({'font_name': 'Calibri', 'font_size': 10, 'border': 1, 'border_color': '#D9E3E9', 'bg_color': '#EAF6EF', 'num_format': '#,##0.00;[Red](#,##0.00)', 'text_wrap': True})
    note = wb.add_format({'font_size': 10, 'font_color': '#435766', 'text_wrap': True, 'valign': 'top'})
    firm = firm_name(r['context']['firm'])
    for table in schedules(r):
        ws = wb.add_worksheet(table['name'])
        width = max(2, len(table['headers']))
        ws.hide_gridlines(2); ws.set_tab_color('#217A55'); ws.set_landscape(); ws.set_paper(9)
        ws.fit_to_pages(1, 0); ws.repeat_rows(0, 4); ws.freeze_panes(5, 1)
        ws.set_column(0, width - 1, 19); ws.set_column(0, 0, 26)
        if width == 2:
            ws.set_column(1, 1, 92)
        ws.merge_range(0, 0, 0, width - 1, table['title'], title); ws.set_row(0, 30)
        ws.merge_range(1, 0, 1, width - 1, firm + ' · ' + r['context']['client'], note)
        ws.merge_range(2, 0, 2, width - 1, f"{r['context']['cutoff']} · {r['context']['currency']} · {r['status']}", note)
        if table['name'] == '00 Portada':
            logo = logo_path(r)
            from PIL import Image
            with Image.open(logo) as im:
                scale = min(180 / im.width, 55 / im.height)
            ws.set_row(3, 48)
            ws.insert_image(3, 0, str(logo), {'x_scale': scale, 'y_scale': scale, 'object_position': 1, 'description': firm})
        ws.write_row(4, 0, table['headers'], header); ws.set_row(4, 32)
        for row_idx, row in enumerate(table['rows'], 5):
            ws.set_row(row_idx, 26)
            for col, v in enumerate(row):
                if isinstance(v, dict):
                    ws.write_formula(row_idx, col, v['formula'], formula, v['value'])
                elif isinstance(v, (int, float)):
                    ws.write_number(row_idx, col, v, numeric)
                else:
                    ws.write_string(row_idx, col, str(v), text)
        end = 5 + len(table['rows'])
        ws.autofilter(4, 0, max(5, end - 1), width - 1)
        ws.merge_range(end + 2, 0, end + 4, width - 1, 'CÓMO SE USA / PROCEDIMIENTO: ' + table['note'], note)
        ws.set_footer('&L' + firm + '&R&P / &N')
        ws.print_area(0, 0, end + 5, width - 1)
    wb.close()
    return buf.getvalue()


def _docx_letter(doc, r, letter):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    blocks = letter['blocks']
    for line in letter['salutation']:
        doc.add_paragraph(line)
    if blocks:
        doc.add_paragraph(blocks[0]).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for it in letter.get('items', []):
        p = doc.add_paragraph(it)
        try:
            p.style = doc.styles['List Bullet']
        except Exception:
            pass
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for block in blocks[1:]:
        doc.add_paragraph(block).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph('')
    for line in letter['signature']:
        doc.add_paragraph(line)
    if letter['response_lines']:
        doc.add_paragraph('')
        rule = doc.add_paragraph()
        rule.add_run('— — — — — — — — — — — — — — — — — — — — — — — — — — —')
        head = doc.add_paragraph()
        head.add_run(letter['response_title']).bold = True
        for line in letter['response_lines']:
            doc.add_paragraph(line)


def build_docx(r):
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.core_properties.title = 'Cartas de confirmación · ' + r['context']['client']
    doc.core_properties.author = r['context']['firm']
    firm = firm_name(r['context']['firm'])
    # Portada
    try:
        doc.add_picture(str(logo_path(r)), width=Inches(2.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass
    h = doc.add_heading('Cartas de confirmación de saldos', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run(f"{firm} · {r['context']['client']} · corte {r['context']['cutoff']} · {r['status']}")
    doc.add_paragraph(f"{r['totals']['count']} cartas · borrador para revisión. Cada carta la firma el cliente y "
                      "se responde directamente al auditor (NIA 505).").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for letter in r['letters']:
        doc.add_page_break()
        _docx_letter(doc, r, letter)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_html(r):
    esc = lambda v: escape(str(v), quote=True)
    ctx = r['context']
    logo = b64encode(logo_path(r).read_bytes()).decode('ascii')
    nav, sections = [], []
    # Cartas con enlace mailto para envío inmediato desde el navegador.
    cards = []
    for l in r['letters']:
        to = next((d['to'] for d in r['dispatch'] if d['id'] == l['id']), '')
        body = l['text']
        mailto = ('<a class="send" href="mailto:' + esc(to) + '?subject=' + esc(l['subject']).replace(' ', '%20') +
                  '&body=' + esc(body).replace('\n', '%0D%0A').replace(' ', '%20') + '">Enviar por correo</a>') if to else '<span class="nomail">Sin correo del contacto</span>'
        cards.append(f'<article class="letter"><h3>{esc(l["id"])} · {esc(l["type_label"])} — {esc(l["entity"])}</h3>'
                     f'<p class="meta">{esc(l["method_label"])} · {mailto}</p><pre>{esc(body)}</pre></article>')
    sections.append('<section id="cartas"><h2>Cartas generadas</h2>' + ''.join(cards) + '</section>')
    nav.append('<a href="#cartas">Cartas</a>')
    for i, table in enumerate(schedules(r)):
        nav.append(f'<a href="#s{i}">{i + 1:02d} {esc(table["title"])}</a>')
        rows_html = ''.join('<tr>' + ''.join('<td>' + esc(cell_value(v)) +
                            (f'<details><summary>Fórmula</summary><code>{esc(v["formula"])}</code></details>' if isinstance(v, dict) else '') +
                            '</td>' for v in row) + '</tr>' for row in table['rows'])
        sections.append(f'<section id="s{i}"><h2>{esc(table["title"])}</h2><div class="scroll"><table><thead><tr>' +
                        ''.join('<th>' + esc(hh) + '</th>' for hh in table['headers']) +
                        '</tr></thead><tbody>' + rows_html + '</tbody></table></div><p><b>Cómo se usa / procedimiento:</b> ' +
                        esc(table['note']) + '</p></section>')
    refs = ''.join(f'<li><a href="{esc(x["url"])}">{esc(x["title"])}</a></li>' for x in REFERENCES)
    return ('<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            '<title>AuditBrain · Confirmaciones de saldos</title><style>'
            'body{margin:0;background:#0d2037;color:#eef5fb;font:15px/1.6 Arial}header,main{max-width:1250px;margin:auto;padding:24px}'
            'header{border-bottom:1px solid #2c4765}h1,h2{color:#75e49c}nav{display:flex;gap:10px;flex-wrap:wrap}a{color:#75e49c}'
            'section{margin:22px 0;padding:22px;background:#142d4b;border:1px solid #2c4765;border-radius:12px}'
            '.scroll{overflow:auto}table{border-collapse:collapse;width:100%}th,td{border:1px solid #38516c;padding:10px;text-align:left;vertical-align:top}'
            'th{background:#17563f}td,pre{white-space:pre-wrap;overflow-wrap:anywhere}.letter{background:#0f2540;border:1px solid #2c4765;border-radius:10px;padding:16px;margin:14px 0}'
            '.letter pre{background:#0b1c30;padding:14px;border-radius:8px;font:13px/1.55 "Courier New",monospace}.meta{color:#9fc4dd}'
            '.send{background:#17563f;color:#eaffef;padding:5px 12px;border-radius:6px;text-decoration:none}.nomail{color:#f5b3b3}'
            'code{font-size:12px}@media print{body,section,.letter{background:white;color:black}h1,h2,a{color:#17563f}nav,.send{display:none}section,.letter{break-inside:avoid}}'
            '</style><header><img alt="Firma auditora" style="max-width:180px;max-height:75px" src="data:image/png;base64,' + logo + '">'
            '<h1>Confirmaciones de saldos</h1><p>' + esc(ctx['client']) + ' · ' + esc(ctx['cutoff']) + ' · ' + esc(ctx['currency']) +
            '</p><p>Borrador para revisión · Los enlaces «Enviar por correo» abren el correo del auditor con la carta lista. '
            'El auditor conserva el control del envío y de la respuesta (NIA 505).</p><nav>' + ''.join(nav) +
            '</nav></header><main>' + ''.join(sections) +
            '<section><h2>Referencias</h2><ul>' + refs + '</ul></section></main></html>')


def letter_email_html(r, letter):
    """HTML de una sola carta para envío por correo (estilos en línea, sin recursos externos)."""
    esc = lambda v: escape(str(v), quote=True)
    firm = firm_name(r['context']['firm'])
    blocks = letter['blocks']
    salut = ''.join(f'<p style="margin:2px 0">{esc(l)}</p>' for l in letter['salutation'])
    intro = f'<p style="margin:14px 0">{esc(blocks[0])}</p>' if blocks else ''
    items = ''.join(f'<li style="margin:4px 0">{esc(it)}</li>' for it in letter.get('items', []))
    items_html = f'<ul style="margin:8px 0 8px 18px;padding:0">{items}</ul>' if items else ''
    rest = ''.join(f'<p style="margin:12px 0">{esc(b)}</p>' for b in blocks[1:])
    sign = ''.join(f'<p style="margin:2px 0">{esc(l)}</p>' for l in letter['signature'])
    resp = ''
    if letter['response_lines']:
        rl = ''.join(f'<p style="margin:3px 0">{esc(l)}</p>' for l in letter['response_lines'])
        resp = ('<div style="margin-top:18px;padding:14px;border:1px dashed #94a3b8;border-radius:8px">'
                f'<p style="margin:0 0 8px;font-weight:bold">{esc(letter["response_title"])}</p>{rl}</div>')
    return (
        '<div style="font-family:Arial,Helvetica,sans-serif;font-size:14px;color:#1f2937;'
        'max-width:680px;margin:auto;line-height:1.55">'
        '<div style="background:#0a2540;color:#fff;padding:14px 18px;border-radius:8px 8px 0 0">'
        f'<strong>{esc(firm)}</strong></div>'
        '<div style="padding:18px;border:1px solid #e2e8f0;border-top:none;border-radius:0 0 8px 8px">'
        f'{salut}{intro}{items_html}{rest}<div style="margin-top:16px">{sign}</div>{resp}'
        '</div></div>'
    )
