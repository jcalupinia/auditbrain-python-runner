"""Papel de trabajo del análisis de estados financieros (NIA 315 / NIA 520).

Arma el `.xlsx` desde la salida de ``analisis.analizar``: hojas ESF y ERI con
los valores homologados (datos) y las columnas Variación, Var % y Vertical %
como **fórmulas**; hoja Ratios cuyos valores son fórmulas que remiten a las
celdas del ESF/ERI; y hoja «NIA 520» con la expectativa, la diferencia (fórmula)
y la marca de lo que supera el umbral. Escapa texto que parece fórmula.
"""
from __future__ import annotations

import base64
import html as _html
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_AZUL = "0A2342"
_BLANCO = "FFFFFF"
_FONDO_TOTAL = "E8EEF6"
_F_TITULO = Font(name="Calibri", size=14, bold=True, color=_AZUL)
_F_SUB = Font(name="Calibri", size=9, italic=True, color="6B7280")
_F_ENCAB = Font(name="Calibri", size=10, bold=True, color=_BLANCO)
_F_DATO = Font(name="Calibri", size=9)
_F_TOTAL = Font(name="Calibri", size=10, bold=True, color=_AZUL)
_RELLENO_ENCAB = PatternFill("solid", fgColor=_AZUL)
_RELLENO_ALERTA = PatternFill("solid", fgColor="F6E0E0")
_THIN = Side(style="thin", color="B8C4D0")
_BORDE = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DER = Alignment(horizontal="right")
_IZQ = Alignment(horizontal="left", vertical="top", wrap_text=True)
_CEN = Alignment(horizontal="center")
_FMT_MONEDA = '#,##0.00'
_FMT_PCT = '0.00%'
_FMT_VECES = '0.00'
_FILA0 = 4  # fila de datos (encabezados en la fila 3)


def _seguro(v) -> str:
    s = "" if v is None else str(v)
    return " " + s if s[:1] in ("=", "+", "-", "@") else s


def _encabezados(ws, fila, titulos):
    for c, t in enumerate(titulos, start=1):
        cel = ws.cell(row=fila, column=c, value=t)
        cel.font = _F_ENCAB
        cel.fill = _RELLENO_ENCAB
        cel.border = _BORDE
        cel.alignment = _CEN


def _hoja_estado(wb, titulo_hoja, titulo, estado, es_activa=False):
    """Escribe una hoja de estado y devuelve (ws, {codigo: fila_excel}, col_ultimo_periodo)."""
    ws = wb.active if es_activa else wb.create_sheet(titulo_hoja)
    if es_activa:
        ws.title = titulo_hoja
    ws.sheet_view.showGridLines = False
    periodos = estado.get("periodos", [])
    lineas = estado.get("lineas", [])
    n = len(periodos)
    ws["A1"] = titulo
    ws["A1"].font = _F_TITULO
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 46
    # Columnas: A Código, B Rubro, C.. períodos, Variación, Var %, Vertical %
    cols = ["Código", "Rubro"] + list(periodos) + ["Variación", "Var %", "Vertical %"]
    _encabezados(ws, _FILA0 - 1, cols)
    for idx in range(3, 3 + n + 3):
        ws.column_dimensions[get_column_letter(idx)].width = 15
    col_p0 = 3                      # primera columna de período
    col_pu = col_p0 + n - 1         # última columna de período
    col_var = col_pu + 1
    col_pct = col_var + 1
    col_ver = col_pct + 1
    L = get_column_letter
    base_codigo = estado.get("base_codigo")
    filas_por_codigo = {l["codigo"]: (_FILA0 + i) for i, l in enumerate(lineas)}
    fila_base = filas_por_codigo.get(base_codigo)

    for i, l in enumerate(lineas):
        r = _FILA0 + i
        es_total = len(str(l["codigo"])) <= 1
        ws.cell(row=r, column=1, value=_seguro(l["codigo"])).font = _F_DATO
        cel_rub = ws.cell(row=r, column=2, value=_seguro(l["etiqueta"]))
        cel_rub.font = _F_TOTAL if es_total else _F_DATO
        cel_rub.alignment = _IZQ
        for j in range(n):
            cel = ws.cell(row=r, column=col_p0 + j, value=round(float(l["valores"][j]), 2))
            cel.number_format = _FMT_MONEDA
            cel.alignment = _DER
            cel.font = _F_DATO
        if n >= 2:
            cprev, cact = L(col_pu - 1), L(col_pu)
            cv = ws.cell(row=r, column=col_var, value=f"={cact}{r}-{cprev}{r}")
            cv.number_format = _FMT_MONEDA; cv.alignment = _DER; cv.font = _F_DATO
            cp = ws.cell(row=r, column=col_pct,
                         value=f"=IF({cprev}{r}=0,\"\",({cact}{r}-{cprev}{r})/{cprev}{r})")
            cp.number_format = _FMT_PCT; cp.alignment = _DER; cp.font = _F_DATO
        if fila_base:
            cu = L(col_pu)
            cver = ws.cell(row=r, column=col_ver,
                           value=f"=IF({cu}${fila_base}=0,\"\",{cu}{r}/{cu}${fila_base})")
            cver.number_format = _FMT_PCT; cver.alignment = _DER; cver.font = _F_DATO
        for c in range(1, col_ver + 1):
            ws.cell(row=r, column=c).border = _BORDE
    ws.freeze_panes = f"A{_FILA0}"
    return ws, filas_por_codigo, col_pu


def _hoja_ratios(wb, ratios, filas_esf, filas_eri, col_pu):
    ws = wb.create_sheet("Ratios")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Ratios financieros de la firma"
    ws["A1"].font = _F_TITULO
    ws.column_dimensions["A"].width = 32
    periodos = ratios.get("periodos", [])
    n = len(periodos)
    _encabezados(ws, _FILA0 - 1, ["Ratio"] + list(periodos))
    for idx in range(2, 2 + n):
        ws.column_dimensions[get_column_letter(idx)].width = 16
    L = get_column_letter
    hoja = {"esf": ("ESF", filas_esf), "eri": ("ERI", filas_eri)}

    def _codigos(spec):
        return [spec[1], spec[2]] if isinstance(spec, (list, tuple)) else [spec]

    def _ref(origen, spec, j):
        """Referencia Excel al componente, o None si algún código no está en la
        hoja (línea en cero, filtrada): el ratio cae a valor literal."""
        nombre_hoja, filas = hoja[origen]
        col = L(3 + j)  # C = primer período en ESF/ERI
        if any(c not in filas for c in _codigos(spec)):
            return None
        if isinstance(spec, (list, tuple)):
            _, a, b = spec
            return f"('{nombre_hoja}'!{col}{filas[a]}-'{nombre_hoja}'!{col}{filas[b]})"
        return f"'{nombre_hoja}'!{col}{filas[spec]}"

    from backend.app.aud.motor_balances.analisis import RATIOS
    spec_por_nombre = {r[0]: r for r in RATIOS}
    for i, fila in enumerate(ratios.get("filas", [])):
        r = _FILA0 + i
        ws.cell(row=r, column=1, value=_seguro(fila["nombre"])).font = _F_TOTAL
        _, num_spec, den_spec, origen, fmt = spec_por_nombre[fila["nombre"]]
        for j in range(n):
            num = _ref(origen, num_spec, j)
            den = _ref(origen, den_spec, j)
            valor = fila["valores"][j] if j < len(fila["valores"]) else None
            if num is not None and den is not None:
                celda_val = f"=IF({den}=0,\"\",{num}/{den})"
            else:  # caso borde: algún componente en cero/ausente → valor calculado
                celda_val = round(valor, 6) if valor is not None else ""
            cel = ws.cell(row=r, column=2 + j, value=celda_val)
            cel.number_format = _FMT_PCT if fmt == "pct" else _FMT_VECES
            cel.alignment = _DER; cel.font = _F_DATO
        for c in range(1, 2 + n):
            ws.cell(row=r, column=c).border = _BORDE


def _hoja_nia520(wb, expectativa, titulo):
    ws = wb.create_sheet("NIA 520")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Expectativa vs. real (NIA 520)"
    ws["A1"].font = _F_TITULO
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 46
    for c in "CDEF":
        ws.column_dimensions[c].width = 15
    ws.column_dimensions["G"].width = 12
    if not expectativa.get("aplicable"):
        ws.cell(row=3, column=1, value="Se requieren al menos dos períodos para la expectativa NIA 520.").font = _F_SUB
        return
    umbral = expectativa.get("umbral_pct", 0.10)
    ws["A2"] = f"{titulo} · umbral de diferencia: {umbral:.0%}"
    ws["A2"].font = _F_SUB
    _encabezados(ws, _FILA0 - 1, ["Código", "Rubro", "Expectativa", "Real", "Diferencia", "Dif %", "¿Explicar?"])
    for i, l in enumerate(expectativa.get("lineas", [])):
        r = _FILA0 + i
        ws.cell(row=r, column=1, value=_seguro(l["codigo"])).font = _F_DATO
        cr = ws.cell(row=r, column=2, value=_seguro(l["etiqueta"])); cr.font = _F_DATO; cr.alignment = _IZQ
        ws.cell(row=r, column=3, value=round(l["expectativa"], 2)).number_format = _FMT_MONEDA
        ws.cell(row=r, column=4, value=round(l["real"], 2)).number_format = _FMT_MONEDA
        cd = ws.cell(row=r, column=5, value=f"=D{r}-C{r}"); cd.number_format = _FMT_MONEDA
        cp = ws.cell(row=r, column=6, value=f"=IF(C{r}=0,\"\",(D{r}-C{r})/C{r})"); cp.number_format = _FMT_PCT
        marca = "Explicar" if l["supera_umbral"] else ""
        cm = ws.cell(row=r, column=7, value=marca); cm.alignment = _CEN; cm.font = _F_TOTAL
        for c in range(1, 8):
            cel = ws.cell(row=r, column=c)
            cel.border = _BORDE
            if l["supera_umbral"]:
                cel.fill = _RELLENO_ALERTA


def generar_papel_estados(analisis: dict) -> bytes:
    """`.xlsx` del análisis de estados financieros."""
    wb = Workbook()
    ws_esf, filas_esf, col_pu = _hoja_estado(
        wb, "ESF", "Estado de Situación Financiera — análisis", analisis["esf"], es_activa=True)
    _hoja_estado(wb, "ERI", "Estado de Resultados Integral — análisis", analisis["eri"])
    filas_eri = {l["codigo"]: (_FILA0 + i) for i, l in enumerate(analisis["eri"]["lineas"])}
    _hoja_ratios(wb, analisis["ratios"], filas_esf, filas_eri, col_pu)
    _hoja_nia520(wb, analisis["expectativa_esf"], "Estado de Situación Financiera")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- HTML autónomo (Excel embebido, imprimible a PDF) -----------------------
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ESTILO_HTML = """
*{box-sizing:border-box}body{margin:0;font-family:'Segoe UI',Calibri,Arial,sans-serif;color:#1a2433;background:#eef1f5}
.hoja{max-width:1200px;margin:0 auto;padding:24px 16px 64px}
.banda{background:#071B2F;color:#fff;padding:18px 22px;border-radius:12px;display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:12px}
.banda h1{margin:0;font-size:20px}.banda .marca{font-size:12px;color:#cdd6e2}.oro{color:#C7A83C;font-weight:700}
.barra{display:flex;gap:10px;margin:16px 0;flex-wrap:wrap}
.btn{border:0;border-radius:8px;padding:9px 16px;font-size:13px;font-weight:600;cursor:pointer;text-decoration:none;display:inline-block}
.btn-oro{background:#C7A83C;color:#20180a}.btn-navy{background:#0A2342;color:#fff}
section.bloque{background:#fff;border-radius:12px;padding:16px 18px;margin-top:16px;box-shadow:0 1px 3px rgba(10,35,66,.08)}
section.bloque h2{margin:0 0 12px;font-size:15px;color:#0A2342}
table{width:100%;border-collapse:collapse;font-size:12px}
th{background:#0A2342;color:#fff;text-align:left;padding:7px 9px}
td{padding:6px 9px;border-bottom:1px solid #e3e8ef}td.num{text-align:right;font-variant-numeric:tabular-nums}
tr.total td{font-weight:700;background:#eef3fa}tr.alerta td{background:#F6E0E0}
.disc{color:#6B7280;font-size:10.5px;font-style:italic;margin-top:24px;text-align:center}
@media print{@page{size:landscape;margin:12mm}body{background:#fff;-webkit-print-color-adjust:exact;print-color-adjust:exact}.barra{display:none}section.bloque{box-shadow:none;break-inside:avoid}}
"""


def _tabla_html(columnas, filas):
    ths = "".join(f"<th>{_html.escape(str(t))}</th>" for t, _ in columnas)
    cuerpo = []
    for fila in filas:
        clase, celdas = "", fila
        if isinstance(fila, dict):
            clase, celdas = fila.get("clase", ""), fila["cells"]
        tds = []
        for i, v in enumerate(celdas):
            cls = ' class="num"' if columnas[i][1] == "num" else ""
            tds.append(f"<td{cls}>{_html.escape('' if v is None else str(v))}</td>")
        attr = f' class="{clase}"' if clase else ""
        cuerpo.append(f"<tr{attr}>{''.join(tds)}</tr>")
    return f"<table><thead><tr>{ths}</tr></thead><tbody>{''.join(cuerpo)}</tbody></table>"


def _money(v):
    try:
        return f"{float(v or 0):,.2f}"
    except (TypeError, ValueError):
        return str(v)


def _pct(v):
    if v in (None, ""):
        return "—"
    try:
        return f"{float(v) * 100:.2f} %"
    except (TypeError, ValueError):
        return str(v)


def _bloque_estado(det, titulo):
    periodos = det.get("periodos", [])
    cols = [("Código", "txt"), ("Rubro", "txt")] + [(p, "num") for p in periodos] + \
           [("Variación", "num"), ("Var %", "num"), ("Vertical %", "num")]
    filas = []
    for l in det.get("lineas", []):
        cells = [l["codigo"], l["etiqueta"]] + [_money(v) for v in l["valores"]]
        cells += [_money(l.get("variacion")) if l.get("variacion") is not None else "—",
                  _pct(l.get("variacion_pct")), _pct(l.get("vertical", [None])[-1])]
        filas.append({"cells": cells, "clase": "total" if len(str(l["codigo"])) <= 1 else ""})
    return {"titulo": titulo, "html": _tabla_html(cols, filas)}


def generar_html_estados(analisis: dict) -> str:
    """HTML autónomo del análisis de estados financieros (Excel embebido)."""
    ratios = analisis.get("ratios", {})
    cols_r = [("Ratio", "txt")] + [(p, "num") for p in ratios.get("periodos", [])]
    filas_r = []
    for f in ratios.get("filas", []):
        vals = [(_pct(v) if f["formato"] == "pct" else (f"{float(v):.2f}" if v is not None else "—"))
                for v in f["valores"]]
        filas_r.append([f["nombre"]] + vals)
    bloques = [{"titulo": "Ratios financieros", "html": _tabla_html(cols_r, filas_r)},
               _bloque_estado(analisis["esf"], "Estado de Situación Financiera"),
               _bloque_estado(analisis["eri"], "Estado de Resultados Integral")]
    exp = analisis.get("expectativa_esf", {})
    if exp.get("aplicable"):
        cols_e = [("Código", "txt"), ("Rubro", "txt"), ("Expectativa", "num"), ("Real", "num"),
                  ("Diferencia", "num"), ("Dif %", "num"), ("¿Explicar?", "txt")]
        filas_e = [{"cells": [l["codigo"], l["etiqueta"], _money(l["expectativa"]), _money(l["real"]),
                              _money(l["diferencia"]), _pct(l["diferencia_pct"]),
                              "Explicar" if l["supera_umbral"] else ""],
                    "clase": "alerta" if l["supera_umbral"] else ""} for l in exp.get("lineas", [])]
        bloques.append({"titulo": f"Expectativa vs. real — NIA 520 (umbral {_pct(exp.get('umbral_pct'))})",
                        "html": _tabla_html(cols_e, filas_e)})

    b64 = base64.b64encode(generar_papel_estados(analisis)).decode("ascii")
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    secciones = "".join(
        f'<section class="bloque"><h2>{_html.escape(b["titulo"])}</h2>{b["html"]}</section>'
        for b in bloques)
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Papel de trabajo — Estados financieros</title><style>{_ESTILO_HTML}</style></head>
<body><div class="hoja"><div class="banda">
<div><h1>Papel de trabajo — Estados financieros</h1><div class="marca">Análisis NIA 315 / NIA 520</div></div>
<div class="marca">AuditConsulting Auditores Cía. Ltda. · <span class="oro">AUDIT-IA</span><br>{fecha}</div></div>
<div class="barra">
<a class="btn btn-oro" href="data:{_XLSX_MIME};base64,{b64}" download="papel-estados-financieros.xlsx">Descargar Excel con fórmulas</a>
<button class="btn btn-navy" onclick="window.print()">Imprimir / Guardar como PDF</button></div>
{secciones}
<div class="disc">Papel de trabajo generado por AUDIT-IA. Las cifras deben ser validadas por el auditor responsable.</div>
</div></body></html>"""


# --- Word / PowerPoint ------------------------------------------------------
def _estructura_estados(analisis: dict) -> dict:
    ratios = analisis.get("ratios", {})
    cols_r = [("Ratio", "txt")] + [(p, "num") for p in ratios.get("periodos", [])]
    filas_r = []
    for f in ratios.get("filas", []):
        vals = [(_pct(v) if f["formato"] == "pct" else (f"{float(v):.2f}" if v is not None else "—"))
                for v in f["valores"]]
        filas_r.append([f["nombre"]] + vals)

    def _bloque(det, titulo):
        periodos = det.get("periodos", [])
        cols = [("Código", "txt"), ("Rubro", "txt")] + [(p, "num") for p in periodos] + \
               [("Variación", "num"), ("Var %", "num"), ("Vertical %", "num")]
        filas = []
        for l in det.get("lineas", []):
            cells = [l["codigo"], l["etiqueta"]] + [_money(v) for v in l["valores"]]
            cells += [_money(l.get("variacion")) if l.get("variacion") is not None else "—",
                      _pct(l.get("variacion_pct")), _pct(l.get("vertical", [None])[-1])]
            filas.append(cells)
        return {"titulo": titulo, "columnas": cols, "filas": filas}

    bloques = [{"titulo": "Ratios financieros", "columnas": cols_r, "filas": filas_r},
               _bloque(analisis["esf"], "Estado de Situación Financiera"),
               _bloque(analisis["eri"], "Estado de Resultados Integral")]
    exp = analisis.get("expectativa_esf", {})
    if exp.get("aplicable"):
        cols_e = [("Código", "txt"), ("Rubro", "txt"), ("Expectativa", "num"), ("Real", "num"),
                  ("Diferencia", "num"), ("Dif %", "num"), ("¿Explicar?", "txt")]
        filas_e = [[l["codigo"], l["etiqueta"], _money(l["expectativa"]), _money(l["real"]),
                    _money(l["diferencia"]), _pct(l["diferencia_pct"]),
                    "Explicar" if l["supera_umbral"] else ""] for l in exp.get("lineas", [])]
        bloques.append({"titulo": "Expectativa vs. real — NIA 520", "columnas": cols_e, "filas": filas_e})
    return {"titulo": "Papel de trabajo — Estados financieros",
            "subtitulo": "Análisis NIA 315 / NIA 520", "bloques": bloques}


def generar_office_estados(analisis: dict, formato: str) -> bytes:
    """`docx` o `pptx` del análisis de estados financieros."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt as DPt, RGBColor as DColor
    from pptx import Presentation
    from pptx.dml.color import RGBColor as PColor
    from pptx.util import Inches, Pt as PPt

    est = _estructura_estados(analisis)
    if formato == "docx":
        doc = Document()
        doc.add_heading(est["titulo"], level=0)
        p = doc.add_paragraph(est["subtitulo"])
        if p.runs:
            p.runs[0].font.size = DPt(9)
        doc.add_paragraph("AuditConsulting Auditores Cía. Ltda. · AUDIT-IA").runs[0].font.color.rgb = DColor(0x6B, 0x72, 0x80)
        for b in est["bloques"]:
            doc.add_heading(b["titulo"], level=1)
            tipos = [t for _, t in b["columnas"]]
            tabla = doc.add_table(rows=1, cols=len(b["columnas"]))
            tabla.style = "Light Grid Accent 1"
            for j, (nombre, _) in enumerate(b["columnas"]):
                run = tabla.rows[0].cells[j].paragraphs[0].add_run(str(nombre)); run.bold = True; run.font.size = DPt(8)
            for fila in b["filas"]:
                celdas = tabla.add_row().cells
                for j, v in enumerate(fila):
                    par = celdas[j].paragraphs[0]
                    par.add_run("" if v is None else str(v)).font.size = DPt(8)
                    if tipos[j] == "num":
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

    prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    caja = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(1.8)).text_frame
    caja.text = est["titulo"]; caja.paragraphs[0].runs[0].font.size = PPt(32)
    caja.paragraphs[0].runs[0].font.color.rgb = PColor(0x0A, 0x23, 0x42)
    sp = caja.add_paragraph(); sp.text = est["subtitulo"] + " · AUDIT-IA"
    sp.runs[0].font.size = PPt(14); sp.runs[0].font.color.rgb = PColor(0xC7, 0xA8, 0x3C)
    for b in est["bloques"]:
        s = prs.slides.add_slide(blank)
        t = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.7)).text_frame
        t.text = b["titulo"]; t.paragraphs[0].runs[0].font.size = PPt(22)
        t.paragraphs[0].runs[0].font.color.rgb = PColor(0x0A, 0x23, 0x42)
        filas = b["filas"][:12]
        n = len(filas) + 1
        tabla = s.shapes.add_table(n, len(b["columnas"]), Inches(0.6), Inches(1.2),
                                   Inches(12.1), Inches(0.3) * n).table
        for j, (nombre, _) in enumerate(b["columnas"]):
            tabla.cell(0, j).text = str(nombre)
            tabla.cell(0, j).text_frame.paragraphs[0].runs[0].font.size = PPt(9)
        for i, fila in enumerate(filas, start=1):
            for j, v in enumerate(fila):
                tabla.cell(i, j).text = "" if v is None else str(v)
                for par in tabla.cell(i, j).text_frame.paragraphs:
                    for run in par.runs:
                        run.font.size = PPt(8)
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
