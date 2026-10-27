"""Lector incremental independiente para ingesta masiva.

No altera las rutas existentes. Cada lote se libera tras consumirlo.
CSV y XLSX se leen sin convertir la poblacion completa a listas.
"""
from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path
from typing import Any

MAX_BATCH_ROWS = 10_000
DEFAULT_BATCH_ROWS = 2_000


def iterar_lotes(
    archivo: str | Path,
    *,
    lote: int = DEFAULT_BATCH_ROWS,
    hoja: str | None = None,
    encoding: str = "utf-8-sig",
) -> Iterator[list[tuple[Any, ...]]]:
    """Entrega lotes acotados. No retiene lotes previos."""
    if not 1 <= lote <= MAX_BATCH_ROWS:
        raise ValueError(f"lote debe estar entre 1 y {MAX_BATCH_ROWS}")
    ruta = Path(archivo)
    ext = ruta.suffix.lower()
    if ext == ".csv":
        with ruta.open("r", encoding=encoding, newline="") as fh:
            lector = csv.reader(fh)
            bloque: list[tuple[Any, ...]] = []
            for fila in lector:
                bloque.append(tuple(fila))
                if len(bloque) == lote:
                    yield bloque
                    bloque = []
            if bloque:
                yield bloque
        return
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            ws = wb[hoja] if hoja else wb.active
            bloque = []
            for fila in ws.iter_rows(values_only=True):
                bloque.append(tuple(fila))
                if len(bloque) == lote:
                    yield bloque
                    bloque = []
            if bloque:
                yield bloque
        finally:
            wb.close()
        return
    raise ValueError("Formato no soportado por lector incremental: use .csv o .xlsx")
