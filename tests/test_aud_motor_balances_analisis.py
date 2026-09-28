"""Análisis de estados financieros (NIA 315/520) y su papel de trabajo.

Importa solo los módulos livianos (motor_balances, analisis, papel_estados);
no levanta la app. Datos homologados sintéticos, recálculo a mano.
"""
import io

from openpyxl import load_workbook

from backend.app.aud.motor_balances import analisis as an
from backend.app.aud.motor_balances.papel_estados import generar_papel_estados
from backend.app.client_portal.flujo import motor_balances as mb


def _homologado():
    esf = {"periodos": ["2023", "2024"], "filas": [
        {"cuenta": "caja", "super_cias": "1010101", "es_hoja": True, "saldos": {"2023": 800.0, "2024": 1200.0}},
        {"cuenta": "ctas", "super_cias": "1010201", "es_hoja": True, "saldos": {"2023": 500.0, "2024": 400.0}},
        {"cuenta": "prov", "super_cias": "2010301", "es_hoja": True, "saldos": {"2023": 300.0, "2024": 650.0}},
        {"cuenta": "cap", "super_cias": "30101", "es_hoja": True, "saldos": {"2023": 1000.0, "2024": 950.0}},
    ]}
    eri = {"periodos": ["2023", "2024"], "filas": [
        {"cuenta": "vta", "super_cias": "40101", "es_hoja": True, "saldos": {"2023": 5000.0, "2024": 6000.0}},
        {"cuenta": "cv", "super_cias": "50101", "es_hoja": True, "saldos": {"2023": 3000.0, "2024": 3500.0}},
    ]}
    return mb.estados_superintendencia(esf, eri)


def test_horizontal_y_vertical():
    res = an.analizar(_homologado())
    por = {l["codigo"]: l for l in res["esf"]["lineas"]}
    activo = por["1"]
    # activo total 2023 = 800+500 = 1300 ; 2024 = 1200+400 = 1600
    assert activo["valores"] == [1300.0, 1600.0]
    assert activo["variacion"] == 300.0
    assert abs(activo["variacion_pct"] - 300 / 1300) < 1e-9
    # vertical del activo sobre sí mismo = 1.0
    assert abs(activo["vertical"][-1] - 1.0) < 1e-9


def test_ratios_conocidos():
    res = an.analizar(_homologado())
    ratios = {r["nombre"]: r["valores"] for r in res["ratios"]["filas"]}
    # 2024: AC=1600, PC=650 → 2.4615 ; Pasivo/Activo=650/1600=0.40625 ; margen bruto=(6000-3500)/6000
    assert abs(ratios["Liquidez corriente"][-1] - 1600 / 650) < 1e-6
    assert abs(ratios["Endeudamiento del activo"][-1] - 650 / 1600) < 1e-6
    assert abs(ratios["Margen bruto"][-1] - (6000 - 3500) / 6000) < 1e-6
    # los ratios de resultado neto/ROA/ROE NO se publican (subtotales derivados)
    assert "Margen neto" not in ratios


def test_expectativa_nia520_marca_lo_que_supera_umbral():
    res = an.analizar(_homologado(), umbral_pct=0.10)
    exp = res["expectativa_esf"]
    assert exp["aplicable"]
    por = {l["codigo"]: l for l in exp["lineas"]}
    # caja pasa de 800 a 1200 → +50% supera 10%
    assert por["1010101"]["supera_umbral"] is True


def test_papel_tiene_hojas_y_formulas():
    res = an.analizar(_homologado())
    wb = load_workbook(io.BytesIO(generar_papel_estados(res)))
    assert wb.sheetnames == ["ESF", "ERI", "Ratios", "NIA 520"]
    ratios = wb["Ratios"]
    formulas = [c.value for fila in ratios.iter_rows() for c in fila
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert formulas and all("'ES" in f or "'ER" in f for f in formulas), \
        "los ratios deben remitir por fórmula a las hojas ESF/ERI"


def test_un_solo_periodo_no_rompe_ni_expectativa():
    esf = {"periodos": ["2024"], "filas": [
        {"cuenta": "caja", "super_cias": "1010101", "es_hoja": True, "saldos": {"2024": 100.0}}]}
    res = an.analizar(mb.estados_superintendencia(esf, {"periodos": [], "filas": []}))
    assert res["expectativa_esf"]["aplicable"] is False
    generar_papel_estados(res)  # no debe lanzar


def test_html_estados_autonomo():
    from backend.app.aud.motor_balances.papel_estados import generar_html_estados
    res = an.analizar(_homologado())
    h = generar_html_estados(res)
    assert h.lstrip().startswith("<!doctype html>")
    assert "http://" not in h and "https://" not in h
    assert "data:application/vnd.openxmlformats" in h
    assert "window.print()" in h
    assert "NIA 520" in h
