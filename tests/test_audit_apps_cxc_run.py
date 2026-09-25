"""AUD-CXC-CARTERA corre de punta a punta por el executor vía adaptadores fn(ctx).

Verifica empíricamente que la Audit App ejecuta los 4 pasos contra el procesador
CXC REAL (``cxc_cartera.ejecutar``/``hojas`` y ``libro.xlsx``/``html``), no que
los salta. Usa el ``EJEMPLO`` del procesador como fuente de datos válida (el
mismo caso recalculado a mano en ``tests/test_proc_cxc_cartera.py``).

Correr sin la app HTTP:
    python -m pytest tests/test_audit_apps_cxc_run.py -q --noconftest
"""
import io
import json

from openpyxl import load_workbook

from backend.app.audit_apps.examples.aud_cxc_cartera import AUD_CXC_CARTERA
from backend.app.audit_apps.executor import AuditAppExecutor
from backend.app.audit_apps.manifest import AuditAppManifest
from backend.app.aud.niif.procesadores import cxc_cartera as m


def _solicitud() -> dict:
    """Solicitud CXC (JSON) con los datos del EJEMPLO del procesador."""
    ej = m.EJEMPLO
    return {
        "datasets": ej["datasets"],
        "parametros": ej["parametros"],
        "corte": ej["corte"],
        "engagement": {
            "client": "Comercial Demo S.A.", "ruc": "1790000000001",
            "framework": m.MARCO_COMPLETAS, "cutoff": ej["corte"],
            "firm": "AuditConsulting Auditores Cía. Ltda.",
            "preparer": "Analista Prueba", "reviewer": "Revisor Prueba", "year": "2025",
        },
    }


def _run(params=None):
    manifest = AuditAppManifest.from_dict(AUD_CXC_CARTERA).validate()
    inputs = {"cartera": json.dumps(_solicitud()).encode("utf-8")}
    p = {"marco": "full"}
    if params:
        p.update(params)
    return manifest, AuditAppExecutor().run(manifest, inputs=inputs, parameters=p, executed_by="a@x.ec")


def test_manifest_valida_y_apunta_a_los_adaptadores():
    m2 = AuditAppManifest.from_dict(AUD_CXC_CARTERA).validate()
    assert m2.id == "AUD-CXC-CARTERA" and m2.version == "1.0.0"
    refs = {s.id: s.engine_ref for s in m2.steps}
    assert refs["calcular_cxc"] == "backend.app.aud.niif.procesadores.cxc_app_adapters.calcular_cxc"
    assert all(r.startswith("backend.app.aud.niif.procesadores.cxc_app_adapters.") for r in refs.values())


def test_corre_los_4_pasos_sin_no_disponibles():
    _, res = _run()
    assert res.no_disponibles == ()            # nada se saltó: todo corrió
    assert [s.step_id for s in res.steps] == [
        "extraer_cartera", "calcular_cxc", "papel_trabajo_excel", "papel_trabajo_html"]
    assert res.input_hashes and res.output_hashes


def test_calcula_con_el_procesador_real():
    _, res = _run()
    # resultado_cxc es la salida real de cxc_cartera.ejecutar (mismos valores del EJEMPLO)
    resultado = next(s.value for s in res.steps if s.produces == "resultado_cxc")
    t = resultado["totals"]
    assert t["saldo"] == "67000.00"
    assert t["costoAmortizado"] == "64337.95" and t["interesNoDevengado"] == "2662.05"
    assert t["deterioroRequerido"] == "6143.38"
    assert t["ajuste"] == "4143.38" and resultado["primary"] == "ajuste"
    assert t["vencidoSinCobro"] == "18000.00" and t["difCircularizacion"] == "500.00"
    assert "hojas" in resultado and [h["name"] for h in resultado["hojas"]] == [n for n, _ in m.CEDULAS]


def test_produce_excel_y_html_reales():
    _, res = _run()
    excel = res.outputs["cxc_excel"]
    assert excel[:2] == b"PK"                    # xlsx = zip válido
    wb = load_workbook(io.BytesIO(excel))
    assert wb.sheetnames                          # abre sin reparaciones
    # las 12 cédulas del procesador viajan en el papel (nombres recortados a 31)
    nombres = set(wb.sheetnames)
    assert "00_Inicio" in nombres and "01_Resumen" in nombres and "10_Ajuste" in nombres
    html = res.outputs["cxc_html"].decode("utf-8")
    assert "<!doctype html" in html.lower() and "Cuentas por cobrar" in html
    # la salida "excepciones_cxc" serializa el resultado calculado
    exc = json.loads(res.outputs["excepciones_cxc"].decode("utf-8"))
    assert exc["totals"]["ajuste"] == "4143.38"
    assert any(e["code"] == "DETERIORO_INSUFICIENTE" for e in exc["exceptions"])


def test_parametro_marco_pymes_cambia_el_modelo_de_deterioro():
    # con marco='sme' el adaptador enruta a PYMES: el deterioro es pérdida incurrida
    _, res = _run({"marco": "sme"})
    resultado = next(s.value for s in res.steps if s.produces == "resultado_cxc")
    assert resultado["labels"]["deterioroRequerido"].startswith("Pérdida incurrida")
    assert resultado["totals"]["deterioroRequerido"] == "6143.38"
    assert any(e["code"] == "TASA_CORRIENTE_PYMES" for e in resultado["exceptions"])
