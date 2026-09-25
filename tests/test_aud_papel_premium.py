"""Papel ejecutivo premium: gráficos SVG (HTML/PDF) y nativos (Excel/PowerPoint)
en las 20 herramientas NIIF (18 de catálogo + pérdidas incurridas + PCE).

Comprueba que los gráficos existen, son autónomos (sin http ni scripts en el
PDF), no mezclan unidades (las tasas % quedan fuera del eje en USD) y que el
Excel sigue abriendo sin «reparar» con las fórmulas intactas.
"""
import io
import re

from openpyxl import load_workbook
from pptx import Presentation

from backend.app.aud.niif import ejercicio_modelo as em
from backend.app.aud.niif.procesadores import PROCESADORES, graficos, libro


def _reg(pid):
    mod = PROCESADORES[pid]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    return d, mod, em._reg(d, mod, ds, par, corte)


def test_html_trae_panorama_svg_en_las_20():
    for pid in PROCESADORES:
        d, _, reg = _reg(pid)
        h = libro.html(d, reg, [], 1, "APROBADO").decode("utf-8")
        assert h.count("<svg") >= 1, pid
        assert 'class="graficos"' in h and 'role="img"' in h, pid
        assert "http://" not in h and "https://" not in h, pid  # autónomo


def test_html_para_pdf_con_graficos_y_sin_scripts():
    for pid in PROCESADORES:
        d, _, reg = _reg(pid)
        h = libro.html(d, reg, [], 1, "APROBADO", para_pdf=True).decode("utf-8")
        assert "<svg" in h and "<script" not in h, pid


def test_resumen_no_mezcla_tasas_con_usd():
    for pid in PROCESADORES:
        mod = PROCESADORES[pid]
        ds, par, corte = em.escenario(mod)
        run = mod.ejecutar(ds, par, corte)
        items = graficos.datos_resumen(mod.hojas(run))
        assert items, pid
        assert not any("%" in e for e, _ in items), pid


def test_hallazgos_top_ordenados_y_otros():
    run = {"exceptions": [{"code": f"C{i}", "amount": 100 * (i + 1)} for i in range(10)]
           + [{"code": "C0", "amount": -50}, {"code": "SIN_IMPORTE", "amount": 0}]}
    items = graficos.datos_hallazgos(run)
    assert len(items) == graficos.TOP_HALLAZGOS + 1
    valores = [v for _, v in items[:-1]]
    assert valores == sorted(valores, reverse=True)
    assert items[-1][0].startswith("Otros (")
    assert not any("Sin importe" in e for e, _ in items)  # sin importe no se grafica
    c0 = [v for e, v in items if e.startswith("C0")] or [items[-1][1]]
    assert sum(v for _, v in items) == sum(100 * (i + 1) for i in range(10)) + 50


def test_barra_con_extremo_redondeado_y_linea_base():
    svg = graficos.svg_barras([("A", 100.0), ("B", -40.0)], "prueba")
    assert svg.count("<path") == 2 and "<line" in svg and "Q" in svg
    assert graficos.svg_barras([("A", 0.0)], "vacío") == ""


def test_cifra_heroe_con_separador_de_miles():
    assert graficos.cifra(-39820) == "−39.820,00"  # es-EC
    assert graficos.cifra({"v": 1234567.891}) == "1.234.567,89"


def _titulo(ch):
    return "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []))


def test_excel_un_solo_dashboard_con_graficos_por_formula():
    """El Excel tiene UNA sola sección de dashboard (00_Inicio): los gráficos viven
    ahí y sus datos son fórmulas a las cédulas; ninguna otra hoja lleva gráficos."""
    for pid in PROCESADORES:
        d, _, reg = _reg(pid)
        wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "APROBADO")))  # abre sin «reparar»
        con_grafico = [ws.title for ws in wb.worksheets if ws._charts]
        assert con_grafico == ["00_Inicio"], (pid, con_grafico)
        # Registrado vs recalculado, composición, distribución y hallazgos: pocas barras con
        # rótulo. Nunca «Cifras del resumen» (todas las filas del Resumen con escalas muy
        # distintas: se veía como un código de barras).
        titulos = [_titulo(ch) for ch in wb["00_Inicio"]._charts]
        assert titulos[0] == "Registrado vs recalculado (USD)", (pid, titulos)
        assert not any("Cifras del resumen" in t for t in titulos), (pid, titulos)
        assert 3 <= len(titulos) <= 4, (pid, titulos)
        # Los datos de los gráficos viven en una hoja oculta (el panel queda limpio) y son fórmulas.
        datos = wb[libro.HOJA_DATOS_GRAFICOS]
        assert datos.sheet_state == "hidden", pid
        formulas = [c.value for row in datos.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        # Títulos que no se montan sobre las barras y rótulos de categoría como texto.
        for ch in wb["00_Inicio"]._charts:
            assert ch.title.overlay is False and ch.series[0].cat.strRef is not None, pid
        assert any(re.match(r"^='00_Inicio'!\$[B-F]\$\d+$", f) for f in formulas), pid  # registrado/recalculado → tarjetas
        assert sum("SUMIFS(" in f and "Problemas" not in f for f in formulas) >= 2, pid  # composición/distribución → cédulas
        assert any("SUMIFS(" in f and "Problemas" in f for f in formulas), pid  # hallazgos → hoja de problemas


def test_excel_tablas_premium():
    """Las demás hojas son tablas premium: cabecera alta con filete dorado, filas
    alternas y filetes horizontales (sin cuadrícula completa)."""
    d, _, reg = _reg("efectivo_equivalentes")
    wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "APROBADO")))
    ws = next(w for w in wb.worksheets if "Resumen" in w.title)
    assert ws.row_dimensions[4].height == 30
    assert ws.cell(row=4, column=1).border.bottom.style == "medium"
    assert ws.cell(row=6, column=1).fill.fgColor.rgb.endswith("F7F9FC")   # fila alterna
    izq = ws.cell(row=5, column=1).border.left
    assert izq is None or izq.style is None                              # sin cuadrícula vertical
    assert ws.cell(row=5, column=1).border.bottom.style == "thin"


def test_pptx_diapositivas_con_grafico_nativo():
    for pid in PROCESADORES:
        d, _, reg = _reg(pid)
        prs = Presentation(io.BytesIO(libro.pptx(d, reg, [], 1, "APROBADO")))
        n = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        assert n >= 1, pid
