"""E7 por HTTP: requerimiento, evidencia, cobertura, mapeo y validación.

La población es un Excel escrito con openpyxl —otro programa, no el sitio—
con las 3 partidas del ejemplo VNR del sitio: costo total 470,00.
"""
import base64
import json
from pathlib import Path

import pytest

from tests.test_aud_ciclo_http import BASE, _accion, _prueba_vnr, _staff_con_proyecto
from tests.test_aud_niif_fichas import _h

ESPEJO = json.loads(
    (Path(__file__).resolve().parents[1] / "backend/app/aud/niif/ciclo/espejo_datos.json").read_text(encoding="utf-8")
)
def _poblacion_xlsx() -> bytes:
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(["Código", "Descripción", "Cantidad", "Costo unitario", "Precio de venta", "Costo terminación", "Costo de venta", "Deterioro registrado"])
    for fila in ([ "0001", "Producto A", 10, 20, 18, 1, 2, 0], ["0002", "Producto B", 5, 30, 40, 1, 2, 0], ["0003", "Producto C", 8, 15, 13, 1, 1, 0]):
        ws.append(fila)
    ws.append([])
    ws.append([None, None, None, None, None, None, None, None])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


XLSX_VNR = _poblacion_xlsx()
MAPA_VNR = {k: i for i, k in enumerate(
    ["id", "description", "quantity", "unit_cost", "selling_price", "completion_cost", "selling_cost", "recorded_allowance"]
)}


@pytest.fixture(autouse=True)
def disco_temporal(tmp_path, monkeypatch):
    # La evidencia de las pruebas nunca toca el disco real de la máquina.
    monkeypatch.setenv("AUD_CICLO_DIR", str(tmp_path / "aud_pruebas"))
    monkeypatch.setenv("AUD_CICLO_MIN_LIBRE_MB", "0")


def _leer(client, tok, p):
    return client.get(f"{BASE}/pruebas/{p['id']}", headers=_h(tok)).json()


def _subir(client, tok, p, req, nombre, contenido, componente=""):
    return client.post(
        f"{BASE}/pruebas/{p['id']}/archivos", headers=_h(tok),
        data={"revision": str(p["revision"]), "requerimiento": req, "componente": componente},
        files={"archivo": (nombre, contenido)},
    )


def _hasta_requerimiento_aprobado(client, tok):
    _, pid = None, None
    tok_pid = _staff_con_proyecto(client)
    tok, pid = tok_pid
    p = _prueba_vnr(client, tok, pid)
    p = _accion(client, tok, p, "research").json()
    p = _accion(client, tok, p, "generate_program").json()
    fuentes = p["registro"]["sources"]
    fuentes[0].update(verified=True, section="par. 9", date="vigente", procedures=["VNR-01", "VNR-02"])
    fuentes[1].update(verified=True, document="NIA 540", section="par. 13", date="vigente", procedures=["VNR-03"])
    p = _accion(client, tok, p, "approve_program", {"program": p["registro"]["program"], "sources": fuentes}).json()
    p = _accion(client, tok, p, "generate_request").json()
    assert p["estado"] == "REQUERIMIENTO_GENERADO"
    reqs = p["registro"]["requests"]
    assert [r["id"] for r in reqs] == ["RQ-001", "RQ-002", "RQ-003", "RQ-VNR-04"]
    # El auditor declara el inventario en dos componentes (dos bodegas).
    reqs[0]["components"] = ["Quito", "Guayaquil"]
    p = _accion(client, tok, p, "approve_request", {"requests": reqs}).json()
    assert p["estado"] == "REQUERIMIENTO_APROBADO", p
    return tok, p


def test_requerimiento_incompleto_se_rechaza(client):
    tok, pid = _staff_con_proyecto(client)
    p = _prueba_vnr(client, tok, pid)
    p = _accion(client, tok, p, "research").json()
    p = _accion(client, tok, p, "generate_program").json()
    fuentes = p["registro"]["sources"]
    fuentes[0].update(verified=True, section="par. 9", date="vigente", procedures=["VNR-01", "VNR-02", "VNR-03"])
    fuentes[1].update(verified=True, document="NIA 540", section="par. 13", date="vigente", procedures=[])
    p = _accion(client, tok, p, "approve_program", {"program": p["registro"]["program"], "sources": fuentes}).json()
    p = _accion(client, tok, p, "generate_request").json()
    malo = [{**p["registro"]["requests"][0], "procedure": "INVENTADO"}]
    r = _accion(client, tok, p, "approve_request", {"requests": malo})
    assert r.status_code == 400 and "procedimiento aprobado" in r.json()["detail"]


def test_evidencia_cobertura_y_validacion_de_punta_a_punta(client):
    tok, p = _hasta_requerimiento_aprobado(client, tok=None)

    # Formato no declarado: el inventario admite XLSX y CSV, no PDF.
    r = _subir(client, tok, p, "RQ-001", "inventario.pdf", b"%PDF", "Quito")
    assert r.status_code == 400 and "XLSX, CSV" in r.json()["detail"]
    # Componente obligatorio cuando el requerimiento los declara.
    r = _subir(client, tok, p, "RQ-001", "inventario.xlsx", XLSX_VNR)
    assert r.status_code == 400 and "por componentes" in r.json()["detail"]

    r = _subir(client, tok, p, "RQ-001", "inventario_quito.xlsx", XLSX_VNR, "Quito")
    assert r.status_code == 201, r.text
    p = _leer(client, tok, p)
    assert p["estado"] == "DOCUMENTACION_RECIBIDA"
    assert any("faltan Guayaquil" in h for h in p["huecos"])
    id_quito = p["archivos"][0]["id"]

    for req, nombre, contenido, comp in [
        ("RQ-001", "inventario_gye.xlsx", XLSX_VNR, "Guayaquil"),
        ("RQ-002", "precios.pdf", b"%PDF-precios", ""),
        ("RQ-003", "gastos.pdf", b"%PDF-gastos", ""),
        ("RQ-VNR-04", "politica.pdf", b"%PDF-politica", ""),
    ]:
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, nombre, contenido, comp).status_code == 201
    p = _leer(client, tok, p)
    assert p["huecos"] == []
    assert len(p["archivos"]) == 5

    # El original se descarga intacto.
    r = client.get(f"{BASE}/pruebas/{p['id']}/archivos/{id_quito}", headers=_h(tok))
    assert r.status_code == 200 and r.content == XLSX_VNR

    # Mapeo: hoja Inventario, encabezado en la fila 1.
    p = _accion(client, tok, p, "map_validate",
                {"fileId": id_quito, "sheet": "Inventario", "header": 1, "mapping": MAPA_VNR}).json()
    assert p["registro"]["validation"]["ok"] is True
    assert p["registro"]["validation"]["records"] == 3
    assert p["registro"]["controlTotal"] == "470.00"

    # Validar sin documentar la revisión de evidencia.
    r = _accion(client, tok, p, "validate", {"ledger": "470.00", "tolerance": "0"})
    assert r.status_code == 400 and "revisión de evidencia" in r.json()["detail"]

    # Rechazar una bodega vuelve a abrir el hueco.
    p = _accion(client, tok, p, "reject_file", {"fileId": id_quito}).json()
    revision = {"evidenceReviewed": True, "evidenceReview": "Cotejé los originales con la extracción.", "ledger": "470.00", "tolerance": "0"}
    r = _accion(client, tok, p, "validate", revision)
    assert r.status_code == 400 and "Cobertura incompleta" in r.json()["detail"] and "Quito" in r.json()["detail"]
    p = _accion(client, tok, _leer(client, tok, p), "reject_file", {"fileId": id_quito}).json()

    # Saldo contable que no concilia y sin aceptación: no se valida.
    r = _accion(client, tok, p, "validate", {**revision, "ledger": "480.00"})
    assert r.status_code == 400 and "conciliación" in r.json()["detail"]

    p = _accion(client, tok, p, "validate", revision).json()
    assert p["estado"] == "DOCUMENTACION_VALIDADA"
    assert p["registro"]["reconciliation"]["within"] is True
    assert {r["status"] for r in p["registro"]["requests"]} == {"RECIBIDO"}

    detalle = _leer(client, tok, p)
    acciones = [e["accion"] for e in detalle["eventos"]]
    assert acciones.count("upload") == 5 and acciones[-1] == "validate"


def test_evidencia_nueva_invalida_el_mapeo(client):
    tok, p = _hasta_requerimiento_aprobado(client, tok=None)
    p = _leer(client, tok, p)
    _subir(client, tok, p, "RQ-001", "q.xlsx", XLSX_VNR, "Quito")
    p = _leer(client, tok, p)
    fid = p["archivos"][0]["id"]
    p = _accion(client, tok, p, "map_validate", {"fileId": fid, "sheet": "Inventario", "header": 1, "mapping": MAPA_VNR}).json()
    assert p["registro"]["rows"]
    _subir(client, tok, p, "RQ-002", "precios.pdf", b"%PDF")
    p = _leer(client, tok, p)
    assert p["registro"]["rows"] == [] and p["registro"]["validation"] is None
    assert "nueva evidencia" in p["registro"]["invalidationReason"]


def test_disco_casi_lleno_rechaza_la_subida(client, monkeypatch):
    tok, p = _hasta_requerimiento_aprobado(client, tok=None)
    monkeypatch.setenv("AUD_CICLO_MIN_LIBRE_MB", str(10 ** 9))
    r = _subir(client, tok, p, "RQ-002", "precios.pdf", b"%PDF")
    assert r.status_code == 400 and "casi lleno" in r.json()["detail"]
    assert _leer(client, tok, p)["archivos"] == []


def test_no_se_sube_antes_de_aprobar_el_requerimiento(client):
    tok, pid = _staff_con_proyecto(client)
    p = _prueba_vnr(client, tok, pid)
    r = _subir(client, tok, p, "RQ-001", "a.xlsx", XLSX_VNR, "Quito")
    assert r.status_code == 400 and "no está habilitada" in r.json()["detail"]


def _inventario(filas):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(["Código", "Descripción", "Cantidad", "Costo unitario", "Precio de venta", "Costo terminación", "Costo de venta", "Deterioro registrado"])
    for f in filas:
        ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_varios_archivos_se_unen_en_una_poblacion(client):
    # E10: Quito (Excel) y Guayaquil (CSV), mismo formato, una sola población del país.
    tok, p = _hasta_requerimiento_aprobado(client, tok=None)
    quito = _inventario([["Q-1", "A", 10, 20, 18, 1, 2, 0], ["Q-2", "B", 5, 30, 40, 1, 2, 0]])
    gye = ("Código,Descripción,Cantidad,Costo unitario,Precio de venta,Costo terminación,Costo de venta,Deterioro registrado\n"
           "G-1,C,8,15,13,1,1,0\n").encode("utf-8")
    assert _subir(client, tok, _leer(client, tok, p), "RQ-001", "quito.xlsx", quito, "Quito").status_code == 201
    assert _subir(client, tok, _leer(client, tok, p), "RQ-001", "gye.csv", gye, "Guayaquil").status_code == 201
    p = _leer(client, tok, p)
    q, g = p["archivos"]
    partes = [{"fileId": q["id"], "sheet": "Inventario", "header": 1, "mapping": MAPA_VNR},
              {"fileId": g["id"], "sheet": "CSV", "header": 1, "mapping": MAPA_VNR}]
    hojas = client.get(f"{BASE}/pruebas/{p['id']}/archivos/{g['id']}", headers=_h(tok))
    assert hojas.status_code == 200
    r = _accion(client, tok, p, "map_validate", {"files": partes})
    assert r.status_code == 200, r.text
    reg = r.json()["registro"]
    assert reg["validation"]["ok"] and reg["validation"]["records"] == 3
    # 10×20 + 5×30 + 8×15 = 470,00: la suma de las dos bodegas.
    assert reg["controlTotal"] == "470.00"
    assert [(f["id"], f["_component"], f["_file"]) for f in reg["rows"]] == [
        ("Q-1", "Quito", "quito.xlsx"), ("Q-2", "Quito", "quito.xlsx"), ("G-1", "Guayaquil", "gye.csv")]
    assert [(m["file"], m["records"]) for m in reg["mappings"]] == [("quito.xlsx", 2), ("gye.csv", 1)]

    # Un archivo rechazado no entra en la población.
    p = _accion(client, tok, r.json(), "reject_file", {"fileId": g["id"]}).json()
    r = _accion(client, tok, p, "map_validate", {"files": partes})
    assert r.status_code == 400 and "rechazado" in r.json()["detail"]


def test_el_modelo_se_llena_y_se_procesa_sin_ajustes(client):
    # E10: el modelo trae las columnas de la ficha; lleno y subido, se mapea por nombre.
    import io
    import openpyxl
    tok, p = _hasta_requerimiento_aprobado(client, tok=None)
    p = _leer(client, tok, p)
    assert p["modelos"] == ["RQ-001"]
    r = client.get(f"{BASE}/pruebas/{p['id']}/modelo/RQ-002", headers=_h(tok))
    assert r.status_code == 400 and "soporte" in r.json()["detail"]
    r = client.get(f"{BASE}/pruebas/{p['id']}/modelo/RQ-001", headers=_h(tok))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Datos", "Instrucciones"]
    etiquetas = [f["label"] for f in p["definicion"]["fields"]]
    assert [c.value for c in wb["Datos"][1]] == etiquetas
    assert any("Quito" in str(c.value) for fila in wb["Instrucciones"].iter_rows() for c in fila), "declara las partes"
    # El cliente llena el modelo (Quito) y el auditor lo sube.
    ws = wb["Datos"]
    assert ws.max_row == 1, "el modelo no trae filas vacías que empujen los datos hacia abajo"
    ws.append(["0001", "Producto A", 10, 20, 18, 1, 2, 0])
    ws.append(["0002", "Producto B", 5, 30, 40, 1, 2, 0])
    b = io.BytesIO()
    wb.save(b)
    assert _subir(client, tok, p, "RQ-001", "quito.xlsx", b.getvalue(), "Quito").status_code == 201
    p = _leer(client, tok, p)
    mapa = {f["key"]: i for i, f in enumerate(p["definicion"]["fields"])}
    p = _accion(client, tok, p, "map_validate", {"files": [{"fileId": p["archivos"][0]["id"], "sheet": "Datos", "header": 1, "mapping": mapa}]}).json()
    assert p["registro"]["validation"]["ok"], p["registro"]["validation"]
    assert p["registro"]["controlTotal"] == "350.00"
    assert [f["_row"] for f in p["registro"]["rows"]] == [2, 3]
