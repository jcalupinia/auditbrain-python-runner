"""Registros del encargo con un clic (decisión del dueño, 2026-09-26): independencia, asistencia a la discusión,
aceptación del socio, carta de encargo firmada y comunicación al gobierno. La planificación los recibe en
``parametros["_encargo"]`` y los muestra en la hoja 00_Registros."""
import datetime

from backend.app.aud.niif.ciclo import servicio
from backend.app.aud.niif.procesadores import planificacion_encargo as enc
from backend.app.aud.niif.procesadores import planificacion_nia as m
from backend.app.context import service as ctx
from backend.app.context.models import Project
from backend.app.db.session import SessionLocal
from tests.test_aud_ciclo_http import BASE, FICHA, _staff_con_proyecto
from tests.test_aud_niif_fichas import _h


def _reg(client, tok, pid, **datos):
    return client.post(f"{BASE}/proyectos/{pid}/registros", headers=_h(tok), json=datos)


def test_registros_reglas_y_anulacion(client):
    tok, pid = _staff_con_proyecto(client)
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    assert _reg(client, tok, pid, tipo="otro").status_code == 400
    assert "rol" in _reg(client, tok, pid, tipo="independencia", rol="Jefe").json()["detail"]
    # La aceptación la hace el socio, después de confirmar su independencia como socio.
    r = _reg(client, tok, pid, tipo="aceptacion")
    assert r.status_code == 400 and "socio" in r.json()["detail"]
    r = _reg(client, tok, pid, tipo="independencia", rol="Socio", nombre="CPA Socia", amenazas="", anio_desde="2020")
    assert r.status_code == 201, r.text
    assert _reg(client, tok, pid, tipo="aceptacion").status_code == 201
    manana = (datetime.date.today() + datetime.timedelta(days=1)).isoformat()
    assert "posterior a hoy" in _reg(client, tok, pid, tipo="carta", fecha=manana).json()["detail"]
    assert _reg(client, tok, pid, tipo="carta", fecha="2025-09-05", limitaciones="Sin acceso a sucursal").status_code == 201
    assert "medio" in _reg(client, tok, pid, tipo="comunicacion").json()["detail"]
    assert _reg(client, tok, pid, tipo="comunicacion", medio="Reunión con el directorio").status_code == 201
    assert _reg(client, tok, pid, tipo="asistencia", rol="Socio", nombre="CPA Socia", fecha="2025-10-10").status_code == 201
    # Un segundo clic reemplaza al anterior (queda anulado, no se borra).
    r = _reg(client, tok, pid, tipo="independencia", rol="Socio", nombre="CPA Socia", amenazas="Préstamo del cliente",
             salvaguardas="Préstamo cancelado")
    assert r.status_code == 201
    datos = client.get(f"{BASE}/proyectos/{pid}/registros", headers=_h(tok)).json()
    e = datos["encargo"]
    assert e["firma"] == "Audit Consulting"
    assert len(e["registros"]["equipo"]) == 1 and e["registros"]["equipo"][0]["amenazas"] == "Préstamo del cliente"
    assert e["registros"]["equipo"][0]["anios"] == 1                    # el reemplazo ya no trae el año declarado
    assert e["registros"]["carta"] == {"actor": "CPA Socia", "fecha": "2025-09-05", "detalle": "Sin acceso a sucursal"}
    assert e["registros"]["aceptacion"]["actor"] == "CPA Socia"
    rid = next(x["id"] for x in datos["registros"] if x["tipo"] == "comunicacion")
    assert client.delete(f"{BASE}/proyectos/{pid}/registros/{rid}", headers=_h(tok)).json() == {"ok": True}
    assert "comunicacion" not in client.get(f"{BASE}/proyectos/{pid}/registros", headers=_h(tok)).json()["encargo"]["registros"]
    assert client.delete(f"{BASE}/proyectos/{pid}/registros/{rid}", headers=_h(tok)).status_code == 400
    # Otra persona no anula lo que registró otro.
    db = SessionLocal()
    try:
        from backend.app.aud.niif.ciclo.models import RegistroEncargo
        otro = db.query(RegistroEncargo).filter(RegistroEncargo.project_id == pid, RegistroEncargo.anulado_en.is_(None)).first()
        assert otro is not None
        try:
            servicio.anular_registro(db, pid, otro.id, "otra@firma.ec", False)
            assert False, "debió negarse"
        except servicio.ReglaIncumplida as ex:
            assert "Solo quien hizo el registro" in str(ex)
    finally:
        db.close()


def test_anios_con_el_cliente_cuentan_los_encargos_anteriores(client):
    tok, pid = _staff_con_proyecto(client)
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    assert _reg(client, tok, pid, tipo="independencia", rol="Gerente", nombre="CPA Gerente").status_code == 201
    db = SessionLocal()
    try:
        pr = db.get(Project, pid)
        nuevo = ctx.create_project(db, client_id=pr.client_id, name="Auditoría 2026", module_code="AUD",
                                   organization_id=pr.organization_id)
        pid2 = nuevo.id
    finally:
        db.close()
    assert client.put(f"{BASE}/proyectos/{pid2}/ficha", headers=_h(tok), json={**FICHA, "year": 2026, "cutoff": "2026-12-31"}).status_code == 200
    assert _reg(client, tok, pid2, tipo="independencia", rol="Gerente", nombre="CPA Gerente").status_code == 201
    eq = client.get(f"{BASE}/proyectos/{pid2}/registros", headers=_h(tok)).json()["encargo"]["registros"]["equipo"]
    assert eq[0]["anios"] == 2
    # El año declarado manda si da más años.
    assert _reg(client, tok, pid2, tipo="independencia", rol="Gerente", nombre="CPA Gerente", anio_desde="2019").status_code == 201
    eq = client.get(f"{BASE}/proyectos/{pid2}/registros", headers=_h(tok)).json()["encargo"]["registros"]["equipo"]
    assert eq[0]["anios"] == 8


def test_la_planificacion_usa_los_registros_de_la_plataforma(client):
    tok, pid = _staff_con_proyecto(client)
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    _reg(client, tok, pid, tipo="independencia", rol="Socio", nombre="CPA Socia")
    _reg(client, tok, pid, tipo="aceptacion")
    _reg(client, tok, pid, tipo="asistencia", rol="Socio", nombre="CPA Socia", fecha="2025-10-10")
    assert m.USA_REGISTROS_ENCARGO is True
    db = SessionLocal()
    try:
        encargo = servicio.registros_encargo(db, pid)
    finally:
        db.close()
    e = m.EJEMPLO
    r = m.ejecutar(e["datasets"], {**{k: v for k, v in e["parametros"].items() if k not in ("socio", "gerente")}, "_encargo": encargo},
                   e["corte"])
    hojas = {h["name"]: h for h in m.hojas(r)}
    filas = hojas[enc.REG]["rows"]
    assert [f[0] for f in filas][:2] == [enc.TIPO_INDEP, enc.TIPO_ASIST]
    assert r["detalle"]["parametros"]["socio"] == "CPA Socia"
    ev = {x["codigo"]: x["estado"] for x in r["detalle"]["evals"]}
    assert ev["ACE-01"] == "Conforme" and ev["CON-02"] == "Pendiente" and ev["COM-01"] == "Pendiente"


def test_documentos_generados_por_la_plataforma(client):
    import io

    from docx import Document

    from backend.app.aud.niif.ciclo import documentos_encargo as docs

    tok, pid = _staff_con_proyecto(client)
    url = f"{BASE}/proyectos/{pid}/documentos"
    assert "ficha" in client.get(f"{url}/carta_encargo", headers=_h(tok)).json()["detail"]
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    assert client.get(f"{url}/otro", headers=_h(tok)).status_code == 400
    _reg(client, tok, pid, tipo="independencia", rol="Socio", nombre="CPA Socia")
    _reg(client, tok, pid, tipo="carta", limitaciones="Sin acceso a la sucursal de Guayaquil")
    r = client.get(f"{url}/carta_encargo", headers=_h(tok))
    assert r.status_code == 200 and r.content.startswith(b"PK")
    texto = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Empresa Ejemplo S.A." in texto and "NIIF completas" in texto and "Sin acceso a la sucursal" in texto
    assert "AuditConsulting Auditores Cía. Ltda." in "\n".join(c.text for t in Document(io.BytesIO(r.content)).tables
                                                               for c in t._cells)
    # Sin planificación ejecutada, el acta y la carta de planificación marcan lo que falta (no inventan).
    for tipo in ("acta_discusion", "carta_planificacion"):
        r = client.get(f"{url}/{tipo}", headers=_h(tok))
        assert r.status_code == 200 and "[PENDIENTE]" in "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    # Con la planificación del ejemplo: riesgos presentes, indicios de fraude y asuntos de la hoja 32.
    e = m.EJEMPLO
    run = {"hojas": m.hojas(m.ejecutar(e["datasets"], e["parametros"], e["corte"]))}
    encargo = e["parametros"]["_encargo"]
    acta = Document(io.BytesIO(docs.acta_discusion(FICHA, encargo, run)))
    celdas = [c.text for t in acta.tables for c in t._cells]
    assert "Ana Torres (ficticio)" in celdas and "RB-01" in celdas and "FRA-03" in celdas and "RB-03" not in celdas
    carta = Document(io.BytesIO(docs.carta_planificacion(FICHA, encargo, run)))
    celdas = [c.text for t in carta.tables for c in t._cells]
    assert "Responsabilidades del auditor" in celdas and "Riesgo significativo · RB-01" in celdas
    assert not any(c.startswith("Envío de la carta") for c in celdas)


def test_los_documentos_llevan_el_marco_del_cliente():
    import io

    from docx import Document

    from backend.app.aud.niif.ciclo import documentos_encargo as docs

    def texto(b):
        return "\n".join(p.text for p in Document(io.BytesIO(b)).paragraphs)

    completas = texto(docs.carta_encargo(FICHA, {}))
    assert "de conformidad con NIIF completas;" in completas and "PYMES" not in completas
    pymes = {**FICHA, "framework": "NIIF para las PYMES", "edition": "2015"}
    assert "NIIF para las PYMES (edición 2015)" in texto(docs.carta_encargo(pymes, {}))
    assert "NIIF para las PYMES (edición 2025)" in texto(docs.carta_planificacion({**pymes, "edition": "2025"}, {}, None))


def test_indagaciones_consultas_y_bloqueo_de_la_aprobacion(client):
    """M3 y M7: indagaciones con un clic (varias vigentes) y consultas que bloquean la aprobación hasta resolverse."""
    from backend.app.aud.niif.ciclo.models import Prueba
    from backend.app.aud.niif.ciclo.reglas import ReglaIncumplida

    tok, pid = _staff_con_proyecto(client)
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    assert "tema" in _reg(client, tok, pid, tipo="indagacion", tema="Fraude", resumen="Sin comentarios relevantes").json()["detail"]
    for tema in ("Partes relacionadas", "Leyes y reglamentos"):
        r = _reg(client, tok, pid, tipo="indagacion", tema=tema, procedimiento="Indagación", persona="Gerente general",
                 resumen="La dirección confirma sus partes y su cumplimiento.")
        assert r.status_code == 201, r.text
    c = _reg(client, tok, pid, tipo="consulta", tema="Provisión del litigio laboral", detalle="Consulta al área técnica")
    assert c.status_code == 201
    e = client.get(f"{BASE}/proyectos/{pid}/registros", headers=_h(tok)).json()["encargo"]
    assert len(e["registros"]["indagaciones"]) == 2 and e["registros"]["consultas"][0]["estado"] == "Abierta"
    assert e["ficha"]["framework"] == FICHA["framework"]
    db = SessionLocal()
    try:
        assert servicio.consultas_abiertas(db, pid) == 1
        p = Prueba(project_id=pid, version=1, estado="EN_REVISION", origen="proc:planificacion_nia",
                   definicion=m.definicion(), registro={}, revision=1, creada_por="x")
        db.add(p)
        db.commit()
        try:
            servicio.aplicar_accion(db, p, "approve", 1, {"conclusion": "Planificación aprobada.", "conclusionReviewed": True}, "yo")
            assert False, "debió bloquear"
        except ReglaIncumplida as ex:
            assert str(ex) == servicio.CONSULTAS_BLOQUEAN
    finally:
        db.close()
    url = f"{BASE}/proyectos/{pid}/registros/{c.json()['id']}/resolver"
    assert "resolución" in client.post(url, headers=_h(tok), json={"resolucion": "ok"}).json()["detail"]
    r = client.post(url, headers=_h(tok), json={"resolucion": "Se provisiona según el criterio del abogado."})
    assert r.status_code == 200 and r.json()["datos"]["estado"] == "Resuelta"
    assert client.post(url, headers=_h(tok), json={"resolucion": "Otra vez resuelta."}).status_code == 400
    db = SessionLocal()
    try:
        assert servicio.consultas_abiertas(db, pid) == 0
    finally:
        db.close()


def test_version_anterior_de_la_planificacion():
    """M4: la versión nueva recibe la materialidad y los riesgos de la anterior (hoja 13 guardada en su ejecución)."""
    from backend.app.aud.niif.ciclo.models import Prueba

    e = m.EJEMPLO
    r = m.ejecutar(e["datasets"], e["parametros"], e["corte"])
    run = {"totals": r["totals"], "hojas": [h for h in m.hojas(r) if h["name"].startswith("13_")]}
    db = SessionLocal()
    try:
        from backend.app.context.models import Project
        pid = db.query(Project.id).first()[0]
        old = Prueba(project_id=pid, version=1, estado="APROBADO", origen="x", definicion={}, registro={"run": run}, revision=1)
        db.add(old)
        db.commit()
        nueva = Prueba(project_id=pid, parent_id=old.id, version=2, estado="PRUEBA_SELECCIONADA", origen="x", definicion={},
                       registro={}, revision=1)
        ant = servicio.version_anterior_run(db, nueva)
    finally:
        db.close()
    assert ant["version"] == 1 and ant["totales"]["materialidad"] == r["totals"]["materialidad"]
    assert ant["riesgos"][0]["cond"] == r["detalle"]["riesgos"][0]["cond"]
    # Con la misma planificación, el comparativo no muestra cambios en la materialidad ni en los riesgos que ya estaban.
    r2 = m.ejecutar(e["datasets"], {**e["parametros"], "_anterior": ant}, e["corte"])
    cam = next(h for h in m.hojas(r2) if h["name"] == "42_Cambios")["rows"]
    estados = [f[3]["v"] if isinstance(f[3], dict) else f[3] for f in cam]
    assert set(estados) == {"Sin cambio"}


def test_enfoque_del_ciclo_lo_confirma_el_socio(client):
    tok, pid = _staff_con_proyecto(client)
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    datos = {"tipo": "enfoque", "ciclo": "Inventarios y costo de ventas", "decision": "Sustantivo",
             "motivo": "Deficiencias en la toma física: no confiamos."}
    assert "socio" in _reg(client, tok, pid, **datos).json()["detail"]
    _reg(client, tok, pid, tipo="independencia", rol="Socio", nombre="CPA Socia")
    assert "Ciclo" in _reg(client, tok, pid, **{**datos, "ciclo": "Otro"}).json()["detail"]
    assert _reg(client, tok, pid, **{**datos, "decision": "Tal vez"}).status_code == 400
    assert _reg(client, tok, pid, **datos).status_code == 201
    assert _reg(client, tok, pid, **{**datos, "decision": "Confiar en controles", "motivo": "Se corrigió la toma física."}).status_code == 201
    enf = client.get(f"{BASE}/proyectos/{pid}/registros", headers=_h(tok)).json()["encargo"]["registros"]["enfoque"]
    assert len(enf) == 1 and enf[0]["decision"] == "Confiar en controles" and enf[0]["actor"] == "CPA Socia"
    r = client.get(f"{BASE}/proyectos/{pid}/documentos/conocimiento_negocio", headers=_h(tok))
    assert r.status_code == 200 and r.content.startswith(b"PK")


def test_memorando_de_conocimiento_del_negocio():
    import io

    from docx import Document

    from backend.app.aud.niif.ciclo import documentos_encargo as docs

    e = m.EJEMPLO
    run = {"hojas": m.hojas(m.ejecutar(e["datasets"], e["parametros"], e["corte"]))}
    d = Document(io.BytesIO(docs.conocimiento_negocio(FICHA, run)))
    textos = [p.text for p in d.paragraphs]
    celdas = [c.text for t in d.tables for c in t._cells]
    assert "Cifras clave" in textos and "Enfoque por ciclo: confianza o no en los controles" in textos
    assert "Entidad auditada" in celdas and "Nómina y beneficios a empleados" in celdas and "R01" in celdas


def test_archivos_de_entrada_con_su_huella_para_el_audit_trail(client):
    """Prioridad baja (NIA 230): la plataforma entrega nombre y SHA-256 de los archivos del cliente vigentes (no los
    rechazados ni los papeles de trabajo) para la hoja 23."""
    from backend.app.aud.niif.ciclo.models import Prueba, PruebaArchivo

    _tok, pid = _staff_con_proyecto(client)
    db = SessionLocal()
    try:
        p = Prueba(project_id=pid, version=1, estado="DOCUMENTACION_RECIBIDA", origen="x", definicion={}, registro={}, revision=1)
        db.add(p)
        db.commit()
        base = {"prueba_id": p.id, "componente": None, "tipo": "text/csv", "tamano": 10, "ruta": "x"}
        db.add_all([PruebaArchivo(**base, requerimiento="RQ-001", nombre="Balance 2024.csv", sha256="a" * 64, subido_por="cpa@x.ec"),
                    PruebaArchivo(**base, requerimiento="RQ-002", nombre="Balance viejo.csv", sha256="b" * 64, estado="rechazado"),
                    PruebaArchivo(**base, requerimiento="RQ-002", nombre="Papel.html", sha256="c" * 64, clase="workpaper")])
        db.commit()
        arch = servicio.archivos_de_entrada(db, p.id)
    finally:
        db.close()
    assert [(x["requerimiento"], x["nombre"], x["sha256"]) for x in arch] == [("RQ-001", "Balance 2024.csv", "a" * 64)]
    assert arch[0]["subido_por"] == "cpa@x.ec" and arch[0]["subido_en"]
