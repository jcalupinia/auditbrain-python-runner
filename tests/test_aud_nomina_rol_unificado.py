"""Normalización de roles en formato «unificado» (sin hoja «normal»).

Regresión del error **«Worksheet normal does not exist»**: un rol cuya columna de
empleado se llama «Nombres» (plural) u otra variante debe reconocerse y reescribirse
a la hoja «normal» del motor, en vez de reventar. Y si de verdad NO hay columnas de
empleado/sueldo, el motor debe AVISAR con un mensaje claro (no tragar el error con un
`except: pass` y dejar que el motor reviente después con el KeyError críptico).

Sin cifras ni nombres de cliente: encabezados genéricos y datos ficticios.
"""
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl import Workbook

from backend.app.aud.niif.procesadores.nomina import mayores
from backend.app.aud.niif.procesadores.nomina import perfil_cuentas as pc


def _rol(tmp_path, headers, filas, hoja="Roles_Unificados"):
    p = tmp_path / "rol.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(headers)
    for f in filas:
        ws.append(f)
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("col_empleado", ["Nombres", "NOMBRE", "Empleado", "Apellidos y Nombres"])
def test_rol_unificado_sin_normal_se_reconoce_y_reescribe(tmp_path, col_empleado):
    # Un rol en otra estructura (sin hoja «normal») con la columna de empleado nombrada
    # de varias formas debe reescribirse a la hoja «normal» en vez de reventar.
    headers = ["MES", "Cédula", col_empleado, "Sueldo", "CARGO"]
    filas = [["ENERO", "1700000000", "FULANO DE TAL", "500", "Operario"]]
    p = _rol(tmp_path, headers, filas)
    dest = str(tmp_path / "adaptado.xlsx")
    out = pc.reescribir_rol(p, dest, mes=1)
    assert "normal" in openpyxl.load_workbook(out).sheetnames


def test_rol_sin_empleado_ni_sueldo_avisa_claro(tmp_path):
    # Un archivo que NO es un rol (sin columnas de empleado ni sueldo) debe fallar con un
    # mensaje claro, no con «Worksheet normal does not exist».
    p = _rol(tmp_path, ["Col1", "Col2", "Col3"], [["a", "b", "c"]])
    with pytest.raises(ValueError, match="No reconozco el rol"):
        pc.reescribir_rol(p, str(tmp_path / "adaptado.xlsx"), mes=1)


def test_normalizar_rol_propaga_el_motivo_en_vez_de_tragarlo(tmp_path):
    # _normalizar_rol ya NO usa «except: pass»: si el rol no se reconoce, PROPAGA el motivo
    # claro en vez de dejar que el motor reviente luego con «Worksheet normal does not exist».
    p = _rol(tmp_path, ["Col1", "Col2"], [["a", "b"]])
    rutas = SimpleNamespace(rol=p, roles_mensuales=[(p, "rol.xlsx")], home=str(tmp_path))
    with pytest.raises(ValueError, match="No reconozco el rol"):
        mayores._normalizar_rol(rutas, pc)
