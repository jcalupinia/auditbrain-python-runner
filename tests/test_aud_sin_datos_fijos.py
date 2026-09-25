"""Papel de trabajo sin cifras calculadas pegadas como valor (decisión del dueño, 2026-09-24).

- Los indicadores de la portada (00_Inicio) son fórmulas a la cédula que los calcula.
- El importe de cada problema remite por fórmula a la celda de la cédula donde se
  origina (procesadores/problemas.py + REF_PROBLEMAS de cada herramienta), en el
  ejercicio modelo, el EJEMPLO y los ESCENARIOS de las 20 herramientas.
- Un mapeo equivocado nunca cambia una cifra: si la celda no tiene ese importe,
  no se enlaza.
- La conclusión no muestra «None» cuando aún no hay conciliación.
"""
import io
import os
import importlib.util

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif import ejercicio_modelo as em
from backend.app.aud.niif.procesadores import PROCESADORES, libro, problemas

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_spec = importlib.util.spec_from_file_location(
    "verificar_problemas_enlazados", os.path.join(RAIZ, "scripts", "verificar_problemas_enlazados.py"))
verificar = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(verificar)


def _libro(pid):
    mod = PROCESADORES[pid]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    reg = em._reg(d, mod, ds, par, corte)
    return load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "APROBADO")))


@pytest.mark.parametrize("pid", list(PROCESADORES))
def test_importe_de_cada_problema_remite_a_su_celda(pid):
    mod = PROCESADORES[pid]
    refs = dict(getattr(mod, "REF_PROBLEMAS", {}) or {})
    for nombre, ds, par, corte in verificar.escenarios(pid, mod):
        run = mod.ejecutar(ds, par, corte)
        _, pend = problemas.enlazar(mod.hojas(run), refs, run.get("exceptions"))
        assert not pend, (pid, nombre, pend)


@pytest.mark.parametrize("pid", list(PROCESADORES))
def test_portada_y_problemas_sin_cifras_pegadas(pid):
    wb = _libro(pid)
    ini = wb["00_Inicio"]
    # Panel (B:F): ninguna cifra escrita como valor; los indicadores son fórmulas.
    fijas = [c.coordinate for r in ini.iter_rows(max_col=6) for c in r
             if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    assert not fijas, (pid, fijas)
    assert any(isinstance(c.value, str) and c.value.startswith("=COUNTA(")
               for r in ini.iter_rows(max_col=6) for c in r), pid
    hp = next(ws for ws in wb.worksheets if [ws.cell(4, j).value for j in (1, 2, 3)] == problemas.COLS)
    fin = next((r for r in range(5, hp.max_row + 1) if not hp.cell(r, 1).value), hp.max_row + 1)
    pegados = [hp.cell(r, 3).coordinate for r in range(5, fin)
               if isinstance(hp.cell(r, 3).value, (int, float)) and abs(hp.cell(r, 3).value) >= 0.005]
    assert not pegados, (pid, pegados)


def test_mapeo_equivocado_no_cambia_la_cifra():
    hojas = [
        {"name": "05_Detalle", "cols": [["Cliente", "t"], ["Pérdida", "n"]],
         "rows": [["Cliente A", {"f": "B1*2", "v": 100.0}], ["Cliente B", {"f": "B2*2", "v": 250.0}]], "total": None},
        {"name": "09_Problemas", "cols": [["Código", "t"], ["Descripción", "t"], ["Importe", "n"]],
         "rows": [["PERDIDA", "Cliente B: pérdida por evaluar", 250.0], ["OTRO", "Sin celda", 999.0]], "total": None},
    ]
    salida, pend = problemas.enlazar(hojas, {"PERDIDA": ("05_Detalle", "Pérdida"), "OTRO": ("05_Detalle", "Pérdida")})
    filas = salida[1]["rows"]
    assert filas[0][2] == {"f": "'05_Detalle'!B6", "v": 250.0}   # fila del cliente citado
    assert filas[1][2] == 999.0 and pend[0]["codigo"] == "OTRO"   # no coincide: queda y se reporta
    # Importe con signo contrario al de la celda: «-ref».
    hojas[1]["rows"] = [["PERDIDA", "Cliente A", -100.0]]
    salida, _ = problemas.enlazar(hojas, {"PERDIDA": ("05_Detalle", "Pérdida")})
    assert salida[1]["rows"][0][2]["f"] == "-'05_Detalle'!B5"


def test_conclusion_sin_none_cuando_aun_no_hay_conciliacion():
    wb = _libro("perdidas_incurridas_s11")
    texto = " ".join(str(c.value) for r in wb["13_Conclusion"].iter_rows() for c in r if c.value)
    assert "None" not in texto and "Pendiente" in texto


# --- Portada («consola») ------------------------------------------------------------------------
# Las tarjetas de 00_Inicio son las del panel del HTML (resultado principal, población,
# recalculado, registrado, problemas) y cada cifra es una fórmula a la cédula que la calcula.
# La tarjeta que no tiene celda de origen no se muestra: nunca un valor pegado.

@pytest.mark.parametrize("pid", list(PROCESADORES))
def test_portada_tarjetas_son_formulas(pid):
    m = PROCESADORES[pid]
    d = m.definicion()
    ds, par, corte = em.escenario(m)
    reg = em._reg(d, m, ds, par, corte)
    ws = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "MUESTRA")))["00_Inicio"]
    r0 = next(c.row for c in ws["B"] if c.value == "INDICADORES CLAVE")
    tarjetas = [(ws[f"{c}{r0 + 1}"].value, ws[f"{c}{r0 + 2}"].value) for c in "BCDEF" if ws[f"{c}{r0 + 1}"].value]
    assert 4 <= len(tarjetas) <= 5, (pid, tarjetas)
    assert tarjetas[-1][0] == "Problemas encontrados", pid
    for rotulo, v in tarjetas:
        assert isinstance(v, str) and v.startswith("="), (pid, rotulo, v)
