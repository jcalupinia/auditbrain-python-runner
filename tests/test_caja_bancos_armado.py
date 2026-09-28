# -*- coding: utf-8 -*-
"""Puente registro de prueba → papel formulado de Caja y Bancos."""
import io

import openpyxl

from backend.app.aud.niif.procesadores import caja_bancos_armado as arm
from backend.app.aud.niif.procesadores import caja_bancos_papel as papel

REG = {
    "engagement": {"client": "LANSEY S.A", "cutoff": "2026-08-31", "period": "ago-2026",
                   "preparer": "WR", "reviewer": "AT"},
    "datasets": {
        "cuentas": [
            {"id": "11010201", "nombre": "BANCO PICHINCHA", "saldo_libros": 5000,
             "saldo_banco": 5120, "saldo_anterior": 4000, "ncuenta": "3033617504"},
        ],
        "partidas": [
            {"cuenta": "11010201", "tipo": "Cheque pendiente", "referencia": "Cheque 101",
             "fecha_origen": "2026-08-20", "importe": 120},
            {"cuenta": "11010201", "tipo": "Nota de crédito", "referencia": "NC banco",
             "fecha_origen": "2026-08-25", "importe": 30},
        ],
    },
}


def test_hay_datos():
    assert arm.hay_datos(REG) is True
    assert arm.hay_datos({}) is False
    assert arm.hay_datos({"datasets": {"cuentas": []}}) is False


def test_arma_papel_valido_con_mapeo_de_categorias():
    data = arm.armar_desde_registro(REG)
    assert data[:2] == b"PK"  # xlsx (zip)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # Cédulas del papel, sin Procedimientos.
    assert "Sumaria" in wb.sheetnames and "Procedimientos" not in wb.sheetnames
    # La partida "Cheque pendiente" del procesador se mapeó a la categoría del papel.
    pc = wb["Partidas conciliatorias"]
    categorias = {pc.cell(r, 3).value for r in range(10, 12)}
    assert papel.CHEQUE in categorias
    assert papel.NC_PENDIENTE in categorias
    # El banco de la partida quedó legible (nombre, no código).
    bancos = {pc.cell(r, 2).value for r in range(10, 12)}
    assert "BANCO PICHINCHA" in bancos


def test_sumaria_toma_saldos_del_dataset():
    wb = openpyxl.load_workbook(io.BytesIO(arm.armar_desde_registro(REG)))
    sm = wb["Sumaria"]
    assert sm["A9"].value == "11010201"
    assert sm["C9"].value == 4000   # saldo anterior
    assert sm["E9"].value == 5000   # saldo actual (saldo_libros)


REG_CON_MAYOR = {
    **REG,
    "datasets": {
        **REG["datasets"],
        "libro_mayor": [
            {"cuenta": "11010201", "descripcion": "BANCO PICHINCHA", "fecha": "2026-08-05",
             "comprobante": "IN-1", "detalle": "Depósito", "debito": 1500, "credito": 0},
            {"cuenta": "11010201", "descripcion": "BANCO PICHINCHA", "fecha": "2026-08-18",
             "comprobante": "CK-9", "detalle": "Pago", "debito": 0, "credito": 400},
        ],
    },
}


def test_libro_mayor_alimenta_movimiento_y_su_hoja():
    wb = openpyxl.load_workbook(io.BytesIO(arm.armar_desde_registro(REG_CON_MAYOR)))
    mv = wb["Movimiento"]
    # Débitos/créditos reales agregados del mayor (1500 y 400), no el neto estimado.
    assert mv["D9"].value == 1500
    assert mv["E9"].value == 400
    # La hoja Libro Mayor trae los asientos cargados.
    lm = wb["Libro Mayor"]
    assert lm["A4"].value == "11010201"
    assert lm["I4"].value == 1500 and lm["J4"].value == 0


REG_REESTRUCTURA = {
    "engagement": {"client": "X", "period": "sep-2026", "cutoff": "2026-09-30"},
    "datasets": {
        "cuentas": [{"id": "11010201", "nombre": "BANCO PICHINCHA", "saldo_libros": 1000,
                     "saldo_banco": 900, "saldo_anterior": 0, "ncuenta": "1"}],
        # cuando hay estado de cuenta, las partidas NO se toman de RQ-002 sino del recálculo
        "partidas": [{"cuenta": "11010201", "tipo": "Otra partida", "referencia": "IGNORAR",
                      "fecha_origen": "2026-09-01", "importe": 999}],
        "libro_mayor": [
            {"cuenta": "11010201", "fecha": "2026-09-28", "debito": 500, "credito": 0},   # depósito en tránsito
            {"cuenta": "11010201", "fecha": "2026-09-10", "debito": 1000, "credito": 0},  # coincide con extracto
        ],
        "estado_cuenta": [
            {"cuenta": "11010201", "fecha": "2026-09-10", "documento": "Depósito", "debito": 0, "credito": 1000},
            {"cuenta": "11010201", "fecha": "2026-09-30", "documento": "Comisión", "debito": 25, "credito": 0},  # ND
        ],
        "conciliacion_anterior": [
            {"cuenta": "11010201", "fecha": "2026-08-01", "categoria": "Cheque sin cobrar",
             "documento": "Cheque 090", "valor": 777, "observacion": "Pendiente"},
        ],
    },
}


def test_reestructuracion_automatica_reemplaza_partidas_manuales():
    wb = openpyxl.load_workbook(io.BytesIO(arm.armar_desde_registro(REG_REESTRUCTURA)))
    pc = wb["Partidas conciliatorias"]
    cats = {}
    for r in range(10, pc.max_row + 1):
        val, cat = pc.cell(r, 6).value, pc.cell(r, 3).value
        if cat:
            cats[cat] = val
    # Depósito en tránsito (500) y ND (25) del recálculo; NO la partida manual de 999.
    assert cats.get(papel.CONSIGNACION) == 500
    assert cats.get(papel.ND_TRANSITO) == 25
    valores = {pc.cell(r, 6).value for r in range(10, pc.max_row + 1)}
    assert 999 not in valores
    # Arrastre del mes anterior: el cheque de 777 sigue abierto.
    assert 777 in valores
