import unittest
from io import BytesIO
from zipfile import ZipFile
import openpyxl
from tests.test_vnr import payload
from backend.app.aud.inventarios_vnr.engine import calculate
from backend.app.aud.inventarios_vnr.parsers import extract
from backend.app.aud.inventarios_vnr.exports import build_xlsx, build_html, schedules

class FileTests(unittest.TestCase):
    def test_csv_and_no_silent_truncation(self):
        e=extract('inventario.csv',b'codigo;cantidad;costo\nA;10;20\n')
        self.assertEqual(e['tables'][0]['rows'][0],['A','10','20'])
        self.assertEqual(len(e['sha256']),64)
        with self.assertRaises(ValueError):extract('large.csv',b'a,b\n'+b'1,2\n'*2001)
    def test_zip_traversal_and_bomb_rejected(self):
        f=BytesIO()
        with ZipFile(f,'w') as z:z.writestr('../escape.csv','a,b\n1,2')
        with self.assertRaises(ValueError):extract('x.zip',f.getvalue())
    def test_xml_entities_rejected(self):
        with self.assertRaises(ValueError):extract('x.xml',b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>')
    def test_image_has_no_fabricated_data(self):
        e=extract('foto.png',b'\x89PNG\r\n\x1a\n')
        self.assertEqual(e['tables'],[]);self.assertIn('manual',e['note'])
    def test_twelve_sheets_formulas_and_cached_results(self):
        r=calculate(payload());buf=build_xlsx(r)
        wb=openpyxl.load_workbook(BytesIO(buf),data_only=False)
        self.assertEqual(len(wb.sheetnames),12)
        self.assertTrue(wb['06 VNR']['H6'].value.startswith('='))
        cached=openpyxl.load_workbook(BytesIO(buf),data_only=True)
        self.assertEqual(cached['06 VNR']['H6'].value,50)
        self.assertEqual(cached['06 VNR']['J6'].value,45)
        self.assertEqual(len(schedules(r)),12)
    def test_formula_injection_and_html_escape(self):
        p=payload();p['rows'][0]['description']='=HYPERLINK("evil")';p['context']['client']='<script>alert(1)</script>'
        r=calculate(p);wb=openpyxl.load_workbook(BytesIO(build_xlsx(r)))
        self.assertEqual(wb['03 Inventario']['B6'].data_type,'s')
        html=build_html(r)
        self.assertNotIn('<script>alert(1)</script>',html)
        self.assertIn('&lt;script&gt;',html)

if __name__=='__main__':unittest.main()

class XmlRegressionTests(unittest.TestCase):
    def test_utf16_dtd_is_rejected(self):
        xml='<!DOCTYPE a [<!ENTITY x "expanded">]><a><row><code>&x;</code></row></a>'.encode('utf-16')
        with self.assertRaises(ValueError):extract('x.xml',xml)

class UnitExportRegressionTests(unittest.TestCase):
    def test_unused_ratio_parameters_do_not_break_unit_export(self):
        p=payload();p['policy'].update(selling_method='unit',eligible_expenses='10,50',sales_base='sin aplicar')
        self.assertTrue(build_xlsx(calculate(p)).startswith(b'PK'))
