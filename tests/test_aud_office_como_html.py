"""El Word y el PowerPoint del papel NIIF se ven igual al HTML (pedido del dueño, 2026-09-25).

- PowerPoint = el HTML en pantalla (tema «Ejecutivo»): fondo azul marino, portada con los
  dos logotipos, las 5 tarjetas del panel y los 4 gráficos del HTML.
- Word = el HTML impreso (tema «Claro», el mismo del PDF): membrete con los logotipos, las
  5 tarjetas, los 4 gráficos y cada cédula con «Cómo se calcula esta hoja».

Los gráficos son los mismos SVG del HTML dibujados como imagen (``svg_png``), y las cifras de
las tarjetas salen de ``html_ejecutivo.kpis_datos``: los tres formatos leen de la misma fuente.
"""
import io

import pytest
from PIL import Image

from backend.app.aud.niif import ejercicio_modelo as em
from backend.app.aud.niif.procesadores import PROCESADORES, graficos, libro, papel_office, svg_png
from backend.app.aud.niif.procesadores import html_ejecutivo as hx
from backend.app.aud.niif.procesadores import graficos_svg as gs


def _reg(pid):
    mod = PROCESADORES[pid]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    return d, mod, em._reg(d, mod, ds, par, corte)


def _es_grafico(blob: bytes) -> bool:
    """Imagen con la proporción de los gráficos del HTML (560×280)."""
    w, h = Image.open(io.BytesIO(blob)).size
    return abs(w / h - gs.ANCHO / gs.ALTO) < 0.01


def test_svg_png_dibuja_el_svg_del_html():
    png = svg_png.a_png(gs.columnas([("A", 100.0), ("B", 40.0)], "barras_linea", "serie", "prueba", gs.CLARO))
    im = Image.open(io.BytesIO(png))
    assert im.size == (int(gs.ANCHO * 2.5), int(gs.ALTO * 2.5))
    # Hay tinta del color de la serie (#2a78d6) donde el SVG dibuja las barras.
    assert (0x2A, 0x78, 0xD6, 255) in {px for px in im.convert("RGBA").getdata()}
    dona = svg_png.a_png(gs.dona([("A", 3.0), ("B", 1.0)], "x", gs.CLARO))
    assert Image.open(io.BytesIO(dona)).size == im.size


@pytest.mark.parametrize("pid", list(PROCESADORES))
def test_powerpoint_como_el_html(pid):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    d, mod, reg = _reg(pid)
    prs = Presentation(io.BytesIO(libro.pptx(d, reg, [], 1, "APROBADO")))
    diaps = list(prs.slides)
    # Fondo azul marino del tema «Ejecutivo» en todas las diapositivas.
    assert all(str(s.background.fill.fore_color.rgb) == papel_office.OSCURO["bg"] for s in diaps), pid
    fotos = [sh for s in diaps for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert sum(_es_grafico(sh.image.blob) for sh in fotos) == 4, pid        # los 4 gráficos del HTML
    assert not any(sh.has_chart for s in diaps for sh in s.shapes), pid     # y ningún otro gráfico
    texto = " ".join(sh.text_frame.text for sh in diaps[1].shapes if sh.has_text_frame)
    p = graficos.panel(mod, reg["run"], reg["run"].get("hojas") or [])
    for k in hx.kpis_datos(p):                                              # las 5 tarjetas del HTML
        assert k["valor"] in texto and k["rotulo"].upper() in texto, (pid, k["valor"])


@pytest.mark.parametrize("pid", ["perdidas_incurridas_s11", "nomina_beneficios", "impuesto_corriente_diferido"])
def test_word_como_el_html_impreso(pid):
    from docx import Document

    d, mod, reg = _reg(pid)
    doc = Document(io.BytesIO(libro.docx(d, reg, [], 1, "APROBADO")))
    cuerpo = [r.target_part.blob for r in doc.part.rels.values() if "image" in r.reltype]
    assert sum(_es_grafico(b) for b in cuerpo) == 4, pid                    # los 4 gráficos del HTML
    textos = [c.text for t in doc.tables for f in t.rows for c in f.cells] + [pa.text for pa in doc.paragraphs]
    todo = " ".join(textos)
    p = graficos.panel(mod, reg["run"], reg["run"].get("hojas") or [])
    for k in hx.kpis_datos(p):
        assert k["valor"] in todo and k["rotulo"].upper() in todo, (pid, k["valor"])
    for h in libro.cedulas(d, reg, [], 1, "APROBADO"):                     # cada cédula con su título
        assert h["label"] in todo, (pid, h["label"])
    assert "ⓘ  Cómo se calcula esta hoja" in todo, pid


def test_word_xml_en_el_orden_del_esquema():
    """Word marca «contenido ilegible» si los hijos de tcPr / pPr van en otro orden."""
    from docx import Document
    from docx.oxml.ns import qn

    d, _, reg = _reg("nomina_beneficios")
    doc = Document(io.BytesIO(libro.docx(d, reg, [], 1, "APROBADO")))
    orden_tc = [qn(f"w:{t}") for t in ("tcW", "gridSpan", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
                                       "textDirection", "tcFitText", "vAlign", "hideMark")]
    partes = [doc.element] + [s.header._element for s in doc.sections] + [s.footer._element for s in doc.sections]
    for raiz in partes:
        for tcPr in raiz.iter(qn("w:tcPr")):
            rango = [orden_tc.index(h.tag) for h in tcPr if h.tag in orden_tc]
            assert rango == sorted(rango), [h.tag for h in tcPr]
        for pPr in raiz.iter(qn("w:pPr")):
            hijos = [h.tag for h in pPr]
            if qn("w:pBdr") in hijos and qn("w:spacing") in hijos:
                assert hijos.index(qn("w:pBdr")) < hijos.index(qn("w:spacing")), hijos
