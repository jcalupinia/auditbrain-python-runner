"""Endpoints de la herramienta de pérdidas esperadas (AUD, require_staff)."""
import io
import json
import uuid
from datetime import date

from openpyxl import Workbook, load_workbook

from backend.app.auth import service as auth_service
from backend.app.auth.models import Role
from backend.app.db.session import SessionLocal, init_db

BASE = "/api/v1/aud/pce-cxc"


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _archivos():
    c23 = _xlsx([("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 100000.0)])
    c24 = _xlsx([("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 20000.0)])
    c25 = _xlsx([("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 10000.0)])
    return [("archivos", ("2023.xlsx", c23)), ("archivos", ("2024.xlsx", c24)),
            ("archivos", ("2025.xlsx", c25))]


def _token(client, role=Role.user):
    init_db()
    tag = uuid.uuid4().hex[:6]
    email, pw = f"pce-{tag}@ex.com", "Sup3rSecret!"
    db = SessionLocal()
    try:
        auth_service.create_user(db, email=email, password=pw, role=role)
    finally:
        db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": pw})
    return r.json()["access_token"]


def test_sin_rol_staff_no_se_puede_calcular(client):
    # El repo solo tiene tres roles: admin, user y client (Role.staff no existe).
    # admin y user SON staff para require_staff; el único rol no-staff es
    # "client" (portal del cliente), y ese lo rechaza get_current_user en
    # defensa en profundidad antes de llegar al chequeo de rol de
    # require_staff ("client-role JWTs must NEVER pass through staff
    # dependencies"), así que la respuesta es 401, no 403. Lo que importa acá
    # es que el cálculo NO se ejecute para un usuario de portal cliente.
    token = _token(client, Role.client)
    r = client.post(f"{BASE}/analizar", files=_archivos(),
                    data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"]})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code in (401, 403), r.text


def test_calcula_y_guarda_la_corrida(client):
    token = _token(client)
    r = client.post(f"{BASE}/analizar", files=_archivos(),
                    data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"],
                                                    "entidad": "PRUEBA S.A."})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["corrida_id"]
    assert cuerpo["ecl_total"] >= 0
    g = client.get(f"{BASE}/corridas/{cuerpo['corrida_id']}", headers={"Authorization": f"Bearer {token}"})
    assert g.status_code == 200
    assert g.json()["resultado"]["ecl_total"] == cuerpo["ecl_total"]


def test_exige_los_tres_cortes(client):
    token = _token(client)
    r = client.post(f"{BASE}/analizar", files=_archivos()[:2],
                    data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31"]})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 400
    assert "tres" in r.json()["detail"].lower()


def test_rechaza_fecha_en_formato_invalido(client):
    """Una fecha en formato inválido (ej: 31/12/2023) debe retornar 400 con mensaje claro."""
    token = _token(client)
    r = client.post(f"{BASE}/analizar", files=_archivos(),
                    data={"parametros": json.dumps({"fechas": ["31/12/2023", "2024-12-31", "2025-12-31"]})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 400, f"Se esperaba 400, se obtuvo {r.status_code}: {r.text}"
    detalle = r.json()["detail"].lower()
    # El mensaje debe mencionar la fecha inválida y el formato esperado
    assert "fecha" in detalle or "formato" in detalle or "aaaa-mm-dd" in detalle, \
        f"El mensaje no explica qué está mal: {r.json()['detail']}"


def test_rechaza_archivo_mayor_al_limite_413(client, monkeypatch):
    """Un archivo que supera MAX_BYTES_POR_ARCHIVO debe retornar 413 con mensaje explicativo."""
    from backend.app.aud.pce_cxc import router

    # Parchea el límite a algo muy pequeño para la prueba (1 KB)
    limite_pequeño = 1 * 1024  # 1 KB
    monkeypatch.setattr(router, "MAX_BYTES_POR_ARCHIVO", limite_pequeño)

    token = _token(client)
    r = client.post(f"{BASE}/analizar", files=_archivos(),
                    data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"]})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 413, f"Se esperaba 413, se obtuvo {r.status_code}: {r.text}"
    detalle = r.json()["detail"].lower()
    # El mensaje debe mencionar el límite en MB
    assert "mb" in detalle, f"El mensaje no menciona el límite en MB: {r.json()['detail']}"


def test_sin_rol_staff_no_puede_consultar_corrida(client):
    """Un usuario sin rol staff no puede leer una corrida con GET /corridas/{id}."""
    # Primero, crear una corrida con un staff
    token_staff = _token(client)
    r_crear = client.post(f"{BASE}/analizar", files=_archivos(),
                          data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"]})},
                          headers={"Authorization": f"Bearer {token_staff}"})
    assert r_crear.status_code == 200, r_crear.text
    corrida_id = r_crear.json()["corrida_id"]

    # Ahora intentar leer con un usuario no-staff
    token_no_staff = _token(client, Role.client)
    r_leer = client.get(f"{BASE}/corridas/{corrida_id}",
                        headers={"Authorization": f"Bearer {token_no_staff}"})
    assert r_leer.status_code in (401, 403), \
        f"Se esperaba 401 o 403 para usuario no-staff, se obtuvo {r_leer.status_code}: {r_leer.text}"


def test_descarga_el_excel_del_papel_de_trabajo(client):
    """GET /corridas/{id}/excel devuelve un .xlsx válido con las trece hojas del papel de trabajo."""
    token = _token(client)
    r_crear = client.post(f"{BASE}/analizar", files=_archivos(),
                          data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"],
                                                          "entidad": "PRUEBA S.A."})},
                          headers={"Authorization": f"Bearer {token}"})
    assert r_crear.status_code == 200, r_crear.text
    corrida_id = r_crear.json()["corrida_id"]

    r = client.get(f"{BASE}/corridas/{corrida_id}/excel", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert ".xlsx" in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["00-Caratula", "01-Parametros", "02-Fuentes", "03-Cohorte", "04-Tasas",
                             "05-Matriz", "06-Individual", "07-Politica", "08-Conciliacion",
                             "09-Tributario", "10-Hallazgos", "11-Pendientes", "12-Bitacora"]


def test_excel_de_corrida_inexistente_da_404(client):
    token = _token(client)
    r = client.get(f"{BASE}/corridas/999999/excel", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 404


def test_sin_rol_staff_no_puede_descargar_el_excel(client):
    """Un usuario de portal cliente no puede bajar el papel de trabajo interno."""
    token_staff = _token(client)
    r_crear = client.post(f"{BASE}/analizar", files=_archivos(),
                          data={"parametros": json.dumps({"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"]})},
                          headers={"Authorization": f"Bearer {token_staff}"})
    assert r_crear.status_code == 200, r_crear.text
    corrida_id = r_crear.json()["corrida_id"]

    token_no_staff = _token(client, Role.client)
    r = client.get(f"{BASE}/corridas/{corrida_id}/excel", headers={"Authorization": f"Bearer {token_no_staff}"})
    assert r.status_code in (401, 403), \
        f"Se esperaba 401 o 403 para usuario no-staff, se obtuvo {r.status_code}: {r.text}"
