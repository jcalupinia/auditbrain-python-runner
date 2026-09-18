"""Recalculador mínimo de fórmulas de Excel para las pruebas del papel de trabajo.

POR QUÉ EXISTE. Un papel de trabajo que no se puede recalcular desde sus
propias celdas no sirve como evidencia de auditoría: el revisor rehace la
cuenta con las cifras impresas y tiene que llegar a lo archivado. Las pruebas
de extremo a extremo hacen exactamente eso -leen las celdas del libro y
evalúan sus fórmulas-, en vez de comprobar que la fórmula "empiece por
=ROUND(", que es verificar la implementación y no el comportamiento.

QUÉ CUBRE. Exactamente el subconjunto que el exportador declara en su
docstring: ``SUM``, ``SUMIFS``, ``COUNTIFS``, ``INDEX``, ``MATCH``, ``IF``,
``ABS``, ``ROUND``, ``MAX`` y ``MIN``, más referencias entre hojas, rangos,
nombres definidos y los operadores aritméticos y de comparación. No pretende
ser Excel: si el libro usara algo fuera de ese subconjunto, aquí levanta un
error en vez de devolver un número inventado.

CRITERIOS QUE IMPORTAN:

- ``ROUND`` redondea medio hacia afuera del cero, igual que Excel y que
  ``motor.redondear`` (``Decimal.quantize`` con ``ROUND_HALF_UP``), así que
  las dos implementaciones coinciden al centavo en ambos signos.
- ``SUM`` ignora el texto, igual que Excel: una banda rotulada "SIN MEDIR" no
  suma cero, sencillamente no suma.
- Una celda vacía vale 0 en aritmética, y el texto en una operación
  aritmética levanta un error (Excel daría ``#¡VALOR!``): el papel no debe
  tener ninguna.
"""
from __future__ import annotations

import re
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

_CELDA = re.compile(r"^\$?[A-Z]{1,3}\$?[0-9]+$")

_TOKEN = re.compile(
    r"""\s*(?:
      (?P<cadena>"(?:[^"]|"")*")
    | (?P<hoja>(?:'[^']+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?[0-9]+(?::\$?[A-Z]{1,3}\$?[0-9]+)?)
    | (?P<numero>[0-9]+(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?)
    | (?P<nombre>[A-Za-z_][\w.]*)
    | (?P<op><=|>=|<>|[-+*/<>=&])
    | (?P<abre>\()
    | (?P<cierra>\))
    | (?P<coma>,)
    | (?P<dospuntos>:)
    )""",
    re.VERBOSE,
)


class ErrorDeFormula(Exception):
    """La fórmula usa algo que este recalculador no cubre, o no cuadra."""


def redondear_excel(valor: float, decimales: int = 2) -> float:
    """``ROUND`` de Excel: medio hacia afuera del cero, como ``motor.redondear``."""
    cuantum = Decimal(1).scaleb(-int(decimales))
    return float(Decimal(str(valor)).quantize(cuantum, rounding=ROUND_HALF_UP))


# ---------------------------------------------------------------------------
# Análisis sintáctico
# ---------------------------------------------------------------------------

def _tokenizar(formula: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    pos = 0
    while pos < len(formula):
        if formula[pos].isspace():
            pos += 1
            continue
        m = _TOKEN.match(formula, pos)
        if not m or m.end() == pos:
            raise ErrorDeFormula(f"No se reconoce '{formula[pos:]}' en «{formula}»")
        tipo = m.lastgroup
        tokens.append((tipo, m.group(tipo).strip()))
        pos = m.end()
    return tokens


class _Analizador:
    """Descenso recursivo sobre los tokens de una fórmula."""

    def __init__(self, tokens: list[tuple[str, str]]):
        self.tokens = tokens
        self.i = 0

    def _mirar(self) -> tuple[str, str] | None:
        return self.tokens[self.i] if self.i < len(self.tokens) else None

    def _tomar(self) -> tuple[str, str]:
        t = self.tokens[self.i]
        self.i += 1
        return t

    def analizar(self) -> Any:
        nodo = self._comparacion()
        if self.i != len(self.tokens):
            raise ErrorDeFormula(f"Sobran tokens: {self.tokens[self.i:]}")
        return nodo

    def _comparacion(self) -> Any:
        izq = self._suma()
        while (t := self._mirar()) and t[0] == "op" and t[1] in ("=", "<>", "<", ">", "<=", ">="):
            self._tomar()
            izq = ("cmp", t[1], izq, self._suma())
        return izq

    def _suma(self) -> Any:
        izq = self._producto()
        while (t := self._mirar()) and t[0] == "op" and t[1] in ("+", "-", "&"):
            self._tomar()
            izq = ("bin", t[1], izq, self._producto())
        return izq

    def _producto(self) -> Any:
        izq = self._unario()
        while (t := self._mirar()) and t[0] == "op" and t[1] in ("*", "/"):
            self._tomar()
            izq = ("bin", t[1], izq, self._unario())
        return izq

    def _unario(self) -> Any:
        t = self._mirar()
        if t and t[0] == "op" and t[1] in ("+", "-"):
            self._tomar()
            return ("neg", self._unario()) if t[1] == "-" else self._unario()
        return self._primario()

    def _primario(self) -> Any:
        t = self._tomar()
        tipo, texto = t
        if tipo == "numero":
            return ("num", float(texto))
        if tipo == "cadena":
            return ("txt", texto[1:-1].replace('""', '"'))
        if tipo == "abre":
            nodo = self._comparacion()
            self._esperar("cierra")
            return nodo
        if tipo == "hoja":
            hoja, rango = texto.split("!", 1)
            return ("ref", hoja.strip("'"), rango.replace("$", ""))
        if tipo == "nombre":
            siguiente = self._mirar()
            if siguiente and siguiente[0] == "abre":
                self._tomar()
                args = self._argumentos()
                return ("fn", texto.upper(), args)
            if _CELDA.match(texto.upper()):
                celda = texto.upper().replace("$", "")
                if (siguiente := self._mirar()) and siguiente[0] == "dospuntos":
                    self._tomar()
                    fin = self._tomar()[1].upper().replace("$", "")
                    return ("ref", None, f"{celda}:{fin}")
                return ("ref", None, celda)
            return ("definido", texto)
        raise ErrorDeFormula(f"Token inesperado: {t}")

    def _argumentos(self) -> list[Any]:
        args: list[Any] = []
        if (t := self._mirar()) and t[0] == "cierra":
            self._tomar()
            return args
        while True:
            args.append(self._comparacion())
            t = self._tomar()
            if t[0] == "cierra":
                return args
            if t[0] != "coma":
                raise ErrorDeFormula(f"Se esperaba ',' o ')' y llegó {t}")

    def _esperar(self, tipo: str) -> None:
        t = self._tomar()
        if t[0] != tipo:
            raise ErrorDeFormula(f"Se esperaba {tipo} y llegó {t}")


# ---------------------------------------------------------------------------
# Evaluación
# ---------------------------------------------------------------------------

class Libro:
    """Evalúa las celdas de un libro de openpyxl abierto CON fórmulas."""

    def __init__(self, wb):
        self.wb = wb
        self._memo: dict[tuple[str, str], Any] = {}
        self._en_curso: set[tuple[str, str]] = set()

    # -- API pública --------------------------------------------------------

    def valor(self, hoja: str, celda: str) -> Any:
        """Valor de una celda: su contenido, o el resultado de su fórmula."""
        clave = (hoja, celda.replace("$", "").upper())
        if clave in self._memo:
            return self._memo[clave]
        if clave in self._en_curso:
            raise ErrorDeFormula(f"Referencia circular en {hoja}!{celda}")
        crudo = self.wb[hoja][clave[1]].value
        self._en_curso.add(clave)
        try:
            valor = self.evaluar(crudo, hoja) if isinstance(crudo, str) and crudo.startswith("=") \
                else crudo
        finally:
            self._en_curso.discard(clave)
        self._memo[clave] = valor
        return valor

    def numero(self, hoja: str, celda: str) -> float:
        """Valor de una celda exigiendo que sea numérico."""
        v = self.valor(hoja, celda)
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise ErrorDeFormula(f"{hoja}!{celda} no es un número: {v!r}")
        return float(v)

    def evaluar(self, formula: str, hoja: str) -> Any:
        """Evalúa una fórmula ('=...') en el contexto de una hoja."""
        texto = formula[1:] if formula.startswith("=") else formula
        nodo = _Analizador(_tokenizar(texto)).analizar()
        return self._eval(nodo, hoja)

    # -- Interno ------------------------------------------------------------

    def _celdas_del_rango(self, hoja: str, rango: str) -> list[tuple[str, str]]:
        if ":" not in rango:
            return [(hoja, rango)]
        min_col, min_fila, max_col, max_fila = range_boundaries(rango)
        return [(hoja, f"{get_column_letter(c)}{f}")
                for f in range(min_fila, max_fila + 1)
                for c in range(min_col, max_col + 1)]

    def _eval(self, nodo, hoja: str) -> Any:
        clase = nodo[0]
        if clase == "num":
            return nodo[1]
        if clase == "txt":
            return nodo[1]
        if clase == "neg":
            return -self._numerico(self._eval(nodo[1], hoja))
        if clase == "ref":
            destino = nodo[1] or hoja
            if ":" in nodo[2]:
                return [self.valor(h, c) for h, c in self._celdas_del_rango(destino, nodo[2])]
            return self.valor(destino, nodo[2])
        if clase == "definido":
            return self._nombre_definido(nodo[1])
        if clase == "bin":
            return self._binario(nodo[1], self._eval(nodo[2], hoja), self._eval(nodo[3], hoja))
        if clase == "cmp":
            return self._comparar(nodo[1], self._eval(nodo[2], hoja), self._eval(nodo[3], hoja))
        if clase == "fn":
            return self._funcion(nodo[1], nodo[2], hoja)
        raise ErrorDeFormula(f"Nodo desconocido: {nodo!r}")

    def _nombre_definido(self, nombre: str) -> Any:
        definidos = self.wb.defined_names
        if nombre not in definidos:
            raise ErrorDeFormula(f"Nombre definido inexistente: {nombre}")
        hoja, celda = next(definidos[nombre].destinations)
        return self.valor(hoja, celda.replace("$", ""))

    @staticmethod
    def _numerico(valor: Any) -> float:
        if valor is None:
            return 0.0
        if isinstance(valor, bool):
            return 1.0 if valor else 0.0
        if isinstance(valor, (int, float)):
            return float(valor)
        raise ErrorDeFormula(f"Se esperaba un número y llegó texto: {valor!r}")

    def _binario(self, op: str, a: Any, b: Any) -> Any:
        if op == "&":
            return f"{'' if a is None else a}{'' if b is None else b}"
        x, y = self._numerico(a), self._numerico(b)
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            if y == 0:
                raise ErrorDeFormula("División por cero (#¡DIV/0!)")
            return x / y
        raise ErrorDeFormula(f"Operador no soportado: {op}")

    @staticmethod
    def _comparar(op: str, a: Any, b: Any) -> bool:
        if isinstance(a, str) or isinstance(b, str):
            a = "" if a is None else (a.upper() if isinstance(a, str) else a)
            b = "" if b is None else (b.upper() if isinstance(b, str) else b)
            if type(a) is not type(b):
                # Excel ordena cualquier texto por encima de cualquier número.
                return {"=": False, "<>": True,
                        "<": isinstance(a, (int, float)), ">": isinstance(b, (int, float)),
                        "<=": isinstance(a, (int, float)), ">=": isinstance(b, (int, float))}[op]
        else:
            a = 0 if a is None else a
            b = 0 if b is None else b
        return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
                "<=": a <= b, ">=": a >= b}[op]

    def _aplanar(self, valores: Any) -> list[Any]:
        return list(valores) if isinstance(valores, list) else [valores]

    def _solo_numeros(self, valores: list[Any]) -> list[float]:
        return [float(v) for v in valores
                if isinstance(v, (int, float)) and not isinstance(v, bool)]

    def _funcion(self, nombre: str, args: list[Any], hoja: str) -> Any:
        if nombre == "IF":
            # Perezosa, como en Excel: la rama no elegida puede ser una división
            # por cero (04-Tasas protege así las bandas sin cartera inicial).
            if len(args) not in (2, 3):
                raise ErrorDeFormula("IF espera 2 o 3 argumentos")
            condicion = self._eval(args[0], hoja)
            if condicion if isinstance(condicion, bool) else bool(self._numerico(condicion)):
                return self._eval(args[1], hoja)
            return self._eval(args[2], hoja) if len(args) == 3 else False

        valores = [self._eval(a, hoja) for a in args]
        if nombre == "SUM":
            planos: list[Any] = []
            for v in valores:
                planos.extend(self._aplanar(v))
            return sum(self._solo_numeros(planos))
        if nombre in ("MAX", "MIN"):
            planos = []
            for v in valores:
                planos.extend(self._aplanar(v))
            numeros = self._solo_numeros(planos)
            if not numeros:
                return 0.0
            return max(numeros) if nombre == "MAX" else min(numeros)
        if nombre == "ROUND":
            if len(valores) != 2:
                raise ErrorDeFormula("ROUND espera 2 argumentos")
            return redondear_excel(self._numerico(valores[0]), int(self._numerico(valores[1])))
        if nombre == "ABS":
            return abs(self._numerico(valores[0]))
        if nombre in ("SUMIFS", "COUNTIFS"):
            return self._sumifs(nombre, valores)
        if nombre == "MATCH":
            objetivo = valores[0]
            lista = self._aplanar(valores[1])
            for pos, v in enumerate(lista, start=1):
                if self._comparar("=", v, objetivo):
                    return pos
            raise ErrorDeFormula(f"MATCH no encontró {objetivo!r}")
        if nombre == "INDEX":
            lista = self._aplanar(valores[0])
            return lista[int(self._numerico(valores[1])) - 1]
        raise ErrorDeFormula(f"Función fuera del subconjunto permitido: {nombre}")

    def _sumifs(self, nombre: str, valores: list[Any]) -> float:
        if nombre == "SUMIFS":
            rango = self._aplanar(valores[0])
            criterios = valores[1:]
        else:
            rango = None
            criterios = valores
        if len(criterios) % 2:
            raise ErrorDeFormula(f"{nombre} espera pares (rango, criterio)")
        pares = [(self._aplanar(criterios[i]), criterios[i + 1])
                 for i in range(0, len(criterios), 2)]
        largo = len(pares[0][0])
        total = 0.0
        for i in range(largo):
            if all(self._comparar("=", c[0][i], c[1]) for c in pares):
                if rango is None:
                    total += 1
                else:
                    v = rango[i]
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        total += float(v)
        return total


def recalcular(binario: bytes, hoja: str, celda: str) -> Any:
    """Atajo: abre el libro en bytes y recalcula una celda desde sus fórmulas."""
    import io

    from openpyxl import load_workbook

    return Libro(load_workbook(io.BytesIO(binario))).valor(hoja, celda)
