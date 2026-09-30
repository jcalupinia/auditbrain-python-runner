# -*- coding: utf-8 -*-
"""Cédulas del papel real DA dentro del procesador «efectivo_equivalentes».

Las 7 cédulas DA (Libro Mayor fuente, Sumaria, Movimiento, Conciliaciones, Partidas,
Arqueo y Hallazgos) se emiten desde ``hojas()`` como botones/cédulas que se llenan al
Procesar, calcadas al papel del cliente con FÓRMULAS VIVAS y referencias cruzadas entre
pestañas (la Sumaria es la fuente de saldos; DA-3 suma las partidas de DA-4 por SUMIFS;
DA-2 pivota el Libro Mayor por SUMIF; DA-5 recuenta por denominación).
"""
import io

import openpyxl

from backend.app.aud.niif.procesadores import efectivo_equivalentes as m
from backend.app.aud.niif.procesadores import datos_cliente, libro

DA = ["DA0_Libro_Mayor", "DA1_Sumaria", "DA2_Movimiento", "DA3_Conciliaciones",
      "DA4_Partidas", "DA5_Arqueo", "DA6_Hallazgos"]


def _run():
    E = m.EJEMPLO
    return E, m.ejecutar(E["datasets"], E["parametros"], E["corte"])


def _hoja(hojas, name):
    return next(h for h in hojas if h["name"] == name)


def test_las_7_cedulas_da_estan_en_cedulas_y_en_hojas():
    _, run = _run()
    nombres = [h["name"] for h in m.hojas(run)]
    # ``hojas()`` y ``CEDULAS`` van en sync (lo exige test_proc_efectivo) y contienen las DA.
    assert nombres == [n for n, _ in m.CEDULAS]
    assert nombres[-7:] == DA


def test_sumaria_variacion_es_formula_viva():
    _, run = _run()
    h = _hoja(m.hojas(run), "DA1_Sumaria")
    cols = [c[0] for c in h["cols"]]
    assert cols == ["Cuenta", "Descripción", "Saldo s/registros anterior", "Variación",
                    "Saldo s/registros actual", "Ref."]
    var = h["rows"][0][cols.index("Variación")]
    assert isinstance(var, dict) and var["f"] == "E5-C5"      # actual − anterior, viva
    assert h["rows"][0][cols.index("Ref.")] == "DA-3"


def test_movimiento_pivota_libro_mayor_por_sumif():
    _, run = _run()
    h = _hoja(m.hojas(run), "DA2_Movimiento")
    cols = [c[0] for c in h["cols"]]
    deb = h["rows"][0][cols.index("Débitos del período")]
    cuadre = h["rows"][0][cols.index("Cuadre s/Sumaria")]
    assert deb["f"].startswith("SUMIF('DA0_Libro_Mayor'!")
    assert "'DA1_Sumaria'!E" in cuadre["f"]                   # cuadra contra la Sumaria


def test_conciliaciones_suma_partidas_de_da4_y_cita_la_sumaria():
    E, run = _run()
    h = _hoja(m.hojas(run), "DA3_Conciliaciones")
    cols = [c[0] for c in h["cols"]]
    # Una fila por cuenta que no es caja (bancos e inversiones).
    conf = [c for c in E["datasets"]["cuentas"] if (c.get("tipo") or "Banco") != "Caja"]
    assert len(h["rows"]) == len(conf)
    fila = h["rows"][0]
    consig = fila[cols.index("(+) Consignaciones no registradas")]
    audit = fila[cols.index("Saldo s/auditoría")]
    registros = fila[cols.index("Saldo s/registros contables")]
    dif = fila[cols.index("Diferencia")]
    assert consig["f"].startswith("SUMIFS('DA4_Partidas'!")
    assert audit["f"] == "D5+E5-F5-G5+H5-I5"                  # extracto ± partidas
    assert registros["f"].startswith("'DA1_Sumaria'!E")       # sin cifra pegada
    assert dif["f"] == "J5-K5"


def test_partidas_dias_vencidos_y_prescripcion_vivos():
    _, run = _run()
    h = _hoja(m.hojas(run), "DA4_Partidas")
    cols = [c[0] for c in h["cols"]]
    fila = h["rows"][0]
    dias = fila[cols.index("Días vencidos")]
    presc = fila[cols.index("Fecha de prescripción")]
    assert "-A5" in dias["f"] and "02_Parametros" in dias["f"]  # corte (Parámetros) − fecha
    assert presc["f"] == 'IF(A5="","",A5+390)'                  # 360 + 30 días


def test_arqueo_por_denominacion_y_diferencia_contra_caja():
    E, run = _run()
    h = _hoja(m.hojas(run), "DA5_Arqueo")
    cols = [c[0] for c in h["cols"]]
    n_arq = len(E["datasets"]["arqueo"])
    total = h["rows"][0][cols.index("Total")]
    assert total["f"] == "B5*C5"                              # cantidad × valor unitario
    # Filas de cierre: total del arqueo, saldo de caja de la Sumaria y diferencia.
    etiquetas = [r[0] for r in h["rows"][n_arq:]]
    assert etiquetas == ["TOTAL ARQUEO", "Saldo de caja según libros (DA-1)", "Diferencia (arqueo − libros)"]


def test_hallazgos_listan_los_problemas_del_calculo():
    _, run = _run()
    h = _hoja(m.hojas(run), "DA6_Hallazgos")
    assert [c[0] for c in h["cols"]] == ["No.", "Observación", "Referencia de PT", "Recomendación"]
    assert len(h["rows"]) == len(run["exceptions"])


def test_da_materializan_en_excel_con_formulas_y_referencias_cruzadas():
    E, run = _run()
    run["hojas"] = datos_cliente.con_datos(m, run, E["datasets"])
    reg = {"engagement": {"client": "LANSEY S.A", "cutoff": E["corte"], "period": E["corte"],
                          "framework": "NIIF completas", "edition": ""},
           "datasets": E["datasets"], "parameters": dict(E["parametros"]), "run": run}
    xb = libro.xlsx(m.definicion(), reg, [], 1, "aprobada")
    wb = openpyxl.load_workbook(io.BytesIO(xb))
    assert all(n in wb.sheetnames for n in DA)
    da3 = wb["DA3_Conciliaciones"]
    # Fórmula viva materializada, no valor pegado, con referencia cruzada a la Sumaria.
    assert str(da3.cell(5, 11).value).startswith("='DA1_Sumaria'!E")
    assert str(da3.cell(5, 5).value).startswith("=SUMIFS('DA4_Partidas'!")
