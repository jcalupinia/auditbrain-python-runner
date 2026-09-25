"""Ficha con procesador especializado (pérdidas incurridas, Sección 11) por HTTP,
de la instalación de la definición al papel aprobado que arma el servidor.

Los anexos se llenan sobre los modelos que el propio portal entrega al cliente.
"""
import io

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif.procesadores import perdidas_incurridas_s11 as pi
from tests.test_aud_ciclo_evidencia import _leer, _subir
from tests.test_aud_ciclo_http import BASE, FICHA, _accion, _staff_con_proyecto
from tests.test_aud_niif_fichas import _crear, _h


@pytest.fixture(autouse=True)
def disco_temporal(tmp_path, monkeypatch):
    monkeypatch.setenv("AUD_CICLO_DIR", str(tmp_path / "aud_pruebas"))
    monkeypatch.setenv("AUD_CICLO_MIN_LIBRE_MB", "0")


def _modelo_lleno(client, tok, p, req, filas):
    r = client.get(f"{BASE}/pruebas/{p['id']}/modelo/{req}", headers=_h(tok))
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    for f in filas:
        wb["Datos"].append(f)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _mapa(tipo):
    return {c["key"]: i for i, c in enumerate(pi.CAMPOS[tipo])}


def test_perdidas_incurridas_de_punta_a_punta(client):
    tok, pid = _staff_con_proyecto(client)
    ficha = _crear(client, tok)
    lista = client.get(f"{BASE}/procesadores", headers=_h(tok)).json()
    assert lista[0]["id"] == "perdidas_incurridas_s11" and lista[0]["ejemplo"]["totales"]["perdida"] == "800.00"
    client.post(f"/api/v1/aud/niif/fichas/{ficha['id']}/estado", headers=_h(tok), json={"estado": "probada"})
    url = f"{BASE}/fichas/{ficha['id']}/definicion"
    mala = {**pi.definicion(), "processor": "otro"}
    assert client.put(url, headers=_h(tok), json={"definicion": mala, "filas": []}).status_code == 400
    r = client.put(url, headers=_h(tok), json={"definicion": pi.definicion(), "filas": []})
    assert r.status_code == 200, r.text
    # El catálogo del módulo la ubica por rubro y distingue probada de enviada.
    h = next(x for x in client.get(f"{BASE}/herramientas", headers=_h(tok)).json() if x["origen"] == f"ficha:{ficha['id']}")
    assert h["estado"] == "probada" and h["marcos"] == ["NIIF para las PYMES"]

    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok),
                      json={**FICHA, "framework": "NIIF para las PYMES"}).status_code == 200
    p = client.post(f"{BASE}/proyectos/{pid}/pruebas", headers=_h(tok), json={"origen": f"ficha:{ficha['id']}"}).json()
    p = _accion(client, tok, p, "research").json()
    p = _accion(client, tok, p, "generate_program").json()
    prog = p["registro"]["program"]
    assert [x["code"] for x in prog][:2] == ["CXCPI-01", "CXCPI-02"]
    codigos = [x["code"] for x in prog]
    fuentes = p["registro"]["sources"]
    fuentes[0].update(verified=True, section="Secc. 11 párr. 11.21–11.26", date="vigente", procedures=codigos)
    fuentes[1].update(verified=True, document="NIA 540", section="párr. 13", date="vigente", procedures=codigos)
    p = _accion(client, tok, p, "approve_program", {"program": prog, "sources": fuentes}).json()
    assert p["estado"] == "PROGRAMA_APROBADO", p
    p = _accion(client, tok, p, "generate_request").json()
    reqs = p["registro"]["requests"]
    assert [r["dataset"] for r in reqs if "dataset" in r] == ["a3", "a2", "a1", "provision", "movimiento"]
    p = _accion(client, tok, p, "approve_request", {"requests": reqs}).json()
    assert p["estado"] == "REQUERIMIENTO_APROBADO", p
    # Un modelo por anexo de cálculo; los de soporte no tienen.
    assert _leer(client, tok, p)["modelos"] == ["RQ-001", "RQ-002", "RQ-003", "RQ-004", "RQ-005"]

    anexos = {
        "RQ-001": [["F-1", "A", "2024-07-01", "2024-08-15", 400], ["F-3", "C", "2025-07-01", "2025-08-15", 2000]],
        "RQ-002": [["F-1", "A", "2024-07-01", "2024-08-15", 1000], ["F-2", "B", "2024-12-15", "2025-01-14", 500]],
        "RQ-004": [["F-1", "A", 200, 0], ["F-9", "Z", 100, 25]],
        "RQ-005": [["2024", 0, 300, 0, 0], ["2025", None, 0, 0, 0]],
    }
    for req, filas in anexos.items():
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, f"{req}.xlsx", _modelo_lleno(client, tok, p, req, filas)).status_code == 201
    for req in ("RQ-006", "RQ-007"):
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, f"{req}.pdf", b"%PDF").status_code == 201
    p = _leer(client, tok, p)
    assert p["huecos"] == [], p["huecos"]

    arch = {a["requerimiento"]: a["id"] for a in p["archivos"]}
    parte = lambda req, tipo: [{"fileId": arch[req], "sheet": "Datos", "header": 1, "mapping": _mapa(tipo)}]
    # Un archivo de otro requerimiento no entra como anexo.
    r = _accion(client, tok, p, "map_validate", {"datasets": {"a3": parte("RQ-002", "cartera")}})
    assert r.status_code == 400 and "no corresponde a RQ-001" in r.json()["detail"]
    p = _accion(client, tok, p, "map_validate", {"datasets": {
        "a3": parte("RQ-001", "cartera"), "a2": parte("RQ-002", "cartera"),
        "provision": parte("RQ-004", "provision"), "movimiento": parte("RQ-005", "movimiento")}}).json()
    assert p["registro"]["validation"]["ok"], p["registro"]["validation"]
    assert p["registro"]["controlTotal"] == "2400.00"
    p = _accion(client, tok, p, "validate", {"evidenceReviewed": True, "evidenceReview": "Cotejé anexos con el mayor.",
                                             "ledger": "2400.00", "tolerance": "0"}).json()
    assert p["estado"] == "DOCUMENTACION_VALIDADA", p

    r = _accion(client, tok, p, "configure", {"basis": "Ficha CXC-PI-01.", "parametros": {"tasas": {"t730": "150"}}})
    assert r.status_code == 400 and "entre 0 y 100" in r.json()["detail"]
    p = _accion(client, tok, p, "configure", {"basis": "Parámetros de la ficha CXC-PI-01; tasa del tramo grave por gestión de cobro.",
                                              "parametros": {"tasaDesc": "0", "tasas": {"t730": "100"}}}).json()
    assert p["registro"]["parameters"]["tasas"] == {"t730": 100}
    p = _accion(client, tok, p, "approve_methodology").json()
    # Sin resultado del navegador: el procesador solo existe en Python.
    p = _accion(client, tok, p, "execute", {}).json()
    assert p["estado"] == "PRUEBA_EJECUTADA", p
    run = p["registro"]["run"]
    # F-3 (91–180, tasa 40 %) → 800; F-1 (361–730, tasa fijada 100 %) → 400.
    assert run["totals"]["perdida"] == "1200.00" and run["totals"]["ajuste"] == "900.00"
    assert [h["name"] for h in run["hojas"]][:4] == ["01_Resumen", "02_Parametros", "03_Evidencia_historica", "04_Matriz_deterioro"]

    # El papel en curso se baja en los cuatro formatos; el Excel lleva fórmulas.
    for fmt, firma in (("xlsx", b"PK"), ("docx", b"PK"), ("pptx", b"PK"), ("html", b"<!doctype html>")):
        r = client.get(f"{BASE}/pruebas/{p['id']}/libro?formato={fmt}", headers=_h(tok))
        assert r.status_code == 200 and r.content.startswith(firma), fmt
    wb = load_workbook(io.BytesIO(client.get(f"{BASE}/pruebas/{p['id']}/libro", headers=_h(tok)).content))
    assert wb["11_Detalle"]["J5"].value.startswith("=IF(")
    assert client.get(f"{BASE}/pruebas/{p['id']}/libro?formato=exe", headers=_h(tok)).status_code == 400

    p = _accion(client, tok, p, "analyze").json()
    p = _accion(client, tok, p, "submit", {"analysis": p["registro"]["analysis"], "conclusion": "Ajuste propuesto de 900,00."}).json()
    assert p["estado"] == "EN_REVISION", p
    p = _accion(client, tok, p, "approve", {"conclusion": "Se propone ajustar la provisión en 900,00.", "conclusionReviewed": True,
                                            "exceptionReview": "Tramos, clientes en mora grave y límites fiscales evaluados."}).json()
    assert p["estado"] == "APROBADO", p
    art = p["registro"]["artifacts"]
    assert set(art) == {"xlsx", "html"}
    x = client.get(f"{BASE}/pruebas/{p['id']}/archivos/{art['xlsx']['id']}", headers=_h(tok))
    wb = load_workbook(io.BytesIO(x.content))
    # El papel ejecutivo antepone el panel "00_Inicio" (marca + KPIs + navegación)
    # a la carátula; "00_Caratula" sigue presente como cédula (libro.xlsx, 2026).
    # Tras "14_Control_Revision" se anexa el panel "16_Evidencia" (cobertura de
    # requerimientos ↔ archivos, del enriquecimiento del ciclo).
    assert wb.sheetnames[0] == "00_Inicio" and "00_Caratula" in wb.sheetnames \
        and "04_Matriz_deterioro" in wb.sheetnames and "14_Control_Revision" in wb.sheetnames \
        and "16_Evidencia" in wb.sheetnames
    detalle = wb["11_Detalle"]
    assert detalle.cell(row=4, column=1).value == "Factura"
    # La hoja ejecutiva anexa una sección "Notas de fórmulas" tras la fila TOTAL,
    # así que TOTAL ya no es la última fila: se busca donde esté (libro.xlsx 2026).
    primeras = [detalle.cell(row=r, column=1).value for r in range(1, detalle.max_row + 1)]
    assert "TOTAL" in primeras
