"""Logo de la compañía auditada subido en la ficha del encargo.

El auditor sube el logo del cliente en la ficha (``datos["logoCliente"]`` como
*data URI*). Reemplaza al logo de ejemplo (LANSEY) en el HTML del papel y, cuando
es raster, aparece en la portada del Excel/Word/PowerPoint. Sin logo, el papel usa
el nombre de la compañía y el HTML muestra el placeholder «Logo del cliente»
(nunca LANSEY en una corrida real)."""
import base64
import io

import pytest
from PIL import Image

from backend.app.aud.niif import ejercicio_modelo as em
from backend.app.aud.niif.ciclo import reglas
from backend.app.aud.niif.procesadores import PROCESADORES, artefacto_html, libro, marca


def _png_uri(w=120, h=48, color=(10, 35, 66)) -> str:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _svg_uri() -> str:
    return "data:image/svg+xml;base64," + base64.b64encode(
        b"<svg xmlns='http://www.w3.org/2000/svg'/>").decode()


_FICHA_BASE = dict(
    client="ACME S.A.", ruc="1790012345001", activity="Comercio", year="2025",
    cutoff="2025-12-31", preparer="JV", reviewer="MC", firm="Audit Consulting",
    framework="NIIF para las PYMES", country="Ecuador", currency="USD",
    visit="Final", edition="2015", reuseScope="one", deferredTax=False,
)


def test_la_ficha_conserva_el_logo_png_y_svg():
    png = _png_uri()
    assert reglas.validar_ficha_encargo({**_FICHA_BASE, "logoCliente": png})["logoCliente"] == png
    assert reglas.validar_ficha_encargo({**_FICHA_BASE, "logoCliente": _svg_uri()})["logoCliente"] == _svg_uri()


def test_la_ficha_sin_logo_queda_vacia():
    assert reglas.validar_ficha_encargo(dict(_FICHA_BASE))["logoCliente"] == ""


def test_la_ficha_rechaza_un_logo_que_no_es_imagen():
    for malo in ("no-es-imagen", "data:text/plain;base64,aaa", "javascript:alert(1)"):
        with pytest.raises(reglas.ReglaIncumplida):
            reglas.validar_ficha_encargo({**_FICHA_BASE, "logoCliente": malo})


def test_la_ficha_rechaza_un_logo_demasiado_grande():
    enorme = "data:image/png;base64," + "A" * (reglas.MAX_LOGO_CLIENTE_CHARS + 1)
    with pytest.raises(reglas.ReglaIncumplida):
        reglas.validar_ficha_encargo({**_FICHA_BASE, "logoCliente": enorme})


def test_marca_decodifica_raster_pero_omite_svg():
    assert marca.bytes_de_uri(_png_uri()) is not None
    assert marca.bytes_de_uri(_svg_uri()) is None      # openpyxl/docx/pptx no incrustan SVG
    assert marca.bytes_de_uri("") is None
    img = marca.imagen_excel_uri(_png_uri(120, 48), 48)
    assert img is not None and img.height == 48 and img.width == 120   # conserva la proporción 120x48
    assert marca.imagen_excel_uri(_svg_uri(), 48) is None


def test_el_cfg_del_artefacto_lleva_el_logo_solo_si_se_cargo():
    png = _png_uri()
    assert artefacto_html.construir_config({}, {"client": "ACME", "logoCliente": png}, {}).get("clientLogo") == png
    assert "clientLogo" not in artefacto_html.construir_config({}, {"client": "ACME"}, {})


def _reg_planificacion(logo):
    mod = PROCESADORES["planificacion_nia"]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    reg = em._reg(d, mod, ds, par, corte)
    reg["engagement"] = {**(reg.get("engagement") or {}), "logoCliente": logo}
    return d, reg


def _imgs_inicio(xlsx: bytes) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))   # si levantara «reparaciones», fallaría aquí
    return len(getattr(wb["00_Inicio"], "_images", []))


def test_el_excel_suma_el_logo_del_cliente_en_la_portada_sin_romperse():
    d, reg = _reg_planificacion(_png_uri())
    d0, reg0 = _reg_planificacion("")
    con = _imgs_inicio(libro.xlsx(d, reg, [], 1, "APROBADO"))
    sin = _imgs_inicio(libro.xlsx(d0, reg0, [], 1, "APROBADO"))
    assert con == sin + 1, (con, sin)


def test_el_excel_con_logo_svg_no_agrega_imagen_raster():
    d, reg = _reg_planificacion(_svg_uri())
    d0, reg0 = _reg_planificacion("")
    assert _imgs_inicio(libro.xlsx(d, reg, [], 1, "APROBADO")) == _imgs_inicio(libro.xlsx(d0, reg0, [], 1, "APROBADO"))


@pytest.mark.skipif(
    __import__("importlib").util.find_spec("matplotlib") is None, reason="sin matplotlib (gráficos de Word/PPT)")
def test_word_y_ppt_suman_el_logo_del_cliente_en_la_portada():
    from docx import Document
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    d, reg = _reg_planificacion(_png_uri())
    d0, reg0 = _reg_planificacion("")

    def imgs_docx(blob):
        dc = Document(io.BytesIO(blob))
        return sum(1 for p in dc.part.package.iter_parts() if "image" in p.content_type)

    assert imgs_docx(libro.docx(d, reg, [], 1, "APROBADO")) == imgs_docx(libro.docx(d0, reg0, [], 1, "APROBADO")) + 1

    def pics_portada(blob):
        prs = Presentation(io.BytesIO(blob))
        return sum(1 for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)

    assert pics_portada(libro.pptx(d, reg, [], 1, "APROBADO")) == pics_portada(libro.pptx(d0, reg0, [], 1, "APROBADO")) + 1
