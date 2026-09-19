"""V3 — lo que no se pudo medir no se netea entre bandas.

POR QUÉ EXISTE. `motor.medir_ecl` acumulaba la exposición sin tasa CON SIGNO
(`sin_medir += saldo`), así que una banda sin tasa con +20.000 y otra sin tasa
con −20.000 se cancelaban: el módulo declaraba `exposicion.sin_medir = 0,00`,
`medicion_completa = True`, ningún hallazgo, `08-Conciliacion` imprimía
«Exposición SIN MEDIR 0,00» y la pantalla la pintaba en verde. Es la regla del
módulo -lo que no se puede medir se declara, nunca se rellena con cero-
incumplida en el AGREGADO, que es donde nadie la estaba mirando: `05-Matriz`
seguía rotulando las dos bandas SIN MEDIR con su exposición a la vista.

CÓMO SE PRESENTA (la decisión, porque las dos magnitudes son reales):

- **Magnitud** (`exposicion.sin_medir`): deudora + |acreedora|, SIN netear. Es
  cuánta exposición quedó sin medición, y es la que decide `medicion_completa`,
  el hallazgo y el KPI de la pantalla.
- **Neta** (`exposicion.sin_medir_neto`): deudora + acreedora, con signo. Es la
  única que puede restarse de la cartera estratificada, porque esa cartera
  también es NETA. `cartera medida = estratificada − sin medir NETO` sigue
  cuadrando al centavo.

Las dos viven en `08-Conciliacion`, cada una en su fila y con su rótulo, y la
hoja explica por qué no coinciden: `medida + magnitud` NO da la estratificada, y
presentarlo de otra forma obligaría a esconder una de las dos.
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import date

from openpyxl import Workbook, load_workbook

from backend.app.auth import service as auth_service
from backend.app.auth.models import Role
from backend.app.db.session import SessionLocal, init_db
from tests.excel_calc import Libro

BASE = "/api/v1/aud/pce-cxc"
CENTAVO = 0.005


def _xlsx(filas: list[tuple]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


#: La cohorte solo tiene cartera en «Más de 730 días», así que esa es la ÚNICA
#: banda con tasa observada. Cualquier otra banda del corte actual queda sin
#: medir.
COHORTE = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 100000.0)]

#: Corte actual: 100.000 medibles y dos bandas sin tasa que se anulan entre sí
#: (+20.000 y −20.000). Es el escenario que reprodujo la revisión de cierre.
ACTUAL = [
    ("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 100000.0),
    ("BETA", "F-2", "NO-RELACIONADOS", date(2024, 1, 1), date(2024, 6, 1), 20000.0),
    ("GAMMA", "NC-1", "NO-RELACIONADOS", date(2025, 1, 1), date(2025, 4, 1), -20000.0),
]

MAGNITUD_SIN_MEDIR = 40000.0
NETO_SIN_MEDIR = 0.0
ESTRATIFICADA = 100000.0


def _archivos():
    return [("archivos", ("cartera_2023.xlsx", _xlsx(COHORTE))),
            ("archivos", ("cartera_2024.xlsx", _xlsx(COHORTE))),
            ("archivos", ("cartera_2025.xlsx", _xlsx(ACTUAL)))]


def _token(client, role=Role.user):
    init_db()
    tag = uuid.uuid4().hex[:6]
    email, pw = f"pce-sm-{tag}@ex.com", "Sup3rSecret!"
    db = SessionLocal()
    try:
        auth_service.create_user(db, email=email, password=pw, role=role)
    finally:
        db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": pw})
    return r.json()["access_token"]


def _corrida(client, token):
    r = client.post(f"{BASE}/analizar", files=_archivos(),
                    data={"parametros": json.dumps(
                        {"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"],
                         "umbral_dias_incumplimiento": 730})},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    return r.json()


def _libro(client, token, corrida_id):
    x = client.get(f"{BASE}/corridas/{corrida_id}/excel",
                   headers={"Authorization": f"Bearer {token}"})
    assert x.status_code == 200, x.text
    return Libro(load_workbook(io.BytesIO(x.content)))


def _fila_de(ws, texto: str) -> int:
    """Número de la fila de `08-Conciliacion` cuyo concepto contiene `texto`."""
    objetivo = texto.lower()
    for f in range(2, ws.max_row + 1):
        if objetivo in str(ws.cell(f, 1).value or "").lower():
            return f
    conceptos = [str(ws.cell(f, 1).value or "") for f in range(2, ws.max_row + 1)]
    raise AssertionError(f"08-Conciliacion no tiene ninguna fila «{texto}»: {conceptos}")


def test_la_exposicion_sin_medir_se_declara_por_su_magnitud(client):
    """Punta a punta: dos bandas sin tasa que se anulan siguen siendo 40.000
    de cartera sin medición."""
    token = _token(client)
    res = _corrida(client, token)
    assert res["exposicion"]["sin_medir"] == MAGNITUD_SIN_MEDIR, res["exposicion"]
    assert res["medicion_completa"] is False, (
        "el módulo declara que midió toda la cartera con 40.000 sin tasa")


def test_la_parte_neta_de_lo_sin_medir_se_conserva_aparte(client):
    """La neta no desaparece: es la única que puede restarse de una cartera que
    también es neta, y por eso el cuadre sigue en pie."""
    token = _token(client)
    res = _corrida(client, token)
    exposicion = res["exposicion"]
    assert exposicion["sin_medir_neto"] == NETO_SIN_MEDIR, exposicion
    assert exposicion["sin_medir_deudora"] == 20000.0, exposicion
    assert exposicion["sin_medir_acreedora"] == -20000.0, exposicion
    # La cartera medida se deriva de la NETA, así que el cuadre es exacto.
    assert exposicion["medida"] == ESTRATIFICADA - NETO_SIN_MEDIR


def test_el_hallazgo_de_cartera_sin_tasa_se_dispara_sobre_la_magnitud(client):
    """El hallazgo se disparaba sobre la suma neteada, así que en este
    escenario no existía."""
    token = _token(client)
    res = _corrida(client, token)
    hallazgo = next((h for h in res["hallazgos"]
                     if h["titulo"] == "Cartera sin tasa histórica"), None)
    assert hallazgo is not None, [h["titulo"] for h in res["hallazgos"]]
    assert "40,000.00" in hallazgo["condicion"] or "40.000,00" in hallazgo["condicion"], (
        hallazgo["condicion"])


def test_08_conciliacion_dice_la_magnitud_y_la_neta_y_sigue_cuadrando(client):
    """La hoja recalcula las dos desde `05-Matriz` y `06-Individual`, y la
    cartera medida sigue saliendo de la neta al centavo."""
    token = _token(client)
    res = _corrida(client, token)
    libro = _libro(client, token, res["corrida_id"])
    ws = libro.wb["08-Conciliacion"]

    f_estratificada = _fila_de(ws, "Exposición estratificada")
    f_deudora = _fila_de(ws, "SIN MEDIR — deudora")
    f_acreedora = _fila_de(ws, "SIN MEDIR — acreedora")
    f_magnitud = _fila_de(ws, "SIN MEDIR (magnitud")
    f_neta = _fila_de(ws, "SIN MEDIR NETA")
    f_medida = _fila_de(ws, "Cartera medida (")

    estratificada = libro.numero("08-Conciliacion", f"B{f_estratificada}")
    magnitud = libro.numero("08-Conciliacion", f"B{f_magnitud}")
    neta = libro.numero("08-Conciliacion", f"B{f_neta}")
    medida = libro.numero("08-Conciliacion", f"B{f_medida}")

    assert estratificada == ESTRATIFICADA
    assert libro.numero("08-Conciliacion", f"B{f_deudora}") == 20000.0
    assert libro.numero("08-Conciliacion", f"B{f_acreedora}") == -20000.0
    # El papel recalcula, desde sus propias celdas, exactamente lo archivado.
    assert magnitud == res["exposicion"]["sin_medir"]
    assert neta == res["exposicion"]["sin_medir_neto"]
    assert medida == res["exposicion"]["medida"]
    # Y el cuadre: la cartera medida más lo sin medir NETO da la estratificada.
    assert abs((medida + neta) - estratificada) < CENTAVO


def test_08_conciliacion_explica_por_que_magnitud_y_neta_no_coinciden(client):
    """Un papel con dos cifras distintas bajo el mismo concepto tiene que decir
    por qué, o el revisor concluye que una de las dos está mal."""
    token = _token(client)
    res = _corrida(client, token)
    libro = _libro(client, token, res["corrida_id"])
    ws = libro.wb["08-Conciliacion"]
    texto = " ".join(str(c.value or "") for f in ws.iter_rows() for c in f)
    assert "no se netea" in texto.lower() or "sin netear" in texto.lower(), texto[:400]
    assert "magnitud" in texto.lower()


def test_una_corrida_sin_nada_sin_medir_sigue_declarando_medicion_completa(client):
    """La guarda no puede convertir toda corrida en incompleta."""
    token = _token(client)
    r = client.post(f"{BASE}/analizar", files=[
        ("archivos", ("cartera_2023.xlsx", _xlsx(COHORTE))),
        ("archivos", ("cartera_2024.xlsx", _xlsx(COHORTE))),
        ("archivos", ("cartera_2025.xlsx", _xlsx(COHORTE)))],
        data={"parametros": json.dumps(
            {"fechas": ["2023-12-31", "2024-12-31", "2025-12-31"],
             "umbral_dias_incumplimiento": 730})},
        headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["exposicion"]["sin_medir"] == 0.0
    assert res["medicion_completa"] is True
