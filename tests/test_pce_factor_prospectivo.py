"""V5 — el componente prospectivo se llama igual y vale lo mismo en todo el libro.

POR QUÉ EXISTE. La misma magnitud tenía dos nombres y dos convenciones
separadas por 1,0:

- `01-Parametros` A5 decía «Ajuste prospectivo» y su celda guardaba `factor − 1`
  con formato de porcentaje, así que «sin ajuste» se veía como «0,00 %»;
- `05-Matriz` rotulaba su columna «Factor prospectivo» y la fórmula compensaba
  con `(1+E)`;
- la pantalla, el JSDoc, el hallazgo y el 400 hablan de «factor: 1,000 = sin
  ajuste», y 0,000 es justo el valor que el módulo declara error fatal de
  entrada.

Ninguna celda de las trece hojas declaraba la convención, y `01-Parametros` es
una celda que el papel INVITA a editar: escribir 12 ahí, creyendo que es el
factor que se envió desde la pantalla, daba un multiplicador de trece.

Se unifica en FACTOR (1,000 = sin ajuste), que es lo que ya dicen la pantalla,
el parámetro que se guarda con la corrida y los dos mensajes de error, y la
celda editable dice qué significa su valor.
"""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.aud.pce_cxc.exporter import construir_excel
from backend.app.aud.pce_cxc.service import analizar
from tests.excel_calc import Libro, columna, fila

CENTAVO = 0.005


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


#: El documento cae en la MISMA banda en el corte t-2 y en el corte actual
#: («Más de 730 días»), que es lo que hace que la banda del corte actual tenga
#: tasa observada y exposición a la vez: sin las dos, la fórmula del factor no
#: se puede comprobar sobre ninguna cifra.
COHORTE_2023 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 100000.0)]
COHORTE_2024 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 40000.0)]
COHORTE_2025 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 20000.0)]

PARAMETROS = {
    "umbral_dias_incumplimiento": 730,
    "factor_prospectivo": {"NO-RELACIONADOS": 1.10, "RELACIONADOS": 1.10},
    "justificacion_prospectivo": "Previsión de deterioro macroeconómico documentada por el cliente.",
}


def _corrida():
    cortes = [{"nombre": f"cartera_{f.year}.xlsx", "contenido": c, "fecha": f}
              for c, f in zip([_xlsx(COHORTE_2023), _xlsx(COHORTE_2024), _xlsx(COHORTE_2025)],
                              [date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)])]
    return analizar(cortes, dict(PARAMETROS))


def _libro(resultado):
    return Libro(load_workbook(io.BytesIO(construir_excel(resultado, dict(PARAMETROS)))))


def _fila_medida(ws) -> int:
    """Primera fila de `05-Matriz` con tasa medida Y exposición.

    Con exposición cero la fórmula da cero cualquiera que sea el factor, así
    que una fila así no comprueba nada sobre la convención.
    """
    col_perdida = columna(ws, "Pérdida esperada")
    col_exposicion = columna(ws, "Exposición")
    for f in range(2, ws.max_row + 1):
        if str(ws.cell(f, 2).value or "") in ("", "TOTAL"):
            break
        perdida = ws[f"{col_perdida}{f}"].value
        exposicion = ws[f"{col_exposicion}{f}"].value
        if (isinstance(perdida, str) and perdida.startswith("=")
                and isinstance(exposicion, (int, float)) and abs(exposicion) > CENTAVO):
            return f
    raise AssertionError("05-Matriz no tiene ninguna banda medida con exposición")


def test_el_libro_usa_una_sola_convencion_y_es_la_de_la_pantalla():
    """Un solo nombre y un solo valor: FACTOR, donde 1,000 es «sin ajuste»."""
    resultado = _corrida()
    libro = _libro(resultado)
    parametros = libro.wb["01-Parametros"]
    f = fila(parametros, "prospectivo — NO-RELACIONADOS")

    assert "Factor prospectivo" in str(parametros.cell(f, 1).value), parametros.cell(f, 1).value
    # El valor es el FACTOR, no el factor menos uno.
    assert parametros.cell(f, 2).value == pytest.approx(1.10, abs=1e-9), (
        f"01-Parametros guarda {parametros.cell(f, 2).value}: el papel y la pantalla usan dos "
        f"convenciones distintas para la misma magnitud")
    # Y no se presenta como porcentaje, que es lo que hacía leer «0,00 %» donde
    # la pantalla dice «1,000».
    assert "%" not in str(parametros.cell(f, 2).number_format)


def test_la_celda_editable_declara_que_significa_su_valor():
    """El papel invita a editarla: tiene que decir qué es 1,000."""
    resultado = _corrida()
    libro = _libro(resultado)
    parametros = libro.wb["01-Parametros"]
    f = fila(parametros, "prospectivo — NO-RELACIONADOS")
    fuente = str(parametros.cell(f, 3).value or "")
    assert "1,000" in fuente, fuente
    assert "sin ajuste" in fuente.lower(), fuente


def test_05_matriz_llama_a_la_columna_lo_mismo_y_guarda_lo_mismo():
    """La columna de `05-Matriz` y la celda de `01-Parametros` son la misma
    magnitud: tienen que dar la misma cifra."""
    resultado = _corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    col = columna(ws, "Factor prospectivo")
    f = _fila_medida(ws)
    assert libro.numero("05-Matriz", f"{col}{f}") == pytest.approx(1.10, abs=1e-9)


def test_poner_1_en_la_celda_del_factor_deja_la_banda_sin_ajuste():
    """La prueba de que la convención es una sola.

    Con la convención anterior, escribir 1,000 en la celda -el valor que la
    pantalla llama «sin ajuste»- daba un multiplicador de DOS. Ahora 1,000 es
    lo que dice que es: la tasa observada, sin tocar.
    """
    resultado = _corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    f = _fila_medida(ws)
    col_exposicion = columna(ws, "Exposición")
    col_tasa = columna(ws, "Tasa aplicada")
    col_perdida = columna(ws, "Pérdida esperada")

    parametros = libro.wb["01-Parametros"]
    destino = fila(parametros, "prospectivo — NO-RELACIONADOS")
    parametros[f"B{destino}"] = 1.0
    parametros[f"B{fila(parametros, 'prospectivo — RELACIONADOS')}"] = 1.0
    libro = Libro(libro.wb)

    exposicion = libro.numero("05-Matriz", f"{col_exposicion}{f}")
    tasa = libro.numero("05-Matriz", f"{col_tasa}{f}")
    perdida = libro.numero("05-Matriz", f"{col_perdida}{f}")
    assert perdida == pytest.approx(round(exposicion * tasa, 2), abs=CENTAVO), (
        "con el factor en 1,000 la banda tendría que valer exposición × tasa observada")


def test_la_banda_recalculada_sigue_dando_lo_archivado_con_factor_1_10():
    """El papel se recalcula desde sus propias celdas y da al centavo lo que
    archivó la corrida, también con la convención nueva."""
    resultado = _corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    f = _fila_medida(ws)
    segmento, banda = ws.cell(f, 1).value, ws.cell(f, 2).value
    archivado = next(t for t in resultado["matriz"]["tramos"]
                     if (t.get("segmento"), t["tramo"]) == (segmento, banda))
    col_perdida = columna(ws, "Pérdida esperada")
    assert libro.numero("05-Matriz", f"{col_perdida}{f}") == pytest.approx(
        archivado["ecl"], abs=CENTAVO)
