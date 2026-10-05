"""Planificación NIA con la carta de control interno subida en PDF: la IA extrae la
tabla (función de chat falsa; en producción usa el servidor de IA local primero), el
auditor la confirma y map_validate la suma al dataset `carta_control_interno`, de modo
que la matriz de riesgos del papel la usa. Flujo completo por HTTP; la extracción se
monkeypatchea para no depender de la red."""
import io
import json

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif.ciclo import extraccion_ia
from backend.app.aud.niif.procesadores import planificacion_nia as m
from tests.test_aud_ciclo_evidencia import _leer, _subir
from tests.test_aud_ciclo_http import BASE, FICHA, _accion, _staff_con_proyecto
from tests.test_aud_niif_fichas import _h


@pytest.fixture(autouse=True)
def disco_temporal(tmp_path, monkeypatch):
    monkeypatch.setenv("AUD_CICLO_DIR", str(tmp_path / "aud_pruebas"))
    monkeypatch.setenv("AUD_CICLO_MIN_LIBRE_MB", "0")


MARCA = "MARCA_EXTRAIDA_IA_XYZ"
CARTA_IA = [
    {"id": "IA1", "proceso": "Inventarios", "hallazgo": f"{MARCA}: no hay tomas físicas periódicas",
     "aseveraciones": "Existencia", "probabilidad": 4, "impacto": 4, "control": 2,
     "respuesta": "Observar la toma física al cierre.", "probar_control": "No"},
    {"id": "IA2", "proceso": "Tesorería", "hallazgo": "Conciliaciones bancarias sin revisar",
     "aseveraciones": "", "probabilidad": 3, "impacto": 3, "control": 3, "respuesta": "", "probar_control": ""},
]


class _Resp:
    def __init__(self, content):
        self.content = content
        self.model = "modelo-falso"


def _chat_falso(filas):
    def chat(messages, system=None):
        return _Resp(json.dumps({"filas": filas}))
    return chat


def _mapa(tipo):
    return {c["key"]: i for i, c in enumerate(m.CAMPOS[tipo])}


def _modelo_lleno(client, tok, p, req, filas):
    r = client.get(f"{BASE}/pruebas/{p['id']}/modelo/{req}", headers=_h(tok))
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    for f in filas:
        wb["Datos"].append(f)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _filas_balance(clave, keys):
    return [[f.get(k, "") for k in keys] for f in m.EJEMPLO["datasets"][clave]]


def _hasta_requerimiento(client, tok, pid):
    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    p = client.post(f"{BASE}/proyectos/{pid}/pruebas", headers=_h(tok),
                    json={"origen": "proc:planificacion_nia"}).json()
    p = _accion(client, tok, p, "research").json()
    p = _accion(client, tok, p, "generate_program").json()
    prog = p["registro"]["program"]
    codigos = [x["code"] for x in prog]
    fuentes = p["registro"]["sources"]
    for s in fuentes:
        s.update(verified=True, section="párr. aplicable", date="vigente", procedures=codigos)
    p = _accion(client, tok, p, "approve_program", {"program": prog, "sources": fuentes}).json()
    assert p["estado"] == "PROGRAMA_APROBADO", p
    p = _accion(client, tok, p, "generate_request").json()
    p = _accion(client, tok, p, "approve_request", {"requests": p["registro"]["requests"]}).json()
    assert p["estado"] == "REQUERIMIENTO_APROBADO", p
    return p


def test_carta_en_pdf_se_extrae_confirma_y_alimenta_la_matriz(client, monkeypatch):
    # La IA no se llama de verdad: texto fijo + función de chat falsa (la cadena de
    # proveedores, local primero, se sustituye por esta).
    monkeypatch.setattr(extraccion_ia, "texto_de_documento", lambda nombre, datos: "Texto de la carta de control interno.")
    monkeypatch.setattr(extraccion_ia, "_chat_por_defecto", lambda: _chat_falso(CARTA_IA))

    tok, pid = _staff_con_proyecto(client)
    p = _hasta_requerimiento(client, tok, pid)

    # Balances (xlsx) por el camino normal; la carta como PDF firmado. Cada subida
    # sube la revisión, así que se relee `p` antes de la siguiente.
    subidas = [
        ("RQ-001", "RQ-001.xlsx", _modelo_lleno(client, tok, p, "RQ-001",
            _filas_balance("balance_anterior", ["codigo", "cuenta", "saldo_anterior"]))),
        ("RQ-002", "RQ-002.xlsx", _modelo_lleno(client, tok, p, "RQ-002",
            _filas_balance("balance_actual", ["codigo", "cuenta", "saldo_actual"]))),
        ("RQ-004", "carta_control_interno.pdf", b"%PDF-1.4 carta"),
    ]
    for req, nombre, contenido in subidas:
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, nombre, contenido).status_code == 201

    p = _leer(client, tok, p)
    arch = {a["requerimiento"]: a["id"] for a in p["archivos"]}
    carta_id = arch["RQ-004"]

    # 1) La IA extrae la tabla de la carta.
    p = _accion(client, tok, p, "extraer_ia", {"fileId": carta_id}).json()
    extr = p["registro"]["extraccion"][str(carta_id)]
    assert extr["dataset"] == "carta_control_interno" and len(extr["rows"]) == 2
    assert extr["rows"][0]["id"] == "IA1" and extr["validation"]["ok"]

    # 2) El auditor revisa y confirma (aquí sin editar).
    p = _accion(client, tok, p, "guardar_extraccion", {"fileId": carta_id, "rows": extr["rows"]}).json()
    assert p["registro"]["extraccion"][str(carta_id)]["validation"]["ok"]

    # 3) map_validate suma las filas de IA al dataset (la carta NO viaja en `datasets`,
    #    porque en la vista el PDF no se lee como hoja de cálculo).
    parte = lambda req, tipo: [{"fileId": arch[req], "sheet": "Datos", "header": 1, "mapping": _mapa(tipo)}]
    p = _accion(client, tok, p, "map_validate", {"datasets": {
        "balance_anterior": parte("RQ-001", "balance_anterior"),
        "balance_actual": parte("RQ-002", "balance_actual")}}).json()
    val = p["registro"]["validation"]
    assert val["ok"], val
    carta = p["registro"]["datasets"]["carta_control_interno"]
    assert [f["id"] for f in carta] == ["IA1", "IA2"]

    # 4) Produce y comprueba que la marca única de la carta extraída llega al papel.
    ct = p["registro"]["controlTotal"]
    p = _accion(client, tok, p, "validate", {"evidenceReviewed": True, "evidenceReview": "Carta extraída y confirmada.",
                                             "ledger": ct, "tolerance": "0"}).json()
    assert p["estado"] == "DOCUMENTACION_VALIDADA", p
    p = _accion(client, tok, p, "configure", {"basis": "Parámetros de la planificación (caso de prueba).",
                                              "parametros": m.EJEMPLO["parametros"]}).json()
    p = _accion(client, tok, p, "approve_methodology").json()
    p = _accion(client, tok, p, "execute", {}).json()
    assert p["estado"] == "PRUEBA_EJECUTADA", p
    assert MARCA in json.dumps(p["registro"]["run"], ensure_ascii=False)


def test_carta_en_pdf_se_extrae_sola_al_procesar(client, monkeypatch):
    """Auto-extracción al procesar: se sube la carta en PDF y NO se pulsa «Extraer con
    IA»; al hacer map_validate (Procesar), la IA la extrae sola y alimenta la matriz,
    con un aviso de revisar."""
    monkeypatch.setattr(extraccion_ia, "texto_de_documento", lambda nombre, datos: "Texto de la carta.")
    monkeypatch.setattr(extraccion_ia, "_chat_por_defecto", lambda: _chat_falso(CARTA_IA))

    tok, pid = _staff_con_proyecto(client)
    p = _hasta_requerimiento(client, tok, pid)
    subidas = [
        ("RQ-001", "RQ-001.xlsx", _modelo_lleno(client, tok, p, "RQ-001",
            _filas_balance("balance_anterior", ["codigo", "cuenta", "saldo_anterior"]))),
        ("RQ-002", "RQ-002.xlsx", _modelo_lleno(client, tok, p, "RQ-002",
            _filas_balance("balance_actual", ["codigo", "cuenta", "saldo_actual"]))),
        ("RQ-004", "carta_control_interno.pdf", b"%PDF-1.4 carta"),
    ]
    for req, nombre, contenido in subidas:
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, nombre, contenido).status_code == 201

    p = _leer(client, tok, p)
    arch = {a["requerimiento"]: a["id"] for a in p["archivos"]}
    # NO se llama extraer_ia: solo se procesa (map_validate) con los balances.
    parte = lambda req, tipo: [{"fileId": arch[req], "sheet": "Datos", "header": 1, "mapping": _mapa(tipo)}]
    p = _accion(client, tok, p, "map_validate", {"datasets": {
        "balance_anterior": parte("RQ-001", "balance_anterior"),
        "balance_actual": parte("RQ-002", "balance_actual")}}).json()
    val = p["registro"]["validation"]
    assert val["ok"], val
    # La IA extrajo la carta sola y la sumó al dataset.
    carta = p["registro"]["datasets"]["carta_control_interno"]
    assert [f["id"] for f in carta] == ["IA1", "IA2"]
    # Quedó marcada como automática y con aviso de revisar.
    entrada = next(iter(p["registro"]["extraccion"].values()))
    assert entrada["auto"] is True and entrada["revisado"] is False
    assert any("extraídas por IA" in w["message"] for w in val["warnings"])


NOTAS_IA = [
    {"nota": "13", "titulo": "Beneficios a empleados", "codigos": "2103", "saldo_auditado": "222400.00"},
]


def _xlsx_notas():
    """Un .xlsx de notas con encabezados que el reconocimiento por alias NO cubre (como
    las notas firmadas del cliente), para forzar el camino de lectura por IA."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Detalle de las notas a los estados financieros 2024"])
    ws.append(["13. Beneficios a empleados (jubilación patronal)", "222,400.00"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_notas_en_xlsx_se_extraen_solas_al_procesar(client, monkeypatch):
    """Regresión: una nota a los estados financieros subida en .xlsx ya NO queda
    atascada en el mapeo manual de columnas. Al Procesar, la IA la lee sola (igual que
    un PDF/Word) y la suma al dataset `notas_estados_financieros`, de modo que alimenta
    la conciliación con el balance y los riesgos de la matriz (NIA 510)."""
    monkeypatch.setattr(extraccion_ia, "texto_de_documento",
                        lambda nombre, datos: "Nota 13 | Beneficios a empleados | 2103 | 222400")
    monkeypatch.setattr(extraccion_ia, "_chat_por_defecto", lambda: _chat_falso(NOTAS_IA))

    tok, pid = _staff_con_proyecto(client)
    p = _hasta_requerimiento(client, tok, pid)
    req_notas = next(r["id"] for r in p["registro"]["requests"] if r.get("dataset") == "notas_estados_financieros")

    subidas = [
        ("RQ-001", "RQ-001.xlsx", _modelo_lleno(client, tok, p, "RQ-001",
            _filas_balance("balance_anterior", ["codigo", "cuenta", "saldo_anterior"]))),
        ("RQ-002", "RQ-002.xlsx", _modelo_lleno(client, tok, p, "RQ-002",
            _filas_balance("balance_actual", ["codigo", "cuenta", "saldo_actual"]))),
        (req_notas, "NOTAS LANSEY 2025.xlsx", _xlsx_notas()),
    ]
    for req, nombre, contenido in subidas:
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, nombre, contenido).status_code == 201

    p = _leer(client, tok, p)
    arch = {a["requerimiento"]: a["id"] for a in p["archivos"]}
    parte = lambda req, tipo: [{"fileId": arch[req], "sheet": "Datos", "header": 1, "mapping": _mapa(tipo)}]
    # Se procesa SOLO con los balances; la nota .xlsx NO viaja en `datasets` (no se mapeó a mano).
    p = _accion(client, tok, p, "map_validate", {"datasets": {
        "balance_anterior": parte("RQ-001", "balance_anterior"),
        "balance_actual": parte("RQ-002", "balance_actual")}}).json()
    val = p["registro"]["validation"]
    assert val["ok"], val
    # La IA leyó la nota .xlsx sola y la sumó al dataset (sin pedir mapeo manual).
    notas = p["registro"]["datasets"]["notas_estados_financieros"]
    assert [f["nota"] for f in notas] == ["13"]
    entrada = p["registro"]["extraccion"][str(arch[req_notas])]
    assert entrada["auto"] is True and entrada["revisado"] is False
    assert any("extraídas por IA" in w["message"] for w in val["warnings"])
