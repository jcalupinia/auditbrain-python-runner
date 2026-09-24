"""Pruebas de la hoja PARÁMETROS del papel de trabajo ICT (P2-F, REP-005)."""
import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from backend.app.ict.parametros_sheet import SHEET_NAME, build_parametros_sheet

PARAMS = {
    "ejercicio_inicio": "2025-01-01", "ejercicio_fin": "2025-12-31",
    "materialidad": "50000.00", "materialidad_ejecucion": "37500.00",
    "umbral_insignificante": "2500.00", "umbral_aprobacion": "5000.00",
    "error_tolerable": "37500.00", "feriados": ["2025-05-01", "2025-08-10"],
    "hora_inicio": 8, "hora_fin": 18, "confianza": 95, "semilla": 20250101,
    "fecha_registro_es_contable": False,
}
SESSION = {"razon_social": "PROPHAR S.A.", "ruc": "1791859596001",
           "ejercicio_fiscal": "2025"}


def test_escribe_los_parametros_presentes():
    wb = Workbook()
    r = build_parametros_sheet(wb, PARAMS, session_data=SESSION)
    assert r.parametros_escritos == 13
    assert SHEET_NAME in wb.sheetnames


def test_materialidad_es_decimal_con_formato():
    wb = Workbook()
    build_parametros_sheet(wb, PARAMS, session_data=SESSION)
    ws = wb[SHEET_NAME]
    montos = [c for row in ws.iter_rows() for c in row
              if isinstance(c.value, Decimal) and c.value == Decimal("50000.00")]
    assert montos and montos[0].number_format == "#,##0.00"


def test_texto_que_parece_formula_se_escapa():
    wb = Workbook()
    inyeccion = {**PARAMS, "feriados": ["=CMD()", "2025-05-01"]}
    build_parametros_sheet(wb, inyeccion, session_data=SESSION)
    ws = wb[SHEET_NAME]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not any("CMD" in str(f) for f in formulas)


def test_parametro_faltante_no_rompe_y_advierte():
    wb = Workbook()
    incompletos = {k: v for k, v in PARAMS.items() if k != "semilla"}
    r = build_parametros_sheet(wb, incompletos, session_data=SESSION)
    assert any("semilla" in a for a in r.advertencias)
    assert r.parametros_escritos == 12
    ws = wb[SHEET_NAME]
    textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("(no informado)" in t for t in textos)


def test_no_corrompe_el_libro():
    wb = Workbook()
    build_parametros_sheet(wb, PARAMS, session_data=SESSION)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    assert SHEET_NAME in load_workbook(buf).sheetnames
