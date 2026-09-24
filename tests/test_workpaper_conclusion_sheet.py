"""Pruebas de la hoja CONCLUSIÓN del papel de trabajo ICT (P2-F, REP-007)."""
import io

from openpyxl import Workbook, load_workbook

from backend.app.ict.conclusion_sheet import (
    DISCLAIMER_IA,
    SHEET_NAME,
    build_conclusion_sheet,
)

CTX = {
    "sintesis": "El ejercicio cuadra; 2 excepciones P0.",
    "total_excepciones": 2, "monto_total_excepciones": "13300.00",
    "cuadra_a1": True, "suficiencia_estado": "Parcial",
    "reglas_no_corridas": ["AST-001", "AST-002"],
}


def _textos(wb):
    return [str(c.value) for row in wb[SHEET_NAME].iter_rows() for c in row if c.value]


def test_escribe_sintesis_y_firma_en_blanco():
    wb = Workbook()
    r = build_conclusion_sheet(wb, CTX, session_data={"ruc": "1791859596001"})
    textos = _textos(wb)
    assert any("El ejercicio cuadra" in t for t in textos)
    assert any("Preparó" in t for t in textos)
    assert any("Revisó" in t for t in textos)
    assert any("Fecha" in t for t in textos)
    assert r.tiene_disclaimer_ia is False


def test_disclaimer_ia_cuando_la_sintesis_es_ia():
    wb = Workbook()
    ctx = {**CTX, "interpretacion": {"sintesis": CTX["sintesis"],
           "confianza_modelo": "baja", "requiere_revision_humana": True}}
    r = build_conclusion_sheet(wb, ctx)
    assert r.tiene_disclaimer_ia
    assert r.confianza_modelo == "baja"
    assert r.requiere_revision_humana is True
    textos = _textos(wb)
    assert any(DISCLAIMER_IA[:20] in t for t in textos)
    assert any("Revisar manualmente" in t for t in textos)


def test_sin_ia_no_hay_disclaimer():
    wb = Workbook()
    r = build_conclusion_sheet(wb, CTX)
    assert not r.tiene_disclaimer_ia
    assert all(DISCLAIMER_IA[:20] not in t for t in _textos(wb))


def test_texto_que_parece_formula_se_escapa():
    wb = Workbook()
    ctx = {**CTX, "sintesis": "=SUMA(A1) resultado peligroso"}
    build_conclusion_sheet(wb, ctx)
    formulas = [c.value for row in wb[SHEET_NAME].iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not any("SUMA" in str(f) for f in formulas)


def test_no_corrompe_el_libro():
    wb = Workbook()
    build_conclusion_sheet(wb, CTX)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    assert SHEET_NAME in load_workbook(buf).sheetnames
