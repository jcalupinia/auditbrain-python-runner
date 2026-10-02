# -*- coding: utf-8 -*-
"""Papel formulado de Caja y Bancos: estructura, fórmulas vivas y cuadre por SUMIFS."""
import io

import openpyxl

from backend.app.aud.niif.procesadores import caja_bancos_papel as m

PAPEL = {
    "engagement": {"client": "LANSEY S.A", "period": "31 de agosto de 2026",
                   "cutoff": "2026-08-31", "preparer": "WR", "reviewer": "AT"},
    "cuentas": [
        # extracto - cheque(120) + NC(0) da 5.000 = saldo actual s/registros -> diferencia 0
        {"cuenta": "11010201", "descripcion": "BANCO PICHINCHA", "saldo_anterior": 4000,
         "saldo_actual": 5000, "ncuenta": "3033617504", "extracto": 5120},
        {"cuenta": "11010202", "descripcion": "BANCO PRODUBANCO", "saldo_anterior": 2000,
         "saldo_actual": 2000, "ncuenta": "1005002670", "extracto": 2000},
    ],
    "partidas": [
        {"fecha": "2026-08-20", "banco": "Pichincha", "documento": "Cheque 123",
         "beneficiario": "Proveedor X", "valor": 120, "observacion": "En tránsito"},
    ],
}


def _recalc(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=False)


def test_estructura_y_hojas():
    wb = _recalc(m.construir(PAPEL))
    assert wb.sheetnames == ["Sumaria", "Movimiento", "Partidas conciliatorias",
                             "Conciliaciones Bancos", "Arqueo Caja", "Hallazgos", "Libro Mayor"]
    # No debe existir la cédula de Procedimientos (el programa vive en la app).
    assert "Procedimientos" not in wb.sheetnames


def test_formulas_vivas_conciliacion():
    wb = _recalc(m.construir(PAPEL))
    cb = wb["Conciliaciones Bancos"]
    # Saldo s/auditoría con Sobregiro incluido en la fórmula.
    assert cb["I9"].value == "=C9+D9-E9-F9-G9-H9"
    # Cheques por SUMIFS a DA-4 con la categoría deducida.
    assert 'SUMIFS' in cb["F9"].value and "Cheque sin cobrar" in cb["F9"].value
    # Sobregiro/ajustes es su propia columna con su SUMIFS.
    assert "Sobregiro/ajuste" in cb["E9"].value
    # Saldo s/registros jala de la Sumaria; diferencia = auditoría - registros.
    assert cb["J9"].value == "=Sumaria!E9"
    assert cb["K9"].value == "=I9-J9"
    # Sumaria: variación = actual - anterior.
    assert wb["Sumaria"]["D9"].value == "=E9-C9"
    # Partidas: días vencidos y prescripción.
    pc = wb["Partidas conciliatorias"]
    assert pc["G10"].value == "=$B$7-A10"
    assert pc["H10"].value == "=A10+390"


def test_clasificar_categorias():
    assert m.clasificar("Cheque 001") == m.CHEQUE
    assert m.clasificar("Nota de crédito", "cliente") == m.NC_PENDIENTE
    assert m.clasificar("Nota de débito") == m.ND_TRANSITO
    assert m.clasificar("Depósito en tránsito") == m.CONSIGNACION
    assert m.clasificar("Sobregiro bancario") == m.SOBREGIRO
    assert m.clasificar("Transferencia rara") == m.OTRA


def test_cuadre_a_cero_por_logica_de_formulas():
    """El caso está armado para que la diferencia de Pichincha sea 0:
    extracto 5.120 − cheque 120 = 5.000 = saldo s/registros. Reproduce la lógica
    de las fórmulas SUMIFS + saldo s/auditoría con los mismos datos del papel."""
    cta = PAPEL["cuentas"][0]  # BANCO PICHINCHA
    cheques = sum(
        p["valor"] for p in PAPEL["partidas"]
        if m.clasificar(p["documento"], p.get("beneficiario", "")) == m.CHEQUE
    )
    saldo_auditoria = cta["extracto"] - cheques
    assert saldo_auditoria == cta["saldo_actual"]  # diferencia = 0
