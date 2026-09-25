"""Puente de contrato motor → hoja EXCEPCIONES del papel de trabajo (P2-F Task 9).

Verifica que `spec_motor_a_espec` traduce la salida real del motor
(`especificar_hoja_excepciones`: columnas planas + filas + resumen) al
`EspecHojaExcepciones` del runner, y que el resultado alimenta sin fricción a
`build_exceptions_sheet` (prueba de extremo a extremo del puente)."""
import io

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.ict.exceptions_sheet import SHEET_NAME, build_exceptions_sheet
from backend.app.ict.motor_adapter import (
    TITULO_POR_DEFECTO,
    spec_motor_a_espec,
)

# Forma EXACTA que hoy devuelve motor.exportar_excepciones.especificar_hoja_excepciones:
# columnas como nombres planos, filas por clave, y resumen con totales por severidad.
SPEC_MOTOR = {
    "hoja": "EXCEPCIONES",
    "columnas": ["regla_id", "nia", "entidad_tipo", "entidad_id",
                 "monto_expuesto", "severidad", "mensaje", "evidencia", "hash"],
    "filas": [
        {"regla_id": "AST-007", "nia": "NIA 240", "entidad_tipo": "asiento",
         "entidad_id": "A1", "monto_expuesto": "12500.00", "severidad": "P0",
         "mensaje": "Gasto no deducible", "evidencia": {"cuenta": "5.01"},
         "hash": "a" * 64},
        {"regla_id": "CMP-003", "nia": "NIA 500", "entidad_tipo": "compra",
         "entidad_id": "C9", "monto_expuesto": "800.00", "severidad": "P2",
         "mensaje": "=OJO diferencia", "evidencia": {"cuenta": "6.02"},
         "hash": "b" * 64},
    ],
    "resumen": {"total_excepciones": 2, "monto_total": "13300.00",
                "por_severidad": {"P0": 1, "P2": 1},
                "monto_por_severidad": {"P0": "12500.00", "P2": "800.00"}},
}


def test_enriquece_columnas_planas_a_objetos():
    espec = spec_motor_a_espec(SPEC_MOTOR, engine_version="motor-1.4.0", run_id="run_9")
    claves = [c["clave"] for c in espec["columnas"]]
    assert claves == SPEC_MOTOR["columnas"]                       # orden preservado
    por_clave = {c["clave"]: c for c in espec["columnas"]}
    assert por_clave["monto_expuesto"]["tipo"] == "monto"
    assert por_clave["severidad"]["tipo"] == "severidad"
    assert por_clave["hash"]["tipo"] == "hash"
    assert por_clave["nia"]["titulo"] == "NIA"


def test_filas_y_resumen_pasan_sin_recalcular():
    espec = spec_motor_a_espec(SPEC_MOTOR)
    assert espec["filas"] == SPEC_MOTOR["filas"]                  # intactas
    assert espec["resumen"] == SPEC_MOTOR["resumen"]             # el motor ya lo calculó
    assert espec["titulo"] == TITULO_POR_DEFECTO


def test_metadatos_de_trazabilidad():
    espec = spec_motor_a_espec(SPEC_MOTOR, engine_version="motor-1.4.0",
                               run_id="run_9", generado_en="2026-09-25T10:00:00Z")
    assert espec["engine_version"] == "motor-1.4.0"
    assert espec["run_id"] == "run_9"
    assert espec["generado_en"] == "2026-09-25T10:00:00Z"


def test_columna_desconocida_cae_a_texto_sin_romper():
    spec = {**SPEC_MOTOR, "columnas": SPEC_MOTOR["columnas"] + ["campo_nuevo"]}
    espec = spec_motor_a_espec(spec)
    nueva = next(c for c in espec["columnas"] if c["clave"] == "campo_nuevo")
    assert nueva["tipo"] == "texto" and nueva["titulo"] == "Campo nuevo"


def test_spec_incompleta_lanza_valueerror():
    for malo in ({"filas": []}, {"columnas": []}, "no-dict"):
        with pytest.raises(ValueError):
            spec_motor_a_espec(malo)


def test_end_to_end_alimenta_build_exceptions_sheet():
    """El puente produce un EspecHojaExcepciones que la capa de presentación
    vuelca a Excel sin fricción (y el libro no se corrompe)."""
    espec = spec_motor_a_espec(SPEC_MOTOR, engine_version="motor-1.4.0", run_id="run_9")
    wb = Workbook()
    r = build_exceptions_sheet(wb, espec, session_data={"razon_social": "PROPHAR S.A.",
                                                        "ruc": "1791859596001",
                                                        "ejercicio_fiscal": "2025"})
    assert r.filas_escritas == 2
    assert r.monto_total == "13300.00"
    # roundtrip: el .xlsx no levanta reparación (la hoja persiste)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = load_workbook(buf)
    assert SHEET_NAME in wb2.sheetnames
    # el texto que parece fórmula no quedó como fórmula
    ws = wb2[SHEET_NAME]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not any("OJO" in str(f) for f in formulas)
