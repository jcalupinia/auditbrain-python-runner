"""Pruebas de la hoja EXCEPCIONES del papel de trabajo ICT (P2-F, REP-006)."""
import io

from openpyxl import Workbook, load_workbook

from backend.app.ict.exceptions_sheet import (
    SHEET_NAME,
    ExceptionsSheetResult,
    build_exceptions_sheet,
)

ESPEC = {
    "titulo": "Excepciones detectadas",
    "columnas": [
        {"clave": "hash", "titulo": "Hash", "tipo": "hash"},
        {"clave": "anexo", "titulo": "Anexo", "tipo": "texto"},
        {"clave": "nia", "titulo": "NIA", "tipo": "nia"},
        {"clave": "monto", "titulo": "Monto", "tipo": "monto"},
        {"clave": "severidad", "titulo": "Severidad", "tipo": "severidad"},
        {"clave": "mensaje", "titulo": "Detalle", "tipo": "texto"},
    ],
    "filas": [
        {"hash": "a1b2c3d4", "anexo": "A5", "nia": "NIA 240",
         "monto": "12500.00", "severidad": "P0", "mensaje": "Gasto no deducible"},
        {"hash": "e5f6a7b8", "anexo": "A2", "nia": "NIA 500",
         "monto": "800.00", "severidad": "P2", "mensaje": "=OJO diferencia"},
    ],
    "resumen": {"total_excepciones": 2, "monto_total": "13300.00",
                "por_severidad": {"P0": 1, "P2": 1},
                "monto_por_severidad": {"P0": "12500.00", "P2": "800.00"}},
    "engine_version": "motor-1.4.0", "generado_en": "2026-09-24T10:00:00Z",
    "run_id": "run_123",
}
SESSION = {"razon_social": "PROPHAR S.A.", "ruc": "1791859596001",
           "ejercicio_fiscal": "2025"}


def _roundtrip(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf)


def test_crea_hoja_con_una_fila_por_excepcion():
    wb = Workbook()
    r = build_exceptions_sheet(wb, ESPEC, session_data=SESSION)
    assert isinstance(r, ExceptionsSheetResult)
    assert r.filas_escritas == 2
    assert SHEET_NAME in wb.sheetnames
    wb2 = _roundtrip(wb)  # regla suprema: no debe corromperse
    assert SHEET_NAME in wb2.sheetnames


def test_encabezados_correctos_desde_columnas():
    wb = Workbook()
    build_exceptions_sheet(wb, ESPEC, session_data=SESSION)
    ws = wb[SHEET_NAME]
    textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    for titulo in ("Hash", "Anexo", "NIA", "Monto", "Severidad", "Detalle"):
        assert any(titulo == t for t in textos), titulo


def test_fila_total_desde_el_resumen():
    wb = Workbook()
    r = build_exceptions_sheet(wb, ESPEC, session_data=SESSION)
    assert r.monto_total == "13300.00"
    assert r.por_severidad == {"P0": 1, "P2": 1}
    ws = wb[SHEET_NAME]
    textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("TOTAL" in t for t in textos)


def test_texto_que_parece_formula_se_escapa():
    """La celda "=OJO diferencia" NO debe quedar como fórmula (inyección)."""
    wb = Workbook()
    build_exceptions_sheet(wb, ESPEC, session_data=SESSION)
    ws = wb[SHEET_NAME]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not any("OJO" in str(f) for f in formulas)
    ojo = [c.value for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and "OJO" in c.value]
    assert ojo and all(v.startswith("'=") for v in ojo)


def test_montos_con_formato_y_como_decimal():
    from decimal import Decimal
    wb = Workbook()
    build_exceptions_sheet(wb, ESPEC, session_data=SESSION)
    ws = wb[SHEET_NAME]
    montos = [c for row in ws.iter_rows() for c in row
              if isinstance(c.value, Decimal) and c.value == Decimal("12500.00")]
    assert montos
    assert montos[0].number_format == "#,##0.00"


def test_espec_vacia_genera_hoja_con_mensaje():
    wb = Workbook()
    vacia = {**ESPEC, "filas": [], "resumen": {"total_excepciones": 0,
             "monto_total": "0.00", "por_severidad": {}}}
    r = build_exceptions_sheet(wb, vacia, session_data=SESSION)
    assert r.filas_escritas == 0
    assert SHEET_NAME in wb.sheetnames
    textos = [str(c.value) for row in wb[SHEET_NAME].iter_rows() for c in row if c.value]
    assert any("Sin excepciones" in t for t in textos)
    _roundtrip(wb)  # no corrompe
