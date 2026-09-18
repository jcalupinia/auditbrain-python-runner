"""Excel del papel de trabajo: hojas, fórmulas y cuadres."""
import io
import re
import zipfile
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from backend.app.aud.pce_cxc import exporter
from backend.app.aud.pce_cxc.exporter import construir_excel
from tests.excel_calc import Libro


class _FechaFija(date):
    """``date`` con un ``today()`` congelado: simula descargar la misma corrida
    en dos días distintos."""

    HOY = "2030-01-01"

    @classmethod
    def today(cls):
        return date(2030, 1, 1)

RESULTADO = {
    "exposicion": {"colectiva": 100000.0, "individual": 0.0, "sin_estratificar": 0.0,
                   "total": 100000.0, "segun_archivo": 100000.0,
                   "factores_anclaje": {"NO-RELACIONADOS": 1.0, "RELACIONADOS": 1.0}},
    "tasas": {"NO-RELACIONADOS": {"Por vencer": 0.01, "0 a 30 días": 0.05}},
    "detalle_cohorte": {"NO-RELACIONADOS": {"Por vencer": {"inicial": 50000.0, "remanente": 500.0, "documentos": 12}}},
    "trazabilidad": 0.99,
    "matriz": {"tramos": [{"tramo": "Por vencer", "exposicion": 80000.0, "tasa_perdida": 0.01,
                           "tasa_ajustada": 0.01, "lgd": 1.0, "horizonte": None,
                           "factor_descuento": 1.0, "ecl": 800.0},
                          {"tramo": "0 a 30 días", "exposicion": 20000.0, "tasa_perdida": 0.05,
                           "tasa_ajustada": 0.05, "lgd": 1.0, "horizonte": None,
                           "factor_descuento": 1.0, "ecl": 1000.0}],
               "exposicion_total": 100000.0, "ecl_total": 1800.0, "exposicion_sin_medir": 0.0,
               "descuento_aplicado": False, "ajuste_prospectivo": 0.0, "justificacion_ajuste": "",
               "fuente_tasas": "Permanencia a 24 meses"},
    "individual": {"casos": [], "saldo_total": 0.0, "ecl_total": 0.0},
    "conciliacion": {"cartera_total": 100000.0, "saldo_contable": 100000.0, "diferencia": 0.0, "cuadra": True},
    "tributario": {"limite_ejercicio_1pct": 1000.0, "tope_acumulado_10pct": 10000.0,
                   "excede_limite_ejercicio": True, "nota": "Límite de deducción"},
    "politica": {"filas": [{"banda": "Por vencer", "banda_origen": "Por vencer", "exposicion": 80000.0,
                            "tasa_politica": 0.0, "provision_politica": 0.0, "ecl": 800.0, "diferencia": 800.0}],
                 "provision_politica_total": 0.0, "diferencia_bruta": 800.0},
    "ecl_total": 1800.0, "hallazgos": [], "pendientes": [],
    "bitacora": {"bandas": ["Por vencer", "0 a 30 días"], "umbral_incumplimiento": 730,
                 "umbral_individual": 0, "cortes": [], "metodo": "Permanencia a 24 meses",
                 "descuento": "No aplicado (NIIF 9 B5.5.44)"},
}


def _abrir(binario, con_formulas=True):
    return load_workbook(io.BytesIO(binario), data_only=not con_formulas)


def test_el_libro_trae_todas_las_hojas():
    wb = _abrir(construir_excel(RESULTADO, {"entidad": "PRUEBA S.A."}))
    assert wb.sheetnames == ["00-Caratula", "01-Parametros", "02-Fuentes", "03-Cohorte", "04-Tasas",
                             "05-Matriz", "06-Individual", "07-Politica", "08-Conciliacion",
                             "09-Tributario", "10-Hallazgos", "11-Pendientes", "12-Bitacora"]


def test_la_matriz_calcula_con_formulas_y_no_con_valores_pegados():
    ws = _abrir(construir_excel(RESULTADO, {}))["05-Matriz"]
    formulas = [c.value for fila in ws.iter_rows() for c in fila
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("*" in f for f in formulas), "la pérdida de cada banda debe ser una fórmula"
    assert any(f.startswith("=SUM(") for f in formulas), "el total debe ser una suma"


def test_los_parametros_estan_en_celdas_con_nombre_y_las_formulas_los_usan():
    wb = _abrir(construir_excel(RESULTADO, {}))
    assert "AjusteProspectivo" in wb.defined_names
    ws = wb["05-Matriz"]
    assert any(isinstance(c.value, str) and "AjusteProspectivo" in c.value
               for fila in ws.iter_rows() for c in fila)


def test_el_libro_abre_sin_reparacion_y_recalcula_al_abrirse():
    wb = _abrir(construir_excel(RESULTADO, {}))
    assert wb.calculation.fullCalcOnLoad is True


def test_banda_sin_medir_no_escribe_cero_ni_formula_de_perdida():
    """Un tramo con ``tasa_perdida`` y ``ecl`` en ``None`` (banda sin tasa
    observada ni sustituta, camino que existe en el motor pero que ninguna
    prueba ejercitaba) debe rotularse "SIN MEDIR" en 05-Matriz, nunca como
    0,00 ni como fórmula: un cero confundiría "no medido" con "pérdida cero
    real"."""
    resultado = {
        "matriz": {"tramos": [
            {"segmento": "NO-RELACIONADOS", "tramo": "Más de 360 días",
             "exposicion": 5000.0, "tasa_perdida": None, "ecl": None},
        ]},
    }
    ws = _abrir(construir_excel(resultado, {}))["05-Matriz"]
    fila = 2
    assert ws.cell(fila, 3).value == 5000.0  # la exposición sí se traslada
    assert ws.cell(fila, 4).value == "SIN MEDIR"
    assert ws.cell(fila, 6).value == "SIN MEDIR"
    assert ws.cell(fila, 6).value != 0
    valor_col_f = ws.cell(fila, 6).value
    assert not (isinstance(valor_col_f, str) and valor_col_f.startswith("=")), \
        "una banda sin medir no debe llevar fórmula de pérdida"


def test_cada_segmento_resuelve_su_propio_factor_prospectivo():
    """El motor mide cada segmento por separado (``service.analizar`` llama a
    ``medir_ecl`` una vez por segmento) y cada uno puede traer un factor
    prospectivo distinto (p. ej. 1,05 para terceros y 1,10 para
    relacionadas). 05-Matriz debe resolver, fila por fila, el nombre
    definido que corresponde al segmento de ESA fila -no un único factor
    global que desconoce el resto de segmentos-."""
    resultado = {
        "matriz": {"tramos": [
            {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
             "exposicion": 80000.0, "tasa_perdida": 0.01, "ecl": 800.0},
            {"segmento": "RELACIONADOS", "tramo": "Por vencer",
             "exposicion": 20000.0, "tasa_perdida": 0.02, "ecl": 400.0},
        ]},
    }
    parametros = {"factor_prospectivo": {"NO-RELACIONADOS": 1.05, "RELACIONADOS": 1.10}}
    wb = _abrir(construir_excel(resultado, parametros))

    # Los dos nombres existen en el libro, cada uno con el valor esperado
    # (factor - 1: la misma convención que ya tenía AjusteProspectivo).
    esperados = {"AjusteProspectivoNoRelacionados": 0.05, "AjusteProspectivoRelacionados": 0.10}
    for nombre, esperado in esperados.items():
        assert nombre in wb.defined_names, f"falta el nombre definido {nombre}"
        dn = wb.defined_names[nombre]
        hoja, celda = next(dn.destinations)
        valor = wb[hoja][celda].value
        assert abs(valor - esperado) < 1e-9, f"{nombre} debería ser {esperado}, es {valor}"

    ws = wb["05-Matriz"]
    fila_no_relacionados, fila_relacionados = 2, 3
    assert ws.cell(fila_no_relacionados, 1).value == "NO-RELACIONADOS"
    assert ws.cell(fila_relacionados, 1).value == "RELACIONADOS"

    # Cada fila resuelve el nombre según SU PROPIO segmento (columna A de esa
    # misma fila), no el de otra fila ni un nombre único compartido.
    formula_no_relacionados = ws.cell(fila_no_relacionados, 5).value
    formula_relacionados = ws.cell(fila_relacionados, 5).value
    esperado_no_relacionados = (f'=IF(A{fila_no_relacionados}="RELACIONADOS",'
                                f'AjusteProspectivoRelacionados,AjusteProspectivoNoRelacionados)')
    esperado_relacionados = (f'=IF(A{fila_relacionados}="RELACIONADOS",'
                             f'AjusteProspectivoRelacionados,AjusteProspectivoNoRelacionados)')
    assert formula_no_relacionados == esperado_no_relacionados, formula_no_relacionados
    assert formula_relacionados == esperado_relacionados, formula_relacionados


def test_el_excel_muestra_el_factor_prospectivo_aplicado_no_el_solicitado():
    """Cuando el motor rechaza un ajuste prospectivo solicitado (p. ej. porque
    falta justificación), aplica 1,0 (sin ajuste) y lo escribe en
    ``resultado["matriz"]["ajuste_prospectivo"]``. El Excel debe reflejar lo
    **realmente aplicado**, no lo solicitado en parámetros.

    Caso: usuario pide 1,15 para NO-RELACIONADOS y 1,20 para RELACIONADOS,
    pero el motor rechaza ambos (sin justificación) y aplica 1,0 para ambos.
    El Excel debe mostrar 0,00% (factor 1,0 = 0% de ajuste), no 15% ni 20%.
    """
    resultado = {
        "matriz": {
            "tramos": [
                {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 80000.0, "tasa_perdida": 0.01, "ecl": 800.0},
                {"segmento": "RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 20000.0, "tasa_perdida": 0.02, "ecl": 400.0},
            ],
            # Lo que el motor realmente aplicó (rechazó los solicitados)
            "ajuste_prospectivo": {"NO-RELACIONADOS": 1.0, "RELACIONADOS": 1.0}
        },
    }
    parametros = {
        # Lo que el usuario pidió pero fue rechazado
        "factor_prospectivo": {"NO-RELACIONADOS": 1.15, "RELACIONADOS": 1.20},
    }
    wb = _abrir(construir_excel(resultado, parametros))

    # Los nombres definidos en 01-Parametros deben traer lo APLICADO (1,0 = 0%)
    # no lo solicitado (1,15 = 15% o 1,20 = 20%).
    esperados = {"AjusteProspectivoNoRelacionados": 0.0, "AjusteProspectivoRelacionados": 0.0}
    for nombre, esperado in esperados.items():
        assert nombre in wb.defined_names, f"falta el nombre definido {nombre}"
        dn = wb.defined_names[nombre]
        hoja, celda = next(dn.destinations)
        valor = wb[hoja][celda].value
        assert abs(valor - esperado) < 1e-9, \
            f"{nombre} debería ser {esperado} (lo aplicado), es {valor}"


def test_el_excel_refleja_factor_prospectivo_aplicado_distinto_de_uno():
    """Caso inverso: cuando el motor SÍ aplica un factor distinto de 1,0,
    el Excel debe reflejarlo. Por ejemplo, factor aplicado 1,08."""
    resultado = {
        "matriz": {
            "tramos": [
                {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 80000.0, "tasa_perdida": 0.01, "ecl": 800.0},
            ],
            # Lo que el motor realmente aplicó
            "ajuste_prospectivo": {"NO-RELACIONADOS": 1.08}
        },
    }
    parametros = {
        # Coincide con lo aplicado
        "factor_prospectivo": {"NO-RELACIONADOS": 1.08},
    }
    wb = _abrir(construir_excel(resultado, parametros))

    dn = wb.defined_names["AjusteProspectivoNoRelacionados"]
    hoja, celda = next(dn.destinations)
    valor = wb[hoja][celda].value
    # 1,08 - 1,0 = 0,08 (8%)
    assert abs(valor - 0.08) < 1e-9, \
        f"AjusteProspectivoNoRelacionados debería ser 0.08 (factor 1,08), es {valor}"


def test_tolera_formato_antiguo_del_ajuste_prospectivo_escalar_no_cero():
    """Formato retrocompatible: corridas guardadas antes de la refactor tenían
    ``ajuste_prospectivo`` como un escalar (p. ej. 0.05), no un diccionario
    por segmento. El exportador debe interpretarlo como el mismo ajuste para
    ambos segmentos, sin revienta al llamar .get() en un float."""
    resultado = {
        "matriz": {
            "tramos": [
                {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 80000.0, "tasa_perdida": 0.01, "ecl": 800.0},
                {"segmento": "RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 20000.0, "tasa_perdida": 0.02, "ecl": 400.0},
            ],
            # Formato antiguo: escalar en lugar de diccionario
            "ajuste_prospectivo": 0.05  # 5% de ajuste para ambos segmentos
        },
    }
    parametros = {}
    wb = _abrir(construir_excel(resultado, parametros))

    # El Excel se genera sin excepción y ambos factores prospectivos quedan en 0.05
    esperados = {"AjusteProspectivoNoRelacionados": 0.05, "AjusteProspectivoRelacionados": 0.05}
    for nombre, esperado in esperados.items():
        assert nombre in wb.defined_names, f"falta el nombre definido {nombre}"
        dn = wb.defined_names[nombre]
        hoja, celda = next(dn.destinations)
        valor = wb[hoja][celda].value
        assert abs(valor - esperado) < 1e-9, \
            f"{nombre} debería ser {esperado}, es {valor}"


def test_tolera_formato_antiguo_del_ajuste_prospectivo_escalar_cero():
    """Formato retrocompatible: cuando el escalar es 0.0 (ajuste cero para
    ambos segmentos, el caso que ya usan varias pruebas existentes del archivo)."""
    resultado = {
        "matriz": {
            "tramos": [
                {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 80000.0, "tasa_perdida": 0.01, "ecl": 800.0},
                {"segmento": "RELACIONADOS", "tramo": "Por vencer",
                 "exposicion": 20000.0, "tasa_perdida": 0.02, "ecl": 400.0},
            ],
            # Formato antiguo: escalar 0.0 (sin ajuste)
            "ajuste_prospectivo": 0.0
        },
    }
    parametros = {}
    wb = _abrir(construir_excel(resultado, parametros))

    # El Excel se genera sin excepción y ambos factores prospectivos quedan en 0.0
    esperados = {"AjusteProspectivoNoRelacionados": 0.0, "AjusteProspectivoRelacionados": 0.0}
    for nombre, esperado in esperados.items():
        assert nombre in wb.defined_names, f"falta el nombre definido {nombre}"
        dn = wb.defined_names[nombre]
        hoja, celda = next(dn.destinations)
        valor = wb[hoja][celda].value
        assert abs(valor - esperado) < 1e-9, \
            f"{nombre} debería ser {esperado}, es {valor}"


def _textos(ws) -> list[str]:
    return [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]


# ---------------------------------------------------------------------------
# I11 — el Excel perdía la marca de PRELIMINAR que la pantalla muestra dos veces
# ---------------------------------------------------------------------------

def test_la_caratula_marca_el_papel_como_preliminar():
    """La pantalla rotula el papel «PRELIMINAR» arriba y abajo; el libro que
    entra al archivo permanente no contenía esa palabra en ninguna de sus
    trece hojas, con «Preparado por» y «Revisado por» en blanco y sin
    advertencia (I11)."""
    wb = _abrir(construir_excel(RESULTADO, {"entidad": "PRUEBA S.A."}))
    textos = _textos(wb["00-Caratula"])
    assert any("PRELIMINAR" in t for t in textos), \
        "la carátula debe declarar que el papel es preliminar"
    assert any("Socio" in t for t in textos), \
        "debe decir de quién depende la aprobación"


def test_la_caratula_no_deja_en_blanco_al_preparador_ni_al_revisor():
    """Una celda vacía se lee como «no aplica»; lo que falta se declara (I11)."""
    wb = _abrir(construir_excel(RESULTADO, {}))
    ws = wb["00-Caratula"]
    valores = {ws.cell(f, 1).value: ws.cell(f, 2).value for f in range(1, 30)}
    assert valores.get("Preparado por") == "(pendiente)"
    assert valores.get("Revisado por") == "(pendiente)"


# ---------------------------------------------------------------------------
# M4 — 01-Parametros atribuía a NIIF 9 B5.5.37 un umbral de 730 días
# ---------------------------------------------------------------------------

def test_el_umbral_de_incumplimiento_no_se_atribuye_a_la_norma():
    """B5.5.37 establece una presunción refutable de 90 días. El umbral de 730
    que usa la herramienta es política de la entidad: la fuente no puede
    presentarlo como el valor por defecto de la norma (M4)."""
    ws = _abrir(construir_excel(RESULTADO, {}))["01-Parametros"]
    assert ws.cell(4, 1).value == "Umbral de incumplimiento (días)"
    fuente = str(ws.cell(4, 3).value)
    assert "730 días por defecto" not in fuente, \
        "la norma no trae 730 días por defecto: eso es política de la entidad"
    assert "90" in fuente, "debe decir cuál es la presunción de la norma (90 días)"
    assert "refutable" in fuente.lower(), "y que esa presunción es refutable"
    assert "olítica de la entidad" in fuente, "y que el plazo aplicado lo fija la entidad"


# ---------------------------------------------------------------------------
# M7 — el papel no era reproducible: se escribía la fecha del día de la descarga
# ---------------------------------------------------------------------------

PARAMETROS_CORRIDA = {
    "entidad": "PRUEBA S.A.",
    "fechas": ["2022-12-31", "2023-12-31", "2024-12-31"],
    "fecha_emision": "2025-03-14",
}


#: Lo único que openpyxl sella al guardar y no se puede fijar por su API:
#: `save_workbook` reescribe `properties.modified` con la hora de ese momento
#: (openpyxl/writer/excel.py). No es contenido del papel -ninguna celda lo
#: muestra-, igual que la hora de escritura de cada entrada del .zip, que fija
#: `zipfile`. Todo lo demás, incluida la fecha de creación, sí es reproducible.
PARTE_SELLADA_AL_GUARDAR = "docProps/core.xml"


def _contenido_del_paquete(binario: bytes) -> dict:
    """Contenido de cada parte del .xlsx, sin el metadato que sella openpyxl."""
    z = zipfile.ZipFile(io.BytesIO(binario))
    return {nombre: z.read(nombre) for nombre in sorted(z.namelist())
            if nombre != PARTE_SELLADA_AL_GUARDAR}


def _creacion_del_paquete(binario: bytes) -> str:
    xml = zipfile.ZipFile(io.BytesIO(binario)).read(PARTE_SELLADA_AL_GUARDAR).decode("utf-8")
    return re.search(r"<dcterms:created[^>]*>([^<]+)<", xml).group(1)


def _celda_por_concepto(ws, concepto: str):
    for fila in range(1, 40):
        if ws.cell(fila, 1).value == concepto:
            return ws.cell(fila, 2).value
    return None


def test_la_fecha_de_emision_sale_de_la_corrida_no_del_dia_de_la_descarga(monkeypatch):
    """``models.py`` dice que la corrida existe para reproducir el papel «tal
    como se emitió», pero 00-Caratula y 12-Bitacora escribían ``date.today()``
    (M7)."""
    monkeypatch.setattr(exporter, "date", _FechaFija)
    wb = _abrir(construir_excel(RESULTADO, PARAMETROS_CORRIDA))

    assert _celda_por_concepto(wb["00-Caratula"], "Fecha de emisión del papel") == "2025-03-14"
    generacion = str(_celda_por_concepto(wb["12-Bitacora"], "Versión y fecha de generación"))
    assert "2025-03-14" in generacion

    for hoja in wb.sheetnames:
        for texto in _textos(wb[hoja]):
            assert _FechaFija.HOY not in texto, \
                f"la hoja {hoja} imprime la fecha de la descarga, no la de la corrida"


def test_sin_fecha_de_emision_registrada_se_declara_en_vez_de_poner_la_de_hoy(monkeypatch):
    """Corridas antiguas que no guardaron la fecha de emisión: lo que no se
    registró se declara, nunca se rellena con el día de la descarga (M7)."""
    monkeypatch.setattr(exporter, "date", _FechaFija)
    wb = _abrir(construir_excel(RESULTADO, {"fechas": ["2022-12-31", "2023-12-31", "2024-12-31"]}))
    valor = str(_celda_por_concepto(wb["00-Caratula"], "Fecha de emisión del papel"))
    assert _FechaFija.HOY not in valor
    assert "no registrada" in valor.lower()


def test_dos_descargas_de_la_misma_corrida_dan_el_mismo_papel(monkeypatch):
    """Dos descargas en días distintos tienen que dar el mismo papel: es la
    única forma de que el revisor recalcule lo que el preparador emitió (M7)."""
    monkeypatch.setattr(exporter, "date", _FechaFija)
    primera = construir_excel(RESULTADO, PARAMETROS_CORRIDA)

    class _OtroDia(_FechaFija):
        HOY = "2031-07-09"

        @classmethod
        def today(cls):
            return date(2031, 7, 9)

    monkeypatch.setattr(exporter, "date", _OtroDia)
    segunda = construir_excel(RESULTADO, PARAMETROS_CORRIDA)

    assert _contenido_del_paquete(primera) == _contenido_del_paquete(segunda)
    # La fecha de creación del archivo también sale de la corrida.
    assert _creacion_del_paquete(primera) == _creacion_del_paquete(segunda)
    assert _creacion_del_paquete(primera).startswith("2025-03-14")


# ---------------------------------------------------------------------------
# M8 — 05-Matriz emitía una fila "SIN MEDIR" con exposición 0,00 por cada
#      combinación segmento x banda inexistente
# ---------------------------------------------------------------------------

BANDAS = ["Por vencer", "0 a 30 días", "31 a 60 días", "61 a 90 días",
          "91 a 180 días", "181 a 360 días", "361 a 730 días", "Más de 730 días"]

#: Lo que arma `service.analizar`: todas las bandas para los dos segmentos.
#: Solo cuatro filas dicen algo; las otras doce son ceros sin tasa.
CON_EXPOSICION = {
    ("NO-RELACIONADOS", "Por vencer"): (80000.0, 0.01, 800.0),
    ("NO-RELACIONADOS", "0 a 30 días"): (20000.0, 0.05, 1000.0),
    ("NO-RELACIONADOS", "Más de 730 días"): (5000.0, None, None),   # con cartera, sin tasa
    ("RELACIONADOS", "Por vencer"): (0.0, 0.02, 0.0),               # cero medido de verdad
}
TRAMOS_COMPLETOS = [
    {"segmento": s, "tramo": b,
     "exposicion": CON_EXPOSICION.get((s, b), (0.0, None, None))[0],
     "tasa_perdida": CON_EXPOSICION.get((s, b), (0.0, None, None))[1],
     "ecl": CON_EXPOSICION.get((s, b), (0.0, None, None))[2]}
    for s in ("NO-RELACIONADOS", "RELACIONADOS") for b in BANDAS
]
RESULTADO_CON_RUIDO = {"matriz": {"tramos": TRAMOS_COMPLETOS}}


def _filas_de_matriz(ws) -> list[tuple]:
    filas = []
    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 2).value == "TOTAL" or ws.cell(fila, 2).value is None:
            break
        filas.append((ws.cell(fila, 1).value, ws.cell(fila, 2).value, ws.cell(fila, 3).value))
    return filas


def test_la_matriz_no_emite_las_combinaciones_sin_exposicion_ni_tasa():
    """De 16 combinaciones segmento x banda, 12 salen en cero y sin tasa: no
    son una pérdida cero medida ni una banda sin medir con cartera, y su ruido
    diluye la señal real de «SIN MEDIR» (M8)."""
    ws = _abrir(construir_excel(RESULTADO_CON_RUIDO, {}))["05-Matriz"]
    filas = _filas_de_matriz(ws)
    assert len(filas) == 4, f"debían quedar 4 filas con información, quedaron {len(filas)}"
    assert {(f[0], f[1]) for f in filas} == set(CON_EXPOSICION)
    for segmento, banda, exposicion in filas:
        assert not (exposicion in (0, 0.0) and ws.cell(2, 4).value == "SIN MEDIR" and banda != "Por vencer")


def test_la_matriz_declara_cuantas_combinaciones_omitio():
    """Omitir no es esconder: el papel dice cuántas combinaciones se dejaron
    fuera y por qué (M8)."""
    ws = _abrir(construir_excel(RESULTADO_CON_RUIDO, {}))["05-Matriz"]
    textos = " ".join(_textos(ws))
    assert "12" in textos, "debe decir cuántas combinaciones se omitieron"
    assert "sin exposición ni tasa" in textos


def test_la_fila_sin_medir_con_cartera_sigue_saliendo():
    """Lo que se filtra es el ruido de cero, nunca una banda con cartera que no
    se pudo medir (M8)."""
    ws = _abrir(construir_excel(RESULTADO_CON_RUIDO, {}))["05-Matriz"]
    filas = {(ws.cell(f, 1).value, ws.cell(f, 2).value): f for f in range(2, ws.max_row + 1)}
    fila = filas[("NO-RELACIONADOS", "Más de 730 días")]
    assert ws.cell(fila, 3).value == 5000.0
    assert ws.cell(fila, 6).value == "SIN MEDIR"


# ---------------------------------------------------------------------------
# I2, I3 y M5 — 04-Tasas: origen, acotamiento y universo de bandas
# ---------------------------------------------------------------------------

def _fila_de_tasas(ws, segmento: str, banda: str) -> int | None:
    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 1).value == segmento and ws.cell(fila, 2).value == banda:
            return fila
    return None


def _tasa_aplicada_de_la_matriz(wb, fila: int):
    """Lo que 05-Matriz aplica en esa fila, siguiendo la referencia si la hay."""
    valor = wb["05-Matriz"].cell(fila, 4).value
    if isinstance(valor, str) and valor.startswith("='04-Tasas'!"):
        celda = valor.split("!", 1)[1]
        return wb["04-Tasas"][celda].value
    return valor


# --- I2 --------------------------------------------------------------------

RESULTADO_SUSTITUTA = {
    "tasas": {"NO-RELACIONADOS": {"0 a 30 días": 0.10}},
    "detalle_cohorte": {"NO-RELACIONADOS": {
        "0 a 30 días": {"inicial": 100000.0, "remanente": 10000.0, "documentos": 8}}},
    "matriz": {"tramos": [{"segmento": "NO-RELACIONADOS", "tramo": "0 a 30 días",
                           "exposicion": 50000.0, "tasa_perdida": 0.42, "ecl": 21000.0}]},
}
JUSTIFICACION_SUSTITUTA = ("Concordato preventivo del deudor principal de la banda, informado por "
                           "la gerencia el 12/01: la historia de la cohorte no lo recoge.")
PARAMETROS_SUSTITUTA = {"tasas_sustitutas": {
    "NO-RELACIONADOS|0 a 30 días": {"tasa": 0.42, "justificacion": JUSTIFICACION_SUSTITUTA}}}


def test_la_tasa_sustituida_se_presenta_como_sustituida_y_con_su_justificacion():
    """El servicio prefiere la sustituta; el exportador hacía lo contrario y,
    si había observada, marcaba origen «Observada» y borraba la justificación.
    Dos hojas del mismo papel daban dos tasas para la misma banda y la única
    evidencia del cambio no se escribía en ninguna parte (I2)."""
    wb = _abrir(construir_excel(RESULTADO_SUSTITUTA, PARAMETROS_SUSTITUTA))
    ws = wb["04-Tasas"]
    fila = _fila_de_tasas(ws, "NO-RELACIONADOS", "0 a 30 días")
    assert fila is not None

    assert ws.cell(fila, 5).value == "Sustituida", ws.cell(fila, 5).value
    assert JUSTIFICACION_SUSTITUTA in str(ws.cell(fila, 6).value)
    assert ws.cell(fila, 4).value == 0.42, "04-Tasas debe traer la tasa que de verdad se aplicó"


def test_las_dos_hojas_dan_la_misma_tasa_para_la_misma_banda():
    """05-Matriz aplicaba 42 % y 04-Tasas mostraba 10 % por fórmula (I2)."""
    wb = _abrir(construir_excel(RESULTADO_SUSTITUTA, PARAMETROS_SUSTITUTA))
    assert _tasa_aplicada_de_la_matriz(wb, 2) == 0.42


# --- I3 --------------------------------------------------------------------

RESULTADO_ACOTADO = {
    # `cohortes.tasas_por_permanencia` acota a 1,0 y deja constancia en anomalías.
    "tasas": {"NO-RELACIONADOS": {"61 a 90 días": 1.0}},
    "detalle_cohorte": {"NO-RELACIONADOS": {
        "61 a 90 días": {"inicial": 10000.0, "remanente": 35000.0, "documentos": 3}}},
    "anomalias": [{"segmento": "NO-RELACIONADOS", "banda": "61 a 90 días", "inicial": 10000.0,
                   "remanente": 35000.0, "tasa_bruta": 3.5, "tipo": "remanente_mayor_que_inicial"}],
    "matriz": {"tramos": [{"segmento": "NO-RELACIONADOS", "tramo": "61 a 90 días",
                           "exposicion": 20000.0, "tasa_perdida": 1.0, "ecl": 20000.0}]},
}


def test_la_tasa_acotada_aparece_acotada_en_el_excel():
    """04-Tasas mostraba E2/D2 = 350 % con origen «Observada» mientras 05-Matriz
    aplicaba 100 %: quien recalculara desde 04-Tasas obtenía 3,5 veces la PCE
    del papel (I3)."""
    wb = _abrir(construir_excel(RESULTADO_ACOTADO, {}))
    ws = wb["04-Tasas"]
    fila = _fila_de_tasas(ws, "NO-RELACIONADOS", "61 a 90 días")

    aplicada = ws.cell(fila, 4).value
    assert isinstance(aplicada, str) and aplicada.startswith("="), \
        "la tasa aplicada debe seguir siendo recalculable desde 03-Cohorte"
    assert "1" in aplicada and "IF" in aplicada, \
        f"la fórmula debe acotar el ratio a 1,0: {aplicada}"
    assert _tasa_aplicada_de_la_matriz(wb, 2) == aplicada, \
        "05-Matriz debe aplicar exactamente la tasa de 04-Tasas"


def test_el_acotamiento_se_explica_en_la_hoja_de_tasas():
    """El acotamiento solo se mencionaba en el texto de un pendiente (I3)."""
    wb = _abrir(construir_excel(RESULTADO_ACOTADO, {}))
    ws = wb["04-Tasas"]
    fila = _fila_de_tasas(ws, "NO-RELACIONADOS", "61 a 90 días")
    assert "acotada" in str(ws.cell(fila, 5).value).lower(), ws.cell(fila, 5).value
    nota = str(ws.cell(fila, 6).value)
    assert "35.000,00" in nota or "35,000.00" in nota, nota
    assert "10.000,00" in nota or "10,000.00" in nota, nota
    # El ratio crudo no se borra: es la evidencia de la que salió el acotamiento.
    crudo = str(ws.cell(fila, 3).value)
    assert crudo.startswith("=") and "03-Cohorte" in crudo, crudo


# --- M5 --------------------------------------------------------------------

def test_la_hoja_de_tasas_cubre_todas_las_bandas_de_la_matriz():
    """Una corrida con 16 filas segmento x banda en 05-Matriz tuvo 1 fila en
    04-Tasas: las bandas sin historia -las que el papel tiene que argumentar-
    no tenían fila en la hoja de tasas (M5)."""
    wb = _abrir(construir_excel(RESULTADO_CON_RUIDO, {}))
    ws04, ws05 = wb["04-Tasas"], wb["05-Matriz"]

    pares_matriz = [(ws05.cell(f, 1).value, ws05.cell(f, 2).value)
                    for f in range(2, 2 + len(CON_EXPOSICION))]
    for segmento, banda in pares_matriz:
        assert _fila_de_tasas(ws04, segmento, banda) is not None, \
            f"falta la fila de {segmento} / {banda} en 04-Tasas"


def test_la_banda_sin_historia_dice_por_que_no_tiene_tasa():
    """Sin fila no hay nada que argumentar; con fila, el papel tiene que decir
    que la cohorte no registra documentos en esa banda (M5)."""
    wb = _abrir(construir_excel(RESULTADO_CON_RUIDO, {}))
    ws = wb["04-Tasas"]
    fila = _fila_de_tasas(ws, "NO-RELACIONADOS", "Más de 730 días")
    assert ws.cell(fila, 4).value == "SIN MEDIR"
    assert ws.cell(fila, 5).value == "SIN MEDIR"
    assert "cohorte" in str(ws.cell(fila, 6).value).lower()


def test_la_hoja_de_tasas_conserva_la_banda_que_solo_existe_en_la_cohorte():
    """Una banda con historia pero sin cartera hoy sigue documentada: su tasa
    es evidencia aunque no se aplique en 05-Matriz (M5)."""
    resultado = {
        "tasas": {"NO-RELACIONADOS": {"181 a 360 días": 0.30}},
        "detalle_cohorte": {"NO-RELACIONADOS": {
            "181 a 360 días": {"inicial": 20000.0, "remanente": 6000.0, "documentos": 4}}},
        "matriz": {"tramos": [{"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
                               "exposicion": 1000.0, "tasa_perdida": 0.01, "ecl": 10.0}]},
    }
    ws = _abrir(construir_excel(resultado, {}))["04-Tasas"]
    fila = _fila_de_tasas(ws, "NO-RELACIONADOS", "181 a 360 días")
    assert fila is not None, "la banda con historia no puede desaparecer de 04-Tasas"
    assert str(ws.cell(fila, 4).value).startswith("=")
    assert "no se aplica" in str(ws.cell(fila, 6).value).lower()


def test_cada_fila_de_tasas_apunta_a_su_propia_fila_de_cohorte():
    """04-Tasas ya no comparte el orden de 03-Cohorte (lista además las bandas
    sin historia), así que la referencia tiene que resolverse por par
    segmento/banda y no por número de fila coincidente."""
    resultado = {
        "tasas": {"NO-RELACIONADOS": {"Por vencer": 0.01, "91 a 180 días": 0.30},
                  "RELACIONADOS": {"Por vencer": 0.02}},
        "detalle_cohorte": {
            "NO-RELACIONADOS": {"91 a 180 días": {"inicial": 10000.0, "remanente": 3000.0, "documentos": 2},
                                "Por vencer": {"inicial": 50000.0, "remanente": 500.0, "documentos": 9}},
            "RELACIONADOS": {"Por vencer": {"inicial": 8000.0, "remanente": 160.0, "documentos": 1}},
        },
        "matriz": {"tramos": [
            {"segmento": "RELACIONADOS", "tramo": "Por vencer", "exposicion": 4000.0,
             "tasa_perdida": 0.02, "ecl": 80.0},
            {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer", "exposicion": 70000.0,
             "tasa_perdida": 0.01, "ecl": 700.0},
            {"segmento": "NO-RELACIONADOS", "tramo": "Más de 730 días", "exposicion": 900.0,
             "tasa_perdida": None, "ecl": None},
        ]},
    }
    wb = _abrir(construir_excel(resultado, {}))
    ws04, ws03 = wb["04-Tasas"], wb["03-Cohorte"]
    referencias = re.compile(r"'03-Cohorte'!D(\d+)")

    revisadas = 0
    for fila in range(2, ws04.max_row + 1):
        formula = str(ws04.cell(fila, 3).value or "")
        m = referencias.search(formula)
        if not m:
            continue
        fila_coh = int(m.group(1))
        assert ws03.cell(fila_coh, 1).value == ws04.cell(fila, 1).value
        assert ws03.cell(fila_coh, 2).value == ws04.cell(fila, 2).value
        revisadas += 1
    assert revisadas == 3, f"debían resolverse las tres bandas con cohorte, se resolvieron {revisadas}"


# ---------------------------------------------------------------------------
# M3 — la fórmula de 06-Individual difería en un centavo de la PCE persistida
# ---------------------------------------------------------------------------

#: `motor.evaluar_individual` con saldo 1000,005 y recuperación 333,333:
#: redondea cada importe por separado y la PCE sobre los valores SIN redondear,
#: así que C - D en Excel daba 666,68 y la base guardaba 666,67.
RESULTADO_CENTAVO = {
    "individual": {"casos": [{"identificacion": "DELTA (NO-RELACIONADOS)", "tramo": None,
                              "saldo": 1000.01, "recuperacion_estimada": 333.33,
                              "saldo_sin_tasa": 0.0, "ecl": 666.67,
                              "sustento": "Acuerdo de pago firmado el 03/02"}],
                   "saldo_total": 1000.01, "ecl_total": 666.67},
    "matriz": {"tramos": []},
}


def test_la_perdida_individual_del_excel_es_la_que_guardo_la_base():
    """Con saldo 1000,005 y recuperación 333,333 la pantalla y la base dan
    666,67 y la fórmula del Excel daba 666,68; 09-Tributario B2 se construye
    sobre la versión del Excel (M3)."""
    wb = _abrir(construir_excel(RESULTADO_CENTAVO, {}))
    ws = wb["06-Individual"]
    assert ws.cell(2, 5).value == 666.67, \
        f"la PCE del papel debe ser la persistida, es {ws.cell(2, 5).value!r}"
    # La recuperación es la cifra derivada (el servicio la calcula como
    # saldo - PCE), así que es ella la que se recalcula en el papel.
    assert ws.cell(2, 4).value == "=C2-E2", ws.cell(2, 4).value


def test_el_total_individual_cuadra_con_el_ecl_total_de_la_corrida():
    """El total de 06-Individual alimenta 09-Tributario: si arrastra el centavo,
    el exceso no deducible también (M3)."""
    wb = _abrir(construir_excel(RESULTADO_CENTAVO, {}))
    ws = wb["06-Individual"]
    valores = [ws.cell(f, 5).value for f in range(2, 3)]
    assert sum(valores) == RESULTADO_CENTAVO["individual"]["ecl_total"]
    assert ws.cell(3, 5).value == "=SUM(E2:E2)"


def test_la_matriz_redondea_la_perdida_como_el_motor():
    """`medir_ecl` redondea la PCE de cada banda a dos decimales antes de
    sumarla; la fórmula del Excel arrastraba todos los decimales del producto
    y el total del papel se apartaba del que guardó la corrida (M3).

    Se comprueba EL NÚMERO que produce la celda, no el texto de la fórmula:
    exigir que "empiece por =ROUND( y termine en ,2)" verifica la
    implementación y no el comportamiento, y era el hueco por el que pasaba
    que el libro no aplicara los acotamientos del motor.
    """
    resultado = {"matriz": {"tramos": [
        {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer",
         "exposicion": 1000.01, "tasa_perdida": 0.3333, "ecl": 333.30},
    ]}}
    libro = Libro(_abrir(construir_excel(resultado, {})))
    # 1000,01 x 0,3333 = 333,303333; el motor archiva 333,30 y la celda tiene
    # que dar exactamente eso, no el producto con todos sus decimales.
    assert libro.numero("05-Matriz", "F2") == 333.30
    assert libro.numero("05-Matriz", "F3") == 333.30  # la fila TOTAL


# ---------------------------------------------------------------------------
# I4 — «Cartera medida» significaba cuatro cosas y el importe sin medir no
#      aparecía en ninguna de las trece hojas
# ---------------------------------------------------------------------------

RESULTADO_SIN_MEDIR = {
    "exposicion": {"colectiva": 100000.0, "individual": 50000.0, "sin_estratificar": 20000.0,
                   "sin_medir": 35000.0, "total": 170000.0, "segun_archivo": 150000.0,
                   "factores_anclaje": {"NO-RELACIONADOS": 1.0, "RELACIONADOS": 1.0}},
    "matriz": {"tramos": [
        {"segmento": "NO-RELACIONADOS", "tramo": "Por vencer", "exposicion": 70000.0,
         "tasa_perdida": 0.01, "ecl": 700.0},
        {"segmento": "NO-RELACIONADOS", "tramo": "Más de 730 días", "exposicion": 30000.0,
         "tasa_perdida": None, "ecl": None},
    ]},
    "individual": {"casos": [
        {"identificacion": "MEGA S.A. (NO-RELACIONADOS)", "tramo": None, "saldo": 50000.0,
         "recuperacion_estimada": 45000.0, "ecl": 5000.0, "saldo_sin_tasa": 5000.0,
         "sustento": "Medido con la tasa de la matriz (provisional); USD 5,000.00 sin medir"},
    ], "saldo_total": 50000.0, "ecl_total": 5000.0},
    "conciliacion": {"cartera_total": 170000.0, "saldo_contable": 170000.0,
                     "diferencia": 0.0, "cuadra": True},
}
#: Lo que la pantalla calcula: total - sin medir - sin estratificar.
CARTERA_MEDIDA_DE_LA_PANTALLA = 170000.0 - 35000.0 - 20000.0


def _suma_sin_medir_del_papel(wb) -> float:
    """Recalcula a mano lo que suma la fórmula de la conciliación."""
    total = 0.0
    ws = wb["05-Matriz"]
    for f in range(2, ws.max_row + 1):
        if ws.cell(f, 2).value == "TOTAL":
            break
        if ws.cell(f, 6).value == "SIN MEDIR":
            total += float(ws.cell(f, 3).value or 0)
    ws = wb["06-Individual"]
    for f in range(2, ws.max_row + 1):
        if ws.cell(f, 1).value == "TOTAL":
            break
        total += float(ws.cell(f, 6).value or 0)
    return total


def test_el_individual_declara_el_saldo_que_no_pudo_medir():
    """`saldo_sin_tasa` viajaba en el resultado y no llegaba a ninguna columna
    del papel: solo se podía leer parseando la frase del sustento (I4)."""
    ws = _abrir(construir_excel(RESULTADO_SIN_MEDIR, {}))["06-Individual"]
    assert "sin medir" in str(ws.cell(1, 6).value).lower(), ws.cell(1, 6).value
    assert ws.cell(2, 6).value == 5000.0
    assert ws.cell(3, 6).value == "=SUM(F2:F2)"


def test_la_conciliacion_totaliza_la_exposicion_sin_medir():
    """Ninguna celda de las trece hojas totalizaba `exposicion.sin_medir`: el
    KPI que la pantalla pinta en rojo no tenía contraparte en el papel (I4)."""
    wb = _abrir(construir_excel(RESULTADO_SIN_MEDIR, {}))
    ws = wb["08-Conciliacion"]
    assert "SIN MEDIR" in str(ws.cell(8, 1).value), ws.cell(8, 1).value
    formula = str(ws.cell(8, 2).value)
    assert formula.startswith("=SUMIFS("), formula
    assert "'05-Matriz'" in formula and "'06-Individual'" in formula, formula
    assert _suma_sin_medir_del_papel(wb) == RESULTADO_SIN_MEDIR["exposicion"]["sin_medir"]


def test_la_cartera_medida_del_papel_es_la_misma_que_la_de_la_pantalla():
    """B5 se rotulaba «Cartera medida» y sumaba la exposición total, incluidas
    las filas SIN MEDIR: excedía a la celda de la pantalla exactamente en el
    importe sin medir, bajo la misma etiqueta (I4)."""
    wb = _abrir(construir_excel(RESULTADO_SIN_MEDIR, {}))
    ws = wb["08-Conciliacion"]

    assert "Exposición estratificada" in str(ws.cell(5, 1).value)
    assert ws.cell(5, 2).value == "='05-Matriz'!C4+'06-Individual'!C3"
    assert "sin estratificar" in str(ws.cell(6, 1).value).lower()
    assert ws.cell(6, 2).value == 20000.0
    assert ws.cell(7, 2).value == "=B5+B6"

    etiqueta = str(ws.cell(9, 1).value)
    assert etiqueta.startswith("Cartera medida"), etiqueta
    assert "sin medir" in etiqueta.lower(), etiqueta
    assert ws.cell(9, 2).value == "=B5-B8"

    estratificada = 100000.0 + 50000.0
    assert estratificada + 20000.0 == RESULTADO_SIN_MEDIR["exposicion"]["total"]
    assert estratificada - RESULTADO_SIN_MEDIR["exposicion"]["sin_medir"] == \
        CARTERA_MEDIDA_DE_LA_PANTALLA


def test_la_conciliacion_sigue_comparando_el_archivo_contra_los_eeff():
    """La partida conciliatoria y el estado no cambian de significado (I4)."""
    ws = _abrir(construir_excel(RESULTADO_SIN_MEDIR, {}))["08-Conciliacion"]
    assert ws.cell(2, 2).value == 150000.0
    assert ws.cell(3, 2).value == "=SaldoContable"
    assert ws.cell(4, 2).value == "=B2-B3"
    assert ws.cell(10, 2).value == '=IF(ABS(B4)<0.01,"CUADRA","DIFERENCIA")'


# ---------------------------------------------------------------------------
# Guardas del libro completo: el papel tiene que abrir sin reparación y sus
# fórmulas tienen que poder recalcular.
# ---------------------------------------------------------------------------

FUNCIONES_PROHIBIDAS = ("INDIRECT(", "OFFSET(", "IFERROR(", "DESREF(", "SI.ERROR(")

RESULTADOS_A_VALIDAR = [
    (RESULTADO, PARAMETROS_CORRIDA),
    (RESULTADO_CON_RUIDO, {}),
    (RESULTADO_SUSTITUTA, PARAMETROS_SUSTITUTA),
    (RESULTADO_ACOTADO, {}),
    (RESULTADO_SIN_MEDIR, {}),
    (RESULTADO_CENTAVO, {}),
    ({}, {}),
]


def _sumar_corridas_nuevas():
    """Las corridas de C2, I5 e I10 entran a las mismas guardas estructurales."""
    RESULTADOS_A_VALIDAR.extend([
        (RESULTADO_POLITICA_PARCIAL, {}),
        ({**RESULTADO, "control_corte_intermedio": CONTROL_INCONSISTENTE}, {}),
        (RESULTADO_TRIBUTARIO, {}),
    ])


def _formulas(wb):
    for hoja in wb.sheetnames:
        for fila in wb[hoja].iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield hoja, c.coordinate, c.value


def test_ninguna_formula_del_libro_apunta_a_una_hoja_inexistente():
    """Una referencia rota es justo lo que hace que Excel pida reparar."""
    referencia = re.compile(r"'([^']+)'!\$?([A-Z]+)\$?(\d+)")
    for resultado, parametros in RESULTADOS_A_VALIDAR:
        wb = _abrir(construir_excel(resultado, parametros))
        for hoja, celda, formula in _formulas(wb):
            for destino, col, fila in referencia.findall(formula):
                assert destino in wb.sheetnames, f"{hoja}!{celda} apunta a '{destino}'"
                assert int(fila) >= 1


def test_el_libro_no_usa_funciones_prohibidas():
    for resultado, parametros in RESULTADOS_A_VALIDAR:
        wb = _abrir(construir_excel(resultado, parametros))
        for hoja, celda, formula in _formulas(wb):
            mayus = formula.upper()
            for prohibida in FUNCIONES_PROHIBIDAS:
                assert prohibida not in mayus, f"{hoja}!{celda} usa {prohibida}"
            assert formula.count('"') % 2 == 0, f"{hoja}!{celda} deja una comilla abierta"
            assert formula.count("(") == formula.count(")"), f"{hoja}!{celda} desbalancea paréntesis"


def test_los_nombres_definidos_resuelven_a_una_celda_existente():
    for resultado, parametros in RESULTADOS_A_VALIDAR:
        wb = _abrir(construir_excel(resultado, parametros))
        for nombre in wb.defined_names:
            hoja, celda = next(wb.defined_names[nombre].destinations)
            assert hoja in wb.sheetnames, f"{nombre} apunta a la hoja '{hoja}'"
            assert wb[hoja][celda] is not None


def test_el_libro_se_construye_con_un_resultado_vacio():
    """El exportador nunca debe reventar: una corrida sin datos da un papel que
    dice que no los hay, no un 500 al descargar."""
    wb = _abrir(construir_excel({}, {}))
    assert len(wb.sheetnames) == 13


# ---------------------------------------------------------------------------
# C2 — La banda cuya política no se ingresó dice SIN COMPARAR, no 0 %.
# ---------------------------------------------------------------------------

RESULTADO_POLITICA_PARCIAL = {
    **RESULTADO,
    "politica": {
        "filas": [
            {"banda": "Por vencer", "banda_origen": "Por vencer", "exposicion": 80000.0,
             "tasa_politica": 0.01, "provision_politica": 800.0, "ecl": 800.0,
             "tasa_observada": 0.01, "sin_comparar": False, "diferencia": 0.0},
            {"banda": "0 a 30 días", "banda_origen": "0 a 30 días", "exposicion": 20000.0,
             "tasa_politica": None, "provision_politica": None, "ecl": 1000.0,
             "tasa_observada": 0.05, "sin_comparar": True, "diferencia": None},
        ],
        "provision_politica_total": 800.0, "diferencia_bruta": 0.0,
        "bandas_sin_politica": ["0 a 30 días"], "politica_declarada": False,
    },
}


def test_la_banda_sin_politica_del_cliente_dice_sin_comparar_y_no_cero():
    ws = _abrir(construir_excel(RESULTADO_POLITICA_PARCIAL, {}))["07-Politica"]
    # Fila 2: la banda que sí tiene política se compara con fórmula.
    assert ws["D2"].value == 0.01
    assert ws["E2"].value == "=C2*D2"
    # Fila 3: sin política ingresada no hay 0 % ni provisión calculada.
    assert ws["D3"].value == "SIN COMPARAR"
    assert ws["E3"].value == "SIN COMPARAR"
    assert ws["G3"].value == "SIN COMPARAR"


def test_la_hoja_de_politica_declara_que_bandas_no_se_pudieron_comparar():
    ws = _abrir(construir_excel(RESULTADO_POLITICA_PARCIAL, {}))["07-Politica"]
    textos = " ".join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value)
    assert "no fue proporcionada" in textos
    assert "0 a 30 días" in textos


# ---------------------------------------------------------------------------
# I5 — El control del corte intermedio llega al papel.
# ---------------------------------------------------------------------------

CONTROL_INCONSISTENTE = {
    "documentos_cohorte": 120, "saldo_cohorte": 500000.0,
    "vivos_en_intermedio": 80, "saldo_en_intermedio": 200000.0, "saldo_en_actual": 150000.0,
    "permanencia_intermedia": 0.6666666666666666,
    "inconsistencias": [
        {"documento": "F-77", "tipo": "reaparece_tras_desaparecer", "saldo_cohorte": 9000.0,
         "saldo_intermedio": 0.0, "saldo_actual": 4000.0},
        {"documento": "F-88", "tipo": "remanente_mayor_que_intermedio", "saldo_cohorte": 3000.0,
         "saldo_intermedio": 1000.0, "saldo_actual": 2500.0},
    ],
    "inconsistencias_total": 2, "inconsistencias_importe": 6500.0, "consistente": False,
}


def _texto_de(ws):
    return " ".join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None)


def test_las_fuentes_declaran_el_control_del_corte_intermedio():
    resultado = {**RESULTADO, "control_corte_intermedio": CONTROL_INCONSISTENTE}
    ws = _abrir(construir_excel(resultado, {}))["02-Fuentes"]
    texto = _texto_de(ws)
    assert "Control del corte intermedio" in texto
    assert "F-77" in texto and "F-88" in texto
    assert "reaparece" in texto


def test_una_cohorte_consistente_tambien_se_declara_en_el_papel():
    consistente = {**CONTROL_INCONSISTENTE, "inconsistencias": [], "inconsistencias_total": 0,
                   "inconsistencias_importe": 0.0, "consistente": True}
    ws = _abrir(construir_excel({**RESULTADO, "control_corte_intermedio": consistente}, {}))["02-Fuentes"]
    texto = _texto_de(ws)
    assert "Control del corte intermedio" in texto
    assert "Sin inconsistencias" in texto


def test_sin_control_registrado_el_papel_lo_dice_en_vez_de_callarlo():
    ws = _abrir(construir_excel(RESULTADO, {}))["02-Fuentes"]
    assert "no se registró" in _texto_de(ws)


# ---------------------------------------------------------------------------
# I10 — 09-Tributario compara contra el tope del 10 % y declara lo que no puede
# calcular.
# ---------------------------------------------------------------------------

RESULTADO_TRIBUTARIO = {
    **RESULTADO,
    "tributario": {"limite_ejercicio_1pct": 1000.0, "tope_acumulado_10pct": 10000.0,
                   "excede_tope_acumulado": False, "exceso_sobre_tope_acumulado": 0.0,
                   "provision_del_ejercicio": None, "limite_ejercicio_verificable": False,
                   "nota": "Límites de deducción (LORTI art. 10 num. 11)."},
}


def test_el_exceso_no_deducible_se_mide_contra_el_tope_acumulado_del_10_por_ciento():
    ws = _abrir(construir_excel(RESULTADO_TRIBUTARIO, {}))["09-Tributario"]
    # La base es la cartera ESTRATIFICADA (08-Conciliacion B5), no el saldo
    # contable completo: el tope se aplica sobre la cartera a la que se refiere
    # la provisión, y sobre lo que no se estratificó no se midió pérdida.
    assert ws["B3"].value == "='08-Conciliacion'!B5*0.1"
    assert "10 %" in ws["A3"].value
    assert "ESTRATIFICADA" in ws["A3"].value
    assert ws["B4"].value == "=MAX(0,B2-B3)"
    assert "10 %" in ws["A4"].value
    assert "no deducible" in ws["A4"].value.lower()


def test_el_uno_por_ciento_queda_como_referencia_del_limite_anual_no_como_comparacion():
    ws = _abrir(construir_excel(RESULTADO_TRIBUTARIO, {}))["09-Tributario"]
    assert ws["B5"].value == "='08-Conciliacion'!B5*0.01"
    assert "ejercicio" in ws["A5"].value.lower()
    # Ninguna fórmula de la hoja resta el 1 % de la PCE acumulada.
    formulas = [c.value for fila in ws.iter_rows() for c in fila
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=MAX(0,B2-B5)" not in formulas


def test_el_tributario_declara_que_el_limite_anual_no_se_puede_verificar():
    ws = _abrir(construir_excel(RESULTADO_TRIBUTARIO, {}))["09-Tributario"]
    texto = _texto_de(ws)
    assert "NO VERIFICABLE" in texto
    assert "movimiento de la provisión del ejercicio" in texto
    # El tratamiento tributario concilia, nunca sustituye la medición contable.
    assert "no la sustituye" in texto


def test_un_umbral_de_incumplimiento_en_cero_se_imprime_como_cero_y_no_como_730():
    """`... or 730` convertía en 730 el cero guardado en una corrida antigua. El
    papel reproduce la corrida tal como se emitió: si se guardó un 0, dice 0."""
    resultado = {**RESULTADO, "bitacora": {**RESULTADO["bitacora"], "umbral_incumplimiento": 0}}
    ws = _abrir(construir_excel(resultado, {}))["01-Parametros"]
    assert ws["B4"].value == 0


def test_sin_umbral_registrado_se_imprime_el_defecto_del_plan():
    resultado = {**RESULTADO, "bitacora": {**RESULTADO["bitacora"], "umbral_incumplimiento": None}}
    ws = _abrir(construir_excel(resultado, {}))["01-Parametros"]
    assert ws["B4"].value == 730


# Las corridas definidas más abajo en el archivo se suman al final, cuando ya
# existen, para que las cuatro guardas estructurales las recorran también.
_sumar_corridas_nuevas()


def test_ninguna_hoja_solapa_celdas_combinadas():
    """Dos rangos combinados que se pisan son una de las causas de que Excel
    pida reparar el archivo al abrirlo."""
    for resultado, parametros in RESULTADOS_A_VALIDAR:
        wb = _abrir(construir_excel(resultado, parametros))
        for hoja in wb.sheetnames:
            ocupadas = {}
            for rango in [str(r) for r in wb[hoja].merged_cells.ranges]:
                c1, f1, c2, f2 = range_boundaries(rango)
                for col in range(c1, c2 + 1):
                    for fil in range(f1, f2 + 1):
                        anterior = ocupadas.get((col, fil))
                        assert anterior is None, f"{hoja}: {rango} se pisa con {anterior}"
                        ocupadas[(col, fil)] = rango


# ---------------------------------------------------------------------------
# T5 — 02-Fuentes declara el IMPORTE de lo descartado, no solo el conteo
# ---------------------------------------------------------------------------

RESULTADO_DESCARTES = {
    **RESULTADO,
    "bitacora": {
        **RESULTADO["bitacora"],
        "cortes": [
            {"archivo": "cartera_2023.xlsx", "fecha": "2023-12-31", "hoja": "Hoja1",
             "fila_encabezado": 1, "mapeo": {"documento": 1}, "formato_fecha": "dmy",
             "documentos": 10, "duplicados_exactos": 0, "documentos_repetidos": 0,
             "descartados": 2, "descartados_importe": 900000.0, "cartera_no_leida": 900000.0,
             "descartados_por_motivo": [
                 {"motivo": "sin número de documento", "filas": 1, "importe": 900000.0},
                 {"motivo": "saldo cero", "filas": 1, "importe": 0.0}],
             "total": 100000.0},
            {"archivo": "cartera_2024.xlsx", "fecha": "2024-12-31", "hoja": "Hoja1",
             "fila_encabezado": 1, "mapeo": {"documento": 1}, "formato_fecha": "dmy",
             "documentos": 8, "duplicados_exactos": 0, "documentos_repetidos": 0,
             "descartados": 0, "descartados_importe": 0.0, "cartera_no_leida": 0.0,
             "descartados_por_motivo": [], "total": 80000.0},
            {"archivo": "cartera_2025.xlsx", "fecha": "2025-12-31", "hoja": "Hoja1",
             "fila_encabezado": 1, "mapeo": {"documento": 1}, "formato_fecha": "dmy",
             "documentos": 6, "duplicados_exactos": 0, "documentos_repetidos": 0,
             "descartados": 1, "descartados_importe": -800000.0, "cartera_no_leida": -800000.0,
             "descartados_por_motivo": [
                 {"motivo": "sin número de documento", "filas": 1, "importe": -800000.0}],
             "total": 100000.0},
        ],
    },
}


def _encabezados_de(ws) -> list[str]:
    return [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]


def test_fuentes_lleva_el_importe_de_la_cartera_que_no_se_pudo_leer():
    """La hoja conservaba «Descartados» como CONTEO y ninguna columna de
    importe: el dinero que el lector no pudo leer no aparecía en ninguna celda
    de las trece hojas."""
    ws = _abrir(construir_excel(RESULTADO_DESCARTES, {}))["02-Fuentes"]
    encabezados = _encabezados_de(ws)
    columna = next((c for c, t in enumerate(encabezados, start=1)
                    if "no leída" in t.lower() or "no leida" in t.lower()), None)
    assert columna, f"02-Fuentes no declara el importe no leído: {encabezados}"
    assert ws.cell(2, columna).value == 900000.0
    assert ws.cell(4, columna).value == -800000.0


def test_fuentes_desglosa_lo_descartado_por_motivo_con_su_importe():
    """Un conteo no dice cuánta cartera se perdió ni por qué."""
    ws = _abrir(construir_excel(RESULTADO_DESCARTES, {}))["02-Fuentes"]
    textos = " ".join(str(c.value or "") for fila in ws.iter_rows() for c in fila)
    assert "sin número de documento" in textos
    importes = [c.value for fila in ws.iter_rows() for c in fila
                if isinstance(c.value, (int, float))]
    assert 900000.0 in importes and -800000.0 in importes
