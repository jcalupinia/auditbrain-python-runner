"""Lote de Audit Apps por rubro: cada procesador RUBRO corre de punta a punta
por el executor vía los adaptadores genéricos fn(ctx) (AUT-002).

Verifica empíricamente, procesador por procesador, que su Audit App generada
ejecuta los 4 pasos contra el motor determinista real (con el EJEMPLO del propio
procesador) y produce Excel/HTML/excepciones reales. Es la prueba de que el
adaptador genérico sirve para TODO el lote, no solo para uno."""
import io
import json

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif import procesadores
from backend.app.audit_apps import proc_manifest as pm
from backend.app.audit_apps.executor import AuditAppExecutor
from backend.app.audit_apps.manifest import AuditAppManifest

IDS = pm.ids_del_lote()


def test_el_lote_no_esta_vacio_y_excluye_cxc():
    assert len(IDS) >= 15
    assert "cxc_cartera" not in IDS  # tiene app dedicada AUD-CXC-CARTERA


@pytest.mark.parametrize("proc_id", IDS)
def test_manifest_valido_y_apunta_a_adaptadores_genericos(proc_id):
    m = AuditAppManifest.from_dict(pm.manifest_de(proc_id)).validate()
    assert m.id.startswith("AUD-")
    refs = [s.engine_ref for s in m.steps]
    assert all(r.startswith("backend.app.aud.niif.procesadores.app_adapters_rubro.") for r in refs)
    assert [s.id for s in m.steps] == ["extraer", "calcular", "papel_excel", "papel_html"]


@pytest.mark.parametrize("proc_id", IDS)
def test_corre_de_punta_a_punta_con_su_ejemplo(proc_id):
    m = AuditAppManifest.from_dict(pm.manifest_de(proc_id)).validate()
    ej = procesadores.PROCESADORES[proc_id].EJEMPLO
    solicitud = json.dumps({
        "datasets": ej["datasets"], "corte": ej["corte"],
        "parametros": ej.get("parametros", {}),
    })
    res = AuditAppExecutor().run(m, inputs={"solicitud": solicitud.encode("utf-8")},
                                parameters={}, executed_by="a@x.ec")
    # los 4 pasos corrieron; nada quedó en no_disponibles
    assert res.no_disponibles == ()
    assert [s.step_id for s in res.steps] == ["extraer", "calcular", "papel_excel", "papel_html"]
    # salidas reales: Excel abre sin reparaciones, HTML autónomo, excepciones presentes
    excel = res.outputs["papel_excel"]
    assert excel[:2] == b"PK"
    assert load_workbook(io.BytesIO(excel)).sheetnames
    assert res.outputs["papel_html"][:6].lower() in (b"<!doct", b"<html>")
    assert res.outputs["excepciones"]


def test_ids_de_app_unicos_en_el_lote():
    ids = [pm.manifest_de(pid)["id"] for pid in IDS]
    assert len(ids) == len(set(ids))
