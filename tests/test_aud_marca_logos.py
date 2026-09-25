"""Los papeles de trabajo NIIF llevan los logotipos de la firma (AuditConsulting)
y de la plataforma (AUDIT-IA) en Excel, Word, PowerPoint y HTML.

Antes la barra del HTML mostraba una letra «A» en un cuadro dorado y el Excel,
el Word y el PowerPoint no tenían logotipo. El HTML los lleva incrustados en
base64: debe seguir funcionando sin internet.
"""
import io
import re
import zipfile

import pytest

from backend.app.aud.niif import ejercicio_modelo as em
from backend.app.aud.niif.procesadores import PROCESADORES, libro, marca


@pytest.fixture(scope="module")
def papel():
    mod = PROCESADORES["perdidas_incurridas_s11"]
    d = mod.definicion()
    ds, par, corte = em.escenario(mod)
    reg = em._reg(d, mod, ds, par, corte)
    return d, reg


def test_logos_existen_y_son_livianos():
    for nombre in marca.ARCHIVOS:
        b = marca.datos(nombre)
        assert 1_000 < len(b) < 80_000, (nombre, len(b))
        w, h = marca.tamano(nombre)
        assert w > h > 0


def test_excel_portada_con_los_dos_logos(papel):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(libro.xlsx(*papel, [], 1, "MUESTRA")))
    ws = wb.worksheets[0]
    assert ws.title == "00_Inicio"
    assert len(ws._images) == 2
    # En la barra superior (fila 2, 0-based 1), como en el HTML: la firma en B y AUDIT-IA en C.
    anclas = sorted((img.anchor._from.col, img.anchor._from.row) for img in ws._images)
    assert anclas == [(1, 1), (2, 1)]
    with zipfile.ZipFile(io.BytesIO(libro.xlsx(*papel, [], 1, "MUESTRA"))) as z:
        assert sum(n.startswith("xl/media/") for n in z.namelist()) == 2
        # Cada logotipo con su marco (<a:xfrm>): sin él varios visores no lo dibujan.
        dib = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/drawings/drawing"))
        assert dib.count("<pic>") == 2 and dib.count("<spPr><a:xfrm") == 2


def test_word_con_logo_en_el_encabezado(papel):
    from docx import Document

    doc = Document(io.BytesIO(libro.docx(*papel, [], 1, "MUESTRA")))
    enc = doc.sections[0].header
    assert enc._element.xpath(".//pic:pic"), "el encabezado no tiene logotipo"


def test_powerpoint_con_logos(papel):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(libro.pptx(*papel, [], 1, "MUESTRA")))
    diaps = list(prs.slides)
    fotos = [[s for s in d.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE] for d in diaps]
    assert len(fotos[0]) == 2  # portada: firma y AUDIT-IA
    assert all(len(f) >= 1 for f in fotos[1:])
    # Título ensanchado al 16:9 (la plantilla por defecto es 4:3).
    # Sin perder su posición vertical ni su alto (python-pptx los deja en 0 si solo se fija el ancho).
    for d in diaps:
        for ph in d.placeholders:
            assert ph.width == prs.slide_width - 2 * ph.left
            assert ph.top > 0 and ph.height > 0, (ph.name, ph.top, ph.height)


def test_html_logos_incrustados_sin_internet(papel):
    h = libro.html(*papel, [], 1, "MUESTRA").decode()
    assert '<div class="logo" aria-hidden="true">A</div>' not in h
    topbar = h[h.index('<header class="topbar">'):h.index("</header>")]
    assert topbar.count('src="data:image/') == 2
    assert 'alt="AuditConsulting Auditores Cía. Ltda."' in topbar and 'alt="AUDIT-IA"' in topbar
    # Membrete para imprimir («Guardar como PDF»): la barra superior no se imprime.
    assert '<div class="membrete">' in h
    assert ".membrete{display:flex!important}" in h
    # Ninguna imagen se carga de afuera.
    assert not re.search(r'<img[^>]+src="(?!data:)', h)


def test_pdf_estatico_con_membrete(papel):
    h = libro.html(*papel, [], 1, "MUESTRA", para_pdf=True).decode()
    assert '<div class="membrete">' in h and "body.pdf .membrete{display:flex}" in h
