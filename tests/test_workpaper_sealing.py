"""Pruebas del sello inmutable del libro ICT (P2-F, REP-013)."""
import io

from openpyxl import Workbook, load_workbook

from backend.app.ict.sealing import (
    SHEET_NAME,
    apply_seal,
    hash_workbook_bytes,
    verify_seal,
)

SELLO = {
    "version_app": "1.0.0", "version_motor": "motor-1.4.0",
    "timestamp": "2026-09-24T10:00:00Z", "hash_salida": "deadbeef",
    "algoritmo": "sha256", "run_id": "run_123",
    "input_hashes": {"F-101": "aaa", "F-103": "bbb"},
}


def test_apply_seal_crea_hoja_y_propiedades():
    wb = Workbook()
    apply_seal(wb, SELLO, session_data={"ruc": "1791859596001"})
    assert SHEET_NAME in wb.sheetnames
    textos = [str(c.value) for row in wb[SHEET_NAME].iter_rows() for c in row if c.value]
    assert any("deadbeef" in t for t in textos)
    assert any("F-101" in t for t in textos)
    assert any("bbb" in t for t in textos)
    assert "deadbeef" in str(wb.properties.keywords)


def test_hash_workbook_es_determinista():
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    h1, h2 = hash_workbook_bytes(data), hash_workbook_bytes(data)
    assert h1 == h2 and len(h1) == 64


def test_verify_seal_detecta_discrepancia():
    wb = Workbook()
    apply_seal(wb, SELLO)
    assert verify_seal(wb, SELLO) is True
    assert verify_seal(wb, {**SELLO, "hash_salida": "otro"}) is False


def test_texto_que_parece_formula_se_escapa():
    wb = Workbook()
    apply_seal(wb, {**SELLO, "input_hashes": {"=RAID": "xyz"}})
    formulas = [c.value for row in wb[SHEET_NAME].iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not formulas


def test_roundtrip_conserva_el_sello():
    wb = Workbook()
    apply_seal(wb, SELLO)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb2 = load_workbook(buf)
    assert SHEET_NAME in wb2.sheetnames
    assert verify_seal(wb2, SELLO) is True
