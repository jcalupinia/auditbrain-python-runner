"""Cuadro No. 2 · Ingresos IVA vs facturación electrónica.

Los XML de tests/fixtures/sri_facturas son las plantillas OFICIALES del SRI
(factura V1.0.0 a V2.1.0, con datos de relleno como "razonSocial0"): son
públicas y no llevan datos de cliente. Prueban que las rutas de nodos que lee
el parser coinciden con el esquema real en las cuatro versiones.

La nota de crédito y el sobre <autorizacion> se arman a mano siguiendo la
ficha técnica del SRI: el zip del usuario sólo traía el esquema de factura.
"""

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from backend.app.aud.obligaciones_fiscales.libro import facturacion as fac
from backend.app.aud.obligaciones_fiscales.libro.fuentes import (
    construir_hojas_de_casilleros,
)

FIXTURES = Path(__file__).parent / "fixtures" / "sri_facturas"
VERSIONES = ["V1.0.0", "V1.1.0", "V2.0.0", "V2.1.0"]
CLAVE_1 = "1" * 49
CLAVE_2 = "2" * 49


def _xml(porcentaje="4", base="100.00", iva="15.00", fecha="15/03/2025",
         clave=CLAVE_1, export=False, raiz="factura", comprador="CLIENTE FICTICIO"):
    info = "infoFactura" if raiz == "factura" else "infoNotaCredito"
    cod_doc = "01" if raiz == "factura" else "04"
    ce = "<comercioExterior>EXPORTADOR</comercioExterior>" if export else ""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<{raiz} id="comprobante" version="1.0.0">
  <infoTributaria><ruc>0000000000001</ruc><claveAcceso>{clave}</claveAcceso>
    <codDoc>{cod_doc}</codDoc><estab>001</estab><ptoEmi>002</ptoEmi>
    <secuencial>000000123</secuencial></infoTributaria>
  <{info}><fechaEmision>{fecha}</fechaEmision>{ce}
    <razonSocialComprador>{comprador}</razonSocialComprador>
    <totalConImpuestos><totalImpuesto><codigo>2</codigo>
      <codigoPorcentaje>{porcentaje}</codigoPorcentaje>
      <baseImponible>{base}</baseImponible><valor>{iva}</valor>
    </totalImpuesto></totalConImpuestos>
  </{info}>
</{raiz}>""".encode("utf-8")


# ---- lectura de comprobantes ----------------------------------------------

@pytest.mark.parametrize("version", VERSIONES)
def test_lee_las_cuatro_plantillas_oficiales_del_sri(version):
    c = fac.leer_comprobante((FIXTURES / f"factura_{version}.xml").read_bytes())
    assert c.errores == []
    assert c.tipo == "Factura"
    assert c.periodo == "2000-01"
    assert c.lineas, "la plantilla trae totalImpuesto de IVA"


def test_clasifica_la_tarifa_por_codigo_porcentaje():
    assert fac.leer_comprobante(_xml("4")).lineas == [("gravada", 100.0, 15.0)]
    assert fac.leer_comprobante(_xml("0", iva="0.00")).lineas == [("cero", 100.0, 0.0)]
    assert fac.leer_comprobante(_xml("7", iva="0.00")).lineas == [("exento", 100.0, 0.0)]


def test_un_codigo_de_tarifa_desconocido_se_reporta_no_se_adivina():
    c = fac.leer_comprobante(_xml("99"))
    assert c.lineas == []
    assert any("99" in e for e in c.errores)


def test_desenvuelve_el_xml_autorizado_que_descarga_el_sri():
    # El portal del SRI entrega el comprobante dentro de <autorizacion>, como
    # texto CDATA que a su vez trae su propia declaración <?xml ...?>.
    interno = _xml().decode("utf-8")
    envuelto = (
        '<?xml version="1.0" encoding="UTF-8"?><autorizacion>'
        "<estado>AUTORIZADO</estado>"
        f"<numeroAutorizacion>{CLAVE_1}</numeroAutorizacion>"
        f"<comprobante><![CDATA[{interno}]]></comprobante></autorizacion>"
    ).encode("utf-8")
    c = fac.leer_comprobante(envuelto)
    assert c.errores == []
    assert c.lineas == [("gravada", 100.0, 15.0)]


def test_nota_de_credito():
    c = fac.leer_comprobante(_xml(raiz="notaCredito"))
    assert c.tipo == "Nota de crédito"
    assert c.lineas == [("gravada", 100.0, 15.0)]


def test_rechaza_xml_con_entidades():
    # Expansión de entidades ("billion laughs"): un comprobante del SRI jamás
    # trae DOCTYPE, así que se rechaza en vez de expandirse.
    bomba = (
        b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
        b'<!ENTITY b "&a;&a;&a;&a;&a;">]><factura>&b;</factura>'
    )
    c = fac.leer_comprobante(bomba)
    assert c.lineas == []
    assert c.errores


def test_zip_con_varias_facturas_ignora_lo_que_no_es_xml_y_no_duplica(tmp_path):
    z = tmp_path / "facturas.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("a.xml", _xml(clave=CLAVE_1))
        zf.writestr("b.xml", _xml(clave=CLAVE_2))
        zf.writestr("b_otra_vez.xml", _xml(clave=CLAVE_2))
        zf.writestr("leeme.txt", "no soy factura")
    comprobantes, errores = fac.leer_facturas([z])
    assert sorted(c.clave for c in comprobantes) == [CLAVE_1, CLAVE_2]
    assert any("duplicad" in e.lower() for e in errores)


def test_cada_comprobante_cae_en_la_fila_del_cuadro_que_le_toca():
    assert fac.fila_del_cuadro("gravada", exportacion=False) == 6
    assert fac.fila_del_cuadro("cero", exportacion=False) == 8
    assert fac.fila_del_cuadro("exento", exportacion=False) == 15
    assert fac.fila_del_cuadro("cero", exportacion=True) == 12


# ---- hojas del libro --------------------------------------------------------

def _libro(comprobantes):
    wb = Workbook()
    dir_f104 = construir_hojas_de_casilleros(
        wb,
        f104_monthly={"2025-03": {"casilleros": {"401": 1000.0, "411": 900.0}}},
        f103_monthly={},
    )["f104"]
    fac.construir_hojas_facturacion(wb, comprobantes, dir_f104)
    return wb


def test_el_cuadro_lee_la_declaracion_de_datos_f104_por_formula():
    ws = _libro([])[fac.SHEET_CUADRO]
    assert ws["C6"].value.startswith("='DATOS F-104'!")  # cas. 401, tarifa ≠ 0
    assert "'DATOS F-104'!" in ws["E6"].value            # NC = bruto − neto


def test_la_facturacion_electronica_se_suma_desde_la_hoja_de_detalle():
    wb = _libro([fac.leer_comprobante(_xml())])
    g6 = wb[fac.SHEET_CUADRO]["G6"].value
    assert g6.startswith("=SUMIFS(") and f"'{fac.SHEET_DATOS}'!" in g6
    claves = [c.value for c in wb[fac.SHEET_DATOS]["A"]]
    assert CLAVE_1 in claves, "cada comprobante queda trazable en el detalle"


def test_la_diferencia_es_declarado_menos_facturado():
    ws = _libro([])[fac.SHEET_CUADRO]
    assert "F14" in ws["N14"].value and "M14" in ws["N14"].value


def test_un_nombre_de_comprador_con_signo_igual_no_se_vuelve_formula():
    wb = _libro([fac.leer_comprobante(_xml(comprador="=HACK()"))])
    textos = [c.value for fila in wb[fac.SHEET_DATOS].iter_rows() for c in fila
              if isinstance(c.value, str) and "HACK" in c.value]
    assert textos and not any(t.startswith("=") for t in textos)
