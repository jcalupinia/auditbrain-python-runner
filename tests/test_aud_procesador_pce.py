"""Pérdida crediticia esperada, enfoque simplificado NIIF 9: ejemplo de la ficha CXC-PCE-01 y ciclo por HTTP."""
import io

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif.procesadores import pce_simplificada_niif9 as pce
from tests.test_aud_ciclo_evidencia import _leer, _subir
from tests.test_aud_ciclo_http import BASE, FICHA, _accion, _staff_con_proyecto
from tests.test_aud_niif_fichas import _crear, _h


def correr(**param):
    return pce.ejecutar(pce.EJEMPLO["datasets"], {**pce.EJEMPLO["parametros"], **param}, pce.EJEMPLO["corte"])


def test_ejemplo_de_la_ficha():
    r = correr()
    m = next(x for x in r["detalle"]["matriz"] if x["k"] == "t60")
    assert m["hist"] == pytest.approx(0.10)                 # (150 impago + 50 castigado) ÷ 2.000
    assert r["detalle"]["factor"] == pytest.approx(1.03)    # 60 %×1 + 20 %×0,9 + 20 %×1,25
    f20 = next(f for f in r["rows"] if f["id"] == "F-20")
    assert f20["tasa"] == "0.103000" and f20["pce"] == "309.00"
    # F-10 ya tiene más de 360 días y ese tramo no tiene historia: no se inventa la tasa.
    assert any(e["code"] == "TASA_FALTANTE" for e in r["exceptions"])
    assert r["totals"]["castigos"] == "50.00" and r["totals"]["dotacion"] == "359.00"   # 309 − 0 + 50


def test_tasa_fijada_individual_y_descuento():
    datos = {**pce.EJEMPLO["datasets"], "actual": [*pce.EJEMPLO["datasets"]["actual"],
                                                   pce._ej("F-30", "D", "2025-12-20", "1000", tasa_individual="40")]}
    r = pce.ejecutar(datos, {**pce.EJEMPLO["parametros"], "tasas": {"tmax": 100}, "tasaDesc": 12, "plazoBase": 12}, "2025-12-31")
    fila = {f["id"]: f for f in r["rows"]}
    assert fila["F-10"]["tasa"] == "1.000000"                # fijada 100 % (el factor no la pasa de 100 %)
    assert fila["F-30"]["pce"] == pce.r2(1000 * 0.40 / 1.12)  # tasa individual, descontada un año al 12 %
    assert fila["F-20"]["pce"] == pce.r2(3000 * 0.103 / 1.12)
    assert not any(e["code"] == "TASA_FALTANTE" for e in r["exceptions"])


def test_sin_escenarios_avisa_y_pesos_en_cero_no_corre():
    r = pce.ejecutar(pce.EJEMPLO["datasets"], {}, "2025-12-31")
    assert any(e["code"] == "SIN_AJUSTE_PROSPECTIVO" for e in r["exceptions"])
    with pytest.raises(ValueError):
        correr(escBasePeso=0, escOptPeso=0, escPesPeso=0)


def test_cedulas_declaradas_y_completas():
    hojas = pce.hojas(correr(provisionRegistrada=100, provisionInicial=80))
    assert [(h["name"], h["label"]) for h in hojas] == pce.CEDULAS
    for h in hojas:
        assert all(len(f) == len(h["cols"]) for f in h["rows"] + ([h["total"]] if h["total"] else [])), h["name"]


# --- ciclo completo por HTTP -----------------------------------------------------

@pytest.fixture
def disco_temporal(tmp_path, monkeypatch):
    monkeypatch.setenv("AUD_CICLO_DIR", str(tmp_path / "aud_pruebas"))
    monkeypatch.setenv("AUD_CICLO_MIN_LIBRE_MB", "0")


def _modelo_lleno(client, tok, p, req, filas):
    r = client.get(f"{BASE}/pruebas/{p['id']}/modelo/{req}", headers=_h(tok))
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    for f in filas:
        wb["Datos"].append(f)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def test_pce_de_punta_a_punta(client, disco_temporal):
    tok, pid = _staff_con_proyecto(client)
    ficha = _crear(client, tok)
    client.post(f"/api/v1/aud/niif/fichas/{ficha['id']}/estado", headers=_h(tok), json={"estado": "probada"})
    lista = {x["id"]: x for x in client.get(f"{BASE}/procesadores", headers=_h(tok)).json()}
    assert lista["pce_simplificada_niif9"]["ejemplo"]["totales"]["pce"] == "309.00"
    r = client.put(f"{BASE}/fichas/{ficha['id']}/definicion", headers=_h(tok),
                   json={"definicion": lista["pce_simplificada_niif9"]["definicion"], "filas": []})
    assert r.status_code == 200, r.text

    assert client.put(f"{BASE}/proyectos/{pid}/ficha", headers=_h(tok), json=FICHA).status_code == 200
    p = client.post(f"{BASE}/proyectos/{pid}/pruebas", headers=_h(tok), json={"origen": f"ficha:{ficha['id']}"}).json()
    p = _accion(client, tok, p, "research").json()
    p = _accion(client, tok, p, "generate_program").json()
    prog = p["registro"]["program"]
    codigos = [x["code"] for x in prog]
    fuentes = p["registro"]["sources"]
    fuentes[0].update(verified=True, section="NIIF 9 5.5.15, B5.5.35", date="vigente", procedures=codigos)
    fuentes[1].update(verified=True, document="NIA 540", section="párr. 13", date="vigente", procedures=codigos)
    p = _accion(client, tok, p, "approve_program", {"program": prog, "sources": fuentes}).json()
    p = _accion(client, tok, p, "generate_request").json()
    p = _accion(client, tok, p, "approve_request", {"requests": p["registro"]["requests"]}).json()
    assert p["estado"] == "REQUERIMIENTO_APROBADO", p
    assert _leer(client, tok, p)["modelos"] == ["RQ-001", "RQ-002", "RQ-003"]

    # Columnas del modelo: factura, cliente, vence, saldo, segmento, tasa individual, RUC.
    anexos = {"RQ-001": [["F-10", "A", "2024-11-15", 150], ["F-20", "C", "2025-11-15", 3000]],
              "RQ-002": [["F-10", "A", "2024-11-15", 1000], ["F-11", "B", "2024-11-20", 1000]],
              "RQ-003": [["F-11", "B", 50]]}
    for req, filas in anexos.items():
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, f"{req}.xlsx", _modelo_lleno(client, tok, p, req, filas)).status_code == 201
    for req in ("RQ-004", "RQ-005"):
        p = _leer(client, tok, p)
        assert _subir(client, tok, p, req, f"{req}.pdf", b"%PDF").status_code == 201
    p = _leer(client, tok, p)
    assert p["huecos"] == [], p["huecos"]
    arch = {a["requerimiento"]: a["id"] for a in p["archivos"]}
    mapa = lambda tipo: {c["key"]: i for i, c in enumerate(pce.CAMPOS[tipo])}
    parte = lambda req, tipo: [{"fileId": arch[req], "sheet": "Datos", "header": 1, "mapping": mapa(tipo)}]
    p = _accion(client, tok, p, "map_validate", {"datasets": {"actual": parte("RQ-001", "cartera"), "anterior": parte("RQ-002", "cartera"),
                                                              "castigos": parte("RQ-003", "castigos")}}).json()
    assert p["registro"]["validation"]["ok"], p["registro"]["validation"]
    assert p["registro"]["controlTotal"] == "3150.00"
    p = _accion(client, tok, p, "validate", {"evidenceReviewed": True, "evidenceReview": "Cotejé la cartera con el mayor.",
                                             "ledger": "3150.00", "tolerance": "0"}).json()
    # El escenario optimista resta: el ajuste negativo es válido.
    p = _accion(client, tok, p, "configure", {"basis": "Escenarios del plan de negocios y del BCE; tasa de 100 % en impago de más de 360 días.",
                                              "parametros": {**pce.EJEMPLO["parametros"], "tasas": {"tmax": "100"},
                                                             "provisionRegistrada": "200"}}).json()
    assert p["estado"] == "PRUEBA_CONFIGURADA", p
    p = _accion(client, tok, p, "approve_methodology").json()
    p = _accion(client, tok, p, "execute", {}).json()
    assert p["estado"] == "PRUEBA_EJECUTADA", p
    run = p["registro"]["run"]
    # F-20: 3.000 × 10,3 % = 309; F-10: 150 × 100 % = 150 → 459; registrada 200 → ajuste 259.
    assert run["totals"]["pce"] == "459.00" and run["totals"]["ajuste"] == "259.00"
    for fmt in ("xlsx", "docx", "pptx", "html"):
        assert client.get(f"{BASE}/pruebas/{p['id']}/libro?formato={fmt}", headers=_h(tok)).status_code == 200, fmt
    wb = load_workbook(io.BytesIO(client.get(f"{BASE}/pruebas/{p['id']}/libro", headers=_h(tok)).content))
    assert wb["09_Detalle"]["K5"].value.startswith("=IF(")


def test_tasa_cero_observada_se_avisa():
    datos = {**pce.EJEMPLO["datasets"], "anterior": [*pce.EJEMPLO["datasets"]["anterior"], pce._ej("F-12", "B", "2025-01-20", "500")],
             "actual": [*pce.EJEMPLO["datasets"]["actual"], pce._ej("F-40", "E", "2026-01-30", "800")]}
    r = pce.ejecutar(datos, pce.EJEMPLO["parametros"], "2025-12-31")
    assert any(e["code"] == "TASA_CERO" and "Corriente" in e["message"] for e in r["exceptions"])
