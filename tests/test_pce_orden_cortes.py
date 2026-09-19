"""V2 — los tres cortes tienen que llegar en orden cronológico estricto.

POR QUÉ EXISTE. Los archivos y sus fechas viajan por POSICIÓN: índice 0 = la
cohorte (t-2), índice 2 = el corte actual (t). Nada comprobaba que esas tres
fechas fueran crecientes, así que subir los archivos al revés -con sus fechas
correctas, tecleadas en el mismo orden invertido- NO fallaba: la herramienta
medía la exposición sobre el corte más ANTIGUO y derivaba las tasas siguiendo
la cartera hacia atrás en el tiempo. El resultado cambiaba en silencio (el
informe de revisión midió 86.206,90 en vez de 208.943,84) y los únicos rastros
eran dos pendientes indirectos.

Estas pruebas son de punta a punta: desde los archivos que sube el auditor
hasta el código HTTP y hasta la celda de `02-Fuentes`.
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.auth import service as auth_service
from backend.app.auth.models import Role
from backend.app.db.session import SessionLocal, init_db

BASE = "/api/v1/aud/pce-cxc"

FECHAS_EN_ORDEN = ["2023-12-31", "2024-12-31", "2025-12-31"]


def _xlsx(saldo: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    ws.append(["ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), saldo])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


#: Los tres cortes, cada uno con su saldo: 100.000 en t-2, 20.000 en t-1 y
#: 10.000 en t. Al revés, la "cohorte" crece en vez de cobrarse.
def _archivos(orden=(0, 1, 2)):
    cortes = [("cartera_2023.xlsx", _xlsx(100000.0)),
              ("cartera_2024.xlsx", _xlsx(20000.0)),
              ("cartera_2025.xlsx", _xlsx(10000.0))]
    return [("archivos", cortes[i]) for i in orden]


def _token(client, role=Role.user):
    init_db()
    tag = uuid.uuid4().hex[:6]
    email, pw = f"pce-orden-{tag}@ex.com", "Sup3rSecret!"
    db = SessionLocal()
    try:
        auth_service.create_user(db, email=email, password=pw, role=role)
    finally:
        db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": pw})
    return r.json()["access_token"]


def _analizar(client, token, orden, fechas):
    return client.post(f"{BASE}/analizar", files=_archivos(orden),
                       data={"parametros": json.dumps({"fechas": fechas})},
                       headers={"Authorization": f"Bearer {token}"})


def test_los_tres_cortes_al_reves_se_rechazan_con_un_400_accionable(client):
    """El escenario del informe: los tres archivos invertidos, con sus fechas
    correctas tecleadas en el mismo orden invertido."""
    token = _token(client)
    r = _analizar(client, token, (2, 1, 0), list(reversed(FECHAS_EN_ORDEN)))
    assert r.status_code == 400, f"se esperaba 400, se obtuvo {r.status_code}: {r.text}"
    detalle = r.json()["detail"]
    assert "cronológico" in detalle, detalle
    # Accionable: dice qué par está mal y con qué fechas, y qué hacer.
    assert "2024-12-31" in detalle and "2025-12-31" in detalle, detalle


@pytest.mark.parametrize("fechas", [
    # El corte intermedio anterior a la cohorte.
    ["2024-12-31", "2023-12-31", "2025-12-31"],
    # El corte actual anterior al intermedio.
    ["2023-12-31", "2025-12-31", "2024-12-31"],
    # Dos fechas iguales: no hay ventana entre esos dos cortes.
    ["2023-12-31", "2023-12-31", "2025-12-31"],
    ["2023-12-31", "2024-12-31", "2024-12-31"],
])
def test_cualquier_par_fuera_de_orden_se_rechaza(client, fechas):
    """Estricto: dos cortes con la misma fecha tampoco abren una ventana."""
    token = _token(client)
    r = _analizar(client, token, (0, 1, 2), fechas)
    assert r.status_code == 400, f"se esperaba 400 con {fechas}: {r.text}"
    assert "cronológico" in r.json()["detail"]


def test_el_orden_correcto_sigue_calculando(client):
    """La guarda no puede estorbar al camino bueno."""
    token = _token(client)
    r = _analizar(client, token, (0, 1, 2), FECHAS_EN_ORDEN)
    assert r.status_code == 200, r.text
    assert r.json()["ecl_total"] >= 0


def test_02_fuentes_afirma_el_rol_junto_a_la_fecha_de_ese_corte(client):
    """El rol lo fija la posición; sin la fecha en la misma fila, el papel
    afirma «corte actual (t)» sobre un archivo que el lector no puede contrastar.
    """
    token = _token(client)
    r = _analizar(client, token, (0, 1, 2), FECHAS_EN_ORDEN)
    assert r.status_code == 200, r.text
    x = client.get(f"{BASE}/corridas/{r.json()['corrida_id']}/excel",
                   headers={"Authorization": f"Bearer {token}"})
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content))["02-Fuentes"]

    encabezados = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    assert "Rol del corte" in encabezados, encabezados
    assert any("Fecha" in e for e in encabezados), (
        f"02-Fuentes no dice la fecha de cada corte: {encabezados}")
    col_rol = encabezados.index("Rol del corte") + 1
    col_fecha = next(i + 1 for i, e in enumerate(encabezados) if "Fecha" in e)

    esperado = [("cohorte (t-2)", "2023-12-31"),
                ("corte intermedio (t-1)", "2024-12-31"),
                ("corte actual (t)", "2025-12-31")]
    obtenido = [(str(ws.cell(f, col_rol).value or ""), str(ws.cell(f, col_fecha).value or ""))
                for f in (2, 3, 4)]
    assert obtenido == esperado, obtenido


def test_el_servicio_tambien_rechaza_el_desorden_sin_pasar_por_el_router(client):
    """La guarda vive en el servicio, no solo en el endpoint: cualquier otro
    camino que llame a `analizar` tiene que encontrarse con la misma regla."""
    from backend.app.aud.pce_cxc import service

    cortes = [{"nombre": n, "contenido": c, "fecha": f, "hoja": None, "mapeo": None}
              for (_, (n, c)), f in zip(_archivos((2, 1, 0)),
                                        [date(2025, 12, 31), date(2024, 12, 31),
                                         date(2023, 12, 31)])]
    with pytest.raises(ValueError, match="cronológico"):
        service.analizar(cortes, {"fechas": list(reversed(FECHAS_EN_ORDEN))})
