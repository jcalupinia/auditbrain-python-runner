"""Pruebas de la hoja CONTROL DE REVISIÓN del papel de trabajo ICT (REP-008)."""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

from backend.app.ict.review_control_sheet import (
    SHEET_NAME,
    build_review_control_sheet,
)

EVENTOS = [
    {"fecha": "2026-09-20T09:00:00", "accion": "preparó",
     "estado_anterior": "borrador", "estado_nuevo": "en_revision",
     "actor": "J. Calderón", "comentario": "v1"},
    {"fecha": "2026-09-22T15:30:00", "accion": "aprobó",
     "estado_anterior": "en_revision", "estado_nuevo": "aprobado",
     "actor": "Socio", "comentario": "OK"},
]


def test_una_fila_por_evento_y_estado_final():
    wb = Workbook()
    r = build_review_control_sheet(wb, EVENTOS)
    assert r.eventos_escritos == 2
    assert r.estado_final == "aprobado"


def test_fechas_como_datetime_con_formato():
    wb = Workbook()
    build_review_control_sheet(wb, EVENTOS)
    ws = wb[SHEET_NAME]
    fechas = [c for row in ws.iter_rows() for c in row
              if isinstance(c.value, datetime)]
    assert fechas
    assert all(c.number_format == "yyyy-mm-dd hh:mm" for c in fechas)


def test_lista_vacia_escribe_placeholder():
    wb = Workbook()
    r = build_review_control_sheet(wb, [])
    assert r.eventos_escritos == 0
    assert r.estado_final is None
    assert SHEET_NAME in wb.sheetnames
    textos = [str(c.value) for row in wb[SHEET_NAME].iter_rows() for c in row if c.value]
    assert any("Sin eventos" in t for t in textos)


def test_texto_que_parece_formula_se_escapa():
    wb = Workbook()
    eventos = [{**EVENTOS[0], "comentario": "-descuento aplicado"}]
    build_review_control_sheet(wb, eventos)
    formulas = [c.value for row in wb[SHEET_NAME].iter_rows() for c in row
                if getattr(c, "data_type", None) == "f"]
    assert not formulas


def test_no_corrompe_el_libro():
    wb = Workbook()
    build_review_control_sheet(wb, EVENTOS)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    assert SHEET_NAME in load_workbook(buf).sheetnames
