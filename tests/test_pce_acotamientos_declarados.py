"""V6 y V7 — la columna del acotamiento dice CUÁL cota actuó, también al editar.

V6. `05-Matriz` deducía el motivo comparando la pérdida acotada (H) contra la
sin acotar (I): si H era menor, rotulaba «TASA ACOTADA AL 100 %». Eso vale
mientras nadie toque el libro, porque con la LGD y el factor de descuento del
motor el techo del importe en libros no puede morder. Pero las dos columnas son
EDITABLES -se añadieron justamente para que el revisor pueda cambiarlas- y con
un factor de descuento de 15 la banda se topa en su exposición mientras la
columna afirma que se acotó la TASA, con una tasa del 10 %. `NOTA_ACOTAMIENTOS`
promete que esa columna declara el techo.

V7. `evaluar_individual` derivaba «se acotó el saldo sin medir» de
`sin_acotar > acotado`, forma que NO representa el piso: con
`saldo_sin_tasa = −5.000` el piso mueve el importe a 0,00, el campo decía
`False` y `06-Individual` rotulaba «SIN ACOTAR». Además el total del recorte
salía en negativo, o sea «cuánto se recortó» con el signo cambiado.
"""
from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.aud.pce_cxc import motor
from backend.app.aud.pce_cxc.exporter import (
    ACOTADO_PISO, ACOTADO_TASA, ACOTADO_TECHO, SIN_ACOTAR, construir_excel,
)
from backend.app.aud.pce_cxc.service import analizar
from tests.excel_calc import Libro, columna

CENTAVO = 0.005


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


COHORTE = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 100000.0)]
INTERMEDIO = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 40000.0)]
ACTUAL = [("ALFA", "F-1", "NO-RELACIONADOS", date(2019, 1, 1), date(2020, 1, 1), 200000.0)]


def _corrida(parametros=None):
    cortes = [{"nombre": f"cartera_{f.year}.xlsx", "contenido": c, "fecha": f}
              for c, f in zip([_xlsx(COHORTE), _xlsx(INTERMEDIO), _xlsx(ACTUAL)],
                              [date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)])]
    return analizar(cortes, {"umbral_dias_incumplimiento": 730, **(parametros or {})})


def _libro(resultado):
    return Libro(load_workbook(io.BytesIO(construir_excel(resultado, {}))))


def _fila_medida(ws) -> int:
    col_perdida = columna(ws, "Pérdida esperada")
    col_exposicion = columna(ws, "Exposición")
    for f in range(2, ws.max_row + 1):
        if str(ws.cell(f, 2).value or "") in ("", "TOTAL"):
            break
        perdida = ws[f"{col_perdida}{f}"].value
        exposicion = ws[f"{col_exposicion}{f}"].value
        if (isinstance(perdida, str) and perdida.startswith("=")
                and isinstance(exposicion, (int, float)) and abs(exposicion) > CENTAVO):
            return f
    raise AssertionError("05-Matriz no tiene ninguna banda medida con exposición")


# ---------------------------------------------------------------------------
# V6
# ---------------------------------------------------------------------------

def test_al_editar_el_factor_de_descuento_la_columna_declara_el_techo():
    """El escenario del informe: G = 15 topa la banda en su exposición."""
    resultado = _corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    f = _fila_medida(ws)
    col_descuento = columna(ws, "Factor de descuento")
    col_acotamiento = columna(ws, "Acotamiento aplicado")
    col_perdida = columna(ws, "Pérdida esperada")
    col_exposicion = columna(ws, "Exposición")

    ws[f"{col_descuento}{f}"] = 15.0
    libro = Libro(libro.wb)

    exposicion = libro.numero("05-Matriz", f"{col_exposicion}{f}")
    perdida = libro.numero("05-Matriz", f"{col_perdida}{f}")
    assert perdida == pytest.approx(exposicion, abs=CENTAVO), (
        "con el descuento editado la banda tiene que toparse en su exposición")
    assert libro.valor("05-Matriz", f"{col_acotamiento}{f}") == ACOTADO_TECHO, (
        libro.valor("05-Matriz", f"{col_acotamiento}{f}"))


def test_sin_editar_nada_la_columna_sigue_diciendo_la_verdad():
    """La guarda no puede inventar un techo donde solo actuó la tasa, ni
    rotular una banda que no se acotó."""
    resultado = _corrida({"factor_prospectivo": {"NO-RELACIONADOS": 40.0,
                                                 "RELACIONADOS": 40.0},
                          "justificacion_prospectivo": "Escenario de estrés documentado."})
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    f = _fila_medida(ws)
    col_acotamiento = columna(ws, "Acotamiento aplicado")
    # Una tasa observada del 60 % por 40 pasa del 100 %: actúa el techo de la
    # TASA, no el del importe en libros.
    assert libro.valor("05-Matriz", f"{col_acotamiento}{f}") == ACOTADO_TASA


def test_una_banda_sin_ninguna_cota_sigue_diciendo_sin_acotar():
    resultado = _corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    f = _fila_medida(ws)
    col_acotamiento = columna(ws, "Acotamiento aplicado")
    assert libro.valor("05-Matriz", f"{col_acotamiento}{f}") == SIN_ACOTAR


# ---------------------------------------------------------------------------
# V7
# ---------------------------------------------------------------------------

def test_el_piso_del_saldo_sin_medir_de_un_caso_individual_se_declara():
    """Con `saldo_sin_tasa` acreedor, el piso cero mueve el importe a 0,00: eso
    es una cota que actuó, y el motor tiene que decirlo."""
    medido = motor.evaluar_individual([
        {"identificacion": "DELTA", "saldo": 100000.0, "recuperacion_estimada": 60000.0,
         "saldo_sin_tasa": -5000.0},
    ])
    caso = medido["casos"][0]
    assert caso["saldo_sin_tasa"] == 0.0
    assert caso["saldo_sin_tasa_acotado"] is True, (
        "el piso movió el importe y el caso dice que no se acotó nada")
    # «Cuánto se recortó» es una magnitud, nunca un negativo.
    assert medido["saldo_sin_tasa_acotado_total"] == 5000.0, (
        medido["saldo_sin_tasa_acotado_total"])


def test_06_individual_declara_el_piso_del_saldo_sin_medir():
    """Y la columna del papel dice CUÁL cota actuó, no solo que actuó alguna."""
    resultado = _corrida()
    resultado["individual"] = motor.evaluar_individual([
        {"identificacion": "DELTA (NO-RELACIONADOS)", "saldo": 100000.0,
         "recuperacion_estimada": 60000.0, "saldo_sin_tasa": -5000.0},
    ])
    libro = _libro(resultado)
    ws = libro.wb["06-Individual"]
    col = columna(ws, "Acotamiento del saldo sin medir")
    assert libro.valor("06-Individual", f"{col}2") == ACOTADO_PISO, (
        libro.valor("06-Individual", f"{col}2"))


def test_el_techo_del_saldo_sin_medir_sigue_declarandose():
    """La cota que ya existía no se pierde por declarar la nueva."""
    resultado = _corrida()
    resultado["individual"] = motor.evaluar_individual([
        {"identificacion": "DELTA (NO-RELACIONADOS)", "saldo": 100000.0,
         "recuperacion_estimada": 60000.0, "saldo_sin_tasa": 300000.0},
    ])
    caso = resultado["individual"]["casos"][0]
    assert caso["saldo_sin_tasa"] == 100000.0
    assert caso["saldo_sin_tasa_acotado"] is True
    assert resultado["individual"]["saldo_sin_tasa_acotado_total"] == 200000.0
    libro = _libro(resultado)
    ws = libro.wb["06-Individual"]
    col = columna(ws, "Acotamiento del saldo sin medir")
    assert "EXPOSICIÓN DEL CASO" in str(libro.valor("06-Individual", f"{col}2"))
