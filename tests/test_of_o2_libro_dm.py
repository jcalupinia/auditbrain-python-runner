"""O2 del Libro DM: el armado aislado en subproceso (``OF_O2_SUBPROCESO``) produce el MISMO
Excel que el armado en proceso (0 diferencias celda por celda).

Patrón calcado de ``test_aud_of_service.test_process_job_end_to_end_with_real_pdf``: siembra un
job con el F-104 real, corre ``process_job`` en foreground con el flag APAGADO (camino actual) y
luego ENCENDIDO (subproceso), y compara los dos ``output.xlsx``.
"""
import uuid
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.auth import service as auth_service
from backend.app.auth.models import Role, User
from backend.app.aud.obligaciones_fiscales import file_storage, jobs, service as of_service
from backend.app.aud.obligaciones_fiscales.models import ToolJob
from backend.app.context import service as ctx_service
from backend.app.db.session import SessionLocal, init_db

from tests._mayor_fixtures import mayor_xlsx

FIXTURES = Path(__file__).parent / "fixtures" / "obligaciones_fiscales"

# Mayor sintético (cifras inventadas): compras con IVA, ventas y una retención, para que haya
# movimientos y clasificación en BD que el subproceso debe re-leer.
_MAYOR = [
    ["1.1.5.1.1", "IVA sobre Compras", "2025-01-05", "COM 1",
     "FAC 001", "9999999999001", "PROVEEDOR DEMO S.A.", "", "COMPRA", 12.00, None, 12.00],
    ["4.1.1.1.1", "Ventas 12%", "2025-01-10", "FAC 2",
     "FAC 100", "9999999999002", "CLIENTE DEMO S.A.", "", "VENTA", None, 500.00, 500.00],
    ["2.1.7.3.2", "Ret. Fuente 1% Bienes", "2025-01-10", "RET 3",
     "RET 100", "9999999999002", "CLIENTE DEMO S.A.", "", "RETENCION", None, 5.00, 5.00],
]


@pytest.fixture(autouse=True)
def _db():
    init_db()
    yield


@pytest.fixture()
def tmp_root(tmp_path, monkeypatch):
    monkeypatch.setenv("AUD_OF_TMP_DIR", str(tmp_path))
    from importlib import reload

    from backend.app.core import config

    reload(config)
    reload(file_storage)
    yield tmp_path


def _mk_admin_project():
    db = SessionLocal()
    try:
        tag = uuid.uuid4().hex[:6]
        u = auth_service.create_user(db, email=f"a-{tag}@ex.com", password="Sup3rSecret!", role=Role.admin)
        u = ctx_service.ensure_user_has_organization(db, u)
        c = ctx_service.create_client(db, organization_id=u.organization_id, name=f"Cliente-{tag}")
        p = ctx_service.create_project(
            db, organization_id=u.organization_id, client_id=c.id, name=f"Aud-{tag}", module_code="AUD",
        )
        ctx_service.add_project_member(db, p.id, u.id, "lead")
        return u.id, p.id
    finally:
        db.close()


def _celdas(xlsx_bytes: bytes) -> dict:
    """{(hoja, fila, col): valor} de todas las celdas no vacías de todas las hojas."""
    wb = load_workbook(BytesIO(xlsx_bytes))
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    out[(name, r, c)] = v
    wb.close()
    return out


def _sembrar_job_con_f104(tmp_root):
    user_id, project_id = _mk_admin_project()
    db = SessionLocal()
    try:
        user = db.get(User, user_id)
        job = of_service.create_job(
            db, user=user, project_id=project_id,
            cliente_name="NEGOCIOS MORACOSTA S.A.", period_label="2025",
        )
        job_id = job.id
    finally:
        db.close()
    job_dir = file_storage.create_job_dir(job_id)
    file_storage.save_input(job_dir, "f104", "f104_enero.pdf", (FIXTURES / "f104_enero.pdf").read_bytes())
    return job_id, job_dir


def _correr_y_leer(job_id, job_dir) -> bytes:
    jobs.process_job(job_id)
    db = SessionLocal()
    try:
        final = db.get(ToolJob, job_id)
        assert final.status == "done", f"status={final.status}, error={final.error_message}"
    finally:
        db.close()
    return file_storage.output_path(job_dir).read_bytes()


def test_o2_libro_dm_identico_a_inline(tmp_root, monkeypatch):
    job_id, job_dir = _sembrar_job_con_f104(tmp_root)

    monkeypatch.setenv("OF_O2_SUBPROCESO", "")        # inline (camino actual)
    assert not jobs._o2_activo()
    celdas_inline = _celdas(_correr_y_leer(job_id, job_dir))

    monkeypatch.setenv("OF_O2_SUBPROCESO", "true")    # subproceso aislado
    assert jobs._o2_activo()
    celdas_o2 = _celdas(_correr_y_leer(job_id, job_dir))

    difs = [k for k in set(celdas_inline) | set(celdas_o2)
            if celdas_inline.get(k) != celdas_o2.get(k)]
    assert not difs, f"{len(difs)} celdas difieren inline vs O2; p.ej. {difs[:5]}"
    assert len(celdas_inline) > 50  # se armó un libro real, no vacío


def test_o2_libro_dm_identico_con_mayor_y_clasificacion(tmp_root, monkeypatch):
    """Con un mayor real: clasifica (llena la clasificación en BD) y verifica que el subproceso
    la re-lee desde su propia sesión y arma un libro idéntico al inline."""
    job_id, job_dir = _sembrar_job_con_f104(tmp_root)
    file_storage.save_input(job_dir, "mayor_general", "mayor.xlsx", mayor_xlsx(_MAYOR))

    jobs.clasificar_mayor_job(job_id)  # llena la clasificación en BD (como el camino real)
    db = SessionLocal()
    try:
        from backend.app.aud.obligaciones_fiscales.mayor.clasificacion_service import clasificacion_de_job
        assert len(clasificacion_de_job(db, job_id=job_id)) >= 1, "no se clasificó ninguna cuenta"
    finally:
        db.close()

    monkeypatch.setenv("OF_O2_SUBPROCESO", "")
    celdas_inline = _celdas(_correr_y_leer(job_id, job_dir))
    monkeypatch.setenv("OF_O2_SUBPROCESO", "true")
    celdas_o2 = _celdas(_correr_y_leer(job_id, job_dir))

    difs = [k for k in set(celdas_inline) | set(celdas_o2)
            if celdas_inline.get(k) != celdas_o2.get(k)]
    assert not difs, f"{len(difs)} celdas difieren inline vs O2 (con mayor); p.ej. {difs[:5]}"


def _clasificacion_resumen(job_id) -> dict:
    """{codigo: (categoria_sugerida, categoria_final, n_movimientos)} de la clasificación en BD."""
    from backend.app.aud.obligaciones_fiscales.mayor.clasificacion_service import clasificacion_de_job
    db = SessionLocal()
    try:
        return {f.codigo_cuenta: (f.categoria_sugerida, f.categoria_final, f.n_movimientos)
                for f in clasificacion_de_job(db, job_id=job_id)}
    finally:
        db.close()


def test_o2_clasificar_mayor_identico_a_inline(tmp_root, monkeypatch):
    """La fase 1 (lee los mayores grandes + clasifica) aislada en subproceso deja la MISMA
    clasificación en BD que en proceso. El payload (resultados/perfiles) debe viajar por la cola
    spawn sin perder nada."""
    job_id, job_dir = _sembrar_job_con_f104(tmp_root)
    file_storage.save_input(job_dir, "mayor_general", "mayor.xlsx", mayor_xlsx(_MAYOR))

    monkeypatch.setenv("OF_O2_SUBPROCESO", "")
    jobs.clasificar_mayor_job(job_id)
    inline = _clasificacion_resumen(job_id)
    assert inline, "la clasificación inline quedó vacía"

    monkeypatch.setenv("OF_O2_SUBPROCESO", "true")
    jobs.clasificar_mayor_job(job_id)   # reemplaza la clasificación del mismo job
    o2 = _clasificacion_resumen(job_id)

    assert o2 == inline, f"la clasificación difiere inline vs O2:\n inline={inline}\n o2={o2}"
