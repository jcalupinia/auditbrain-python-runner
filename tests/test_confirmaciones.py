import copy
import unittest
from io import BytesIO

from backend.app.aud.confirmaciones_saldos.engine import calculate
from backend.app.aud.confirmaciones_saldos.exports import build_xlsx, build_docx, build_html, schedules, letter_email_html


def payload():
    return {
        'context': {
            'client': 'Comercial Andina S.A.', 'client_ruc': '1790011110001', 'country': 'Ecuador',
            'currency': 'USD', 'year': 2026, 'cutoff': '2026-12-31', 'visit': 'Final',
            'preparer': 'J. Pérez', 'reviewer': 'M. Torres', 'firm': 'audit_consulting',
            'signatory': 'Ana Vaca', 'signatory_role': 'Gerente Financiera',
            'auditor_email': 'confirmaciones@auditconsulting.ec',
            'auditor_address': 'Av. Amazonas N30-01, Quito', 'auditor_phone': '+593 2 111 2222',
            'response_deadline': '2027-01-20',
        },
        'defaults': {},
        'items': [
            {'id': 'B-01', 'type': 'bancos', 'entity': 'Banco Pichincha C.A.',
             'contact_email': 'confirma@pichincha.com', 'account_ref': '2100', 'amount': '152340.55'},
            {'id': 'C-01', 'type': 'cuentas_por_cobrar', 'entity': 'Distribuidora Sur',
             'contact_email': 'pagos@dsur.ec', 'amount': '48200.00', 'reference': 'Mayor fila 12'},
            {'id': 'P-01', 'type': 'proveedores', 'entity': 'Importadora Global',
             'contact_email': 'cxp@global.com'},
        ],
        'coverage': [{'type': 'cuentas_por_cobrar', 'ledger_balance': '60000'}],
        'tolerance': '0.01',
    }


class ConfirmacionesTests(unittest.TestCase):
    def test_basic_counts_and_totals(self):
        r = calculate(payload())
        self.assertEqual(r['totals']['count'], 3)
        self.assertEqual(r['totals']['with_email'], 3)
        self.assertEqual(r['totals']['without_email'], 0)
        # 152340.55 + 48200.00 (proveedor en blanco no suma importe)
        self.assertEqual(r['totals']['total_sampled'], '200540.55')
        self.assertEqual(r['status'], 'BORRADOR PARA REVISION')
        self.assertEqual(len(r['letters']), 3)
        self.assertEqual(len(r['dispatch']), 3)
        self.assertEqual(len(r['register']), 3)

    def test_default_method_per_type_and_override(self):
        r = calculate(payload())
        methods = {it['id']: it['method'] for it in r['items']}
        # el formato de la firma es «en blanco» (el tercero declara el saldo)
        self.assertEqual(methods['B-01'], 'en_blanco')
        self.assertEqual(methods['P-01'], 'en_blanco')
        p = payload(); p['defaults']['method'] = 'negativa'; p['items'][2]['amount'] = '5000'
        r2 = calculate(p)
        # el default global manda sobre el default del tipo
        self.assertEqual(r2['items'][0]['method'], 'negativa')
        self.assertEqual(r2['items'][2]['method'], 'negativa')
        p2 = payload(); p2['items'][0]['method'] = 'en_blanco'
        self.assertEqual(calculate(p2)['items'][0]['method'], 'en_blanco')

    def test_blank_method_hides_amount(self):
        p = payload()
        p['items'][1]['method'] = 'en_blanco'   # CxC en blanco
        p['defaults']['include_response_slip'] = True
        r = calculate(p)
        cxc = next(l for l in r['letters'] if l['id'] == 'C-01')
        self.assertFalse(any('48,200' in b or '48200' in b for b in cxc['blocks']))
        self.assertIsNotNone(cxc['response_lines'])
        self.assertTrue(any('Saldo/valor a la fecha de corte' in x for x in cxc['response_lines']))

    def test_response_slip_off_by_default(self):
        # el formato de la firma no lleva recuadro; solo aparece si se solicita
        r = calculate(payload())
        self.assertTrue(all(l['response_lines'] is None for l in r['letters']))

    def test_positive_reveals_amount_negative_has_no_response_block(self):
        p = payload(); p['items'][1]['method'] = 'positiva'; p['defaults']['include_response_slip'] = True
        r = calculate(p)
        cxc = next(l for l in r['letters'] if l['id'] == 'C-01')
        self.assertTrue(any('USD 48,200.00' in b for b in cxc['blocks']))
        self.assertIsNotNone(cxc['response_lines'])
        p2 = payload(); p2['items'][1]['method'] = 'negativa'; p2['defaults']['include_response_slip'] = True
        neg = next(l for l in calculate(p2)['letters'] if l['id'] == 'C-01')
        self.assertIsNone(neg['response_lines'])   # negativa nunca lleva recuadro
        self.assertTrue(any('no coincide' in b for b in neg['blocks']))

    def test_amount_required_for_cxc_and_proveedor_when_not_blank(self):
        p = payload(); del p['items'][1]['amount']; p['items'][1]['method'] = 'positiva'
        with self.assertRaises(ValueError):
            calculate(p)
        # proveedor positivo sin importe también se rechaza
        p2 = payload(); p2['items'][2]['method'] = 'positiva'
        with self.assertRaises(ValueError):
            calculate(p2)
        # bancos/seguros/abogados no exigen importe
        p3 = payload(); del p3['items'][0]['amount']
        self.assertEqual(calculate(p3)['items'][0]['amount'], None)

    def test_duplicate_and_invalid_type_rejected(self):
        p = payload(); p['items'][1]['id'] = 'B-01'
        with self.assertRaises(ValueError):
            calculate(p)
        p2 = payload(); p2['items'][0]['type'] = 'otro'
        with self.assertRaises(ValueError):
            calculate(p2)

    def test_context_validation(self):
        for key in ('client', 'firm', 'signatory', 'auditor_email', 'response_deadline'):
            p = payload(); p['context'][key] = ''
            with self.assertRaises(ValueError):
                calculate(p)
        p = payload(); p['context']['firm'] = 'otra'
        with self.assertRaises(ValueError):
            calculate(p)
        p = payload(); p['context']['visit'] = 'X'
        with self.assertRaises(ValueError):
            calculate(p)

    def test_email_validation(self):
        p = payload(); p['context']['auditor_email'] = 'no-es-correo'
        with self.assertRaises(ValueError):
            calculate(p)
        p2 = payload(); p2['items'][0]['contact_email'] = 'malo@'
        with self.assertRaises(ValueError):
            calculate(p2)

    def test_missing_contact_email_flags_dispatch(self):
        p = payload(); p['items'][2]['contact_email'] = ''
        r = calculate(p)
        d = next(x for x in r['dispatch'] if x['id'] == 'P-01')
        self.assertEqual(d['to'], '')
        self.assertEqual(d['status'], 'Falta correo del contacto')
        self.assertEqual(r['totals']['without_email'], 1)

    def test_date_and_deadline_rules(self):
        p = payload(); p['context']['response_deadline'] = '2026-12-01'  # anterior al corte
        with self.assertRaises(ValueError):
            calculate(p)
        p2 = payload(); p2['context']['year'] = 2025  # no coincide con el corte
        with self.assertRaises(ValueError):
            calculate(p2)
        p3 = payload(); p3['context']['currency'] = 'US'
        with self.assertRaises(ValueError):
            calculate(p3)

    def test_coverage_and_percentage(self):
        r = calculate(payload())
        cov = next(c for c in r['coverage'] if c['type'] == 'cuentas_por_cobrar')
        self.assertEqual(cov['sampled'], '48200.00')
        self.assertEqual(cov['ledger'], '60000.00')
        self.assertEqual(cov['difference'], '11800.00')
        self.assertEqual(cov['coverage_pct'], '80.33')
        self.assertFalse(cov['within_tolerance'])
        # rubro sin saldo del mayor queda declarado sin cobertura
        bancos = next(c for c in r['coverage'] if c['type'] == 'bancos')
        self.assertIsNone(bancos['ledger'])

    def test_regional_number(self):
        p = payload()
        p['items'][1]['amount'] = '1.234,50'   # formato europeo
        r = calculate(p)
        self.assertEqual(r['items'][1]['amount'], '1234.50')

    def test_fingerprint_is_deterministic_and_input_sensitive(self):
        a = calculate(payload())['input_sha256']
        b = calculate(payload())['input_sha256']
        self.assertEqual(a, b)
        p = payload(); p['items'][0]['amount'] = '1'
        self.assertNotEqual(a, calculate(p)['input_sha256'])

    def test_all_eight_types_render(self):
        types = ['bancos', 'cuentas_por_cobrar', 'proveedores', 'relacionados', 'seguros',
                 'abogados', 'inventarios_terceros', 'inversiones']
        p = payload()
        p['items'] = [{'id': f'T{i}', 'type': t, 'entity': f'Entidad {t}',
                       'contact_email': f'c{i}@x.co', 'amount': '100'} for i, t in enumerate(types)]
        r = calculate(p)
        self.assertEqual(len(r['letters']), 8)
        self.assertTrue(all(l['subject'] for l in r['letters']))

    def test_language_selection(self):
        # por defecto español
        es = calculate(payload())
        self.assertEqual(es['context']['language'], 'es')
        self.assertTrue(any('Señores' in x for x in es['letters'][0]['salutation']))
        # inglés
        p = payload(); p['context']['language'] = 'en'
        en = calculate(p)
        letter = en['letters'][0]
        self.assertEqual(letter['language'], 'en')
        self.assertTrue(any('Dear' in x for x in letter['salutation']))
        self.assertIn('independent auditors', letter['blocks'][0])
        # francés y portugués disponibles y renderizando
        from backend.app.aud.confirmaciones_saldos.plantillas import languages
        self.assertEqual(set(languages()), {'es', 'en', 'fr', 'pt'})
        for lang, needle in (('fr', 'auditeurs indépendants'), ('pt', 'auditores independentes')):
            pl = payload(); pl['context']['language'] = lang
            rr = calculate(pl)
            self.assertEqual(rr['letters'][0]['language'], lang)
            self.assertIn(needle, rr['letters'][0]['blocks'][0])
        # idioma no disponible se rechaza
        p2 = payload(); p2['context']['language'] = 'zz'
        with self.assertRaises(ValueError):
            calculate(p2)

    def test_letter_matches_firm_format(self):
        p = payload()
        p['context'].update(place='Quito', letter_date='2026-09-21', cutoff='2026-08-31', year=2026,
                            response_deadline='2026-09-30')
        r = calculate(p)
        banco = next(l for l in r['letters'] if l['id'] == 'B-01')
        text = banco['text']
        self.assertIn('Quito, 21 de septiembre de 2026', text)
        self.assertIn('Presente.-', text)
        self.assertIn('auditores independientes', text)
        self.assertIn('AUDITCONSULTING AUDITORES CIA. LTDA.', text)
        self.assertIn('con corte al 31 de agosto de 2026', text)

    def test_control_register_grouped_by_rubro(self):
        import openpyxl
        r = calculate(payload())
        wb = openpyxl.load_workbook(BytesIO(build_xlsx(r)))
        ws = wb['04 Control']
        headers = [c.value for c in ws[5]]
        self.assertEqual(headers[0], 'Descripción')
        self.assertIn('Libros', headers)
        self.assertIn('Confirmación', headers)
        self.assertIn('Diferencia', headers)

    def test_letter_email_html_render(self):
        r = calculate(payload())
        html = letter_email_html(r, r['letters'][0])
        self.assertIn(r['items'][0]['entity'], html)
        self.assertIn('<ul', html)                    # los puntos van como lista
        self.assertIn(r['context']['signatory'], html)
        self.assertNotIn('<script', html.lower())

    def test_exports_open_and_no_leak(self):
        p = payload(); p['items'][1]['method'] = 'en_blanco'
        r = calculate(p)
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(build_xlsx(r)))
        self.assertEqual(wb.sheetnames[0], '00 Portada')
        self.assertEqual(len(schedules(r)), 8)
        from docx import Document
        Document(BytesIO(build_docx(r)))   # no lanza => documento válido
        html = build_html(r)
        self.assertIn('Confirmaciones de saldos', html)
        self.assertIn('mailto:', html)


if __name__ == '__main__':
    unittest.main()
