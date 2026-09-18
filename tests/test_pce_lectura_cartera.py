"""Lectura del análisis de antigüedad a nivel de documento."""
import io
from datetime import date

import pytest
from openpyxl import Workbook

from backend.app.aud.pce_cxc.bandas import BANDAS_POR_DEFECTO
from backend.app.aud.pce_cxc.lectura import leer_cartera

CORTE = date(2025, 12, 31)


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cartera"
    ws.append(["Cliente", "N° Documento", "Tipo de cliente", "Fecha emisión",
               "Fecha vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_lee_documentos_y_los_clasifica_por_mora():
    datos = _xlsx([
        ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0),
        ("BETA CIA", "F-2", "RELACIONADOS", date(2025, 1, 1), date(2026, 3, 1), 500.0),
    ])
    r = leer_cartera(datos, "cartera.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert r["fila_encabezado"] == 1
    assert r["formato_fecha"] == "nativo"
    assert [f["banda"] for f in r["filas"]] == ["0 a 30 días", "Por vencer"]
    assert [f["segmento"] for f in r["filas"]] == ["NO-RELACIONADOS", "RELACIONADOS"]
    assert r["total_saldo"] == 1500.0


def test_descarta_filas_identicas_y_conserva_documentos_repetidos_distintos():
    fila = ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0)
    distinta = ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 250.0)
    r = leer_cartera(_xlsx([fila, fila, distinta]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert r["duplicados_exactos"] == 1
    assert r["documentos_repetidos"] == 1
    assert len(r["filas"]) == 2
    assert r["total_saldo"] == 1250.0


def test_reporta_lo_descartado_sin_perderlo_en_silencio():
    r = leer_cartera(_xlsx([
        ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), None, 800.0),
        ("BETA CIA", "", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 900.0),
    ]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert len(r["filas"]) == 0
    assert [d["motivo"] for d in r["descartados"]] == ["sin fecha de vencimiento", "sin número de documento"]
    assert sum(d["saldo"] for d in r["descartados"]) == 1700.0


def test_el_mapeo_manual_manda_sobre_la_deteccion():
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B", "C", "D", "E", "F"])
    ws.append(["ALFA", "F-9", "NO-RELACIONADOS", date(2025, 1, 1), date(2025, 6, 1), 700.0])
    bio = io.BytesIO()
    wb.save(bio)
    r = leer_cartera(bio.getvalue(), "c.xlsx", CORTE, BANDAS_POR_DEFECTO,
                     mapeo={"cliente": 0, "documento": 1, "tipo": 2, "emision": 3,
                            "vencimiento": 4, "saldo": 5})
    assert len(r["filas"]) == 1
    assert r["filas"][0]["banda"] == "181 a 360 días"


def test_la_fila_identica_repetida_queda_registrada_para_poder_cuadrar():
    fila = ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0)
    r = leer_cartera(_xlsx([fila, fila]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert r["duplicados_exactos"] == 1
    descartes = [d for d in r["descartados"] if d["motivo"] == "fila idéntica repetida"]
    assert len(descartes) == 1
    assert descartes[0]["saldo"] == 1000.0
    assert descartes[0]["fila_origen"] == 3
    # El archivo trae 2.000 y el resultado los explica: 1.000 medidos + 1.000 descartados.
    assert r["total_saldo"] + sum(d["saldo"] for d in r["descartados"]) == 2000.0


def test_las_dos_filas_de_un_documento_repetido_quedan_marcadas():
    a = ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0)
    b = ("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 250.0)
    r = leer_cartera(_xlsx([a, b]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert [f["repetido"] for f in r["filas"]] == [True, True]
    assert r["documentos_repetidos"] == 1


def test_el_encabezado_se_puede_indicar_a_mano():
    wb = Workbook()
    ws = wb.active
    ws.append(["ANÁLISIS DE ANTIGÜEDAD AL 31/12/2025", None, None, None, None, None])
    ws.append([None, None, None, None, None, None])
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    ws.append(["ALFA", "F-9", "NO-RELACIONADOS", date(2025, 1, 1), date(2025, 6, 1), 700.0])
    bio = io.BytesIO()
    wb.save(bio)
    r = leer_cartera(bio.getvalue(), "c.xlsx", CORTE, BANDAS_POR_DEFECTO,
                     mapeo={"cliente": 0, "documento": 1, "tipo": 2, "emision": 3,
                            "vencimiento": 4, "saldo": 5},
                     fila_encabezado=3)
    assert r["fila_encabezado"] == 3
    assert len(r["filas"]) == 1
    assert r["filas"][0]["documento"] == "F-9"


def test_una_fila_de_encabezado_fuera_de_rango_se_rechaza():
    """Valida que fila_encabezado se rechace si está fuera del rango válido."""
    datos = _xlsx([("ALFA S.A.", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0)])
    mapeo = {"cliente": 0, "documento": 1, "tipo": 2, "emision": 3, "vencimiento": 4, "saldo": 5}

    # Con mapeo manual y fila_encabezado fuera de rango: debe rechazarse
    for valor in (50, 0, -3):
        with pytest.raises(ValueError, match="fuera de rango"):
            leer_cartera(datos, "c.xlsx", CORTE, BANDAS_POR_DEFECTO, mapeo=mapeo, fila_encabezado=valor)

    # Sin mapeo y con fila_encabezado fuera de rango: debe rechazarse
    for valor in (50, 0, -3):
        with pytest.raises(ValueError, match="fuera de rango"):
            leer_cartera(datos, "c.xlsx", CORTE, BANDAS_POR_DEFECTO, fila_encabezado=valor)


# ---------------------------------------------------------------------------
# I7 - La clasificación de segmento no depende del separador que use el cliente
# ---------------------------------------------------------------------------

def test_no_relacionados_escrito_de_cualquier_forma_no_cae_en_relacionados():
    """El efecto de fallar aquí no es de una fila: toda la cartera de terceros
    se muda al segmento de relacionadas, el total de NO-RELACIONADOS queda en 0
    y el anclaje de ambos segmentos se descuadra a la vez."""
    filas = [
        ("A", "F-1", "NO_RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("B", "F-2", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("C", "F-3", "NO RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("D", "F-4", "NORELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("E", "F-5", "no.relacionados", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("F", "F-6", "TERCEROS NO RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("G", "F-7", "RELACIONADAS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
        ("H", "F-8", "PARTES RELACIONADAS", date(2025, 9, 1), date(2025, 12, 1), 100.0),
    ]
    r = leer_cartera(_xlsx(filas), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert [f["segmento"] for f in r["filas"]] == ["NO-RELACIONADOS"] * 6 + ["RELACIONADOS"] * 2


# ---------------------------------------------------------------------------
# C1 - El importe descartado se totaliza, no solo se cuenta
# ---------------------------------------------------------------------------

def test_el_importe_descartado_se_totaliza_y_se_desglosa_por_motivo():
    """Guardar solo el conteo de descartados hace desaparecer el importe: con
    anclaje a los EEFF ese hueco se reparte sobre las filas que sí entraron."""
    r = leer_cartera(_xlsx([
        ("ALFA", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0),
        ("BETA", "", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 750000.0),
        ("GAMA", "F-3", "NO-RELACIONADOS", date(2025, 9, 1), None, 1250000.0),
    ]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)

    assert r["total_saldo"] == pytest.approx(1000.0)
    assert r["descartados_importe"] == pytest.approx(2000000.0)

    por_motivo = {d["motivo"]: d for d in r["descartados_por_motivo"]}
    assert por_motivo["sin número de documento"]["filas"] == 1
    assert por_motivo["sin número de documento"]["importe"] == pytest.approx(750000.0)
    assert por_motivo["sin fecha de vencimiento"]["filas"] == 1
    assert por_motivo["sin fecha de vencimiento"]["importe"] == pytest.approx(1250000.0)
    # El desglose cuadra contra el total, sin residuos.
    assert sum(d["importe"] for d in r["descartados_por_motivo"]) == pytest.approx(
        r["descartados_importe"])


def test_la_fila_identica_repetida_no_cuenta_como_cartera_no_leida():
    """Descartar una fila idéntica repetida es una depuración deliberada, no
    cartera perdida: su importe se informa igual, pero aparte."""
    fila = ("ALFA", "F-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 1000.0)
    r = leer_cartera(_xlsx([fila, fila]), "c.xlsx", CORTE, BANDAS_POR_DEFECTO)
    assert r["descartados_importe"] == pytest.approx(1000.0)
    assert r["cartera_no_leida"] == pytest.approx(0.0)
