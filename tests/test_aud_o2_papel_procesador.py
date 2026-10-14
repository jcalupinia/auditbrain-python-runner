"""O2 (AUD_O2_SUBPROCESO): el papel armado en subproceso aislado debe ser IDÉNTICO al inline.

Siembra una prueba PCE (pce_cohortes_niif9) en la BD de test, genera el papel por
``servicio.papel_procesador`` —la función que usa el endpoint de descarga— con el flag
APAGADO (inline) y ENCENDIDO (subproceso), y compara el Excel celda por celda: 0 diferencias.
Verifica así que aislar el armado del papel (anti-OOM) no cambia el resultado contable.
"""
from __future__ import annotations

import datetime as dt
import io

import openpyxl

from backend.app.aud.niif.ciclo import servicio
from backend.app.aud.niif.ciclo.models import Prueba
from backend.app.aud.niif.procesadores import pce_cohortes_niif9 as m
from backend.app.db.session import SessionLocal, init_db

CORTE = "2025-12-31"
PARMS = {"umbral": 730, "desdoblar": "No", "relKey": "RELACIONAD", "fT": 1.0, "fR": 1.0,
         "castiga": "No", "eNR_t": 0, "eR_t": 0, "pol_t360": 5}


def _d(off_days: int) -> str:
    return (dt.date(2025, 12, 31) - dt.timedelta(days=off_days)).isoformat()


def _cartera(n: int, corte_idx: int) -> list[dict]:
    filas = []
    for i in range(n):
        if i % 10 < 6:  # cohorte persistente entre cortes
            doc = f"C{i}"; saldo = 1000.0 * (1.0 - 0.2 * corte_idx); off = 400 + (i % 300)
        else:
            doc = f"N{corte_idx}-{i}"; saldo = 500.0 + (i % 40) * 37.0; off = (i % 360)
        seg = "RELACIONADO" if i % 7 == 0 else "NO-RELACIONADO"
        filas.append({"id": doc, "cliente": f"Cliente {i % 50}", "tipo": seg,
                      "factura": _d(off + 90), "vence": _d(off), "saldo": f"{saldo:.2f}", "_row": i + 2})
    return filas


def _datasets(n: int) -> dict:
    return {
        "cartera_t2": _cartera(n, 0), "cartera_t1": _cartera(n, 1), "cartera_t": _cartera(n, 2),
        "provision": [{"codigo": "1.1.03.02", "descripcion": "(-) Provisión cuentas incobrables",
                       "saldo_anterior": "1000", "saldo_actual": "1000", "_row": 2}],
        "mayor_t2": [{"concepto": "Sin movimiento", "constitucion": 0, "reversion": 0, "castigos": 0, "_row": 2}],
        "mayor_t1": [{"concepto": "Sin movimiento", "constitucion": 0, "reversion": 0, "castigos": 0, "_row": 2}],
        "mayor_t": [{"concepto": "Sin movimiento", "constitucion": 0, "reversion": 0, "castigos": 0, "_row": 2}],
    }


def _norm(c) -> str:
    v = c.value
    if v is None:
        return ""
    if isinstance(v, str) and v.startswith("="):
        return "F:" + v[1:].replace(" ", "")
    if isinstance(v, (dt.date, dt.datetime)):
        return "D:" + v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        return "N:" + format(round(float(v), 6), ".6f")
    return "T:" + str(v)


def _difs(xa: bytes, xb: bytes) -> int:
    wa = openpyxl.load_workbook(io.BytesIO(xa)); wb = openpyxl.load_workbook(io.BytesIO(xb))
    assert wa.sheetnames == wb.sheetnames, (wa.sheetnames, wb.sheetnames)
    d = 0
    for name in wa.sheetnames:
        sa, sb = wa[name], wb[name]
        R = max(sa.max_row, sb.max_row); C = max(sa.max_column, sb.max_column)
        for r in range(1, R + 1):
            for c in range(1, C + 1):
                if _norm(sa.cell(r, c)) != _norm(sb.cell(r, c)):
                    d += 1
    return d


def _sembrar(db, n: int) -> Prueba:
    defi = m.definicion(); defi["processor"] = "pce_cohortes_niif9"
    reg = {"parameters": dict(PARMS),
           "engagement": {"client": "Test O2", "cutoff": CORTE, "framework": "", "edition": ""},
           "datasets": _datasets(n),
           # reg["run"] truthy: _run_fresco recomputa desde los datasets al descargar (guard "no ejecutada").
           "run": {"_sembrado": True}, "program": [], "sources": []}
    p = Prueba(project_id=1, version=1, estado="APROBADO", origen="pce", definicion=defi, registro=reg, revision=1)
    db.add(p); db.commit(); db.refresh(p)
    return p


def test_o2_papel_identico_a_inline(monkeypatch):
    init_db()
    db = SessionLocal()
    try:
        p = _sembrar(db, 150)
        monkeypatch.setenv("AUD_O2_SUBPROCESO", "false")
        xi, hi = servicio.papel_procesador(db, p)
        monkeypatch.setenv("AUD_O2_SUBPROCESO", "true")
        xo, ho = servicio.papel_procesador(db, p)
    finally:
        db.close()
    assert len(xi) > 0 and len(xo) > 0, "ambos papeles deben generarse"
    assert len(ho) > 0, "el HTML debe generarse en el subproceso (fix spawn/Agg)"
    assert _difs(xi, xo) == 0, "el Excel del subproceso debe ser idéntico al inline (0 diferencias)"
