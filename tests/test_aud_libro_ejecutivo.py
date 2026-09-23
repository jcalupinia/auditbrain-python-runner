"""Papel de trabajo EJECUTIVO (Excel) — identidad SIGMANSERVICE, panel de inicio,
navegación, semáforo y bitácora en español, con las fórmulas intactas.

El verificador celda a celda contra Excel real (scripts/verificar_formulas_pi.py)
usa win32com (solo Windows); aquí se prueba, sin Excel, que cada fórmula de la
definición de la cédula se escribe EXACTA en su celda (layout intacto) y que el
libro abre sin «reparar».
"""
import io

from openpyxl import load_workbook

from backend.app.aud.niif.procesadores import estilo_ejecutivo as est
from backend.app.aud.niif.procesadores import libro
from backend.app.aud.niif.procesadores import PROCESADORES
from backend.app.aud.niif.ciclo import reglas


def _reg(pid):
    mod = PROCESADORES[pid]
    d = mod.definicion()
    esc = getattr(mod, "ESCENARIOS", None)
    if esc:
        _, datasets, param, corte = esc[0]
    else:
        E = mod.EJEMPLO
        datasets, param, corte = E["datasets"], E.get("parametros", {}), E["corte"]
    run = mod.ejecutar(datasets, param, corte)
    run["hojas"] = mod.hojas(run)
    reg = {"run": run, "engagement": {"client": "Cliente de muestra S.A.", "ruc": "1790000000001",
           "cutoff": corte, "framework": (param.get("_marco") or d["frameworks"][0]),
           "firm": "AuditConsulting Auditores Cía. Ltda.", "preparer": "AA", "reviewer": "BB"},
           "program": [{**x, "reference": x.get("source", "")} for x in d["program"]], "sources": []}
    return d, mod, reg


def test_excel_ejecutivo_panel_navegacion_y_formatos():
    d, mod, reg = _reg("perdidas_incurridas_s11")
    eventos = [{"fecha": "2026-01-15T10:00:00", "accion": "create", "estado_anterior": None,
                "estado_nuevo": "PRUEBA_SELECCIONADA", "actor": "aa@x", "comentario": ""},
               {"fecha": "2026-01-15T11:00:00", "accion": "execute", "estado_anterior": "METODOLOGIA_APROBADA",
                "estado_nuevo": "PRUEBA_EJECUTADA", "actor": "aa@x", "comentario": ""}]
    wb = load_workbook(io.BytesIO(libro.xlsx(d, mod, reg, eventos, 1, "APROBADO") if False else libro.xlsx(d, reg, eventos, 1, "APROBADO")))
    assert wb.sheetnames[0] == "00_Inicio"
    ini = wb["00_Inicio"]
    assert ini.sheet_view.showGridLines is False
    # navegación por hipervínculo interno a cada cédula
    destinos = [c.hyperlink.target for row in ini.iter_rows() for c in row if c.hyperlink]
    assert len(destinos) >= 10 and all(t.startswith("#'") for t in destinos)
    # impresión horizontal, ajustar a ancho
    detalle = wb[wb.sheetnames[1]]
    assert detalle.page_setup.orientation == "landscape" and detalle.page_setup.fitToWidth == 1
    # bitácora en español, sin códigos crudos
    ctrl = wb[next(n for n in wb.sheetnames if "Control" in n)]
    vals = [c.value for row in ctrl.iter_rows() for c in row]
    assert "Prueba creada" in vals and "Prueba ejecutada" in vals
    assert "PRUEBA_EJECUTADA" not in vals


def test_formulas_intactas_en_su_celda():
    d, mod, reg = _reg("perdidas_incurridas_s11")
    hojas = libro.cedulas(d, reg, [], 1, "APROBADO")
    titulos = libro._titulos_unicos(hojas)
    wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "APROBADO")))
    total = comprobadas = 0
    for h, t in zip(hojas, titulos):
        ws = wb[t]
        filas = h["rows"] + ([h["total"]] if h.get("total") else [])
        for i, fila in enumerate(filas):
            for j, v in enumerate(fila):
                if isinstance(v, dict) and "f" in v:
                    total += 1
                    if ws.cell(row=5 + i, column=1 + j).value == "=" + v["f"]:
                        comprobadas += 1
    assert total > 0 and comprobadas == total


def test_todas_las_herramientas_generan_excel_que_abre_sin_reparar():
    for pid, mod in PROCESADORES.items():
        if not getattr(mod, "RUBRO", None):
            continue
        d, _, reg = _reg(pid)
        wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "APROBADO")))  # abre sin reparar
        assert wb.sheetnames[0] == "00_Inicio", pid


def test_tokens_y_etiquetas_cubren_el_vocabulario():
    for e in reglas.ESTADOS:
        assert est.estado_es(e) != e, e  # todos traducidos
    for a in ("create", "execute", "approve_program", "approve"):
        assert est.accion_es(a)
    assert est.color_semaforo(0) == est.GREEN
    assert est.color_semaforo(2) == est.AMBER
    assert est.color_semaforo(9) == est.RED
