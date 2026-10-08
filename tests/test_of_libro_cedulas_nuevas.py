"""Cédulas DM Programa, DM2 Sumaria y DM10 Hallazgos (papel de trabajo)."""

import datetime

from openpyxl import Workbook

from backend.app.aud.obligaciones_fiscales.libro.cedulas.dm2_sumaria import (
    SHEET_DM2, build_dm2,
)
from backend.app.aud.obligaciones_fiscales.libro.cedulas.dm10_hallazgos import (
    SHEET_DM10, build_dm10,
)
from backend.app.aud.obligaciones_fiscales.libro.cedulas.dm_programa import (
    SHEET_DM_PROGRAMA, build_dm_programa,
)
from backend.app.aud.obligaciones_fiscales.mayor.tipos import Movimiento


class _Fila:
    def __init__(self, codigo, nombre, categoria):
        self.codigo_cuenta = codigo
        self.nombre_cuenta = nombre
        self.categoria_final = categoria


def _celdas(ws):
    return {
        (c.row, c.column): c.value
        for fila in ws.iter_rows() for c in fila if c.value is not None
    }


# ------------------------------------------------- DM Programa ------------

def test_dm_programa_tiene_cabecera_objetivos_y_procedimientos():
    wb = Workbook()
    build_dm_programa(wb, cliente="CLIENTE X", periodo="2026")
    ws = wb[SHEET_DM_PROGRAMA]
    textos = " ".join(str(v) for v in _celdas(ws).values())
    assert "CLIENTE X" in textos
    assert "Objetivos" in textos
    assert "Procedimientos" in textos
    # Referencia cada cédula que el sistema genera, y NO DM1 (no se genera).
    assert "DM2" in textos and "DM8" in textos and "DM10" in textos
    assert "DM1 " not in textos


# ------------------------------------------------- DM2 Sumaria ------------

def test_dm2_calcula_saldo_anterior_variacion_y_corte():
    """2.3.2.01: apertura -30000 (cierre anterior) + movimientos del período
    por 24000 netos → corte -6000, variación 24000."""
    clasificacion = [_Fila("2.3.2.01", "FISCO", "RET_RENTA")]
    movs = [
        Movimiento(codigo="2.3.2.01", fecha=datetime.date(2026, 1, 1),
                   descripcion="REG. SALDOS INICIALES", haber=30000.0),
        Movimiento(codigo="2.3.2.01", fecha=datetime.date(2026, 3, 10),
                   descripcion="REG. PAGO", debe=24000.0),
    ]
    wb = Workbook()
    build_dm2(wb, clasificacion=clasificacion, movimientos=movs,
              cliente="C", periodo="2026")
    ws = wb[SHEET_DM2]
    # Fila de la cuenta (código en col 1).
    fila = next(r for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value == "2.3.2.01")
    assert ws.cell(fila, 3).value == -30000.0   # saldo cierre anterior
    assert ws.cell(fila, 4).value == 24000.0    # variación
    assert ws.cell(fila, 6).value == -6000.0    # saldo al corte


def test_dm2_sin_apertura_toda_la_variacion_es_del_periodo():
    clasificacion = [_Fila("1.1.5.1.1", "IVA Compras", "IVA_COMPRAS")]
    movs = [Movimiento(codigo="1.1.5.1.1", fecha=datetime.date(2026, 2, 1),
                       descripcion="REG. FACT", debe=100.0)]
    wb = Workbook()
    build_dm2(wb, clasificacion=clasificacion, movimientos=movs, periodo="2026")
    ws = wb[SHEET_DM2]
    fila = next(r for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value == "1.1.5.1.1")
    assert ws.cell(fila, 3).value == 0.0      # sin cierre anterior
    assert ws.cell(fila, 4).value == 100.0    # variación
    assert ws.cell(fila, 6).value == 100.0    # corte


def test_dm2_no_deja_division_por_cero_en_el_porcentaje():
    clasificacion = [_Fila("1.1.5.1.1", "IVA Compras", "IVA_COMPRAS")]
    movs = [Movimiento(codigo="1.1.5.1.1", fecha=datetime.date(2026, 2, 1),
                       descripcion="x", debe=100.0)]
    wb = Workbook()
    build_dm2(wb, clasificacion=clasificacion, movimientos=movs, periodo="2026")
    ws = wb[SHEET_DM2]
    valores = [c.value for fila in ws.iter_rows() for c in fila]
    assert not any(isinstance(v, str) and "DIV/0" in v for v in valores)


def test_dm2_incluye_una_fila_total():
    clasificacion = [_Fila("1.1.5.1.1", "IVA Compras", "IVA_COMPRAS")]
    movs = [Movimiento(codigo="1.1.5.1.1", fecha=datetime.date(2026, 2, 1),
                       descripcion="x", debe=100.0)]
    wb = Workbook()
    build_dm2(wb, clasificacion=clasificacion, movimientos=movs, periodo="2026")
    ws = wb[SHEET_DM2]
    textos = [str(v) for v in _celdas(ws).values()]
    assert "TOTAL" in textos


# ------------------------------------------------- DM10 Hallazgos ---------

def test_dm10_lista_las_cedulas_con_diferencias():
    wb = Workbook()
    build_dm10(wb, cliente="C", periodo="2026")
    ws = wb[SHEET_DM10]
    refs = {ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    for ref in ("DM3", "DM4", "DM5", "DM6", "DM7", "DM8"):
        assert ref in refs


def test_dm10_tiene_columnas_de_observacion_y_recomendacion():
    wb = Workbook()
    build_dm10(wb, periodo="2026")
    ws = wb[SHEET_DM10]
    textos = " ".join(str(v) for v in _celdas(ws).values())
    assert "Observación" in textos
    assert "Recomendación" in textos


def test_dm2_publica_las_direcciones_del_saldo_al_corte():
    """DM2 devuelve {("saldo_corte", código): dirección} para que DM3 cruce
    contra la misma cifra de la sumaria."""
    clasificacion = [_Fila("1.1.5.1.2", "Crédito Tributario", "IVA_COMPRAS")]
    movs = [Movimiento(codigo="1.1.5.1.2", fecha=datetime.date(2026, 2, 1),
                       descripcion="x", debe=100.0)]
    wb = Workbook()
    salida = build_dm2(wb, clasificacion=clasificacion, movimientos=movs, periodo="2026")
    assert ("saldo_corte", "1.1.5.1.2") in salida
    addr = salida[("saldo_corte", "1.1.5.1.2")]
    assert addr.startswith("'DM2 Cédula Sumaria'!")
    # La celda referida tiene el saldo al corte (100).
    ws = wb[SHEET_DM2]
    assert ws[addr.split("!", 1)[1]].value == 100.0
