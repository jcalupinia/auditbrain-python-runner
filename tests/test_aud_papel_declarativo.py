"""Papel de las pruebas DECLARATIVAS con el diseño de las pruebas con procesador
(pedido del dueño, 2026-09-25: «haz las pruebas declarativas con el diseño nuevo»).

El navegador envía las cédulas con fórmulas del exportador del sitio (``cargaPapel``) y el
servidor arma Excel (portada = panel del HTML), HTML, Word y PowerPoint con ``libro``.
Los datos de ejemplo (``tests/fixtures/papel_declarativo``) los genera
``frontend/scripts/fixture_papel_declarativo.mjs``; ``papelDeclarativo.test.js`` falla si
quedan desactualizados. La igualdad fórmula = valor se verifica con LibreOffice en
``scripts/verificar_papel_declarativo.py`` (debe dar «DIFERENCIAS: 0»).
"""
import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.aud.niif.procesadores import declarativo, graficos, libro
from backend.app.aud.niif.procesadores import html_ejecutivo as hx
from tests.test_aud_office_como_html import _es_grafico

FIX = Path(__file__).parent / "fixtures" / "papel_declarativo"
EJEMPLOS = ("vnr", "niif16")


def _carga(nombre):
    return json.loads((FIX / f"{nombre}.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module", params=EJEMPLOS)
def papel(request):
    carga = _carga(request.param)
    args = declarativo.armar(carga)
    return request.param, carga, args, declarativo.papel(carga)


def _formulas(hojas):
    for h in hojas:
        for i, fila in enumerate(h["rows"]):
            for j, c in enumerate(fila):
                if isinstance(c, dict) and "f" in c:
                    yield h["name"], i, j, c["f"]


def test_mismas_celdas_que_el_sitio(papel):
    """La fila k de la cédula del sitio es la fila k del Excel: toda fórmula del sitio queda
    en la misma celda (salvo las que el adaptador corrige, que se prueban aparte)."""
    _, carga, (d, reg, *_), _ = papel
    hojas = {h["name"]: h for h in reg["run"]["hojas"]}
    assert "01_Caratula" not in hojas, "la portada del sitio la reemplaza 00_Inicio"
    corregidas = 0
    for nombre, filas in zip(carga["cedulas"]["nombres"], carga["cedulas"]["hojas"]):
        if nombre == "01_Caratula":
            continue
        h = hojas[nombre]
        for k, fila in enumerate(filas[4:]):
            for j, c in enumerate(fila):
                if isinstance(c, dict) and "f" in c:
                    nueva = h["rows"][k][j]
                    assert isinstance(nueva, dict), (nombre, k, j)
                    if nueva["f"] != c["f"]:
                        corregidas += 1
                        assert nombre == "07_Calculos" and "'13_Cuadro'!" in nueva["f"], (nombre, k, j, nueva["f"])
    if d.get("series"):
        assert corregidas, "las reglas con agregados de la serie deben leer el cuadro"


def test_formulas_solo_a_hojas_del_libro(papel):
    _, _, (d, reg, *_), _ = papel
    hojas = libro.cedulas(d, reg, [], 1, "APROBADO")
    nombres = {h["name"] for h in hojas}
    for hoja, i, j, f in _formulas(hojas):
        assert f.count("(") == f.count(")"), (hoja, i, j, f)
        for ref in re.findall(r"'([^']+)'!", f):
            assert ref in nombres, (hoja, i, j, ref)


def test_sin_cifras_calculadas_pegadas(papel):
    """Registros por COUNTIF, importe de cada excepción enlazado a su celda de origen y, en una
    serie, el cuadro de períodos con fórmulas."""
    nombre, _, (d, reg, *_), _ = papel
    h = {x["name"]: x for x in reg["run"]["hojas"]}
    registros = next(f for f in h["08_Pruebas"]["rows"] if f[0] == "Registros")
    assert registros[1]["f"].startswith("COUNTIF('05_Data_Original'!")
    exc = h["09_Excepciones"]
    assert exc.get("problemas") and exc["rows"]
    for f in exc["rows"]:
        assert isinstance(f[4], dict) and f[4]["f"].startswith(("'07_Calculos'!", "'13_Cuadro'!")), f
        assert f[2] in declarativo.CODIGOS.values(), "el código va en español"
    if nombre == "niif16":
        cuadro = h["13_Cuadro"]
        calculadas = [c for f in cuadro["rows"] for c in f[3:] if c is not None]
        assert calculadas and all(isinstance(c, dict) and c["f"].startswith("ROUND(ROUND(") for c in calculadas)


def test_cada_columna_calculada_se_explica(papel):
    _, _, (d, reg, *_), _ = papel
    hojas = libro.cedulas(d, reg, [], 1, "APROBADO")
    for h in hojas:
        for b in libro.como_se_calcula(h, hojas):
            assert b["escrita"], (h["name"], b["columna"])
            assert libro.PLANTILLA_GENERICA not in b["explicacion"]


def test_panel_completo(papel):
    _, _, (d, reg, *_), _ = papel
    p = graficos.panel(graficos.modulo(d), reg["run"], reg["run"]["hojas"])
    assert not p["faltan"], p["faltan"]
    assert p["problemas"]["valor"] == len(reg["run"]["exceptions"])
    assert len(hx.kpis_datos(p)) == 5


def test_excel_portada_es_el_panel(papel):
    _, _, (d, reg, *_), archivos = papel
    wb = load_workbook(io.BytesIO(archivos["xlsx"]))
    ini = wb["00_Inicio"]
    assert wb.sheetnames[0] == "00_Inicio" and len(ini._images) == 2        # los dos logotipos
    assert len(ini._charts) == 4                                             # los 4 gráficos del HTML
    valores = [ini.cell(row=14, column=c).value for c in range(2, 7) if ini.cell(row=13, column=c).value]
    assert valores and all(isinstance(v, str) and v.startswith("=") for v in valores), valores
    for ws in wb.worksheets:                                                 # sin texto que Excel lea como fórmula
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value[:1] in "+-@" and c.data_type != "f":
                    pytest.fail(f"{ws.title}!{c.coordinate}: {c.value[:40]}")
    assert {"05_Data_Original", "07_Calculos", "09_Excepciones", "00_Anexo_tecnico"} <= set(wb.sheetnames)


def test_html_word_y_powerpoint_con_el_diseno_nuevo(papel):
    from docx import Document
    from pptx import Presentation

    from backend.app.aud.niif.procesadores import papel_office

    _, _, (d, reg, *_), archivos = papel
    h = archivos["html"].decode()
    assert h.count('<div class="kpi k-') == 5                              # las 5 tarjetas del panel
    topbar = h[h.index('<header class="topbar">'):h.index("</header>")]
    assert topbar.count('src="data:image/') == 2
    for ext in ("xlsx", "docx", "pptx"):
        assert re.search(rf'download="[^"]+\.{ext}" href="data:', h), ext
    assert not re.search(r'<(img|script|link)[^>]+(src|href)="https?:', h)

    prs = Presentation(io.BytesIO(archivos["pptx"]))
    diaps = list(prs.slides)
    assert all(str(s.background.fill.fore_color.rgb) == papel_office.OSCURO["bg"] for s in diaps)
    fotos = [sh.image.blob for s in diaps for sh in s.shapes if sh.shape_type == 13]
    assert sum(_es_grafico(b) for b in fotos) == 4

    doc = Document(io.BytesIO(archivos["docx"]))
    cuerpo = [r.target_part.blob for r in doc.part.rels.values() if "image" in r.reltype]
    assert sum(_es_grafico(b) for b in cuerpo) == 4
    todo = " ".join([c.text for t in doc.tables for f in t.rows for c in f.cells] + [p.text for p in doc.paragraphs])
    p = graficos.panel(graficos.modulo(d), reg["run"], reg["run"]["hojas"])
    for k in hx.kpis_datos(p):
        assert k["valor"] in todo and k["rotulo"].upper() in todo, k


@pytest.mark.parametrize("roto", [
    {},
    {"herramienta": {"definition": {}}, "cedulas": {}},
    "nombres_distintos",
    "hoja_ajena",
])
def test_carga_invalida(roto):
    carga = _carga("vnr")
    if roto == "nombres_distintos":
        carga["cedulas"]["etiquetas"] = carga["cedulas"]["etiquetas"][:-1]
    elif roto == "hoja_ajena":
        carga["cedulas"]["nombres"][2] = "99_Inventada"
    else:
        carga = roto
    with pytest.raises(declarativo.CargaInvalida):
        declarativo.armar(carga)


def test_procesadores_no_cambian():
    """Las piezas compartidas siguen igual para una prueba con procesador."""
    from backend.app.aud.niif import ejercicio_modelo as em
    from backend.app.aud.niif.procesadores import PROCESADORES

    mod = PROCESADORES["nomina_beneficios"]
    d = mod.definicion()
    assert graficos.modulo(d) is mod
    ds, par, corte = em.escenario(mod)
    reg = em._reg(d, mod, ds, par, corte)
    nombres = [h["name"] for h in libro.cedulas(d, reg, [], 1, "APROBADO")]
    assert nombres[:3] == ["00_Caratula", "00_Programa", "00_Fuentes"] and nombres[-1] == "14_Control_Revision"
