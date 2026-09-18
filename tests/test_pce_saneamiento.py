"""V1 — ningún texto del cliente puede entrar al papel como FÓRMULA de Excel.

POR QUÉ EXISTE. `openpyxl` marca una celda como fórmula (`data_type == "f"`) en
cuanto el texto empieza por «=». Todo lo que el papel imprime de nombre de
cliente, nombre de archivo, número de documento, entidad, justificación o
sustento lo escribe el cliente, así que un nombre como ``=cmd|' /C calc'!A0``
entraba al libro como fórmula. Tiene dos caras y las dos son graves:

1. **Inyección de fórmula o DDE**: el libro lo abre el auditor en SU máquina.
2. **El libro pide reparación**: una fórmula inválida es exactamente lo que hace
   que Excel ofrezca «recuperar» el archivo, y que el libro abra sin pedir
   reparación es la garantía central de esta rama.

QUÉ SE COMPRUEBA, de punta a punta: desde el archivo que sube el auditor -con
datos hostiles en TODOS los campos que el cliente controla- hasta la celda del
libro descargado. La prueba no mira los cuatro sumideros que el informe nombra:
recorre LAS TRECE HOJAS y falla si cualquier celda es fórmula y contiene la
marca hostil, así que un sumidero nuevo que alguien olvide sanear la rompe.
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.auth import service as auth_service
from backend.app.auth.models import Role
from backend.app.db.session import SessionLocal, init_db

BASE = "/api/v1/aud/pce-cxc"

#: Nombre de cliente hostil del informe de revisión: DDE clásico y, a la vez,
#: una fórmula que Excel no puede evaluar (el libro pediría repararse).
CLIENTE_HOSTIL = "=cmd|' /C calc'!A0"
#: Los otros cuatro caracteres con los que Excel (y LibreOffice) empiezan una
#: fórmula o un comando.
DOCUMENTO_HOSTIL = "+1+1"
ENTIDAD_HOSTIL = '=HYPERLINK("http://ejemplo.invalido","haga clic")'
JUSTIFICACION_HOSTIL = "@SUM(1+1)*cmd|' /C calc'!A0"
ARCHIVO_HOSTIL = "=1+1.xlsx"
SUSTITUTA_HOSTIL = "-2+3+cmd|' /C calc'!A0"

#: Todo lo hostil junto: si cualquiera de estos textos aparece en una celda
#: marcada como fórmula, el papel quedó expuesto.
TEXTOS_HOSTILES = (CLIENTE_HOSTIL, DOCUMENTO_HOSTIL, ENTIDAD_HOSTIL,
                   JUSTIFICACION_HOSTIL, ARCHIVO_HOSTIL, SUSTITUTA_HOSTIL)


def _xlsx(filas: list[tuple]) -> bytes:
    """Un análisis de antigüedad donde los textos son TEXTO, no fórmulas.

    Es el escenario real: el ERP del cliente exporta el nombre tal cual y la
    celda de origen es de texto. Sin forzarlo, `openpyxl` guardaría
    ``=cmd|…`` como fórmula en el propio archivo de entrada y el lector
    (`data_only=True`) leería `None`, que no es lo que se quiere probar.
    """
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    for fila in ws.iter_rows():
        for c in fila:
            if isinstance(c.value, str):
                c.data_type = "s"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _archivos_hostiles():
    """Los tres cortes, con el nombre del cliente y el del archivo hostiles."""
    def corte(saldo):
        return _xlsx([(CLIENTE_HOSTIL, DOCUMENTO_HOSTIL, "NO-RELACIONADOS",
                       date(2023, 9, 1), date(2023, 12, 1), saldo)])

    return [("archivos", (ARCHIVO_HOSTIL, corte(300000.0))),
            ("archivos", ("cartera 2024.xlsx", corte(200000.0))),
            ("archivos", ("cartera 2025.xlsx", corte(150000.0)))]


def _token(client, role=Role.user):
    init_db()
    tag = uuid.uuid4().hex[:6]
    email, pw = f"pce-san-{tag}@ex.com", "Sup3rSecret!"
    db = SessionLocal()
    try:
        auth_service.create_user(db, email=email, password=pw, role=role)
    finally:
        db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": pw})
    return r.json()["access_token"]


def _parametros_hostiles() -> dict:
    return {
        "fechas": ["2023-12-31", "2024-12-31", "2025-12-31"],
        "entidad": ENTIDAD_HOSTIL,
        "ruc": "=1791240154001",
        "preparado_por": "=A1",
        "revisado_por": "+A1",
        "referencia": "@PT-PCE-CXC",
        "moneda": "=USD",
        "marco": "-NIIF 9",
        "justificacion_prospectivo": JUSTIFICACION_HOSTIL,
        # Umbral bajo a propósito: fuerza que el cliente hostil salga de la
        # matriz colectiva y llegue a 06-Individual, que es el sumidero que el
        # informe reprodujo.
        "umbral_individual": 1000,
        "tasas_sustitutas": {"NO-RELACIONADOS|Más de 730 días":
                             {"tasa": 0.5, "justificacion": SUSTITUTA_HOSTIL}},
    }


def _libro_de_la_corrida(client, token, parametros=None):
    r = client.post(f"{BASE}/analizar", files=_archivos_hostiles(),
                    data={"parametros": json.dumps(parametros or _parametros_hostiles())},
                    headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    corrida_id = r.json()["corrida_id"]
    x = client.get(f"{BASE}/corridas/{corrida_id}/excel",
                   headers={"Authorization": f"Bearer {token}"})
    assert x.status_code == 200, x.text
    return load_workbook(io.BytesIO(x.content)), r.json()


def _celdas_formula(wb):
    """Todas las celdas que Excel evaluaría como fórmula, hoja por hoja."""
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if c.data_type == "f":
                    yield ws.title, c.coordinate, str(c.value)


def test_ningun_texto_del_cliente_llega_al_papel_como_formula(client):
    """De punta a punta: archivo hostil subido -> celda del libro descargado.

    Recorre las trece hojas. Si CUALQUIER celda es fórmula y contiene alguno de
    los textos hostiles, el papel quedó expuesto a inyección y, además, a que
    Excel pida repararlo.
    """
    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)

    expuestas = [(hoja, coord, valor)
                 for hoja, coord, valor in _celdas_formula(wb)
                 if any(t in valor for t in TEXTOS_HOSTILES)]
    assert not expuestas, (
        "Texto del cliente escrito como FÓRMULA de Excel en el papel de trabajo: "
        + "; ".join(f"{h}!{c} -> {v!r}" for h, c, v in expuestas))


def test_el_nombre_hostil_del_cliente_se_conserva_como_texto_en_06_individual(client):
    """Sanear no es perder el dato: la celda tiene que decir EXACTAMENTE el
    nombre que traía el archivo, pero como texto, no como fórmula."""
    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)
    ws = wb["06-Individual"]
    assert ws["A2"].value == CLIENTE_HOSTIL, ws["A2"].value
    assert ws["A2"].data_type == "s", (
        f"06-Individual!A2 es {ws['A2'].data_type!r}: el nombre del cliente entra como fórmula")


def test_el_nombre_hostil_del_archivo_se_conserva_como_texto_en_02_fuentes(client):
    """El nombre del archivo lo pone el cliente y también se imprime."""
    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)
    ws = wb["02-Fuentes"]
    celda = ws["A2"]
    assert celda.value == ARCHIVO_HOSTIL, celda.value
    assert celda.data_type == "s", (
        f"02-Fuentes!A2 es {celda.data_type!r}: el nombre del archivo entra como fórmula")


@pytest.mark.parametrize("hoja,columna,texto", [
    ("00-Caratula", "B", ENTIDAD_HOSTIL),
    ("01-Parametros", "C", JUSTIFICACION_HOSTIL),
    ("04-Tasas", "F", SUSTITUTA_HOSTIL),
])
def test_los_demas_sumideros_tambien_escriben_texto(client, hoja, columna, texto):
    """Los sumideros que el informe NO nombra: entidad, justificación del ajuste
    prospectivo y justificación de una tasa sustituta."""
    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)
    ws = wb[hoja]
    encontradas = [c for f in ws.iter_rows() for c in f
                   if isinstance(c.value, str) and texto in c.value]
    assert encontradas, f"{hoja}: no se encontró el texto hostil en ninguna celda"
    for c in encontradas:
        assert c.data_type == "s", f"{hoja}!{c.coordinate} es fórmula: {c.value!r}"


def test_todo_texto_que_empieza_por_un_inicio_de_formula_queda_marcado_como_literal(client):
    """`+`, `-` y `@` no son fórmula DENTRO del .xlsx, pero sí al pegarlos.

    En un .xlsx solo es fórmula lo que lleva `<f>`, así que un texto que empieza
    por `+` se guarda como texto. El riesgo llega después: el auditor copia la
    celda y la pega en otra hoja, y Excel vuelve a interpretar la ENTRADA -ahí
    `+1+1`, `-2+3` y `@SUM(…)` sí son fórmulas-. Marcar la celda con el prefijo
    de literal de Excel (`quotePrefix`) cierra también ese camino, y es lo que
    hace que el nombre se siga leyendo tal cual.
    """
    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)
    sin_marcar = []
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if (isinstance(c.value, str) and c.data_type == "s"
                        and c.value[:1] in ("=", "+", "-", "@", "\t", "\r")
                        and not c.quotePrefix):
                    sin_marcar.append(f"{ws.title}!{c.coordinate} -> {c.value!r}")
    assert not sin_marcar, (
        "Texto que empieza por un inicio de fórmula y no está marcado como literal: "
        + "; ".join(sin_marcar))


def test_el_libro_con_datos_hostiles_no_tiene_ninguna_formula_invalida(client):
    """La otra cara de V1: una fórmula que Excel no puede evaluar es lo que hace
    que pida reparar el archivo.

    Todas las fórmulas del papel las escribe el exportador y usan solo el
    subconjunto que declara su docstring. Si aparece una fórmula que empieza por
    algo que no es una función permitida, una referencia o un paréntesis, es
    texto del cliente que se coló.
    """
    import re

    token = _token(client)
    wb, _ = _libro_de_la_corrida(client, token)
    permitido = re.compile(
        r"^=(SUM|SUMIFS|COUNTIFS|INDEX|MATCH|IF|ABS|ROUND|MAX|MIN)\(|"
        r"^=[A-Z][A-Za-z]*[0-9]|^='[0-9]{2}-|^=[A-Z][A-Za-z]+$")
    sospechosas = [(h, c, v) for h, c, v in _celdas_formula(wb) if not permitido.match(v)]
    assert not sospechosas, (
        "Fórmulas que el exportador no escribe (texto del cliente evaluado como fórmula): "
        + "; ".join(f"{h}!{c} -> {v!r}" for h, c, v in sospechosas))
