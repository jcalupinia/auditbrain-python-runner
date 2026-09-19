"""El papel de trabajo se recalcula desde sus propias celdas y da lo archivado.

Éstas son las pruebas de extremo a extremo que faltaban: parten de un resultado
REAL de ``service.analizar`` sobre tres cortes, construyen el libro con
``exporter.construir_excel`` y vuelven a hacer la cuenta LEYENDO LAS CELDAS con
openpyxl, igual que haría el revisor que recibe el archivo.

Sin ellas, los acotamientos que protegen la medición (piso cero ante notas de
crédito, techo del importe en libros bruto ante un factor prospectivo
desbocado) vivían solo en el motor: la pantalla y la base decían una cosa y las
fórmulas del libro, otra. Un papel de trabajo que no se puede recalcular desde
sus propias celdas no sirve como evidencia de auditoría.
"""
import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.aud.pce_cxc.exporter import construir_excel
from backend.app.aud.pce_cxc.motor import ParametrosECL, resumen_deterioro
from backend.app.aud.pce_cxc.service import SEGMENTOS, analizar
from tests.excel_calc import Libro, columna, fila

CENTAVO = 0.005


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Documento", "Tipo", "Emisión", "Vencimiento", "Saldo"])
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _cortes(filas_2023, filas_2024, filas_2025):
    contenidos = [_xlsx(filas_2023), _xlsx(filas_2024), _xlsx(filas_2025)]
    fechas = [date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)]
    return [{"nombre": f"cartera_{f.year}.xlsx", "contenido": c, "fecha": f}
            for c, f in zip(contenidos, fechas)]


def _col_perdida(ws):
    """Columna de la pérdida esperada de 05-Matriz, resuelta por su rótulo.

    La hoja gana columnas (la LGD y el factor de descuento que el motor sí
    aplica), y una prueba que clava la letra deja de comprobar lo que su
    docstring dice y pasa a comprobar la columna de al lado.
    """
    return columna(ws, "Pérdida esperada")


def _libro(resultado, parametros=None):
    binario = construir_excel(resultado, parametros or {})
    return Libro(load_workbook(io.BytesIO(binario)))


def _filas_medidas(ws):
    """Filas de 05-Matriz con pérdida calculada (las SIN MEDIR no se recalculan)."""
    filas = []
    for i in range(2, ws.max_row + 1):
        banda = ws.cell(i, 2).value
        if banda in (None, "TOTAL"):
            break
        filas.append(i)
    return filas


# ---------------------------------------------------------------------------
# Cohorte compartida: ALFA/F-1 pasa de 100.000 (t-2, banda "0 a 30 días") a
# 10.000 en el corte actual, así que esa banda observa una tasa del 10 %.
# ---------------------------------------------------------------------------

COHORTE_2023 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 100000.0)]
COHORTE_2024 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 20000.0)]
COHORTE_2025 = [("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 10000.0)]


def _corrida_con_nota_de_credito():
    """C3: una nota de crédito deja la banda con exposición acreedora.

    El motor le aplica el piso cero (la corrección de valor no puede ser
    negativa); si la fórmula del libro no lo aplica, el papel imprime una
    pérdida esperada negativa que neutraliza pérdidas reales.
    """
    actual = COHORTE_2025 + [
        ("GAMMA", "F-9", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 200000.0),
        ("GAMMA", "NC-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), -250000.0),
    ]
    return analizar(_cortes(COHORTE_2023, COHORTE_2024, actual),
                    {"umbral_dias_incumplimiento": 730})


def _corrida_con_factor_desbocado():
    """I1: un factor prospectivo de 12,0 lleva la tasa observada sobre el 100 %.

    El motor acota la tasa al 100 % y la pérdida al importe en libros bruto
    (B5.5.35); sin esos topes el libro provisiona más cartera de la que existe.
    """
    actual = COHORTE_2025 + [
        ("GAMMA", "F-9", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 200100.0),
    ]
    return analizar(
        _cortes(COHORTE_2023, COHORTE_2024, actual),
        {"umbral_dias_incumplimiento": 730, "factor_prospectivo": 12.0,
         "justificacion_prospectivo": "Contracción del sector prevista por el BCE para 2026."})


def _corrida_con_nota_de_credito_individual():
    """C4 + I8: un cliente de evaluación individual con una nota de crédito.

    DELTA supera el umbral individual con 300.000 en una banda SIN tasa y una
    nota de crédito de 100.000 en una banda CON tasa. Su saldo sin medir no
    puede superar su propia exposición, y la nota de crédito tiene que quedar
    declarada igual que en la matriz colectiva.
    """
    actual = COHORTE_2025 + [
        ("DELTA", "F-5", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), -100000.0),
        ("DELTA", "F-6", "NO-RELACIONADOS", date(2023, 1, 1), date(2023, 6, 1), 300000.0),
        ("OMEGA", "F-7", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 50000.0),
    ]
    return analizar(_cortes(COHORTE_2023, COHORTE_2024, actual),
                    {"umbral_dias_incumplimiento": 730, "umbral_individual": 100000})


def _corrida_con_cartera_medida_acotada():
    """U1: lo que no se pudo medir supera a la cartera estratificada.

    BETA sale a evaluación individual con 500.000 en una banda SIN tasa, y
    GAMMA deja la banda «0 a 30 días» con exposición acreedora por una nota de
    crédito. La cartera total sigue siendo POSITIVA (160.000), pero
    «estratificada menos sin medir» da −350.000: el motor acota esa resta a
    [0, cartera total] y el papel tiene que hacer lo mismo y declararlo.
    """
    actual = COHORTE_2025 + [
        ("BETA", "F-2", "NO-RELACIONADOS", date(2020, 1, 1), date(2020, 6, 1), 500000.0),
        ("GAMMA", "F-9", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 100000.0),
        ("GAMMA", "NC-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), -450000.0),
    ]
    return analizar(_cortes(COHORTE_2023, COHORTE_2024, actual),
                    {"umbral_dias_incumplimiento": 730, "umbral_individual": 100000})


def _corrida_con_cartera_sin_estratificar():
    """T3: los EEFF traen cartera de relacionadas que el archivo no tiene.

    El tope tributario del 10 % se calculaba sobre dos bases distintas: el
    motor sobre la exposición estratificada y el Excel sobre el saldo contable
    completo. La pantalla decía que la provisión excede el tope y el papel
    decía que no.
    """
    intermedio = [("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 50000.0)]
    actual = [
        ("ALFA", "F-1", "NO-RELACIONADOS", date(2023, 9, 1), date(2023, 12, 1), 30000.0),
        ("GAMMA", "F-9", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), 200000.0),
    ]
    return analizar(
        _cortes(COHORTE_2023, intermedio, actual),
        {"umbral_dias_incumplimiento": 730,
         "eeff": {"no_relacionados": 100000.0, "relacionados": 300000.0}})


# ---------------------------------------------------------------------------
# T1 — 05-Matriz recalcula exactamente la pérdida esperada archivada
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("corrida", [
    _corrida_con_nota_de_credito,
    _corrida_con_factor_desbocado,
    _corrida_con_nota_de_credito_individual,
])
def test_cada_banda_de_la_matriz_recalcula_la_perdida_archivada(corrida):
    """La celda recalculada tiene que dar, al centavo, la pérdida que archivó la
    corrida: si el libro no aplica el piso cero ni el techo del importe en
    libros bruto, el revisor que rehaga la cuenta obtiene otra cifra."""
    resultado = corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    tramos = {(t.get("segmento"), t["tramo"]): t for t in resultado["matriz"]["tramos"]}
    medidas = 0
    for i in _filas_medidas(ws):
        clave = (ws.cell(i, 1).value, ws.cell(i, 2).value)
        archivada = tramos[clave]["ecl"]
        if archivada is None:
            assert ws[f"{_col_perdida(ws)}{i}"].value == "SIN MEDIR"
            continue
        medidas += 1
        assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{i}") == pytest.approx(archivada, abs=CENTAVO), \
            f"{clave} recalcula distinto de lo archivado"
    assert medidas, "la corrida de prueba tiene que traer bandas medidas"


@pytest.mark.parametrize("corrida", [
    _corrida_con_nota_de_credito,
    _corrida_con_factor_desbocado,
    _corrida_con_nota_de_credito_individual,
])
def test_el_total_de_la_matriz_recalcula_el_total_archivado(corrida):
    resultado = corrida()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    fila_total = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 2).value == "TOTAL")
    assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{fila_total}") == pytest.approx(
        resultado["matriz"]["ecl_total"], abs=CENTAVO)


@pytest.mark.parametrize("corrida", [
    _corrida_con_nota_de_credito,
    _corrida_con_factor_desbocado,
    _corrida_con_nota_de_credito_individual,
])
def test_la_hoja_tributaria_recalcula_la_perdida_esperada_de_la_corrida(corrida):
    """09-Tributario B2 se construye sobre 05-Matriz y 06-Individual: si la
    matriz arrastra el negativo, el efecto tributario también."""
    resultado = corrida()
    libro = _libro(resultado)
    assert libro.numero("09-Tributario", "B2") == pytest.approx(
        resultado["ecl_total"], abs=CENTAVO)


def test_la_perdida_esperada_del_papel_nunca_es_negativa():
    """C3: la corrección de valor no puede ser negativa (NIIF 9 5.5.15)."""
    resultado = _corrida_con_nota_de_credito()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    for i in _filas_medidas(ws):
        if ws[f"{_col_perdida(ws)}{i}"].value == "SIN MEDIR":
            continue
        assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{i}") >= -CENTAVO
    assert libro.numero("09-Tributario", "B2") >= -CENTAVO


def test_la_perdida_esperada_del_papel_nunca_supera_la_exposicion_de_su_banda():
    """I1: B5.5.35 mide sobre el importe en libros bruto."""
    resultado = _corrida_con_factor_desbocado()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    for i in _filas_medidas(ws):
        if ws[f"{_col_perdida(ws)}{i}"].value == "SIN MEDIR":
            continue
        exposicion = libro.numero("05-Matriz", f"C{i}")
        assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{i}") <= max(exposicion, 0.0) + CENTAVO


def test_el_papel_declara_cuando_un_acotamiento_actuo():
    """Un acotamiento que actúa en silencio es justo lo que C3 pedía evitar: el
    papel tiene que decir cuál actuó y cuánto separó del cálculo puro."""
    for corrida, esperado in ((_corrida_con_nota_de_credito, "PISO CERO"),
                              (_corrida_con_factor_desbocado, "TASA ACOTADA")):
        resultado = corrida()
        libro = _libro(resultado)
        ws = libro.wb["05-Matriz"]
        declarados = []
        for i in _filas_medidas(ws):
            valor = ws[f"{columna(ws, 'Acotamiento aplicado')}{i}"].value
            declarados.append(str(libro.evaluar(valor, "05-Matriz")
                                  if isinstance(valor, str) and valor.startswith("=") else valor))
        assert any(esperado in d.upper() for d in declarados), \
            f"05-Matriz no declara el acotamiento {esperado}: {declarados}"


def test_la_politica_recalcula_la_misma_perdida_que_la_matriz():
    """07-Politica suma 05-Matriz con SUMIFS: hereda sus fórmulas, y con ellas
    el piso cero."""
    resultado = _corrida_con_nota_de_credito()
    libro = _libro(resultado)
    ws = libro.wb["07-Politica"]
    for i in range(2, ws.max_row + 1):
        banda = ws.cell(i, 1).value
        if banda in (None, "TOTAL"):
            break
        fila = next(f for f in resultado["politica"]["filas"] if f["banda"] == banda)
        celda = ws.cell(i, 6).value
        if fila["ecl"] is None:
            assert celda in ("SIN MEDIR", "SIN COMPARAR")
            continue
        assert libro.numero("07-Politica", f"F{i}") == pytest.approx(fila["ecl"], abs=CENTAVO)


# ---------------------------------------------------------------------------
# T2 — La cartera medida cuadra sobre la misma base
# ---------------------------------------------------------------------------

def test_la_cartera_medida_no_es_negativa_ni_supera_la_cartera_total():
    """El saldo sin medir de un caso individual sumaba solo sus bandas positivas
    mientras su exposición era el neto, así que «Cartera medida» salía negativa
    sobre una cartera positiva."""
    resultado = _corrida_con_nota_de_credito_individual()
    exposicion = resultado["exposicion"]
    assert exposicion["total"] > 0
    assert exposicion["medida"] >= 0, "la cartera medida no puede ser negativa"
    assert exposicion["medida"] <= exposicion["total"] + CENTAVO
    assert exposicion["sin_medir"] <= exposicion["total"] + CENTAVO
    assert resultado["porcentaje_sobre_cartera"] >= 0


def test_el_saldo_sin_medir_de_un_caso_no_supera_su_propia_exposicion():
    """06-Individual mostraba «Saldo sin medir» por encima de la «Exposición»
    del mismo cliente."""
    resultado = _corrida_con_nota_de_credito_individual()
    for caso in resultado["individual"]["casos"]:
        assert caso["saldo_sin_tasa"] <= max(caso["saldo"], 0.0) + CENTAVO, caso["identificacion"]


def test_la_conciliacion_recalcula_la_cartera_medida_de_la_pantalla():
    """La «Cartera medida» de 08-Conciliacion y el KPI de la pantalla son la
    misma cifra, y lo sin medir del papel es la MAGNITUD que archiva la corrida
    (la deudora y la acreedora no se netean entre sí)."""
    resultado = _corrida_con_nota_de_credito_individual()
    libro = _libro(resultado)
    ws = libro.wb["08-Conciliacion"]
    assert libro.numero("08-Conciliacion", f"B{fila(ws, 'Cartera medida (')}") == pytest.approx(
        resultado["exposicion"]["medida"], abs=CENTAVO)
    assert libro.numero(
        "08-Conciliacion", f"B{fila(ws, 'SIN MEDIR (magnitud')}") == pytest.approx(
        resultado["exposicion"]["sin_medir"], abs=CENTAVO)
    assert libro.numero("08-Conciliacion", f"B{fila(ws, 'SIN MEDIR NETA')}") == pytest.approx(
        resultado["exposicion"]["sin_medir_neto"], abs=CENTAVO)


# ---------------------------------------------------------------------------
# T7 — La nota de crédito de un cliente individual se declara
# ---------------------------------------------------------------------------

def test_la_nota_de_credito_de_un_cliente_individual_se_declara():
    """Un cliente con neto positivo escondía su nota de crédito: el piso cero
    actuaba en silencio sobre 10.000,00 de pérdida."""
    resultado = _corrida_con_nota_de_credito_individual()
    assert resultado["exposicion"]["negativa"] == pytest.approx(-100000.0, abs=CENTAVO)
    assert resultado["individual"]["ecl_acotada_por_piso"] == pytest.approx(10000.0, abs=CENTAVO)
    titulos = [h["titulo"] for h in resultado["hallazgos"]]
    assert "Saldos acreedores en la cartera medida" in titulos


def test_el_papel_muestra_el_saldo_acreedor_del_cliente_individual():
    """06-Individual tiene que poder sumarse: el saldo acreedor va en su propia
    columna, no dentro de una frase."""
    resultado = _corrida_con_nota_de_credito_individual()
    libro = _libro(resultado)
    ws = libro.wb["06-Individual"]
    encabezados = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    columna = next((c for c, t in enumerate(encabezados, start=1) if "acreedor" in t.lower()), None)
    assert columna, f"06-Individual no declara el saldo acreedor: {encabezados}"
    acreedores = [ws.cell(i, columna).value for i in range(2, ws.max_row + 1)
                  if ws.cell(i, 1).value not in (None, "TOTAL")]
    assert any(isinstance(v, (int, float)) and v < -CENTAVO for v in acreedores)


# ---------------------------------------------------------------------------
# T3 — El tope acumulado del 10 % se calcula sobre una sola base
# ---------------------------------------------------------------------------

def test_el_tope_del_diez_por_ciento_usa_la_misma_base_en_el_papel_y_en_el_motor():
    """El tope se aplica sobre la cartera A LA QUE SE REFIERE LA PROVISIÓN: la
    exposición estratificada, que es la única sobre la que se midió pérdida.
    `=SaldoContable*0.1` incluía la cartera que no se ubicó en ninguna banda y
    sobre la que no se midió nada, así que inflaba el tope."""
    resultado = _corrida_con_cartera_sin_estratificar()
    tributario = resultado["tributario"]
    # El escenario tiene cartera sin estratificar: las dos bases no coinciden.
    assert resultado["exposicion"]["sin_estratificar"] > 0
    assert resultado["conciliacion"]["saldo_contable"] > resultado["exposicion"]["medida"]

    libro = _libro(resultado)
    assert libro.numero("09-Tributario", "B3") == pytest.approx(
        tributario["tope_acumulado_10pct"], abs=CENTAVO)
    assert libro.numero("09-Tributario", "B4") == pytest.approx(
        tributario["exceso_sobre_tope_acumulado"], abs=CENTAVO)
    assert libro.numero("09-Tributario", "B5") == pytest.approx(
        tributario["limite_ejercicio_1pct"], abs=CENTAVO)


def test_el_papel_dice_lo_mismo_que_la_pantalla_sobre_si_se_excede_el_tope():
    """La pantalla decía «excede» y el papel decía que no: el mismo dato, dos
    conclusiones opuestas."""
    resultado = _corrida_con_cartera_sin_estratificar()
    libro = _libro(resultado)
    excede_en_el_papel = (libro.numero("09-Tributario", "B2")
                          > libro.numero("09-Tributario", "B3") + CENTAVO)
    assert excede_en_el_papel is resultado["tributario"]["excede_tope_acumulado"]


def test_la_hoja_tributaria_declara_sobre_que_cartera_se_aplica_el_tope():
    """La base elegida no puede quedar implícita en una fórmula."""
    resultado = _corrida_con_cartera_sin_estratificar()
    ws = _libro(resultado).wb["09-Tributario"]
    textos = " ".join(str(c.value or "") for fila in ws.iter_rows() for c in fila).lower()
    assert "estratificada" in textos
    assert "sin estratificar" in textos


# ---------------------------------------------------------------------------
# U1 — La cota de la «cartera medida» vive en la fórmula y se declara
# ---------------------------------------------------------------------------

def test_la_conciliacion_recalcula_la_cartera_medida_cuando_la_cota_actua():
    """El defecto crítico de la tanda anterior: la cota se puso en el motor y la
    fórmula siguió siendo la resta cruda, así que `08-Conciliacion` B9 imprimía
    −350.000,00 bajo el rótulo «Cartera medida» mientras la corrida archivaba
    0,00 y la pantalla pintaba 0,00."""
    resultado = _corrida_con_cartera_medida_acotada()
    # El escenario es el que hace actuar la cota: sin esto la prueba no prueba.
    assert resultado["exposicion"]["total"] > 0
    assert resultado["exposicion"]["sin_medir"] > resultado["exposicion"]["total"]
    assert resultado["exposicion"]["medida"] == pytest.approx(0.0, abs=CENTAVO)

    libro = _libro(resultado)
    assert libro.numero("08-Conciliacion", "B9") == pytest.approx(
        resultado["exposicion"]["medida"], abs=CENTAVO)


def test_la_conciliacion_declara_que_la_cartera_medida_se_acoto():
    """Ninguna cota puede actuar en silencio: el papel imprime la resta cruda,
    la acotada y cuál cota actuó."""
    resultado = _corrida_con_cartera_medida_acotada()
    libro = _libro(resultado)
    ws = libro.wb["08-Conciliacion"]
    textos = {str(ws.cell(f, 1).value or "").lower(): f for f in range(1, ws.max_row + 1)}
    fila_cruda = next((f for t, f in textos.items() if "sin acotar" in t), None)
    assert fila_cruda, f"08-Conciliacion no imprime la cartera medida sin acotar: {list(textos)}"
    assert libro.numero("08-Conciliacion", f"B{fila_cruda}") == pytest.approx(
        resultado["exposicion"]["medida_sin_acotar"], abs=CENTAVO)

    fila_decl = next((f for t, f in textos.items() if "acotamiento" in t), None)
    assert fila_decl, f"08-Conciliacion no declara el acotamiento: {list(textos)}"
    declarado = libro.valor("08-Conciliacion", f"B{fila_decl}")
    assert "PISO CERO" in str(declarado).upper(), declarado


def test_la_corrida_declara_el_acotamiento_de_la_cartera_medida():
    """La pantalla y la base también lo reciben: `exposicion.medida_acotada`."""
    resultado = _corrida_con_cartera_medida_acotada()
    assert resultado["exposicion"]["medida_acotada"] == "piso_cero"
    assert resultado["exposicion"]["medida_sin_acotar"] == pytest.approx(-350000.0, abs=CENTAVO)
    titulos = [h["titulo"] for h in resultado["hallazgos"]]
    assert "Cartera medida acotada" in titulos, titulos


# ---------------------------------------------------------------------------
# U3 — El recorte del «saldo sin medir» de un caso individual se declara
# ---------------------------------------------------------------------------

def test_el_recorte_del_saldo_sin_medir_llega_al_resultado():
    """El servicio acotaba ANTES de entregar el caso al motor, así que los tres
    campos del motor se calculaban sobre un valor ya recortado y siempre decían
    que ninguna cota había actuado: la declaración era inalcanzable por
    construcción."""
    resultado = _corrida_con_nota_de_credito_individual()
    caso = next(c for c in resultado["individual"]["casos"]
                if c["identificacion"].startswith("DELTA"))
    # DELTA: 300.000 en una banda SIN tasa y una nota de crédito de 100.000 en
    # otra. Lo realmente sin medir eran 300.000 sobre una exposición de 200.000.
    assert caso["saldo"] == pytest.approx(200000.0, abs=CENTAVO)
    assert caso["saldo_sin_tasa_sin_acotar"] == pytest.approx(300000.0, abs=CENTAVO)
    assert caso["saldo_sin_tasa"] == pytest.approx(200000.0, abs=CENTAVO)
    assert caso["saldo_sin_tasa_acotado"] is True
    assert resultado["individual"]["saldo_sin_tasa_acotado_total"] == pytest.approx(
        100000.0, abs=CENTAVO)


def _fila_de_delta(ws):
    return next(i for i in range(2, ws.max_row + 1)
                if str(ws.cell(i, 1).value or "").startswith("DELTA"))


def test_el_papel_recalcula_el_saldo_sin_medir_acotado_del_caso():
    """06-Individual no leía ninguno de los tres campos: el saldo sin medir era
    un número pegado. Ahora la cota está en la fórmula y la celda recalcula, al
    centavo, lo que archivó la corrida."""
    resultado = _corrida_con_nota_de_credito_individual()
    libro = _libro(resultado)
    ws = libro.wb["06-Individual"]
    fila = _fila_de_delta(ws)
    caso = next(c for c in resultado["individual"]["casos"]
                if c["identificacion"].startswith("DELTA"))
    assert libro.numero("06-Individual", f"F{fila}") == pytest.approx(
        caso["saldo_sin_tasa"], abs=CENTAVO)
    assert libro.numero("06-Individual", f"G{fila}") == pytest.approx(
        caso["saldo_sin_tasa_sin_acotar"], abs=CENTAVO)


def test_el_saldo_sin_medir_del_caso_sigue_a_las_celdas_que_lo_sostienen():
    """Y RECALCULA de verdad, que es lo que distingue un papel de trabajo de un
    listado.

    Comparar solo la cifra no separa una fórmula de un número pegado: mientras
    nadie toque el libro, los dos dan lo mismo. La diferencia aparece cuando el
    revisor corrige una celda, que es exactamente para lo que existe el papel.
    Aquí se corrigen las dos celdas que sostienen la cota y se comprueba que la
    columna las sigue.
    """
    resultado = _corrida_con_nota_de_credito_individual()

    # 1) Se baja la exposición del cliente por debajo de lo sin medir: el techo
    #    de la cota tiene que morder más.
    libro = _libro(resultado)
    fila = _fila_de_delta(libro.wb["06-Individual"])
    libro.wb["06-Individual"][f"C{fila}"] = 120000.0
    libro = Libro(libro.wb)
    assert libro.numero("06-Individual", f"F{fila}") == pytest.approx(120000.0, abs=CENTAVO)

    # 2) Se corrige lo sin medir a un importe acreedor: el piso cero manda.
    libro = _libro(resultado)
    libro.wb["06-Individual"][f"G{fila}"] = -5000.0
    libro = Libro(libro.wb)
    assert libro.numero("06-Individual", f"F{fila}") == pytest.approx(0.0, abs=CENTAVO)

    # 3) Y sin que ninguna cota muerda, la columna es el importe de al lado.
    libro = _libro(resultado)
    libro.wb["06-Individual"][f"G{fila}"] = 30000.0
    libro = Libro(libro.wb)
    assert libro.numero("06-Individual", f"F{fila}") == pytest.approx(30000.0, abs=CENTAVO)


def test_el_papel_declara_que_el_saldo_sin_medir_se_acoto():
    """Ninguna cota en silencio: la columna dice que el saldo sin medir se
    recortó a la exposición del propio caso."""
    resultado = _corrida_con_nota_de_credito_individual()
    libro = _libro(resultado)
    ws = libro.wb["06-Individual"]
    encabezados = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    columna = next((c for c, t in enumerate(encabezados, start=1)
                    if "acotamiento" in t.lower() and "sin medir" in t.lower()), None)
    assert columna, f"06-Individual no declara el acotamiento del saldo sin medir: {encabezados}"
    fila = next(i for i in range(2, ws.max_row + 1)
                if str(ws.cell(i, 1).value or "").startswith("DELTA"))
    declarado = str(libro.valor("06-Individual", f"{chr(64 + columna)}{fila}")).upper()
    assert "ACOTADO" in declarado and "SIN ACOTAR" not in declarado, declarado


def test_la_corrida_declara_el_recorte_del_saldo_sin_medir_como_hallazgo():
    """Y llega a la pantalla y al papel de hallazgos, no solo a un campo."""
    resultado = _corrida_con_nota_de_credito_individual()
    assert resultado["exposicion"]["sin_medir_recortado"] == pytest.approx(100000.0, abs=CENTAVO)
    titulos = [h["titulo"] for h in resultado["hallazgos"]]
    assert "Saldo sin medir acotado a la exposición del caso" in titulos, titulos


# ---------------------------------------------------------------------------
# U8 — La fórmula de 05-Matriz no puede quedarse atrás del motor
# ---------------------------------------------------------------------------

def _corrida_con_lgd_y_descuento():
    """U8: la misma corrida real, medida con una LGD del 40 % y descuento.

    El servicio fija hoy `lgd = 1.0` y no descuenta, así que la fórmula podía
    omitir los dos factores sin que ninguna prueba lo notara. El motor SÍ los
    soporta: aquí se vuelve a medir la exposición REAL de la corrida con una
    LGD del 40 % y una tasa de descuento del 8 % a un año, que es exactamente
    lo que ocurriría el día que el módulo deje de fijar la LGD en 1.

    La corrida base tiene una banda medida con exposición POSITIVA (86.956,52
    al 30 %): sobre una banda acreedora el piso cero llevaría la pérdida a 0,00
    con LGD y sin ella, y la prueba no probaría nada.
    """
    resultado = _corrida_con_cartera_sin_estratificar()
    exposiciones = {s: {} for s in SEGMENTOS}
    tasas = {s: {} for s in SEGMENTOS}
    for t in resultado["matriz"]["tramos"]:
        exposiciones[t["segmento"]][t["tramo"]] = t["exposicion"]
        if t["tasa_perdida"] is not None:
            tasas[t["segmento"]][t["tramo"]] = t["tasa_perdida"]
    parametros = {
        s: ParametrosECL(tasas_perdida=tasas[s], lgd=0.4, tasa_descuento=0.08,
                         horizontes={t: 1.0 for t in exposiciones[s]},
                         fuente_tasas="Permanencia a 24 meses")
        for s in SEGMENTOS
    }
    nuevo = resumen_deterioro(exposiciones, parametros)
    return {**resultado, "matriz": nuevo["colectivo"],
            "individual": {"casos": [], "saldo_total": 0.0, "ecl_total": 0.0,
                           "saldo_sin_tasa_total": 0.0, "saldo_acreedor_total": 0.0,
                           "ecl_acotada_por_piso": 0.0, "ecl_acotada_por_techo": 0.0,
                           "saldo_sin_tasa_acotado_total": 0.0},
            "ecl_total": nuevo["ecl_total"]}


def test_la_matriz_recalcula_la_perdida_con_lgd_y_factor_de_descuento():
    """La fórmula omitía la LGD y el factor de descuento que el motor sí aplica:
    con una LGD del 40 % el papel imprimía 2,5 veces la pérdida archivada."""
    resultado = _corrida_con_lgd_y_descuento()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    tramos = {(t.get("segmento"), t["tramo"]): t for t in resultado["matriz"]["tramos"]}
    medidas = 0
    for i in _filas_medidas(ws):
        clave = (ws.cell(i, 1).value, ws.cell(i, 2).value)
        archivada = tramos[clave]["ecl"]
        if archivada is None:
            continue
        medidas += 1
        assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{i}") == pytest.approx(archivada, abs=CENTAVO), \
            f"{clave}: el papel no aplica la LGD ni el factor de descuento del motor"
    assert medidas, "la corrida de prueba tiene que traer bandas medidas"
    fila_total = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 2).value == "TOTAL")
    assert libro.numero("05-Matriz", f"{_col_perdida(ws)}{fila_total}") == pytest.approx(
        resultado["matriz"]["ecl_total"], abs=CENTAVO)


def test_la_matriz_imprime_la_lgd_y_el_factor_de_descuento_de_cada_banda():
    """No basta con que la fórmula los aplique: tienen que estar en una celda
    que el revisor pueda cambiar, como el factor prospectivo."""
    resultado = _corrida_con_lgd_y_descuento()
    libro = _libro(resultado)
    ws = libro.wb["05-Matriz"]
    encabezados = [str(ws.cell(1, c).value or "").lower() for c in range(1, ws.max_column + 1)]
    col_lgd = next((c for c, t in enumerate(encabezados, start=1) if "lgd" in t), None)
    col_fd = next((c for c, t in enumerate(encabezados, start=1) if "descuento" in t), None)
    assert col_lgd and col_fd, encabezados
    tramos = {(t.get("segmento"), t["tramo"]): t for t in resultado["matriz"]["tramos"]}
    for i in _filas_medidas(ws):
        t = tramos[(ws.cell(i, 1).value, ws.cell(i, 2).value)]
        if t["ecl"] is None:
            continue
        assert ws.cell(i, col_lgd).value == pytest.approx(t["lgd"])
        assert ws.cell(i, col_fd).value == pytest.approx(t["factor_descuento"])


def test_la_matriz_no_rotula_un_acotamiento_que_no_puede_alcanzarse():
    """`ACOTADO_TECHO` era inalcanzable desde el motor: con la tasa acotada al
    100 %, la LGD en [0, 1] y el factor de descuento en (0, 1], el producto
    nunca supera la exposición. Un rótulo que nunca se alcanza es ruido."""
    from backend.app.aud.pce_cxc import exporter
    ws = _libro(_corrida_con_factor_desbocado()).wb["05-Matriz"]
    formulas = [str(ws.cell(i, c).value or "") for i in range(2, ws.max_row + 1)
                for c in range(1, ws.max_column + 1)]
    assert not any(exporter.ACOTADO_TECHO in f for f in formulas), \
        "05-Matriz sigue rotulando el techo del importe en libros bruto, que no puede morder"


# ---------------------------------------------------------------------------
# U7 — El tope tributario del 10 % sobre una cartera estratificada negativa
# ---------------------------------------------------------------------------

def _corrida_con_cartera_estratificada_negativa():
    """U7: las notas de crédito dejan la cartera estratificada NETA acreedora.

    El tope del 10 % salía en -4.000,00 y «excede el tope» decía que sí con una
    pérdida esperada de 0,00. El papel y la pantalla coincidían, pero los dos
    en un absurdo.
    """
    actual = COHORTE_2025 + [
        ("GAMMA", "NC-1", "NO-RELACIONADOS", date(2025, 9, 1), date(2025, 12, 1), -50000.0),
    ]
    return analizar(_cortes(COHORTE_2023, COHORTE_2024, actual),
                    {"umbral_dias_incumplimiento": 730})


def test_el_tope_tributario_no_concluye_sobre_una_cartera_no_positiva():
    """Un tope negativo no es un límite: lo que no se puede contrastar se
    declara, igual que el límite anual del 1 %."""
    resultado = _corrida_con_cartera_estratificada_negativa()
    tributario = resultado["tributario"]
    assert resultado["exposicion"]["colectiva"] < 0, "el escenario tiene que ser neto acreedor"
    assert tributario["tope_acumulado_10pct"] < 0
    assert tributario["tope_acumulado_verificable"] is False
    assert tributario["excede_tope_acumulado"] is False
    assert tributario["exceso_sobre_tope_acumulado"] == 0.0


def test_el_papel_tributario_declara_que_el_tope_no_es_contrastable():
    """Y la celda del exceso no puede imprimir un número: `MAX(0;B2-B3)` sobre
    un tope negativo daba un «exceso» de 4.000,00 con provisión 0,00."""
    resultado = _corrida_con_cartera_estratificada_negativa()
    libro = _libro(resultado)
    exceso = libro.valor("09-Tributario", "B4")
    assert isinstance(exceso, str) and "NO" in exceso.upper(), exceso
    ws = libro.wb["09-Tributario"]
    textos = " ".join(str(c.value or "") for fila in ws.iter_rows() for c in fila).lower()
    assert "acreedor" in textos or "no positiva" in textos, textos[:400]


def test_el_tope_tributario_sigue_concluyendo_con_cartera_positiva():
    """La guarda no puede apagar el contraste cuando sí procede."""
    resultado = _corrida_con_factor_desbocado()
    tributario = resultado["tributario"]
    assert tributario["tope_acumulado_verificable"] is True
    libro = _libro(resultado)
    excede_en_el_papel = (libro.numero("09-Tributario", "B2")
                          > libro.numero("09-Tributario", "B3") + CENTAVO)
    assert excede_en_el_papel is tributario["excede_tope_acumulado"]
    assert libro.numero("09-Tributario", "B4") == pytest.approx(
        tributario["exceso_sobre_tope_acumulado"], abs=CENTAVO)
