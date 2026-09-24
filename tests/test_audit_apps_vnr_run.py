"""AUD-INV-VNR corre de punta a punta por el executor vía adaptadores fn(ctx).

Verifica empíricamente que /audit-apps/run ejecuta los 4 pasos contra el motor
VNR REAL (parsers/engine/exports), no que los salta. Reusa el request válido de
tests/test_vnr.py::payload (mismo que valida el motor de producción)."""
import io
import json

from openpyxl import load_workbook

from backend.app.audit_apps.examples import AUD_INV_VNR
from backend.app.audit_apps.executor import AuditAppExecutor
from backend.app.audit_apps.manifest import AuditAppManifest
from tests.test_vnr import payload


def _run(params=None):
    manifest = AuditAppManifest.from_dict(AUD_INV_VNR).validate()
    inputs = {"inventario": json.dumps(payload()).encode("utf-8")}
    p = {"framework": "full", "selling_method": "unit"}
    if params:
        p.update(params)
    return manifest, AuditAppExecutor().run(manifest, inputs=inputs, parameters=p, executed_by="a@x.ec")


def test_manifest_apunta_a_los_adaptadores():
    m = AuditAppManifest.from_dict(AUD_INV_VNR).validate()
    refs = {s.id: s.engine_ref for s in m.steps}
    assert refs["calcular_vnr"] == "backend.app.aud.inventarios_vnr.app_adapters.calcular_vnr"
    assert all(r.startswith("backend.app.aud.inventarios_vnr.app_adapters.") for r in refs.values())


def test_corre_los_4_pasos_sin_no_disponibles():
    _, res = _run()
    assert res.no_disponibles == ()            # nada se saltó: todo corrió
    assert [s.step_id for s in res.steps] == [
        "extraer_inventario", "calcular_vnr", "papel_trabajo_excel", "papel_trabajo_html"]
    assert res.input_hashes and res.output_hashes


def test_calcula_con_el_motor_real():
    _, res = _run()
    # resultado_vnr es la salida real de engine.calculate (mismos valores que test_vnr)
    resultado = next(s.value for s in res.steps if s.produces == "resultado_vnr")
    assert resultado["totals"]["impairment"] == "50.00"
    assert resultado["totals"]["adjustment"] == "45.00"
    assert resultado["rows"][0]["nrv_unit"] == "15.00"


def test_produce_excel_y_html_reales():
    _, res = _run()
    excel = res.outputs["vnr_excel"]
    assert excel[:2] == b"PK"                    # xlsx = zip válido
    wb = load_workbook(io.BytesIO(excel))
    assert wb.sheetnames                          # abre sin reparaciones
    html = res.outputs["vnr_html"].decode("utf-8")
    assert "<!doctype html" in html.lower() and "Valor neto de realización" in html
    # la salida "excepciones_vnr" serializa el resultado calculado
    exc = json.loads(res.outputs["excepciones_vnr"].decode("utf-8"))
    assert exc["totals"]["impairment"] == "50.00"


def test_parametro_tax_rate_activa_diferidos():
    # con tax_rate el adaptador enciende tax y el motor calcula impuesto diferido
    req = payload()
    req["rows"][0]["tax_base"] = "200"
    req["tax"] = {"enabled": True, "rate": "0.25", "reference": "Norma fiscal verificada",
                  "recognize_dta": True, "recoverability": "Proyección revisada",
                  "recorded_dta": "0", "recorded_dtl": "0"}
    manifest = AuditAppManifest.from_dict(AUD_INV_VNR).validate()
    inputs = {"inventario": json.dumps(req).encode("utf-8")}
    res = AuditAppExecutor().run(manifest, inputs=inputs,
                                 parameters={"framework": "full", "selling_method": "unit", "tax_rate": "0.25"},
                                 executed_by="a@x.ec")
    resultado = next(s.value for s in res.steps if s.produces == "resultado_vnr")
    assert resultado["totals"]["dta"] == "12.50"
