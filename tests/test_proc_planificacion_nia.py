"""Planificación de la auditoría con los documentos de entrada del cronograma (balances de comprobación, carta de control
interno, informe y notas del año anterior): cifras de control resueltas a mano sobre el ejemplo ficticio."""
import pytest

from backend.app.aud.niif.procesadores import planificacion_nia as m


def _run(**param):
    e = m.EJEMPLO
    return m.ejecutar(e["datasets"], {**e["parametros"], **param}, e["corte"])


def _esc(nombre):
    _, ds, par, corte = next(x for x in m.ESCENARIOS if x[0] == nombre)
    return m.ejecutar(ds, par, corte)


def _cuenta(r, codigo):
    return next(x for x in r["detalle"]["cuentas"] if x["codigo"] == codigo)


def test_mapa_jerarquia_y_signo_automatico():
    r = _run()
    d = r["detalle"]
    # Los acreedores vienen negativos en el balance de comprobación: el signo de presentación los vuelve positivos.
    assert d["signo"] == {"Activo": 1, "Pasivo": -1, "Patrimonio": -1, "Ingresos": -1, "Costos": 1, "Gastos": 1, "Otros": 1}
    c = _cuenta(r, "110302")
    assert (c["nivel"], c["detalle"], c["clas"], c["sec"]) == (4, "Sí", "Activo corriente", "Activo")
    assert _cuenta(r, "1103")["detalle"] == "No" and _cuenta(r, "1")["nivel"] == 1
    # El prefijo más largo gana: 12 = activo no corriente; 101/102 (plan de la Superintendencia) también se reconocen.
    mapa = m._mapa(m.MAPA_DEFECTO)
    assert m._clasificar("1201", mapa) == "Activo no corriente" and m._clasificar("10201", mapa) == "Activo no corriente"
    assert m._clasificar("9101", mapa) == "Otros"
    # Códigos con puntos: 1.1.10 no es subcuenta de 1.1.1.
    assert not m._debajo("1.1.10", "1.1.1") and m._debajo("1.1.1.05", "1.1.1") and m._debajo("110101", "1101")


def test_cuadre_estados_resumidos_e_indices():
    r = _run()
    e = r["detalle"]["est9"]["act"]
    # Activo 3.204.200 = pasivo 1.679.700 + patrimonio 1.156.850 + resultado 367.650.
    assert (e["TOTAL ACTIVO"], e["TOTAL PASIVO"], e["PATRIMONIO TOTAL"]) == pytest.approx((3204200, 1679700, 1524500))
    assert e["Diferencia de cuadre"] == pytest.approx(0)
    # Cartera comercial con su provisión (rubro contado una sola vez en la cuenta superior): 812.300 − 48.700 = 763.600.
    # «Otras cuentas por cobrar» (anticipos a empleados) no es cartera comercial; proveedores sin impuestos ni beneficios.
    assert e["Cuentas por cobrar"] == pytest.approx(763600) and e["Cuentas por pagar"] == pytest.approx(604300)
    # Ventas 4.860.500 (código 41) y otros ingresos 18.400; utilidad antes de participación e impuestos 490.200.
    assert (e["Ventas netas"], e["(+) Otros ingresos"], e[m.UAI]) == pytest.approx((4860500, 18400, 490200))
    i = r["detalle"]["ind"]["act"]
    # Razón corriente = activo corriente ÷ pasivo corriente; días de cartera = 763.600 × 365 ÷ 4.860.500 = 57,34.
    ac = 3204200 - (900000 + 480000 - 412600 + 96000 + 38500)
    pc = 1679700 - (420000 + 146200 + 64800)
    assert i["razonCorriente"] == round(ac / pc, 2)
    assert i["diasCartera"] == pytest.approx(57.34)
    assert i["endTotal"] == pytest.approx(round(1679700 / 3204200 * 100, 2))
    assert m._semaforo("razonCorriente", i["razonCorriente"]) == "Verde · Cómodo"


def test_r1_saldo_propio_de_la_cuenta_y_jerarquia_que_no_suma():
    # El cliente trae la cuenta 1103 con 780.000 aunque sus subcuentas suman 812.300 − 48.700 = 763.600: se usa el saldo
    # propio (R1), la diferencia de 16.400 se reporta y no se fuerza nada; el activo total sale de la cuenta 1 del archivo.
    ds = dict(m.EJEMPLO["datasets"])
    ds["balance_actual"] = [{**x, "saldo_actual": "780000.00"} if x["codigo"] == "1103" else x for x in ds["balance_actual"]]
    r = m.ejecutar(ds, m.EJEMPLO["parametros"], m.EJEMPLO["corte"])
    assert _cuenta(r, "1103")["act"] == pytest.approx(780000)
    e = r["detalle"]["est9"]["act"]
    assert e["TOTAL ACTIVO"] == pytest.approx(3204200) and e["Cuentas por cobrar"] == pytest.approx(780000)
    # Se reportan 1103 (780.000 − 763.600 = 16.400) y su cuenta superior 11, que ya no suma sus subcuentas (−16.400).
    jer = {x["message"].split(":")[0]: x["amount"] for x in r["exceptions"] if x["code"] == "JERARQUIA_NO_SUMA"}
    assert jer == {"Balance al corte · 11 ACTIVO CORRIENTE": "-16400.00", "Balance al corte · 1103 CUENTAS POR COBRAR CLIENTES": "16400.00"}


def test_r4_r5_dias_sobre_365_y_dupont_que_reconcilia():
    for nombre in ("base", "preliminar_eri", "perdida_pymes"):
        i = _esc(nombre)["detalle"]["ind"]["act"]
        assert i["dias"] == 365
        assert (i["dupontRoi"], i["dupont"]) == (i["roi"], i["roe"])
    # Corte de 8 meses: días de cartera sobre 365 con ventas de 8 meses (salen mayores; se declara en la lectura).
    r = _esc("preliminar_eri")
    e = r["detalle"]["est9"]["act"]
    assert r["detalle"]["ind"]["act"]["diasCartera"] == round(e["Cuentas por cobrar"] * 365 / e["Ventas netas"], 2)
    assert m.LECTURA_DIAS[1].startswith("Corte parcial")


def test_origenes_y_aplicaciones_cuadran_con_el_efectivo():
    for nombre in ("base", "preliminar_eri", "preliminar_prorrateo", "perdida_pymes"):
        pu = _esc(nombre)["detalle"]["puente"]
        assert pu["dif"] == pytest.approx(0, abs=0.005), nombre
    o = {x["codigo"]: x for x in _run()["detalle"]["origenes"]}
    # Inventarios sube 301.400: aplicación de efectivo; proveedores sube 43.100: origen.
    assert o["1104"]["efecto"] == pytest.approx(-301400) and o["2101"]["efecto"] == pytest.approx(43100)
    assert "1101" not in o          # el efectivo no se explica a sí mismo


def test_materialidad_del_ejemplo_y_periodo_de_la_base():
    t = _run()["totals"]
    # Ingresos 4.878.900 × 1 % = 48.789,00; desempeño 50 % = 24.394,50; trivial 5 % = 2.439,45.
    assert (t["materialidad"], t["desempeno"], t["trivial"]) == ("48789.00", "24394.50", "2439.45")
    # En la preliminar la base automática es el año anterior auditado (diciembre 2025), no el corte de 8 meses.
    r = _esc("preliminar_eri")
    assert r["detalle"]["periodo"] == "Año anterior" and r["totals"]["materialidad"] == "48789.00"
    _, ds, par, corte = next(x for x in m.ESCENARIOS if x[0] == "preliminar_eri")
    r = m.ejecutar(ds, {**par, "periodoBase": "Corte actual"}, corte)
    assert r["totals"]["materialidad"] == "34310.00"        # (3.420.000 + 11.000) × 1 %


def test_revision_preliminar_con_eri_y_prorrateo():
    # Con el ERI al mismo corte: ventas agosto 2025 = 2.980.000.
    assert _cuenta(_esc("preliminar_eri"), "4101")["ant"] == pytest.approx(2980000)
    # Sin él, se prorratea diciembre 2025: 4.860.500 × 8 ÷ 12.
    assert _cuenta(_esc("preliminar_prorrateo"), "4101")["ant"] == pytest.approx(4860500 * 8 / 12)
    # El balance compara el cierre anterior (diciembre 2025) con el corte en ambos casos.
    assert _cuenta(_esc("preliminar_prorrateo"), "110301")["ant"] == pytest.approx(812300)


def test_matriz_de_la_carta_de_control_interno():
    carta = {x["id"]: x for x in _run()["detalle"]["carta"]}
    # R01: 4 × 5 = 20, control 1 → residual 20 → Alto. R02: 16 × (6 − 2) ÷ 5 = 12,8 → Medio. R04: 12 × 3 ÷ 5 = 7,2 → Bajo.
    assert (carta["R01"]["inh"], carta["R01"]["res"], carta["R01"]["nivel"]) == (20, 20, "Alto")
    assert (carta["R02"]["res"], carta["R02"]["nivel"]) == (pytest.approx(12.8), "Medio")
    assert carta["R04"]["nivel"] == "Bajo" and carta["R06"]["nivel"] == "Pendiente de calificación"
    assert carta["R02"]["herramienta"] == m.HERRAMIENTAS["Inventarios"]


def test_riesgos_notas_y_problemas():
    r = _run()
    codigos = [e["code"] for e in r["exceptions"]]
    assert "RIESGO_FRAUDE_INGRESOS" in codigos and "ELUSION_CONTROLES" in codigos
    assert codigos.count("RIESGO_CCI_ALTO") == 2 and "RIESGO_CCI_PENDIENTE" in codigos
    # Salvedad del año anterior (jubilación patronal 18.500) → riesgo alto a verificar.
    assert any(e["code"] == "INFORME_ANTERIOR" and e["amount"] == "18500.00" for e in r["exceptions"])
    # Nota 13: 83.600 + 137.800 = 221.400 en el balance contra 222.400 auditado → −1.000 (NIA 510).
    n13 = next(n for n in r["detalle"]["notas"] if n["nota"] == "13")
    assert n13["dif"] == pytest.approx(-1000)
    assert [e["amount"] for e in r["exceptions"] if e["code"] == "NOTA_NO_CONCILIA"] == ["-1000.00"]
    assert "ESF_NO_CUADRA" not in codigos and "BASE_NO_VALIDA" not in codigos


def test_refutar_la_presuncion_exige_motivo():
    codigos = [e["code"] for e in _run(refutarIngresos="Sí")["exceptions"]]
    assert "REFUTACION_SIN_MOTIVO" in codigos and "RIESGO_FRAUDE_INGRESOS" not in codigos
    codigos = [e["code"] for e in _run(refutarIngresos="Sí", motivoRefutacion="Ventas de contado con precio regulado")["exceptions"]]
    assert "REFUTACION_SIN_MOTIVO" not in codigos


def test_patrimonio_en_deficit_se_presenta_negativo():
    # Revisor A1: resultados acumulados pasan de −360.450 (acreedor) a +1.800.000 (deudor, pérdidas acumuladas): toda la rama
    # de patrimonio sube 2.160.450 y la de bancos baja lo mismo para que el balance siga cuadrando. El patrimonio queda en
    # déficit y debe verse negativo, con su indicio NIA 570, sin descuadres ficticios.
    D = 2160450
    delta = {"3301": D, "33": D, "3": D, "110102": -D, "1101": -D, "11": -D, "1": -D}
    ds = dict(m.EJEMPLO["datasets"])
    ds["balance_actual"] = [{**x, "saldo_actual": f"{float(x['saldo_actual']) + delta[x['codigo']]:.2f}"} if x["codigo"] in delta else x
                            for x in ds["balance_actual"]]
    r = m.ejecutar(ds, m.EJEMPLO["parametros"], m.EJEMPLO["corte"])
    codigos = [e["code"] for e in r["exceptions"]]
    assert r["detalle"]["signo"]["Patrimonio"] == -1
    assert r["detalle"]["est9"]["act"]["PATRIMONIO TOTAL"] < 0 and "PATRIMONIO_NEGATIVO" in codigos
    assert "ESF_NO_CUADRA" not in codigos and "JERARQUIA_NO_SUMA" not in codigos


def test_escenario_de_perdida_sin_documentos_del_anio_anterior():
    r = _esc("perdida_pymes")
    codigos = [e["code"] for e in r["exceptions"]]
    # Base ≤ 0: sin materialidad y nada queda marcado como material (revisor A3).
    assert r["detalle"]["materialidad"]["global"] is None and r["totals"]["materialidad"] == "0.00"
    assert all(x["material"] == "No" for x in r["detalle"]["cuentas"])
    # Costo +700.000 y sin impuesto: utilidad antes de participación e impuestos 490.200 − 700.000 + 122.550 − 122.550 = −209.800.
    assert r["detalle"]["bases"][m.UAI] == pytest.approx(-209800)
    assert {"BASE_NO_VALIDA", "PERDIDA_EJERCICIO", "ENCARGO_INICIAL", "SIN_CARTA_CI", "SIN_INFORME_ANTERIOR",
            "SIN_NOTAS_ANTERIOR"} <= set(codigos)
    assert "ESF_NO_CUADRA" not in codigos and r["detalle"]["marco"] == m.MARCO_PYMES


def test_balance_que_no_cuadra():
    # Bancos +1.000 sin contrapartida en toda su rama (110102, 1101, 11 y 1): el balance al corte deja de cuadrar por 1.000.
    ds = dict(m.EJEMPLO["datasets"])
    rama = {"110102", "1101", "11", "1"}
    ds["balance_actual"] = [{**x, "saldo_actual": f"{float(x['saldo_actual']) + 1000:.2f}"} if x["codigo"] in rama else x
                            for x in ds["balance_actual"]]
    codigos = [e["code"] for e in m.ejecutar(ds, {}, "2025-12-31")["exceptions"]]
    r = m.ejecutar(ds, {}, "2025-12-31")
    assert [e["amount"] for e in r["exceptions"] if e["code"] == "ESF_NO_CUADRA"] == ["1000.00"]
    assert "ESF_ANTERIOR_NO_CUADRA" not in codigos and "JERARQUIA_NO_SUMA" not in codigos


@pytest.mark.parametrize("param, mensaje", [
    ({"baseMaterialidad": "Utilidad bruta"}, "Base de la materialidad"),
    ({"pctIngresos": 0}, "Porcentaje sobre ingresos"),
    ({"pctDesempeno": 120}, "Materialidad de desempeño"),
    ({"tipoRevision": "Intermedia"}, "Tipo de revisión"),
    ({"mesesTranscurridos": 13}, "Meses transcurridos"),
    ({"mapaCuentas": "1 Activo"}, "Mapa de cuentas"),
    ({"mapaCuentas": "1=Bancos"}, "Mapa de cuentas"),
    ({"umbralMedio": 20, "umbralAlto": 15}, "Matriz de riesgos"),
    ({"encargoInicial": "Tal vez"}, "Encargo inicial"),
    ({"fechaFinal": "31/02/2026"}, "Fecha de la visita final"),
])
def test_parametros_invalidos(param, mensaje):
    with pytest.raises(ValueError, match=mensaje):
        _run(**param)


def test_anexos_obligatorios_y_validaciones():
    with pytest.raises(ValueError, match="fecha de corte"):
        m.ejecutar({"balance_anterior": m.EJEMPLO["datasets"]["balance_anterior"]}, {}, "2025-12-31")
    with pytest.raises(ValueError, match="cierre del año anterior"):
        m.ejecutar({"balance_actual": m.EJEMPLO["datasets"]["balance_actual"]}, {}, "2025-12-31")
    dup = [{"codigo": "1101", "cuenta": "A", "saldo_actual": "1", "_row": 2}, {"codigo": "1101.0", "cuenta": "B", "saldo_actual": "2", "_row": 3}]
    assert not m.validar_filas("balance_actual", dup)["ok"]
    assert not m.validar_filas("carta_control_interno", [{"id": "R1", "proceso": "x", "hallazgo": "y", "probabilidad": "7", "_row": 2}])["ok"]
    assert not m.validar_filas("informe_anterior", [{"concepto": "x", "tipo": "Comentario", "detalle": "y", "_row": 2}])["ok"]
    assert m._cod("1101.0") == "1101" and m._cod(" 1.1.01 ") == "1.1.01"


def test_hojas_con_las_cedulas_y_el_ancho_de_columnas():
    r = _run()
    hs = m.hojas(r)
    assert [h["name"] for h in hs] == [n for n, _ in m.CEDULAS]
    for h in hs:
        for fila in h["rows"]:
            assert len(fila) == len(h["cols"]), h["name"]
    horizontal = next(h for h in hs if h["name"] == "08_Horizontal")
    assert len(horizontal["rows"]) == len(r["detalle"]["cuentas"])
    # Programa: un PT por hallazgo de la carta, uno por posible riesgo presente, uno por cuenta a revisar sin riesgo que
    # la cubra y los procedimientos de todo encargo.
    programa = next(h for h in hs if h["name"] == "19_Programa")
    presentes = sum(1 for x in r["detalle"]["riesgos"] if x["presenta"] == "Sí")
    propias = sum(1 for f in programa["rows"] if str(f[2]).startswith("Cuenta "))
    assert len(programa["rows"]) == len(r["detalle"]["carta"]) + presentes + propias + len(m.PROC_ENCARGO)


def test_panel_con_textos_propios_y_las_demas_herramientas_sin_cambio():
    from backend.app.aud.niif.procesadores import PROCESADORES, graficos
    from backend.app.aud.niif.procesadores import html_ejecutivo as hx

    r = _run()
    r["hojas"] = m.hojas(r)
    p = graficos.panel(m, r, r["hojas"])
    assert not p["faltan"] and p["comparativo"]["rotulo"] == "Umbrales de la planificación"
    assert p["problemas"]["rotulo"] == "Asuntos para la planificación"
    tarjetas = hx.kpis_datos(p)
    assert tarjetas[2]["nota"] == "NIA 320 párr. 11" and "var" not in tarjetas[2]
    otro = PROCESADORES["cxc_cartera"]
    assert graficos.textos(otro.PANEL)["comparativo"] == "Registrado vs recalculado"
    assert graficos.textos(otro.PANEL)["nota_registrado"] == "según el cliente"


def test_tableros_del_artefacto_en_html_excel_word_y_ppt():
    """Los gráficos del artefacto de análisis que el prompt no produjo: índices por grupo (liquidez, actividad,
    endeudamiento, rentabilidad) y analítico (estructura del balance, estado de resultados), anterior frente a actual,
    en los cuatro formatos del papel."""
    import io

    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    from backend.app.aud.niif.procesadores import datos_cliente, graficos, libro

    e = m.EJEMPLO
    r = m.ejecutar(e["datasets"], e["parametros"], e["corte"])
    r["hojas"] = datos_cliente.con_datos(m, r, e["datasets"])
    p = graficos.panel(m, r, r["hojas"])
    assert p["faltan"] == []
    t = {x["rotulo"]: x for x in p["tableros"]}
    assert list(t) == ["Liquidez", "Actividad", "Endeudamiento", "Rentabilidad", "Estructura del balance", "Estado de resultados"]
    assert t["Liquidez"]["categorias"] == ["Razón corriente", "Prueba ácida"]
    assert t["Liquidez"]["series"] == [("Anterior", [1.81, 1.1]), ("Actual", [2.0, 1.09])]
    assert t["Actividad"]["series"][1][1] == [57.34, 102.65, 64.83, 95.16]
    assert t["Estructura del balance"]["categorias"][2:] == ["Pasivo total", "Patrimonio total"]
    assert t["Estado de resultados"]["series"][1][1][-1] == 367650.0

    reg = {"run": r, "engagement": {"client": "Comercial Andina de Ejemplo S.A.", "cutoff": e["corte"]}, "program": [], "sources": []}
    d = m.definicion()
    html = libro.html(d, reg, [], 1, "Borrador").decode("utf-8")
    assert "Tableros del análisis" in html and html.count('class="graficos tableros"') == 1
    # Premium: degradado, variación ▲▼ coloreada por sentido favorable y semáforo de la cédula.
    from backend.app.aud.niif.procesadores import graficos_svg as gs

    liq = gs.agrupadas(t["Liquidez"]["categorias"], t["Liquidez"]["series"], "Liquidez", "veces", None,
                       t["Liquidez"]["mejor"], t["Liquidez"]["estados"])
    assert "<linearGradient" in liq and 'fill="url(#' in liq and "Semáforo: Verde · Cómodo" in liq
    assert "▲ 10,5 %" in liq and "var(--c-baja)" in liq      # razón corriente sube y es favorable → verde
    assert "▼ 0,9 %" in liq and "var(--c-alta)" in liq       # prueba ácida baja y es desfavorable → rojo
    assert gs.variacion_tablero(10.28, 10.91, "%") == ("▲ 0,63 pp", 1)
    # Ningún color se repite en la misma lámina: cada tablero tiene su propia familia de color.
    assert len({x["color"] for x in p["tableros"]}) == len(p["tableros"])
    # Paleta ejecutiva (sin celeste ni verde) y relieve 3D: techo iluminado y lateral en sombra.
    assert {x["color"] for x in p["tableros"]} <= set(gs.TABLERO_HEX)
    assert 'fill="#FFFFFF" fill-opacity="0.34"' in liq and 'fill="#000000" fill-opacity="0.38"' in liq
    assert 'data-s="s-tableros"' in html      # los tableros van en su propia lámina
    wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "Borrador")))
    assert len(wb["00_Inicio"]._charts) == 4 + 6
    datos = [c.value for row in wb[libro.HOJA_DATOS_GRAFICOS].iter_rows() for c in row]
    assert "='10_Indices'!$E$6" in datos and "='09_Estados'!$D$30" in datos   # razón corriente actual, utilidad neta actual
    # El rótulo de cada indicador lleva su variación por fórmula (FIXED: separador decimal del equipo).
    assert any(isinstance(x, str) and x.startswith('="Razón corriente"&CHAR(10)') and "FIXED(" in x for x in datos)
    tab = wb["00_Inicio"]._charts[4]
    assert all(s_.graphicalProperties.gradFill is not None for s_ in tab.series)   # barras con degradado
    colores = [str(g.srgbClr) for ch in wb["00_Inicio"]._charts[4:] for s_ in ch.series
               for g in s_.graphicalProperties.gradFill.gsLst]
    assert len(colores) == len(set(colores)), "un color de tablero se repite en la portada del Excel"
    doc = Document(io.BytesIO(libro.docx(d, reg, [], 1, "Borrador")))
    textos = [c.text for tb in doc.tables for fila in tb.rows for c in fila.cells]
    assert any(x.startswith("Rentabilidad") for x in textos) and any(p_.text == "Tableros del análisis" for p_ in doc.paragraphs)
    prs = Presentation(io.BytesIO(libro.pptx(d, reg, [], 1, "Borrador")))
    diaps = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text] for s in prs.slides]
    diaps = [x for x in diaps if x[0] == "Tableros del análisis"]
    assert [x[1] for x in diaps] == ["Liquidez y actividad", "Endeudamiento y rentabilidad",
                                     "Estructura del balance y estado de resultados"]


def test_sumarias_por_rubro_con_subcuentas_ajustes_y_cuadre():
    """Cédulas sumarias (reclamo del dueño: «no veo las sumarias»): un bloque por rubro con la cuenta del rubro,
    sus subcuentas con sangría, el total de las cuentas de detalle y el cuadre, todo por fórmula a 08_Horizontal."""
    r = _run()
    hs = {h["name"]: h for h in m.hojas(r)}
    h = hs["08S_Sumarias"]
    v = lambda c: c.get("v") if isinstance(c, dict) else c  # noqa: E731
    filas = [[v(c) for c in f] for f in h["rows"]]
    rubros = [f for f, e in zip(filas, h["estilos"]) if e.get("tipo") == "titulo"]
    # Un bloque por cada cuenta de nivel 3 (o superior sin subcuentas) del plan del ejemplo.
    assert len(rubros) == sum(1 for x in r["detalle"]["cuentas"] if m._es_rubro(x)) == 28
    # CxC: rubro, dos subcuentas con sangría, total y cuadre en cero; nota del año anterior y marca de umbral.
    i = next(k for k, f in enumerate(filas) if f[1] == "1103")
    assert [f[1] for f in filas[i:i + 3]] == ["1103", "110301", "110302"]
    assert h["estilos"][i + 1] == {"sangria": 1, "col": "Cuenta"}
    assert filas[i + 3][2] == m.TXT_TOTAL_SUMARIA and filas[i + 3][5:7] == [644600.0, 763600.0]
    assert filas[i + 4][2] == m.TXT_CUADRE_SUMARIA and filas[i + 4][5:7] == [0.0, 0.0] and filas[i + 4][12] == "Cuadra"
    assert filas[i][11] == "Nota 4" and filas[i][12] == "Supera el umbral"
    # Todo cuadra en el ejemplo; los saldos vienen de 08_Horizontal y los ajustes suben de las cuentas de detalle.
    assert all(f[12] == "Cuadra" for f, e in zip(filas, h["estilos"]) if e.get("tipo") == "control")
    fila_rubro = h["rows"][i]
    assert fila_rubro[5]["f"].startswith("'08_Horizontal'!G") and fila_rubro[7]["f"] == f"N(H{m.FILA0 + i + 1})+N(H{m.FILA0 + i + 2})"
    assert h["rows"][i + 1][7] is None                     # en la cuenta de detalle el ajuste lo escribe el auditor
    # Si la jerarquía no suma (1103 con saldo propio distinto de sus subcuentas), el cuadre lo muestra.
    e = m.EJEMPLO
    ds = {k: [dict(x) for x in v_] for k, v_ in e["datasets"].items()}
    for x in ds["balance_actual"]:
        if str(x["codigo"]) == "1103":
            x["saldo_actual"] = 780000.0
    r2 = m.ejecutar(ds, e["parametros"], e["corte"])
    h2 = next(x for x in m.hojas(r2) if x["name"] == "08S_Sumarias")
    j = next(k for k, f in enumerate(h2["rows"]) if v(f[1]) == "1103")
    assert v(h2["rows"][j + 4][12]) == "Revisar la jerarquía" and abs(v(h2["rows"][j + 4][6]) - 16400.0) < 0.01


def test_notas_desglose_por_cuenta_composicion_auditada_y_rubros_sin_nota():
    """Notas a los EEFF (reclamo del dueño): detalle por cuenta de cada nota con su conciliación (NIA 510), composición
    auditada línea por línea que cuadra con el saldo de la nota, y los rubros del balance que ninguna nota cubre."""
    r = _run()
    hs = {h["name"]: h for h in m.hojas(r)}
    v = lambda c: c.get("v") if isinstance(c, dict) else c  # noqa: E731
    det = [[v(c) for c in f] for f in hs["15D_Notas_Detalle"]["rows"]]
    # Nota 6: el rubro y sus tres subcuentas, total del balance 929.100 y coincide con la nota auditada.
    i = det.index(next(f for f in det if f[0] == "Nota 6" and f[1] == "1201"))
    assert [f[1] for f in det[i:i + 4]] == ["1201", "120101", "120102", "120103"]
    tot = next(f for f in det[i:] if f[2] == m.TXT_TOTAL_NOTA)
    assert tot[4:6] == [929100.0, 967400.0]
    assert next(f for f in det[i:] if f[2] == m.TXT_DIF_NOTA)[8] == "Coincide"
    # Nota 13: el balance anterior es 1.000 menor que la nota auditada.
    d13 = next(f for f in det if f[0] == "Nota 13" and f[2] == m.TXT_DIF_NOTA)
    assert d13[4] == -1000.0 and d13[8] == "Revisar (NIA 510)"
    # Rubros sin nota al final (p. ej. capital social e impuestos por pagar).
    sin = [f[1] for f in det if f[8] == "Sin nota"]
    assert len(sin) == 10 and "3101" in sin and "2103" in sin
    # Composición: líneas de saldo que suman el total de la nota; la nota 6 trae su movimiento del año (informativo).
    comp = [[v(c) for c in f] for f in hs["15C_Composicion"]["rows"]]
    assert any(f[1] == "Depreciación del año" and f[2] == "Movimiento" for f in comp)
    assert all(f[4] == "Coincide" for f in comp if f[1] in (m.TXT_TOTAL_COMP, m.TXT_DIF_COMP))
    ctl = {v(f[0]): [v(c) for c in f] for f in hs["16_Control"]["rows"]}
    assert ctl["Composición de las notas que no suma el saldo auditado"][2:4] == [0, "Conforme"]
    assert ctl["Rubros del balance sin nota del año anterior"][2:4] == [10, "Revisar"]
    # Si una línea de la composición no cuadra, el control lo marca.
    e = m.EJEMPLO
    ds = {k: [dict(x) for x in v_] for k, v_ in e["datasets"].items()}
    next(x for x in ds["notas_detalle"] if x["nota"] == "5" and x["tipo"] == "Saldo")["importe"] = "650000.00"
    r2 = m.ejecutar(ds, e["parametros"], e["corte"])
    hs2 = {h["name"]: h for h in m.hojas(r2)}
    ctl2 = {v(f[0]): [v(c) for c in f] for f in hs2["16_Control"]["rows"]}
    assert ctl2["Composición de las notas que no suma el saldo auditado"][2:4] == [2, "Revisar"]   # total y saldo auditado
    # Sin notas cargadas (encargo inicial) los controles quedan «No evaluado».
    rp = _esc("perdida_pymes")
    ctl3 = {v(f[0]): [v(c) for c in f] for f in next(h for h in m.hojas(rp) if h["name"] == "16_Control")["rows"]}
    assert ctl3["Rubros del balance sin nota del año anterior"][3] == "No evaluado"
    assert m.validar_filas("notas_detalle", [{"nota": "4", "concepto": "x", "importe": "1", "tipo": "Ajuste", "_row": 2}])["ok"] is False


def test_semaforo_no_significativo_con_patrimonio_negativo():
    """Con patrimonio en déficit, los índices que dividen para el patrimonio no pueden salir en verde («Conservador»,
    «Bajo»): van en rojo como no significativos y la lectura remite a empresa en marcha (hallazgo de los agentes)."""
    v = lambda c: c.get("v") if isinstance(c, dict) else c  # noqa: E731
    for esc, esperado in (("patrimonio_deficit", True), ("base", False)):
        h = next(x for x in m.hojas(_esc(esc)) if x["name"] == "10_Indices")
        filas = {v(f[0]): f for f in h["rows"]}
        for nombre in ("Endeudamiento financiero (veces)", "Endeudamiento patrimonial (veces)",
                       "Multiplicador de apalancamiento (veces)", "ROE (%)", "ROE por DuPont (%)"):
            f = filas[nombre]
            assert (v(f[6]) == m.NO_SIGNIFICATIVO) is esperado, (esc, nombre, v(f[6]))
            assert (v(f[7]) == m.LECTURA_PATRIMONIO) is esperado
            assert "PATRIMONIO" not in f[6]["f"] and "<=0" in f[6]["f"]    # la condición va por fórmula al patrimonio de la hoja 09
        assert v(filas["Razón corriente (veces)"][6]) != m.NO_SIGNIFICATIVO


def test_programa_cubre_todas_las_cuentas_a_revisar_y_procedimientos_de_todo_encargo():
    """NIA 330 párr. 18 (hallazgo de los agentes): toda cuenta marcada en la hoja 18 tiene un procedimiento en la 19, sea
    por un riesgo de su área o por su propio procedimiento sustantivo; y el programa trae los procedimientos de todo encargo."""
    v = lambda c: c.get("v") if isinstance(c, dict) else c  # noqa: E731
    for esc in ("base", "preliminar_eri", "preliminar_prorrateo", "patrimonio_deficit", "perdida_pymes"):
        r = _esc(esc)
        hs = {h["name"]: h for h in m.hojas(r)}
        prog = [[v(c) for c in f] for f in hs["19_Programa"]["rows"]]
        areas = {m._area(f[1]) for f in prog if f[3] != "Todo encargo"} | {m._area(f[1] + " ") for f in prog}
        codigos = {str(f[2]).replace("Cuenta ", "") for f in prog} | {str(f[1]).split(" ")[0] for f in prog}
        for x in r["detalle"]["revisar"]:
            if x["revisa"] != "Sí":
                continue
            c = x["x"]
            cubierta = (c["codigo"] in codigos or m._area(c["cuenta"], c["sec"]) in areas
                        or any(m._area(k["proceso"] + " " + k["hallazgo"]) == m._area(c["cuenta"], c["sec"]) for k in r["detalle"]["carta"]))
            assert cubierta, (esc, c["codigo"], c["cuenta"])
        normas = {f[2] for f in prog if f[3] == "Todo encargo"}
        assert normas == {"NIA 560", "NIA 550", "NIA 501", "NIA 570", "NIA 250", "NIA 580"}
        assert all(f[7] and f[8] and f[9] and f[10] for f in prog)          # aseveraciones, evidencia, responsable, aplica
    # PPE y proveedores (sin riesgo propio en el ejemplo) tienen ahora su procedimiento sustantivo.
    prog = [[v(c) for c in f] for f in next(h for h in m.hojas(_run()) if h["name"] == "19_Programa")["rows"]]
    assert {"Cuenta 1201", "Cuenta 2101"} <= {f[2] for f in prog}
    # Sin materialidad, las anomalías no se evalúan (antes decía «Conforme») y el fraude en ingresos marca Ventas netas.
    rp = _esc("perdida_pymes")
    hs = {h["name"]: h for h in m.hojas(rp)}
    ctl = {v(f[0]): [v(c) for c in f] for f in hs["16_Control"]["rows"]}
    assert ctl["Anomalías de severidad alta"][3] == "No evaluado"
    rev = {v(f[0]): [v(c) for c in f] for f in hs["18_Cuentas_Revisar"]["rows"]}
    assert rev["4101"][8] == "RB-01" and rev["4101"][10] == "Sí"


def test_riesgo_significativo_sobre_el_inherente_y_colores_por_nivel():
    """NIA 315 párr. 32 y NIA 330 párr. 21: el riesgo significativo se juzga sobre el riesgo INHERENTE; un control fuerte baja
    el residual pero no lo vuelve «Bajo». Los niveles se pintan por color en Excel (formato condicional), HTML, Word y PPT."""
    import io

    from openpyxl import load_workbook

    from backend.app.aud.niif.procesadores import base, datos_cliente, libro

    v = lambda c: c.get("v") if isinstance(c, dict) else c  # noqa: E731
    e = m.EJEMPLO
    ds = dict(e["datasets"])
    # R07: 5 × 5 = 25 (significativo) con control 5 → residual 25 × 1 ÷ 5 = 5, que sin la regla sería «Bajo».
    ds["carta_control_interno"] = list(ds["carta_control_interno"]) + [
        m._ci("R07", "Tesorería", "Pagos a proveedores del exterior por montos altos.", "Ocurrencia", "5", "5", "5", "")]
    r = m.ejecutar(ds, e["parametros"], e["corte"])
    carta = {x["id"]: x for x in r["detalle"]["carta"]}
    assert (carta["R07"]["inh"], carta["R07"]["res"], carta["R07"]["sig"], carta["R07"]["nivel"]) == (25, 5, "Sí", "Alto")
    assert carta["R01"]["sig"] == "Sí" and carta["R02"]["sig"] == "No" and carta["R06"]["sig"] == "No"
    hs = {h["name"]: h for h in m.hojas(r)}
    h12 = hs["12_Riesgos_CCI"]
    fila = next(f for f in h12["rows"] if f[0] == "R07")
    assert h12["cols"][12][0] == "¿Riesgo significativo?" and v(fila[12]) == "Sí" and v(fila[9]) == "Alto"
    assert fila[12]["f"] == f'IF(H{m.FILA0 + 6}="","",IF(H{m.FILA0 + 6}>={m._par("umbralSignificativo")},"Sí","No"))'
    assert f'M{m.FILA0 + 6}="Sí"' in fila[9]["f"]
    prog = next(f for f in hs["19_Programa"]["rows"] if f[2] == "R07")
    assert v(prog[3]) == "Significativo" and v(prog[6]) == m.OPORTUNIDAD_ALTO
    assert prog[3]["f"].startswith(f"IF('12_Riesgos_CCI'!M{m.FILA0 + 6}=\"Sí\",\"Significativo\"")
    # Parámetro con su sustento (VERIFICAR) en la hoja 02.
    par = {v(f[0]): f for f in hs["02_Parametros"]["rows"]}
    etq = m.ETIQUETAS_PARAM["umbralSignificativo"]
    assert v(par[etq][1]) == 20 and "NIA 315" in par[etq][2]
    # Umbral más alto que el inherente → deja de ser significativo y el nivel vuelve al residual.
    r2 = m.ejecutar(ds, {**e["parametros"], "umbralSignificativo": 25}, e["corte"])
    assert {x["id"]: x["sig"] for x in r2["detalle"]["carta"]}["R01"] == "No"
    # Colores por nivel: columnas declaradas y reglas de formato condicional en el Excel.
    assert hs["12_Riesgos_CCI"]["colores"] == ["Nivel"] and hs["10_Indices"]["colores"] == ["Semáforo"]
    assert hs["16_Control"]["colores"] == ["Estado"] and hs["19_Programa"]["colores"] == ["Nivel"]
    assert base.rol_color(h12, "Nivel", "Pendiente de calificación") == "info"
    assert base.rol_color(hs["10_Indices"], "Semáforo", m.NO_SIGNIFICATIVO) == "alta"
    assert base.rol_color(hs["19_Programa"], "Nivel", "Significativo") == "sig"
    assert base.rol_color(h12, "Proceso o área", "Alto") is None
    r["hojas"] = datos_cliente.con_datos(m, r, ds)
    reg = {"run": r, "engagement": {"client": "Comercial Andina de Ejemplo S.A.", "cutoff": e["corte"]}, "program": [], "sources": []}
    d = m.definicion()
    wb = load_workbook(io.BytesIO(libro.xlsx(d, reg, [], 1, "borrador")))
    ws = wb[next(n for n in wb.sheetnames if n.startswith("12_"))]
    reglas = [(str(rg.sqref), x.formula[0]) for rg in ws.conditional_formatting for x in rg.rules]
    assert any(s.startswith("J") and '"Alto"' in f for s, f in reglas)
    assert any(s.startswith("J") and '"Significativo"' in f for s, f in reglas)
    html = libro.html(d, reg, [], 1, "borrador").decode("utf-8")
    assert 'class="nivel n-sig"' in html and 'class="nivel n-alta"' in html and ".nivel.n-sig{" in html
